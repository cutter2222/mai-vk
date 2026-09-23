"""Structured model content reaches editable PPTX without model-selected layouts."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

import pytest
from pptx import Presentation

from presentation_designer.generation import variants as vr
from presentation_designer.layout.shapes import iter_shapes
from presentation_designer.shared.settings import get_settings
from tests.layout.test_compose import _compose, _plan_with


@pytest.mark.parametrize("count", [2, 3])
@pytest.mark.parametrize("missing_image", [False, True])
def test_content_first_pptx_preserves_blocks_and_notes(
    count: int,
    missing_image: bool,
    rich_profile: dict[str, Any],
    rich_template_path: Path,
    example_package: dict[str, Any],
    tmp_path: Path,
) -> None:
    ctx = vr.build_context({}, rich_profile, example_package, "balanced", {}, get_settings(), 1)
    ctx.theses = [
        vr.Thesis("t1", 1, "claim", "Проверяем пилот", "", True, None, [], [], [], [], None)
    ]
    packet = vr.Packet(0, ctx.theses, 1, 1, 1, {})
    items = [
        {"sub": "Согласие", "text": "Участие в пилоте добровольное"},
        {"sub": "Поддержка", "text": "Нагрузку проверяем до расширения"},
        {"sub": "Результат", "text": "Решение принимаем после проверки"},
    ][:count]
    answer = {
        "slides": [
            {
                "theses": ["t1"],
                "title": "Сначала проверяем условия пилота",
                "visual": "image" if missing_image else "cards",
                "image": "nonexistent" if missing_image else None,
                "items": items,
                "notes": "Не расширять пилот без проверки условий",
            }
        ],
    }
    draft = vr.make_validator(ctx, packet)(answer)["drafts"][0]
    assert draft.image is None
    assert draft.pattern.role not in ("screenshot", "mockup", "image_full")
    fitted = vr.fit_draft(ctx, draft)
    slide = vr.slide_from_draft(ctx, fitted, slide_id="s1", order=1)
    plan = _plan_with(rich_profile, [slide])
    result = _compose(
        plan, rich_profile, rich_template_path, example_package, tmp_path / "content-first.pptx"
    )
    assert result.integrity.ok
    presentation = Presentation(result.pptx_path)
    assert len(presentation.slides) == 1
    actual = presentation.slides[0]
    visible = re.sub(r"\s+", " ", " ".join(s.text for s in iter_shapes(actual) if s.has_text_frame))
    assert draft.title in visible
    for item in items:
        assert item["sub"] in visible
        assert item["text"] in visible
    assert draft.notes in actual.notes_slide.notes_text_frame.text
