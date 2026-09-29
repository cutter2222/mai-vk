"""Объёмная гистограмма в шаблоне не роняет сборку (finansy.pptx, bar3DChart)."""

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.ns import qn
from pptx.util import Emu

from presentation_designer.generation.office_chart import read_data
from presentation_designer.layout.charts import chart_family_of, describe_chart


def _bar3d_chart():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["A", "B", "C"]
    data.add_series("Ряд 1", (1, 2, 3))
    data.add_series("Ряд 2", (4, 5, 6))
    frame = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Emu(0), Emu(0), Emu(4_000_000), Emu(3_000_000), data
    )
    xchart = frame.chart._chartSpace.find(".//" + qn("c:barChart"))
    xchart.tag = qn("c:bar3DChart")
    for tag in ("c:overlap", "c:serLines"):
        for el in xchart.findall(qn(tag)):
            xchart.remove(el)
    return frame.chart


def test_bar3d_chart_is_read() -> None:
    chart = _bar3d_chart()
    info = describe_chart(chart)
    assert info["type"] == "three_d_column_clustered"
    assert info["series_count"] == 2
    assert info["categories_count"] == 3
    assert chart_family_of(chart) == "category"
    assert read_data(chart) == (
        ["A", "B", "C"],
        [("Ряд 1", [1.0, 2.0, 3.0]), ("Ряд 2", [4.0, 5.0, 6.0])],
    )
