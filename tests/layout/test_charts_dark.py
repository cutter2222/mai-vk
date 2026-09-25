"""Новая диаграмма: прозрачный фон, на тёмном слайде — светлые подписи."""

from __future__ import annotations

from pptx import Presentation

from presentation_designer.layout import charts


def _spec() -> charts.ChartSpec:
    return charts.chart_spec(
        {"type": "column", "dataset_id": "d1"},
        {
            "dataset_id": "d1",
            "columns": [{"name": "Район"}, {"name": "Доля"}],
            "rows": [["Северный", 68], ["Речной", 71]],
        },
    )


def test_new_chart_is_transparent_and_light_on_dark() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    style = charts.ChartStyle(text_color="#1A1A1A").on_dark()
    frame = charts.add_chart(slide, (0, 0, 4000000, 3000000), _spec(), style)
    xml = frame.chart._chartSpace.xml
    assert xml.count("<a:noFill/>") >= 4, "фон и рамка диаграммы и области построения"
    assert 'val="1A1A1A"' not in xml and 'val="FFFFFF"' in xml


def test_light_chart_keeps_text_color() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    frame = charts.add_chart(
        slide, (0, 0, 4000000, 3000000), _spec(), charts.ChartStyle(text_color="#1A1A1A")
    )
    assert 'val="1A1A1A"' in frame.chart._chartSpace.xml
