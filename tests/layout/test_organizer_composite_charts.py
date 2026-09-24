"""Диаграммы из фигур в шаблонах организаторов: VK Tech (слайды 41, 44–46: полосы-«бабочка»,
наложенные столбцы с легендой, кольца с градиентом, кольцо с клином) и VK Education (слайд 46:
кольца с числом в центре). На остальных слайдах всех четырёх шаблонов ничего не находится."""

from __future__ import annotations

import pathlib
from collections import Counter

import pytest
from pptx import Presentation

from presentation_designer.layout.composite_charts import find_charts, swap_composites

pytestmark = pytest.mark.organizer_data

VKTECH = "VK Tech шаблон.pptx"
VKEDU = "Шаблон презентации VK Education.pptx"
EXPECTED = {
    VKTECH: {41: ["bar", "bar"], 44: ["column"], 45: ["doughnut"] * 3, 46: ["doughnut"]},
    VKEDU: {46: ["doughnut"] * 4},
    "VK_WorkSpace_Клиентская_конференция_Шаблон_03.pptx": {},
    "ЛЦТ2026 Шаблон презентации.pptx": {},
}


@pytest.mark.parametrize("name", sorted(EXPECTED))
def test_only_real_charts_are_found(organizer_dir: pathlib.Path, name: str) -> None:
    prs = Presentation(str(organizer_dir / name))
    found = {
        n: sorted(spec.kind for spec in specs)
        for n, slide in enumerate(prs.slides, start=1)
        if (specs := find_charts(slide))
    }
    assert found == {n: sorted(kinds) for n, kinds in EXPECTED[name].items()}


def test_vk_tech_charts_keep_their_shape(organizer_dir: pathlib.Path) -> None:
    slides = Presentation(str(organizer_dir / VKTECH)).slides
    (columns,) = find_charts(slides[43])
    # 11 мест по два наложенных ряда; «10» на всех столбцах — значения по рисунку.
    assert len(columns.categories) == 11 and len(columns.values) == 2
    assert columns.overlap == 100 and columns.basis == "measured"
    assert max(max(row) for row in columns.values) == 10
    assert columns.legend is not None and columns.category_style is not None
    rings = find_charts(slides[44])
    assert sorted(spec.values[0][0] for spec in rings) == [7, 24, 71]
    (wedge,) = find_charts(slides[45])
    assert wedge.values == [[98, 2]]
    butterfly = find_charts(slides[40])
    # Подписи категорий посередине общие для двух половин — остаются надписями слайда.
    assert [spec.category_style for spec in butterfly] == [None, None]
    assert sorted(spec.reverse for spec in butterfly) == [False, True]


def test_vk_education_rings_take_the_printed_numbers(organizer_dir: pathlib.Path) -> None:
    rings = find_charts(Presentation(str(organizer_dir / VKEDU)).slides[45])
    assert all(spec.basis == "label" for spec in rings)
    assert sorted(spec.values[0][0] for spec in rings) == [12, 22, 32, 42]


def test_vk_tech_swap_replaces_every_found_chart(organizer_dir: pathlib.Path) -> None:
    prs = Presentation(str(organizer_dir / VKTECH))
    swaps = swap_composites(prs, slides=[prs.slides[k] for k in (40, 43, 44, 45)])
    assert Counter(s.status for s in swaps) == {"replaced": 7}
    assert sorted({s.slide for s in swaps}) == [41, 44, 45, 46]
