"""Контур и скругление фигуры: что из PPTX доезжает до ComposedDeck.

Ради этого тест и написан: карточки шаблонов VK не хранят контур у самой фигуры — там пустой
`a:ln`, а цвет и толщина живут в стиле темы (`p:style/a:lnRef`). Пока это не читалось, рамки
карточек пропадали и на холсте редактора, и в HTML: слайд выглядел иначе, чем в PowerPoint.
"""

from __future__ import annotations

from typing import Any

from lxml import etree
from pptx import Presentation
from pptx.util import Emu

from presentation_designer.layout.composed import _line_entry
from presentation_designer.parsing.template.geometry import walk_shapes
from presentation_designer.parsing.template.styles import Theme, parse_theme

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
}


def _shape_with(sp_pr_xml: str, style_xml: str = "") -> Any:
    """Слайд с одной фигурой: spPr и p:style задаются как в файле шаблона."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Emu(914400), Emu(914400), Emu(2000000), Emu(1000000))
    element = box._element
    old = element.find("p:spPr", NS)
    element.replace(old, etree.fromstring(sp_pr_xml))
    if style_xml:
        element.insert(
            list(element).index(element.find("p:spPr", NS)) + 1, etree.fromstring(style_xml)
        )
    infos = walk_shapes(slide, slide.part, prs.slide_width, prs.slide_height)
    return next(i for i in infos if i.element_id == str(box.shape_id)), prs


SP_PR_CARD = (
    '<p:spPr xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
    ' xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
    '<a:xfrm><a:off x="914400" y="914400"/><a:ext cx="2360141" cy="1876461"/></a:xfrm>'
    '<a:prstGeom prst="roundRect"><a:avLst><a:gd name="adj" fmla="val 9821"/></a:avLst>'
    "</a:prstGeom><a:noFill/><a:ln/></p:spPr>"
)
STYLE_ACCENT2 = (
    '<p:style xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
    ' xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
    '<a:lnRef idx="2"><a:schemeClr val="accent2"/></a:lnRef>'
    '<a:fillRef idx="0"><a:schemeClr val="accent2"/></a:fillRef>'
    '<a:effectRef idx="0"><a:schemeClr val="accent2"/></a:effectRef>'
    '<a:fontRef idx="minor"><a:schemeClr val="lt1"/></a:fontRef></p:style>'
)


def test_adjust_value_is_read_from_the_file() -> None:
    """Скругление берётся из «adj», а не угадывается: 9821/100000 у карточки VK Tech."""
    info, _ = _shape_with(SP_PR_CARD)
    assert info.geometry == "roundRect"
    assert info.geometry_adjust == 0.09821


def test_outline_from_theme_style_reaches_composed_deck() -> None:
    """У фигуры контур не записан — он в стиле темы; в описании колоды он должен быть."""
    info, _ = _shape_with(SP_PR_CARD, STYLE_ACCENT2)
    assert info.line_hex is None, "у самой фигуры цвета нет"
    assert info.line_ref_idx == 2

    theme = Theme(
        name="t",
        colors={"accent2": "#E4002B"},
        major_font="Play",
        minor_font="Play",
        line_widths_pt=[0.75, 1.5, 3.0],
    )
    assert _line_entry(info, theme) == {"color": "#E4002B", "width_pt": 1.5}


def test_own_outline_wins_over_theme() -> None:
    """Свой цвет и своя толщина важнее стиля: их и записываем."""
    sp_pr = (
        '<p:spPr xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
        ' xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<a:xfrm><a:off x="0" y="0"/><a:ext cx="100000" cy="100000"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
        '<a:ln w="25400"><a:solidFill><a:srgbClr val="123456"/></a:solidFill></a:ln></p:spPr>'
    )
    info, _ = _shape_with(sp_pr, STYLE_ACCENT2)
    theme = Theme(name="t", colors={"accent2": "#E4002B"}, major_font="P", minor_font="P")
    assert _line_entry(info, theme) == {"color": "#123456", "width_pt": 2.0}
    assert info.geometry_adjust is None, "у прямоугольника скруглять нечего"


def test_no_line_means_no_entry() -> None:
    """Явный `a:noFill` в контуре — это «линии нет», а не «возьми из темы»."""
    sp_pr = (
        '<p:spPr xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
        ' xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<a:xfrm><a:off x="0" y="0"/><a:ext cx="100000" cy="100000"/></a:xfrm>'
        '<a:prstGeom prst="roundRect"><a:avLst/></a:prstGeom>'
        "<a:ln><a:noFill/></a:ln></p:spPr>"
    )
    info, _ = _shape_with(sp_pr, STYLE_ACCENT2)
    theme = Theme(name="t", colors={"accent2": "#E4002B"}, major_font="P", minor_font="P")
    assert _line_entry(info, theme) is None


def test_theme_line_widths_are_parsed() -> None:
    """Толщины линий темы читаются из fmtScheme: по ним считается контур из стиля."""
    xml = (
        '<a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" name="T">'
        '<a:themeElements><a:clrScheme name="c"/><a:fontScheme name="f"/>'
        '<a:fmtScheme name="s"><a:lnStyleLst>'
        '<a:ln w="6350"/><a:ln w="19050"/><a:ln w="38100"/>'
        "</a:lnStyleLst></a:fmtScheme></a:themeElements></a:theme>"
    )
    theme = parse_theme(etree.fromstring(xml))
    assert theme.line_widths_pt == [0.5, 1.5, 3.0]
