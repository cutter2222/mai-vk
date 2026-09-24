"""Диаграммы из фигур слайда → нативные: столбцы-картинки с числами у концов, подписями
категорий и легендой; кольца с числом-процентом в центре (в том числе из двух картинок).
Числа, сходящиеся с рисунком, берутся точно; «рыба» уступает рисунку. Ряд одинаковых
карточек и кольцо без числа не трогаются; сорвавшаяся постройка не портит слайд; сборка
варианта original повторяет замены."""

from __future__ import annotations

import io
import pathlib
import zipfile
from copy import deepcopy
from typing import Any

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from presentation_designer.generation.original import original_plan, original_story
from presentation_designer.layout import composite_charts
from presentation_designer.layout.chart_images import ChartSwap, charts_report
from presentation_designer.layout.compose import compose_deck
from presentation_designer.layout.composite_charts import find_charts, swap_composites
from presentation_designer.parsing.content.importer import import_content
from presentation_designer.parsing.content.parsers import ParseCache
from presentation_designer.shared.settings import Settings
from tests.layout.conftest import MINI_TEMPLATE, own_profile
from tests.parsing.content.conftest import material
from tests.parsing.raster_charts import synth

BASE = Inches(4.5)  # основание столбцов
BAR_W = Inches(0.6)
PITCH = Inches(0.9)


def _deck() -> tuple[Any, Any]:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(9144000), Emu(5143500)
    return prs, prs.slides.add_slide(prs.slide_layouts[6])


def _text(slide: Any, text: str, box: tuple[int, int, int, int], *, size: int = 14) -> Any:
    tb = slide.shapes.add_textbox(*box)
    tf = tb.text_frame
    tf.text = text
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.runs[0]
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    return tb


def _columns(
    slide: Any,
    heights: list[float],
    labels: list[str],
    *,
    color: tuple[int, int, int] = synth.BLUE,
    names: list[str] | None = None,
    x0: int = Inches(0.8),
) -> None:
    """Столбцы-картинки на общем основании, числа над ними, подписи категорий под ними."""
    for k, (h, label) in enumerate(zip(heights, labels, strict=True)):
        x = x0 + k * PITCH
        top = BASE - Inches(h)
        data = synth.bar_piece(60, round(60 * h / 0.6), color, 1.0, 0.2)
        slide.shapes.add_picture(io.BytesIO(data), x, top, BAR_W, Inches(h))
        box = (x - Inches(0.2), top - Inches(0.4), BAR_W + Inches(0.4), Inches(0.35))
        _text(slide, label, box)
        if names:
            _text(
                slide,
                names[k],
                (x - Inches(0.1), BASE + Inches(0.05), BAR_W + Inches(0.2), Inches(0.3)),
                size=10,
            )


def _chart_frames(slide: Any) -> list[Any]:
    return [s for s in slide.shapes if s.has_chart]


def _pictures(slide: Any) -> list[Any]:
    return [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]


def test_columns_with_matching_numbers_become_a_chart_with_those_numbers() -> None:
    prs, slide = _deck()
    names = ["Янв", "Фев", "Мар", "Апр", "Май"]
    _columns(slide, [1.0, 2.0, 3.0, 2.5, 1.5], ["10", "20", "30", "25", "15"], names=names)
    (spec,) = find_charts(slide)
    assert spec.kind == "column" and spec.basis == "label"
    assert spec.values == [[10, 20, 30, 25, 15]]
    assert spec.categories == names
    assert spec.label_position == "outEnd"

    swaps = swap_composites(prs)
    assert [(s.status, s.kind) for s in swaps] == [("replaced", "pieces")]
    (frame,) = _chart_frames(slide)
    assert not _pictures(slide)
    # Числа и подписи категорий ушли в диаграмму: отдельных надписей на слайде не осталось.
    assert not [s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
    chart = frame.chart
    assert list(chart.plots[0].series[0].values) == [10, 20, 30, 25, 15]
    assert list(chart.plots[0].categories) == names
    assert chart.plots[0].has_data_labels
    assert "из фигур" in frame.name
    descr = frame._element.nvGraphicFramePr.cNvPr.get("descr")
    assert "значения перенесены из подписей" in descr


def test_placeholder_numbers_give_way_to_the_drawing() -> None:
    prs, slide = _deck()
    _columns(slide, [1.0, 2.0, 3.0, 1.5], ["10", "10", "10", "10"])
    (spec,) = find_charts(slide)
    assert spec.basis == "measured"
    # Самый длинный столбец равен самому большому числу, остальные — пропорционально.
    assert spec.values == [[pytest.approx(3.3, abs=0.1), 6.7, 10.0, 5.0]]
    swap_composites(prs)
    (frame,) = _chart_frames(slide)
    assert "сняты по рисунку" in frame._element.nvGraphicFramePr.cNvPr.get("descr")


def test_two_overlapping_series_take_names_from_the_drawn_legend() -> None:
    prs, slide = _deck()
    dark, light = (0, 119, 255), (150, 220, 255)
    # На одном месте два столбца: короткий — спереди, как в шаблоне VK Tech.
    tall = [2.0, 2.4, 1.8, 2.2]
    short = [1.2, 1.9, 1.0, 1.6]
    for k in range(4):
        x = Inches(0.8) + k * PITCH
        for h, color in ((tall[k], light), (short[k], dark)):
            data = synth.bar_piece(60, round(60 * h / 0.6), color, 1.0, 1.0)
            slide.shapes.add_picture(io.BytesIO(data), x, BASE - Inches(h), BAR_W, Inches(h))
            _text(
                slide, str(round(h * 10)), (x, BASE - Inches(h) + Inches(0.02), BAR_W, Inches(0.3))
            )
    for k, (color, name) in enumerate(((dark, "План"), (light, "Факт"))):
        y = Inches(0.4) + k * Inches(0.3)
        marker = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(6), y, Inches(0.15), Inches(0.15))
        marker.fill.solid()
        marker.fill.fore_color.rgb = RGBColor(*color)
        marker.line.fill.background()
        _text(slide, name, (Inches(6.2), y - Inches(0.05), Inches(1), Inches(0.25)), size=10)
    (spec,) = find_charts(slide)
    assert spec.overlap == 100 and spec.basis == "label"
    assert sorted(spec.names) == ["План", "Факт"]
    assert spec.legend is not None
    # Спереди — ряд, который короче в большинстве категорий.
    assert spec.names[-1] == "План"
    swap_composites(prs)
    (frame,) = _chart_frames(slide)
    assert frame.chart.has_legend
    assert not [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE]


def _ring(slide: Any, percent: float, label: str, *, x: int = Inches(1)) -> None:
    arcs = [(0.0, percent * 3.6, synth.BLUE), (percent * 3.6, 360 - percent * 3.6, synth.TRACK)]
    slide.shapes.add_picture(
        io.BytesIO(synth.ring_piece(300, 0.66, arcs)), x, Inches(1), Inches(2.4), Inches(2.4)
    )
    _text(slide, label, (x + Inches(0.4), Inches(1.95), Inches(1.6), Inches(0.5)), size=28)


def test_ring_with_a_percent_in_the_middle_becomes_a_doughnut_with_its_label() -> None:
    prs, slide = _deck()
    _ring(slide, 42, "42%")
    (spec,) = find_charts(slide)
    assert spec.kind == "doughnut" and spec.basis == "label"
    assert spec.values == [[42, 58]]
    assert spec.number_format == '0"%"'
    swap_composites(prs)
    (frame,) = _chart_frames(slide)
    assert not _pictures(slide)
    assert not [s for s in slide.shapes if s.has_text_frame]
    visible, center = frame.chart._chartSpace.xpath(".//c:ser")
    # Число в центре несёт второе кольцо — невидимый полный круг: его подпись не ездит за
    # серединой дуги, когда значение правят. У видимой дуги подписей нет.
    assert visible.find(qn("c:dLbls")) is None
    labels = center.find(qn("c:dLbls")).findall(qn("c:dLbl"))
    assert labels[0].find(qn("c:showVal")).get("val") == "1"
    assert labels[0].find(qn("c:layout")) is not None
    assert labels[1].find(qn("c:delete")).get("val") == "1"
    assert list(frame.chart.plots[0].series[1].values) == [42, 0]
    # Остаток и число в центре во встроенной книге — формулы: правится одно число B2.
    blob = frame.chart.part.chart_workbook.xlsx_part.blob
    with zipfile.ZipFile(io.BytesIO(blob)) as zf:
        sheet = zf.read("xl/worksheets/sheet1.xml").decode()
    assert "<f>100-B2</f>" in sheet
    assert "<f>MAX(B2,0.001)</f>" in sheet


def test_ring_of_two_pictures_is_measured_together() -> None:
    prs, slide = _deck()
    full = synth.ring_piece(300, 0.66, [(0.0, 360.0, synth.BLUE)])
    slide.shapes.add_picture(io.BytesIO(full), Inches(1), Inches(1), Inches(2.4), Inches(2.4))
    # Серая дуга поверх градиентного кольца — «остаток», нарисованный отдельной картинкой.
    arc = synth.ring_piece(300, 0.66, [(0.0, 90.0, synth.TRACK)])
    slide.shapes.add_picture(io.BytesIO(arc), Inches(1), Inches(1), Inches(2.4), Inches(2.4))
    _text(slide, "10%", (Inches(1.4), Inches(1.95), Inches(1.6), Inches(0.5)), size=28)
    (spec,) = find_charts(slide)
    # На картинке 75 %, а в центре «10%» — значение по рисунку.
    assert spec.basis == "measured"
    assert spec.values == [[75, 25]]
    assert len(spec.pieces) == 3
    swap_composites(prs)
    assert not _pictures(slide)


def test_row_of_equal_cards_with_numbers_is_not_a_chart() -> None:
    _, slide = _deck()
    _columns(slide, [1.5, 1.5, 1.5, 1.5], ["1", "2", "3", "4"])
    assert find_charts(slide) == []


def test_single_ring_without_a_number_is_left_to_the_picture_reader() -> None:
    _, slide = _deck()
    arcs = [(0.0, 150.0, synth.BLUE), (150.0, 210.0, synth.TRACK)]
    slide.shapes.add_picture(
        io.BytesIO(synth.ring_piece(300, 0.66, arcs)),
        Inches(1),
        Inches(1),
        Inches(2.4),
        Inches(2.4),
    )
    assert find_charts(slide) == []


def test_failed_build_leaves_the_slide_as_it_was(monkeypatch: pytest.MonkeyPatch) -> None:
    prs, slide = _deck()
    _ring(slide, 42, "42%")
    before = [s.shape_id for s in slide.shapes]
    rels = set(slide.part.rels)
    real = composite_charts._build_doughnut

    def broken(slide: Any, spec: Any) -> Any:
        real(slide, spec)
        raise RuntimeError("сбой после добавления диаграммы")

    monkeypatch.setattr(composite_charts, "_build_doughnut", broken)
    (swap,) = swap_composites(prs)
    assert swap.status == "kept" and "сбой" in swap.reason
    assert [s.shape_id for s in slide.shapes] == before
    assert not _chart_frames(slide)
    assert set(slide.part.rels) == rels


def test_report_counts_assembled_charts_for_the_full_build() -> None:
    swaps = [
        ChartSwap(44, "столбцы из фигур", "", "replaced", measured=22, kind="pieces"),
        ChartSwap(45, "кольцо из фигур", "", "kept", "не построена", kind="pieces"),
    ]
    report = charts_report({}, {}, swaps, None)
    assert report["composites"] == 1
    assert [s["kind"] for s in report["swaps"]] == ["pieces", "pieces"]


def _mini_with_ring(tmp_path: pathlib.Path) -> pathlib.Path:
    prs = Presentation(str(MINI_TEMPLATE))
    slide = prs.slides[len(prs.slides) - 1]
    _ring(slide, 42, "42%", x=Inches(6))
    path = tmp_path / "mini_ring.pptx"
    prs.save(str(path))
    return path


def test_original_build_repeats_assembled_charts(
    tmp_path: pathlib.Path, import_settings: Settings
) -> None:
    source = _mini_with_ring(tmp_path)
    profile = own_profile(source, "tpl_ring")
    package = dict(
        import_content(
            [material(source)],
            {"purpose": "other", "title": "Мини с кольцом"},
            package_id="pkg_ring",
            settings=import_settings,
            cache=ParseCache(tmp_path / "import-cache"),
            use_model=False,
        ).package
    )
    story = original_story(package, profile, {"language": "ru"})
    plan = original_plan(story, profile, package, plan_id="plan_ring")

    plain = compose_deck(plan, profile, source, package, out_pptx=tmp_path / "plain.pptx")
    assert not plain.chart_swaps
    assert (tmp_path / "plain.pptx").read_bytes() == source.read_bytes()

    swapped = compose_deck(
        deepcopy(plan), profile, source, package, out_pptx=tmp_path / "ring.pptx", composites=True
    )
    assert [(s.status, s.kind) for s in swapped.chart_swaps] == [("replaced", "pieces")]
    last = Presentation(str(tmp_path / "ring.pptx")).slides[-1]
    assert _chart_frames(last)
    assert not _pictures(last)


@pytest.mark.parametrize(
    ("top", "end"),
    [(10, 12), (9.1, 10), (7.3, 8), (11, 12), (58, 70), (100, 120), (1234, 1400), (3.3, 3.5)],
)
def test_axis_end_matches_the_onlyoffice_auto_scale(top: float, end: float) -> None:
    # Замер рендера ONLYOFFICE на сервере 24.09.2026: конец оси при минимуме 0.
    assert composite_charts.auto_max(top) == pytest.approx(end)


def test_value_axis_end_is_automatic_and_the_drawing_is_kept() -> None:
    prs, slide = _deck()
    _columns(slide, [1.0, 2.0, 3.0, 1.5], ["10", "20", "30", "15"])
    (spec,) = find_charts(slide)
    # Самый длинный столбец (30) на рисунке — 3 дюйма; конец оси 35 — область на 3,5 дюйма.
    assert spec.plot[3] == pytest.approx(Inches(3.0) * 35 / 30, rel=0.01)
    assert spec.plot[1] + spec.plot[3] == pytest.approx(BASE, abs=Inches(0.01))
    swap_composites(prs)
    (frame,) = _chart_frames(slide)
    scaling = frame.chart._chartSpace.xpath(".//c:valAx/c:scaling")[0]
    assert scaling.find(qn("c:max")) is None
    assert scaling.find(qn("c:min")).get("val") == "0.0"
