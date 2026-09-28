"""Записанные ответы проходят настоящий планировщик и сборку редактируемого PPTX."""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

import pytest
from pptx import Presentation

from presentation_designer.generation import capacity as cap
from presentation_designer.generation import variants as vr
from presentation_designer.layout.compose import compose_deck
from presentation_designer.layout.shapes import iter_shapes
from presentation_designer.llm.skills import get_skill
from presentation_designer.pipeline.run import resolve_slide_count
from presentation_designer.shared.slide_text import plain
from tests.generation.test_variants import _assert_overflow_reported, _assert_valid
from tests.layout.conftest import MINI_TEMPLATE


def _normalized(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


@pytest.mark.parametrize("template", ["mini", "rich"])
@pytest.mark.parametrize("variant", vr.VARIANTS)
def test_replay_plan_content_survives_pptx_composition(
    template: str,
    variant: str,
    mini_profile: dict[str, Any],
    rich_profile: dict[str, Any],
    rich_template_path: Path,
    example_package: dict[str, Any],
    example_story: dict[str, Any],
    replay_client: Any,
    tmp_path: Path,
) -> None:
    profile = mini_profile if template == "mini" else rich_profile
    source = MINI_TEMPLATE if template == "mini" else rich_template_path
    package = {k: v for k, v in example_package.items() if not k.startswith("_")}
    result = vr.build_variant_plan(
        example_story,
        profile,
        package,
        variant,
        {"language": "ru"},
        slide_count=resolve_slide_count(variant, {"language": "ru"}),
        client=replay_client,
        skill=get_skill("variant_planner"),
        use_model=True,
    )
    plan = result.plan
    _assert_valid(plan, profile, package, example_story)
    _assert_overflow_reported(plan, result.report)
    # Повтор пакета по вместимости зависит от записанного ответа: если он был и терял
    # содержание, повтор отклонён (сама логика — в test_capacity_retry_content.py).
    codes = {f["code"] for f in result.report["fixes"]}
    if "packet_retried" in codes and "packet_retry_rejected_content" not in codes:
        assert result.plan["coverage"]["missing"] == []
    composed = compose_deck(
        plan,
        profile,
        source,
        package,
        out_pptx=tmp_path / "deck.pptx",
        package_dir=Path(example_package["_assets_dir"]),
        job_id="job_replay_content",
    )
    # Эти артефакты можно конвертировать отдельно: тест не подменяет PDF картинками.
    for name, data in {
        "plan": plan,
        "profile": profile,
        "package": package,
        "planning-report": result.report,
        "composed": composed.deck,
    }.items():
        (tmp_path / f"{name}.json").write_text(
            json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8"
        )
    assert composed.integrity.ok
    presentation = Presentation(composed.pptx_path)
    assert len(presentation.slides) == len(plan["slides"])
    facts = {f["fact_id"]: f for f in package["facts"]}
    for slide, planned in zip(presentation.slides, plan["slides"], strict=True):
        visible = _normalized(
            "\n".join(s.text_frame.text for s in iter_shapes(slide) if s.has_text_frame)
        )
        assert "{fact:" not in visible
        for block in planned["blocks"]:
            texts = [str(block.get("text") or "")]
            texts.extend(
                str(item.get(key) or "")
                for item in block.get("items", [])
                for key in ("text", "sub")
            )
            for text in texts:
                # Разметка **жирного** в плане становится начертанием, в тексте PPTX её нет.
                expected = _normalized(plain(cap.substitute_facts(text, facts)))
                if expected:
                    assert expected in visible, (planned["slide_id"], block["slot_id"], expected)
        if planned.get("notes"):
            assert _normalized(planned["notes"]) in _normalized(
                slide.notes_slide.notes_text_frame.text
            )
