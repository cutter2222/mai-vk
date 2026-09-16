"""Операции пакета на шаблонах организаторов: без data/organizers пропускаются."""

from __future__ import annotations

import json
import pathlib

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from presentation_designer.layout import ooxml

pytestmark = pytest.mark.organizer_data


def _templates(organizer_dir: pathlib.Path) -> list[pathlib.Path]:
    manifest = json.loads((organizer_dir / "manifest.json").read_text())
    root = organizer_dir.parent.parent
    return [root / e["path"] for e in manifest["templates"]]


def test_templates_pass_package_check(organizer_dir: pathlib.Path) -> None:
    for path in _templates(organizer_dir):
        report = ooxml.check_package(path)
        assert report.ok, (path.name, report.errors)


def test_roundtrip_keeps_parts_and_binaries(
    organizer_dir: pathlib.Path, tmp_path: pathlib.Path
) -> None:
    for path in _templates(organizer_dir):
        out = tmp_path / path.name
        Presentation(str(path)).save(str(out))
        diff = ooxml.compare_packages(path, out)
        assert diff["only_in_a"] == [] and diff["only_in_b"] == [], path.name
        assert diff["changed_binary"] == [], path.name
        assert all(n == "[Content_Types].xml" for n in diff["changed_xml"]), (path.name, diff)
        assert ooxml.check_package(out).ok, path.name


def test_clone_group_slide_and_prune(organizer_dir: pathlib.Path, tmp_path: pathlib.Path) -> None:
    for path in _templates(organizer_dir):
        prs = Presentation(str(path))
        source = next(
            s for s in prs.slides if any(sh.shape_type == MSO_SHAPE_TYPE.GROUP for sh in s.shapes)
        )
        clone = ooxml.clone_slide(prs, source)
        assert ooxml.shape_tree_signature(clone) == ooxml.shape_tree_signature(source), path.name
        ooxml.keep_only_slides(prs, [clone])
        out = tmp_path / path.name
        prs.save(str(out))
        report = ooxml.check_package(out)
        assert report.ok and report.slides == 1, (path.name, report.errors)
        assert len(Presentation(str(out)).slides) == 1
