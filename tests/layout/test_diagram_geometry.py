"""Native diagram geometry regressions found during ONLYOFFICE visual acceptance."""

import math

import pytest
from pptx import Presentation
from pptx.util import Inches

from presentation_designer.layout.diagrams import DiagramStyle, content_fits, draw_diagram


@pytest.mark.parametrize("count", [2, 3])
def test_venn_labels_are_above_all_circles_and_inside_slot(count: int) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = (Inches(1), Inches(1), Inches(10), Inches(4))
    block = {"kind": "venn", "items": [{"text": "Старт", "sub": "Согласие"}] * count}
    style = DiagramStyle.from_profile({})
    assert content_fits(block, box, style)
    result = draw_diagram(slide, box, block, style)
    shapes = list(slide.shapes[-1].shapes)
    circles, labels = shapes[:count], shapes[count:]
    assert len(labels) == count
    assert all(not circle.text for circle in circles)
    assert all(label.text == "Старт\nСогласие" for label in labels)
    assert all(str(label.shape_id) in result.node_ids for label in labels)
    for label in shapes:
        assert box[0] <= label.left < label.left + label.width <= box[0] + box[2]
        assert box[1] <= label.top < label.top + label.height <= box[1] + box[3]
    for i, label in enumerate(labels):
        for other in labels[i + 1 :]:
            assert (
                label.left + label.width <= other.left
                or other.left + other.width <= label.left
                or label.top + label.height <= other.top
                or other.top + other.height <= label.top
            )


@pytest.mark.parametrize("count", range(2, 7))
def test_cycle_has_visible_arrow_shafts(count: int) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    draw_diagram(
        slide,
        (Inches(1), Inches(1), Inches(10), Inches(4)),
        {"kind": "cycle", "items": [{"text": str(i)} for i in range(count)]},
        DiagramStyle.from_profile({}),
    )
    shapes = list(slide.shapes[-1].shapes)
    arrows = shapes[count:]
    assert len(arrows) == count
    assert all(math.hypot(arrow.width, arrow.height) / 12700 > 12 for arrow in arrows)


@pytest.mark.parametrize("count", range(2, 7))
def test_matrix_cards_stay_inside_pane_and_preserve_text(count: int) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = (Inches(1), Inches(1), Inches(10), Inches(4))
    block = {
        "kind": "matrix",
        "items": [{"text": f"Риск {i}", "sub": "Резерв"} for i in range(count)],
    }
    style = DiagramStyle.from_profile({})
    assert content_fits(block, box, style)
    result = draw_diagram(slide, box, block, style)
    nodes = list(slide.shapes[-1].shapes)
    assert len(result.node_ids) == count
    for i, node in enumerate(nodes):
        assert node.text == f"Риск {i}\nРезерв"
        assert box[1] <= node.top < node.top + node.height <= box[1] + box[3]
        assert box[0] <= node.left < node.left + node.width <= box[0] + box[2]
    if count == 2:
        assert all(node.height == box[3] // 2 for node in nodes)
        assert all(node.top == box[1] + box[3] // 4 for node in nodes)
    else:
        assert min(node.top for node in nodes) == box[1]


def test_matrix_long_text_expands_cards_without_shrinking_font() -> None:
    box = (Inches(1), Inches(1), Inches(10), Inches(4))
    style = DiagramStyle.from_profile({})
    block = {"kind": "matrix", "items": [{"text": "Риск", "sub": "Резервный канал " * 20}] * 2}
    assert content_fits(block, box, style)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    draw_diagram(slide, box, block, style)
    for node in slide.shapes[-1].shapes:
        assert box[3] // 2 < node.height <= box[3]
        assert node.text == "Риск\n" + block["items"][0]["sub"].strip()
        assert node.text_frame.paragraphs[1].runs[0].font.size.pt == style.sub_size_pt
    block["items"] = [{"text": "Риск", "sub": "Резервный канал " * 200}] * 2
    assert not content_fits(block, box, style)
