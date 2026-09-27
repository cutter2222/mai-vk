"""Presentation emphasis is native DrawingML and shares plain text with measurement."""

from lxml import etree
from pptx import Presentation
from pptx.util import Inches

from presentation_designer.generation import capacity
from presentation_designer.layout import text as tx
from presentation_designer.layout.shapes import NS
from presentation_designer.shared.slide_text import plain, spans
from tests.generation.test_page_density import pattern


def test_only_balanced_strong_markers_are_interpreted():
    assert spans("Оценка **после проверки** рисков") == [
        ("Оценка ", False),
        ("после проверки", True),
        (" рисков", False),
    ]
    for literal in ["2 * 3", "**не закрыто", "file_name", "# Title", "***nested***"]:
        assert plain(literal) == literal


def test_native_runs_preserve_fact_references_and_survive_pptx(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(9), Inches(4))
    facts = {"f1": {"raw": "12 млн ₽"}}
    result = tx.fill_text(
        shape._element,
        "Оценка **{fact:f1}** только после проверки",
        facts=facts,
        size_pt=24,
        markup=True,
    )
    assert result.plain == "Оценка 12 млн ₽ только после проверки"
    assert result.fact_refs == ["f1"]
    path = tmp_path / "emphasis.pptx"
    prs.save(path)
    actual = Presentation(path).slides[0].shapes[0]
    assert actual.text == result.plain
    runs = actual.text_frame.paragraphs[0].runs
    assert runs[1].font.bold and runs[1].text == "12 млн ₽"
    assert all(r.font.size.pt == 24 for r in runs)


def test_lists_have_separate_paragraphs_and_markup_is_opt_in():
    prs = Presentation()
    shape = prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(
        Inches(1), Inches(1), Inches(9), Inches(4)
    )
    result = tx.fill_bullets(shape._element, ["**Риск** нагрузки", "Условие согласия"], markup=True)
    assert result.plain == "Риск нагрузки\nУсловие согласия"
    assert len(shape.text_frame.paragraphs) == 2
    tx.fill_text(shape._element, "**literal**")
    assert shape.text == "**literal**"


def test_capacity_uses_visible_chars_and_conservative_bold_metrics():
    slot = pattern().slots["body"]
    slot.size_pt = 24
    marked = capacity.measure("Оценка **после проверки**", slot, 12192000, 6858000)
    assert marked.chars == len("Оценка после проверки")
    slot.kind = "code"
    literal = capacity.measure("**literal**", slot, 12192000, 6858000)
    assert literal.chars == len("**literal**")


def test_native_bullet_marker_survives_emphasis_and_reopen(tmp_path):
    prs = Presentation()
    shape = prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(
        Inches(1), Inches(1), Inches(9), Inches(4)
    )
    paragraph = shape.text_frame.paragraphs[0]
    marker = etree.SubElement(paragraph._p.get_or_add_pPr(), f"{{{NS['a']}}}buChar")
    marker.set("char", "•")
    tx.fill_bullets(shape._element, ["**Риск** нагрузки", "Условие согласия"], markup=True)
    path = tmp_path / "bullets.pptx"
    prs.save(path)
    paragraphs = Presentation(path).slides[0].shapes[0].text_frame.paragraphs
    assert len(paragraphs) == 2
    assert paragraphs[0].runs[0].font.bold
    for p in paragraphs:
        assert p._p.find("a:pPr/a:buChar", NS).get("char") == "•"
        assert "•" not in p.text and "**" not in p.text
