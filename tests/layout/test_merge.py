"""Слияние пакетов: id мастеров и макетов не сталкиваются, пакет остаётся целым."""

from __future__ import annotations

import collections
import pathlib

from pptx import Presentation

from presentation_designer.audit.deterministic import check_package
from presentation_designer.layout.merge import (
    _SLD_LAYOUT_ID,
    insert_slide,
    merge_presentation,
    replace_slide,
)


def _ids(prs) -> list[int]:
    lst = prs.part._element.get_or_add_sldMasterIdLst()
    ids = [int(e.get("id")) for e in lst.sldMasterId_lst]
    for master in prs.slide_masters:
        ids += [int(e.get("id")) for e in master.part._element.iter(_SLD_LAYOUT_ID)]
    return ids


def test_merged_masters_get_distinct_layout_ids(tmp_path: pathlib.Path) -> None:
    base = Presentation()
    source = Presentation()
    source.slides.add_slide(source.slide_layouts[1]).shapes.title.text = "Из второго файла"
    result = merge_presentation(base, source)
    assert result.masters == 1 and result.slides == 1
    assert len(base.slide_masters) == 2
    ids = _ids(base)
    assert [v for v in collections.Counter(ids).values() if v > 1] == []
    assert min(ids) >= 2147483648
    out = tmp_path / "merged.pptx"
    base.save(out)
    assert check_package(out) == []


FIXTURE = pathlib.Path(__file__).resolve().parents[1] / "fixtures" / "pptx" / "mini_template.pptx"


def _slide_ids(prs) -> list[str]:
    return [e.get("id") for e in prs.slides._sldIdLst.sldId_lst]


def test_replace_slide_keeps_manual_edits_and_slide_ids(tmp_path: pathlib.Path) -> None:
    base = Presentation(FIXTURE)
    base.slides[0].shapes.title.text = "Ручная правка"
    ids = _slide_ids(base)
    source = Presentation(FIXTURE)
    # Слайд с диаграммой встаёт на место второго: часть диаграммы и её книга переносятся.
    result = replace_slide(base, 1, source, 2)
    assert result.slides == 1 and result.masters == 0
    out = tmp_path / "replaced.pptx"
    base.save(out)
    assert check_package(out) == []
    saved = Presentation(out)
    assert _slide_ids(saved) == ids
    assert saved.slides[0].shapes.title.text == "Ручная правка"
    assert [s.name for s in saved.slides[1].shapes] == ["TextBox 1", "Table 2", "Chart 3"]
    assert saved.slides[1].shapes[2].has_chart
    assert len(saved.slides) == 4


def test_replace_slide_brings_a_missing_layout_with_its_master(tmp_path: pathlib.Path) -> None:
    base = Presentation(FIXTURE)
    for layout in base.slide_masters[0].slide_layouts:
        if layout.name == "Title Slide":
            layout.name = "Обложка копии"
    source = Presentation(FIXTURE)
    result = replace_slide(base, 3, source, 0)
    assert result.masters == 1 and result.warnings
    out = tmp_path / "replaced.pptx"
    base.save(out)
    assert check_package(out) == []
    saved = Presentation(out)
    assert saved.slides[3].slide_layout.name == "Title Slide"
    assert len(saved.slide_masters) == 2


def test_replace_slide_drops_notes_and_links_to_other_slides(tmp_path: pathlib.Path) -> None:
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    base = Presentation(FIXTURE)
    source = Presentation(FIXTURE)
    slide = source.slides[1]
    slide.notes_slide.notes_text_frame.text = "Заметка докладчика"
    rid = slide.part.relate_to(source.slides[3].part, RT.SLIDE)
    run = slide.shapes[0].text_frame.paragraphs[0].runs[0]
    link = run._r.get_or_add_rPr().add_hlinkClick(rid)
    link.set("action", "ppaction://hlinksldjump")
    replace_slide(base, 0, source, 1)
    out = tmp_path / "replaced.pptx"
    base.save(out)
    assert check_package(out) == []
    saved = Presentation(out)
    assert not saved.slides[0].has_notes_slide
    assert saved.slides[0].shapes[0].text_frame.text == slide.shapes[0].text_frame.text


def test_replace_slide_rejects_a_missing_slide() -> None:
    import pytest

    from presentation_designer.layout.merge import MergeError

    base = Presentation(FIXTURE)
    with pytest.raises(MergeError):
        replace_slide(base, 9, Presentation(FIXTURE), 0)


def test_insert_slide_puts_a_copy_after_the_current_one(tmp_path: pathlib.Path) -> None:
    import pytest

    from presentation_designer.layout.merge import MergeError

    base = Presentation(FIXTURE)
    base.slides[0].shapes.title.text = "Ручная правка"
    ids = _slide_ids(base)
    source = Presentation(FIXTURE)
    # Слайд шаблона с диаграммой встаёт вторым: у остальных слайдов прежние записи и правки.
    result = insert_slide(base, 1, source, 2)
    assert result.slides == 1 and result.masters == 0
    out = tmp_path / "inserted.pptx"
    base.save(out)
    assert check_package(out) == []
    saved = Presentation(out)
    got = _slide_ids(saved)
    assert len(got) == 5 and [got[0], *got[2:]] == ids and got[1] not in ids
    assert saved.slides[0].shapes.title.text == "Ручная правка"
    assert [s.name for s in saved.slides[1].shapes] == ["TextBox 1", "Table 2", "Chart 3"]
    assert saved.slides[1].shapes[2].has_chart
    # В начало и в самый конец — тоже; места за концом и слайда за концом шаблона нет.
    insert_slide(saved, 0, source, 0)
    insert_slide(saved, len(saved.slides), source, 3)
    assert len(saved.slides) == 7
    saved.save(out)
    assert check_package(out) == []
    with pytest.raises(MergeError):
        insert_slide(saved, 9, source, 0)
    with pytest.raises(MergeError):
        insert_slide(saved, 0, source, 9)
