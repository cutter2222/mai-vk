"""Ручные правки объектов слайда (этап 22): текст и стиль на фрагментах, положение внутри
групп, замена картинок и иконок из шаблона, пакета и файла, фон слайда, охрана адреса,
предупреждения о значениях вне токенов и о переполнении, медиа ревизии, детерминизм id."""

from __future__ import annotations

import io
import pathlib
from typing import Any

from pptx import Presentation

from presentation_designer.layout.ooxml import NS_A
from presentation_designer.layout.shapes import iter_shapes
from tests.layout.test_compose import _compose, _package, _plan_with

NS = {"a": NS_A, "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}


def _cards(profile: dict[str, Any]) -> dict[str, Any]:
    return next(p for p in profile["patterns"] if p["role"] == "cards")


def _ref(pattern: dict[str, Any], slot_id: str) -> str:
    return str(next(s for s in pattern["slots"] if s["slot_id"] == slot_id)["element_ref"])


def _cards_plan(profile: dict[str, Any], overrides: list[dict[str, Any]]) -> dict[str, Any]:
    cards = _cards(profile)
    return _plan_with(
        profile,
        [
            {
                "pattern_id": cards["pattern_id"],
                "title": "Три карточки",
                "blocks": [
                    {"slot_id": "title_1", "kind": "title", "text": "Три карточки"},
                    {"slot_id": "label_1", "kind": "label", "text": "Первая"},
                    {"slot_id": "body_1", "kind": "body", "text": "Описание первой"},
                    {"slot_id": "label_2", "kind": "label", "text": "Вторая"},
                    {"slot_id": "body_2", "kind": "body", "text": "Описание второй"},
                    {"slot_id": "label_3", "kind": "label", "text": "Третья"},
                    {"slot_id": "body_3", "kind": "body", "text": "Описание третьей"},
                ],
                "overrides": overrides,
            }
        ],
    )


def _objects(result: Any, index: int = 0) -> dict[str, dict[str, Any]]:
    return {o["object_id"]: o for o in result.deck["slides"][index]["objects"]}


def _shape(result: Any, object_id: str, index: int = 0) -> Any:
    prs = Presentation(str(result.pptx_path))
    slide = list(prs.slides)[index]
    return next(s for s in iter_shapes(slide) if str(s.shape_id) == object_id)


def _png(color: tuple[int, int, int], size: int = 40, alpha: bool = False) -> bytes:
    from PIL import Image

    mode = "RGBA" if alpha else "RGB"
    img = Image.new(mode, (size, size), (*color, 255) if alpha else color)
    if alpha:
        for x in range(size):
            for y in range(size):
                if (x - size // 2) ** 2 + (y - size // 2) ** 2 > (size // 2) ** 2:
                    img.putpixel((x, y), (0, 0, 0, 0))
    out = io.BytesIO()
    img.save(out, format="PNG")
    return out.getvalue()


# ---------- текст и стиль ----------


def test_text_and_style_written_to_runs(
    rich_profile: dict[str, Any],
    rich_template_path: pathlib.Path,
    example_package: dict[str, Any],
    tmp_path: Any,
) -> None:
    cards = _cards(rich_profile)
    title_id = _ref(cards, "title_1")
    body_id = _ref(cards, "body_1")
    plan = _cards_plan(
        rich_profile,
        [
            {
                "op": "text",
                "target": {
                    "object_id": title_id,
                    "source_object_id": title_id,
                    "slot_id": "title_1",
                },
                "text": "Новый заголовок\nвторая строка",
            },
            {
                "op": "style",
                "target": {"object_id": title_id, "slot_id": "title_1"},
                "style": {
                    "font": {
                        "size_pt": 27,
                        "bold": False,
                        "italic": True,
                        "color": "#FF3985",
                        "family": "Comic Sans MS",
                    },
                    "align": "center",
                },
            },
            {
                "op": "style",
                "target": {"object_id": body_id},
                "style": {"font": {"size_pt": 14, "color": "#000000"}},
            },
        ],
    )
    result = _compose(plan, rich_profile, rich_template_path, example_package, tmp_path / "t.pptx")
    title = _shape(result, title_id)
    xml = title._element
    runs = xml.findall(".//a:r/a:rPr", NS)
    assert runs, "фрагменты на месте"
    for rpr in runs:
        assert rpr.get("sz") == "2700" and rpr.get("b") == "0" and rpr.get("i") == "1"
        fill = rpr.find("a:solidFill/a:srgbClr", NS)
        assert fill is not None and fill.get("val") == "FF3985"
        assert rpr.find("a:schemeClr", NS) is None
        assert rpr.find("a:latin", NS).get("typeface") == "Comic Sans MS"
        assert rpr.find("a:cs", NS).get("typeface") == "Comic Sans MS"
        names = [child.tag.split("}")[-1] for child in rpr]
        assert names.index("solidFill") < names.index("latin") < names.index("cs"), (
            "порядок детей rPr по схеме"
        )
    for p in xml.findall(".//a:p", NS):
        ppr = p.find("a:pPr", NS)
        assert ppr is not None and ppr.get("algn") == "ctr" and next(iter(p)) is ppr
        end = p.find("a:endParaRPr", NS)
        assert end is not None and end.get("sz") == "2700"
    assert title.text_frame.text.replace("\v", "\n") == "Новый заголовок\nвторая строка"
    objs = _objects(result)
    title_obj = objs[title_id]
    assert title_obj["content_source"] == "user"
    assert [o["op"] for o in title_obj["user_overrides"]] == ["text", "style"]
    style = title_obj["text"]["computed_style"]["font"]
    assert style["size_pt"] == 27 and style["italic"] is True and style["color"] == "#FF3985"
    assert style["family"] == "Comic Sans MS" and style["bold"] is False
    assert title_obj["text"]["paragraphs"][0]["align"] == "center"
    body_obj = objs[body_id]
    assert body_obj["content_source"] == "plan", "стиль без замены текста источник не меняет"
    assert body_obj["text"]["computed_style"]["font"]["size_pt"] == 14
    codes = [w["code"] for w in result.warnings]
    assert codes.count("override_off_template") == 1, (
        "черный цвет и 14 пт — из шаблона, 27 пт и Comic Sans — нет"
    )
    outside = next(w for w in result.warnings if w["code"] == "override_off_template")
    assert "Comic Sans" in outside["message"] and "27" in outside["message"]
    deck_slide = result.deck["slides"][0]
    assert deck_slide["overrides"] == plan["slides"][0]["overrides"]
    assert deck_slide["overrides_dropped"] == []
    assert title_obj["geometry"] == "rect"


# ---------- геометрия и группы ----------


def test_geometry_in_group_uses_inverse_transform(
    rich_profile: dict[str, Any],
    rich_template_path: pathlib.Path,
    example_package: dict[str, Any],
    tmp_path: Any,
) -> None:
    cards = _cards(rich_profile)
    icon_id = _ref(cards, "icon_1")
    label_id = _ref(cards, "label_2")
    icon_slot = next(s for s in cards["slots"] if s["slot_id"] == "icon_1")
    assert icon_slot.get("group_path"), "иконка образца лежит в группе"
    wanted_icon = {"x": 0.2, "y": 0.3, "width": 0.05, "height": 0.08}
    wanted_label = {"x": 0.4, "y": 0.5, "width": 0.25, "height": 0.1}
    plan = _cards_plan(
        rich_profile,
        [
            {"op": "geometry", "target": {"object_id": icon_id}, "geometry": {"bbox": wanted_icon}},
            {
                "op": "geometry",
                "target": {"object_id": label_id},
                "geometry": {"bbox": wanted_label},
            },
        ],
    )
    result = _compose(plan, rich_profile, rich_template_path, example_package, tmp_path / "g.pptx")
    objs = _objects(result)
    for object_id, wanted in ((icon_id, wanted_icon), (label_id, wanted_label)):
        got = objs[object_id]["bbox"]
        for key, value in wanted.items():
            assert abs(got[key] - value) < 0.002, (object_id, key, got, wanted)
        assert objs[object_id]["user_overrides"][0]["op"] == "geometry"
    assert objs[icon_id]["group_path"], "объект остался в группе"
    assert objs[icon_id]["content_source"] != "user", "перемещение содержимое не меняет"


# ---------- картинки, иконки и фон ----------


def test_picture_from_template_package_and_file(
    rich_profile: dict[str, Any],
    rich_template_path: pathlib.Path,
    example_package: dict[str, Any],
    tmp_path: Any,
) -> None:
    package, _ = _package(example_package)
    cards = _cards(rich_profile)
    icon_ids = [_ref(cards, f"icon_{i}") for i in (1, 2, 3)]
    template_icon = next(a for a in rich_profile["assets"] if a["kind"] == "icon")
    photo = next(a for a in package["assets"] if a["width_px"] != a["height_px"])
    own = tmp_path / "own.png"
    own.write_bytes(_png((10, 200, 30), 64))
    plan = _cards_plan(
        rich_profile,
        [
            {
                "op": "picture",
                "target": {"object_id": icon_ids[0], "slot_id": "icon_1"},
                "picture": {
                    "source": {"kind": "template", "asset_id": template_icon["asset_id"]},
                    "color": "#FF3985",
                },
            },
            {
                "op": "picture",
                "target": {"object_id": icon_ids[1]},
                "picture": {
                    "source": {"kind": "package", "asset_id": photo["asset_id"]},
                    "fit": "cover",
                    "color": "#123456",
                },
            },
            {
                "op": "picture",
                "target": {"object_id": icon_ids[2]},
                "picture": {
                    "source": {"kind": "file", "file_id": "file_own", "name": "own.png"},
                    "fit": "contain",
                },
            },
        ],
    )
    result = _compose(
        plan,
        rich_profile,
        rich_template_path,
        example_package,
        tmp_path / "p.pptx",
        extra_assets={"file_own": own},
    )
    objs = _objects(result)
    first = objs[icon_ids[0]]
    assert first["content_source"] == "user" and first["picture"]["recolored"] is True
    assert first["picture"]["asset_id"] == template_icon["asset_id"] + ":recolored"
    assert first["picture"]["origin"] == "template" and first["picture"]["fit"] == "contain"
    from PIL import Image

    with Image.open(io.BytesIO(_shape(result, icon_ids[0]).image.blob)) as img:
        center = img.convert("RGBA").getpixel((img.width // 2, img.height // 2))
        assert center[:3] == (0xFF, 0x39, 0x85)
    second = objs[icon_ids[1]]
    assert (
        second["picture"]["asset_id"] == photo["asset_id"]
        and second["picture"]["origin"] == "content"
    )
    assert second["picture"].get("crop"), "cover: фото 16:9 в квадратной рамке обрезано"
    assert second["picture"]["recolored"] is False
    third = objs[icon_ids[2]]
    assert third["content_source"] == "user" and third["picture"]["origin"] == "content"
    assert third["picture"]["fit"] == "contain" and third["picture"]["asset_id"].startswith(
        "media_"
    )
    with Image.open(io.BytesIO(_shape(result, icon_ids[2]).image.blob)) as img:
        assert img.convert("RGB").getpixel((5, 5)) == (10, 200, 30)
    codes = [w["code"] for w in result.warnings]
    assert codes.count("override_recolor_unsupported") == 1, "фото не маска: перекраска пропущена"
    assert "override_asset_missing" not in codes and "override_file_missing" not in codes
    assert result.deck["slides"][0]["overrides_dropped"] == []


def test_background_solid_image_inherited(
    rich_profile: dict[str, Any],
    rich_template_path: pathlib.Path,
    example_package: dict[str, Any],
    tmp_path: Any,
) -> None:
    package, _ = _package(example_package)
    cards = _cards(rich_profile)
    photo = next(a for a in package["assets"] if a["width_px"] != a["height_px"])
    logo = next(a for a in rich_profile["assets"] if a["kind"] == "logo")
    slide = {
        "pattern_id": cards["pattern_id"],
        "title": "Фон",
        "blocks": [{"slot_id": "title_1", "kind": "title", "text": "Фон"}],
    }
    plan = _plan_with(
        rich_profile,
        [
            {
                **slide,
                "overrides": [
                    {"op": "background", "background": {"kind": "solid", "color": "#F5F7FA"}}
                ],
            },
            {
                **slide,
                "overrides": [
                    {
                        "op": "background",
                        "background": {
                            "kind": "image",
                            "source": {"kind": "package", "asset_id": photo["asset_id"]},
                            "fit": "cover",
                        },
                    }
                ],
            },
            {
                **slide,
                "overrides": [
                    {
                        "op": "background",
                        "background": {
                            "kind": "image",
                            "source": {"kind": "template", "asset_id": logo["asset_id"]},
                            "fit": "contain",
                        },
                    }
                ],
            },
            {**slide, "overrides": [{"op": "background", "background": {"kind": "inherited"}}]},
        ],
    )
    result = _compose(plan, rich_profile, rich_template_path, example_package, tmp_path / "b.pptx")
    kinds = [s["background"] for s in result.deck["slides"]]
    assert kinds[0] == {"kind": "solid", "color": "#F5F7FA"}
    assert kinds[1]["kind"] == "image" and kinds[1]["asset_id"] == photo["asset_id"]
    assert kinds[2]["kind"] == "image" and kinds[2]["asset_id"] == logo["asset_id"]
    assert kinds[3] == {"kind": "inherited"}
    asset_ids = {a["asset_id"] for a in result.deck["assets"]}
    assert photo["asset_id"] in asset_ids and logo["asset_id"] in asset_ids
    prs = Presentation(str(result.pptx_path))
    slides = list(prs.slides)
    bg = slides[0]._element.find("p:cSld/p:bg", NS)
    assert bg is not None and next(iter(slides[0]._element.find("p:cSld", NS))) is bg, (
        "p:bg первым в cSld"
    )
    assert bg.find(".//a:srgbClr", NS).get("val") == "F5F7FA"
    blip_fill = slides[1]._element.find("p:cSld/p:bg//a:blipFill", NS)
    assert blip_fill is not None and blip_fill.find("a:srcRect", NS) is not None, "cover: обрезка"
    fill_rect = slides[2]._element.find("p:cSld/p:bg//a:stretch/a:fillRect", NS)
    assert fill_rect is not None and (fill_rect.get("l") or fill_rect.get("t")), "contain: поля"
    assert slides[3]._element.find("p:cSld/p:bg", NS) is None
    assert not [w for w in result.warnings if w["code"].startswith("override_")]


# ---------- охрана адреса и отказ ----------


def test_guard_mismatch_and_missing_target_dropped(
    rich_profile: dict[str, Any],
    rich_template_path: pathlib.Path,
    example_package: dict[str, Any],
    tmp_path: Any,
) -> None:
    cards = _cards(rich_profile)
    title_id = _ref(cards, "title_1")
    icon_id = _ref(cards, "icon_1")
    plan = _cards_plan(
        rich_profile,
        [
            {"op": "text", "target": {"object_id": "999"}, "text": "нет такого"},
            {
                "op": "text",
                "target": {"object_id": title_id, "source_object_id": "1"},
                "text": "чужой образец",
            },
            {
                "op": "text",
                "target": {"object_id": title_id, "slot_id": "body_1"},
                "text": "чужой слот",
            },
            {"op": "text", "target": {"object_id": icon_id}, "text": "текст в картинку"},
            {"op": "style", "target": {"object_id": icon_id}, "style": {"font": {"bold": True}}},
            {
                "op": "picture",
                "target": {"object_id": title_id},
                "picture": {"source": {"kind": "template", "asset_id": "asset_2"}},
            },
            {
                "op": "picture",
                "target": {"object_id": icon_id},
                "picture": {"source": {"kind": "template", "asset_id": "asset_nope"}},
            },
            {
                "op": "picture",
                "target": {"object_id": icon_id},
                "picture": {"source": {"kind": "file", "file_id": "file_nope"}},
            },
            {
                "op": "text",
                "target": {
                    "object_id": title_id,
                    "source_object_id": title_id,
                    "slot_id": "title_1",
                },
                "text": "Правильный адрес",
            },
        ],
    )
    result = _compose(plan, rich_profile, rich_template_path, example_package, tmp_path / "d.pptx")
    dropped = result.deck["slides"][0]["overrides_dropped"]
    codes = [d["code"] for d in dropped]
    assert codes == [
        "override_target_missing",
        "override_guard_mismatch",
        "override_guard_mismatch",
        "override_unsupported",
        "override_unsupported",
        "override_unsupported",
        "override_asset_missing",
        "override_file_missing",
    ]
    assert dropped[0]["target"] == {"object_id": "999"}
    assert _shape(result, title_id).text_frame.text == "Правильный адрес"
    objs = _objects(result)
    assert [o["op"] for o in objs[title_id]["user_overrides"]] == ["text"]
    assert "user_overrides" not in objs[icon_id]
    warned = [w["code"] for w in result.warnings]
    assert all(code in warned for code in set(codes))


# ---------- переполнение и медиа ----------


def test_overflow_warning_and_media_exported(
    rich_profile: dict[str, Any],
    rich_template_path: pathlib.Path,
    example_package: dict[str, Any],
    tmp_path: Any,
) -> None:
    cards = _cards(rich_profile)
    label_id = _ref(cards, "label_1")
    plan = _cards_plan(
        rich_profile,
        [
            {
                "op": "text",
                "target": {"object_id": label_id},
                "text": "Очень длинная подпись карточки, которая точно не поместится в узкую рамку "
                * 4,
            },
            {"op": "style", "target": {"object_id": label_id}, "style": {"font": {"size_pt": 54}}},
        ],
    )
    media_dir = tmp_path / "media"
    result = _compose(
        plan,
        rich_profile,
        rich_template_path,
        example_package,
        tmp_path / "o.pptx",
        media_dir=media_dir,
        media_prefix="balanced/r2/",
    )
    overflow = [w for w in result.warnings if w["code"] == "override_overflow"]
    assert overflow and label_id in overflow[0]["message"]
    assets = result.deck["assets"]
    assert assets, "логотип образца — ресурс колоды"
    for asset in assets:
        assert asset["artifact"].startswith("balanced/r2/media/")
        assert (media_dir / asset["artifact"].split("media/", 1)[1]).is_file()
    assert result.report["media_files"] == len(assets)


def test_object_ids_stable_across_recompose(
    rich_profile: dict[str, Any],
    rich_template_path: pathlib.Path,
    example_package: dict[str, Any],
    tmp_path: Any,
) -> None:
    package, _ = _package(example_package)
    cards = _cards(rich_profile)
    photo = next(a for a in package["assets"] if a["width_px"] != a["height_px"])
    slides = [
        {
            "pattern_id": cards["pattern_id"],
            "title": "Первый",
            "blocks": [
                {"slot_id": "title_1", "kind": "title", "text": "Первый"},
                {
                    "slot_id": "icon_1",
                    "kind": "image",
                    "image": {"asset_id": photo["asset_id"], "fit": "cover"},
                },
                {"slot_id": "label_1", "kind": "label", "text": "Подпись"},
                {"slot_id": "body_1", "kind": "body", "text": "Текст"},
            ],
        },
        {
            "pattern_id": cards["pattern_id"],
            "title": "Второй",
            "blocks": [
                {"slot_id": "title_1", "kind": "title", "text": "Второй"},
                {
                    "slot_id": "body_1",
                    "kind": "diagram",
                    "diagram": {"kind": "process", "items": [{"text": "Раз"}, {"text": "Два"}]},
                },
            ],
        },
    ]
    plan = _plan_with(rich_profile, slides)
    first = _compose(plan, rich_profile, rich_template_path, example_package, tmp_path / "a.pptx")
    second = _compose(plan, rich_profile, rich_template_path, example_package, tmp_path / "b.pptx")
    ids_first = [[o["object_id"] for o in s["objects"]] for s in first.deck["slides"]]
    ids_second = [[o["object_id"] for o in s["objects"]] for s in second.deck["slides"]]
    assert ids_first == ids_second
    assert any(o["content_source"] == "generated" for o in first.deck["slides"][1]["objects"])
