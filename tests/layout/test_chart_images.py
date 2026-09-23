"""Диаграммы-картинки → нативные диаграммы: замена на месте по чтению с картинки (тип, ряды,
цвета, сглаживание, вторая ось, кольцо, полосы без шкалы), имена рядов из легенды слайда,
отчёт для чата и сборка варианта original, которая повторяет замены по готовым чтениям."""

from __future__ import annotations

import io
import pathlib
from copy import deepcopy
from types import SimpleNamespace
from typing import Any

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from presentation_designer.generation.original import original_plan, original_story
from presentation_designer.layout.chart_images import (
    ChartSwap,
    charts_report,
    deck_pictures,
    picture_sha,
    report_readings,
    swap_pictures,
    swap_summary,
    unread_charts,
)
from presentation_designer.layout.compose import compose_deck
from presentation_designer.parsing.content.importer import import_content
from presentation_designer.parsing.content.parsers import ParseCache
from presentation_designer.parsing.raster_charts.model import (
    AxisInfo,
    ChartReading,
    ChartStructure,
    SeriesInfo,
)
from presentation_designer.parsing.raster_charts.pixels import load
from presentation_designer.parsing.raster_charts.rebuild import read_with_structure
from presentation_designer.shared.settings import Settings
from tests.layout.conftest import MINI_TEMPLATE, own_profile
from tests.parsing.content.conftest import material
from tests.parsing.raster_charts import synth

CATS = ["Ряд 1", "Ряд 2", "Ряд 3"]
DATES = ["1/1/19", "2/1/19", "3/1/19", "4/1/19", "5/1/19"]
BOX = (Inches(1), Inches(1.5), Inches(8), Inches(4.5))


def _read(data: bytes, st: ChartStructure) -> ChartReading:
    outcome = read_with_structure(load(data), st)
    assert outcome.status == "read", outcome.reason
    assert outcome.reading is not None
    return outcome.reading


def _deck(data: bytes) -> tuple[Any, Any]:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(io.BytesIO(data), *BOX)
    return prs, slide


def _lines() -> tuple[bytes, ChartReading]:
    data, ticks = synth.lines(
        [[1.5, 13.6, 9.5, 18.8, 27.3], [0.5, 16.7, 8.5, 24.6, 18.8]],
        [synth.BLUE, synth.PINK],
        DATES,
    )
    st = ChartStructure(
        status="chart",
        kind="line",
        categories=DATES,
        series=[
            SeriesInfo(color="#1A73E8", type="line", smooth=True),
            SeriesInfo(color="#E8457A", type="line", smooth=True),
        ],
        primary_axis=AxisInfo(ticks=ticks),
    )
    return data, _read(data, st)


def _chart_frames(slide: Any) -> list[Any]:
    return [s for s in slide.shapes if s.has_chart]


def _values(chart: Any) -> list[list[float]]:
    return [list(s.values) for plot in chart.plots for s in plot.series]


def test_line_picture_becomes_a_smooth_native_chart_in_the_same_place(
    tmp_path: pathlib.Path,
) -> None:
    data, reading = _lines()
    prs, slide = _deck(data)
    sha = picture_sha(slide.shapes[0])
    swaps = swap_pictures(prs, {sha: reading})
    assert [(s.slide, s.status) for s in swaps] == [(1, "replaced")]
    assert swaps[0].measured == 10
    assert not [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    # Картинки больше нет — связь слайда с ней удалена, в файле она не остаётся.
    assert not [r for r in slide.part.rels.values() if r.reltype.endswith("/image")]
    (frame,) = _chart_frames(slide)
    assert (frame.left, frame.top, frame.width, frame.height) == BOX
    chart = frame.chart
    assert list(chart.plots[0].categories) == DATES
    for got, want in zip(_values(chart), [s.points for s in reading.series], strict=True):
        assert got == pytest.approx([p.value for p in want])
    xml = chart._chartSpace.xml
    # Сглаживание и прочие логические флаги записаны явно: ONLYOFFICE пустой элемент
    # читает как «нет» — линия ломается, скрытая ось появляется.
    assert '<c:smooth val="1"/>' in xml and "<c:smooth/>" not in xml
    assert "<c:delete/>" not in xml and "<c:varyColors/>" not in xml
    descr = frame._element.nvGraphicFramePr.cNvPr.get("descr")
    assert "восстановлена по картинке" in descr and "10 значений" in descr
    # Файл открывается снова и диаграмма читается.
    out = tmp_path / "deck.pptx"
    prs.save(str(out))
    again = Presentation(str(out)).slides[0]
    assert _values(_chart_frames(again)[0].chart)[0] == pytest.approx(
        [p.value for p in reading.series[0].points]
    )


def test_combo_puts_the_line_on_its_own_right_axis() -> None:
    cols, line = [[26, 43, 20], [53, 70, 18]], [800, 2400, 1500]
    data, ticks, right = synth.columns(
        cols, [synth.BLUE, synth.CYAN], CATS, line=line, line_hi=3000, line_step=500
    )
    st = ChartStructure(
        status="chart",
        kind="combo",
        categories=CATS,
        series=[
            SeriesInfo(name="A", color="#1E88E5", type="column"),
            SeriesInfo(name="B", color="#80E8F0", type="column"),
            SeriesInfo(name="C", color="#E0508A", type="line", axis="secondary"),
        ],
        primary_axis=AxisInfo(ticks=ticks),
        secondary_axis=AxisInfo(ticks=right),
    )
    prs, slide = _deck(data)
    swap_pictures(prs, {picture_sha(slide.shapes[0]): _read(data, st)})
    plot_area = _chart_frames(slide)[0].chart._chartSpace.chart.plotArea
    bar = plot_area.find(qn("c:barChart"))
    lines = plot_area.find(qn("c:lineChart"))
    assert len(bar.findall(qn("c:ser"))) == 2 and len(lines.findall(qn("c:ser"))) == 1
    vals = plot_area.findall(qn("c:valAx"))
    assert len(vals) == 2
    primary, secondary = vals
    line_axes = {el.get("val") for el in lines.findall(qn("c:axId"))}
    assert secondary.find(qn("c:axId")).get("val") in line_axes
    assert secondary.find(qn("c:axPos")).get("val") == "r"
    assert float(secondary.find(qn("c:scaling")).find(qn("c:max")).get("val")) == 3000
    assert float(primary.find(qn("c:scaling")).find(qn("c:max")).get("val")) == 100
    # Основная ось — без сетки: у картинки её не было.
    assert primary.find(qn("c:majorGridlines")) is None


def test_doughnut_keeps_hole_first_angle_and_slice_colors() -> None:
    colors = [synth.BLUE, synth.PINK, synth.CYAN]
    data = synth.doughnut([22, 50, 28], colors, hole=0.72)
    st = ChartStructure(
        status="chart",
        kind="doughnut",
        categories=["А", "Б", "В"],
        category_colors=["#1E88E5", "#E0508A", "#80E8F0"],
    )
    reading = _read(data, st)
    prs, slide = _deck(data)
    swap_pictures(prs, {picture_sha(slide.shapes[0]): reading})
    chart = _chart_frames(slide)[0].chart
    group = chart._chartSpace.chart.plotArea.find(qn("c:doughnutChart"))
    assert 68 <= int(group.find(qn("c:holeSize")).get("val")) <= 76
    assert group.find(qn("c:firstSliceAng")).get("val") == str(round(reading.first_angle or 0))
    fills = [str(p.format.fill.fore_color.rgb) for p in chart.plots[0].series[0].points]
    assert [f"#{f}" for f in fills] == reading.slice_colors


def test_bars_without_a_scale_hide_the_axis_and_keep_the_top_down_order() -> None:
    values = [79, 63, 43, 38, 10]
    names = ["Первый", "Второй", "Третий", "Четвёртый", "Пятый"]
    st = ChartStructure(
        status="chart",
        kind="bar",
        categories=names,
        series=[SeriesInfo(color="#1E88E5", type="bar", labels=["79", "63", "43", None, "10"])],
        primary_axis=AxisInfo(ticks=[]),
    )
    data = synth.hbars(values, names)
    reading = _read(data, st)
    assert reading.geometry.value_max and reading.geometry.value_max >= 79
    prs, slide = _deck(data)
    swap_pictures(prs, {picture_sha(slide.shapes[0]): reading})
    chart = _chart_frames(slide)[0].chart
    plot_area = chart._chartSpace.chart.plotArea
    cat, val = plot_area.find(qn("c:catAx")), plot_area.find(qn("c:valAx"))
    assert cat.find(qn("c:scaling")).find(qn("c:orientation")).get("val") == "maxMin"
    assert val.find(qn("c:delete")).get("val") == "1"
    assert val.find(qn("c:crosses")).get("val") == "max"
    assert chart.plots[0].has_data_labels


def test_highlighted_columns_keep_their_colors_in_the_native_chart() -> None:
    names = ["VK", "Конкурент 1", "Конкурент 2"]
    st = ChartStructure(
        status="chart",
        kind="column",
        categories=names,
        series=[SeriesInfo(color="#007BFF", type="column", labels=["43%", "23%", "18%"])],
        primary_axis=AxisInfo(ticks=[]),
    )
    data = synth.highlighted([43, 23, 18], names)
    reading = _read(data, st)
    prs, slide = _deck(data)
    swap_pictures(prs, {picture_sha(slide.shapes[0]): reading})
    series = _chart_frames(slide)[0].chart.plots[0].series[0]
    fills = [
        str(pt.find(qn("c:spPr")).find(qn("a:solidFill")).find(qn("a:srgbClr")).get("val"))
        for pt in series._element.findall(qn("c:dPt"))
    ]
    light = "{:02X}{:02X}{:02X}".format(*synth.LIGHT)
    assert fills == ["0077FF", light, light]


def test_generic_series_names_come_from_the_legend_drawn_on_the_slide() -> None:
    data, reading = _lines()
    assert [s.name for s in reading.series] == ["Ряд 1", "Ряд 2"]
    prs, slide = _deck(data)
    for k, (name, color) in enumerate((("DAU", "1A73E8"), ("MAU", "E8457A"))):
        x = Inches(1 + 1.5 * k)
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, Inches(1), Inches(0.2), Inches(0.2))
        dot.fill.solid()
        dot.fill.fore_color.rgb = RGBColor.from_string(color)
        label = slide.shapes.add_textbox(x + Inches(0.3), Inches(0.95), Inches(1), Inches(0.3))
        label.text_frame.text = name
        label.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
    picture = next(s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE)
    swap_pictures(prs, {picture_sha(picture): reading})
    chart = _chart_frames(slide)[0].chart
    assert [s.name for s in chart.plots[0].series] == ["DAU", "MAU"]


def test_legend_in_groups_with_theme_colors_names_the_series() -> None:
    """Как на слайде 48 VK Education: каждая пара «кружок — надпись» — своя группа с теми же
    координатами детей, кружки залиты цветами темы."""
    data, reading = _lines()
    theme = {"accent1": "#4F81BD", "accent2": "#C0504D"}  # тема шаблона python-pptx
    reading = reading.model_copy(
        update={
            "series": [
                s.model_copy(update={"color": c})
                for s, c in zip(reading.series, theme.values(), strict=True)
            ]
        }
    )
    prs, slide = _deck(data)
    for k, (name, accent) in enumerate(
        (("DAU", MSO_THEME_COLOR.ACCENT_1), ("MAU", MSO_THEME_COLOR.ACCENT_2))
    ):
        group = slide.shapes.add_group_shape()
        dot = group.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1), Inches(1), Inches(0.2), Inches(0.2))
        dot.fill.solid()
        dot.fill.fore_color.theme_color = accent
        label = group.shapes.add_textbox(Inches(1.3), Inches(0.95), Inches(1), Inches(0.3))
        label.text_frame.text = name
        group.left = Inches(1 + 2 * k)  # сдвигается рамка группы, координаты детей те же
    picture = next(s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE)
    swap_pictures(prs, {picture_sha(picture): reading})
    chart = _chart_frames(slide)[0].chart
    assert [s.name for s in chart.plots[0].series] == ["DAU", "MAU"]


def test_cropped_picture_stays_a_picture() -> None:
    data, reading = _lines()
    prs, slide = _deck(data)
    picture = slide.shapes[0]
    picture.crop_left = 0.2
    swaps = swap_pictures(prs, {picture_sha(picture): reading})
    assert [(s.status, s.reason) for s in swaps] == [("kept", "картинка обрезана на слайде")]
    assert not _chart_frames(slide)


def test_report_for_the_chat_names_slides_and_reasons() -> None:
    swaps = [
        ChartSwap(3, "a", "a" * 64, "replaced", measured=12),
        ChartSwap(5, "b", "b" * 64, "replaced"),
        ChartSwap(6, "c", "c" * 64, "kept", "картинка обрезана на слайде"),
    ]
    outcomes = {
        "d": SimpleNamespace(status="rejected", reason="модель не ответила: timeout"),
        "e": SimpleNamespace(status="skipped", reason="время на чтение диаграмм вышло"),
        "f": SimpleNamespace(status="skipped", reason="мало фона — не похоже на диаграмму"),
        "g": SimpleNamespace(status="not_chart", reason="фотография"),
    }
    unread = unread_charts(outcomes, {"d": [7], "e": [9, 11], "f": [2], "g": [4]})
    assert unread == [(7, "модель не ответила"), (9, "не успел прочитать")]
    text = swap_summary(swaps, unread)
    assert text == (
        "2 диаграммы-картинки стали редактируемыми (слайды 3, 5): значения сняты по "
        "картинке, проверьте их. 3 остались картинкой: слайд 6 — картинка обрезана на "
        "слайде; слайд 7 — модель не ответила; слайд 9 — не успел прочитать."
    )
    assert swap_summary([], []) is None
    one = swap_summary(swaps[:1], [])
    assert one is not None and one.startswith("1 диаграмма-картинка стала редактируемой (слайд 3)")


def test_report_keeps_readings_for_the_full_build() -> None:
    data, reading = _lines()
    prs, _ = _deck(data)
    blobs, where = deck_pictures(prs)
    (sha,) = blobs
    outcome = SimpleNamespace(status="read", reason="", reading=reading)
    report = charts_report({sha: outcome}, where, [], None)
    assert report["pictures"][0]["slides"] == [1]
    assert report_readings(report) == {sha: reading}


def _mini_with_chart(tmp_path: pathlib.Path) -> pathlib.Path:
    data, _ = _lines()
    prs = Presentation(str(MINI_TEMPLATE))
    slide = prs.slides[len(prs.slides) - 1]
    slide.shapes.add_picture(io.BytesIO(data), Inches(1), Inches(4.2), Inches(5), Inches(2.5))
    path = tmp_path / "mini_chart.pptx"
    prs.save(str(path))
    return path


def test_original_build_swaps_charts_only_with_readings(
    tmp_path: pathlib.Path, import_settings: Settings
) -> None:
    source = _mini_with_chart(tmp_path)
    profile = own_profile(source, "tpl_chart")
    package = dict(
        import_content(
            [material(source)],
            {"purpose": "other", "title": "Мини с диаграммой"},
            package_id="pkg_chart",
            settings=import_settings,
            cache=ParseCache(tmp_path / "import-cache"),
            use_model=False,
        ).package
    )
    story = original_story(package, profile, {"language": "ru"})
    plan = original_plan(story, profile, package, plan_id="plan_chart")
    _, reading = _lines()
    sha = picture_sha(
        next(
            s
            for s in Presentation(str(source)).slides[-1].shapes
            if s.shape_type == MSO_SHAPE_TYPE.PICTURE
        )
    )

    plain = compose_deck(plan, profile, source, package, out_pptx=tmp_path / "plain.pptx")
    assert not plain.chart_swaps
    # Без чтений файл не тронут: копия исходника байт в байт.
    assert (tmp_path / "plain.pptx").read_bytes() == source.read_bytes()

    swapped = compose_deck(
        deepcopy(plan),
        profile,
        source,
        package,
        out_pptx=tmp_path / "swapped.pptx",
        chart_readings={sha: reading.model_dump()},
    )
    assert [s.status for s in swapped.chart_swaps] == ["replaced"]
    last = Presentation(str(tmp_path / "swapped.pptx")).slides[-1]
    assert _chart_frames(last)
    assert not [s for s in last.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    # Описание колоды видит диаграмму на месте картинки.
    kinds = [o["kind"] for o in swapped.deck["slides"][-1]["objects"]]
    assert "chart" in kinds
