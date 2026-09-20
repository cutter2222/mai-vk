"""Сборка на четырёх целевых шаблонах (закрытые материалы, маркер organizer_data): профили
анализатором без рендера и VLM, планы трёх вариантов без модели, контент-пакет и смысловой
план — собственные. Проверяются свойства, не зависящие от формулировок: PPTX открывается
python-pptx, число и порядок слайдов по плану, пакет проходит проверку целостности без
недостижимых частей, факты подставлены, таблицы и диаграммы плана стали нативными объектами,
ComposedDeck сверяется с файлом по идентификаторам объектов и частям, макеты и мастера
сохранены. Отчёт по каждому шаблону — в runs/compose-tests/<имя>.json.
"""

from __future__ import annotations

import json
import pathlib
import time
import zipfile
from typing import Any

import pytest
from pptx import Presentation

from presentation_designer.contracts import ComposedDeck
from presentation_designer.generation.variants import VARIANTS
from presentation_designer.layout.compose import compose_deck
from presentation_designer.layout.shapes import iter_shapes
from presentation_designer.layout.text import FACT_REF
from presentation_designer.parsing.template.analyzer import analyze_template
from presentation_designer.shared.settings import ROOT

pytestmark = pytest.mark.organizer_data

REPORT_DIR = ROOT / "runs" / "compose-tests"
TEMPLATES = (
    "VK Tech шаблон.pptx",
    "VK_WorkSpace_Клиентская_конференция_Шаблон_03.pptx",
    "Шаблон презентации VK Education.pptx",
    "ЛЦТ2026 Шаблон презентации.pptx",
)


@pytest.mark.parametrize("name", TEMPLATES, ids=lambda n: n.split(".")[0][:24])
def test_compose_three_variants_on_organizer_template(
    organizer_dir: pathlib.Path,
    name: str,
    example_package: dict[str, Any],
    make_plan: Any,
    tmp_path: pathlib.Path,
) -> None:
    path = organizer_dir / name
    if not path.exists():
        pytest.fail(f"нет файла {path}: прогон организаторов не должен пропускаться молча")
    profile = analyze_template(
        path,
        template_id="tpl_org",
        name=name,
        size_bytes=path.stat().st_size,
        render=False,
        use_vlm=False,
    ).profile
    package = {k: v for k, v in example_package.items() if not k.startswith("_")}
    assets = pathlib.Path(example_package["_assets_dir"])
    with zipfile.ZipFile(path) as zf:
        source_layouts = {n for n in zf.namelist() if n.startswith("ppt/slideLayouts/slideLayout")}
        source_masters = {n for n in zf.namelist() if n.startswith("ppt/slideMasters/slideMaster")}
    report: dict[str, Any] = {"template": name, "variants": {}}
    for variant_id in VARIANTS:
        plan = make_plan(profile, variant_id)
        started = time.perf_counter()
        result = compose_deck(
            plan,
            profile,
            path,
            package,
            out_pptx=tmp_path / f"{variant_id}.pptx",
            package_dir=assets,
            job_id="job_org",
            variant_id=variant_id,
        )
        elapsed = int((time.perf_counter() - started) * 1000)
        deck = result.deck
        ComposedDeck.model_validate(deck)
        assert result.integrity.ok, result.integrity.errors
        assert result.integrity.unreachable_parts == []
        prs = Presentation(str(result.pptx_path))
        assert len(prs.slides) == len(plan["slides"]) == deck["stats"]["slides"]
        planned_tables = sum(1 for s in plan["slides"] for b in s["blocks"] if b["kind"] == "table")
        planned_charts = sum(1 for s in plan["slides"] for b in s["blocks"] if b["kind"] == "chart")
        tables = charts = 0
        for slide, plan_slide, deck_slide in zip(
            prs.slides, plan["slides"], deck["slides"], strict=True
        ):
            assert deck_slide["slide_id"] == plan_slide["slide_id"]
            shapes = list(iter_shapes(slide))
            by_id = {str(s.shape_id): s for s in shapes}
            texts = "\n".join(s.text_frame.text for s in shapes if s.has_text_frame)
            assert FACT_REF.search(texts) is None
            # Заголовок плана присутствует на слайде (текст слота title).
            title_obj = next(
                (o for o in deck_slide["objects"] if o.get("slot_kind") == "title"), None
            )
            if title_obj is not None:
                assert title_obj["text"]["plain"] == plan_slide["title"]
            for obj in deck_slide["objects"]:
                assert obj["object_id"] in by_id, (name, deck_slide["slide_id"], obj["object_id"])
            for removed in deck_slide["removed_object_ids"]:
                assert removed not in by_id
            tables += sum(1 for s in shapes if getattr(s, "has_table", False) and s.has_table)
            charts += sum(1 for s in shapes if getattr(s, "has_chart", False) and s.has_chart)
        assert tables >= planned_tables and charts >= planned_charts
        with zipfile.ZipFile(result.pptx_path) as zf:
            names = set(zf.namelist())
            # Макеты и мастера исходного шаблона сохранены (макеты не удаляются по умолчанию).
            assert source_layouts <= names and source_masters <= names
            for deck_slide in deck["slides"]:
                for obj in deck_slide["objects"]:
                    if obj["kind"] == "chart":
                        assert obj["chart"]["chart_part"] in names
                        assert obj["chart"].get("built") in ("replaced", "rebuilt", "added")
        codes = sorted({w["code"] for w in result.warnings})
        report["variants"][variant_id] = {
            "slides": len(plan["slides"]),
            "objects": deck["stats"]["objects"],
            "removed_objects": deck["stats"]["removed_objects"],
            "tables": tables,
            "charts": charts,
            "file_size_kb": deck["stats"]["file_size_bytes"] // 1024,
            "compose_ms": elapsed,
            "timings_ms": result.report["timings_ms"],
            "warnings": codes,
        }
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    (REPORT_DIR / f"{name.split('.')[0][:40]}.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=1), encoding="utf-8"
    )


def test_vkedu_overrides_deterministic(
    organizer_dir: pathlib.Path,
    example_package: dict[str, Any],
    make_plan: Any,
    tmp_path: pathlib.Path,
) -> None:
    """Ручные правки редактора на настоящем шаблоне: id объектов совпадают между двумя
    сборками одного плана, правки текста, стиля, положения и фона на трёх слайдах
    применяются без отброшенных, PPTX проходит проверку целостности."""
    name = "Шаблон презентации VK Education.pptx"
    path = organizer_dir / name
    if not path.exists():
        pytest.fail(f"нет файла {path}: прогон организаторов не должен пропускаться молча")
    profile = analyze_template(
        path,
        template_id="tpl_org",
        name=name,
        size_bytes=path.stat().st_size,
        render=False,
        use_vlm=False,
    ).profile
    package = {k: v for k, v in example_package.items() if not k.startswith("_")}
    assets = pathlib.Path(example_package["_assets_dir"])
    plan = make_plan(profile, "balanced")
    base = compose_deck(
        plan, profile, path, package, out_pptx=tmp_path / "base.pptx", package_dir=assets
    )
    again = compose_deck(
        plan, profile, path, package, out_pptx=tmp_path / "again.pptx", package_dir=assets
    )
    ids = lambda deck: [[o["object_id"] for o in s["objects"]] for s in deck["slides"]]  # noqa: E731
    assert ids(base.deck) == ids(again.deck)
    # Правки по объектам ComposedDeck первой сборки: заголовки трёх содержательных слайдов.
    edited: list[str] = []
    for deck_slide, plan_slide in zip(
        base.deck["slides"], sorted(plan["slides"], key=lambda s: s["order"]), strict=True
    ):
        title = next(
            (o for o in deck_slide["objects"] if o.get("slot_kind") == "title" and o.get("text")),
            None,
        )
        if title is None or len(edited) >= 3:
            continue
        picture = next((o for o in deck_slide["objects"] if o["kind"] == "picture"), None)
        overrides = [
            {
                "op": "text",
                "target": {
                    "object_id": title["object_id"],
                    "source_object_id": title.get("source_object_id"),
                    "slot_id": title["slot_id"],
                },
                "text": "Правка редактора",
            },
            {
                "op": "style",
                "target": {"object_id": title["object_id"]},
                "style": {"font": {"bold": True, "color": "#0077FF"}, "align": "left"},
            },
            {"op": "background", "background": {"kind": "solid", "color": "#F5F7FA"}},
        ]
        if picture is not None:
            box = dict(picture["bbox"])
            box["x"] = round(max(0.0, box["x"] - 0.02), 4)
            overrides.append(
                {
                    "op": "geometry",
                    "target": {"object_id": picture["object_id"]},
                    "geometry": {"bbox": box},
                }
            )
        plan_slide["overrides"] = overrides
        edited.append(plan_slide["slide_id"])
    assert len(edited) == 3
    result = compose_deck(
        plan, profile, path, package, out_pptx=tmp_path / "edited.pptx", package_dir=assets
    )
    assert result.integrity.ok
    for deck_slide in result.deck["slides"]:
        if deck_slide["slide_id"] not in edited:
            assert "overrides" not in deck_slide
            continue
        assert deck_slide["overrides_dropped"] == [], deck_slide["overrides_dropped"]
        assert deck_slide["background"] == {"kind": "solid", "color": "#F5F7FA"}
        title = next(o for o in deck_slide["objects"] if o.get("slot_kind") == "title")
        assert title["text"]["plain"] == "Правка редактора"
        assert title["content_source"] == "user"
        assert title["text"]["computed_style"]["font"]["color"] == "#0077FF"
    ComposedDeck.model_validate(result.deck)
    assert ids(result.deck) == ids(base.deck), "правки не меняют идентификаторы объектов"

def test_vktech_drawn_chart_and_logo_removal(
    organizer_dir: pathlib.Path,
    example_package: dict[str, Any],
    tmp_path: pathlib.Path,
) -> None:
    """Две ловушки шаблона VK Tech: диаграмма, нарисованная картинками, и логотип на макетах.

    Первая: слайд-образец 44 — это ряд столбиков из картинок с подписями значений. Заполнять
    его числами нельзя, высоты нарисованы заранее; на его месте строится нативная диаграмма,
    а части ряда убираются. Вторая: знак VK лежит на макетах, объекта на слайде нет, и
    правкой слайда его не снять — снимает флаг колоды `template_logo: drop`.
    """
    name = "VK Tech шаблон.pptx"
    path = organizer_dir / name
    if not path.exists():
        pytest.fail(f"нет файла {path}: прогон организаторов не должен пропускаться молча")
    profile = analyze_template(
        path,
        template_id="tpl_org",
        name=name,
        size_bytes=path.stat().st_size,
        render=False,
        use_vlm=False,
    ).profile
    package = {k: v for k, v in example_package.items() if not k.startswith("_")}
    assets = pathlib.Path(example_package["_assets_dir"])
    pattern = next(p for p in profile["patterns"] if p.get("chart_parts"))
    title_slot = next(sl["slot_id"] for sl in pattern["slots"] if sl["kind"] == "title")
    chart_slot = next(sl["slot_id"] for sl in pattern["slots"] if sl["kind"] == "chart")
    dataset = package["datasets"][0]

    def plan_with(logo: str) -> dict[str, Any]:
        return {
            "schema_version": "1.3",
            "plan_id": "plan_org",
            "template_id": profile["template_id"],
            "package_id": package["package_id"],
            "story_id": "story_org",
            "language": "ru",
            "variant": {"variant_id": "balanced", "axis": "density", "value": "balanced"},
            "slide_count": {"exact": 1},
            "template_logo": logo,
            "slides": [
                {
                    "slide_id": "s1",
                    "order": 1,
                    "pattern_id": pattern["pattern_id"],
                    "title": "Динамика",
                    "blocks": [
                        {"slot_id": title_slot, "kind": "title", "text": "Динамика"},
                        {
                            "slot_id": chart_slot,
                            "kind": "chart",
                            "chart": {
                                "dataset_id": dataset["dataset_id"],
                                "chart_type": "column",
                                "category_column": dataset["columns"][0]["name"],
                                "series_columns": [dataset["columns"][1]["name"]],
                                "legend": True,
                            },
                        },
                    ],
                }
            ],
            "coverage": {"required_thesis_ids": [], "covered": []},
            "generation_meta": {"skills": [], "models": []},
        }

    kept = compose_deck(
        plan_with("keep"),
        profile,
        path,
        package,
        out_pptx=tmp_path / "keep.pptx",
        package_dir=assets,
    )
    ComposedDeck.model_validate(kept.deck)
    slide = kept.deck["slides"][0]
    charts = [o for o in slide["objects"] if o["kind"] == "chart"]
    assert len(charts) == 1, "на месте нарисованного ряда одна нативная диаграмма"
    assert charts[0]["chart"]["built"] == "added"
    assert charts[0]["bbox"]["width"] >= 0.5, "диаграмма занимает область ряда, а не один столбик"
    assert not [o for o in slide["objects"] if o["kind"] == "picture"], "столбики-картинки убраны"
    assert len(slide["removed_object_ids"]) >= len(pattern["chart_parts"])

    logos = [f for f in profile["fixed_elements"] if f["kind"] == "logo"]
    assert logos, "у шаблона VK Tech знак на макетах"

    def logo_shapes(pptx: pathlib.Path) -> int:
        prs = Presentation(str(pptx))
        refs = {str(f["element_ref"]) for f in logos}
        found = 0
        for master in prs.slide_masters:
            for layout in master.slide_layouts:
                found += sum(1 for sh in iter_shapes(layout) if str(sh.shape_id) in refs)
        return found

    assert logo_shapes(tmp_path / "keep.pptx") > 0, "с флагом keep знак остаётся"
    dropped = compose_deck(
        plan_with("drop"),
        profile,
        path,
        package,
        out_pptx=tmp_path / "drop.pptx",
        package_dir=assets,
    )
    assert dropped.deck["template_logo"] == "drop"
    assert dropped.report["counts"].get("logos_removed", 0) > 0
    assert logo_shapes(tmp_path / "drop.pptx") == 0, "знак снят со всех макетов"

