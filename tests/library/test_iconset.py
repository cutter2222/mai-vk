"""Набор иконок: разбор путей SVG в команды DrawingML, подбор по тексту, фигура на слайде."""

from __future__ import annotations

import math

import pytest
from pptx import Presentation

from presentation_designer.library import iconset
from presentation_designer.library.iconset.geometry import (
    arc_to_cubic,
    custgeom_xml,
    element_ops,
    parse_path,
)


def test_relative_path_with_packed_arc_flags() -> None:
    # Флаги дуги пишутся слитно («0 01»), числа — без пробелов («.5.5»).
    ops = parse_path("M10 13H6m4 2v-4a2 2 0 0 0-4 0v4M14 14.5a.5.5 0 01.5.5z")
    assert ops[0] == ("M", 10.0, 13.0)
    assert ops[1] == ("L", 6.0, 13.0)
    assert ops[2] == ("M", 10.0, 15.0)
    assert ops[3] == ("L", 10.0, 11.0)
    curves = [op for op in ops if op[0] == "C"]
    assert curves, "дуга стала кубическими кривыми"
    assert ops[-1] == ("Z",)


def test_arc_ends_exactly_at_target_and_stays_on_circle() -> None:
    # Полуокружность радиуса 2 от (0, 0) до (4, 0): две четверти.
    ops = arc_to_cubic(0, 0, 2, 2, 0, False, True, 4, 0)
    assert len(ops) == 2
    assert ops[-1][5:] == (4, 0)
    mid = ops[0][5:]
    assert math.isclose(math.dist(mid, (2, 0)), 2, abs_tol=1e-6)


def test_shapes_become_closed_paths() -> None:
    circle = element_ops("circle", {"cx": "12", "cy": "12", "r": "10"})
    assert circle[0][0] == "M" and circle[-1] == ("Z",)
    rect = element_ops("rect", {"x": "2", "y": "5", "width": "20", "height": "14", "rx": "2"})
    assert rect[-1] == ("Z",) and sum(op[0] == "C" for op in rect) == 4
    line = element_ops("line", {"x1": "1", "y1": "2", "x2": "3", "y2": "4"})
    assert line == [("M", 1.0, 2.0), ("L", 3.0, 4.0)]


def test_every_icon_converts() -> None:
    # Весь набор разбирается без ошибок: неизвестная команда уронила бы вёрстку слайда.
    for name in iconset.icon_names():
        xml = custgeom_xml(iconset.icon_nodes(name) or [])
        assert "<a:pathLst>" in xml, name


@pytest.mark.parametrize(
    ("text", "icon"),
    [
        ("Безопасность данных", "shield-check"),
        ("Рост выручки", "trending-up"),
        ("Текучесть кадров снизилась", "user-minus"),
        ("rocket", "rocket"),
        ("team users", "users"),
    ],
)
def test_find_icon_by_meaning(text: str, icon: str) -> None:
    assert iconset.find_icon(text) == icon


def test_unknown_russian_text_gets_no_random_icon() -> None:
    # Обрывки латиницы в русском тексте («IT») не должны тянуть случайную иконку по тегам.
    assert iconset.find_icon("Практика IT-сектора") is None


def test_row_icons_do_not_repeat() -> None:
    icons = iconset.pick_icons(["Рост продаж", "Рост выручки", "Безопасность"])
    assert icons[0] == "trending-up"
    assert icons[1] != "trending-up"
    assert icons[2] == "shield-check"


def test_icon_is_one_recolorable_shape() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = iconset.add_icon(slide, "rocket", 100, 200, 457200, "#FF3985")
    xml = shape._element.xml
    assert shape.name == "Icon rocket"
    assert "a:custGeom" in xml and "a:prstGeom" not in xml
    assert 'val="FF3985"' in xml
    assert "p:style" not in xml, "вид не зависит от стиля фигур темы"
