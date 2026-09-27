"""Content-aware emphasis preserves native text, geometry and the template hierarchy."""

from __future__ import annotations

from io import BytesIO
from typing import Any

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from presentation_designer.library.build import build_slide
from presentation_designer.library.dress import LOOKS, dress_slide
from presentation_designer.library.spec import find_composition
from presentation_designer.library.tokens import DesignCode
from presentation_designer.library.typography import emphasize, required_height


def _box(slide: Any, text: str, *, width: float = 3, height: float = 1) -> Any:
    shape = slide.shapes.add_textbox(0, 0, Inches(width), Inches(height))
    shape.text = text
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(12)
    return shape


def test_peers_use_one_measured_size_without_changing_runs_or_geometry() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shapes = [
        _box(slide, "Рост"),
        _box(slide, "Экономия на обслуживании региональных подразделений", height=0.7),
    ]
    shapes[1].text_frame.paragraphs[0].runs[0].font.bold = True
    before = [(s.shape_id, s.text, s.left, s.top, s.width, s.height) for s in shapes]
    assert emphasize(shapes, 24, "Arial")
    sizes = {r.font.size.pt for s in shapes for p in s.text_frame.paragraphs for r in p.runs}
    assert len(sizes) == 1
    size = sizes.pop()
    assert 12 <= size < 24
    assert all(required_height(s, size, "Arial") <= s.height for s in shapes)
    assert shapes[1].text_frame.paragraphs[0].runs[0].font.bold
    assert before == [(s.shape_id, s.text, s.left, s.top, s.width, s.height) for s in shapes]


def test_overflow_is_not_hidden_by_shrinking_or_deleting_text() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shapes = [_box(slide, "Рост"), _box(slide, "Длинная подпись " * 80, height=0.2)]
    before = [s._element.xml for s in shapes]
    assert not emphasize(shapes, 24, "Arial")
    assert before == [s._element.xml for s in shapes]


def test_measurement_counts_soft_breaks_blank_paragraphs_and_spacing() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = _box(slide, "Один")
    one = required_height(shape, 16, "Arial")
    shape.text = "Один\vДва\n\nТри"
    four = required_height(shape, 16, "Arial")
    assert four > one * 3
    shape.text_frame.paragraphs[0].space_after = Pt(12)
    shape.text_frame.paragraphs[0].line_spacing = 1.5
    assert required_height(shape, 16, "Arial") > four + Pt(12)


def test_long_kpi_label_is_not_blindly_promoted_to_body_size() -> None:
    prs = Presentation()
    code = DesignCode(body_pt=24, caption_pt=12, body_font="Arial")
    composition = find_composition("kpi_row@card=True,count=3")
    assert composition is not None
    slide, refs, cards = build_slide(prs, prs.slide_layouts[6], composition, code)
    shapes = {str(s.shape_id): s for s in slide.shapes}
    label = shapes[refs["kpi_1_label"]]
    label.text_frame.paragraphs[0].runs[0].text = "Экономия региональных подразделений"
    label.height = Inches(0.65)
    value = shapes[refs["kpi_1_value"]]
    value.text_frame.paragraphs[0].runs[0].text = "40 %"
    assert required_height(label, code.body_pt, "Arial") > label.height
    assert required_height(label, code.caption_pt, "Arial") <= label.height
    dress_slide(
        slide,
        composition,
        refs,
        cards,
        code,
        LOOKS["balanced"],
        width=prs.slide_width,
        height=prs.slide_height,
        backdrop=code.background,
    )
    size = label.text_frame.paragraphs[0].runs[0].font.size.pt
    assert code.caption_pt <= size < code.body_pt
    assert required_height(label, size, "Arial") <= label.height


def test_unknown_or_empty_text_is_left_alone() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = _box(slide, "Без явного кегля")
    shape.text_frame.paragraphs[0].runs[0].font.size = None
    before = shape._element.xml
    assert not emphasize([None, shape], 24, "Arial")
    assert shape._element.xml == before
    assert not emphasize([], 24, "Arial")


@pytest.mark.parametrize("inherited", [True, False])
def test_mixed_run_sizes_do_not_lose_their_hierarchy(inherited: bool) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = _box(slide, "Основной текст", height=2)
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.font.size = Pt(32)
    note = paragraph.add_run()
    note.text = " с уточнением"
    note.font.size = None if inherited else Pt(10)
    before = shape._element.xml
    assert not emphasize([shape], 24, "Arial")
    assert shape._element.xml == before


def test_one_unknown_peer_keeps_entire_group_unchanged() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    peers = [_box(slide, "Явный размер"), _box(slide, "Размер из шаблона")]
    peers[1].text_frame.paragraphs[0].runs[0].font.size = None
    before = [s._element.xml for s in peers]
    assert not emphasize(peers, 24, "Arial")
    assert [s._element.xml for s in peers] == before


@pytest.mark.parametrize("canvas", [(13.333, 7.5), (10, 7.5)])
@pytest.mark.parametrize("dark", [False, True])
@pytest.mark.parametrize(
    ("composition_id", "slot_id", "text", "target_role"),
    [
        ("diagram_pane@lead=True", "lead", "План расширения сервиса на четвёртый квартал", "lead"),
        ("title_slide", "title", "Умные уведомления", "cover"),
        ("closing", "title", "Следующий шаг", "cover"),
        ("section@numbered=False", "title", "Результаты пилота", "section"),
        ("section@numbered=True", "lead", "Проверенные результаты и ограничения", "lead"),
        ("agenda@cols=1", "agenda_1", "Проблема, решение, результаты и следующий шаг", "agenda"),
        ("kpi_row@card=True,count=3", "kpi_1_label", "Экономия времени", "label"),
    ],
)
def test_builtin_emphasis_roundtrip(
    canvas: tuple[float, float],
    dark: bool,
    composition_id: str,
    slot_id: str,
    text: str,
    target_role: str,
) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = (Inches(v) for v in canvas)
    code = DesignCode(
        title_font="Arial",
        body_font="Arial",
        title_pt=30,
        subtitle_pt=22,
        body_pt=16,
        caption_pt=12,
        background="#162238" if dark else "#FFFFFF",
        text_color="#FFFFFF" if dark else "#111111",
    )
    composition = find_composition(composition_id)
    assert composition is not None
    slide, refs, cards = build_slide(prs, prs.slide_layouts[6], composition, code)
    shapes = {str(s.shape_id): s for s in slide.shapes}
    target = shapes[refs[slot_id]]
    target.text_frame.paragraphs[0].runs[0].text = text
    if target_role == "label":
        shapes[refs["kpi_1_value"]].text_frame.paragraphs[0].runs[0].text = "40 %"
    before = {s.shape_id: (s.left, s.top, s.width, s.height, s.text) for s in slide.shapes}
    original_size = target.text_frame.paragraphs[0].runs[0].font.size.pt
    done = dress_slide(
        slide,
        composition,
        refs,
        cards,
        code,
        LOOKS["balanced"],
        width=prs.slide_width,
        height=prs.slide_height,
        backdrop=code.background,
    )
    assert done
    size = target.text_frame.paragraphs[0].runs[0].font.size.pt
    assert size > original_size
    assert required_height(target, size, "Arial") <= target.height
    assert {
        s.shape_id: (s.left, s.top, s.width, s.height, s.text)
        for s in slide.shapes
        if s.shape_id in before
    } == before
    stream = BytesIO()
    prs.save(stream)
    stream.seek(0)
    reopened = Presentation(stream)
    saved = next(s for s in reopened.slides[0].shapes if s.shape_id == target.shape_id)
    assert saved.text == text
    assert saved.text_frame.paragraphs[0].runs[0].font.size.pt == size


@pytest.mark.parametrize("composition_id", ["title_slide", "closing", "diagram_pane@lead=True"])
def test_long_service_titles_and_leads_are_not_rewritten(composition_id: str) -> None:
    prs = Presentation()
    code = DesignCode(title_font="Arial", body_font="Arial")
    composition = find_composition(composition_id)
    assert composition is not None
    slide, refs, cards = build_slide(prs, prs.slide_layouts[6], composition, code)
    slot = "lead" if composition.family == "diagram_pane" else "title"
    target = next(s for s in slide.shapes if str(s.shape_id) == refs[slot])
    target.text_frame.paragraphs[0].runs[0].text = "Условия масштабирования пилотного проекта " * 40
    before = target._element.xml
    done = dress_slide(
        slide,
        composition,
        refs,
        cards,
        code,
        LOOKS["balanced"],
        width=prs.slide_width,
        height=prs.slide_height,
        backdrop=code.background,
    )
    assert "cover_typography" not in done and "lead_typography" not in done
    assert target._element.xml == before
