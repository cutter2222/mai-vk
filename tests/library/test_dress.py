"""Отделка собственной композиции после заполнения: значки, кегль по роли, высота по тексту."""

from __future__ import annotations

import pathlib
from typing import Any

import pytest
from pptx import Presentation

from presentation_designer.layout.package import layout_by_id
from presentation_designer.library.build import build_slide
from presentation_designer.library.dress import LOOKS, dress_slide, split_breaks
from presentation_designer.library.spec import find_composition
from presentation_designer.library.tokens import DesignCode
from tests.library.test_library import analyze

CARD_TEXTS = {
    "card_1_title": "Безопасность данных",
    "card_1_body": "Шифрование и контроль доступа к личным кабинетам",
    "card_2_title": "Рост выручки",
    "card_2_body": "Продажи выросли на 9 % за четыре месяца пилота",
    "card_3_title": "Команда",
    "card_3_body": "Двенадцать человек: учёные и инженеры",
}


def _built(mini_template: pathlib.Path, composition_id: str) -> tuple[Any, ...]:
    profile = analyze(mini_template)
    code = DesignCode.from_profile(profile)
    composition = find_composition(composition_id)
    assert composition is not None
    prs = Presentation(str(mini_template))
    layout = layout_by_id(prs, str(profile["layouts"][0]["layout_id"]))
    slide, refs, card_ids = build_slide(prs, layout, composition, code)
    return prs, slide, composition, refs, card_ids, code


def _fill(slide: Any, refs: dict[str, str], texts: dict[str, str]) -> None:
    shapes = {str(s.shape_id): s for s in slide.shapes}
    for slot_id, text in texts.items():
        shapes[refs[slot_id]].text_frame.paragraphs[0].runs[0].text = text


def _overlap(a: Any, b: Any) -> bool:
    return (
        a.left < b.left + b.width
        and b.left < a.left + a.width
        and a.top < b.top + b.height
        and b.top < a.top + a.height
    )


@pytest.mark.parametrize("variant", ["balanced", "compact", "detailed"])
def test_cards_get_icon_badges_and_readable_text(mini_template: pathlib.Path, variant: str) -> None:
    prs, slide, composition, refs, card_ids, code = _built(
        mini_template, "cards_grid@cols=3,numbered=False,rows=1"
    )
    _fill(slide, refs, {"title": "Итоги пилота", **CARD_TEXTS})
    done = dress_slide(
        slide, composition, refs, card_ids, code, LOOKS[variant],
        width=int(prs.slide_width), height=int(prs.slide_height), backdrop=code.background,
    )  # fmt: skip
    assert "cards" in done
    names = [s.name for s in slide.shapes]
    assert names.count("Badge") == 3
    assert {"Icon shield-check", "Icon trending-up", "Icon users"} <= set(names)
    shapes = {str(s.shape_id): s for s in slide.shapes}
    texts = [shapes[refs[k]] for k in CARD_TEXTS]
    for shape in texts:
        size = shape.text_frame.paragraphs[0].runs[0].font.size.pt
        assert size >= code.caption_pt, "текст карточки не мельче подписи шаблона"
        assert 0 <= shape.left and shape.left + shape.width <= prs.slide_width
        assert 0 <= shape.top and shape.top + shape.height <= prs.slide_height
    for i, a in enumerate(texts):
        for b in texts[i + 1 :]:
            assert not _overlap(a, b), "текстовые рамки карточек не заходят друг на друга"
    title = shapes[refs["card_1_title"]].text_frame.paragraphs[0].runs[0].font.size.pt
    body = shapes[refs["card_1_body"]].text_frame.paragraphs[0].runs[0].font.size.pt
    assert title > body, "заголовок карточки крупнее её текста"


def test_numbered_cards_keep_numbers_instead_of_icons(mini_template: pathlib.Path) -> None:
    prs, slide, composition, refs, card_ids, code = _built(
        mini_template, "cards_grid@cols=3,numbered=True,rows=1"
    )
    numbers = {f"card_{i}_number": f"0{i}" for i in (1, 2, 3)}
    _fill(slide, refs, {"title": "Шаги", **CARD_TEXTS, **numbers})
    dress_slide(
        slide, composition, refs, card_ids, code, LOOKS["balanced"],
        width=int(prs.slide_width), height=int(prs.slide_height), backdrop=code.background,
    )  # fmt: skip
    names = [s.name for s in slide.shapes]
    assert not any(n.startswith("Icon ") for n in names)
    shapes = {str(s.shape_id): s for s in slide.shapes}
    assert shapes[refs["card_2_number"]].text_frame.text == "02", "номер из плана сохранён"


def test_line_breaks_become_list_items_without_changing_text() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(0, 0, 3000000, 2000000)
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = "План на квартал.\v• Октябрь: два продукта\v• Ноябрь: тихие часы"
    before = box.text_frame.text.replace("\v", "")
    assert split_breaks(box) == 3
    assert len(box.text_frame.paragraphs) == 3
    assert box.text_frame.text.replace("\n", "") == before
