"""Шаблон, где оформление нарисовано на слайдах, а не в макете: фон и полоса под заголовком
на каждом образце, макет пустой. Собственная композиция берёт у образцов фон и декор
(`library.skin`), а дизайн-код — цвета слайдов, если тема файла с ними не связана."""

from __future__ import annotations

import pathlib
from typing import Any

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt

from presentation_designer.library.build import build_slide
from presentation_designer.library.skin import template_skin
from presentation_designer.library.spec import find_composition
from presentation_designer.library.tokens import DesignCode, theme_detached

W, H = 12192000, 6858000
BACKGROUND = "F4F7F8"
TITLE = "142B3A"
BAR_Y = 0.21


def _solid_background(slide: Any, color: str) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor.from_string(color)


def _text(
    slide: Any,
    box: tuple[float, float, float, float],
    text: str,
    size: float,
    color: str,
    bold: bool = False,
) -> None:
    x, y, w, h = box
    shape = slide.shapes.add_textbox(
        Emu(int(x * W)), Emu(int(y * H)), Emu(int(w * W)), Emu(int(h * H))
    )
    run = shape.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def _bar(slide: Any) -> None:
    for i, color in enumerate(("D66E62", "E7A844", "7665A7", "4D8EB3")):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Emu(int((0.056 + i * 0.22) * W)),
            Emu(int(BAR_Y * H)),
            Emu(int(0.2 * W)),
            Emu(int(0.008 * H)),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(color)
        shape.line.fill.background()


@pytest.fixture(scope="module")
def drawn_template(tmp_path_factory: pytest.TempPathFactory) -> pathlib.Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    blank = prs.slide_layouts[6]
    cover = prs.slides.add_slide(blank)
    _solid_background(cover, "254E66")
    _text(cover, (0.06, 0.2, 0.6, 0.2), "Этапы исследования", 54, "FFFFFF", bold=True)
    for n in range(3):
        slide = prs.slides.add_slide(blank)
        _solid_background(slide, BACKGROUND)
        _text(slide, (0.056, 0.064, 0.887, 0.083), f"Заголовок образца {n + 1}", 36, TITLE, True)
        _bar(slide)
        _text(slide, (0.06, 0.4, 0.8, 0.3), "Текст образца под полосой. " * 4, 16, "607681")
    path = tmp_path_factory.mktemp("drawn") / "drawn.pptx"
    prs.save(str(path))
    return path


def _profile(path: pathlib.Path) -> dict[str, Any]:
    """Профиль как у анализатора, но с ролями, известными заранее (без VLM)."""
    prs = Presentation(str(path))
    patterns = []
    for index, slide in enumerate(prs.slides, start=1):
        cover = index == 1
        ids = [str(s.shape_id) for s in slide.shapes]
        texts = [s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
        slots = []
        for shape in texts:
            run = shape.text_frame.paragraphs[0].runs[0]
            slots.append(
                {
                    "slot_id": "title_1" if shape is texts[0] else "body_1",
                    "kind": "title" if shape is texts[0] else "body",
                    "bbox": {
                        "x": shape.left / W,
                        "y": shape.top / H,
                        "width": shape.width / W,
                        "height": shape.height / H,
                    },
                    "element_ref": str(shape.shape_id),
                    "sample_text": shape.text_frame.text,
                    "font": {
                        "family": "Arial",
                        "size_pt": run.font.size.pt,
                        "bold": bool(run.font.bold),
                        "color": f"#{run.font.color.rgb}",
                    },
                }
            )
        static = [i for i in ids if i not in {s["element_ref"] for s in slots}]
        patterns.append(
            {
                "pattern_id": f"pat_s{index}",
                "role": "title" if cover else "cards",
                "source": {"kind": "sample_slide", "slide_index": index, "layout_id": "L7"},
                "slots": slots,
                "static_object_ids": static,
                "tone": {
                    "background": "dark" if cover else "light",
                    "luminance": 0.07 if cover else 0.93,
                    "source": "slide_fill",
                },
            }
        )
    return {
        "slide_size": {"width_emu": W, "height_emu": H},
        "patterns": patterns,
        "design_tokens": {
            "colors": {
                "theme": {
                    "dk1": "#000000",
                    "lt1": "#FFFFFF",
                    "accent1": "#156082",
                    "accent2": "#E97132",
                    "accent3": "#196B24",
                },
                "palette": [
                    {"hex": "#FFFFFF", "role": "background", "usage_count": 9, "source": "theme"},
                    {"hex": "#156082", "role": "primary", "usage_count": 0, "source": "theme"},
                    {"hex": "#E97132", "role": "secondary", "usage_count": 0, "source": "theme"},
                    {"hex": "#41677A", "role": "accent", "usage_count": 30, "source": "slides"},
                    {"hex": "#D66E62", "role": "accent", "usage_count": 12, "source": "slides"},
                    {"hex": "#E7A844", "role": "accent", "usage_count": 12, "source": "slides"},
                    {"hex": "#7665A7", "role": "accent", "usage_count": 9, "source": "slides"},
                ],
            }
        },
    }


def test_design_code_takes_slide_colors_when_theme_is_unused(drawn_template: pathlib.Path) -> None:
    profile = _profile(drawn_template)
    assert theme_detached(profile)
    code = DesignCode.from_profile(profile)
    assert code.accents_from_slides
    assert code.title_color == f"#{TITLE}" and code.title_bold
    assert "#156082" not in code.accents, "синий Office из неиспользуемой темы не акцент"
    assert code.accents[0] == "#41677A"


def test_design_code_keeps_theme_that_slides_use(drawn_template: pathlib.Path) -> None:
    profile = _profile(drawn_template)
    palette = profile["design_tokens"]["colors"]["palette"]
    palette[1]["usage_count"] = 40  # accent1 темы стоит на объектах: тема связана
    assert not theme_detached(profile)
    code = DesignCode.from_profile(profile)
    assert not code.accents_from_slides
    assert code.accents[0] == "#156082"


def test_skin_reads_slide_background_and_repeated_decor(drawn_template: pathlib.Path) -> None:
    profile = _profile(drawn_template)
    prs = Presentation(str(drawn_template))
    skin = template_skin(list(prs.slides), profile)
    assert skin.background == f"#{BACKGROUND}", "фон содержательных образцов, не обложки"
    assert len(skin.decor) == 4, "четыре сегмента полосы стоят на всех образцах"
    assert skin.title_box is not None and skin.title_box[1] == pytest.approx(0.064, abs=0.002)


def test_builtin_slide_wears_template_skin(drawn_template: pathlib.Path) -> None:
    profile = _profile(drawn_template)
    prs = Presentation(str(drawn_template))
    skin = template_skin(list(prs.slides), profile)
    composition = find_composition("cards_grid@cols=3,numbered=False,rows=1")
    assert composition is not None
    slide, refs, _ = build_slide(
        prs, prs.slide_layouts[6], composition, DesignCode.from_profile(profile), skin
    )
    fill = slide.background.fill
    assert str(fill.fore_color.rgb) == BACKGROUND, "фон образцов на новом слайде"
    bars = [s for s in slide.shapes if abs(s.top / H - BAR_Y) < 0.002 and s.height / H < 0.02]
    assert len(bars) == 4, "полоса под заголовком перенесена"
    title = next(s for s in slide.shapes if str(s.shape_id) == refs["title"])
    assert (title.top + title.height) / H <= BAR_Y, "заголовок поднят над полосой"
    run = title.text_frame.paragraphs[0].runs[0]
    assert run.font.bold and str(run.font.color.rgb) == TITLE


def test_section_slide_gets_background_without_header_decor(drawn_template: pathlib.Path) -> None:
    profile = _profile(drawn_template)
    prs = Presentation(str(drawn_template))
    skin = template_skin(list(prs.slides), profile)
    composition = find_composition("section@numbered=False")
    assert composition is not None
    slide, _, _ = build_slide(
        prs, prs.slide_layouts[6], composition, DesignCode.from_profile(profile), skin
    )
    assert str(slide.background.fill.fore_color.rgb) == BACKGROUND
    bars = [s for s in slide.shapes if abs(s.top / H - BAR_Y) < 0.002 and s.height / H < 0.02]
    assert not bars, "шапка образцов над заголовком посередине не ставится"
