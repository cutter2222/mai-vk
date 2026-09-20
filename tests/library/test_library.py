"""Библиотека собственных композиций: дизайн-код, разворачивание семейств, регистрация
паттернов, построение слайда в пакете шаблона и приоритет паттернов автора в отборе."""

from __future__ import annotations

import pathlib
from typing import Any

import pytest
from pptx import Presentation

from presentation_designer.generation import matching as mt
from presentation_designer.layout.package import layout_by_id
from presentation_designer.library.build import ANCHOR_KINDS, build_slide
from presentation_designer.library.register import builtin_patterns
from presentation_designer.library.spec import find_composition, load_families
from presentation_designer.library.tokens import DesignCode
from presentation_designer.parsing.template.analyzer import analyze_template


def analyze(path: pathlib.Path) -> dict[str, Any]:
    return analyze_template(
        path,
        template_id="tpl_test",
        name=path.name,
        size_bytes=path.stat().st_size,
        render=False,
        use_vlm=False,
    ).profile


def test_families_expand_to_distinct_compositions() -> None:
    families = load_families()
    assert families, "реестр семейств не пуст"
    compositions = [c for f in families for c in f.expand()]
    ids = [c.composition_id for c in compositions]
    assert len(ids) == len(set(ids)), "идентификаторы композиций уникальны"
    assert len(compositions) > len(families), "параметры дают больше композиций, чем семейств"
    for c in compositions:
        assert c.slots, f"{c.composition_id}: есть слоты"
        assert any(s.kind == "title" for s in c.slots), f"{c.composition_id}: есть заголовок"
        for slot in c.slots:
            x, y, w, h = slot.bbox
            assert 0 <= x < 1 and 0 <= y < 1, f"{c.composition_id}/{slot.slot_id}: начало на холсте"
            assert 0 < w <= 1 and 0 < h <= 1 and x + w <= 1.001 and y + h <= 1.001, (
                f"{c.composition_id}/{slot.slot_id}: не выходит за холст"
            )
        assert find_composition(c.composition_id) is not None, "композиция находится по id"


def test_design_code_reads_tokens_of_template(mini_template: pathlib.Path) -> None:
    profile = analyze(mini_template)
    code = DesignCode.from_profile(profile)
    assert code.title_pt >= code.body_pt > 0
    assert code.accents and all(a.startswith("#") for a in code.accents)
    assert code.text_color.startswith("#") and code.background.startswith("#")
    assert profile["design_tokens"]["shape"]["card_geometry"] in ("rect", "roundRect")


def test_design_code_falls_back_without_tokens() -> None:
    code = DesignCode.from_profile({})
    assert code.missing, "пустой профиль отмечен как без токенов"
    assert code.accents and code.title_pt > 0, "умолчания не ломают построение"


def test_builtin_patterns_join_profile_with_capacity(mini_template: pathlib.Path) -> None:
    profile = analyze(mini_template)
    builtin = [p for p in profile["patterns"] if p["source"]["kind"] == "builtin"]
    assert builtin, "композиции попали в профиль"
    for p in builtin:
        assert p["source"]["composition_id"] and p["source"]["layout_id"]
        text_slots = [s for s in p["slots"] if s["kind"] not in ANCHOR_KINDS]
        assert text_slots, f"{p['pattern_id']}: есть текстовые слоты"
        for slot in text_slots:
            cap = slot["capacity"]
            assert cap["max_chars"] > 0 and cap["max_lines"] >= 1
            assert cap["measured_with"]["method"] in ("font_metrics", "heuristic")


def test_template_patterns_win_when_suitable(mini_template: pathlib.Path) -> None:
    """Паттерн шаблона идёт первым; своя композиция — следом, без него — первой."""
    profile = analyze(mini_template)
    patterns = mt.profile_patterns(profile)
    need = mt.Need("cards", items=3)
    candidates = mt.candidates_for(patterns, need, "balanced", limit=4)
    assert candidates, "кандидаты нашлись"
    assert not candidates[0].builtin, "первым идёт паттерн загруженного шаблона"
    only_builtin = [p for p in patterns if p.builtin]
    assert mt.candidates_for(only_builtin, need, "balanced", limit=2), (
        "без паттернов шаблона своя композиция становится кандидатом"
    )


def test_builtin_patterns_without_layouts_are_empty() -> None:
    assert builtin_patterns({"layouts": []}) == []


def test_build_slide_puts_shapes_on_template_layout(mini_template: pathlib.Path) -> None:
    profile = analyze(mini_template)
    pattern = next(
        p
        for p in profile["patterns"]
        if p["source"]["kind"] == "builtin" and p["source"]["composition_id"].startswith("kpi_row")
    )
    composition = find_composition(pattern["source"]["composition_id"])
    assert composition is not None
    prs = Presentation(str(mini_template))
    layout = layout_by_id(prs, pattern["source"]["layout_id"])
    assert layout is not None, "макет композиции есть в шаблоне"
    before = len(prs.slides._sldIdLst)
    slide, refs, card_ids = build_slide(prs, layout, composition, DesignCode.from_profile(profile))
    assert len(prs.slides._sldIdLst) == before + 1
    assert not list(slide.placeholders), "подсказки макета убраны"
    assert set(refs) == {s.slot_id for s in composition.slots}, "у каждого слота есть объект"
    assert len(set(refs.values())) == len(refs), "идентификаторы объектов не повторяются"
    assert len(card_ids) == len(composition.cards)
    shape_ids = {str(s.shape_id) for s in slide.shapes}
    assert set(refs.values()) <= shape_ids and set(card_ids) <= shape_ids


@pytest.mark.parametrize(
    "composition_id", ["kpi_row@card=True,count=3", "cards_grid@cols=2,numbered=False,rows=1"]
)
def test_cards_are_painted_in_template_colors(
    mini_template: pathlib.Path, composition_id: str
) -> None:
    profile = analyze(mini_template)
    code = DesignCode.from_profile(profile)
    composition = find_composition(composition_id)
    assert composition is not None and composition.cards
    prs = Presentation(str(mini_template))
    layout = layout_by_id(prs, str(profile["layouts"][0]["layout_id"]))
    slide, _, card_ids = build_slide(prs, layout, composition, code)
    cards = [s for s in slide.shapes if str(s.shape_id) in set(card_ids)]
    assert cards, "плашки созданы"
    palette = {c.lstrip("#").upper() for c in [*code.accents, code.background, code.text_color]}
    for card in cards:
        assert card.fill.type is not None, "заливка задана явно, а не унаследована"
        rgb = str(card.fill.fore_color.rgb).upper()
        # Цвет — либо акцент шаблона, либо акцент, разбавленный его фоном: чужих красок нет.
        assert rgb in palette or _between(rgb, palette), f"цвет {rgb} вне дизайн-кода"


def _between(rgb: str, palette: set[str]) -> bool:
    """Цвет лежит между акцентом и фоном: мягкая подложка карточки."""
    value = tuple(int(rgb[i : i + 2], 16) for i in (0, 2, 4))
    bounds = [tuple(int(p[i : i + 2], 16) for i in (0, 2, 4)) for p in palette]
    return all(min(b[k] for b in bounds) <= value[k] <= max(b[k] for b in bounds) for k in range(3))
