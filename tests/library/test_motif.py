"""Фоновый мотив шаблона: встаёт на пустоту слайда под содержание, бледно и заметно на фоне."""

from __future__ import annotations

import io
from typing import Any

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt

from presentation_designer.library.motif import (
    NAME,
    OPACITY,
    Motif,
    motifs_for,
    occupied,
    place_motif,
)

W, H = 12192000, 6858000


def _motif_png(color: tuple[int, int, int]) -> bytes:
    img = Image.new("RGBA", (400, 600), (0, 0, 0, 0))
    draw = ImageDraw.Draw(img)
    draw.ellipse((0, 0, 400, 400), fill=(*color, 255))
    draw.rectangle((100, 350, 400, 600), fill=(*color, 255))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _deck(
    color: tuple[int, int, int] = (0, 110, 255), crop: bool = False
) -> tuple[Any, list[Motif]]:
    """Презентация, где мотив лежит картинкой на первом слайде, и мотивы по профилю."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    blank = prs.slide_layouts[6]
    source = prs.slides.add_slide(blank)
    picture = source.shapes.add_picture(io.BytesIO(_motif_png(color)), 0, 0)
    part = source.part.related_part(picture._element.xpath(".//a:blip/@r:embed")[0])
    profile = {
        "assets": [
            {
                "asset_id": "asset_1",
                "kind": "image",
                "media_path": str(part.partname).lstrip("/"),
                "sha256": "x",
                "tags": ["decor", "decor:8"] + (["decor-crop"] if crop else []),
                "bbox_on_source": {"x": 0.72, "y": 0.0, "width": 0.28, "height": 1.0},
            }
        ]
    }
    return prs, motifs_for(profile, prs)


def _slide(prs: Any) -> Any:
    return prs.slides.add_slide(prs.slide_layouts[6])


def _text(slide: Any, box: tuple[float, float, float, float], text: str, size: float) -> Any:
    x, y, w, h = box
    shape = slide.shapes.add_textbox(
        Emu(int(x * W)), Emu(int(y * H)), Emu(int(w * W)), Emu(int(h * H))
    )
    shape.text_frame.word_wrap = True
    run = shape.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(size)
    return shape


def _overlaps(a: tuple[float, ...], b: tuple[float, ...]) -> bool:
    return a[0] < b[0] + b[2] and b[0] < a[0] + a[2] and a[1] < b[1] + b[3] and b[1] < a[1] + a[3]


def _frac(shape: Any) -> tuple[float, float, float, float]:
    return (
        int(shape.left) / W,
        int(shape.top) / H,
        int(shape.width) / W,
        int(shape.height) / H,
    )


def test_motif_from_profile_takes_the_template_edge() -> None:
    _prs, motifs = _deck()
    assert len(motifs) == 1
    assert motifs[0].edge == "right" and motifs[0].score == 8


def test_motif_fills_the_empty_side_behind_content() -> None:
    prs, motifs = _deck()
    slide = _slide(prs)
    _text(slide, (0.05, 0.06, 0.9, 0.12), "Заголовок слайда", 32)
    # Рамка списка тянется до низа, а текст занимает её верх: место считается по тексту.
    body = _text(slide, (0.05, 0.25, 0.6, 0.65), "Короткий пункт списка", 18)
    assert place_motif(slide, motifs, width=W, height=H, backdrop=1.0)
    decor = next(s for s in slide.shapes if s.name == NAME)
    # Под всем содержанием и на пустоте: ни заголовок, ни текст списка не задеты.
    assert slide.shapes._spTree[2] is decor._element
    boxes = occupied(slide, W, H) or []
    others = [b for b in boxes if b != _frac(decor)]
    assert not any(_overlaps(_frac(decor), b) for b in others)
    x, y, w, h = _frac(decor)
    assert x + w >= 0.999, "мотив прижат к краю, как в шаблоне"
    assert y <= 0.001 or y + h >= 0.999, "и держится за угол, а не висит посередине края"
    assert int(body.top) < int(decor.top) + int(decor.height)
    # Прозрачность запечена в картинку: бледно, но видно.
    with Image.open(io.BytesIO(decor.image.blob)) as img:
        peak = img.getchannel("A").getextrema()[1]
    assert 0 < peak <= round(255 * OPACITY[1])


def test_full_slide_gets_no_motif() -> None:
    prs, motifs = _deck()
    slide = _slide(prs)
    _text(slide, (0.05, 0.06, 0.9, 0.12), "Заголовок", 32)
    card = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Emu(int(0.03 * W)),
        Emu(int(0.22 * H)),
        Emu(int(0.95 * W)),
        Emu(int(0.75 * H)),
    )
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(0xEE, 0xF2, 0xFF)
    assert not place_motif(slide, motifs, width=W, height=H, backdrop=1.0)
    assert not any(s.name == NAME for s in slide.shapes)


def test_slide_with_photo_gets_no_motif() -> None:
    prs, motifs = _deck()
    slide = _slide(prs)
    photo = Image.new("RGB", (800, 600), (120, 140, 160))
    buf = io.BytesIO()
    photo.save(buf, format="PNG")
    buf.seek(0)
    slide.shapes.add_picture(buf, Emu(int(0.5 * W)), Emu(int(0.2 * H)), Emu(int(0.45 * W)))
    assert occupied(slide, W, H) is None
    assert not place_motif(slide, motifs, width=W, height=H, backdrop=1.0)


def test_white_motif_is_skipped_on_white_and_used_on_dark() -> None:
    prs, motifs = _deck(color=(255, 255, 255))
    light = _slide(prs)
    _text(light, (0.05, 0.06, 0.5, 0.12), "Заголовок", 32)
    assert not place_motif(light, motifs, width=W, height=H, backdrop=1.0)
    dark = _slide(prs)
    _text(dark, (0.05, 0.06, 0.5, 0.12), "Заголовок", 32)
    assert place_motif(dark, motifs, width=W, height=H, backdrop=0.01)


def test_paler_motif_gets_more_opacity() -> None:
    _prs, blue = _deck(color=(0, 90, 255))
    _prs, grey = _deck(color=(200, 204, 210))
    assert blue[0].opacity(1.0) < grey[0].opacity(1.0) <= OPACITY[1]


def test_pattern_goes_further_off_edge_than_a_whole_figure() -> None:
    """Узкая пустота у края: узор (можно резать) выдвигается наполовину, цельная фигура
    не ставится — от неё остался бы обрубок."""

    def narrow(prs: Any) -> Any:
        slide = _slide(prs)
        _text(slide, (0.05, 0.06, 0.9, 0.12), "Заголовок", 32)
        card = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Emu(int(0.03 * W)),
            Emu(int(0.22 * H)),
            Emu(int(0.8 * W)),
            Emu(int(0.75 * H)),
        )
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(0xEE, 0xF2, 0xFF)
        return slide

    prs, figure = _deck()
    assert not place_motif(narrow(prs), figure, width=W, height=H, backdrop=1.0)
    prs, pattern = _deck(crop=True)
    slide = narrow(prs)
    assert place_motif(slide, pattern, width=W, height=H, backdrop=1.0)
    decor = next(s for s in slide.shapes if s.name == NAME)
    assert _frac(decor)[0] >= 0.83 + 0.02


def test_gradient_only_picture_is_not_a_motif() -> None:
    """Картинка из одного полупрозрачного свечения — не мотив: на слайде вышло бы пятно."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    source = prs.slides.add_slide(prs.slide_layouts[6])
    img = Image.new("RGBA", (400, 400), (0, 0, 0, 0))
    glow = Image.linear_gradient("L").resize((400, 400)).point(lambda v: v // 3)
    img.putalpha(glow)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    picture = source.shapes.add_picture(buf, 0, 0)
    part = source.part.related_part(picture._element.xpath(".//a:blip/@r:embed")[0])
    profile = {
        "assets": [
            {
                "asset_id": "asset_1",
                "kind": "image",
                "media_path": str(part.partname).lstrip("/"),
                "sha256": "x",
                "tags": ["decor", "decor:8"],
                "bbox_on_source": {"x": 0.7, "y": 0.0, "width": 0.3, "height": 1.0},
            }
        ]
    }
    assert motifs_for(profile, prs) == []
