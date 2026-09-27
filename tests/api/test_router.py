"""Роутер чата и отмена правок через API (этап 37)."""

from __future__ import annotations

import io

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from presentation_designer.contracts import models as m
from tests.api.test_onlyoffice import (
    SECRET,  # noqa: F401
    office,  # noqa: F401  # фикстура копии
)


def say(client: TestClient, project: str, text: str, file_ids: list[str] | None = None) -> str:
    return client.post(
        f"/api/projects/{project}/events",
        json={"role": "user", "kind": "message", "text": text, "file_ids": file_ids or []},
    ).json()["event_id"]


def test_route_decides_by_rules_and_writes_questions_to_the_feed(client: TestClient) -> None:
    project = client.post("/api/projects", json={}).json()["project_id"]
    undo = client.post(
        "/api/chat/route", json={"project_id": project, "event_id": say(client, project, "отмени")}
    ).json()
    assert undo["kind"] == "run" and undo["steps"][0]["action"] == "undo"
    assert undo["source"] == "rules" and "event" not in undo

    png = io.BytesIO()
    from PIL import Image

    Image.new("RGB", (20, 20)).save(png, "PNG")
    (photo,) = client.post(
        f"/api/projects/{project}/files",
        files=[("files", ("фото.png", png.getvalue(), "image/png"))],
    ).json()
    event = say(client, project, "вставь картинку", [photo["file_id"]])
    asked = client.post("/api/chat/route", json={"project_id": project, "event_id": event}).json()
    assert asked["kind"] == "question" and "слайд" in asked["text"]
    # Вопрос записан в ленту, как ответ ассистента.
    events = client.get(f"/api/projects/{project}").json()["events"]
    assert events[-1]["role"] == "assistant" and events[-1]["text"] == asked["text"]

    # Со слайдом, открытым в редакторе, картинка идёт на него; плашка объекта — полем chip,
    # строка плашки в тексте сообщения разбору не мешает.
    placed = client.post(
        "/api/chat/route",
        json={"project_id": project, "event_id": event, "current_slide": 4},
    ).json()
    assert placed["steps"][0] == {
        "action": "image",
        "slides": [4],
        "instruction": "вставь картинку",
        "file_id": photo["file_id"],
    }
    chip = say(client, project, "Слайд 2 · «Заголовок»\nсократи вдвое")
    edit = client.post(
        "/api/chat/route",
        json={
            "project_id": project,
            "event_id": chip,
            "chip": {"slide": 2, "object": {"name": "Title 1", "label": "«Заголовок»"}},
        },
    ).json()
    assert edit["steps"][0]["action"] == "object_edit"
    assert edit["steps"][0]["instruction"] == "сократи вдвое"
    assert edit["steps"][0]["target"]["name"] == "Title 1"

    bad = client.post("/api/chat/route", json={"project_id": project, "event_id": "evt_missing"})
    assert bad.status_code == 422


@pytest.mark.parametrize("attached", [False, True])
def test_route_keeps_questions_out_of_edit_executors(client: TestClient, attached: bool) -> None:
    project = client.post("/api/projects", json={}).json()["project_id"]
    file_ids = []
    if attached:
        from PIL import Image

        png = io.BytesIO()
        Image.new("RGB", (20, 20)).save(png, "PNG")
        (photo,) = client.post(
            f"/api/projects/{project}/files",
            files=[("files", ("фото.png", png.getvalue(), "image/png"))],
        ).json()
        file_ids.append(photo["file_id"])
    text = "Что на этой картинке?" if attached else "Как удалить логотип?"
    event_id = say(client, project, text, file_ids)
    before = client.get(f"/api/projects/{project}").json()
    response = client.post(
        "/api/chat/route",
        json={"project_id": project, "event_id": event_id, "current_slide": 1},
    )
    assert response.status_code == 200, response.text
    decision = response.json()
    assert decision["kind"] == "answer"
    assert decision["steps"] == []
    assert decision["source"] == "rules"
    assert "event" not in decision
    assert client.get(f"/api/projects/{project}").json() == before


def test_office_undo_restores_bytes_as_a_new_revision(client: TestClient, office) -> None:  # noqa: F811
    store, doc_id = office
    original = store.read(doc_id, 0)
    prs = Presentation(io.BytesIO(original))
    prs.slides[0].shapes.title.text = "Правка"
    edited = io.BytesIO()
    prs.save(edited)
    token = store.begin_edit(doc_id, 0)
    store.commit_edit(doc_id, token, 0, edited.getvalue())
    store.end_edit(doc_id, token)
    assert store.get(doc_id)["revision"] == 1

    url = f"/api/office/documents/{doc_id}/undo"
    assert client.post(url, json={"revision": 1, "to_revision": 1}).status_code == 422
    done = client.post(url, json={"revision": 1, "to_revision": 0})
    assert done.status_code == 200, done.text
    assert done.json()["document"]["revision"] == 2 and done.json()["changed"] is True
    assert store.read(doc_id, 2) == original
    # После правки копию уже меняли (ревизия 2) — отмена правки 1 отдельно невозможна.
    stale = client.post(url, json={"revision": 1, "to_revision": 0})
    assert stale.status_code == 409 and stale.json()["error"]["code"] == "office_undo_stale"


def test_office_sequential_undo_preserves_revision_history(client: TestClient, office) -> None:  # noqa: F811
    store, doc_id = office
    original = store.read(doc_id, 0)
    revisions = [original]
    for base, title in enumerate(["Первая правка", "Вторая правка"]):
        deck = Presentation(io.BytesIO(revisions[-1]))
        deck.slides[0].shapes.title.text = title
        output = io.BytesIO()
        deck.save(output)
        revisions.append(output.getvalue())
        token = store.begin_edit(doc_id, base)
        store.commit_edit(doc_id, token, base, revisions[-1])
        store.end_edit(doc_id, token)
    url = f"/api/office/documents/{doc_id}/undo"
    assert client.post(url, json={"revision": 1, "to_revision": 0}).status_code == 409
    assert client.post(url, json={"revision": 2, "to_revision": 1}).status_code == 200
    assert store.read(doc_id, 3) == revisions[1]
    response = client.post(url, json={"revision": 1, "to_revision": 0})
    assert response.status_code == 200, response.text
    assert store.read(doc_id, 4) == original
    assert store.read(doc_id, 2) == revisions[2]
    assert client.post(url, json={"revision": 1, "to_revision": 0}).status_code == 409


def test_revert_restores_the_slide_as_a_new_revision(
    client: TestClient, pptx_bytes: bytes, xlsx_bytes: bytes
) -> None:
    from tests.pipeline.helpers import run_generation

    run = run_generation(client, pptx_bytes, xlsx_bytes, title="Отмена правки")
    job_id = run["job_id"]
    base = f"/api/generations/{job_id}/variants/balanced"
    edit = client.post(
        f"{base}/edits",
        json={"base_revision": 1, "slide_index": 1, "instruction": "Новый заголовок"},
    )
    assert edit.status_code == 202
    result = m.GenerationResult.model_validate(client.get(f"/api/generations/{job_id}").json())
    balanced = next(v for v in result.variants if v.variant_id == "balanced")
    assert balanced.revision == 2

    stale = client.post(
        f"{base}/revert", json={"base_revision": 1, "to_revision": 1, "slide_index": 1}
    )
    assert stale.status_code == 409 and stale.json()["error"]["code"] == "revision_stale"
    reverted = client.post(
        f"{base}/revert", json={"base_revision": 2, "to_revision": 1, "slide_index": 1}
    )
    assert reverted.status_code == 202, reverted.text
    result = m.GenerationResult.model_validate(client.get(f"/api/generations/{job_id}").json())
    balanced = next(v for v in result.variants if v.variant_id == "balanced")
    assert balanced.revision == 3
    entry = next(e for e in result.edits if e.edit_job_id == reverted.json()["edit_job_id"])
    assert entry.result == "applied" and entry.new_revision == 3 and entry.slide_index == 1
    names = client.get(f"/api/generations/{job_id}").json()["artifacts_manifest"]
    r1 = client.get(f"/api/generations/{job_id}/artifacts/balanced/r1/deck.pptx").content
    r3 = client.get(f"/api/generations/{job_id}/artifacts/balanced/r3/deck.pptx").content
    assert "balanced/r3/deck.pptx" in names and r1 == r3
    status = client.get(f"/api/jobs/{reverted.json()['edit_job_id']}").json()
    assert status["status"] == "succeeded"


@pytest.mark.parametrize("tamper", [None, "plan.json", "deck.pptx"])
def test_revert_chain_requires_identical_revision_files(
    client: TestClient, pptx_bytes: bytes, xlsx_bytes: bytes, orchestrator, tamper
) -> None:
    from tests.pipeline.helpers import run_generation

    job_id = run_generation(client, pptx_bytes, xlsx_bytes)["job_id"]
    base = f"/api/generations/{job_id}/variants/balanced"
    for revision, slide in [(1, 1), (2, 2)]:
        response = client.post(
            f"{base}/edits",
            json={
                "base_revision": revision,
                "slide_index": slide,
                "instruction": "Новый заголовок",
            },
        )
        assert response.status_code == 202, response.text
    first = {"base_revision": 2, "to_revision": 1, "slide_index": 1}
    assert client.post(f"{base}/revert", json=first).status_code == 409
    # Нельзя подменить цель отмены или адрес слайда.
    for to_revision, slide_index in [(1, 2), (2, 1)]:
        response = client.post(
            f"{base}/revert",
            json={"base_revision": 3, "to_revision": to_revision, "slide_index": slide_index},
        )
        assert response.status_code == 422, response.text
    response = client.post(
        f"{base}/revert", json={"base_revision": 3, "to_revision": 2, "slide_index": 2}
    )
    assert response.status_code == 202, response.text
    root = orchestrator.artifacts.revision_dir(job_id, "balanced", 4)
    original = orchestrator.artifacts.revision_dir(job_id, "balanced", 1)
    if tamper:
        path = root / tamper
        path.write_bytes(path.read_bytes() + b" ")
    response = client.post(f"{base}/revert", json=first)
    if tamper:
        assert response.status_code == 409, response.text
        assert orchestrator.state.get_variant(job_id, "balanced")["revision"] == 4
        return
    assert response.status_code == 202, response.text
    restored = orchestrator.artifacts.revision_dir(job_id, "balanced", 5)
    for path in original.rglob("*"):
        if path.is_file() and path.name != "manifest.json":
            assert (restored / path.relative_to(original)).read_bytes() == path.read_bytes()
    entry = orchestrator.state.get_repair(response.json()["edit_job_id"])
    assert entry["base_revision"] == 4 and entry["new_revision"] == 5
    assert client.post(f"{base}/revert", json=first).status_code == 409


def test_edit_result_events_pass_the_contract(client: TestClient) -> None:
    project = client.post("/api/projects", json={}).json()["project_id"]
    event = client.post(
        f"/api/projects/{project}/events",
        json={
            "role": "assistant",
            "kind": "edit_result",
            "text": "Слайд 3: заголовок сокращён.",
            "document_id": "doc_1",
            "revision": 2,
            "base_revision": 1,
            "slides": [3],
        },
    )
    assert event.status_code == 201, event.text
    patched = client.patch(
        f"/api/projects/{project}/events/{event.json()['event_id']}", json={"undone": True}
    )
    assert patched.status_code == 200 and patched.json()["undone"] is True
    doc = client.get(f"/api/projects/{project}").json()
    assert doc["schema_version"] == "1.6" and doc["events"][-1]["undone"] is True
