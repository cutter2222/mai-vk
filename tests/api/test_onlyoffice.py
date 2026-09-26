"""Signed office callbacks, immutable revisions and isolated AI source."""

import io
import time
from concurrent.futures import ThreadPoolExecutor
from urllib.parse import urlsplit

import httpx
import pytest
from pptx import Presentation

from presentation_designer.api.errors import ApiError
from presentation_designer.api.routes import onlyoffice
from presentation_designer.pipeline.office import OfficeStore, sign, verify

SECRET = "test-onlyoffice-secret-not-for-production-123456"


def test_project_template_copy_is_isolated_and_persistent(client, office, pptx_bytes):
    store, _ = office
    template_id = client.post("/api/templates", files={"file": ("Шаблон.pptx", pptx_bytes)}).json()[
        "template_id"
    ]
    copies = []
    for _ in range(2):
        project = client.post("/api/projects", json={}).json()["project_id"]
        client.patch(f"/api/projects/{project}", json={"template_id": template_id})
        url = f"/api/office/projects/{project}/template"
        before = client.get(f"/api/projects/{project}").json()
        response = client.post(url, json={"template_id": template_id})
        assert response.status_code == 200, response.text
        doc = response.json()
        assert store.read(doc["id"], 0) == pptx_bytes
        token = store.begin_edit(doc["id"], 0)
        store.commit_edit(doc["id"], token, 0, pptx_bytes + b"copy-only")
        store.end_edit(doc["id"], token)
        reopened = client.post(url, json={"template_id": template_id}).json()
        assert reopened["id"] == doc["id"]
        assert reopened["revision"] == 1
        assert client.get(f"/api/projects/{project}").json() == before
        assert client.post(url, json={"template_id": "other"}).status_code == 409
        copies.append(doc["id"])
    assert copies[0] != copies[1]
    assert client.get(f"/api/templates/{template_id}/source").content == pptx_bytes
    assert (
        client.post(
            "/api/office/projects/missing/template", json={"template_id": template_id}
        ).status_code
        == 404
    )


def test_multiple_object_edit_is_atomic_and_revision_scoped(client, office, monkeypatch):
    from pptx.util import Inches

    from presentation_designer.generation.office_object_edit import ObjectsEditPlan

    store, doc_id = office
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    for i in range(4):
        slide.shapes.add_textbox(
            Inches(1), Inches(i + 1), Inches(2), Inches(0.5)
        ).text = f"Объект {i}"
    output = io.BytesIO()
    deck.save(output)
    opened = store.open(doc_id)
    store.callback(doc_id, opened["active_key"], 2, output.getvalue())
    base = f"/api/office/documents/{doc_id}"
    objects = client.get(base + "/objects/1").json()["objects"]
    targets = [{"slide": obj["slide"], "shape_id": obj["shape_id"]} for obj in objects[:3]]
    invalid = True

    async def propose(data, instruction, settings, selected, tokens=None):
        assert [target.model_dump() for target in selected] == targets
        assert instruction == "Все три правее"
        return ObjectsEditPlan.model_validate(
            {
                "explanation": "Все три правее",
                "edits": [
                    {
                        "target": target,
                        "plan": {
                            "explanation": "",
                            "patches": [],
                            "position": {
                                "x": 0.9 if invalid and i == 2 else 0.2,
                                "y": objects[i]["bbox"]["y"],
                            },
                        },
                    }
                    for i, target in enumerate(targets)
                ],
            }
        )

    monkeypatch.setattr(onlyoffice.office_object_edit, "propose_many", propose)
    body = {"revision": 1, "instruction": "Все три правее", "targets": targets}
    assert client.post(base + "/edit", json=body).status_code == 422
    assert store.get(doc_id)["revision"] == 1
    assert store.read(doc_id, 1) == output.getvalue()
    invalid = False
    response = client.post(base + "/edit", json=body)
    assert response.status_code == 200, response.text
    assert response.json()["document"]["revision"] == 2
    changed = client.get(base + "/objects/2").json()["objects"]
    assert [obj["bbox"]["x"] for obj in changed] == [0.2, 0.2, 0.2, 0.1]
    assert client.post(base + "/edit", json=body).status_code == 409
    for extra in [{"targets": []}, {"targets": targets + targets[:1]}, {"target": targets[0]}]:
        assert client.post(base + "/edit", json=body | extra).status_code == 422


def test_object_edit_endpoint_is_revision_scoped(client, office, monkeypatch):
    from pptx.util import Inches

    from presentation_designer.generation.office_object_edit import ObjectEditPlan, Position

    store, doc_id = office
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6]).shapes.add_textbox(
        Inches(1), Inches(1), Inches(2), Inches(1)
    ).text = "Цель"
    output = io.BytesIO()
    deck.save(output)
    opened = store.open(doc_id)
    store.callback(doc_id, opened["active_key"], 2, output.getvalue())
    base = f"/api/office/documents/{doc_id}"
    response = client.get(base + "/objects/1")
    assert response.status_code == 200
    obj = response.json()["objects"][0]
    assert obj["label"] == "Цель"
    target = {"slide": obj["slide"], "shape_id": obj["shape_id"]}

    async def propose(data, instruction, settings, selected, tokens=None):
        assert selected.model_dump() == target
        assert data == output.getvalue()
        return ObjectEditPlan(
            explanation="Правее", patches=[], position=Position(x=0.3, y=obj["bbox"]["y"])
        )

    monkeypatch.setattr(onlyoffice.office_object_edit, "propose", propose)
    response = client.post(
        base + "/edit", json={"revision": 1, "instruction": "Правее", "target": target}
    )
    assert response.status_code == 200, response.text
    assert response.json()["document"]["revision"] == 2
    assert client.get(base + "/objects/2").json()["objects"][0]["bbox"]["x"] == 0.3
    assert client.get(base + "/objects/1").json()["objects"][0]["bbox"]["x"] == 0.1
    assert (
        client.post(
            base + "/edit", json={"revision": 1, "instruction": "Правее", "target": target}
        ).status_code
        == 409
    )
    assert store.get(doc_id)["active_key"] is None


def test_saved_revision_preview_without_editor_session(client, office, monkeypatch, pptx_bytes):
    import json

    store, doc_id = office
    sources = []

    def preview(content, settings):
        sources.append(content)
        root = onlyoffice.cache_path(content, settings)
        root.mkdir(parents=True, exist_ok=True)
        manifest = {"slides": ["slide-01.png"], "ratio": 16 / 9}
        (root / "manifest.json").write_text(json.dumps(manifest))
        (root / "slide-01.png").write_bytes(b"png")
        return root, manifest

    monkeypatch.setattr(onlyoffice, "preview_revision", preview)
    base = f"/api/office/documents/{doc_id}/preview"
    assert client.get(base + "/0/slide-01.png").status_code == 404
    assert client.get(base + "/0").json()["revision"] == 0
    image = client.get(base + "/0/slide-01.png")
    assert image.content == b"png"
    assert "immutable" in image.headers["cache-control"]
    assert client.get(base + "/0/manifest.json").status_code == 404
    assert client.get(base + "/999").status_code == 404
    assert sources == [pptx_bytes]
    assert store.get(doc_id)["active_key"] is None

    opened = store.open(doc_id)
    changed = pptx_bytes + b"new-revision"
    store.callback(doc_id, opened["active_key"], 6, changed)
    assert client.get(base + "/1").json()["revision"] == 1
    assert sources[-1] == changed
    assert store.get(doc_id)["active_key"] == opened["active_key"]


def test_preview_conversion_failure_is_retryable(client, office, monkeypatch):
    def fail(*args):
        raise onlyoffice.ConversionError("not ready")

    monkeypatch.setattr(onlyoffice, "preview_revision", fail)
    _, doc_id = office
    response = client.get(f"/api/office/documents/{doc_id}/preview/0")
    assert response.status_code == 503
    assert response.json()["error"]["code"] == "office_preview_failed"


def test_deferred_preview_page_failure_is_retryable(client, office, orchestrator, monkeypatch):
    import json

    store, doc_id = office
    root = onlyoffice.cache_path(store.read(doc_id, 0), orchestrator.settings)
    root.mkdir(parents=True)
    (root / "manifest.json").write_text(json.dumps({"slides": ["slide-02.png"]}))

    def fail(*args):
        raise onlyoffice.ConversionError("failed page")

    monkeypatch.setattr(onlyoffice, "preview_page", fail)
    url = f"/api/office/documents/{doc_id}/preview/0/slide-02.png"
    response = client.get(url)
    assert response.status_code == 503
    assert response.json()["error"]["code"] == "office_preview_failed"
    assert "immutable" not in response.headers.get("cache-control", "")
    assert store.get(doc_id)["active_key"] is None

    def ready(path, name):
        image = path / name
        image.write_bytes(b"ready-png")
        return image

    monkeypatch.setattr(onlyoffice, "preview_page", ready)
    response = client.get(url)
    assert response.status_code == 200
    assert response.content == b"ready-png"
    assert "immutable" in response.headers["cache-control"]


def test_edit_exclusive_lease_and_revision(office, pptx_bytes):
    store, doc_id = office
    opened = store.open(doc_id)
    with pytest.raises(ApiError):
        store.begin_edit(doc_id, 0)
    store.callback(doc_id, opened["active_key"], 4, None)
    token = store.begin_edit(doc_id, 0)
    with pytest.raises(ApiError):
        store.open(doc_id)
    with pytest.raises(ApiError):
        store.begin_edit(doc_id, 0)
    store.commit_edit(doc_id, token, 0, pptx_bytes)
    assert store.get(doc_id)["revision"] == 0
    store.end_edit(doc_id, token)
    assert store.open(doc_id)["active_key"] != opened["active_key"]
    with pytest.raises(ApiError):
        store.commit_edit(doc_id, token, 0, pptx_bytes)
    assert not store.callback(doc_id, opened["active_key"], 2, pptx_bytes)


def test_expired_edit_cannot_commit(office, pptx_bytes):
    store, doc_id = office
    token = store.begin_edit(doc_id, 0)
    with store.connect() as db:
        db.execute("UPDATE edit_locks SET expires=0")
    replacement = store.begin_edit(doc_id, 0)
    store.end_edit(doc_id, token)
    with pytest.raises(ApiError):
        store.commit_edit(doc_id, token, 0, pptx_bytes)
    with pytest.raises(ApiError):
        store.open(doc_id)
    store.end_edit(doc_id, replacement)


def test_edit_endpoint_uses_saved_pptx(client, office, monkeypatch):
    from presentation_designer.generation.office_edit import EditPlan, TextPatch

    store, doc_id = office
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[0]).shapes.title.text = "Ручная правка"
    output = io.BytesIO()
    deck.save(output)
    opened = store.open(doc_id)
    store.callback(doc_id, opened["active_key"], 2, output.getvalue())

    async def propose(data, instruction, settings, only=None):
        assert data == output.getvalue()
        assert instruction == "Измени заголовок"
        with pytest.raises(ApiError):
            store.open(doc_id)
        return EditPlan(
            explanation="Готово",
            patches=[TextPatch(slide=1, run=0, before="Ручная правка", after="После ИИ")],
        )

    monkeypatch.setattr(onlyoffice.office_edit, "propose", propose)
    response = client.post(
        f"/api/office/documents/{doc_id}/edit",
        json={"revision": 1, "instruction": "Измени заголовок"},
    )
    assert response.status_code == 200, response.text
    assert response.json()["changed"] is True
    assert response.json()["document"]["revision"] == 2
    assert store.read(doc_id, 1) == output.getvalue()
    assert Presentation(io.BytesIO(store.read(doc_id, 2))).slides[0].shapes.title.text == "После ИИ"
    assert (
        client.post(
            f"/api/office/documents/{doc_id}/edit",
            json={"revision": 1, "instruction": "Старая база"},
        ).status_code
        == 409
    )
    assert store.open(doc_id)["seed_revision"] == 2


def test_edit_failure_unlocks_without_writing(client, office, monkeypatch):
    store, doc_id = office

    async def fail(*args):
        raise ValueError("Тестовый отказ")

    monkeypatch.setattr(onlyoffice.office_edit, "propose", fail)
    response = client.post(
        f"/api/office/documents/{doc_id}/edit", json={"revision": 0, "instruction": "Правка"}
    )
    assert response.status_code == 422
    assert store.get(doc_id)["revision"] == 0
    assert store.open(doc_id)["seed_revision"] == 0


@pytest.fixture
def office(orchestrator, pptx_bytes):
    orchestrator.settings.onlyoffice.enabled = True
    orchestrator.settings.onlyoffice.jwt_secret = SECRET
    store = OfficeStore(orchestrator.settings.paths.data_dir)
    doc = store.create("job_test/balanced/r1/deck.pptx", "test.pptx", pptx_bytes)
    return store, doc["id"]


def test_disabled(client):
    assert client.get("/api/office/capabilities").json() == {"enabled": False}
    assert client.post("/api/office/documents/no/config").status_code == 503
    assert client.post("/api/office/templates/no/config").status_code == 503


def test_capabilities_give_the_sdk_address_for_warm_up(client, office):
    # Интерфейс прогревает редактор, пока презентация готовится, а не после её появления.
    body = client.get("/api/office/capabilities").json()
    assert body["enabled"] is True
    assert body["script_url"].endswith("/web-apps/apps/api/documents/api.js?rendering=pd16v2")


def test_logo_edit_explains_what_is_missing(client, office):
    _, doc_id = office
    # Документ без задания-источника: шаблона нет, логотип менять не по чему.
    response = client.post(
        f"/api/office/documents/{doc_id}/edit",
        json={"revision": 0, "instruction": "убери логотип", "logo": {"action": "remove"}},
    )
    assert response.status_code == 422
    assert "шаблон" in response.json()["error"]["message"]
    both = client.post(
        f"/api/office/documents/{doc_id}/edit",
        json={
            "revision": 0,
            "instruction": "убери логотип",
            "logo": {"action": "remove"},
            "target": {"slide": 1, "shape_id": "2"},
        },
    )
    assert both.status_code == 422


@pytest.mark.parametrize("format,media", [("pdf", "application/pdf"), ("html", "text/html")])
def test_export_saved_revision(client, office, monkeypatch, format, media):
    store, doc_id = office
    original = store.read(doc_id, 0)

    def export(data, requested, settings):
        assert data == original
        assert requested == format
        return b"exported-revision"

    monkeypatch.setattr(onlyoffice, "export_revision", export)
    response = client.get(f"/api/office/documents/{doc_id}/download/0?format={format}")
    assert response.status_code == 200
    assert response.content == b"exported-revision"
    assert response.headers["content-type"].startswith(media)
    assert response.headers["content-disposition"].endswith(f'.{format}"')
    assert store.read(doc_id, 0) == original
    assert store.get(doc_id)["active_key"] is None
    missing = client.get(f"/api/office/documents/{doc_id}/download/99?format={format}")
    assert missing.status_code == 404


def test_export_failure_and_invalid_format(client, office, monkeypatch):
    _, doc_id = office

    def fail(*args):
        raise onlyoffice.ConversionError("unavailable")

    monkeypatch.setattr(onlyoffice, "export_revision", fail)
    url = f"/api/office/documents/{doc_id}/download/0"
    assert client.get(url + "?format=pdf").status_code == 503
    assert client.get(url + "?format=exe").status_code == 422
    assert client.get(url).status_code == 200


def test_template_viewer_is_signed_read_only_and_does_not_change_source(client, office, pptx_bytes):
    from urllib.parse import parse_qs

    uploaded = client.post("/api/templates", files={"file": ("Шаблон.pptx", pptx_bytes)})
    assert uploaded.status_code == 202, uploaded.text
    template_id = uploaded.json()["template_id"]
    before = client.get(f"/api/templates/{template_id}").json()
    response = client.post(f"/api/office/templates/{template_id}/config")
    assert response.status_code == 200, response.text
    config = response.json()["config"]
    signed = verify(config["token"], SECRET)
    assert signed == {k: v for k, v in config.items() if k != "token"}
    assert config["editorConfig"]["mode"] == "view"
    assert config["editorConfig"]["customization"] == {
        "anonymous": {"request": False},
        "compactHeader": True,
        "toolbarHideFileName": True,
        "uiTheme": "theme-white",
        "compactToolbar": True,
        "hideRightMenu": True,
        "hideRulers": True,
        "hideNotes": True,
        "comments": False,
        "features": {"featuresTips": False},
        "zoom": -1,
        "plugins": False,
        "macros": False,
        "help": False,
        "feedback": {"visible": False},
        "suggestFeature": False,
    }
    assert config["editorConfig"]["user"]["name"] == "Просмотр шаблона"
    assert "callbackUrl" not in config["editorConfig"]
    assert config["document"]["permissions"]["edit"] is False
    assert config["document"]["permissions"]["chat"] is False
    url = urlsplit(config["document"]["url"])
    assert client.get(url.path + "?" + url.query).content == pptx_bytes
    assert client.get(url.path, params={"token": "bad"}).status_code == 403
    claims = verify(parse_qs(url.query)["token"][0], SECRET)
    claims["scope"] = "office-file"
    assert client.get(url.path, params={"token": sign(claims, SECRET)}).status_code == 403
    claims["scope"] = "office-template"
    claims["sha256"] = "wrong"
    assert client.get(url.path, params={"token": sign(claims, SECRET)}).status_code == 403
    assert client.get(f"/api/templates/{template_id}/source").content == pptx_bytes
    assert client.get(f"/api/templates/{template_id}").json() == before
    assert client.post("/api/office/templates/missing/config").status_code == 404
    assert client.get("/api/templates/missing/source").status_code == 404
    # Download remains useful with ONLYOFFICE disabled.
    client.app.state.orchestrator.settings.onlyoffice.enabled = False
    assert client.get(f"/api/templates/{template_id}/source").content == pptx_bytes


def test_tokens():
    token = sign({"key": "hello", "exp": time.time() + 60}, SECRET)
    assert verify(token, SECRET)["key"] == "hello"
    for invalid in ("", "none", token + "x", sign({"exp": 1}, SECRET)):
        with pytest.raises(ApiError):
            verify(invalid, SECRET)
    with pytest.raises(ApiError):
        verify(token, "wrong secret")


def test_config_simplifies_ui_without_disabling_editing_or_saving(client, office):
    _, doc_id = office
    response = client.post(f"/api/office/documents/{doc_id}/config")
    assert response.status_code == 200, response.text
    config = response.json()["config"]
    editor = config["editorConfig"]
    assert editor["mode"] == "edit"
    assert config["document"]["permissions"] == {
        "edit": True,
        "download": False,
        "print": True,
        "chat": False,
        "comment": False,
    }
    assert editor["customization"] == {
        "forcesave": True,
        "autosave": True,
        "compactHeader": True,
        "toolbarHideFileName": True,
        "uiTheme": "theme-white",
        "compactToolbar": True,
        "hideRightMenu": True,
        "hideRulers": True,
        "hideNotes": True,
        "comments": False,
        "features": {"featuresTips": False},
        "zoom": -1,
        "plugins": False,
        "macros": False,
        "help": False,
        "feedback": {"visible": False},
        "suggestFeature": False,
    }
    # Keep all UI settings inside the signed configuration, not a browser-side override.
    assert verify(config["token"], SECRET)["editorConfig"] == editor


def test_config_source_and_session_key(client, office, pptx_bytes):
    store, doc_id = office
    url = f"/api/office/documents/{doc_id}"
    config = client.post(url + "/config").json()["config"]
    assert verify(config["token"], SECRET)["document"] == config["document"]
    assert config["documentType"] == "slide"
    customization = config["editorConfig"]["customization"]
    assert customization["compactHeader"] is True
    assert customization["toolbarHideFileName"] is True
    assert customization["forcesave"] is True
    assert verify(config["token"], SECRET)["editorConfig"] == config["editorConfig"]
    assert (
        config["document"]["key"]
        == client.post(url + "/config").json()["config"]["document"]["key"]
    )
    source = urlsplit(config["document"]["url"])
    assert client.get(source.path + "?" + source.query).content == pptx_bytes
    assert client.get(source.path + "?token=bad").status_code == 403
    assert client.get(source.path.replace("/0", "/1") + "?" + source.query).status_code == 403
    assert (
        client.post(
            url + "/callback", json={"key": config["document"]["key"], "status": 4}
        ).status_code
        == 403
    )
    claims = {"key": config["document"]["key"], "status": 4}
    assert client.post(url + "/callback", json={"token": sign(claims, SECRET)}).json() == {
        "error": 0
    }
    assert store.get(doc_id)["active_key"] is None
    assert client.post(url + "/config").json()["config"]["document"]["key"] != claims["key"]


@pytest.mark.parametrize("value", [None, True, "NaN", float("nan"), float("inf"), 10**400])
@pytest.mark.parametrize("claim", ["exp", "nbf"])
def test_invalid_token_time_claim(claim, value):
    with pytest.raises(ApiError) as error:
        verify(sign({claim: value}, SECRET), SECRET)
    assert error.value.status == 403


@pytest.mark.parametrize("token", [None, 123, [], {}, "bad.token.signature"])
def test_malformed_callback_token(client, office, token):
    _, doc_id = office
    assert (
        client.post(f"/api/office/documents/{doc_id}/callback", json={"token": token}).status_code
        == 403
    )


@pytest.mark.parametrize("value", [None, 123, [], {}])
def test_malformed_signed_save_url(client, office, value):
    store, doc_id = office
    key = store.open(doc_id)["active_key"]
    response = client.post(
        f"/api/office/documents/{doc_id}/callback",
        json={"token": sign({"key": key, "status": 2, "url": value}, SECRET)},
    )
    assert response.json() == {"error": 1}
    assert store.get(doc_id)["revision"] == 0
    assert store.get(doc_id)["error"]


def test_save_round_trip(client, office, monkeypatch, pptx_bytes):
    store, doc_id = office
    url = f"/api/office/documents/{doc_id}"
    config = client.post(url + "/config").json()["config"]
    key = config["document"]["key"]
    presentation = Presentation(io.BytesIO(pptx_bytes))
    presentation.slides[0].shapes.add_textbox(0, 0, 1000000, 1000000).text = "Manual office edit"
    buffer = io.BytesIO()
    presentation.save(buffer)
    edited = buffer.getvalue()
    real_client = httpx.AsyncClient
    requests = []

    def handle(request):
        requests.append(str(request.url))
        return httpx.Response(200, content=edited)

    monkeypatch.setattr(
        onlyoffice.httpx,
        "AsyncClient",
        lambda **kw: real_client(
            transport=httpx.MockTransport(handle),
            **kw,
        ),
    )
    claims = {"key": key, "status": 6, "url": "http://onlyoffice/cache/files/test/output.pptx"}
    # Unsigned siblings are ignored; only signed URL and status are used.
    body = {"token": sign(claims, SECRET), "url": "http://evil.invalid/", "status": 4}
    assert client.post(url + "/callback", json=body).json() == {"error": 0}
    assert requests == [claims["url"]]
    assert client.get(url + "/download/1").content == edited
    assert client.get(url + "/download/0").content == pptx_bytes
    assert client.post(url + "/callback", json=body).json() == {"error": 0}
    assert store.get(doc_id)["revision"] == 1  # idempotent repeated force-save
    # Force-save must NOT rotate a key while users are still connected.
    assert client.post(url + "/config").json()["config"]["document"]["key"] == key
    claims["status"] = 2
    assert client.post(url + "/callback", json={"token": sign(claims, SECRET)}).json() == {
        "error": 0
    }
    assert store.get(doc_id)["active_key"] is None
    reopened = client.post(url + "/config").json()["config"]
    source = urlsplit(reopened["document"]["url"])
    assert client.get(source.path + "?" + source.query).content == edited
    assert reopened["document"]["key"] != key
    assert client.post(url + "/callback", json={"token": sign(claims, SECRET)}).json() == {
        "error": 1
    }
    assert store.create("job_test/balanced/r1/deck.pptx", "test.pptx", pptx_bytes)["id"] == doc_id


@pytest.mark.parametrize("status", [2, 6])
def test_invalid_layout_callback_is_not_accepted(client, office, monkeypatch, pptx_bytes, status):
    from zipfile import ZipFile

    store, doc_id = office
    key = store.open(doc_id)["active_key"]
    output = io.BytesIO()
    with ZipFile(io.BytesIO(pptx_bytes)) as source, ZipFile(output, "w") as target:
        for info in source.infolist():
            target.writestr(
                info,
                b"<broken" if info.filename.endswith("slideLayout3.xml") else source.read(info),
            )
    real_client = httpx.AsyncClient
    monkeypatch.setattr(
        onlyoffice.httpx,
        "AsyncClient",
        lambda **kw: real_client(
            transport=httpx.MockTransport(
                lambda request: httpx.Response(200, content=output.getvalue())
            ),
            **kw,
        ),
    )
    url = f"/api/office/documents/{doc_id}/callback"
    body = {"key": key, "status": status, "url": "http://onlyoffice/cache/files/test"}
    assert client.post(url, json={"token": sign(body, SECRET)}).json() == {"error": 1}
    assert store.get(doc_id)["active_key"] == key
    assert store.get(doc_id)["revision"] == 0
    assert store.read(doc_id, 0) == pptx_bytes
    for notification in (1, 4):
        store.callback(doc_id, key, notification, None)
        assert store.get(doc_id)["error"]


def test_normalization_audit_is_atomic_and_idempotent(office, pptx_bytes):
    store, doc_id = office
    key = store.open(doc_id)["active_key"]
    for _ in range(2):
        assert store.callback(
            doc_id,
            key,
            6,
            pptx_bytes + b"normalized",
            raw_pptx=b"raw",
            normalized_parts=["layout.xml"],
        )
    doc = store.get(doc_id)
    assert doc["revision"] == 1
    assert len(doc["normalizations"]) == 1
    assert doc["normalizations"][0]["parts"] == ["layout.xml"]
    with store.connect() as db:
        assert bytes(db.execute("SELECT raw_pptx FROM save_normalizations").fetchone()[0]) == b"raw"
    assert store.read(doc_id, 0) == pptx_bytes


def test_callback_normalizes_known_layout_name_and_retains_raw(
    client, office, monkeypatch, pptx_bytes
):
    from zipfile import ZipFile

    from lxml import etree

    store, _ = office
    part = "ppt/slideLayouts/slideLayout3.xml"
    output = io.BytesIO()
    with ZipFile(io.BytesIO(pptx_bytes)) as source, ZipFile(output, "w") as target:
        for info in source.infolist():
            xml = source.read(info)
            if info.filename == part:
                root = etree.fromstring(xml)
                root.set("matchingName", 'Слайд "Спасибо!"')
                xml = etree.tostring(root, encoding="UTF-8")
            target.writestr(info, xml)
    original = output.getvalue()
    doc_id = store.create("test/quoted-layout", "test.pptx", original)["id"]
    key = store.open(doc_id)["active_key"]
    output = io.BytesIO()
    with ZipFile(io.BytesIO(original)) as source, ZipFile(output, "w") as target:
        for info in source.infolist():
            xml = source.read(info)
            if info.filename == part:
                xml = xml.replace(b"&quot;", b'"')
            target.writestr(info, xml)
    raw = output.getvalue()
    real_client = httpx.AsyncClient
    monkeypatch.setattr(
        onlyoffice.httpx,
        "AsyncClient",
        lambda **kw: real_client(
            transport=httpx.MockTransport(lambda request: httpx.Response(200, content=raw)),
            **kw,
        ),
    )
    response = client.post(
        f"/api/office/documents/{doc_id}/callback",
        json={
            "token": sign(
                {"key": key, "status": 2, "url": "http://onlyoffice/cache/files/test"},
                SECRET,
            )
        },
    )
    assert response.json() == {"error": 0}
    doc = store.get(doc_id)
    assert doc["active_key"] is None
    assert doc["error"] is None
    assert doc["normalizations"][0]["parts"] == [part]
    saved = store.read(doc_id, doc["revision"])
    with ZipFile(io.BytesIO(saved)) as archive:
        assert etree.fromstring(archive.read(part)).get("matchingName") == 'Слайд "Спасибо!"'
    assert store.read(doc_id, 0) == original
    with store.connect() as db:
        assert (
            bytes(
                db.execute(
                    "SELECT raw_pptx FROM save_normalizations WHERE document_id=?", (doc_id,)
                ).fetchone()[0]
            )
            == raw
        )


@pytest.mark.parametrize(
    "url",
    [
        "http://169.254.169.254/cache/files/x",
        "http://evil.invalid/cache/files/x",
        "http://onlyoffice/healthcheck",
        "http://onlyoffice/cache/files/../x",
        "http://onlyoffice/cache/files/%2e%2e/x",
        "http://user@onlyoffice/cache/files/x",
        "http://onlyoffice:8000/cache/files/x",
    ],
)
def test_reject_save_ssrf(url):
    with pytest.raises(ValueError):
        onlyoffice.saved_url(url, "http://onlyoffice", "http://localhost:8080/onlyoffice")


def test_public_save_url():
    assert (
        onlyoffice.saved_url(
            "http://localhost:8080/onlyoffice/cache/files/key/output.pptx?token=abc",
            "http://onlyoffice",
            "http://localhost:8080/onlyoffice",
        )
        == "http://onlyoffice/cache/files/key/output.pptx?token=abc"
    )


def test_failed_save_does_not_ack_or_replace(client, office, monkeypatch, pptx_bytes):
    store, doc_id = office
    doc = store.open(doc_id)
    real_client = httpx.AsyncClient
    monkeypatch.setattr(
        onlyoffice.httpx,
        "AsyncClient",
        lambda **kw: real_client(
            transport=httpx.MockTransport(
                lambda request: httpx.Response(200, content=b"not a pptx")
            ),
            **kw,
        ),
    )
    body = {
        "token": sign(
            {"key": doc["active_key"], "status": 2, "url": "http://onlyoffice/cache/files/x"},
            SECRET,
        )
    }
    assert client.post(f"/api/office/documents/{doc_id}/callback", json=body).json() == {"error": 1}
    assert store.get(doc_id)["error"]
    assert store.read(doc_id, 0) == pptx_bytes
    assert store.get(doc_id)["revision"] == 0


def test_concurrent_creation_and_keys(office, pptx_bytes):
    store, doc_id = office

    def open_copy(_):
        doc = store.create("job_test/balanced/r1/deck.pptx", "test.pptx", pptx_bytes)
        return store.open(doc["id"])["active_key"]

    with ThreadPoolExecutor(max_workers=4) as pool:
        keys = list(pool.map(open_copy, range(8)))
    assert len(set(keys)) == 1
    assert store.get(doc_id)["revision"] == 0


def test_insert_image_signs_command_for_project_file_and_template_asset(client, office, tmp_path):
    from PIL import Image

    from tests.fixtures.rich_template import build_rich_template

    _, doc_id = office
    url = f"/api/office/documents/{doc_id}/images"
    project = client.post("/api/projects", json={}).json()["project_id"]
    png = io.BytesIO()
    Image.new("RGB", (40, 20), (200, 30, 30)).save(png, format="PNG")
    photo, notes = client.post(
        f"/api/projects/{project}/files",
        files=[
            ("files", ("фото.png", png.getvalue(), "image/png")),
            ("files", ("notes.txt", b"text", "text/plain")),
        ],
    ).json()

    response = client.post(url, json={"project_id": project, "file_id": photo["file_id"]})
    assert response.status_code == 200, response.text
    command = response.json()
    image_url = f"http://api:8000/api/projects/{project}/files/{photo['file_id']}/content"
    assert command["c"] == "add"
    assert command["images"] == [{"fileType": "png", "url": image_url}]
    # Document Server проверяет подпись браузерной команды тем же секретом.
    assert verify(command["token"], SECRET) == {"c": "add", "images": command["images"]}
    # Ссылка для Document Server отдаёт ровно эти байты.
    assert client.get(urlsplit(image_url).path).content == png.getvalue()

    def error(body):
        r = client.post(url, json=body)
        return r.status_code, r.json()["error"]["code"]

    assert error({"project_id": project, "file_id": notes["file_id"]}) == (422, "file_not_image")
    other = client.post("/api/projects", json={}).json()["project_id"]
    assert error({"project_id": other, "file_id": photo["file_id"]}) == (404, "file_not_found")
    assert client.post(url, json={"file_id": photo["file_id"]}).status_code == 422
    assert (
        client.post(
            "/api/office/documents/missing/images",
            json={"project_id": project, "file_id": photo["file_id"]},
        ).status_code
        == 404
    )

    rich = build_rich_template(tmp_path / "rich.pptx").read_bytes()
    template_id = client.post("/api/templates", files={"file": ("Шаблон.pptx", rich)}).json()[
        "template_id"
    ]
    assets = client.get(f"/api/templates/{template_id}").json()["profile"]["assets"]
    asset = next(a for a in assets if a["media_path"].startswith("ppt/media/"))
    command = client.post(url, json={"template_id": template_id, "asset_id": asset["asset_id"]})
    assert command.status_code == 200, command.text
    image = command.json()["images"][0]
    assert image["url"] == f"http://api:8000/api/templates/{template_id}/media/{asset['asset_id']}"
    assert client.get(urlsplit(image["url"]).path).status_code == 200
    assert error({"template_id": template_id, "asset_id": "asset_nope"}) == (
        404,
        "asset_not_found",
    )


def _saved(store, doc_id, deck) -> bytes:
    """Сохранение из ONLYOFFICE: колбэк открытой сессии кладёт следующую ревизию."""
    output = io.BytesIO()
    deck.save(output)
    opened = store.open(doc_id)
    store.callback(doc_id, opened["active_key"], 2, output.getvalue())
    return output.getvalue()


def test_live_selection_edit_finds_the_object_by_name_and_position(client, office, monkeypatch):
    from pptx.util import Mm

    from presentation_designer.generation.office_object_edit import ObjectEditPlan, Position

    store, doc_id = office
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    for x, text in [(20, "Левая"), (120, "Правая")]:
        box = slide.shapes.add_textbox(Mm(x), Mm(30), Mm(60), Mm(20))
        box.text = text
        box.name = "Карточка"  # у клонов образца имена часто совпадают
    data = _saved(store, doc_id, deck)
    objects = client.get(f"/api/office/documents/{doc_id}/objects/1").json()["objects"]
    right = next(o for o in objects if o["label"] == "Правая")
    assert right["name"] == "Карточка"

    async def propose(original, instruction, settings, selected, tokens=None):
        assert original == data
        assert selected.model_dump() == {"slide": 1, "shape_id": right["shape_id"]}
        return ObjectEditPlan(
            explanation="Сдвинул", patches=[], position=Position(x=0.5, y=right["bbox"]["y"])
        )

    monkeypatch.setattr(onlyoffice.office_object_edit, "propose", propose)
    base = f"/api/office/documents/{doc_id}/edit"
    live = {
        "slide": 1,
        "name": "Карточка",
        "box": {"x": 119.5, "y": 30.2, "width": 60, "height": 20},
    }
    response = client.post(base, json={"revision": 1, "instruction": "правее", "live_target": live})
    assert response.status_code == 200, response.text
    assert response.json()["document"]["revision"] == 2
    missing = client.post(
        base,
        json={"revision": 2, "instruction": "правее", "live_target": live | {"name": "Нет такой"}},
    )
    assert (
        missing.status_code == 422 and "выделите его заново" in missing.json()["error"]["message"]
    )
    both = client.post(
        base,
        json={
            "revision": 2,
            "instruction": "x",
            "live_target": live,
            "target": {"slide": 1, "shape_id": "2"},
        },
    )
    assert both.status_code == 422


def test_apply_slide_puts_the_rebuilt_slide_into_the_edited_copy(
    client, office, orchestrator, monkeypatch, pptx_bytes, tmp_path
):
    store, doc_id = office
    edited = Presentation(io.BytesIO(pptx_bytes))
    edited.slides[0].shapes.title.text = "Ручная правка"
    _saved(store, doc_id, edited)
    rebuilt = Presentation(io.BytesIO(pptx_bytes))
    rebuilt.slides[1].shapes[0].text_frame.text = "Пересобранный слайд"
    artifact = tmp_path / "deck.pptx"
    rebuilt.save(artifact)
    monkeypatch.setattr(
        onlyoffice,
        "build_generation_result",
        lambda state, job: {"artifacts_manifest": {"balanced/r2/deck.pptx": {}}},
    )
    monkeypatch.setattr(orchestrator.artifacts, "resolve", lambda job, name: artifact)
    url = f"/api/office/documents/{doc_id}/apply-slide"
    body = {
        "revision": 1,
        "job_id": "job_test",
        "variant_id": "balanced",
        "artifact_revision": 2,
        "slide": 2,
        "source_slide": 2,
    }
    response = client.post(url, json=body)
    assert response.status_code == 200, response.text
    assert response.json()["document"]["revision"] == 2
    saved = Presentation(io.BytesIO(store.read(doc_id, 2)))
    assert saved.slides[0].shapes.title.text == "Ручная правка"
    assert saved.slides[1].shapes[0].text_frame.text == "Пересобранный слайд"
    assert len(saved.slides) == len(rebuilt.slides)
    assert client.post(url, json=body).status_code == 409  # ревизия копии уже 2
    assert client.post(url, json=body | {"revision": 2, "job_id": "job_other"}).status_code == 409
    assert client.post(url, json=body | {"revision": 2, "artifact_revision": 3}).status_code == 404
    assert client.post(url, json=body | {"revision": 2, "slide": 9}).status_code == 422
    assert store.get(doc_id)["active_key"] is None


def _group_and_placeholder_deck() -> Presentation:
    from pptx.util import Mm

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])  # заголовок с рамкой из макета
    slide.shapes.title.text = "Команда проекта"
    group = slide.shapes.add_group_shape()
    for x, text in [(20, "Имя Фамилия"), (120, "Должность")]:
        box = group.shapes.add_textbox(Mm(x), Mm(90), Mm(60), Mm(20))
        box.text = text
        box.name = f"Google Shape;{x};p54"
    return deck


@pytest.mark.parametrize(
    ("live", "before"),
    [
        # Фигура в группе: ONLYOFFICE отдаёт её имя и рамку от угла группы.
        (
            {
                "name": "Google Shape;120;p54",
                "in_group": True,
                "box": {"x": 100, "y": 0, "width": 60, "height": 20},
            },
            "Должность",
        ),
        (
            {"name": "Title 1", "box": {"x": 12.7, "y": 7.6, "width": 228.6, "height": 31.75}},
            "Команда проекта",
        ),
    ],
)
def test_live_selection_edit_reaches_group_child_and_placeholder(
    client, office, monkeypatch, live, before
):
    from presentation_designer.generation.office_edit import TextPatch
    from presentation_designer.generation.office_object_edit import ObjectEditPlan

    store, doc_id = office
    _saved(store, doc_id, _group_and_placeholder_deck())
    objects = client.get(f"/api/office/documents/{doc_id}/objects/1").json()["objects"]
    wanted = next(o for o in objects if o["label"] == before)

    async def propose(original, instruction, settings, selected, tokens=None):
        assert selected.model_dump() == {"slide": 1, "shape_id": wanted["shape_id"]}
        return ObjectEditPlan(
            explanation="Готово",
            patches=[
                TextPatch(slide=1, run=wanted["runs"][0], before=before, after=before + " (новое)")
            ],
        )

    monkeypatch.setattr(onlyoffice.office_object_edit, "propose", propose)
    response = client.post(
        f"/api/office/documents/{doc_id}/edit",
        json={"revision": 1, "instruction": "допиши", "live_target": {"slide": 1, **live}},
    )
    assert response.status_code == 200, response.text
    after = client.get(f"/api/office/documents/{doc_id}/objects/2").json()["objects"]
    assert any(o["label"] == before + " (новое)" for o in after)


def test_snapshot_of_a_saved_copy_marks_manually_edited_slides(client, office):
    store, doc_id = office
    original = store.read(doc_id, 0)
    deck = Presentation(io.BytesIO(original))
    first = next(s for s in deck.slides[0].shapes if s.has_text_frame and s.text_frame.text)
    first.text_frame.text = "Ручная правка"
    _saved(store, doc_id, deck)
    seed = client.get(f"/api/office/documents/{doc_id}/snapshot/0").json()
    saved = client.get(f"/api/office/documents/{doc_id}/snapshot/1").json()
    assert seed["source"] == {
        "kind": "office",
        "document_id": doc_id,
        "revision": 0,
        "pptx_sha256": seed["source"]["pptx_sha256"],
    }
    assert "edited" not in seed["outline"][0]
    assert [e["edited"] for e in saved["outline"]] == [True] + [False] * (len(saved["outline"]) - 1)
    assert any(
        (o.get("text") or {}).get("plain") == "Ручная правка" for o in saved["slides"][0]["objects"]
    )
    missing = client.get(f"/api/office/documents/{doc_id}/snapshot/9")
    assert missing.status_code == 404


def test_picture_from_the_message_goes_to_the_named_place(client, office, monkeypatch):
    """«Вставь картинку в правый блок» с вложением: место называет модель, картинка встаёт
    новой ревизией копии; без снимка слайда (рендер недоступен) — по рамке модели."""
    from PIL import Image

    from presentation_designer.export.pdf import ConversionError
    from presentation_designer.generation.office_image import Area, Placement

    store, doc_id = office
    project = client.post("/api/projects", json={}).json()["project_id"]
    png = io.BytesIO()
    Image.new("RGB", (400, 300), (30, 120, 200)).save(png, format="PNG")
    (photo,) = client.post(
        f"/api/projects/{project}/files",
        files=[("files", ("фото.png", png.getvalue(), "image/png"))],
    ).json()
    asked: list[tuple[str, int, object]] = []

    async def propose(data, instruction, settings, slide, snapshot):
        asked.append((instruction, slide, snapshot))
        return Placement(
            explanation="Поставил в блок справа.", area=Area(x=0.55, y=0.2, width=0.4, height=0.6)
        )

    def no_render(*args):
        raise ConversionError("рендер недоступен в тестах")

    monkeypatch.setattr(onlyoffice.office_image, "propose", propose)
    monkeypatch.setattr(onlyoffice, "preview_revision", no_render)
    response = client.post(
        f"/api/office/documents/{doc_id}/edit",
        json={
            "revision": 0,
            "instruction": "вставь картинку в правый блок",
            "image": {"file_id": photo["file_id"], "slide": 1},
        },
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["changed"] is True and body["message"] == "Поставил в блок справа."
    assert asked == [("вставь картинку в правый блок", 1, None)]
    prs = Presentation(io.BytesIO(store.read(doc_id, 1)))
    pictures = [s for s in prs.slides[0].shapes if s.shape_type == 13]
    assert [p.name for p in pictures][-1] == "Картинка из чата"
    assert pictures[-1].left == round(0.55 * prs.slide_width)

    # Картинка ставится по словам или в выделенный объект, не вместе с логотипом.
    both = client.post(
        f"/api/office/documents/{doc_id}/edit",
        json={
            "revision": 1,
            "instruction": "замени логотип",
            "logo": {"action": "remove"},
            "image": {"file_id": photo["file_id"], "slide": 1},
        },
    )
    assert both.status_code == 422


def test_table_from_an_attached_file_goes_to_the_named_place(client, office, monkeypatch):
    """«Вот таблица — на слайд справа» с xlsx: место называет модель, таблица — в стиле
    шаблона новой ревизией копии."""
    from openpyxl import Workbook

    from presentation_designer.export.pdf import ConversionError
    from presentation_designer.generation.office_image import Area, Placement

    store, doc_id = office
    project = client.post("/api/projects", json={}).json()["project_id"]
    book = Workbook()
    for row in [["Год", "Выручка"], [2023, 75], [2024, 120]]:
        book.active.append(row)
    buf = io.BytesIO()
    book.save(buf)
    (sheet,) = client.post(
        f"/api/projects/{project}/files",
        files=[("files", ("Выручка.xlsx", buf.getvalue(), "application/octet-stream"))],
    ).json()

    async def propose(data, instruction, settings, slide, snapshot, what="картинку"):
        assert what == "таблицу"
        return Placement(explanation="Справа.", area=Area(x=0.5, y=0.2, width=0.45, height=0.5))

    def no_render(*args):
        raise ConversionError("рендер недоступен в тестах")

    monkeypatch.setattr(onlyoffice.office_image, "propose", propose)
    monkeypatch.setattr(onlyoffice, "preview_revision", no_render)
    response = client.post(
        f"/api/office/documents/{doc_id}/edit",
        json={
            "revision": 0,
            "instruction": "поставь справа",
            "table": {"file_id": sheet["file_id"], "slide": 1},
        },
    )
    assert response.status_code == 200, response.text
    assert "2 строк, 2 столбцов" in response.json()["message"]
    prs = Presentation(io.BytesIO(store.read(doc_id, 1)))
    frame = next(s for s in prs.slides[0].shapes if s.name == "Таблица из чата")
    assert frame.table.cell(2, 0).text == "2024"
