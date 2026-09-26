"""Знак шаблона в офисной копии: замена и удаление на всех слайдах одним действием, по
картинкам из профиля; чужие логотипы на слайдах не трогаются."""

from __future__ import annotations

import hashlib
import io

import pytest
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
from pptx.oxml.shapes.picture import CT_Picture
from pptx.util import Emu, Inches

from presentation_designer.generation.office_logo import (
    file_logo_hashes,
    logo_hashes,
    promote_file,
    replace_logo,
)
from presentation_designer.layout.ooxml import check_package


def _png(size: tuple[int, int], color: tuple[int, int, int], *, opaque: bool = False) -> bytes:
    """Знак — цветная фигура на прозрачном поле, как у настоящих логотипов; `opaque` —
    непрозрачный прямоугольник, как декор макетов VK WorkSpace."""
    buf = io.BytesIO()
    if opaque:
        Image.new("RGB", size, color).save(buf, "PNG")
        return buf.getvalue()
    image = Image.new("RGBA", size, (0, 0, 0, 0))
    w, h = size
    ImageDraw.Draw(image).rectangle((w // 8, h // 8, w - w // 8, h - h // 8), fill=(*color, 255))
    image.save(buf, "PNG")
    return buf.getvalue()


BRAND = _png((400, 100), (0, 119, 255))
PARTNER = _png((200, 200), (255, 56, 133))
NEW = _png((300, 150), (255, 90, 0))


def _deck() -> bytes:
    """Знак шаблона слева сверху на макете, логотип партнёра — картинкой на слайде."""
    prs = Presentation()
    layout = prs.slide_layouts[6]
    _, rid = layout.part.get_or_add_image_part(io.BytesIO(BRAND))
    pic = CT_Picture.new_pic(90, "Logo", "", rid, Inches(0.5), Inches(0.4), Inches(2), Inches(0.5))
    layout.shapes._spTree.append(pic)
    slide = prs.slides.add_slide(layout)
    slide.shapes.add_picture(io.BytesIO(PARTNER), Inches(6), Inches(3), Inches(1), Inches(1))
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def _pictures(data: bytes) -> dict[str, list[tuple[str, Emu, Emu, Emu, Emu]]]:
    prs = Presentation(io.BytesIO(data))
    out: dict[str, list[tuple[str, Emu, Emu, Emu, Emu]]] = {"layout": [], "slide": []}
    for where, shapes in (("layout", prs.slide_layouts[6].shapes), ("slide", prs.slides[0].shapes)):
        for sh in shapes:
            if sh.shape_type == 13:
                sha = hashlib.sha256(sh.image.blob).hexdigest()
                out[where].append((sha, sh.left, sh.top, sh.width, sh.height))
    return out


def sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def test_brand_mark_is_replaced_on_the_layout_and_keeps_its_place() -> None:
    data, n = replace_logo(_deck(), {sha(BRAND)}, NEW)
    assert n == 1
    pics = _pictures(data)
    ((new_sha, left, top, width, height),) = pics["layout"]
    assert new_sha == sha(NEW)
    # Та же высота и левый край, ширина — по пропорции новой картинки (2:1).
    assert (left, top, height) == (Inches(0.5), Inches(0.4), Inches(0.5))
    assert width == Inches(1)
    assert [p[0] for p in pics["slide"]] == [sha(PARTNER)]
    assert replace_logo(data, {sha(BRAND)}, NEW)[1] == 0


def test_brand_mark_is_removed_and_partner_logos_stay() -> None:
    data, n = replace_logo(_deck(), {sha(BRAND)}, None)
    assert n == 1
    pics = _pictures(data)
    assert pics["layout"] == []
    assert [p[0] for p in pics["slide"]] == [sha(PARTNER)]


def test_only_pinned_template_logos_count_and_vector_images_are_refused() -> None:
    profile = {
        "fixed_elements": [{"kind": "logo", "asset_id": "a1"}],
        "assets": [
            {"asset_id": "a1", "kind": "logo", "sha256": "brand"},
            {"asset_id": "a2", "kind": "logo", "sha256": "partner"},
        ],
    }
    assert logo_hashes(profile) == {"brand"}
    assert logo_hashes(None) == set()
    with pytest.raises(ValueError, match="PNG или JPG"):
        replace_logo(_deck(), {sha(BRAND)}, b"<svg xmlns='http://www.w3.org/2000/svg'/>")


def test_mark_of_two_pictures_becomes_one_logo_in_their_frame() -> None:
    # Как у VK Tech: квадрат VK и надпись tech — две картинки вплотную на одном макете.
    icon, word = _png((100, 100), (0, 119, 255)), _png((300, 100), (20, 20, 20))
    prs = Presentation()
    layout = prs.slide_layouts[6]
    for n, (data, left, width) in enumerate(((icon, 0.5, 0.5), (word, 1.05, 1.5))):
        _, rid = layout.part.get_or_add_image_part(io.BytesIO(data))
        layout.shapes._spTree.append(
            CT_Picture.new_pic(
                80 + n, f"Logo {n}", "", rid, Inches(left), Inches(0.4), Inches(width), Inches(0.5)
            )
        )
    prs.slides.add_slide(layout)
    buf = io.BytesIO()
    prs.save(buf)
    data, n = replace_logo(buf.getvalue(), {sha(icon), sha(word)}, NEW)
    assert n == 1
    ((new_sha, left, top, width, height),) = _pictures(data)["layout"]
    assert new_sha == sha(NEW)
    # Общая рамка двух частей: левый край первой, высота прежняя, ширина — по пропорции.
    assert (left, top, height) == (Inches(0.5), Inches(0.4), Inches(0.5))
    assert width == Inches(1)


def _layout_pic(layout, data: bytes, shape_id: int, left: float, top: float, width: float):
    _, rid = layout.part.get_or_add_image_part(io.BytesIO(data))
    pic = CT_Picture.new_pic(
        shape_id, f"Logo {shape_id}", "", rid, Inches(left), Inches(top), Inches(width), Inches(0.5)
    )
    layout.shapes._spTree.append(pic)
    return pic


def _rect(layout, shape_id: int, left: float, fill: str) -> None:
    """Прямоугольник 3×1 дюйма поверх полосы знаков внизу макета."""
    layout.shapes._spTree.append(
        parse_xml(
            f'<p:sp {nsdecls("p", "a")}><p:nvSpPr><p:cNvPr id="{shape_id}" name="Rect"/>'
            "<p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr>"
            f'<a:xfrm><a:off x="{Inches(left)}" y="{Inches(6.6)}"/>'
            f'<a:ext cx="{Inches(3)}" cy="{Inches(1)}"/></a:xfrm>'
            f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>{fill}</p:spPr></p:sp>'
        )
    )


def _saved(prs) -> bytes:
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_decor_with_logo_marks_is_not_replaced() -> None:
    # Непрозрачная вырезка фона (точки VK WorkSpace) по месту и размеру похожа на знак.
    decor = _png((400, 100), (10, 10, 30), opaque=True)
    prs = Presentation()
    layout = prs.slide_layouts[6]
    _layout_pic(layout, BRAND, 90, 0.5, 0.4, 2)
    _layout_pic(layout, decor, 91, 6, 0.4, 2)
    prs.slides.add_slide(layout)
    data, n = replace_logo(_saved(prs), {sha(BRAND), sha(decor)}, None)
    assert n == 1
    assert [p[0] for p in _pictures(data)["layout"]] == [sha(decor)]


def test_template_mark_moves_from_the_layout_onto_every_slide(tmp_path) -> None:
    icon = _png((100, 100), (0, 119, 255))
    decor = _png((400, 100), (10, 10, 30), opaque=True)
    prs = Presentation()
    layout = prs.slide_layouts[6]
    _layout_pic(layout, BRAND, 90, 0.5, 6.8, 2)
    # Google Slides кладёт знак дважды в одну рамку: на слайд попадает один.
    _layout_pic(layout, BRAND, 91, 0.5, 6.8, 2)
    _layout_pic(layout, icon, 92, 8.5, 6.8, 0.5)
    _layout_pic(layout, decor, 93, 5, 0.2, 2)
    for _ in range(2):
        prs.slides.add_slide(layout)
    hidden = prs.slides.add_slide(layout)
    hidden._element.set("showMasterSp", "0")
    path = tmp_path / "deck.pptx"
    prs.save(str(path))

    assert file_logo_hashes(path) >= {sha(BRAND), sha(icon)}
    data, placed = promote_file(path)
    assert data is not None and placed == 4
    out = Presentation(io.BytesIO(data))
    for slide in list(out.slides)[:2]:
        pics = [(sha(p.image.blob), p.left, p.top, p.width) for p in slide.shapes]
        assert pics == [
            (sha(BRAND), Inches(0.5), Inches(6.8), Inches(2)),
            (sha(icon), Inches(8.5), Inches(6.8), Inches(0.5)),
        ]
    # Слайд со скрытым оформлением макета знака не видел и не получает.
    assert list(out.slides[2].shapes) == []
    # С макета знак снят, декор остался.
    assert [sha(p.image.blob) for p in out.slide_layouts[6].shapes if p.shape_type == 13] == [
        sha(decor)
    ]
    (tmp_path / "out.pptx").write_bytes(data)
    assert check_package(tmp_path / "out.pptx").errors == []
    # Повторный проход ничего не находит: знак уже на слайдах.
    assert promote_file(tmp_path / "out.pptx") == (None, 0)


def test_mark_hidden_under_a_filled_layout_shape_stays_put(tmp_path) -> None:
    # VK WorkSpace прячет тёмный знак чёрной плашкой; пустая рамка без заливки знак не прячет.
    prs = Presentation()
    layout = prs.slide_layouts[6]
    dark = _png((400, 100), (20, 20, 20))
    _layout_pic(layout, dark, 90, 0.5, 6.8, 2)
    _layout_pic(layout, BRAND, 91, 7, 6.8, 2)
    _rect(layout, 92, 0.3, '<a:solidFill><a:srgbClr val="000000"/></a:solidFill>')
    _rect(layout, 93, 6.8, "<a:noFill/>")
    prs.slides.add_slide(layout)
    path = tmp_path / "deck.pptx"
    prs.save(str(path))
    data, placed = promote_file(path)
    assert data is not None and placed == 1
    out = Presentation(io.BytesIO(data))
    assert [sha(p.image.blob) for p in out.slides[0].shapes] == [sha(BRAND)]
    layout_pics = [p for p in out.slide_layouts[6].shapes if p.shape_type == 13]
    assert [sha(p.image.blob) for p in layout_pics] == [sha(dark)]
