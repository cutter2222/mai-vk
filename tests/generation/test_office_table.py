"""Таблицы в копии (этап 39): ячейка, строки и столбцы с прежней шириной, сортировка,
выделение, выравнивание; таблица из xlsx и csv в стиле шаблона."""

from __future__ import annotations

import io

import pytest
from openpyxl import Workbook
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches

from presentation_designer.generation.office_image import Area, Placement
from presentation_designer.generation.office_objects import objects
from presentation_designer.generation.office_ops import ObjectOp, Tokens, apply
from presentation_designer.generation.office_table import place_table, read_table
from presentation_designer.layout.tables import TableStyle

TOKENS = Tokens(palette=["#0077FF", "#1A1A1A"], scale=[12.0, 16.0], fonts=["Arial"])
ROWS = [["Тариф", "Цена"], ["Базовый", "990"], ["Школа", "29 000"], ["Семья", "4 900"]]


def _deck() -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(16), Inches(9)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    frame = slide.shapes.add_table(4, 2, Inches(1), Inches(1), Inches(10), Inches(4))
    frame.name = "Тарифы"
    for r, row in enumerate(ROWS):
        for c, value in enumerate(row):
            frame.table.cell(r, c).text = value
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def _table(data: bytes):
    frame = next(s for s in Presentation(io.BytesIO(data)).slides[0].shapes if s.has_table)
    return frame, [[c.text for c in row.cells] for row in frame.table.rows]


def _run(data: bytes, *ops: ObjectOp) -> tuple[bytes, list[str]]:
    obj = next(o for o in objects(data) if o.name == "Тарифы")
    return apply(data, obj, list(ops), TOKENS, None)


def test_cells_rows_and_sorting() -> None:
    data = _deck()
    out, _ = _run(
        data,
        ObjectOp(op="table.set_cell", row=2, column=2, text="1 490"),
        ObjectOp(op="table.add_row", row=4, cells=["Премиум", "49 000"]),
        ObjectOp(op="table.delete_row", row=4),
        ObjectOp(op="table.sort", column=2, descending=True),
    )
    _, cells = _table(out)
    assert cells == [
        ["Тариф", "Цена"],
        ["Премиум", "49 000"],
        ["Школа", "29 000"],
        ["Базовый", "1 490"],
    ]
    with pytest.raises(ValueError, match="строки 9 нет"):
        _run(data, ObjectOp(op="table.delete_row", row=9))


def test_columns_keep_the_table_width() -> None:
    data = _deck()
    width = sum(c.width for c in _table(data)[0].table.columns)
    out, _ = _run(
        data, ObjectOp(op="table.add_column", column=2, cells=["Скидка", "5%", "10%", "0%"])
    )
    frame, cells = _table(out)
    assert cells[0] == ["Тариф", "Цена", "Скидка"] and cells[2][2] == "10%"
    assert abs(sum(c.width for c in frame.table.columns) - width) <= 2
    out, _ = _run(out, ObjectOp(op="table.delete_column", column=1))
    assert _table(out)[1][0] == ["Цена", "Скидка"]


def test_highlight_and_align_use_template_style() -> None:
    data = _deck()
    out, notes = _run(
        data,
        ObjectOp(op="table.highlight", row=3),
        ObjectOp(op="table.align", column=2, align="right"),
    )
    frame, _ = _table(out)
    run = frame.table.cell(2, 0).text_frame.paragraphs[0].runs[0]
    assert run.font.bold and str(run.font.color.rgb) == "0077FF"
    assert frame.table.cell(1, 1).text_frame.paragraphs[0].alignment == 3  # по правому краю
    assert "акцентом" in notes[0]


def test_read_xlsx_and_csv() -> None:
    book = Workbook()
    sheet = book.active
    for row in [["Год", "Выручка"], [2023, 75.0], [2024, 120.5], [None, None]]:
        sheet.append(row)
    buf = io.BytesIO()
    book.save(buf)
    header, rows, truncated = read_table(buf.getvalue(), "Отчёт.xlsx")
    assert header == ["Год", "Выручка"] and rows[0][0] == "2023" and len(rows) == 2
    assert not truncated
    header, rows, _ = read_table("Тариф;Цена\nБазовый;990\n".encode(), "t.csv")
    assert header == ["Тариф", "Цена"] and rows == [["Базовый", "990"]]
    with pytest.raises(ValueError, match="xlsx или csv"):
        read_table(b"", "t.pdf")


def test_table_from_file_is_placed_in_the_named_area() -> None:
    data = _deck()
    placement = Placement(explanation="справа", area=Area(x=0.7, y=0.2, width=0.27, height=0.5))
    out = place_table(
        data,
        1,
        ["Год", "Выручка"],
        [["2023", "75"], ["2024", "120,5"]],
        placement,
        TableStyle(font_family="Arial"),
    )
    frame = next(
        s for s in Presentation(io.BytesIO(out)).slides[0].shapes if s.name == "Таблица из чата"
    )
    assert frame.left == round(0.7 * Inches(16)) and frame.table.cell(0, 1).text == "Выручка"


def test_table_from_file_keeps_rows_compact_and_numbers_right() -> None:
    data = _deck()
    placement = Placement(explanation="справа", area=Area(x=0.7, y=0.2, width=0.27, height=0.7))
    out = place_table(
        data,
        1,
        ["Год", "Учеников"],
        [["2023", "4800"], ["2024", "15 200"]],
        placement,
        TableStyle(font_family="Arial"),
    )
    prs = Presentation(io.BytesIO(out))
    frame = next(s for s in prs.slides[0].shapes if s.name == "Таблица из чата")
    area_top, area_h = 0.2 * prs.slide_height, 0.7 * prs.slide_height
    # Три строки не растягиваются на всё место и стоят по его центру.
    assert frame.height < area_h / 2
    assert abs(frame.top + frame.height / 2 - (area_top + area_h / 2)) < Inches(0.05)
    cell = frame.table.cell(2, 1)
    assert cell.text_frame.paragraphs[0].alignment == PP_ALIGN.RIGHT
    assert cell.vertical_anchor == MSO_ANCHOR.MIDDLE
