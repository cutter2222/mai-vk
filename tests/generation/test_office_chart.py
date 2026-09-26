"""Диаграммы в копии (этап 40): данные, тип, оформление, выделение, сортировка и смена подачи
на месте; числа только из просьбы или из диаграммы."""

from __future__ import annotations

import io

import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Inches

from presentation_designer.generation.office_chart import (
    apply_chart_ops,
    check_numbers,
    kind_of,
    make_editable,
    place_chart,
    read_data,
    spec_from_table,
)
from presentation_designer.generation.office_object_edit import ObjectEditPlan, apply_plan
from presentation_designer.generation.office_objects import ObjectTarget, objects
from presentation_designer.generation.office_ops import ChartSeries, ObjectOp, Tokens

TOKENS = Tokens(palette=["#FFFFFF", "#0077FF", "#FF3985", "#1A1A1A"], fonts=["Arial"])


def _deck(*, table: bool = False) -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(16), Inches(9)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1))
    title.text_frame.text = "Выручка"
    data = CategoryChartData()
    data.categories = ["2023", "2024", "2025"]
    data.add_series("Выручка", (75, 100, 140))
    frame = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(8), Inches(5), data
    )
    frame.name = "Диаграмма выручки"
    if table:
        t = slide.shapes.add_table(3, 3, Inches(10), Inches(2), Inches(5), Inches(3))
        t.name = "Тарифы"
        for r, row in enumerate([["Год", "Школ", "Учеников"], ["2023", "12", "4 800"],
                                 ["2024", "35", "15 200"]]):  # fmt: skip
            for c, value in enumerate(row):
                t.table.cell(r, c).text = value
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def _obj(data: bytes, name: str):
    return next(o for o in objects(data) if o.name == name)


def _chart(data: bytes, name: str = "Диаграмма выручки"):
    return next(s for s in Presentation(io.BytesIO(data)).slides[0].shapes if s.name == name)


def _run(data: bytes, *ops: ObjectOp, name: str = "Диаграмма выручки"):
    obj = _obj(data, name)
    return apply_chart_ops(data, obj.slide, obj.shape_id, list(ops), TOKENS)


def test_value_change_rewrites_data_in_place() -> None:
    data = _deck()
    out, notes = _run(
        data,
        ObjectOp(
            op="chart.set_data",
            categories=["2023", "2024", "2025"],
            series=[ChartSeries(name="Выручка", values=[75, 120, 140])],
        ),
    )
    frame = _chart(out)
    assert read_data(frame.chart) == (["2023", "2024", "2025"], [("Выручка", [75, 120, 140])])
    assert "по вашему сообщению" in notes[0]
    # Та же фигура: имя и место прежние, тип прежний.
    assert frame.left == Inches(1) and kind_of(frame.chart) == "column"


def test_invented_numbers_are_rejected() -> None:
    op = ObjectOp(
        op="chart.set_data",
        categories=["2023", "2024", "2025"],
        series=[ChartSeries(name="Выручка", values=[75, 110, 140])],
    )
    with pytest.raises(ValueError, match="не пересчитывай"):
        check_numbers("увеличь 2024 на 10%", [75, 100, 140], [op])
    check_numbers("2024 — 110, а не 100", [75, 100, 140], [op])


def test_type_change_keeps_frame_name_and_order() -> None:
    data = _deck(table=True)
    out, notes = _run(data, ObjectOp(op="chart.type", chart_type="bar"))
    shapes = list(Presentation(io.BytesIO(out)).slides[0].shapes)
    names = [s.name for s in shapes]
    assert names == ["TextBox 1", "Диаграмма выручки", "Тарифы"]
    frame = shapes[1]
    assert kind_of(frame.chart) == "bar" and frame.top == Inches(2)
    assert read_data(frame.chart)[1] == [("Выручка", [75, 100, 140])]
    assert "полосы" in notes[0]
    # Старая часть диаграммы не висит в связях слайда.
    slide = Presentation(io.BytesIO(out)).slides[0]
    charts = [r for r in slide.part.rels.values() if r.reltype.endswith("/chart")]
    assert len(charts) == 1


def test_legend_labels_title_highlight_and_sort() -> None:
    data = _deck()
    out, notes = _run(
        data,
        ObjectOp(op="chart.legend", on=True, place="right"),
        ObjectOp(op="chart.labels", on=True),
        ObjectOp(op="chart.title", text="Выручка, млн ₽"),
        ObjectOp(op="chart.highlight", category="2024"),
        ObjectOp(op="chart.sort", descending=True),
    )
    chart = _chart(out).chart
    assert chart.has_legend and chart.legend.position == XL_LEGEND_POSITION.RIGHT
    assert chart.plots[0].has_data_labels
    assert chart.chart_title.text_frame.text == "Выручка, млн ₽"
    assert read_data(chart)[0] == ["2025", "2024", "2023"]
    assert any("2024" in n for n in notes)
    with pytest.raises(ValueError, match="категории «2030»"):
        _run(data, ObjectOp(op="chart.highlight", category="2030"))


def test_chart_to_table_and_back() -> None:
    data = _deck(table=True)
    out, notes = _run(data, ObjectOp(op="object.to_table"))
    frame = _chart(out)
    assert frame.has_table and "таблицей" in notes[0]
    cells = [[c.text for c in row.cells] for row in frame.table.rows]
    assert cells == [["", "Выручка"], ["2023", "75"], ["2024", "100"], ["2025", "140"]]
    out, notes = _run(data, ObjectOp(op="object.to_chart", chart_type="line"), name="Тарифы")
    frame = _chart(out, "Тарифы")
    assert frame.has_chart and kind_of(frame.chart) == "line"
    assert read_data(frame.chart) == (
        ["2023", "2024"],
        [("Школ", [12.0, 35.0]), ("Учеников", [4800.0, 15200.0])],
    )


def test_object_edit_plan_runs_chart_ops_after_text_ops() -> None:
    data = _deck()
    obj = _obj(data, "Диаграмма выручки")
    plan = ObjectEditPlan(
        explanation="Сделал круговой",
        patches=[],
        ops=[ObjectOp(op="chart.type", chart_type="pie")],
    )
    out, notes = apply_plan(data, ObjectTarget(slide=1, shape_id=obj.shape_id), plan, TOKENS)
    assert kind_of(_chart(out).chart) == "pie" and notes


def test_chart_from_file_rows_is_placed_in_template_style() -> None:
    spec = spec_from_table(["Год", "Школ"], [["2023", "12"], ["2024", "35"]])
    assert spec.chart_type == "column" and spec.series == [("Школ", [12.0, 35.0])]
    out = place_chart(_deck(), 1, spec, (0.6, 0.2, 0.35, 0.5), TOKENS)
    frame = _chart(out, "Диаграмма из чата")
    assert frame.has_chart and frame.left == round(0.6 * Inches(16))
    fill = frame.chart.plots[0].series[0].format.fill
    assert str(fill.fore_color.rgb) == "0077FF"  # первый яркий цвет палитры, не белый


def test_make_editable_without_composites_changes_nothing() -> None:
    data = _deck()
    assert make_editable(data, 1) == (data, [])


def _shrink_table_frame(data: bytes) -> bytes:
    """Рамка таблицы меньше её столбцов и строк — как у таблиц из Google Slides."""
    prs = Presentation(io.BytesIO(data))
    frame = next(s for s in prs.slides[0].shapes if s.name == "Тарифы")
    frame.width, frame.height = Inches(1), Inches(1)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def test_table_to_chart_takes_the_real_table_size() -> None:
    data = _shrink_table_frame(_deck(table=True))
    out, _ = _run(data, ObjectOp(op="object.to_chart", chart_type="column"), name="Тарифы")
    frame = _chart(out, "Тарифы")
    assert frame.width == Inches(5) and frame.height >= Inches(1.5)


def test_labels_are_rounded_and_missing_legend_is_explained() -> None:
    data = _deck()
    out, notes = _run(
        data, ObjectOp(op="chart.legend", on=False), ObjectOp(op="chart.labels", on=True)
    )
    labels = _chart(out).chart.plots[0].data_labels
    assert labels.number_format == "#,##0" and not labels.number_format_is_linked
    assert any("легенды нет" in n for n in notes)


def test_chart_to_table_rows_are_compact_and_centered() -> None:
    from pptx.enum.text import MSO_ANCHOR

    out, _ = _run(_deck(), ObjectOp(op="object.to_table"))
    frame = _chart(out)
    assert frame.height < Inches(5)
    assert frame.table.cell(1, 1).vertical_anchor == MSO_ANCHOR.MIDDLE


def test_wide_table_turns_rows_into_series() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    t = slide.shapes.add_table(3, 8, Inches(0.5), Inches(1), Inches(9), Inches(2))
    t.name = "Кварталы"
    header = ["Продукт"] + [f"{q} кв." for q in range(1, 8)]
    for c, value in enumerate(header):
        t.table.cell(0, c).text = value
    for r, name in enumerate(["Карты", "Вклады"], start=1):
        t.table.cell(r, 0).text = name
        for c in range(1, 8):
            t.table.cell(r, c).text = str(r * 10 + c)
    out = io.BytesIO()
    prs.save(out)
    data = out.getvalue()
    out_data, notes = _run(data, ObjectOp(op="object.to_chart"), name="Кварталы")
    chart = _chart(out_data, "Кварталы").chart
    categories, series = read_data(chart)
    assert [n for n, _ in series] == ["Карты", "Вклады"] and categories[0] == "1 кв."
    assert "данные те же" in notes[-1]
