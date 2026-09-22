"""The live SDK fixture never resolves or overwrites a user's office document."""

import hashlib
import importlib.util
import io
from pathlib import Path

import pytest
from pptx import Presentation

from presentation_designer.pipeline.artifacts import ArtifactStore
from presentation_designer.pipeline.office import OfficeStore
from presentation_designer.pipeline.results import build_generation_result
from presentation_designer.pipeline.state import State


def test_seed_is_unique_and_refuses_output_reuse(tmp_path):
    path = Path(__file__).resolve().parents[2] / "scripts/seed_office_sdk.py"
    spec = importlib.util.spec_from_file_location("seed_office_sdk", path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    data = tmp_path / "data"
    store = OfficeStore(data)
    user = store.create("user/source", "Original", b"user-original")
    before = store.get(user["id"])
    first = module.seed(data, tmp_path / "first")
    second = module.seed(data, tmp_path / "second")
    assert first["document_id"] != second["document_id"]
    assert first["source"] != second["source"]
    for report, name in [(first, "first"), (second, "second")]:
        doc = store.get(report["document_id"])
        assert doc["revision"] == 0 and doc["active_key"] is None
        assert doc["source"].startswith("sdk-sandbox/")
        content = store.read(doc["id"], 0)
        assert content == (tmp_path / name / "before.pptx").read_bytes()
        assert hashlib.sha256(content).hexdigest() == report["sha256"]
        assert len(Presentation(io.BytesIO(content)).slides) == 1
    with pytest.raises(FileExistsError):
        module.seed(data, tmp_path / "first")
    assert store.get(user["id"]) == before
    assert store.read(user["id"], 0) == b"user-original"


def test_project_seed_publishes_only_new_synthetic_artifacts(tmp_path):
    path = Path(__file__).resolve().parents[2] / "scripts/seed_office_sdk.py"
    spec = importlib.util.spec_from_file_location("seed_office_sdk", path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    data, artifacts = tmp_path / "data", tmp_path / "artifacts"
    state = State(data / "state.sqlite3")
    user = state.create_project("User project", {}, {})
    reports = [module.seed(data, tmp_path / name, artifacts) for name in ("first", "second")]
    assert reports[0]["project_id"] != reports[1]["project_id"]
    for report in reports:
        project = state.get_project(report["project_id"])
        assert project["job_id"] == report["job_id"]
        result = build_generation_result(state, report["job_id"])
        assert result["status"] == "succeeded"
        assert result["variants"][0]["artifacts"]["pptx"] == report["artifact"]
        content = (
            ArtifactStore(artifacts).resolve(report["job_id"], report["artifact"]).read_bytes()
        )
        assert content == OfficeStore(data).read(report["document_id"], 0)
        assert report["source"] == f"{report['job_id']}/{report['artifact']}"
        assert report["job_id"].startswith("job_sdk_sandbox_")
    with pytest.raises(FileExistsError):
        module.seed(data, tmp_path / "first", artifacts)
    assert len(state.list_projects()) == 3
    assert state.get_project(user["project_id"]) == user
    assert state.list_active_jobs() == []


def test_project_seed_can_be_opened_through_real_api(client, orchestrator, tmp_path):
    orchestrator.settings.onlyoffice.enabled = True
    orchestrator.settings.onlyoffice.jwt_secret = "sdk-seed-test-secret-not-for-production-123456"
    path = Path(__file__).resolve().parents[2] / "scripts/seed_office_sdk.py"
    spec = importlib.util.spec_from_file_location("seed_office_sdk", path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    report = module.seed(
        orchestrator.settings.data_dir, tmp_path / "seed", orchestrator.settings.artifacts_dir
    )
    project = client.get(f"/api/projects/{report['project_id']}")
    assert project.status_code == 200, project.text
    assert project.json()["job_id"] == report["job_id"]
    generation = client.get(f"/api/generations/{report['job_id']}")
    assert generation.status_code == 200, generation.text
    opened = client.post(
        "/api/office/documents", json={"job_id": report["job_id"], "artifact": report["artifact"]}
    )
    assert opened.status_code == 200, opened.text
    assert opened.json()["id"] == report["document_id"]
    assert opened.json()["revision"] == 0


@pytest.mark.parametrize("project", [False, True])
def test_source_seed_copies_bytes_without_modifying_original(tmp_path, project):
    path = Path(__file__).resolve().parents[2] / "scripts/seed_office_sdk.py"
    spec = importlib.util.spec_from_file_location("seed_office_sdk", path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    source = tmp_path / "brand.pptx"
    deck = Presentation()
    for _ in range(2):
        deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(source)
    original = source.read_bytes()
    data = tmp_path / "data"
    artifacts = tmp_path / "artifacts" if project else None
    reports = [
        module.seed(data, tmp_path / name, artifacts, source_path=source)
        for name in ("first", "second")
    ]
    assert reports[0]["document_id"] != reports[1]["document_id"]
    for report in reports:
        assert report["source"].startswith("job_sdk_sandbox_" if project else "sdk-brand/")
        assert report["synthetic"] is False and report["slides"] == 2
        assert OfficeStore(data).read(report["document_id"], 0) == original
        if project:
            result = build_generation_result(State(data / "state.sqlite3"), report["job_id"])
            assert result["variants"][0]["slide_count"] == 2
            assert (
                ArtifactStore(artifacts).resolve(report["job_id"], report["artifact"]).read_bytes()
                == original
            )
    assert source.read_bytes() == original
    with pytest.raises(FileExistsError):
        module.seed(data, tmp_path / "first", source_path=source)
