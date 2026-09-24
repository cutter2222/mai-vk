"""Снимок колоды (этап 36): одна форма для ревизии варианта и офисной копии.

Копия после сохранения в ONLYOFFICE получает новые `cNvPr id` и `p:sldId` (проверено на
сохранённых ревизиях izbox.ru), имена фигур остаются. Снимок копии с ComposedDeck ревизии-
источника должен совпасть со снимком самой ревизии всем, кроме номеров.
"""

from __future__ import annotations

import io
import json
import pathlib
import re
import zipfile
from typing import Any

import jsonschema
import pytest
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
from referencing import Registry, Resource

from presentation_designer.generation.snapshot import deck_snapshot, outline_lines, slide_lines
from presentation_designer.layout.composed import SlideRecord, SlotFill, build_composed_deck
from presentation_designer.parsing.template.package import open_template

ROOT = pathlib.Path(__file__).resolve().parents[2]
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"


def validator() -> jsonschema.Draft202012Validator:
    registry = Registry()
    schemas = {}
    for path in (ROOT / "contracts" / "schemas").glob("*.schema.json"):
        schema = json.loads(path.read_text())
        schemas[path.name] = schema
        registry = registry.with_resource(schema["$id"], Resource.from_contents(schema))
    return jsonschema.Draft202012Validator(schemas["deck_snapshot.schema.json"], registry=registry)


def assert_valid(snapshot: dict[str, Any]) -> None:
    errors = [f"{list(e.path)}: {e.message}" for e in validator().iter_errors(snapshot)]
    assert not errors, errors[:5]


def build_deck(*, drop_chart_cache: bool = False) -> bytes:
    deck = Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[1])  # заголовок и текст из макета
    first.shapes.title.text = "Выручка растёт"
    first.placeholders[1].text = "Первый пункт\nВторой пункт"
    first.notes_slide.notes_text_frame.text = "Сказать про рост"
    outer = first.shapes.add_group_shape()
    inner = outer.shapes.add_group_shape()
    inner.shapes.add_textbox(Inches(1), Inches(5), Inches(2), Inches(1)).text = "В группе"
    outer.shapes.add_textbox(Inches(4), Inches(5), Inches(2), Inches(1)).text = "Рядом"
    turned = first.shapes.add_textbox(Inches(7), Inches(1), Inches(2), Inches(1))
    turned.text = "Повёрнут"
    turned.rotation = 30
    second = deck.slides.add_slide(deck.slide_layouts[5])
    second.shapes.title.text = "Данные"
    table = second.shapes.add_table(3, 3, Inches(0.5), Inches(1.5), Inches(5), Inches(2)).table
    for r, row in enumerate([["Год", "Выручка", ""], ["2023", "100", "+5%"], ["2024", "120", ""]]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    table.cell(0, 1).merge(table.cell(0, 2))
    data = CategoryChartData()
    data.categories = ["2023", "2024"]
    data.add_series("Выручка", (100, 120))
    frame = second.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6), Inches(1.5), Inches(3.5), Inches(3), data
    )
    frame.name = "Диаграмма выручки"
    if drop_chart_cache:
        chart = frame.chart._chartSpace
        for cache in list(chart.iter(f"{C}numCache")) + list(chart.iter(f"{C}strCache")):
            cache.getparent().remove(cache)
    out = io.BytesIO()
    deck.save(out)
    return out.getvalue()


def objects_by_text(slide: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {o["text"]["plain"]: o for o in slide["objects"] if o.get("text")}


def test_snapshot_sees_groups_placeholders_rotation_notes_table_and_chart() -> None:
    snapshot = deck_snapshot(build_deck())
    assert_valid(snapshot)
    first, second = snapshot["slides"]
    assert first["title"] == "Выручка растёт" and first["notes"] == "Сказать про рост"
    texts = objects_by_text(first)
    title = texts["Выручка растёт"]
    assert title["role"] == "title"
    assert title["placeholder"] == {"type": "title", "idx": 0, "inherited_geometry": True}
    assert title["bbox"]["width"] > 0.5  # рамка из макета, а не нули
    assert texts["Первый пункт\nВторой пункт"]["role"] == "body"
    nested = texts["В группе"]
    assert len(nested["address"]["group_path"]) == 2
    groups = [o for o in first["objects"] if o["kind"] == "group"]
    assert [o["address"]["object_id"] for o in groups] == nested["address"]["group_path"]
    assert texts["Рядом"]["address"]["group_path"] == nested["address"]["group_path"][:1]
    assert texts["Повёрнут"]["rotation_deg"] == 30
    assert "style" in title and title["style"]["size_pt"] > 0
    table = next(o for o in second["objects"] if o["kind"] == "table")
    assert table["role"] == "table"
    assert table["table"]["rows"] == [
        ["Год", "Выручка", ""],
        ["2023", "100", "+5%"],
        ["2024", "120", ""],
    ]
    assert table["table"]["merged"] == [{"row": 1, "col": 2, "rows": 1, "cols": 2}]
    chart = next(o for o in second["objects"] if o["kind"] == "chart")
    assert chart["chart"] == {
        "type": "column_clustered",
        "categories": ["2023", "2024"],
        "series": [{"name": "Выручка", "values": [100.0, 120.0]}],
        "values_from": "cache",
    }
    assert snapshot["outline"][1]["kinds"] == ["text", "table", "chart"]
    assert outline_lines(snapshot)[1] == "2. «Данные» — текст, таблица, диаграмма"
    lines = slide_lines(second)
    assert "2023 | 100 | +5%" in lines and "  Выручка: 100, 120" in lines


def test_chart_without_cache_reads_embedded_workbook() -> None:
    snapshot = deck_snapshot(build_deck(drop_chart_cache=True))
    chart = next(o for o in snapshot["slides"][1]["objects"] if o["kind"] == "chart")["chart"]
    assert chart["values_from"] == "workbook"
    assert chart["series"] == [{"name": "Выручка", "values": [100.0, 120.0]}]
    assert chart["categories"] == ["2023", "2024"]


def renumber_like_onlyoffice(data: bytes) -> bytes:
    """Сохранение ONLYOFFICE: новые номера фигур на слайдах и новые `p:sldId`, имена те же."""
    source = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    counter = iter(range(1_000_000_001, 2_000_000_000, 7919))
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for entry in source.infolist():
            blob = source.read(entry)
            if re.fullmatch(r"ppt/slides/slide\d+\.xml", entry.filename):
                root = etree.fromstring(blob)
                for node in root.iter(f"{P}cNvPr"):
                    node.set("id", str(next(counter)))
                blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            elif entry.filename == "ppt/presentation.xml":
                root = etree.fromstring(blob)
                for n, node in enumerate(root.iter(f"{P}sldId")):
                    node.set("id", str(9000 + n))
                blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            target.writestr(entry, blob)
    return out.getvalue()


def composed_for(path: pathlib.Path) -> tuple[dict[str, Any], dict[str, Any]]:
    """ComposedDeck ревизии по самому файлу, как у варианта original: слайд плана на каждый
    слайд, заголовок — слот title."""
    pkg = open_template(path)
    records = []
    plan_slides = []
    for slide in pkg.slides:
        title = next(
            (s for s in slide.shapes if s.placeholder_type in ("title", "ctrTitle") and s.text),
            None,
        )
        fills = (
            [SlotFill("title", "title", "title", title.element_id, title.element_id)]
            if title
            else []
        )
        slide_id = f"s{slide.index:02d}"
        records.append(
            SlideRecord(
                slide_id=slide_id,
                order=slide.index,
                pattern_id=f"p{slide.index}",
                source_slide_index=slide.index,
                source_slide_part=slide.part,
                layout_id=slide.layout_id,
                title=title.text if title else "",
                fills=fills,
            )
        )
        plan_slides.append({"slide_id": slide_id, "role": "freeform", "order": slide.index})
    composed = build_composed_deck(
        path,
        plan={"plan_id": "plan", "slides": plan_slides},
        profile={"template_id": "tpl"},
        package={},
        records=records,
        job_id="job",
        variant_id="original",
        revision=1,
        pptx_artifact="original/r1/deck.pptx",
        warnings=[],
        composer={"name": "test", "version": "0"},
    )
    return composed, {"slides": plan_slides}


def comparable(snapshot: dict[str, Any]) -> list[Any]:
    """Всё, кроме номеров ревизии: адреса — через имена и глубину групп."""
    out = []
    for slide in snapshot["slides"]:
        objects = []
        for obj in slide["objects"]:
            item = {k: v for k, v in obj.items() if k != "address"}
            item["depth"] = len(obj["address"]["group_path"])
            objects.append(item)
        rest = {k: v for k, v in slide.items() if k not in ("objects", "sld_id")}
        out.append((rest, objects))
    return out


def check_same_form(path: pathlib.Path) -> None:
    data = path.read_bytes()
    composed, plan = composed_for(path)
    revision = deck_snapshot(data, composed=composed, plan=plan)
    copy = deck_snapshot(renumber_like_onlyoffice(data), composed=composed, plan=plan)
    assert_valid(revision)
    assert_valid(copy)
    assert [s["sld_id"] for s in revision["slides"]] != [s["sld_id"] for s in copy["slides"]]
    assert comparable(revision) == comparable(copy)
    assert all(s.get("slide_id") == f"s{s['index']:02d}" for s in copy["slides"])
    titled = [
        o
        for s in copy["slides"]
        for o in s["objects"]
        if (o.get("slot") or {}).get("slot_id") == "title"
    ]
    assert titled and all(o["role"] == "title" for o in titled)


def test_revision_and_office_copy_give_the_same_snapshot(tmp_path: pathlib.Path) -> None:
    path = tmp_path / "deck.pptx"
    path.write_bytes(build_deck())
    check_same_form(path)


@pytest.mark.organizer_data
@pytest.mark.parametrize(
    "name",
    [
        "VK Tech шаблон.pptx",
        "VK_WorkSpace_Клиентская_конференция_Шаблон_03.pptx",
        "ЛЦТ2026 Шаблон презентации.pptx",
        "Шаблон презентации VK Education.pptx",
    ],
)
def test_organizer_templates_revision_and_copy_match(
    organizer_dir: pathlib.Path, name: str
) -> None:
    path = organizer_dir / name
    check_same_form(path)
    snapshot = deck_snapshot(path.read_bytes())
    objects = [o for s in snapshot["slides"] for o in s["objects"]]
    kinds = {o["kind"] for o in objects}
    assert "group" in kinds and any(o["address"]["group_path"] for o in objects)
    if name.startswith("ЛЦТ"):
        charts = [o["chart"] for o in objects if o.get("chart")]
        assert len(charts) == 3 and all(c["series"] and c["categories"] for c in charts)
        assert any(o.get("placeholder", {}).get("inherited_geometry") for o in objects)
    if name.startswith(("Шаблон презентации VK Education", "VK_WorkSpace")):
        tables = [o["table"]["rows"] for o in objects if o.get("table")]
        assert tables and all(rows and rows[0] for rows in tables)


def test_edited_marks_only_changed_slides() -> None:
    data = build_deck()
    deck = Presentation(io.BytesIO(data))
    deck.slides[1].shapes.title.text = "Данные за два года"
    out = io.BytesIO()
    deck.save(out)
    base = deck_snapshot(data)
    edited = deck_snapshot(renumber_like_onlyoffice(out.getvalue()), base=base)
    assert [s["edited"] for s in edited["slides"]] == [False, True]
    assert [e["edited"] for e in edited["outline"]] == [False, True]


def test_example_matches_schema() -> None:
    assert_valid(json.loads((ROOT / "contracts/examples/deck_snapshot.example.json").read_text()))
