"""Object scope and geometry must validate inside the model repair boundary."""

import io
from copy import deepcopy
from zipfile import ZipFile

import pytest
from pptx import Presentation
from pptx.util import Inches

from presentation_designer.generation import office_object_edit as edit
from presentation_designer.generation.office_objects import ObjectTarget, objects
from presentation_designer.llm.cache import ResponseCache
from presentation_designer.llm.client import LlmClient
from presentation_designer.llm.limiter import LocalLimiter, Quota
from presentation_designer.llm.retry import RetryPolicy
from presentation_designer.llm.stub import StubTransport
from presentation_designer.llm.types import ResponseError
from presentation_designer.shared.settings import Settings, get_models_config


@pytest.fixture
def deck_bytes():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    for index, text in enumerate(["Выбранный текст", "Соседний текст"]):
        slide.shapes.add_textbox(Inches(1), Inches(index + 1), Inches(3), Inches(1)).text = text
    deck.slides.add_slide(deck.slide_layouts[6]).shapes.add_textbox(
        Inches(1), Inches(1), Inches(3), Inches(1)
    ).text = "Другой слайд"
    output = io.BytesIO()
    deck.save(output)
    return output.getvalue()


@pytest.mark.parametrize(
    ("scope", "fault", "hint"),
    [
        ("object", "run", "вне выбранного объекта"),
        ("object", "position", "границы слайда"),
        ("many", "run", "вне выбранного объекта"),
        ("many", "position", "границы слайда"),
        ("many", "target", "вне выбора"),
        ("many", "duplicate", "повторно"),
    ],
)
@pytest.mark.parametrize("recover", [True, False])
async def test_invalid_object_plan_is_repaired_before_acceptance(
    deck_bytes, scope, fault, hint, recover, tmp_path, monkeypatch
):
    chosen, neighbour, _ = objects(deck_bytes)
    target = ObjectTarget(slide=chosen.slide, shape_id=chosen.shape_id)
    good = {
        "explanation": "Изменила текст",
        "patches": [{"slide": 1, "run": 0, "before": "Выбранный текст", "after": "Новый текст"}],
    }
    bad = deepcopy(good)
    if fault == "run":
        bad["patches"] = [{"slide": 1, "run": 1, "before": "Соседний текст", "after": "Нельзя"}]
    elif fault == "position":
        bad["position"] = {"x": 0.99, "y": 0.1}
    if scope == "many":
        good = {
            "explanation": "Изменила текст",
            "edits": [{"target": target.model_dump(), "plan": good}],
        }
        bad = {
            "explanation": "Изменила текст",
            "edits": [{"target": target.model_dump(), "plan": bad}],
        }
        if fault == "target":
            bad["edits"][0]["target"]["shape_id"] = neighbour.shape_id
            bad["edits"][0]["plan"]["patches"] = []
        elif fault == "duplicate":
            bad["edits"].append(deepcopy(bad["edits"][0]))

    stub = StubTransport()
    stub.answer(bad, times=1 if recover else None)
    stub.answer(good)
    settings = Settings()
    client = LlmClient(
        settings=settings,
        models=get_models_config(),
        transport=stub,
        limiter=LocalLimiter(Quota(2, 1000, 1_000_000)),
        cache=ResponseCache("off", tmp_path / "cache", tmp_path / "fixtures"),
        retry_policy=RetryPolicy(max_retries=1, base_s=0, max_s=0),
    )
    monkeypatch.setattr(edit, "build_client", lambda settings: client)

    async def propose():
        if scope == "object":
            return await edit.propose(deck_bytes, "Измени выбранный текст", settings, target)
        return await edit.propose_many(deck_bytes, "Измени выбранный текст", settings, [target])

    if not recover:
        with pytest.raises(ResponseError, match=hint):
            await propose()
    else:
        plan = await propose()
        assert plan.model_dump(exclude_none=True) == good
        updated = (
            edit.patch_object(deck_bytes, target, plan)
            if scope == "object"
            else edit.patch_objects(deck_bytes, [target], plan)
        )
        deck = Presentation(io.BytesIO(updated))
        assert deck.slides[0].shapes[0].text == "Новый текст"
        assert deck.slides[0].shapes[1].text == "Соседний текст"
        with ZipFile(io.BytesIO(deck_bytes)) as before, ZipFile(io.BytesIO(updated)) as after:
            assert before.namelist() == after.namelist()
            for name in before.namelist():
                if name != "ppt/slides/slide1.xml":
                    assert before.read(name) == after.read(name), name
    assert len(stub.calls) == 2
    assert hint in stub.calls[1].messages[-1].text
    assert Presentation(io.BytesIO(deck_bytes)).slides[0].shapes[0].text == "Выбранный текст"


@pytest.mark.parametrize("fault", ["empty", "duplicate", "missing"])
async def test_invalid_selection_never_calls_model(deck_bytes, fault, monkeypatch):
    chosen = objects(deck_bytes)[0]
    target = ObjectTarget(slide=chosen.slide, shape_id=chosen.shape_id)
    targets = {
        "empty": [],
        "duplicate": [target, target],
        "missing": [ObjectTarget(slide=1, shape_id="9999")],
    }[fault]

    def unexpected_client(settings):
        pytest.fail("Invalid selection must fail before constructing the LLM client")

    monkeypatch.setattr(edit, "build_client", unexpected_client)
    with pytest.raises(ValueError, match=r"выбор объектов|не найден"):
        await edit.propose_many(deck_bytes, "Правее", Settings(), targets)
