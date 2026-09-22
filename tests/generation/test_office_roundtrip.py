"""Round-trip diagnostics must reject malformed editor output, never repair it."""

import importlib.util
import json
import sys
from pathlib import Path
from zipfile import ZipFile

import pytest
from pptx import Presentation
from pptx.util import Inches


@pytest.fixture
def audit():
    path = Path(__file__).resolve().parents[2] / "scripts/check_office_roundtrip.py"
    spec = importlib.util.spec_from_file_location("check_office_roundtrip", path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_invalid_layout_attribute_is_reported_without_repair(audit, tmp_path, monkeypatch):
    before, after = tmp_path / "before.pptx", tmp_path / "after.pptx"
    Presentation().save(before)
    with ZipFile(before) as source, ZipFile(after, "w") as target:
        for name in source.namelist():
            content = source.read(name)
            if name == "ppt/slideLayouts/slideLayout1.xml":
                content = b'<layout matchingName="Slide "Thanks!""/>'
            target.writestr(name, content)
    original, invalid = before.read_bytes(), after.read_bytes()
    assert audit.xml_errors(before) == []
    assert audit.xml_errors(after)[0]["part"] == "ppt/slideLayouts/slideLayout1.xml"
    out = tmp_path / "audit"
    monkeypatch.setattr(
        sys,
        "argv",
        [
            "check",
            str(before),
            str(after),
            "--out",
            str(out),
            "--slide",
            "1",
            "--old",
            "Old",
            "--new",
            "New",
        ],
    )
    with pytest.raises(SystemExit, match="Invalid OOXML"):
        audit.main()
    report = json.loads((out / "report.json").read_text())
    assert report["accepted"] is False
    assert report["visual_comparison_performed"] is False
    assert report["xml_errors"]["before"] == []
    assert before.read_bytes() == original and after.read_bytes() == invalid
    with pytest.raises(FileExistsError):
        audit.main()


def test_snapshot_retains_text_geometry_and_table(audit, tmp_path):
    path = tmp_path / "deck.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(2), Inches(3), Inches(1)).text = "Brand text"
    slide.shapes.add_table(1, 1, Inches(1), Inches(4), Inches(3), Inches(1)).table.cell(
        0, 0
    ).text = "Cell"
    deck.save(path)
    snapshot = audit.snapshot(path)
    assert snapshot["size"] == [deck.slide_width, deck.slide_height]
    assert snapshot["slides"][0][0]["text"] == "Brand text"
    assert snapshot["slides"][0][0]["box"][:2] == [Inches(1), Inches(2)]
    assert snapshot["slides"][0][1]["table"] == [["Cell"]]


def test_numeric_formatting_drift_fails_even_when_pixels_match(audit, tmp_path, monkeypatch):
    from types import SimpleNamespace

    from lxml import etree
    from PIL import Image
    from pptx.util import Pt

    before, after = tmp_path / "before.pptx", tmp_path / "after.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    box.text = "Old"
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(14.06)
    deck.save(before)
    with ZipFile(before) as source, ZipFile(after, "w") as target:
        for info in source.infolist():
            data = source.read(info)
            if info.filename == "ppt/slides/slide1.xml":
                root = etree.fromstring(data)
                root.find(
                    ".//{http://schemas.openxmlformats.org/drawingml/2006/main}t"
                ).text = "New"
                root.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}rPr").set(
                    "sz", "1400"
                )
                data = etree.tostring(root)
            target.writestr(info, data)
    image = tmp_path / "slide.png"
    Image.new("RGB", (10, 10), "white").save(image)
    monkeypatch.setattr(audit, "convert_to_pdf", lambda *a, **kw: SimpleNamespace(pdf_path=image))
    monkeypatch.setattr(audit, "render_thumbnails", lambda *a, **kw: SimpleNamespace(paths=[image]))
    monkeypatch.setattr(audit, "get_settings", lambda: None)
    out = tmp_path / "audit"
    monkeypatch.setattr(
        sys,
        "argv",
        [
            "check",
            str(before),
            str(after),
            "--out",
            str(out),
            "--slide",
            "1",
            "--old",
            "Old",
            "--new",
            "New",
        ],
    )
    with pytest.raises(SystemExit) as error:
        audit.main()
    assert error.value.code == 1
    report = json.loads((out / "report.json").read_text())
    assert report["expected_text_and_geometry_only"] is True
    assert report["pixels"][0]["identical"] is True
    assert report["numeric_xml_deltas"][0]["delta"] == "-6"
