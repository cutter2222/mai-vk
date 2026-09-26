"""Картинка из сообщения — на слайд офисной копии: блок, нарисованный фоном, находится по
снимку слайда вплоть до скругления; объект слайда даёт свою рамку; картинка на месте
картинки заменяет её."""

from __future__ import annotations

import hashlib
import io

import pytest
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

from presentation_designer.generation.office_image import (
    Area,
    Placement,
    image_bytes,
    place_image,
    snap_area,
)
from presentation_designer.generation.office_objects import objects


def _png(size: tuple[int, int], color: tuple[int, int, int], fmt: str = "PNG") -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, fmt)
    return buf.getvalue()


def _snapshot() -> bytes:
    """Слайд 1600×900: светлый фон и белый скруглённый блок справа (как «Спасибо» VK Tech)."""
    image = Image.new("RGB", (1600, 900), (245, 246, 250))
    ImageDraw.Draw(image).rounded_rectangle((928, 173, 1480, 727), radius=40, fill="white")
    buf = io.BytesIO()
    image.save(buf, "PNG")
    return buf.getvalue()


BLOCK = (928 / 1600, 173 / 900, 552 / 1600, 554 / 900)


def _deck(*, picture: bool = False) -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(16), Inches(9)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if picture:
        slide.shapes.add_picture(
            io.BytesIO(_png((100, 100), (0, 0, 255))), Inches(9), Inches(2), Inches(5), Inches(5)
        )
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_block_drawn_on_the_background_is_found_by_its_pixels() -> None:
    # Модель назвала место грубо — внутри блока и меньше его: рамка берётся по самому блоку.
    for rough in ((0.62, 0.3, 0.2, 0.3), (0.55, 0.15, 0.4, 0.7)):
        box, radius = snap_area(_snapshot(), rough)
        assert box == pytest.approx(BLOCK, abs=0.01)
        assert 0.03 < radius < 0.12
    # Место без однородного блока («справа» на узоре) остаётся тем, что назвала модель.
    noise = Image.effect_noise((1600, 900), 80).convert("RGB")
    buf = io.BytesIO()
    noise.save(buf, "PNG")
    assert snap_area(buf.getvalue(), (0.6, 0.2, 0.3, 0.5)) == ((0.6, 0.2, 0.3, 0.5), 0.0)


def test_picture_fills_the_block_without_distortion_and_takes_its_corners() -> None:
    photo = _png((1200, 600), (200, 40, 40), "JPEG")
    placement = Placement(
        explanation="в белый блок справа", area=Area(x=0.6, y=0.3, width=0.2, height=0.3)
    )
    out = Presentation(io.BytesIO(place_image(_deck(), 1, photo, placement, snapshot=_snapshot())))
    (pic,) = [s for s in out.slides[0].shapes if s.shape_type == 13]
    frame = (
        pic.left / Inches(16),
        pic.top / Inches(9),
        pic.width / Inches(16),
        pic.height / Inches(9),
    )
    assert frame == pytest.approx(BLOCK, abs=0.01)
    # Картинка 2:1 в почти квадратный блок: обрезана по бокам поровну, по высоте целиком.
    assert pic.crop_left == pytest.approx(pic.crop_right) and pic.crop_left > 0.2
    assert pic.crop_top == pic.crop_bottom == 0
    geom = pic._element.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}prstGeom")
    assert geom.get("prst") == "roundRect"


def test_contain_keeps_the_whole_picture_inside_the_place() -> None:
    photo = _png((1200, 600), (200, 40, 40))
    placement = Placement(
        explanation="целиком", area=Area(x=0.5, y=0.2, width=0.4, height=0.6), fit="contain"
    )
    out = Presentation(io.BytesIO(place_image(_deck(), 1, photo, placement)))
    (pic,) = [s for s in out.slides[0].shapes if s.shape_type == 13]
    assert pic.crop_left == pic.crop_top == 0
    assert pic.width / pic.height == pytest.approx(2, abs=0.01)
    assert pic.width == pytest.approx(0.4 * Inches(16), rel=0.01)


def test_picture_in_place_of_a_picture_replaces_it() -> None:
    deck = _deck(picture=True)
    (old,) = [o for o in objects(deck) if o.kind == "pic"]
    photo = _png((600, 600), (10, 200, 10))
    placement = Placement(explanation="заменил")
    out = Presentation(io.BytesIO(place_image(deck, 1, photo, placement, target=old)))
    (pic,) = [s for s in out.slides[0].shapes if s.shape_type == 13]
    assert hashlib.sha256(pic.image.blob).hexdigest() == hashlib.sha256(photo).hexdigest()
    assert (pic.left, pic.top, pic.width, pic.height) == (
        Inches(9),
        Inches(2),
        Inches(5),
        Inches(5),
    )


def test_images_are_normalized_and_bad_input_is_explained() -> None:
    gif = _png((30, 20), (1, 2, 3), "GIF")
    data, size = image_bytes(gif)
    assert size == (30, 20) and Image.open(io.BytesIO(data)).format == "PNG"
    with pytest.raises(ValueError, match="PNG или JPG"):
        image_bytes(b"not an image")
    with pytest.raises(ValueError, match="слайда 3 нет"):
        place_image(
            _deck(), 3, gif, Placement(explanation="", area=Area(x=0, y=0, width=1, height=1))
        )
    with pytest.raises(ValueError, match="куда поставить"):
        place_image(_deck(), 1, gif, Placement(explanation=""))


def test_free_area_avoids_content_and_keeps_margins() -> None:
    from presentation_designer.generation.office_image import free_area

    # Заголовок сверху и текст слева: свободна правая часть под заголовком.
    boxes = [(0.05, 0.05, 0.9, 0.12), (0.05, 0.25, 0.4, 0.6)]
    x, y, w, h = free_area(boxes)
    assert x >= 0.45 and y >= 0.17 and x + w <= 0.96 + 1e-9 and w >= 0.4 and h >= 0.6


def test_chart_area_falls_back_to_free_space_or_refuses() -> None:
    from presentation_designer.generation.office_image import area_for

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(16), Inches(9)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.8), Inches(1), Inches(7), Inches(7))
    box.text_frame.text = "Текст слева"
    out = io.BytesIO()
    prs.save(out)
    data = out.getvalue()
    # Модель назвала узкую полоску у правого края: берётся свободная правая часть.
    strip = Placement(explanation="справа", area=Area(x=0.97, y=0.1, width=0.03, height=0.8))
    x, _, w, _ = area_for(data, 1, strip, what="диаграмму")
    assert x >= 0.5 and w >= 0.3
    # Место закрыло бы текст, а свободного нет — честный отказ.
    wide = slide.shapes.add_textbox(Inches(8.5), Inches(1), Inches(7), Inches(7))
    wide.text_frame.text = "Текст справа"
    out = io.BytesIO()
    prs.save(out)
    with pytest.raises(ValueError, match="нет свободного места под диаграмму"):
        area_for(out.getvalue(), 1, strip, what="диаграмму")
