import io
import json
import subprocess
import sys
from pathlib import Path
from zipfile import ZipFile

import pytest
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt

from presentation_designer.generation.office_edit import NS, EditPlan, TextPatch, patch_pptx
from presentation_designer.generation.office_objects import ObjectTarget, objects
from presentation_designer.generation.office_preservation import verify_selected_text_edit


@pytest.fixture
def deck():
    presentation = Presentation()
    for _ in range(2):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        for text in ("Подробный заголовок презентации", "Не менять"):
            shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
            run = shape.text_frame.paragraphs[0].add_run()
            run.text = text
            run.font.name = "Arial"
            run.font.size = Pt(24)
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def rewrite(data, name, transform):
    output = io.BytesIO()
    with ZipFile(io.BytesIO(data)) as source, ZipFile(output, "w") as target:
        for entry in source.infolist():
            content = source.read(entry)
            target.writestr(entry, transform(content) if entry.filename == name else content)
    return output.getvalue()


def test_translation_and_shortening_preserve_selected_text_only(deck):
    obj = objects(deck)[0]
    for text in ("Detailed presentation title", "Заголовок"):
        result = patch_pptx(
            deck,
            EditPlan(
                explanation="",
                patches=[
                    TextPatch(slide=1, run=0, before="Подробный заголовок презентации", after=text),
                ],
            ),
        )
        report = verify_selected_text_edit(deck, result, [obj])
        assert report["changed_text_runs"] == 1
        assert report["changed_slide_parts"] == ["ppt/slides/slide1.xml"]
        assert report["slides"] == 2
        assert Presentation(io.BytesIO(result)).slides[0].shapes[0].text == text


@pytest.mark.parametrize(
    "name",
    [
        "ppt/slideMasters/slideMaster1.xml",
        "ppt/slideLayouts/slideLayout7.xml",
        "ppt/theme/theme1.xml",
        "ppt/slides/_rels/slide1.xml.rels",
        "ppt/presentation.xml",
        "ppt/slides/slide2.xml",
    ],
)
def test_rejects_changes_to_protected_parts(deck, name):
    result = rewrite(deck, name, lambda content: content + b"\n")
    with pytest.raises(ValueError, match="защищённая часть"):
        verify_selected_text_edit(deck, result, [objects(deck)[0]])


@pytest.mark.parametrize(
    "xpath,attribute,value",
    [
        (".//a:rPr", "sz", "3600"),
        (".//a:off", "x", "0"),
        (".//a:latin", "typeface", "Another Font"),
    ],
)
def test_rejects_formatting_or_geometry_change(deck, xpath, attribute, value):
    def change(content):
        root = etree.fromstring(content)
        root.find(xpath, NS).set(attribute, value)
        return etree.tostring(root)

    result = rewrite(deck, "ppt/slides/slide1.xml", change)
    with pytest.raises(ValueError, match="оформление"):
        verify_selected_text_edit(deck, result, [objects(deck)[0]])


def test_rejects_other_text_on_same_slide(deck):
    result = patch_pptx(
        deck,
        EditPlan(
            explanation="",
            patches=[
                TextPatch(slide=1, run=1, before="Не менять", after="Changed"),
            ],
        ),
    )
    with pytest.raises(ValueError, match="вне выбора"):
        verify_selected_text_edit(deck, result, [objects(deck)[0]])


def test_rejects_added_removed_or_duplicate_zip_parts(deck):
    for mode in ("add", "remove", "duplicate"):
        output = io.BytesIO()
        with ZipFile(io.BytesIO(deck)) as source, ZipFile(output, "w") as target:
            for entry in source.infolist():
                if mode != "remove" or entry.filename != "ppt/theme/theme1.xml":
                    target.writestr(entry, source.read(entry))
            if mode == "add":
                target.writestr("unexpected.xml", b"<root/>")
            if mode == "duplicate":
                with pytest.warns(UserWarning, match="Duplicate name"):
                    target.writestr("ppt/theme/theme1.xml", source.read("ppt/theme/theme1.xml"))
        with pytest.raises(ValueError, match="состав PPTX"):
            verify_selected_text_edit(deck, output.getvalue(), [objects(deck)[0]])


def test_rejects_run_removal(deck):
    def remove(content):
        root = etree.fromstring(content)
        node = root.find(".//a:r", NS)
        node.getparent().remove(node)
        return etree.tostring(root)

    result = rewrite(deck, "ppt/slides/slide1.xml", remove)
    with pytest.raises(ValueError, match="структура"):
        verify_selected_text_edit(deck, result, [objects(deck)[0]])


def test_noop_and_invalid_selection(deck):
    obj = objects(deck)[0]
    assert verify_selected_text_edit(deck, deck, [obj])["changed_text_runs"] == 0
    for selection in ([], [obj, obj], [ObjectTarget(slide=1, shape_id="99999")]):
        with pytest.raises(ValueError):
            verify_selected_text_edit(deck, deck, selection)


def test_sandbox_cli_offline_and_no_overwrite(deck, tmp_path):
    source = tmp_path / "source.pptx"
    source.write_bytes(deck)
    plan = tmp_path / "plan.json"
    plan.write_text(
        json.dumps(
            {
                "explanation": "offline fixture",
                "patches": [
                    {
                        "slide": 1,
                        "run": 0,
                        "before": "Подробный заголовок презентации",
                        "after": "Заголовок",
                    }
                ],
            }
        ),
        encoding="utf-8",
    )
    output = tmp_path / "sandbox"
    script = Path(__file__).resolve().parents[2] / "scripts/check_office_ai.py"
    command = [
        sys.executable,
        str(script),
        str(source),
        "--out",
        str(output),
        "--slide",
        "1",
        "--shape-id",
        objects(deck)[0].shape_id,
        "--instruction",
        "Сократи",
        "--plan",
        str(plan),
    ]
    completed = subprocess.run(command, capture_output=True, text=True, timeout=30)
    assert completed.returncode == 0, completed.stderr
    report = json.loads((output / "report.json").read_text())
    assert report["mode"] == "offline-plan"
    assert report["source_unchanged"]
    assert report["preservation"]["changed_text_runs"] == 1
    assert not report["official_ai_plugin_tested"]
    assert source.read_bytes() == deck
    candidate = (output / "after.pptx").read_bytes()
    repeated = subprocess.run(command, capture_output=True, text=True, timeout=30)
    assert repeated.returncode != 0
    assert (output / "after.pptx").read_bytes() == candidate

    # Offline plans must not bypass the same shortening guard used for live responses.
    invalid = json.loads(plan.read_text())
    invalid["patches"][0]["after"] = "Подробный заголовок презентации с лишними повторами"
    plan.write_text(json.dumps(invalid), encoding="utf-8")
    rejected_output = tmp_path / "rejected"
    command[command.index("--out") + 1] = str(rejected_output)
    rejected = subprocess.run(command, capture_output=True, text=True, timeout=30)
    assert rejected.returncode != 0
    assert "увеличило длину абзаца" in rejected.stderr
    assert not (rejected_output / "after.pptx").exists()
    assert not (rejected_output / "report.json").exists()
    assert source.read_bytes() == deck


@pytest.mark.parametrize("position", [None, {"x": 0.2, "y": 0.2}])
def test_sandbox_cli_rejects_noop_or_movement(deck, tmp_path, position):
    source = tmp_path / "source.pptx"
    source.write_bytes(deck)
    plan = tmp_path / "plan.json"
    plan.write_text(json.dumps({"explanation": "", "patches": [], "position": position}))
    output = tmp_path / "sandbox"
    completed = subprocess.run(
        [
            sys.executable,
            str(Path(__file__).resolve().parents[2] / "scripts/check_office_ai.py"),
            str(source),
            "--out",
            str(output),
            "--slide",
            "1",
            "--shape-id",
            objects(deck)[0].shape_id,
            "--instruction",
            "Сократи",
            "--plan",
            str(plan),
        ],
        capture_output=True,
        text=True,
        timeout=30,
    )
    assert completed.returncode != 0
    assert not (output / "after.pptx").exists()
    assert not (output / "report.json").exists()
    assert source.read_bytes() == deck
