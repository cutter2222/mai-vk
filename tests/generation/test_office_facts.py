"""Numeric shortening regression, including the actual LLM retry boundary."""

import io
import json
from zipfile import ZipFile

import pytest
from pptx import Presentation
from pptx.util import Inches

from presentation_designer.generation import office_edit, office_object_edit
from presentation_designer.generation.office_facts import (
    validate_shortening,
    validate_shortening_length,
)
from presentation_designer.generation.office_objects import objects
from presentation_designer.generation.office_preservation import verify_selected_text_edit
from presentation_designer.llm.cache import ResponseCache
from presentation_designer.llm.client import LlmClient
from presentation_designer.llm.limiter import LocalLimiter, Quota
from presentation_designer.llm.retry import RetryPolicy
from presentation_designer.llm.stub import StubTransport
from presentation_designer.llm.types import ResponseError
from presentation_designer.shared.settings import Settings, get_models_config

ORIGINAL = "Пользователи теряют ключевые сигналы в потоке из 27 сообщений в день."
GOOD = "В потоке 27 сообщений в день теряются важные уведомления."
BAD = "В потоке 27 сообщений теряются важные уведомления."


@pytest.mark.parametrize(
    "after",
    [
        BAD,
        GOOD.replace("27", "28"),
        GOOD.replace("сообщений", "писем"),
        GOOD.replace("в день", "в месяц"),
        GOOD + " Ещё 2 сигнала.",
        "Сигналы теряются.",
    ],
)
def test_shortening_rejects_changed_number_unit_or_period(after):
    with pytest.raises(ValueError, match="число, единицу или период"):
        validate_shortening("Сократи текст", ORIGINAL, after)


@pytest.mark.parametrize("instruction", ["Сократи текст", "Сделай короче", "Shorten this"])
def test_shortening_preserves_complete_expression(instruction):
    validate_shortening(instruction, ORIGINAL, GOOD)
    validate_shortening(instruction, ORIGINAL, GOOD.replace("27 сообщений", "27\u00a0сообщений"))


def test_non_shortening_edit_can_change_facts():
    validate_shortening("Замени 27 на 28", ORIGINAL, GOOD.replace("27", "28"))
    validate_shortening("Переведи на английский", ORIGINAL, "27 messages per day")


def test_period_cannot_move_to_another_quantity():
    with pytest.raises(ValueError):
        validate_shortening(
            "Сократи", "27 сообщений в день и 5 в месяц", "27 сообщений в месяц и 5 в день"
        )
    with pytest.raises(ValueError):
        validate_shortening("Сократи", "Рост 60% и 60%", "Рост 60%")
    with pytest.raises(ValueError):
        validate_shortening("Сократи", "Рост 2,5% в год", "Рост 2,5%")


@pytest.mark.parametrize("scope", ["deck", "object", "many"])
@pytest.mark.parametrize("recover", [True, False])
@pytest.mark.parametrize("split", [False, True])
async def test_invalid_shortening_retried_before_acceptance(
    scope, recover, split, tmp_path, monkeypatch
):
    deck = Presentation()
    shape = deck.slides.add_slide(deck.slide_layouts[6]).shapes.add_textbox(
        Inches(1), Inches(1), Inches(5), Inches(1)
    )
    texts = ["Пользователи теряют ключевые сигналы в потоке из ", "27", " сообщений", " в день."]
    for text in texts if split else [ORIGINAL]:
        shape.text_frame.paragraphs[0].add_run().text = text
    output = io.BytesIO()
    deck.save(output)
    original = output.getvalue()
    target = objects(original)[0]

    def answer(text):
        patches = [{"slide": 1, "run": 0, "before": ORIGINAL, "after": text}]
        if split:
            patches = (
                [{"slide": 1, "run": 3, "before": texts[3], "after": "."}]
                if text == BAD
                else [{"slide": 1, "run": 0, "before": texts[0], "after": "Потеря сигналов: "}]
            )
        plan = {
            "explanation": "Сократил",
            "patches": patches,
        }
        if scope == "many":
            return {
                "explanation": "Сократил",
                "edits": [
                    {"target": {"slide": target.slide, "shape_id": target.shape_id}, "plan": plan}
                ],
            }
        return plan

    settings = Settings()
    stub = StubTransport()
    stub.answer(answer(BAD), times=1)
    stub.answer(answer(GOOD if recover else BAD))
    client = LlmClient(
        settings=settings,
        models=get_models_config(),
        transport=stub,
        limiter=LocalLimiter(Quota(2, 1000, 1_000_000)),
        cache=ResponseCache("off", tmp_path / "cache", tmp_path / "fixtures"),
        retry_policy=RetryPolicy(max_retries=1, base_s=0, max_s=0),
    )
    module = office_edit if scope == "deck" else office_object_edit
    monkeypatch.setattr(module, "build_client", lambda settings: client)

    async def propose():
        if scope == "deck":
            return await module.propose(original, "Сократи текст", settings)
        if scope == "object":
            return await module.propose(original, "Сократи текст", settings, target)
        return await module.propose_many(original, "Сократи текст", settings, [target])

    if recover:
        plan = await propose()
        if scope == "many":
            plan = plan.edits[0].plan
        assert plan.patches[0].after == ("Потеря сигналов: " if split else GOOD)
        updated = office_edit.patch_pptx(original, plan)
        assert verify_selected_text_edit(original, updated, [target])["changed_text_runs"] == 1
        result = Presentation(io.BytesIO(updated)).slides[0].shapes[0]
        assert result.text == ("Потеря сигналов: 27 сообщений в день." if split else GOOD)
        assert len(result.text_frame.paragraphs[0].runs) == (4 if split else 1)
    else:
        with pytest.raises(ResponseError, match="число, единицу или период"):
            await propose()
    assert len(stub.calls) == 2
    assert "число, единицу или период" in stub.calls[1].messages[-1].text
    assert Presentation(io.BytesIO(original)).slides[0].shapes[0].text == ORIGINAL


@pytest.mark.parametrize(
    ("before", "after"),
    [
        ("Диапазон 10–20% в год", "10 и 20% в год"),
        ("Срок 2026-09-21", "21-09-2026"),
        ("Дата 21.09.2026", "21.09 и 2026"),
        ("Сумма 1 200 000 рублей в месяц", "1 200 000 рублей"),
        ("Бюджет 2,5 млн рублей", "2,5 млн"),
        ("Бюджет $500", "500"),
        ("Рост 60%", "60"),
    ],
)
def test_compound_numeric_expression_cannot_be_split_or_lose_unit(before, after):
    with pytest.raises(ValueError, match="число, единицу или период"):
        validate_shortening("Сократи", before, after)
    validate_shortening("Сократи", before, before.replace(" ", "\u00a0"))


@pytest.mark.parametrize(
    "quantity",
    [
        "10–20% в год",
        "21.09.2026",
        "2026-09-21",
        "1 200 000 рублей в месяц",
        "$500",
        "2,5 млн рублей",
        "60%",
    ],
)
def test_compound_expressions_allow_shorter_surrounding_text(quantity):
    validate_shortening("Сократи", f"Подробное описание показателя: {quantity}", quantity)


@pytest.mark.parametrize(
    ("prefix", "suffix"),
    [
        ("10", "–20% в год"),
        ("21", ".09.2026"),
        ("1 ", "200 000 рублей в месяц"),
        ("2,5", " млн рублей"),
        ("60", "%"),
    ],
)
def test_compound_expression_split_across_runs_keeps_suffix(prefix, suffix):
    deck = Presentation()
    shape = deck.slides.add_slide(deck.slide_layouts[6]).shapes.add_textbox(
        Inches(1), Inches(1), Inches(5), Inches(1)
    )
    for text in [prefix, suffix]:
        shape.text_frame.paragraphs[0].add_run().text = text
    output = io.BytesIO()
    deck.save(output)
    plan = office_edit.EditPlan(
        explanation="Сократил",
        patches=[
            office_edit.TextPatch(slide=1, run=1, before=suffix, after=""),
        ],
    )
    with pytest.raises(ValueError, match="число, единицу или период"):
        plan.validate_facts("Сократи", output.getvalue())


@pytest.mark.parametrize("fragment", ["сообщений ", "в день", "7 сообщений в день"])
def test_paragraph_context_rejects_loss_in_non_numeric_or_partial_numeric_run(fragment):
    deck = Presentation()
    shape = deck.slides.add_slide(deck.slide_layouts[6]).shapes.add_textbox(
        Inches(1), Inches(1), Inches(5), Inches(1)
    )
    texts = {
        "сообщений ": ["27 ", "сообщений ", "в день"],
        "в день": ["27", " сообщений ", "в день"],
        "7 сообщений в день": ["2", "7 сообщений в день"],
    }[fragment]
    for text in texts:
        shape.text_frame.paragraphs[0].add_run().text = text
    output = io.BytesIO()
    deck.save(output)
    plan = office_edit.EditPlan(
        explanation="Сократил",
        patches=[
            office_edit.TextPatch(
                slide=1,
                run=texts.index(fragment),
                before=fragment,
                after="7 сообщений" if fragment.startswith("7") else "",
            )
        ],
    )
    with pytest.raises(ValueError, match="число, единицу или период"):
        plan.validate_facts("Сократи", output.getvalue())


def test_period_cannot_be_transferred_between_paragraphs():
    deck = Presentation()
    shape = deck.slides.add_slide(deck.slide_layouts[6]).shapes.add_textbox(
        Inches(1), Inches(1), Inches(5), Inches(1)
    )
    for paragraph, text in [
        (shape.text_frame.paragraphs[0], "27 сообщений"),
        (shape.text_frame.add_paragraph(), "5 сообщений"),
    ]:
        paragraph.add_run().text = text
        paragraph.add_run().text = " в день" if text.startswith("27") else " в месяц"
    output = io.BytesIO()
    deck.save(output)
    plan = office_edit.EditPlan(
        explanation="Сократил",
        patches=[
            office_edit.TextPatch(slide=1, run=1, before=" в день", after=" в месяц"),
            office_edit.TextPatch(slide=1, run=3, before=" в месяц", after=" в день"),
        ],
    )
    with pytest.raises(ValueError, match="число, единицу или период"):
        plan.validate_facts("Сократи", output.getvalue())


def test_soft_break_is_whitespace_and_numeric_free_wording_can_shorten():
    deck = Presentation()
    shape = deck.slides.add_slide(deck.slide_layouts[6]).shapes.add_textbox(
        Inches(1), Inches(1), Inches(5), Inches(1)
    )
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.add_run().text = "Очень длинное пояснение: 27"
    paragraph.add_line_break()
    paragraph.add_run().text = "сообщений в день"
    output = io.BytesIO()
    deck.save(output)
    original = output.getvalue()
    plan = office_edit.EditPlan(
        explanation="Сократил",
        patches=[
            office_edit.TextPatch(
                slide=1, run=0, before="Очень длинное пояснение: 27", after="Поток: 27"
            ),
        ],
    )
    plan.validate_facts("Сократи", original)
    plan.patches = [
        office_edit.TextPatch(slide=1, run=1, before="сообщений в день", after="сообщений")
    ]
    with pytest.raises(ValueError, match="число, единицу или период"):
        plan.validate_facts("Сократи", original)


def test_shortening_cannot_expand_a_list_item_by_duplicating_adjacent_run():
    deck = Presentation()
    shape = deck.slides.add_slide(deck.slide_layouts[6]).shapes.add_textbox(
        Inches(1), Inches(1), Inches(5), Inches(3)
    )
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.add_run().text = "Меняйте скриншоты "
    paragraph.add_line_break()
    paragraph.add_run().text = "на том же месте. Остальные элементы не смещайте."
    shape.text_frame.add_paragraph().text = "Для навигации используем точки."
    output = io.BytesIO()
    deck.save(output)
    plan = office_edit.EditPlan(
        explanation="Сократил",
        patches=[
            office_edit.TextPatch(
                slide=1,
                run=0,
                before="Меняйте скриншоты ",
                after="Меняйте скриншоты на том же месте. Остальные элементы не смещайте.",
            ),
        ],
    )
    with pytest.raises(ValueError, match="увеличило длину абзаца"):
        plan.validate_facts("Сократи список", output.getvalue())
    plan.validate_facts("Переформулируй список", output.getvalue())


def test_shortening_length_normalizes_whitespace_and_allows_safe_refusal():
    validate_shortening_length("Сократи", "27 сообщений в день", "27  сообщений\nв день")
    validate_shortening_length("Сократи", ORIGINAL, ORIGINAL)
    validate_shortening_length("Сократи", ORIGINAL, GOOD)
    with pytest.raises(ValueError, match="увеличило длину абзаца"):
        validate_shortening_length("Сократи", GOOD, ORIGINAL)


def test_paragraph_context_keeps_run_ids_spaces_breaks_and_skips_empty_paragraphs():
    from lxml import etree

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(0), Inches(0), Inches(1), Inches(1)).text = "Вне выбора"
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(3))
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.add_run().text = "Первый пункт "
    paragraph.add_line_break()
    paragraph.add_run().text = "продолжается."
    shape.text_frame.add_paragraph()
    shape.text_frame.add_paragraph().text = "Второй пункт."
    output = io.BytesIO()
    deck.save(output)
    with ZipFile(io.BytesIO(output.getvalue())) as archive:
        root = etree.fromstring(archive.read(office_edit.slides(archive)[0]))
    assert office_edit.paragraph_context(root, [1, 2, 3]) == [
        {
            "parts": [
                {"run": 1, "text": "Первый пункт "},
                {"break": "\n"},
                {"run": 2, "text": "продолжается."},
            ]
        },
        {"parts": [{"run": 3, "text": "Второй пункт."}]},
    ]


@pytest.mark.parametrize("scope", ["deck", "object", "many"])
async def test_stale_run_repair_gets_exact_text_and_paragraph_context(scope, tmp_path, monkeypatch):
    deck = Presentation()
    shape = deck.slides.add_slide(deck.slide_layouts[6]).shapes.add_textbox(
        Inches(1), Inches(1), Inches(5), Inches(3)
    )
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.add_run().text = "Длинное пояснение: "
    paragraph.add_line_break()
    paragraph.add_run().text = "27 сообщений в день."
    output = io.BytesIO()
    deck.save(output)
    original = output.getvalue()
    target = objects(original)[0]

    def answer(before):
        plan = {
            "explanation": "Сократил",
            "patches": [{"slide": 1, "run": 0, "before": before, "after": "Поток: "}],
        }
        return (
            {
                "explanation": "Сократил",
                "edits": [{"target": {"slide": 1, "shape_id": target.shape_id}, "plan": plan}],
            }
            if scope == "many"
            else plan
        )

    stub = StubTransport()
    stub.answer(answer("Длинное пояснение:"), times=1)
    stub.answer(answer("Длинное пояснение: "))
    settings = Settings()
    client = LlmClient(
        settings=settings,
        models=get_models_config(),
        transport=stub,
        limiter=LocalLimiter(Quota(2, 1000, 1_000_000)),
        cache=ResponseCache("off", tmp_path / "cache", tmp_path / "fixtures"),
        retry_policy=RetryPolicy(max_retries=1, base_s=0, max_s=0),
    )
    module = office_edit if scope == "deck" else office_object_edit
    monkeypatch.setattr(module, "build_client", lambda settings: client)
    if scope == "deck":
        plan = await module.propose(original, "Сократи", settings)
    elif scope == "object":
        plan = await module.propose(original, "Сократи", settings, target)
    else:
        plan = (await module.propose_many(original, "Сократи", settings, [target])).edits[0].plan
    assert len(stub.calls) == 2
    assert '"Длинное пояснение: "' in stub.calls[1].messages[-1].text
    context = json.loads(stub.calls[0].messages[-1].text)
    content = (
        context["slides"][0]
        if scope == "deck"
        else (context["objects"][0] if scope == "many" else context)
    )
    assert content["paragraphs"][0]["parts"] == [
        {"run": 0, "text": "Длинное пояснение: "},
        {"break": "\n"},
        {"run": 1, "text": "27 сообщений в день."},
    ]
    updated = office_edit.patch_pptx(original, plan)
    assert verify_selected_text_edit(original, updated, [target])["changed_text_runs"] == 1
    assert Presentation(io.BytesIO(updated)).slides[0].shapes[0].text == (
        "Поток: \v27 сообщений в день."
    )
