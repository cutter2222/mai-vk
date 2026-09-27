"""Patches preserve manual work without reconstructing a deck from its old plan."""

import io
from zipfile import ZipFile

import pytest
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt

from presentation_designer.generation.office_edit import NS, EditPlan, TextPatch, patch_pptx


@pytest.fixture
def manual_deck():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(2), Inches(3), Inches(4), Inches(1))
    run = shape.text_frame.paragraphs[0].add_run()
    run.text = "Ручной заголовок"
    run.font.bold = True
    run.font.size = Pt(27)
    shape.text_frame.add_paragraph().text = "Не менять вручную добавленный текст"
    deck.slides.add_slide(deck.slide_layouts[6]).shapes.add_textbox(
        Inches(1), Inches(1), Inches(3), Inches(1)
    ).text = "Ручной соседний слайд"
    output = io.BytesIO()
    deck.save(output)
    return output.getvalue()


def test_only_requested_text_changes(manual_deck):
    plan = EditPlan(
        explanation="Заголовок изменён",
        patches=[TextPatch(slide=1, run=0, before="Ручной заголовок", after="Новый заголовок")],
    )
    result = patch_pptx(manual_deck, plan)
    with ZipFile(io.BytesIO(manual_deck)) as before, ZipFile(io.BytesIO(result)) as after:
        assert before.namelist() == after.namelist()
        for name in before.namelist():
            if name != "ppt/slides/slide1.xml":
                assert before.read(name) == after.read(name), name
        old = etree.fromstring(before.read("ppt/slides/slide1.xml"))
        new = etree.fromstring(after.read("ppt/slides/slide1.xml"))
        new.find(".//a:t", NS).text = "Ручной заголовок"
        assert etree.tostring(old, method="c14n") == etree.tostring(new, method="c14n")
    deck = Presentation(io.BytesIO(result))
    shape = deck.slides[0].shapes[0]
    assert shape.left == Inches(2)
    assert shape.top == Inches(3)
    assert shape.text_frame.paragraphs[0].runs[0].font.bold
    assert shape.text_frame.paragraphs[0].runs[0].font.size == Pt(27)
    assert deck.slides[1].shapes[0].text == "Ручной соседний слайд"


@pytest.mark.parametrize(
    "changes",
    [
        {"before": "Устаревший текст"},
        {"slide": 3},
        {"run": 100},
    ],
)
def test_wrong_patch_rejected(manual_deck, changes):
    patch = {"slide": 1, "run": 0, "before": "Ручной заголовок", "after": "Новый"}
    with pytest.raises(ValueError):
        patch_pptx(manual_deck, EditPlan(explanation="", patches=[TextPatch(**(patch | changes))]))


def test_noop_is_byte_identical(manual_deck):
    assert (
        patch_pptx(manual_deck, EditPlan(explanation="Не поддерживается", patches=[]))
        == manual_deck
    )


@pytest.mark.parametrize("round_coordinates", [False, True])
def test_selected_object_move_preserves_other_bytes(manual_deck, round_coordinates):
    from presentation_designer.generation.office_object_edit import (
        ObjectEditPlan,
        Position,
        patch_object,
    )
    from presentation_designer.generation.office_objects import objects

    obj = objects(manual_deck)[0]
    assert obj.runs == [0, 1]
    assert obj.label.startswith("Ручной заголовок")
    updated = patch_object(
        manual_deck,
        obj,
        ObjectEditPlan(
            explanation="Правее",
            patches=[],
            position=Position(x=0.4, y=round(obj.bbox.y, 6) if round_coordinates else obj.bbox.y),
        ),
    )
    with ZipFile(io.BytesIO(manual_deck)) as before, ZipFile(io.BytesIO(updated)) as after:
        for name in before.namelist():
            if name != "ppt/slides/slide1.xml":
                assert before.read(name) == after.read(name)
        old = etree.fromstring(before.read("ppt/slides/slide1.xml"))
        new = etree.fromstring(after.read("ppt/slides/slide1.xml"))
        new.find(".//a:off", NS).set("x", old.find(".//a:off", NS).get("x"))
        assert etree.tostring(old, method="c14n") == etree.tostring(new, method="c14n")
    shape = Presentation(io.BytesIO(updated)).slides[0].shapes[0]
    assert shape.left == Inches(4)
    assert shape.text_frame.paragraphs[0].runs[0].font.bold


def test_selected_object_rejects_other_text_and_off_slide_position(manual_deck):
    from presentation_designer.generation.office_object_edit import (
        ObjectEditPlan,
        Position,
        patch_object,
    )
    from presentation_designer.generation.office_objects import ObjectTarget, objects

    obj = objects(manual_deck)[0]
    with pytest.raises(ValueError, match="вне выбранного"):
        patch_object(
            manual_deck,
            obj,
            ObjectEditPlan(
                explanation="",
                patches=[
                    TextPatch(
                        slide=2,
                        run=0,
                        before="Ручной соседний слайд",
                        after="Нельзя",
                    )
                ],
            ),
        )
    with pytest.raises(ValueError, match="границы"):
        patch_object(
            manual_deck,
            obj,
            ObjectEditPlan(explanation="", patches=[], position=Position(x=0.9, y=0.2)),
        )
    with pytest.raises(ValueError, match="не найден"):
        patch_object(
            manual_deck,
            ObjectTarget(slide=1, shape_id="999"),
            ObjectEditPlan(explanation="", patches=[]),
        )
    assert (
        patch_object(
            manual_deck,
            obj,
            ObjectEditPlan(
                explanation="", patches=[], position=Position(x=obj.bbox.x, y=obj.bbox.y)
            ),
        )
        == manual_deck
    )


@pytest.fixture
def structured_deck():
    """Фигура во вложенной группе с масштабом (система координат группы вдвое крупнее),
    заголовок-плейсхолдер с рамкой из макета, повёрнутая надпись и соседний слайд."""
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Заголовок из макета"
    outer = slide.shapes.add_group_shape()
    inner = outer.shapes.add_group_shape()
    box = inner.shapes.add_textbox(Inches(2), Inches(0), Inches(2), Inches(1))
    box.name = "Шаг 2"
    box.text = "Тестирование"
    inner.shapes.add_textbox(Inches(0), Inches(0), Inches(2), Inches(1)).text = "Анализ"
    # Группа занимает на слайде 2×1 дюйма, её дети живут в системе 4×2: масштаб 0,5.
    transform = outer._element.find("p:grpSpPr/a:xfrm", NS)
    transform.find("a:off", NS).attrib.update({"x": str(Inches(1)), "y": str(Inches(4))})
    transform.find("a:ext", NS).attrib.update({"cx": str(Inches(2)), "cy": str(Inches(1))})
    transform.find("a:chOff", NS).attrib.update({"x": "0", "y": "0"})
    transform.find("a:chExt", NS).attrib.update({"cx": str(Inches(4)), "cy": str(Inches(2))})
    turned = slide.shapes.add_textbox(Inches(6), Inches(1), Inches(2), Inches(1))
    turned.text = "Повёрнут"
    turned.rotation = 30
    deck.slides.add_slide(deck.slide_layouts[6]).shapes.add_textbox(
        Inches(1), Inches(1), Inches(3), Inches(1)
    ).text = "Соседний слайд"
    output = io.BytesIO()
    deck.save(output)
    return output.getvalue()


def test_object_map_sees_groups_placeholders_and_rotated_shapes(structured_deck):
    from presentation_designer.generation.office_objects import objects

    found = {obj.label: obj for obj in objects(structured_deck)}
    title = found["Заголовок из макета"]
    assert title.placeholder == "title" and title.bbox.width > 0.5
    step = found["Тестирование"]
    assert step.name == "Шаг 2" and len(step.group_path) == 2 and step.kind == "sp"
    # 2 дюйма в системе группы — 1 дюйм на слайде, от левого края группы (1 дюйм).
    assert step.bbox.x * 10 == pytest.approx(2) and step.bbox.width * 10 == pytest.approx(1)
    assert found["Повёрнут"].rotation == 30
    groups = [obj for obj in objects(structured_deck) if obj.kind == "grpSp"]
    assert len(groups) == 2
    assert groups[0].label == "Тестирование Анализ" and sorted(groups[0].runs) == sorted(
        step.runs + found["Анализ"].runs
    )


def assert_only_slide1_changed(before_bytes, after_bytes):
    with ZipFile(io.BytesIO(before_bytes)) as before, ZipFile(io.BytesIO(after_bytes)) as after:
        assert before.namelist() == after.namelist()
        for name in before.namelist():
            if name != "ppt/slides/slide1.xml":
                assert before.read(name) == after.read(name), name


@pytest.mark.parametrize("label", ["Тестирование", "Заголовок из макета", "Повёрнут"])
def test_text_edit_and_move_of_group_child_placeholder_and_rotated(structured_deck, label):
    from presentation_designer.generation.office_object_edit import (
        ObjectEditPlan,
        Position,
        patch_object,
    )
    from presentation_designer.generation.office_objects import objects

    obj = next(o for o in objects(structured_deck) if o.label == label)
    text = patch_object(
        structured_deck,
        obj,
        ObjectEditPlan(
            explanation="",
            patches=[TextPatch(slide=1, run=obj.runs[0], before=label, after=label + "!")],
        ),
    )
    assert_only_slide1_changed(structured_deck, text)
    assert any(o.label == label + "!" for o in objects(text))
    x, y = min(0.5, 0.99 - obj.bbox.width), min(0.6, 0.99 - obj.bbox.height)
    moved = patch_object(
        structured_deck,
        obj,
        ObjectEditPlan(explanation="", patches=[], position=Position(x=x, y=y)),
    )
    assert_only_slide1_changed(structured_deck, moved)
    after = next(o for o in objects(moved) if o.shape_id == obj.shape_id)
    assert (after.bbox.x, after.bbox.y) == (pytest.approx(x), pytest.approx(y))
    assert (after.bbox.width, after.bbox.height) == (
        pytest.approx(obj.bbox.width),
        pytest.approx(obj.bbox.height),
    )
    assert after.rotation == obj.rotation and after.group_path == obj.group_path
    others = [o for o in objects(moved) if o.shape_id != obj.shape_id and not o.group_path]
    assert others == [
        o for o in objects(structured_deck) if o.shape_id != obj.shape_id and not o.group_path
    ]


def test_live_selection_finds_group_child_by_group_relative_box(structured_deck):
    from presentation_designer.generation.office_objects import LiveBox, LiveTarget, resolve_live

    deck = Presentation(io.BytesIO(structured_deck))
    # Второе «Шаг 2» верхнего уровня: имя совпадает, но рамка от угла слайда.
    deck.slides[0].shapes.add_textbox(Inches(5), Inches(4), Inches(1), Inches(0.5)).name = "Шаг 2"
    output = io.BytesIO()
    deck.save(output)
    data = output.getvalue()
    # ONLYOFFICE отдаёт у фигуры в группе положение от левого верхнего угла группы.
    live = LiveTarget(
        slide=1,
        name="Шаг 2",
        in_group=True,
        box=LiveBox(x=25.4, y=0, width=25.4, height=12.7),
    )
    target = resolve_live(data, live)
    from presentation_designer.generation.office_objects import objects

    chosen = next(o for o in objects(data) if o.shape_id == target.shape_id)
    assert chosen.group_path and chosen.label == "Тестирование"
    top = resolve_live(
        data,
        LiveTarget(slide=1, name="Шаг 2", box=LiveBox(x=127, y=101.6, width=25.4, height=12.7)),
    )
    assert next(o for o in objects(data) if o.shape_id == top.shape_id).group_path == []


@pytest.mark.parametrize("fault", ["moved", "duplicate", "no_box", "wrong_group"])
def test_live_selection_rejects_stale_or_ambiguous_object(fault):
    from presentation_designer.generation.office_objects import LiveBox, LiveTarget, resolve_live

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    shape.name = "Selected"
    shape.text = "Neighbour must not change"
    if fault in ("duplicate", "no_box"):
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).name = shape.name
    output = io.BytesIO()
    deck.save(output)
    live = LiveTarget(
        slide=1,
        name=shape.name,
        in_group=fault == "wrong_group",
        box=None
        if fault == "no_box"
        else LiveBox(x=80 if fault == "moved" else 25.4, y=25.4, width=50.8, height=25.4),
    )
    with pytest.raises(ValueError, match=r"выделите.*заново"):
        resolve_live(output.getvalue(), live)


def test_live_selection_survives_shape_renumbering():
    from presentation_designer.generation.office_objects import LiveBox, LiveTarget, resolve_live

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    shape.name = "Selected"
    shape._element.find(".//{*}cNvPr").set("id", "99")
    output = io.BytesIO()
    deck.save(output)
    target = resolve_live(
        output.getvalue(),
        LiveTarget(
            slide=1,
            name=shape.name,
            box=LiveBox(x=25.4, y=25.4, width=50.8, height=25.4),
        ),
    )
    assert target.shape_id == "99"


async def test_object_proposal_receives_only_selected_text(manual_deck, monkeypatch):
    import json
    from types import SimpleNamespace

    from presentation_designer.generation import office_object_edit
    from presentation_designer.generation.office_objects import objects
    from presentation_designer.shared.settings import Settings

    obj = objects(manual_deck)[0]
    calls = []

    class Client:
        async def complete(self, request, validator):
            context = json.loads(request.messages[-1].text)
            assert context["object"]["shape_id"] == obj.shape_id
            assert context["instruction"] == "Правее"
            assert context["runs"] == [
                {"run": 0, "text": "Ручной заголовок"},
                {"run": 1, "text": "Не менять вручную добавленный текст"},
            ]
            plan = validator(
                {"explanation": "Правее", "patches": [], "position": {"x": 0.3, "y": 0.4}}
            )
            return SimpleNamespace(parsed=plan.model_dump())

        async def aclose(self):
            calls.append("closed")

    monkeypatch.setattr(office_object_edit, "build_client", lambda settings: Client())
    plan = await office_object_edit.propose(manual_deck, "Правее", Settings(), obj)
    assert plan.position.x == 0.3
    assert calls == ["closed"]


async def test_multi_object_proposal_has_shared_context_and_rejects_escape(
    manual_deck, monkeypatch
):
    import json
    from types import SimpleNamespace

    from presentation_designer.generation import office_object_edit as edit
    from presentation_designer.generation.office_objects import ObjectTarget, objects
    from presentation_designer.shared.settings import Settings

    chosen = [ObjectTarget(slide=obj.slide, shape_id=obj.shape_id) for obj in objects(manual_deck)]
    closed = []

    class Client:
        async def complete(self, request, validator):
            content = json.loads(request.messages[-1].text)
            assert content["instruction"] == "Оба правее"
            assert len(content["objects"]) == 2
            assert content["objects"][0]["runs"][0]["text"] == "Ручной заголовок"
            return SimpleNamespace(parsed=validator({"explanation": "Без изменений", "edits": []}))

        async def aclose(self):
            closed.append(True)

    monkeypatch.setattr(edit, "build_client", lambda settings: Client())
    plan = await edit.propose_many(manual_deck, "Оба правее", Settings(), chosen)
    assert edit.patch_objects(manual_deck, chosen, plan) == manual_deck
    assert closed == [True]
    change = edit.SelectedObjectEdit(
        target=chosen[1], plan=edit.ObjectEditPlan(explanation="", patches=[])
    )
    with pytest.raises(ValueError, match="вне выбора"):
        edit.patch_objects(
            manual_deck, chosen[:1], edit.ObjectsEditPlan(explanation="", edits=[change])
        )
    with pytest.raises(ValueError, match="повторно"):
        edit.patch_objects(
            manual_deck, chosen, edit.ObjectsEditPlan(explanation="", edits=[change, change])
        )
