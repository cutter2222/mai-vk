"""Вариант original: загруженная презентация как готовый результат. Смысловой план и план
слайдов строятся без модели по профилю и пакету, композер сохраняет объекты образца: тексты
и число слайдов первой ревизии совпадают с файлом, правка одного блока меняет только его."""

from __future__ import annotations

import pathlib
import shutil
from copy import deepcopy
from typing import Any

import pytest
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER as PP
from pptx.oxml.ns import qn

from presentation_designer.contracts import ContentPackage, SlidePlan, StoryPlan, TemplateProfile
from presentation_designer.contracts.validators import check_slide_plan, check_story_plan
from presentation_designer.generation.original import (
    deck_patterns,
    original_plan,
    original_profile,
    original_story,
    skipped_slides,
)
from presentation_designer.layout.compose import compose_deck
from presentation_designer.layout.text import fill_empty_placeholders
from presentation_designer.parsing.content.importer import import_content
from presentation_designer.parsing.content.parsers import ParseCache
from presentation_designer.shared.settings import Settings
from tests.layout.conftest import MINI_TEMPLATE, own_profile
from tests.parsing.content.conftest import material


@pytest.fixture(scope="module")
def mini_profile() -> dict[str, Any]:
    return own_profile(MINI_TEMPLATE, "tpl_mini")


def _texts(path: pathlib.Path) -> list[list[str]]:
    prs = Presentation(str(path))
    out: list[list[str]] = []
    for slide in prs.slides:
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                texts.append(" ".join(shape.text_frame.text.split()))
        out.append(sorted(texts))
    return out


def _package(tmp_path: pathlib.Path, import_settings: Settings) -> dict[str, Any]:
    result = import_content(
        [material(MINI_TEMPLATE)],
        {"purpose": "other", "title": "Мини-шаблон"},
        package_id="pkg_orig",
        settings=import_settings,
        cache=ParseCache(tmp_path / "import-cache"),
        use_model=False,
    )
    return dict(result.package)


def test_original_story_and_plan_are_valid_and_cover_every_slide(
    mini_profile: dict[str, Any], tmp_path: pathlib.Path, import_settings: Settings
) -> None:
    package = _package(tmp_path, import_settings)
    story = original_story(package, mini_profile, {"language": "ru"})
    assert not check_story_plan(
        StoryPlan.model_validate(story), ContentPackage.model_validate(package)
    )
    plan = original_plan(story, mini_profile, package, plan_id="plan_orig")
    violations = check_slide_plan(
        SlidePlan.model_validate(plan),
        TemplateProfile.model_validate(mini_profile),
        ContentPackage.model_validate(package),
        StoryPlan.model_validate(story),
    )
    # Обязательные слоты таблицы и диаграммы заполнены самим образцом — композер их сохраняет.
    assert {v.code for v in violations} <= {"slot_required"}, violations
    slides = plan["slides"]
    # Слайд на каждый образец шаблона; собственные композиции библиотеки образцами не считаются.
    samples = [p for p in mini_profile["patterns"] if p["source"]["kind"] == "sample_slide"]
    assert len(slides) == len(story["theses"]) == len(samples)
    assert plan["variant"]["variant_id"] == "original" and plan["slide_count"]["exact"] == len(
        slides
    )
    assert all(b["fit"]["action"] == "as_is" for s in slides for b in s["blocks"])
    assert plan["coverage"]["missing"] == []
    assert not plan["generation_meta"]["models"], "модель не участвует"


def test_original_compose_keeps_texts_and_edit_changes_only_its_block(
    mini_profile: dict[str, Any], tmp_path: pathlib.Path, import_settings: Settings
) -> None:
    package = _package(tmp_path, import_settings)
    story = original_story(package, mini_profile, {"language": "ru"})
    plan = original_plan(story, mini_profile, package, plan_id="plan_orig")
    out = tmp_path / "r1.pptx"
    result = compose_deck(plan, mini_profile, MINI_TEMPLATE, package, out_pptx=out, job_id="job_o")
    assert result.integrity.ok
    assert out.read_bytes() == MINI_TEMPLATE.read_bytes()
    assert _texts(out) == _texts(MINI_TEMPLATE)
    assert all(not s.get("removed_object_ids") for s in result.deck["slides"])
    assert all(
        o.get("content_source") in ("sample", "template", "static")
        for s in result.deck["slides"]
        for o in s.get("objects", [])
        if o.get("content_source")
    )

    # Правка из чата: один блок с новым текстом, остальное остаётся объектами образца.
    edited = {**plan, "slides": [dict(s) for s in plan["slides"]]}
    target = next(s for s in edited["slides"] if s["blocks"])
    block = dict(target["blocks"][0])
    block["text"] = "Новый заголовок после правки"
    target["blocks"] = [block, *target["blocks"][1:]]
    out2 = tmp_path / "r2.pptx"
    result2 = compose_deck(
        edited, mini_profile, MINI_TEMPLATE, package, out_pptx=out2, job_id="job_o"
    )
    assert result2.integrity.ok
    before, after = _texts(MINI_TEMPLATE), _texts(out2)
    changed = [i for i, (a, b) in enumerate(zip(before, after, strict=True)) if a != b]
    assert changed == [target["order"] - 1]
    assert any("Новый заголовок после правки" in t for t in after[target["order"] - 1])


def test_real_layers_plan_and_compose_original_without_model(
    mini_profile: dict[str, Any], tmp_path: pathlib.Path, import_settings: Settings
) -> None:
    """Слой плана строит original без провайдера моделей, слой вёрстки пропускает правку по
    фактам и кладёт ревизию в артефакты; тексты совпадают с файлом."""
    from presentation_designer.pipeline.artifacts import ArtifactStore
    from presentation_designer.pipeline.real import RealLayers
    from presentation_designer.pipeline.run import ComposeInput, PlanInput

    package = _package(tmp_path, import_settings)
    story = original_story(package, mini_profile, {"language": "ru"})
    settings = Settings()
    settings.paths.artifacts_dir = tmp_path / "artifacts"
    layers = RealLayers(settings)
    layers._llm_failed = True  # провайдер не настроен: original обязан обходиться без него
    plan = layers.plan(
        PlanInput("job_o", "original", mini_profile, package, story, {"language": "ru"}, None)
    )
    assert plan["variant"]["variant_id"] == "original" and len(plan["slides"]) == 4
    store = ArtifactStore(settings.artifacts_dir)
    with store.stage_revision("job_o", "original", 1) as staging:
        out = layers.compose(
            ComposeInput(
                "job_o", "original", 1, plan, mini_profile, MINI_TEMPLATE, package, story, staging
            )
        )
    assert layers.last_feedback_report is None or not layers.last_feedback_report.get("applied")
    deck_path = store.revision_dir("job_o", "original", 1) / "deck.pptx"
    assert out.slide_count == 4 and _texts(deck_path) == _texts(MINI_TEMPLATE)


def test_edit_validation_tolerates_sample_filled_required_slots(
    mini_profile: dict[str, Any], tmp_path: pathlib.Path, import_settings: Settings
) -> None:
    """Проверка плана после правки: у original обязательный слот без блока (таблица,
    диаграмма, пустой в файле заголовок) — не нарушение, у вариантов вёрстки — нарушение."""
    from presentation_designer.generation.edit import EditError, validate_plan

    package = _package(tmp_path, import_settings)
    story = original_story(package, mini_profile, {"language": "ru"})
    plan = original_plan(story, mini_profile, package, plan_id="plan_orig")
    plan["slides"][0]["blocks"] = [b for b in plan["slides"][0]["blocks"] if b["kind"] != "title"]
    validate_plan(plan, mini_profile, package, story)
    as_balanced = {**plan, "variant": {**plan["variant"], "variant_id": "balanced"}}
    with pytest.raises(EditError):
        validate_plan(as_balanced, mini_profile, package, story)


def test_unrecognized_blank_hidden_slides_survive_import_and_edit(
    tmp_path: pathlib.Path, import_settings: Settings
) -> None:
    from presentation_designer.generation.edit import validate_plan
    from presentation_designer.pipeline.real import RealLayers
    from presentation_designer.pipeline.run import PlanInput

    source = tmp_path / "source.pptx"
    prs = Presentation(str(MINI_TEMPLATE))
    prs.slides[1]._element.set("show", "0")
    prs.slides[1].notes_slide.notes_text_frame.text = "Сохранить заметки скрытого слайда"
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(str(source))
    profile = own_profile(source, "tpl_import")
    # Reproduce classification exclusions deterministically, including an interior page.
    profile["patterns"] = [
        p for p in profile["patterns"] if p["source"].get("slide_index") not in (2, 5)
    ]
    before = deepcopy(profile)
    package = _package(tmp_path, import_settings)
    story = original_story(package, profile, {"language": "ru"})
    layers = RealLayers(Settings())
    layers._llm_failed = True
    plan = layers.plan(PlanInput("job_o", "original", profile, package, story, {}, None))
    assert len(plan["slides"]) == len(story["theses"]) == 5
    assert [p["source"]["slide_index"] for p in deck_patterns(profile)] == [1, 2, 3, 4, 5]
    assert skipped_slides(profile) == []
    validate_plan(plan, profile, package, story)
    TemplateProfile.model_validate(original_profile(profile))
    out = tmp_path / "original.pptx"
    result = compose_deck(plan, profile, source, package, out_pptx=out, prune_layouts=True)
    assert result.integrity.ok
    assert len(result.deck["slides"]) == 5
    assert out.read_bytes() == source.read_bytes()
    assert profile == before, "обычная генерация не должна получать служебные композиции"

    edited = deepcopy(plan)
    edited["slides"][0]["blocks"][0]["text"] = "Точечная правка"
    revised = tmp_path / "edited.pptx"
    compose_deck(edited, profile, source, package, out_pptx=revised)
    actual = Presentation(str(revised))
    assert len(actual.slides) == 5
    assert _texts(revised)[1:] == _texts(source)[1:]
    assert actual.slides[1]._element.get("show") == "0"
    assert (
        actual.slides[1].notes_slide.notes_text_frame.text
        == prs.slides[1].notes_slide.notes_text_frame.text
    )
    for old, new in zip(prs.slides, actual.slides, strict=True):
        if old == prs.slides[0]:
            continue
        assert old._element.xml == new._element.xml


def _has_empty_text_placeholder(slide: Any) -> bool:
    return any(
        ph.has_text_frame
        and not ph.text_frame.text.strip()
        and ph.placeholder_format.type
        in (PP.TITLE, PP.CENTER_TITLE, PP.SUBTITLE, PP.BODY, PP.OBJECT)
        for ph in slide.placeholders
    )


@pytest.mark.organizer_data
def test_vk_education_original_keeps_slides_and_turns_prompts_into_text(
    tmp_path: pathlib.Path,
) -> None:
    source = (
        pathlib.Path(__file__).resolve().parents[2]
        / "data/organizers/Шаблон презентации VK Education.pptx"
    )
    if not source.is_file():
        pytest.skip("закрытый шаблон организаторов отсутствует")
    profile = own_profile(source, "tpl_vk")
    package = {"package_id": "pkg_vk", "facts": [], "blocks": []}
    story = original_story(package, profile, {"language": "ru"})
    plan = original_plan(story, profile, package, plan_id="plan_vk")
    total = len(Presentation(str(source)).slides)
    assert len(plan["slides"]) == len(story["theses"]) == total
    assert [p["source"]["slide_index"] for p in deck_patterns(profile)] == list(range(1, total + 1))
    out = tmp_path / "original.pptx"
    result = compose_deck(plan, profile, source, package, out_pptx=out)
    assert result.integrity.ok and len(result.deck["slides"]) == total
    # Титул VK Education — пустые заголовок и текст: редактор подписывал их подсказкой, которая
    # исчезала по щелчку и не попадала в превью. Теперь это текст с оформлением макета.
    before, after = Presentation(str(source)), Presentation(str(out))
    first = {ph.placeholder_format.type: ph.text_frame.text for ph in after.slides[0].placeholders}
    assert first == {PP.TITLE: "Заголовок слайда", PP.BODY: "Текст слайда"}
    assert not any(_has_empty_text_placeholder(s) for s in after.slides)
    first_texts = [(o.get("text") or {}).get("plain") for o in result.deck["slides"][0]["objects"]]
    assert "Заголовок слайда" in first_texts
    # Остальные страницы переносятся как есть.
    kept = [
        (old, new)
        for old, new in zip(before.slides, after.slides, strict=True)
        if not _has_empty_text_placeholder(old)
    ]
    assert len(kept) > total // 2
    assert all(old._element.xml == _without_new_pictures(new, old) for old, new in kept)
    # Знак шаблона перенесён с макетов на слайды: в редакторе его можно двигать.
    assert result.report["counts"].get("logos_placed", 0) > 0


def _without_new_pictures(new: Any, old: Any) -> str:
    """XML слайда без картинок, которых не было в исходном: знака, перенесённого с макета."""
    ids = {c.get("id") for c in old._element.iter(qn("p:cNvPr"))}
    element = deepcopy(new._element)
    for pic in list(element.iter(qn("p:pic"))):
        if pic.find(f"{qn('p:nvPicPr')}/{qn('p:cNvPr')}").get("id") not in ids:
            pic.getparent().remove(pic)
    return str(element.xml)


def test_empty_placeholders_become_the_prompt_text_the_editor_shows() -> None:
    prs = Presentation(str(MINI_TEMPLATE))
    samples = [s._element.xml for s in prs.slides]
    title = prs.slides.add_slide(prs.slide_layouts[0])
    caption = prs.slides.add_slide(prs.slide_layouts[8])
    caption.shapes.title.text = "Свой заголовок"
    assert fill_empty_placeholders(prs.slides) == 3
    assert [ph.text_frame.text for ph in title.placeholders] == [
        "Заголовок слайда",
        "Подзаголовок слайда",
    ]
    texts = {ph.placeholder_format.type: ph.text_frame.text for ph in caption.placeholders}
    # Заполненный заголовок не трогается, пустое место под картинку остаётся пустым.
    assert texts == {PP.TITLE: "Свой заголовок", PP.PICTURE: "", PP.BODY: "Текст слайда"}
    assert [s._element.xml for s in list(prs.slides)[: len(samples)]] == samples
    english = Presentation(str(MINI_TEMPLATE))
    fresh = english.slides.add_slide(english.slide_layouts[5])
    assert fill_empty_placeholders([fresh], "en") == 1
    assert fresh.shapes.title.text == "Slide title"


def test_original_compose_and_preview_fill_empty_placeholders_alike(
    tmp_path: pathlib.Path, import_settings: Settings
) -> None:
    from presentation_designer.pipeline.jobs import _preview_prompts

    source = tmp_path / "source.pptx"
    prs = Presentation(str(MINI_TEMPLATE))
    prs.slides.add_slide(prs.slide_layouts[5])
    prs.save(str(source))
    profile = own_profile(source, "tpl_prompts")
    package = _package(tmp_path, import_settings)
    story = original_story(package, profile, {"language": "ru"})
    plan = original_plan(story, profile, package, plan_id="plan_prompts")
    out = tmp_path / "original.pptx"
    result = compose_deck(plan, profile, source, package, out_pptx=out)
    assert result.integrity.ok
    before, after = Presentation(str(source)), Presentation(str(out))
    assert after.slides[-1].shapes.title.text == "Заголовок слайда"
    for old, new in zip(list(before.slides)[:-1], list(after.slides)[:-1], strict=True):
        assert old._element.xml == new._element.xml

    # Предварительная ревизия (её открывает ONLYOFFICE) заполняет ту же копию так же, а файл
    # без пустых мест не пересохраняет.
    preview = tmp_path / "preview.pptx"
    shutil.copyfile(source, preview)
    _preview_prompts(preview, {"language": "ru"})
    shown = Presentation(str(preview))
    assert [s._element.xml for s in shown.slides] == [s._element.xml for s in after.slides]
    untouched = tmp_path / "untouched.pptx"
    shutil.copyfile(MINI_TEMPLATE, untouched)
    _preview_prompts(untouched, {})
    assert untouched.read_bytes() == MINI_TEMPLATE.read_bytes()


def test_original_reorder_delete_and_duplicate(
    mini_profile: dict[str, Any], tmp_path: pathlib.Path, import_settings: Settings
) -> None:
    package = _package(tmp_path, import_settings)
    story = original_story(package, mini_profile, {})
    plan = original_plan(story, mini_profile, package, plan_id="plan_orig")
    texts = _texts(MINI_TEMPLATE)
    for sequence in ([3, 0, 2], [2, 0, 2, 1]):
        edited = deepcopy(plan)
        edited["slides"] = []
        for order, source_index in enumerate(sequence, 1):
            slide = deepcopy(plan["slides"][source_index])
            slide.update(order=order, slide_id=f"s{order}")
            edited["slides"].append(slide)
        out = tmp_path / f"reordered-{len(sequence)}.pptx"
        result = compose_deck(edited, mini_profile, MINI_TEMPLATE, package, out_pptx=out)
        assert result.integrity.ok
        assert _texts(out) == [texts[i] for i in sequence]


def test_original_duplicate_preserves_notes_links_and_independent_edits(
    tmp_path: pathlib.Path, import_settings: Settings
) -> None:
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    source = tmp_path / "source.pptx"
    prs = Presentation(str(MINI_TEMPLATE))
    for index, slide in enumerate(prs.slides):
        slide.notes_slide.notes_text_frame.text = f"Заметки {index}"
    prs.slides[2]._element.set("show", "0")
    prs.slides[0].shapes[0].click_action.target_slide = prs.slides[1]
    prs.save(str(source))
    profile = own_profile(source, "tpl_duplicate")
    package = _package(tmp_path, import_settings)
    plan = original_plan(original_story(package, profile, {}), profile, package, plan_id="p")
    duplicate = deepcopy(plan["slides"][0])
    duplicate.update(slide_id="duplicate", order=5)
    plan["slides"].append(duplicate)
    plan["slides"][0]["blocks"][0]["text"] = "Правка только оригинала"
    out = tmp_path / "duplicated.pptx"
    compose_deck(plan, profile, source, package, out_pptx=out)
    actual = Presentation(str(out))
    assert all(slide.has_notes_slide for slide in actual.slides)
    assert [slide.notes_slide.notes_text_frame.text for slide in actual.slides] == [
        "Заметки 0",
        "Заметки 1",
        "Заметки 2",
        "Заметки 3",
        "Заметки 0",
    ]
    assert _texts(out)[4] == _texts(source)[0]
    assert "Правка только оригинала" in _texts(out)[0]
    assert actual.slides[2]._element.get("show") == "0"
    for index in (0, 4):
        slide = actual.slides[index]
        assert slide.shapes[0].click_action.target_slide.part is actual.slides[1].part
        back_links = [
            rel.target_part
            for rel in slide.notes_slide.part.rels.values()
            if rel.reltype == RT.SLIDE
        ]
        assert back_links == [slide.part]
    assert len({s.notes_slide.part.partname for s in actual.slides}) == 5
    actual.slides[4].notes_slide.notes_text_frame.text = "Независимые заметки"
    assert actual.slides[0].notes_slide.notes_text_frame.text == "Заметки 0"


def test_original_notes_only_edit_is_not_discarded(
    mini_profile: dict[str, Any], tmp_path: pathlib.Path, import_settings: Settings
) -> None:
    package = _package(tmp_path, import_settings)
    plan = original_plan(
        original_story(package, mini_profile, {}), mini_profile, package, plan_id="p"
    )
    plan["slides"][0]["notes"] = "Новые заметки докладчика"
    out = tmp_path / "notes.pptx"
    compose_deck(plan, mini_profile, MINI_TEMPLATE, package, out_pptx=out)
    slide = Presentation(str(out)).slides[0]
    assert slide.has_notes_slide
    assert slide.notes_slide.notes_text_frame.text == "Новые заметки докладчика"
    assert _texts(out) == _texts(MINI_TEMPLATE)
