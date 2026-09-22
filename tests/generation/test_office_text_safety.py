"""Live failure regressions: Unicode identity and conservative semantic anchors."""

import io
import json

import pytest
from pptx import Presentation
from pptx.util import Inches

from presentation_designer.generation import office_edit, office_object_edit
from presentation_designer.generation.office_objects import objects
from presentation_designer.generation.office_preservation import verify_selected_text_edit
from presentation_designer.llm.cache import ResponseCache
from presentation_designer.llm.client import LlmClient
from presentation_designer.llm.limiter import LocalLimiter, Quota
from presentation_designer.llm.retry import RetryPolicy
from presentation_designer.llm.stub import StubTransport
from presentation_designer.llm.types import ResponseError
from presentation_designer.shared.settings import Settings, get_models_config


def deck_with_runs(texts, line_break=False):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(2))
    paragraph = shape.text_frame.paragraphs[0]
    for i, text in enumerate(texts):
        if i and line_break:
            paragraph.add_line_break()
        paragraph.add_run().text = text
    slide.shapes.add_textbox(Inches(1), Inches(4), Inches(5), Inches(1)).text = "Сосед"
    deck.slides.add_slide(deck.slide_layouts[6])
    output = io.BytesIO()
    deck.save(output)
    return output.getvalue()


@pytest.mark.parametrize(
    ("texts", "afters"),
    [
        (["Оси можно использовать, если данных много."], ["Оси нужны при множестве данных."]),
        (["Если возможно, делайте столбики."], ["Делайте столбики."]),
        (["Синий, но можно свободно использовать дополнительные."], ["Синий, дополнительные."]),
        (["В заголовках уже задан", "нужный размер шрифта"], ["В заголовках", "размер шрифта"]),
        (["Выравниваем по левому краю (исключения — схемы)."], ["Текст по левому краю."]),
        (["Не удаляйте подписи."], ["Удаляйте подписи."]),
        (["Данные можно использовать."], ["Данные можно использовать, но"]),
        (["Необязательные оси можно использовать."], [""]),
        (["Перекрытие = ±50%."], ["Перекрытие = 50%."]),
        (
            ["Основной цвет — синий, можно", " использовать дополнительные."],
            ["Основной цвет — синий,", " дополнительные."],
        ),
    ],
)
def test_shortening_rejects_live_semantic_losses(texts, afters):
    data = deck_with_runs(texts, line_break=True)
    plan = office_edit.EditPlan(
        explanation="Сокращено",
        patches=[
            office_edit.TextPatch(slide=1, run=i, before=before, after=after)
            for i, (before, after) in enumerate(zip(texts, afters, strict=True))
        ],
    )
    with pytest.raises(ValueError):
        plan.validate_facts("Сократи, сохрани смысл", data)
    # These guards must not block explicit changes outside shortening.
    plan.validate_facts("Замени текст", data)


def test_safe_shortening_retains_complete_split_sentence():
    texts = ["В заголовках уже задан", "нужный размер шрифта"]
    data = deck_with_runs(texts, line_break=True)
    plan = office_edit.EditPlan(
        explanation="Сокращено",
        patches=[
            office_edit.TextPatch(slide=1, run=0, before=texts[0], after="В заголовках задан"),
            office_edit.TextPatch(slide=1, run=1, before=texts[1], after="размер шрифта"),
        ],
    )
    plan.validate_facts("Сократи", data)
    verify_selected_text_edit(data, office_edit.patch_pptx(data, plan), [objects(data)[0]])


@pytest.mark.parametrize("scope", ["deck", "object", "many"])
@pytest.mark.parametrize("unicode_text", [True, False])
@pytest.mark.parametrize("recover", [True, False])
async def test_bad_text_repaired_with_exact_context(
    scope, unicode_text, recover, tmp_path, monkeypatch
):
    before = (
        "С\u00a02011\u202fгода мы создаём проекты VK."
        if unicode_text
        else ("Оси можно использовать, если данных много.")
    )
    after = (
        "С\u00a02011\u202fгода создаём проекты VK."
        if unicode_text
        else ("Оси можно использовать, если много данных.")
    )
    data = deck_with_runs([before])
    target = objects(data)[0]

    def answer(bad):
        plan = {
            "explanation": "Сокращено",
            "patches": [
                {
                    "slide": 1,
                    "run": 0,
                    "before": before.replace("\u00a0", " ").replace("\u202f", " ")
                    if bad and unicode_text
                    else before,
                    "after": "Оси нужны при множестве данных."
                    if bad and not unicode_text
                    else after,
                }
            ],
        }
        return (
            {
                "explanation": "Сокращено",
                "edits": [
                    {
                        "target": {"slide": 1, "shape_id": target.shape_id},
                        "plan": plan,
                    }
                ],
            }
            if scope == "many"
            else plan
        )

    stub = StubTransport()
    stub.answer(answer(True), times=1 if recover else None)
    stub.answer(answer(False))
    settings = Settings()
    client = LlmClient(
        settings=settings,
        models=get_models_config(),
        transport=stub,
        limiter=LocalLimiter(Quota(2, 1000, 1_000_000)),
        cache=ResponseCache("off", tmp_path / "cache", tmp_path / "fixtures"),
        retry_policy=RetryPolicy(max_retries=1, base_s=0, max_s=0),
    )
    module = office_edit if scope == "deck" else office_object_edit
    monkeypatch.setattr(module, "build_client", lambda settings: client)

    async def propose():
        if scope == "deck":
            return await module.propose(data, "Сократи", settings)
        if scope == "object":
            return await module.propose(data, "Сократи", settings, target)
        return (await module.propose_many(data, "Сократи", settings, [target])).edits[0].plan

    if not recover:
        with pytest.raises(ResponseError, match=r"базой правки|модальность"):
            await propose()
        assert len(stub.calls) == 2
        assert Presentation(io.BytesIO(data)).slides[0].shapes[0].text == before
        return
    plan = await propose()
    assert len(stub.calls) == 2
    if unicode_text:
        prompt = stub.calls[0].messages[-1].text
        assert "\\u00a0" in prompt and "\\u202f" in prompt
        assert "\u00a0" not in prompt and "\u202f" not in prompt
        assert before in json.dumps(json.loads(prompt), ensure_ascii=False)
        assert "\\u00a0" in stub.calls[1].messages[-1].text
    updated = office_edit.patch_pptx(data, plan)
    assert Presentation(io.BytesIO(updated)).slides[0].shapes[0].text == after
    verify_selected_text_edit(data, updated, [target])


def test_meaning_check_joins_styled_words_without_inserting_spaces():
    texts = ["Оси мож", "но использовать, если данных много."]
    data = deck_with_runs(texts)
    plan = office_edit.EditPlan(
        explanation="Сократил",
        patches=[
            office_edit.TextPatch(
                slide=1, run=1, before=texts[1], after="но брать, если данных много."
            ),
        ],
    )
    plan.validate_facts("Сократи", data)
    plan.patches[0].after = "но брать при множестве данных."
    with pytest.raises(ValueError, match="модальность"):
        plan.validate_facts("Сократи", data)


@pytest.mark.parametrize("char", ["\u00a0", "\u202f", "\u2009", "\u200b", "\u2060", "\ufeff"])
def test_context_roundtrips_invisible_characters_without_normalizing(char):
    text = f"Текст{char}VK \\u00a0"
    encoded = office_edit.text_context_json({"text": text})
    assert char not in encoded
    assert json.loads(encoded) == {"text": text}
