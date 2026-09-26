"""Операции над объектом копии на месте (этап 38): стиль из токенов шаблона, абзацы, геометрия,
подпись и копия карточки; остальные части PPTX — байт в байт."""

from __future__ import annotations

import io
from zipfile import ZipFile

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from presentation_designer.generation.office_objects import objects
from presentation_designer.generation.office_ops import (
    ObjectOp,
    Tokens,
    apply,
    nearest_color,
    next_size,
    overflow_note,
)

TOKENS = Tokens(
    palette=["#0077FF", "#1A1A1A", "#FFFFFF", "#FF7A00"],
    scale=[12.0, 16.0, 20.0, 28.0, 40.0],
    fonts=["Arial"],
    text_color="#1A1A1A",
    caption_pt=12.0,
    body_pt=16.0,
)


def _deck() -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(16), Inches(9)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(3))
    box.name = "Пункты"
    frame = box.text_frame
    frame.text = "Первый пункт"
    for text in ("Второй пункт", "Третий пункт"):
        frame.add_paragraph().text = text
    for run in (p.runs[0] for p in frame.paragraphs):
        run.font.size = Pt(16)
    for n in range(3):
        card = slide.shapes.add_textbox(Inches(1 + n * 5), Inches(5), Inches(4), Inches(2))
        card.name = f"Карточка {n + 1}"
        card.text_frame.text = f"Карточка {n + 1}"
    prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(0, 0, Inches(1), Inches(1))
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def _obj(data: bytes, name: str):
    return next(o for o in objects(data) if o.slide == 1 and o.name == name)


def _shape(data: bytes, name: str):
    return next(s for s in Presentation(io.BytesIO(data)).slides[0].shapes if s.name == name)


def test_tokens_pick_scale_steps_and_palette_colors() -> None:
    assert next_size(16, 1, TOKENS.scale) == 20
    assert next_size(16, -1, TOKENS.scale) == 12
    assert next_size(12, -1, TOKENS.scale) == 10  # ниже шкалы — ×1,15, не меньше 9 пт
    assert next_size(9.5, -1, TOKENS.scale) == 9
    assert nearest_color("#0077FF", TOKENS.palette) == ("#0077FF", True)
    assert nearest_color("#E00000", TOKENS.palette) == ("#FF7A00", False)


def test_style_ops_change_only_the_object_and_keep_other_parts() -> None:
    data = _deck()
    obj = _obj(data, "Пункты")
    ops = [
        ObjectOp(op="style.size", step=1),
        ObjectOp(op="style.bold", on=True),
        ObjectOp(op="style.color", color="#E00000"),
        ObjectOp(op="style.align", align="center"),
        ObjectOp(op="style.font", font="Comic Sans"),
    ]
    out, notes = apply(data, obj, ops, TOKENS, 16.0)
    shape = _shape(out, "Пункты")
    runs = [r for p in shape.text_frame.paragraphs for r in p.runs]
    assert all(r.font.size == Pt(20) and r.font.bold for r in runs)
    assert all(str(r.font.color.rgb) == "FF7A00" for r in runs)
    assert all(p.alignment == 2 for p in shape.text_frame.paragraphs)  # по центру
    assert any("ближайший #FF7A00" in n for n in notes)
    assert any("Comic Sans" in n and "Arial" in n for n in notes)
    # Меняется только XML слайда 1.
    with ZipFile(io.BytesIO(data)) as before, ZipFile(io.BytesIO(out)) as after:
        changed = {n for n in before.namelist() if before.read(n) != after.read(n)}
    assert changed == {"ppt/slides/slide1.xml"}
    # Карточки на том же слайде не тронуты.
    assert _shape(out, "Карточка 1").text_frame.paragraphs[0].runs[0].font.bold is None


def test_paragraphs_are_added_with_neighbour_style_and_removed_by_number() -> None:
    data = _deck()
    obj = _obj(data, "Пункты")
    out, _ = apply(
        data,
        obj,
        [
            ObjectOp(op="text.insert_paragraph", paragraph=3, text="Сроки: до конца года"),
            ObjectOp(op="text.delete_paragraph", paragraph=2),
        ],
        TOKENS,
        16.0,
    )
    paragraphs = _shape(out, "Пункты").text_frame.paragraphs
    assert [p.text for p in paragraphs] == ["Первый пункт", "Третий пункт", "Сроки: до конца года"]
    assert paragraphs[-1].runs[0].font.size == Pt(16)
    with pytest.raises(ValueError, match="абзаца 9 нет"):
        apply(data, obj, [ObjectOp(op="text.delete_paragraph", paragraph=9)], TOKENS, 16.0)


def test_geometry_delete_and_order() -> None:
    data = _deck()
    out, _ = apply(
        data, _obj(data, "Пункты"), [ObjectOp(op="object.resize", width=0.75)], TOKENS, None
    )
    assert _shape(out, "Пункты").width == Inches(12)
    out, notes = apply(data, _obj(data, "Карточка 2"), [ObjectOp(op="object.delete")], TOKENS, None)
    assert "Карточка 2" not in [s.name for s in Presentation(io.BytesIO(out)).slides[0].shapes]
    assert notes == ["объект удалён"]
    out, _ = apply(
        data, _obj(data, "Пункты"), [ObjectOp(op="object.z_order", to="front")], TOKENS, None
    )
    assert [s.name for s in Presentation(io.BytesIO(out)).slides[0].shapes][-1] == "Пункты"


def test_caption_below_in_template_style() -> None:
    data = _deck()
    obj = _obj(data, "Пункты")
    out, notes = apply(
        data,
        obj,
        [ObjectOp(op="object.add_text", text="Источник: Росстат", role="caption")],
        TOKENS,
        None,
    )
    caption = _shape(out, "Подпись")
    assert caption.text_frame.text == "Источник: Росстат"
    run = caption.text_frame.paragraphs[0].runs[0]
    assert (
        run.font.size == Pt(12) and str(run.font.color.rgb) == "1A1A1A" and run.font.name == "Arial"
    )
    assert caption.top > _shape(out, "Пункты").top + _shape(out, "Пункты").height
    assert "надпись" in notes[0]


def test_another_card_joins_the_row_within_its_bounds() -> None:
    data = _deck()
    out, notes = apply(
        data,
        _obj(data, "Карточка 3"),
        [ObjectOp(op="object.add_block", text="Карточка 4")],
        TOKENS,
        None,
    )
    cards = sorted(
        (
            s
            for s in Presentation(io.BytesIO(out)).slides[0].shapes
            if s.name.startswith("Карточка")
        ),
        key=lambda s: s.left,
    )
    assert len(cards) == 4 and cards[-1].text_frame.text == "Карточка 4"
    assert cards[0].left == Inches(1)
    assert abs(cards[-1].left + cards[-1].width - Inches(15)) < Inches(0.01)
    widths = {c.width for c in cards}
    assert max(widths) - min(widths) <= 1
    ids = [s.shape_id for s in Presentation(io.BytesIO(out)).slides[0].shapes]
    assert len(ids) == len(set(ids))
    assert "их теперь 4" in notes[0]


def test_overflow_is_reported_after_the_edit() -> None:
    data = _deck()
    obj = _obj(data, "Карточка 1")
    long_text = "Очень длинный текст карточки, который не поместится. " * 12
    out, _ = apply(
        data, obj, [ObjectOp(op="text.insert_paragraph", paragraph=1, text=long_text)], TOKENS, 18.0
    )
    assert overflow_note(out, 1, "Карточка 1") is not None
    assert overflow_note(data, 1, "Пункты") is None
