"""Готовая презентация («Открыть как презентацию»): копия публикуется ревизией r1 прямо в
запросе создания задания, подписи пустых плейсхолдеров и диаграммы-картинки в фоне дают
ревизию r2 (r1 не меняется), полная сборка пишет в последнюю ревизию и повторяет замены по её
charts.json; ready_at — время первой публикации, и сборка его не переписывает."""

from __future__ import annotations

import hashlib
import io
import pathlib
from collections.abc import Callable
from typing import Any

from fastapi.testclient import TestClient
from pptx import Presentation

from presentation_designer.api.app import create_app
from presentation_designer.pipeline import jobs
from presentation_designer.pipeline.artifacts import ArtifactStore
from presentation_designer.pipeline.files import FileStore
from presentation_designer.pipeline.jobs import InlineExecutor, Orchestrator, _preview_chart_report
from presentation_designer.pipeline.real import RealLayers
from presentation_designer.pipeline.run import ChartImagesOutput, ComposeInput, ComposeOutput
from presentation_designer.pipeline.state import State
from presentation_designer.pipeline.stubs import StubLayers
from presentation_designer.shared.settings import Settings
from tests.pipeline.helpers import Recorder

PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
REPORT: dict[str, Any] = {
    "version": 1,
    "pictures": [
        {"sha256": "a" * 64, "slides": [2], "status": "read", "reason": "", "reading": None}
    ],
    "swaps": [{"slide": 2, "sha256": "a" * 64, "status": "replaced", "reason": ""}],
    "message": "Сделал диаграммы редактируемыми: 1 (слайд 2).",
}


class ChartLayers(StubLayers):
    """Заглушки с заменой диаграмм: запоминают, что увидела полная сборка."""

    def __init__(self, replaced: int = 1) -> None:
        super().__init__(stage_delay_ms=0)
        self.replaced = replaced
        self.progress: list[str] = []
        self.reports: list[Any] = []
        self.composed: list[int] = []

    def chart_images(
        self, pptx_path: pathlib.Path, progress: Callable[[str, float], None] | None = None
    ) -> ChartImagesOutput:
        assert pptx_path.is_file()
        if progress is not None:
            progress("Делаю диаграммы редактируемыми: 1 картинка похожа на диаграммы", 0.0)
            progress("Делаю диаграммы редактируемыми: 1 из 1", 1.0)
            self.progress.append("called")
        if not self.replaced:
            return ChartImagesOutput()
        # Замена меняет файл: в r2 лежит уже не копия r1.
        prs = Presentation(str(pptx_path))
        prs.core_properties.comments = "диаграммы заменены"
        prs.save(str(pptx_path))
        return ChartImagesOutput(report=REPORT, message=REPORT["message"], replaced=self.replaced)

    def compose(self, inp: ComposeInput) -> ComposeOutput:
        self.reports.append(inp.chart_report)
        self.composed.append(inp.revision)
        return super().compose(inp)


def _orchestrator(tmp_path: pathlib.Path, layers: StubLayers, executor: Any = None) -> Orchestrator:
    settings = Settings()
    settings.paths.data_dir = tmp_path / "data"
    settings.paths.artifacts_dir = tmp_path / "artifacts"
    settings.paths.runs_dir = tmp_path / "runs"
    settings.paths.backups_dir = tmp_path / "backups"
    return Orchestrator(
        settings,
        State(settings.db_path),
        FileStore(settings.uploads_dir, max_upload_mb=2, max_unzipped_mb=64),
        ArtifactStore(settings.artifacts_dir),
        layers,
        executor or InlineExecutor(),
    )


def _open_deck(client: TestClient, data: bytes) -> str:
    """Тот же путь, что у интерфейса: файл проекта → шаблон → содержание → задание original."""
    project_id = client.post("/api/projects", json={}).json()["project_id"]
    file_id = client.post(
        f"/api/projects/{project_id}/files",
        files=[("files", ("deck.pptx", data, PPTX_MIME))],
    ).json()[0]["file_id"]
    template_id = client.post("/api/templates", json={"file_id": file_id}).json()["template_id"]
    package_id = client.post("/api/content", json={"file_ids": [file_id]}).json()["package_id"]
    r = client.post(
        "/api/generations",
        json={
            "schema_version": "1.2",
            "template_id": template_id,
            "package_id": package_id,
            "settings": {"variants": ["original"]},
        },
    )
    assert r.status_code == 202, r.text
    return str(r.json()["job_id"])


def _sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _with_empty_title(data: bytes) -> bytes:
    """Колода с пустым плейсхолдером заголовка: подпись дописывается только в r2."""
    prs = Presentation(io.BytesIO(data))
    prs.slides.add_slide(prs.slide_layouts[0])
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def test_charts_make_r2_and_r1_stays_a_copy(tmp_path: pathlib.Path, pptx_bytes: bytes) -> None:
    layers = ChartLayers()
    orch = _orchestrator(tmp_path, layers)
    with TestClient(create_app(orch, reconcile=False)) as client:
        job_id = _open_deck(client, pptx_bytes)
        result = client.get(f"/api/generations/{job_id}").json()
        r1 = client.get(f"/api/generations/{job_id}/artifacts/original/r1/deck.pptx").content
        r2 = client.get(f"/api/generations/{job_id}/artifacts/original/r2/deck.pptx").content
    assert layers.progress == ["called"]
    variant = result["variants"][0]
    assert [r["revision"] for r in variant["revisions"]] == [1, 2]
    assert variant["revision"] == 2 and variant["artifacts"]["pptx"] == "original/r2/deck.pptx"
    # r1 — копия файла байт в байт: офисная копия r1 создаётся по ней и больше не меняется.
    assert _sha(r1) == _sha(pptx_bytes) and _sha(r2) != _sha(r1)
    names = set(result["artifacts_manifest"])
    assert "original/r1/charts.json" not in names
    # Полная сборка — в r2 и по её charts.json (заглушка вёрстки его не переносит, настоящая —
    # да, см. последний тест); план есть только у r2.
    assert layers.composed == [2] and layers.reports == [REPORT]
    assert "original/r2/plan.json" in names and "original/r1/plan.json" not in names
    assert "original/r2/deck.pdf" in names
    # Фраза для чата — одна, из момента публикации r2.
    assert [w for w in result["warnings"] if w["code"] == "chart_images"] == [
        {"code": "chart_images", "message": REPORT["message"]}
    ]


def test_without_changes_there_is_no_r2(tmp_path: pathlib.Path, pptx_bytes: bytes) -> None:
    layers = ChartLayers(replaced=0)
    orch = _orchestrator(tmp_path, layers)
    with TestClient(create_app(orch, reconcile=False)) as client:
        job_id = _open_deck(client, pptx_bytes)
        result = client.get(f"/api/generations/{job_id}").json()
    variant = result["variants"][0]
    assert [r["revision"] for r in variant["revisions"]] == [1]
    assert layers.composed == [1]
    assert {"original/r1/deck.pdf", "original/r1/plan.json"} <= set(result["artifacts_manifest"])
    assert not [w for w in result["warnings"] if w["code"] == "chart_images"]


def test_prompts_alone_make_r2(tmp_path: pathlib.Path, pptx_bytes: bytes) -> None:
    """Подписи пустых плейсхолдеров не входят в синхронный путь (сохранение пакета на сервере
    дольше секунды): их получает r2, а r1 остаётся копией."""
    data = _with_empty_title(pptx_bytes)
    orch = _orchestrator(tmp_path, ChartLayers(replaced=0))
    with TestClient(create_app(orch, reconcile=False)) as client:
        job_id = _open_deck(client, data)
        result = client.get(f"/api/generations/{job_id}").json()
        r1 = client.get(f"/api/generations/{job_id}/artifacts/original/r1/deck.pptx").content
        r2 = client.get(f"/api/generations/{job_id}/artifacts/original/r2/deck.pptx").content
    assert [r["revision"] for r in result["variants"][0]["revisions"]] == [1, 2]
    assert _sha(r1) == _sha(data)
    title = Presentation(io.BytesIO(r2)).slides[-1].shapes.title
    assert title is not None and title.text.strip()
    assert not Presentation(io.BytesIO(r1)).slides[-1].shapes.title.text.strip()


def test_ready_at_is_the_first_publication(tmp_path: pathlib.Path, pptx_bytes: bytes) -> None:
    """Копия опубликована в запросе: ready_at ставится тогда, фон и полная сборка его не
    переписывают, и first_file_ready_ms считает показ, а не конец разбора."""
    recorder = Recorder()
    layers = ChartLayers()
    orch = _orchestrator(tmp_path, layers, recorder)
    with TestClient(create_app(orch, reconcile=False)) as client:
        job_id = _open_deck(client, pptx_bytes)
        shown = client.get(f"/api/generations/{job_id}").json()
        variant = shown["variants"][0]
        assert variant["status"] == "running" and variant["slide_count"] == 4
        assert variant["artifacts"]["pptx"] == "original/r1/deck.pptx"
        ready_at = variant["ready_at"]
        # Задачи ещё не выполнялись: ни одна из них не нужна для показа.
        for func in (
            jobs.task_analyze,
            jobs.task_import,
            jobs.task_original_preview,
            jobs.task_story,
        ):
            recorder.run(func)
        recorder.run(jobs.task_variant)
        recorder.run(jobs.task_finalize)
        done = client.get(f"/api/generations/{job_id}").json()
    assert done["status"] in ("succeeded", "needs_review")
    assert done["variants"][0]["ready_at"] == ready_at
    first = done["metrics"]["timeline"]["first_file_ready_ms"]
    assert first is not None and first < 5000


def test_full_build_follows_the_preview_or_reads_itself(tmp_path: pathlib.Path) -> None:
    assert _preview_chart_report(None) is None  # предварительной ревизии нет: читать самим
    preview = tmp_path / "r1"
    preview.mkdir()
    # Ревизия показана без отчёта (чтения не было или оно сломалось): не заменять ничего.
    assert _preview_chart_report(preview) == {}
    (preview / "charts.json").write_text('{"version": 1, "pictures": []}', encoding="utf-8")
    assert _preview_chart_report(preview) == {"version": 1, "pictures": []}


def test_real_compose_carries_the_preview_report_into_the_revision(tmp_path: pathlib.Path) -> None:
    layers = RealLayers(Settings())
    store = ArtifactStore(tmp_path / "artifacts")
    with store.stage_revision("job_x", "original", 2) as staging:
        inp = ComposeInput(
            "job_x", "original", 2, {}, {}, None, {}, {}, staging, chart_report=REPORT
        )
        assert layers._chart_report(inp, None, None) is None
    manifest = store.read_manifest("job_x", "original", 2)
    assert "original/r2/charts.json" in manifest
