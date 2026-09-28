from pathlib import Path

import httpx
import pytest
from pptx import Presentation
from pptx.util import Inches
from scripts.benchmark_presenton import build_request, download_url, generate, inspect_deck


def test_request_preserves_long_source_and_tables():
    text = "подробность " * 250
    request = build_request(
        {
            "brief": {"title": "Проверка"},
            "blocks": [{"text": text, "items": ["Оговорка"]}],
            "datasets": [
                {
                    "dataset_id": "ds1",
                    "columns": [{"name": "Экономия", "unit": "млн ₽"}],
                    "rows": [[12.5], [31]],
                }
            ],
        },
        "general",
        15,
    )
    assert text in request["content"]
    assert "Оговорка" in request["content"]
    assert "Экономия (млн ₽)" in request["content"]
    assert "12.5" in request["content"]
    assert request["web_search"] is False
    assert "slides_markdown" not in request


@pytest.mark.parametrize("count", [0, 16])
def test_slide_budget(count):
    with pytest.raises(ValueError):
        build_request({}, "general", count)


@pytest.mark.parametrize("url", ["https://evil.test/app_data/x", "//evil.test/x", "/etc/passwd"])
def test_reject_foreign_downloads(url):
    with pytest.raises(ValueError):
        download_url("http://localhost:5051", url)


def test_download_path():
    assert download_url("http://localhost:5051", "/app_data/x.pptx") == (
        "http://localhost:5051/app_data/x.pptx"
    )


def test_generate_preserves_existing_output(tmp_path, monkeypatch):
    out = tmp_path / "run"
    out.mkdir()
    report = out / "report.json"
    report.write_text("previous evidence", encoding="utf-8")

    def unexpected_client(**kwargs):
        pytest.fail("Existing output must be rejected before any network call")

    monkeypatch.setattr("scripts.benchmark_presenton.httpx.Client", unexpected_client)
    with pytest.raises(FileExistsError):
        generate("http://localhost:5051", {}, out, 10)
    assert report.read_text(encoding="utf-8") == "previous evidence"
    assert not (out / "request.json").exists()


def make_deck(path: Path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text = "На слайде"
    slide.notes_slide.notes_text_frame.text = "Только заметки"
    table = slide.shapes.add_table(1, 1, Inches(1), Inches(3), Inches(3), Inches(1)).table
    table.cell(0, 0).text = "12,5 млн ₽"
    prs.save(path)


def test_inspect_separates_notes_and_includes_tables(tmp_path):
    path = tmp_path / "deck.pptx"
    make_deck(path)
    report = inspect_deck(path)
    assert report["within_15_slides"]
    slide = report["slides"][0]
    assert "12,5 млн ₽" in slide["slide_xml_text"]
    assert "Только заметки" not in slide["slide_xml_text"]
    assert slide["notes_text"] == "Только заметки"
    assert slide["table_count"] == 1
    assert slide["chart_count"] == 0


def test_generate_downloads_pptx_and_does_not_claim_acceptance(tmp_path, monkeypatch):
    path = tmp_path / "input.pptx"
    make_deck(path)

    def handler(request):
        if request.method == "POST":
            return httpx.Response(200, json={"path": "/app_data/deck.pptx"})
        return httpx.Response(200, content=path.read_bytes())

    client = httpx.Client(transport=httpx.MockTransport(handler))
    monkeypatch.setattr("scripts.benchmark_presenton.httpx.Client", lambda **kw: client)
    report = generate("http://localhost:5051", {"content": "Материал"}, tmp_path / "run", 10)
    assert report["status"] == "generated_not_accepted"
    assert report["cost_usd"] is None
    assert report["deck"]["slide_count"] == 1


def test_failed_generation_leaves_report_without_secrets(tmp_path, monkeypatch):
    def handler(request):
        return httpx.Response(500, text="secret-key")

    client = httpx.Client(transport=httpx.MockTransport(handler))
    monkeypatch.setattr("scripts.benchmark_presenton.httpx.Client", lambda **kw: client)
    out = tmp_path / "run"
    with pytest.raises(httpx.HTTPStatusError):
        generate("http://localhost:5051", {}, out, 10)
    text = (out / "report.json").read_text()
    assert '"status": "failed"' in text
    assert "secret-key" not in text
