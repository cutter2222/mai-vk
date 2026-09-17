"""Синтетический шаблон для стилей служебных слайдов и групп образцов (этап 14): два титула
одного состава на светлом и тёмном фоне, два разделителя разного тона, два карточных образца
одной сигнатуры, два финала разного тона. Фон задаётся заливкой слайда, поэтому тон читается
из `p:bg` самого слайда (`source: slide_fill`)."""

from __future__ import annotations

import io
import pathlib

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt
from tests.fixtures.rich_template import _png

LIGHT = RGBColor(0xEB, 0xF3, 0xF9)
DARK = RGBColor(0x0B, 0x1E, 0x3A)
INK = RGBColor(0x20, 0x20, 0x20)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def build_variety_template(path: pathlib.Path) -> pathlib.Path:
    prs = Presentation()
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)
    title_layout = prs.slide_layouts[0]  # Title Slide
    section_layout = prs.slide_layouts[2]  # Section Header
    blank = prs.slide_layouts[6]
    logo = _png((0, 119, 255), 48)
    icon = _png((0, 119, 255), 40, alpha=True)

    def background(slide: object, color: RGBColor) -> None:
        fill = slide.background.fill  # type: ignore[attr-defined]
        fill.solid()
        fill.fore_color.rgb = color

    def add_logo(slide: object) -> None:
        slide.shapes.add_picture(
            io.BytesIO(logo), Emu(11200000), Emu(6200000), Emu(500000), Emu(500000)
        )  # type: ignore[attr-defined]

    def textbox(
        slide: object,
        text: str,
        x: int,
        y: int,
        w: int,
        h: int,
        size: int,
        *,
        bold: bool = False,
        color: RGBColor = INK,
    ) -> object:
        tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))  # type: ignore[attr-defined]
        tb.text_frame.text = text
        run = tb.text_frame.paragraphs[0].runs[0]
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        return tb

    # 1–2. титулы одного состава: светлый и тёмный
    for color in (LIGHT, DARK):
        s = prs.slides.add_slide(title_layout)
        background(s, color)
        s.shapes.title.text = "Название презентации"
        s.placeholders[1].text = "Имя Фамилия, должность"
        add_logo(s)
    # 3–4. разделители разного тона (макет Section Header → роль по имени макета)
    for color, text_color in (
        (LIGHT, INK),
        (DARK, WHITE),
    ):
        s = prs.slides.add_slide(section_layout)
        background(s, color)
        s.shapes.title.text = "Название раздела"
        s.placeholders[1].text = "Короткое пояснение к разделу"
        for ph in s.placeholders:
            for p in ph.text_frame.paragraphs:
                for r in p.runs:
                    r.font.color.rgb = text_color
        add_logo(s)
    # 5–6. два карточных образца одной сигнатуры: иконка + заголовок + текст × 3
    for variant in range(2):
        s = prs.slides.add_slide(blank)
        background(s, LIGHT)
        textbox(s, "Заголовок в одну строчку", 600000, 400000, 11000000, 900000, 32, bold=True)
        for i in range(3):
            x = 600000 + i * 3700000
            grp = s.shapes.add_group_shape()
            grp.shapes.add_picture(io.BytesIO(icon), Emu(x), Emu(1700000), Emu(400000), Emu(400000))
            t = grp.shapes.add_textbox(Emu(x), Emu(2200000), Emu(3300000), Emu(600000))
            t.text_frame.text = "Заголовок"
            t.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
            b = grp.shapes.add_textbox(Emu(x), Emu(2900000), Emu(3300000), Emu(1800000))
            b.text_frame.text = (
                "Текст описания карточки в несколько строк"
                if variant == 0
                else "Пояснение к карточке, вторая композиция того же состава"
            )
            b.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
        add_logo(s)
    # 7. список: заголовок и маркированный текст (кандидат для содержания без карточек)
    s = prs.slides.add_slide(blank)
    background(s, LIGHT)
    textbox(s, "Заголовок слайда", 600000, 400000, 11000000, 900000, 32, bold=True)
    tb = textbox(s, "Пункт первый", 600000, 1600000, 11000000, 3600000, 18)
    for text in ("Пункт второй", "Пункт третий"):
        para = tb.text_frame.add_paragraph()  # type: ignore[attr-defined]
        para.text = text
        para.runs[0].font.size = Pt(18)
    for para in tb.text_frame.paragraphs:  # type: ignore[attr-defined]
        p_pr = para._p.get_or_add_pPr()
        p_pr.set("marL", "342900")
        p_pr.set("indent", "-342900")
        bu = p_pr.makeelement(
            "{http://schemas.openxmlformats.org/drawingml/2006/main}buChar", {"char": "•"}
        )
        p_pr.append(bu)
    add_logo(s)
    # 8–9. финалы разного тона
    for color, text_color in (
        (DARK, WHITE),
        (LIGHT, INK),
    ):
        s = prs.slides.add_slide(blank)
        background(s, color)
        textbox(
            s,
            "Спасибо за внимание!",
            600000,
            2400000,
            11000000,
            1200000,
            44,
            bold=True,
            color=text_color,
        )
        textbox(s, "Имя Фамилия · почта", 600000, 3800000, 11000000, 600000, 18, color=text_color)
        add_logo(s)

    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    return path
