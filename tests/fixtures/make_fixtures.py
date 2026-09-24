"""Создаёт небольшие собственные фикстуры для обычных тестов: PPTX, документы каждого
формата импорта (docx, xlsx, csv, md, txt, pdf с текстовым слоем и скан без него, картинки)
и демонстрационный контент-пакет в examples/content.

Закрытые материалы организаторов сюда не попадают.
Запуск: uv run python tests/fixtures/make_fixtures.py
"""

from __future__ import annotations

import csv
import io
import json
import pathlib
import zipfile

from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Emu, Pt

HERE = pathlib.Path(__file__).resolve().parent
PPTX_DIR = HERE / "pptx"
DOCS_DIR = HERE / "content"
EXAMPLES_DIR = HERE.parents[1] / "examples" / "content"


def make_template() -> pathlib.Path:
    """Мини-шаблон: титул, карточки, таблица с диаграммой, финальный слайд."""
    prs = Presentation()
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)
    blank = prs.slide_layouts[6]

    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Название презентации"
    s.placeholders[1].text = "Имя Фамилия, должность"

    s = prs.slides.add_slide(blank)
    tb = s.shapes.add_textbox(Emu(600000), Emu(400000), Emu(11000000), Emu(800000))
    tb.text_frame.text = "Заголовок в две или одну строчку"
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(32)
    for i in range(3):
        card = s.shapes.add_textbox(
            Emu(600000 + i * 3700000), Emu(1600000), Emu(3400000), Emu(2600000)
        )
        card.text_frame.text = "Заголовок\nТекст"
        card.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
        card.text_frame.paragraphs[1].runs[0].font.size = Pt(14)

    s = prs.slides.add_slide(blank)
    tb = s.shapes.add_textbox(Emu(600000), Emu(400000), Emu(11000000), Emu(800000))
    tb.text_frame.text = "Таблица и диаграмма"
    rows, cols = 4, 3
    table = s.shapes.add_table(
        rows, cols, Emu(600000), Emu(1500000), Emu(5000000), Emu(2000000)
    ).table
    for r in range(rows):
        for c in range(cols):
            table.cell(r, c).text = "Заголовок" if r == 0 else "Текст"
    data = CategoryChartData()
    data.categories = ["Май", "Июнь", "Июль"]
    data.add_series("Открываемость", (31, 38, 44))
    s.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Emu(6200000), Emu(1500000), Emu(5400000), Emu(3500000), data
    )

    s = prs.slides.add_slide(blank)
    tb = s.shapes.add_textbox(Emu(600000), Emu(2500000), Emu(11000000), Emu(1200000))
    tb.text_frame.text = "Спасибо за внимание!"
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(40)

    PPTX_DIR.mkdir(parents=True, exist_ok=True)
    out = PPTX_DIR / "mini_template.pptx"
    prs.save(out)
    # Копия с другим sha256 для сценария «вариант падает»: шаблоны дедуплицируются по байтам,
    # поэтому имя со словом fail должно приходить с отдельным файлом.
    twin = PPTX_DIR / "mini_template_fail.pptx"
    with zipfile.ZipFile(out) as zin, zipfile.ZipFile(twin, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
        zout.comment = b"fixture: template whose detailed variant fails in stub compose"
    return out


def make_content() -> list[pathlib.Path]:
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    outs: list[pathlib.Path] = []

    doc = Document()
    doc.add_heading("Проблема", level=1)
    doc.add_paragraph(
        "Пользователи пропускают до 40 % важных уведомлений, "
        "потому что получают их без приоритизации."
    )
    doc.add_heading("Решение", level=1)
    for item in ("Приоритизация по контексту", "Единый центр уведомлений", "Тихие часы"):
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("Результат пилота", level=1)
    doc.add_paragraph(
        "Экономия составила 12,5 млн ₽ за год при росте открываемости с 31 % до 44 %."
    )
    p = DOCS_DIR / "product_description.docx"
    doc.save(p)
    outs.append(p)

    wb = Workbook()
    ws = wb.active
    ws.title = "Метрики"
    ws.append(["Месяц", "Открываемость, %", "Отписки, %"])
    for row in (["Май", 31, 4.1], ["Июнь", 38, 3.2], ["Июль", 44, 2.7]):
        ws.append(row)
    p = DOCS_DIR / "metrics.xlsx"
    wb.save(p)
    outs.append(p)

    p = DOCS_DIR / "metrics.csv"
    with p.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["Месяц", "Открываемость", "Отписки"])
        w.writerows([["Май", 31, 4.1], ["Июнь", 38, 3.2], ["Июль", 44, 2.7]])
    outs.append(p)

    p = DOCS_DIR / "brief.md"
    p.write_text(
        "# Запуск сервиса умных уведомлений\n\n"
        "Назначение: продукт. Аудитория: руководители продуктовых направлений.\n\n"
        "Цель: получить одобрение на пилот.\n\n"
        "## Обязательно включить\n\n- метрики пилота\n- план на квартал\n"
    )
    outs.append(p)

    p = DOCS_DIR / "notes.txt"
    p.write_text(
        "Заметки по пилоту\n\n"
        "Пилот шёл с мая по июль 2026 года в трёх регионах.\n"
        "Команда пилота — 12 человек, бюджет 2,4 млн ₽.\n\n"
        "Риски:\n- зависимость от push-провайдера\n- нагрузка на поддержку в первые недели\n"
    )
    outs.append(p)

    p = DOCS_DIR / "report.pdf"
    p.write_bytes(
        make_text_pdf(
            [
                ("Pilot report 2026", 20),
                ("Revenue grew by 25 % to 340 mln RUB in Q2 2026.", 11),
                ("Churn dropped from 4.1 % to 2.7 % during the pilot.", 11),
                ("- Priority inbox", 11),
                ("- Quiet hours", 11),
            ]
        )
    )
    outs.append(p)

    p = DOCS_DIR / "scan.pdf"
    scan = make_png(800, 600, "Скан без текстового слоя", (240, 240, 240))
    Image.open(io.BytesIO(scan)).convert("RGB").save(p, format="PDF")
    outs.append(p)

    p = DOCS_DIR / "chart.png"
    p.write_bytes(make_png(640, 400, "Открываемость по месяцам", (255, 255, 255), bars=True))
    outs.append(p)
    p = DOCS_DIR / "logo.png"
    p.write_bytes(make_logo(128))
    outs.append(p)

    doc = Document()
    doc.add_heading("Отчёт с таблицей и рисунком", level=1)
    doc.add_paragraph("Выручка компании «Ромашка» за 2025 год выросла на 25 % год к году.")
    table = doc.add_table(rows=3, cols=3)
    for r, row in enumerate(
        (
            ["Квартал", "Выручка, млн ₽", "Клиенты"],
            ["1 кв.", "80", "1 200"],
            ["2 кв.", "95", "1 450"],
        )
    ):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    doc.add_paragraph("Рисунок 1. Динамика открываемости", style="Caption")
    doc.add_picture(
        io.BytesIO(make_png(640, 400, "Динамика", (255, 255, 255), bars=True)), width=Inches(4)
    )
    p = DOCS_DIR / "report_with_table.docx"
    doc.save(p)
    outs.append(p)
    return outs


def make_png(
    width: int, height: int, title: str, background: tuple[int, int, int], bars: bool = False
) -> bytes:
    img = Image.new("RGB", (width, height), background)
    draw = ImageDraw.Draw(img)
    draw.text((24, 20), title, fill=(30, 30, 30))
    if bars:
        values = [31, 38, 44]
        for i, value in enumerate(values):
            x0 = 80 + i * 160
            draw.rectangle([x0, height - 40 - value * 6, x0 + 100, height - 40], fill=(0, 119, 255))
            draw.text((x0 + 30, height - 32), str(value), fill=(30, 30, 30))
    buf = io.BytesIO()
    img.save(buf, format="PNG", optimize=True)
    return buf.getvalue()


def make_mockup_png(width: int, height: int) -> bytes:
    """Схематичный экран «центра уведомлений»: шапка и строки с метками приоритета. Картинка
    с содержанием, а не однотонная заготовка: импорт отсеивает почти пустые изображения."""
    img = Image.new("RGB", (width, height), (244, 247, 252))
    draw = ImageDraw.Draw(img)
    draw.rectangle([0, 0, width, 96], fill=(0, 119, 255))
    draw.rounded_rectangle([40, 28, 360, 68], radius=12, fill=(255, 255, 255))
    colors = [(255, 72, 72), (255, 170, 0), (0, 119, 255), (160, 170, 185), (160, 170, 185)]
    for i, color in enumerate(colors):
        top = 140 + i * 124
        draw.rounded_rectangle(
            [40, top, width - 40, top + 100],
            radius=16,
            fill=(255, 255, 255),
            outline=(220, 226, 235),
        )
        draw.ellipse([64, top + 30, 104, top + 70], fill=color)
        draw.rounded_rectangle(
            [130, top + 28, 130 + 420 - i * 40, top + 46], radius=8, fill=(40, 44, 52)
        )
        draw.rounded_rectangle(
            [130, top + 58, 130 + 640 - i * 60, top + 72], radius=7, fill=(190, 198, 210)
        )
    buf = io.BytesIO()
    img.save(buf, format="PNG", optimize=True)
    return buf.getvalue()


def make_logo(size: int) -> bytes:
    img = Image.new("RGBA", (size, size), (0, 0, 0, 0))
    draw = ImageDraw.Draw(img)
    draw.ellipse((8, 8, size - 8, size - 8), fill=(0, 119, 255, 255))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def make_text_pdf(lines: list[tuple[str, int]]) -> bytes:
    """Минимальный PDF с текстовым слоем (Helvetica, латиница): без внешних библиотек."""
    content_lines = ["BT"]
    y = 780
    for text, size in lines:
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        content_lines.append(f"/F1 {size} Tf 1 0 0 1 60 {y} Tm ({escaped}) Tj")
        y -= size + 14
    content_lines.append("ET")
    content = "\n".join(content_lines).encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return out.getvalue()


def make_examples() -> list[pathlib.Path]:
    """Демонстрационный контент-пакет: текст с числами, таблица, две картинки, бриф."""
    EXAMPLES_DIR.mkdir(parents=True, exist_ok=True)
    outs: list[pathlib.Path] = []

    doc = Document()
    doc.add_heading("Умные уведомления: итоги пилота", level=1)
    doc.add_heading("Проблема", level=2)
    doc.add_paragraph(
        "Пользователи пропускают до 40 % важных уведомлений, потому что получают их "
        "без приоритизации: в среднем 27 уведомлений в день на человека."
    )
    doc.add_heading("Решение", level=2)
    for item in (
        "Приоритизация по контексту: важные уведомления поднимаются наверх",
        "Единый центр уведомлений вместо трёх разрозненных каналов",
        "Тихие часы с отложенной доставкой",
    ):
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("Результаты пилота", level=2)
    doc.add_paragraph(
        "Пилот шёл с мая по июль 2026 года в трёх регионах. Открываемость выросла с 31 % "
        "до 44 %, доля отписок снизилась с 4,1 % до 2,7 %. Экономия на поддержке составила "
        "12,5 млн ₽ за год при затратах на пилот 2,4 млн ₽."
    )
    doc.add_paragraph(
        "«Впервые уведомления стали помогать, а не отвлекать», — руководитель службы "
        "поддержки региона «Север».",
        style="Quote",
    )
    doc.add_heading("План на квартал", level=2)
    for item in (
        "Октябрь: подключение двух новых продуктов",
        "Ноябрь: персональные тихие часы",
        "Декабрь: расширение на все регионы, цель — 60 % открываемости",
    ):
        doc.add_paragraph(item, style="List Number")
    doc.add_paragraph("Рисунок 1. Центр уведомлений", style="Caption")
    doc.add_picture(io.BytesIO(make_mockup_png(1280, 800)), width=Inches(5))
    p = EXAMPLES_DIR / "overview.docx"
    doc.save(p)
    outs.append(p)

    wb = Workbook()
    ws = wb.active
    ws.title = "Метрики"
    ws.append(["Месяц", "Открываемость, %", "Отписки, %", "Активные пользователи"])
    for row in (["Май", 31, 4.1, 120500], ["Июнь", 38, 3.2, 131200], ["Июль", 44, 2.7, 140800]):
        ws.append(row)
    ws2 = wb.create_sheet("Экономика")
    ws2.append(["Статья", "Сумма, млн ₽"])
    for row in (
        ["Затраты на пилот", 2.4],
        ["Экономия на поддержке за год", 12.5],
        ["Ожидаемая экономия при масштабировании", 31],
    ):
        ws2.append(row)
    p = EXAMPLES_DIR / "metrics.xlsx"
    wb.save(p)
    outs.append(p)

    p = EXAMPLES_DIR / "notes.md"
    p.write_text(
        "# Заметки к выступлению\n\n"
        "Аудитория уже знакома с продуктом, детали инфраструктуры не нужны.\n\n"
        "## Риски\n\n- зависимость от push-провайдера\n- нагрузка на поддержку в первые недели\n\n"
        "## Просьба\n\nОдобрить расширение пилота на все продукты в четвёртом квартале.\n"
    )
    outs.append(p)

    p = EXAMPLES_DIR / "openrate_chart.png"
    p.write_bytes(make_png(960, 540, "Открываемость по месяцам, %", (255, 255, 255), bars=True))
    outs.append(p)
    p = EXAMPLES_DIR / "logo.png"
    p.write_bytes(make_logo(256))
    outs.append(p)

    p = EXAMPLES_DIR / "brief.json"
    p.write_text(
        json.dumps(
            {
                "purpose": "product",
                "title": "Запуск сервиса умных уведомлений",
                "audience": "руководители продуктовых направлений",
                "goal": "получить одобрение на расширение пилота",
                "language": "ru",
                "tone": "деловой, без технических деталей",
                "must_include": ["метрики пилота", "план на квартал"],
                "avoid": ["технические детали инфраструктуры"],
                "slide_count": {"min": 10, "max": 15},
            },
            ensure_ascii=False,
            indent=2,
        )
        + "\n"
    )
    outs.append(p)
    return outs


if __name__ == "__main__":
    print(make_template())
    for path in make_content():
        print(path)
    for path in make_examples():
        print(path)
