"""Сборка ContentPackage: валидность, кэш разбора (смена брифа не перечитывает файлы,
удаление одного материала — тоже), режим брифа без выдуманных показателей, недостающие
данные, изображения, порядок файлов в ключе; настоящий слой через API и CLI."""

from __future__ import annotations

import json
import pathlib
from typing import Any

from presentation_designer.contracts import ContentPackage
from presentation_designer.contracts.validators import check_content_package
from presentation_designer.parsing.content import importer as imp
from presentation_designer.parsing.content.parsers import ParseCache
from presentation_designer.shared.settings import Settings
from tests.parsing.content.conftest import CONTENT, EXAMPLES


def _import(files: Any, brief: Any, settings: Settings, cache: ParseCache, **kw: Any) -> Any:
    return imp.import_content(
        files, brief, package_id="pkg_t", settings=settings, cache=cache, use_model=False, **kw
    )


def test_package_from_fixtures_is_valid(
    materials: Any, cache: ParseCache, import_settings: Settings
) -> None:
    files = materials("product_description.docx", "metrics.xlsx", "brief.md", "chart.png")
    brief = {
        "purpose": "product",
        "title": "Запуск сервиса умных уведомлений",
        "must_include": ["метрики пилота", "план на квартал"],
    }
    result = _import(files, brief, import_settings, cache)
    pkg = result.package
    doc = ContentPackage.model_validate(pkg)
    assert not check_content_package(doc)
    assert pkg["schema_version"] == "1.3" and pkg["mode"] == "mixed"
    assert [s["source_id"] for s in pkg["sources"]] == [
        "src_brief",
        "src_1",
        "src_2",
        "src_3",
        "src_4",
    ]
    assert pkg["sources"][1]["parser"] == {"name": "docx", "version": "0.1.0"}
    assert pkg["sources"][1]["units"]["chars"] > 100 and pkg["sources"][1]["file_id"] == "file_1"
    assert pkg["sources"][2]["units"] == {"sheets": 1, "tables": 1}
    kinds = [b["kind"] for b in pkg["blocks"]]
    assert kinds.count("table") == 1 and kinds.count("figure") == 1
    table_block = next(b for b in pkg["blocks"] if b["kind"] == "table")
    assert table_block["dataset_id"] == "ds1" and table_block["importance"] == "must"
    assert table_block["source_location"] == {"sheet": "Метрики", "cell_range": "A1:C4"}
    figure = next(b for b in pkg["blocks"] if b["kind"] == "figure")
    asset = pkg["assets"][0]
    assert figure["asset_id"] == asset["asset_id"] == "img_1"
    assert asset["path"] == "assets/img_1.png" and asset["mime"] == "image/png"
    assert asset["sha256"] and result.assets["assets/img_1.png"].startswith(b"\x89PNG")
    must = [f for f in pkg["facts"] if f["must_keep"]]
    assert {f["raw"] for f in must} >= {"40 %", "12,5 млн ₽", "31 %", "44 %"}
    assert all(f["source_location"].get("fragment") for f in must)
    derived = [f for f in pkg["facts"] if f.get("derived")]
    assert derived and all(not f["must_keep"] for f in derived)
    paragraph = next(b for b in pkg["blocks"] if b["block_id"] == "b2")
    assert paragraph["importance"] == "must"
    assert pkg["missing_data"] == []
    meta = pkg["import_meta"]
    assert meta["importer"] == {"name": "content_importer", "version": imp.IMPORTER_VERSION}
    assert meta["parsers"] == {
        "docx": "0.1.0",
        "image": "0.1.0",
        "markdown": "0.1.0",
        "xlsx": "0.1.0",
    }
    assert meta["cache"] == {"files_hit": 0, "files_missed": 4} and meta["model_calls"] == 0
    assert meta["import_key"].startswith("sha256:")


def test_cache_reuse_on_brief_change_and_file_removal(
    materials: Any, cache: ParseCache, import_settings: Settings
) -> None:
    files = materials("product_description.docx", "metrics.xlsx")
    first = _import(files, {"purpose": "product", "title": "А"}, import_settings, cache)
    assert first.report["cache"] == {"files_hit": 0, "files_missed": 2}
    second = _import(
        files, {"purpose": "report", "title": "Б", "goal": "цель"}, import_settings, cache
    )
    assert second.report["cache"] == {"files_hit": 2, "files_missed": 0}
    assert second.package["brief"]["purpose"] == "report"
    # Содержание не зависит от брифа, ключ импорта — тоже.
    assert second.package["blocks"] == first.package["blocks"]
    assert second.package["import_meta"]["import_key"] == first.package["import_meta"]["import_key"]
    third = _import(files[:1], {"purpose": "product", "title": "А"}, import_settings, cache)
    assert third.report["cache"] == {"files_hit": 1, "files_missed": 0}
    assert third.package["import_meta"]["import_key"] != first.package["import_meta"]["import_key"]
    # Порядок файлов входит в ключ и меняет нумерацию блоков.
    swapped = _import(list(reversed(files)), None, import_settings, cache)
    assert (
        swapped.package["import_meta"]["import_key"] != first.package["import_meta"]["import_key"]
    )
    assert swapped.package["blocks"][0]["kind"] == "table"
    assert swapped.package["mode"] == "package" and swapped.package["sources"][0]["kind"] == "xlsx"
    assert any(w["code"] == "brief_incomplete" for w in first.package["warnings"]) is False
    assert any(w["code"] == "brief_incomplete" for w in swapped.package["warnings"]) is False


def test_brief_mode_marks_missing_data_and_invents_nothing(
    cache: ParseCache, import_settings: Settings
) -> None:
    brief = {
        "purpose": "product",
        "title": "Новый продукт",
        "goal": "одобрить пилот",
        "must_include": ["финансовые результаты"],
        "notes": "Пилот на 3 квартала.",
    }
    result = _import([], brief, import_settings, cache)
    pkg = result.package
    assert pkg["mode"] == "brief" and pkg["sources"] == [
        {
            "source_id": "src_brief",
            "kind": "user_input",
            "name": "бриф",
            "extracted": True,
            "parser": {"name": "brief", "version": "0.1.0"},
        }
    ]
    assert [b["kind"] for b in pkg["blocks"]] == ["heading", "paragraph", "bullets", "paragraph"]
    assert all(b["source_id"] == "src_brief" for b in pkg["blocks"])
    assert pkg["facts"] == []  # «3 квартала» — не показатель
    assert {m["what"] for m in pkg["missing_data"]} == {
        "финансовые результаты",
        "показатели и цифры",
    }
    assert not check_content_package(ContentPackage.model_validate(pkg))


def test_incomplete_brief_and_unsupported_or_broken_files(
    materials: Any, cache: ParseCache, import_settings: Settings, tmp_path: pathlib.Path
) -> None:
    broken = tmp_path / "broken.docx"
    broken.write_bytes(b"not a zip")
    video = tmp_path / "clip.mp4"
    video.write_bytes(b"\x00" * 10)
    files = [
        *materials("scan.pdf"),
        *materials("broken.docx", root=tmp_path),
        *materials("clip.mp4", root=tmp_path),
    ]
    result = _import(files, {"title": "Только тема"}, import_settings, cache)
    pkg = result.package
    assert pkg["brief"]["purpose"] == "other" and pkg["brief"]["language"] == "ru"
    assert any(
        w["code"] == "brief_incomplete" and "purpose" in w["message"] for w in pkg["warnings"]
    )
    scan, docx, mp4 = pkg["sources"][1:]
    assert any(w["code"] == "no_text_layer" for w in scan["warnings"])
    assert docx["extracted"] is False and docx["warnings"][0]["code"] == "docx_unreadable"
    assert mp4["extracted"] is False and mp4["warnings"][0]["code"] == "format_unsupported"
    assert ContentPackage.model_validate(pkg)


def test_examples_package_has_text_numbers_table_and_two_images(
    materials: Any, cache: ParseCache, import_settings: Settings
) -> None:
    files = materials(
        "overview.docx", "metrics.xlsx", "notes.md", "openrate_chart.png", "logo.png", root=EXAMPLES
    )
    brief = json.loads((EXAMPLES / "brief.json").read_text())
    result = _import(files, brief, import_settings, cache)
    pkg = result.package
    assert not check_content_package(ContentPackage.model_validate(pkg))
    assert len(pkg["datasets"]) == 2 and len(pkg["assets"]) == 3
    assert pkg["brief"]["must_include"] == ["метрики пилота", "план на квартал"]
    assert pkg["brief"]["slide_count"] == {"min": 10, "max": 15}
    assert sum(1 for f in pkg["facts"] if f["must_keep"]) >= 6
    assert pkg["missing_data"] == []


def test_api_real_import_publishes_assets(tmp_path: pathlib.Path, monkeypatch: Any) -> None:
    """Настоящий слой импорта через API на встроенном исполнителе: пакет валиден,
    изображения доступны по манифесту, тот же набор файлов даёт тот же пакет."""
    from fastapi.testclient import TestClient

    from presentation_designer.api.app import create_app
    from presentation_designer.pipeline.artifacts import ArtifactStore
    from presentation_designer.pipeline.files import FileStore
    from presentation_designer.pipeline.jobs import InlineExecutor, Orchestrator
    from presentation_designer.pipeline.real import RealLayers
    from presentation_designer.pipeline.state import State

    monkeypatch.setenv("PD_QWEN_BASE_URL", "https://example.invalid/v1")
    monkeypatch.setenv("PD_QWEN_API_KEY", "replace-me")
    settings = Settings()
    settings.paths.data_dir = tmp_path / "data"
    settings.paths.artifacts_dir = tmp_path / "artifacts"
    settings.paths.runs_dir = tmp_path / "runs"
    settings.content_import.cache_dir = tmp_path / "import-cache"
    state = State(settings.db_path)
    orch = Orchestrator(
        settings,
        state,
        FileStore(settings.uploads_dir, max_upload_mb=2, max_unzipped_mb=64),
        ArtifactStore(settings.artifacts_dir),
        RealLayers(settings),
        InlineExecutor(),
    )
    docx = (CONTENT / "report_with_table.docx").read_bytes()
    with TestClient(create_app(orch, reconcile=False)) as client:
        health = client.get("/api/health").json()
        assert health["execution_mode"]["layers"]["parsing.content"] == "real"
        assert health["execution_mode"]["layers"]["brief"] == "real"
        assert health["execution_mode"]["layers"]["generation.story"] == "real"
        r = client.post(
            "/api/content",
            files=[("files", ("report.docx", docx, "application/octet-stream"))],
            data={"brief": json.dumps({"purpose": "report", "title": "Отчёт", "language": "ru"})},
        )
        assert r.status_code == 202, r.text
        package_id = r.json()["package_id"]
        detail = client.get(f"/api/content/{package_id}").json()
        assert detail["status"] == "succeeded", detail.get("error")
        pkg = ContentPackage.model_validate(detail["package"])
        assert pkg.mode == "mixed" and len(pkg.assets) == 1 and len(pkg.datasets) == 1
        asset = detail["package"]["assets"][0]["path"]
        png = client.get(f"/api/content/{package_id}/assets/{asset}")
        assert png.status_code == 200 and png.content.startswith(b"\x89PNG")
        assert client.get(f"/api/content/{package_id}/assets/../x").status_code == 404
        r2 = client.post(
            "/api/content",
            files=[("files", ("report.docx", docx, "application/octet-stream"))],
            data={"brief": json.dumps({"purpose": "report", "title": "Отчёт", "language": "ru"})},
        )
        assert r2.json()["cached"] is True and r2.json()["package_id"] == package_id
        # Бриф без модели: провайдер не настроен, отвечает эвристика с отметкой источника.
        brief = client.post("/api/brief", json={"text": "Сделай отчёт про итоги квартала"}).json()
        assert brief["source"] == "heuristic" and brief["brief"]["purpose"] == "report"


def test_cli_import_writes_package_assets_manifest(
    tmp_path: pathlib.Path, monkeypatch: Any
) -> None:
    from presentation_designer.cli.main import main

    monkeypatch.setenv("PD_CONTENT_IMPORT__CACHE_DIR", str(tmp_path / "cache"))
    out = tmp_path / "out"
    code = main(["import", str(EXAMPLES), "--out", str(out), "--no-model"])
    assert code == 0
    package = json.loads((out / "package.json").read_text())
    assert ContentPackage.model_validate(package)
    assert package["brief"]["title"] == "Запуск сервиса умных уведомлений"
    manifest = json.loads((out / "manifest.json").read_text())
    assert manifest["command"] == "import" and len(manifest["inputs"]["files"]) == 5
    assert manifest["inputs"]["brief"].endswith("brief.json")
    assert sorted(p.name for p in (out / "assets").iterdir()) == [
        "img_1.png",
        "img_2.png",
        "img_3.png",
    ]
    report = json.loads((out / "report.json").read_text())
    assert report["counts"]["datasets"] == 2 and report["model_used"] is False
    assert main(["import", str(tmp_path / "нет"), "--out", str(out)]) == 2


def test_pptx_material_with_numbers_gives_valid_facts(
    tmp_path: pathlib.Path, cache: ParseCache, import_settings: Settings
) -> None:
    """Факт из презентации-материала: у блока есть номер слайда, у факта — только страница,
    лист и фрагмент, как требует контракт (18.09: пакет не проходил проверку из-за slide)."""
    from pptx import Presentation

    from tests.parsing.content.conftest import material

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Итоги пилота"
    slide.placeholders[1].text = "Выручка выросла на 25 % за 2025 год, 340 клиентов подключены"
    path = tmp_path / "deck.pptx"
    prs.save(path)
    result = _import([material(path)], {"purpose": "product"}, import_settings, cache)
    pkg = result.package
    doc = ContentPackage.model_validate(pkg)
    assert not check_content_package(doc)
    assert pkg["facts"], "числа со слайда стали фактами"
    assert all("slide" not in (f.get("source_location") or {}) for f in pkg["facts"])
    assert any((b.get("source_location") or {}).get("slide") == 1 for b in pkg["blocks"])
