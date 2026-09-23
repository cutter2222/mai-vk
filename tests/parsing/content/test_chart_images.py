"""Vision boundary is mocked; parsing, contracts and native PPTX/workbook are real."""

from __future__ import annotations

import asyncio
import copy
import io
import pathlib
import zipfile
from typing import Any

import openpyxl
import pytest
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

from presentation_designer.contracts import ContentPackage
from presentation_designer.contracts.validators import check_content_package
from presentation_designer.layout import charts
from presentation_designer.llm.types import Response
from presentation_designer.parsing.content.chart_images import Extraction, extract_charts
from presentation_designer.parsing.content.importer import import_content
from presentation_designer.parsing.content.parsers.image import parse
from presentation_designer.shared.settings import Settings
from tests.parsing.content.conftest import material


def readable(kind: str = "column") -> dict[str, Any]:
    return {
        "status": "readable",
        "reason": "Все значения подписаны",
        "type": kind,
        "title": "Конверсия",
        "unit": "%",
        "categories": ["2024", "2025"],
        "series": [
            {
                "name": "Пилот",
                "points": [
                    {"value": 12.5, "label": "12,5%", "basis": "label"},
                    {"value": 40, "label": "40%", "basis": "label"},
                ],
            }
        ],
        "axis_minimum": 0,
        "axis_maximum": 50,
    }


class Vision:
    def __init__(self, answer: Any) -> None:
        self.answer = answer
        self.requests: list[Any] = []

    async def complete(self, req: Any) -> Response:
        self.requests.append(req)
        if isinstance(self.answer, Exception):
            raise self.answer
        return Response(text="", parsed=copy.deepcopy(self.answer))


@pytest.fixture
def chart_png(tmp_path: pathlib.Path) -> pathlib.Path:
    path = tmp_path / "chart.png"
    image = Image.new("RGB", (1200, 900), "white")
    draw = ImageDraw.Draw(image)
    draw.line((100, 100, 100, 800, 1100, 800), fill="black", width=3)
    for x, top, label, category in [(250, 625, "12.5%", "2024"), (650, 240, "40%", "2025")]:
        draw.rectangle((x, top, x + 180, 800), fill="#0077FF")
        draw.text((x, top - 40), label, fill="black", font_size=30)
        draw.text((x, 820), category, fill="black", font_size=30)
    image.save(path)
    return path


def run_import(path: pathlib.Path, answer: Any, settings: Settings, **kw: Any) -> Any:
    return import_content(
        [material(path)],
        None,
        package_id="pkg_chart",
        settings=settings,
        llm_client=Vision(answer),
        **kw,
    )


@pytest.mark.parametrize("embedded", [False, True])
def test_import_raster_and_pptx_with_provenance(
    chart_png: pathlib.Path,
    import_settings: Settings,
    embedded: bool,
) -> None:
    path = chart_png
    if embedded:
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(str(chart_png), Inches(1), Inches(1))
        path = chart_png.with_suffix(".pptx")
        prs.save(path)
    client = Vision(readable())
    result = import_content(
        [material(path)], None, package_id="pkg_chart", settings=import_settings, llm_client=client
    )
    package = result.package
    model = ContentPackage.model_validate(package)
    assert not check_content_package(model)
    assert model.datasets[0].source_chart is not None
    ds = package["datasets"][0]
    assert ds["rows"] == [["2024", 12.5], ["2025", 40]]
    assert ds["columns"][0]["type"] == "string"
    assert ds["source_chart"] == {
        "type": "column",
        "asset_id": "img_1",
        "axis_minimum": 0,
        "axis_maximum": 50,
    }
    assert package["blocks"][0]["dataset_id"] == ds["dataset_id"]
    assert package["assets"][0]["kind"] == "chart_image"
    assert result.assets["assets/img_1.png"] == chart_png.read_bytes()
    assert all(f["uncertainty"]["extracted_by"] == "model" for f in package["facts"])
    if embedded:
        assert ds["source_location"] == {"slide": 1}
    req = client.requests[0]
    assert req.role == "vlm" and req.images[0].detail == "high"
    assert Image.open(io.BytesIO(req.images[0].data)).size == (1200, 900)
    assert package["import_meta"]["model_calls"] == 1
    assert result.report["chart_images"]["items"]
    # Vision results never contaminate deterministic parse cache.
    offline = run_import(path, readable(), import_settings, use_model=False)
    assert not offline.package["datasets"]
    assert offline.report["cache"]["files_hit"] == 1


@pytest.mark.parametrize(
    "case",
    [
        "estimated",
        "unknown",
        "mismatch",
        "nan",
        "bool",
        "short",
        "units",
        "axis",
        "uncertain",
        "unsupported",
    ],
)
def test_untrustworthy_values_retain_source(
    chart_png: pathlib.Path,
    import_settings: Settings,
    case: str,
) -> None:
    answer = readable()
    point = answer["series"][0]["points"][0]
    if case == "estimated":
        point["basis"] = "estimated"
    elif case == "unknown":
        point.update(value=None, label=None, basis="unreadable")
    elif case == "mismatch":
        point["value"] = 99
    elif case == "nan":
        point["value"] = float("nan")
    elif case == "bool":
        point["value"] = True
    elif case == "short":
        answer["series"][0]["points"].pop()
    elif case == "units":
        answer["unit"] = "млн"
    elif case == "axis":
        answer["axis_minimum"] = 10
    else:
        answer["status"] = case
    result = run_import(chart_png, answer, import_settings)
    assert not result.package["datasets"] and not result.package["facts"]
    assert result.package["blocks"][0]["kind"] == "figure"
    assert result.package["warnings"][0]["code"] == "chart_image_retained"
    assert result.package["missing_data"]
    assert result.assets["assets/img_1.png"] == chart_png.read_bytes()


def test_provider_error_limits_and_nonchart(
    chart_png: pathlib.Path, import_settings: Settings
) -> None:
    for answer in (RuntimeError("provider unavailable"), {"invalid": "response"}):
        result = run_import(chart_png, answer, import_settings)
        assert not result.package["datasets"] and result.package["warnings"]
    answer = readable()
    answer["status"] = "not_chart"
    result = run_import(chart_png, answer, import_settings)
    assert not result.package["datasets"] and not result.package["warnings"]
    import_settings.content_import.chart_image_max_images = 0
    result = run_import(chart_png, readable(), import_settings)
    assert result.report["chart_images"]["calls"] == 0
    assert result.package["warnings"]


@pytest.mark.parametrize("kind", ["column", "bar", "line"])
def test_native_chart_has_exact_workbook_and_design_style(
    chart_png: pathlib.Path,
    import_settings: Settings,
    kind: str,
) -> None:
    ds = run_import(chart_png, readable(kind), import_settings).package["datasets"][0]
    # Planner must not be able to silently change the transcribed chart's type or series.
    spec = charts.chart_spec({"type": "pie", "series": ["nonexistent"], "dataset_id": "ds1"}, ds)
    assert spec.chart_type == kind and spec.units == "%"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    style = charts.ChartStyle(font_family="Arial", accents=["#AA2255"])
    box = (Inches(1), Inches(1), Inches(8), Inches(5))
    # Even same-family native template charts cannot impose the wrong chart type/axis.
    old = charts.add_chart(slide, box, charts.chart_spec({"type": "column"}, ds), style)
    _frame, how = charts.replace_or_add_chart(slide, old, box, spec, style)
    assert how == "rebuilt"
    target = chart_png.with_suffix(".pptx")
    prs.save(target)
    chart = Presentation(target).slides[0].shapes[0].chart
    assert chart.chart_type == charts.CHART_TYPES[kind]
    assert chart.series[0].values == (12.5, 40.0)
    assert chart.value_axis.minimum_scale == 0 and chart.value_axis.maximum_scale == 50
    assert chart.value_axis.axis_title.text_frame.text == "%"
    assert "AA2255" in chart._chartSpace.xml and "Arial" in chart._chartSpace.xml
    with zipfile.ZipFile(target) as zf:
        workbook = next(n for n in zf.namelist() if n.startswith("ppt/embeddings/"))
        wb = openpyxl.load_workbook(io.BytesIO(zf.read(workbook)))
    assert list(wb.active.values) == [(None, "Пилот"), ("2024", 12.5), ("2025", 40)]


def test_unknown_values_are_gaps_not_zero(chart_png: pathlib.Path) -> None:
    ds = {
        "columns": [{"name": "Период", "type": "string"}, {"name": "Ряд", "type": "number"}],
        "rows": [["A", 0], ["B", None], ["C", "н/д"], ["D", float("inf")], ["E"]],
    }
    spec = charts.chart_spec({"type": "line"}, ds)
    assert spec.series == [("Ряд", [0.0, None, None, None, None])]
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    charts.add_chart(slide, (0, 0, Inches(8), Inches(5)), spec, charts.ChartStyle())
    target = chart_png.with_suffix(".pptx")
    prs.save(target)
    assert Presentation(target).slides[0].shapes[0].chart.series[0].values == (
        0.0,
        None,
        None,
        None,
        None,
    )


def test_scale_not_recomputed() -> None:
    answer = readable()
    answer["unit"] = "млн"
    for point in answer["series"][0]["points"]:
        point["label"] = point["label"].replace("%", " млн")
    Extraction.model_validate(answer).validate_readable()


def test_dedup_deadline_and_row_limit(chart_png: pathlib.Path, import_settings: Settings) -> None:
    image = parse(chart_png).images[0]
    client = Vision(readable())
    accepted, report = extract_charts([image, image], client, budget_s=1, max_images=8)
    assert len(accepted) == report["calls"] == len(client.requests) == 1

    class SlowVision:
        async def complete(self, req: Any) -> Response:
            await asyncio.sleep(1)
            return Response(text="", parsed=readable())

    accepted, report = extract_charts([image], SlowVision(), budget_s=0.02, max_images=8)
    assert not accepted and report["calls"] == 1
    assert next(iter(report["items"].values()))["status"] == "rejected"
    import_settings.content_import.max_dataset_rows = 1
    result = run_import(chart_png, readable(), import_settings)
    assert not result.package["datasets"]
    assert result.package["warnings"] and result.package["blocks"][0]["kind"] == "figure"


def test_missing_chart_source_is_invalid(
    chart_png: pathlib.Path, import_settings: Settings
) -> None:
    package = run_import(chart_png, readable(), import_settings).package
    package["datasets"][0]["source_chart"]["asset_id"] = "missing"
    assert any(
        v.code == "asset_missing"
        for v in check_content_package(ContentPackage.model_validate(package))
    )
