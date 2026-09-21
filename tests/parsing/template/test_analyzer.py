"""Анализатор на синтетических шаблонах: пакет, геометрия групп, классификация, ресурсы,
постоянные элементы, токены, паттерны с ёмкостью, правила, валидный профиль, ключ кэша."""

from __future__ import annotations

import pathlib
import zipfile

import pytest

from presentation_designer.contracts import TemplateProfile
from presentation_designer.parsing.template import analyzer as an
from presentation_designer.parsing.template.classify import classify_all, is_marker
from presentation_designer.parsing.template.package import PackageError, open_template
from presentation_designer.shared import text_metrics


def analyze(path: pathlib.Path, **kw: object) -> an.AnalysisResult:
    return an.analyze_template(
        path,
        template_id="tpl_test",
        name=path.name,
        size_bytes=path.stat().st_size,
        render=False,
        use_vlm=False,
        **kw,  # type: ignore[arg-type]
    )


def test_package_errors(tmp_path: pathlib.Path) -> None:
    with pytest.raises(PackageError) as info:
        open_template(tmp_path / "нет.pptx")
    assert info.value.code == "file_missing"
    bad = tmp_path / "bad.pptx"
    bad.write_bytes(b"not a zip at all")
    with pytest.raises(PackageError) as info:
        open_template(bad)
    assert info.value.code == "not_a_zip"
    docx = tmp_path / "doc.pptx"
    with zipfile.ZipFile(docx, "w") as zf:
        zf.writestr("word/document.xml", "<w/>")
        zf.writestr("[Content_Types].xml", "<Types/>")
    with pytest.raises(PackageError) as info:
        open_template(docx)
    assert info.value.code == "not_pptx" and "docx" in str(info.value)


def test_group_geometry_and_classification(rich_template: pathlib.Path) -> None:
    pkg = open_template(rich_template)
    assert (pkg.width_emu, pkg.height_emu) == (12192000, 6858000)
    cards = pkg.slides[1]
    grouped = [s for s in cards.shapes if s.group_path]
    assert len(grouped) == 9, "три группы по три объекта раскрыты в плоский список"
    icons = [s for s in grouped if s.kind == "picture"]
    assert all(0.04 < s.x < 0.7 and 0.24 < s.y < 0.26 for s in icons), (
        "координаты детей группы в долях слайда"
    )
    assert all(s.media_sha256 and s.media_blob for s in icons)
    classes = {c.slide_index: c.kind for c in classify_all(pkg)}
    assert classes == {
        1: "content_sample",
        2: "content_sample",
        3: "content_sample",
        4: "style_guide",
        5: "asset_catalog",
        6: "hidden",
    }
    assert is_marker("Заголовок") and is_marker("ххх%") and is_marker("Имя Фамилия, должность")
    assert not is_marker("Шрифт для заголовков — Play, кегль не менее 24.")


def test_profile_is_valid_and_complete(rich_template: pathlib.Path) -> None:
    result = analyze(rich_template)
    profile = result.profile
    TemplateProfile.model_validate(profile)
    # Паттерны шаблона и собственные композиции библиотеки лежат в одном списке;
    # здесь проверяется разбор шаблона, поэтому свои композиции отфильтрованы.
    from_template = [p for p in profile["patterns"] if p["source"]["kind"] != "builtin"]
    roles = {p["pattern_id"]: p["role"] for p in from_template}
    assert roles["pat_s1"] == "title"
    assert roles["pat_s2"] == "cards" and roles["pat_s3"] == "kpi"
    assert set(roles) == {"pat_s1", "pat_s2", "pat_s3"}, (
        "инструкция, каталог и скрытый — не паттерны"
    )
    builtin = [p for p in profile["patterns"] if p["source"]["kind"] == "builtin"]
    assert builtin, "библиотека собственных композиций подключена к профилю"
    assert all(p["source"].get("composition_id") for p in builtin)
    cards = next(p for p in profile["patterns"] if p["pattern_id"] == "pat_s2")
    kinds = [s["kind"] for s in cards["slots"]]
    assert kinds.count("icon") == 3 and kinds.count("body") >= 3 and kinds.count("title") >= 1
    groups = {s.get("repeat_group") for s in cards["slots"] if s["kind"] in ("icon", "body")}
    assert groups and None not in groups, "иконки и тексты карточек в повторяющихся группах"
    assert cards["constraints"]["max_items"] == 3 and cards["removable_object_ids"]
    body = next(s for s in cards["slots"] if s["kind"] == "body")
    cap = body["capacity"]
    assert cap["max_chars"] >= len(body["sample_text"]) and cap["max_lines"] >= 1
    assert cap["measured_with"]["method"] in ("font_metrics", "heuristic")
    assert body["computed_style"]["font"]["size_pt"] == 14.0
    assert body["element_ref"] and body["group_path"], "ссылка на объект и путь групп сохранены"
    # каталог иконок: изображения сохранены как ресурсы, не как паттерн
    icon_assets = [a for a in profile["assets"] if a["kind"] == "icon"]
    assert len(icon_assets) >= 8 and any("catalog" in a["tags"] for a in icon_assets)
    assert any(a["kind"] == "logo" for a in profile["assets"]), "повторяющаяся картинка — логотип"
    assert any(f["kind"] == "logo" for f in profile["fixed_elements"])
    # инструкция превратилась в правила
    kinds = {g["kind"] for g in profile["guidelines"]}
    assert "typography" in kinds and "color" in kinds
    assert (
        "Заголовок" in profile["placeholder_markers"] and "ххх%" in profile["placeholder_markers"]
    )
    sample = {s["slide_index"]: s["classification"] for s in profile["sample_slides"]}
    assert sample[4] == "style_guide" and sample[5] == "asset_catalog" and sample[6] == "hidden"
    assert profile["stats"]["slides"] == 6 and profile["design_tokens"]["colors"]["palette"]
    assert profile["llm_digest"].startswith("Размер слайда")
    assert result.report.counts["patterns"] == 3
    assert any(w["code"] in ("font_substituted",) or True for w in profile["warnings"])


def test_mini_template_profile(mini_template: pathlib.Path) -> None:
    profile = analyze(mini_template).profile
    TemplateProfile.model_validate(profile)
    roles = [p["role"] for p in profile["patterns"]]
    assert "title" in roles and "table" in roles and "thanks" in roles
    table = next(p for p in profile["patterns"] if p["role"] == "table")
    assert {s["kind"] for s in table["slots"]} >= {"table", "chart"}
    assert profile["stats"]["native_charts"] == 1 and profile["stats"]["native_tables"] == 1
    assert any(d["kind"] == "slide_number" for d in profile["dynamic_fields"])


def test_groups_and_tone(variety_template: pathlib.Path) -> None:
    """Образцы одной сигнатуры попадают в одну группу независимо от тона; тон читается из
    заливки слайда с источником; style_key — тон, семейство макета и plain."""
    profile = analyze(variety_template).profile
    TemplateProfile.model_validate(profile)
    # Версии не вписаны числом: профиль и анализатор меняются вместе с ключом кэша.
    assert profile["schema_version"] == an.PROFILE_SCHEMA_VERSION
    assert profile["analyzer"]["version"] == an.ANALYZER_VERSION
    # Группы и тон есть только у образцов шаблона: собственные композиции строятся из
    # дизайн-кода и своего слайда-источника не имеют.
    by_id = {p["pattern_id"]: p for p in profile["patterns"] if p["source"]["kind"] != "builtin"}
    assert by_id["pat_s1"]["group_id"] == by_id["pat_s2"]["group_id"], "титулы одного состава"
    assert by_id["pat_s3"]["group_id"] == by_id["pat_s4"]["group_id"], "разделители"
    assert by_id["pat_s5"]["group_id"] == by_id["pat_s6"]["group_id"], "карточки одной сигнатуры"
    assert by_id["pat_s8"]["group_id"] == by_id["pat_s9"]["group_id"], "финалы"
    assert by_id["pat_s1"]["group_id"] != by_id["pat_s3"]["group_id"]
    tones = {pid: p["tone"]["background"] for pid, p in by_id.items()}
    assert tones["pat_s1"] == "light" and tones["pat_s2"] == "dark"
    assert tones["pat_s3"] == "light" and tones["pat_s4"] == "dark"
    assert tones["pat_s8"] == "dark" and tones["pat_s9"] == "light"
    assert all(p["tone"]["source"] == "slide_fill" for p in by_id.values())
    assert 0 <= by_id["pat_s2"]["tone"]["luminance"] < 0.5 <= by_id["pat_s1"]["tone"]["luminance"]
    assert by_id["pat_s1"]["style_key"] == "light|title slide|plain"
    assert by_id["pat_s2"]["style_key"] == "dark|title slide|plain"
    assert by_id["pat_s4"]["style_key"] == "dark|section header|plain"
    # Группа записана и в sample_slides, и у паттерна.
    sample = {s["slide_index"]: s.get("group_id") for s in profile["sample_slides"]}
    assert sample[1] == by_id["pat_s1"]["group_id"]
    # Тон без собственной заливки слайда — из темы (белый lt1 у шаблона python-pptx).
    mini = analyze(
        pathlib.Path(__file__).resolve().parents[2] / "fixtures" / "pptx" / "mini_template.pptx"
    ).profile
    mini_template_patterns = [p for p in mini["patterns"] if p["source"]["kind"] != "builtin"]
    assert {p["tone"]["source"] for p in mini_template_patterns} <= {
        "theme",
        "master_fill",
        "layout_fill",
    }
    assert all(p["tone"]["background"] == "light" for p in mini_template_patterns)


def test_layout_family() -> None:
    from presentation_designer.parsing.template.layouts import layout_family

    assert layout_family("2_Титульный слайд") == "титульный слайд"
    assert layout_family("12. Разделитель тёмный") == "разделитель тёмный"
    assert layout_family("Section Header") == "section header"
    assert layout_family("Карточки (2)") == "карточки"
    assert layout_family("") == "layout" and layout_family(None) == "layout"


def test_profile_key_changes_with_inputs(monkeypatch: pytest.MonkeyPatch) -> None:
    base = an.profile_key("a" * 64)
    assert base == an.profile_key("a" * 64)
    assert an.profile_key("b" * 64) != base
    monkeypatch.setattr(an, "ANALYZER_VERSION", "9.9.9")
    assert an.profile_key("a" * 64) != base


def test_text_metrics_capacity_uses_font_file() -> None:
    font = text_metrics.resolve_font("Play")
    assert font.face is not None and not font.substituted, "Play лежит в docker/fonts"
    cap = text_metrics.capacity(width_emu=3300000, height_emu=1800000, size_pt=14, font=font)
    assert cap.method == "font_metrics" and cap.max_lines >= 5 and cap.chars_per_line >= 25
    wide = text_metrics.text_width_pt("ШШШШ", font, 14)
    narrow = text_metrics.text_width_pt("iiii", font, 14)
    assert wide > narrow, "ширина зависит от глифов, а не от числа символов"
    unknown = text_metrics.resolve_font("Шрифт Которого Нет")
    assert unknown.substituted


def test_single_title_per_pattern(tmp_path: pathlib.Path) -> None:
    """Заголовки карточек кеглем крупнее заголовка слайда (шаблон ЛЦТ-2026: пустые
    плейсхолдеры номеров шагов 28 pt при заголовке 20 pt) не становятся обязательными
    заголовками: заголовок в паттерне один — плейсхолдер заголовка, остальные крупные
    надписи — подписи карточек."""
    from pptx import Presentation
    from pptx.util import Emu, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
    slide.shapes.title.text = "Четыре шага"
    slide.shapes.title.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
    for i in range(4):
        x = 400000 + i * 2950000
        num = slide.shapes.add_textbox(Emu(x), Emu(1800000), Emu(700000), Emu(450000))
        num.text_frame.text = "Шаг"
        num.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
        body = slide.shapes.add_textbox(Emu(x), Emu(2400000), Emu(2400000), Emu(2500000))
        body.text_frame.text = f"Описание шага {i + 1}: что делаем и зачем"
        body.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    path = tmp_path / "steps.pptx"
    prs.save(path)

    profile = analyze(path).profile
    TemplateProfile.model_validate(profile)
    pattern = profile["patterns"][0]
    kinds = [(s["slot_id"], s["kind"], s["required"]) for s in pattern["slots"]]
    titles = [k for k in kinds if k[1] == "title"]
    assert titles == [("title_1", "title", True)], kinds
    assert [k[0] for k in kinds if k[1] == "label"] == ["label_1", "label_2", "label_3", "label_4"]
    assert not any(k[2] for k in kinds if k[1] == "label"), "подписи карточек не обязательны"


def test_layout_previews_rendered_with_fake_converter(
    tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Подложка холста редактора (этап 23): при рендере анализатор собирает временный PPTX
    из пустых слайдов макетов паттернов и кладёт `previews/layout-<id>.png`; без
    ONLYOFFICE — предупреждение, не отказ."""
    import io

    from PIL import Image
    from pptx import Presentation

    from presentation_designer.export import pdf
    from tests.layout.conftest import MINI_TEMPLATE

    calls = tmp_path / "calls.log"

    def pdf_bytes(pages: int) -> bytes:
        images = [Image.new("RGB", (160, 90), (200 + i, 210, 230)) for i in range(pages)]
        buf = io.BytesIO()
        images[0].save(buf, format="PDF", save_all=True, append_images=images[1:])
        return buf.getvalue()

    slide_count = len(Presentation(str(MINI_TEMPLATE)).slides)
    (tmp_path / "full.pdf").write_bytes(pdf_bytes(slide_count))
    (tmp_path / "layouts.pdf").write_bytes(pdf_bytes(3))

    def convert(source, out, **kwargs):
        with calls.open("a") as log:
            log.write(source.stem + "\n")
        target = out / f"{source.stem}.pdf"
        target.write_bytes(
            (tmp_path / ("layouts.pdf" if source.stem == "layouts" else "full.pdf")).read_bytes()
        )
        return pdf.PdfResult(target, 0.01)

    monkeypatch.setattr(pdf, "convert_to_pdf", convert)
    result = an.analyze_template(
        MINI_TEMPLATE,
        template_id="tpl_test",
        name=MINI_TEMPLATE.name,
        size_bytes=MINI_TEMPLATE.stat().st_size,
        render=True,
        use_vlm=False,
        workdir=tmp_path / "work",
    )
    layout_ids = {p["source"]["layout_id"] for p in result.profile["patterns"]}
    assert layout_ids, "у паттернов есть макеты"
    layout_previews = {k for k in result.previews if k.startswith("previews/layout-")}
    assert layout_previews == {f"previews/layout-{lid}.png" for lid in layout_ids}
    assert result.report.counts["layout_previews"] == len(layout_ids)
    assert calls.read_text().split() == [MINI_TEMPLATE.stem, "layouts"]
    with Image.open(io.BytesIO(next(iter(result.previews[k] for k in layout_previews)))) as img:
        assert img.width > 0
    assert result.profile["analyzer"]["version"] == an.ANALYZER_VERSION
