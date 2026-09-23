"""CLI `read-charts`: диаграммы-картинки из PPTX или файлов картинок → чтение с оверлеями.

    uv run -m presentation_designer.cli read-charts \\
        "data/organizers/Шаблон презентации VK Education.pptx" --slides 45-50 \\
        --out runs/charts/vkedu/
    … image.png other.png            картинки напрямую
    … --no-model                     только предфильтр «похоже на диаграмму»
    … --structure runs/charts/vkedu/ взять устройство из сохранённых <ключ>.structure.json
                                     (правка руками и перемер без модели)

Выходы: `<out>/readings.json` (для каждой картинки: слайд, sha256, статус, причина, чтение,
устройство), `<out>/<ключ>.overlay.png` (найденная шкала, вершины, точки, сегменты; точки,
восстановленные по соседям или ребру, — оранжевые со звёздочкой) и `<ключ>.structure.json`.
Сводка — в stdout. Модель вызывается только для картинок, прошедших предфильтр.

CLI `rebuild-charts`: те же чтения — и диаграммы-картинки заменяются нативными диаграммами,
как при «Открыть как презентацию»:

    uv run -m presentation_designer.cli rebuild-charts deck.pptx --out runs/charts/deck.pptx
    … --structure runs/charts/vkedu/  без модели, по сохранённым устройствам
"""

from __future__ import annotations

import argparse
import json
import pathlib
import sys
import time
from typing import Any

from pptx import Presentation

from presentation_designer.layout.chart_images import (
    deck_pictures,
    swap_pictures,
    swap_summary,
    unread_charts,
)
from presentation_designer.parsing.raster_charts.model import ChartStructure
from presentation_designer.parsing.raster_charts.overlay import draw_overlay
from presentation_designer.parsing.raster_charts.pixels import chart_likeness, load
from presentation_designer.parsing.raster_charts.rebuild import (
    Outcome,
    image_sha,
    read_chart_images,
    read_with_structure,
)
from presentation_designer.shared.settings import get_settings

IMAGE_SUFFIXES = (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp")


def build_parser(parser: argparse.ArgumentParser | None = None) -> argparse.ArgumentParser:
    parser = parser or argparse.ArgumentParser(prog="presentation-designer read-charts")
    parser.add_argument("inputs", nargs="+", help="PPTX или файлы картинок")
    parser.add_argument("--slides", default=None, help="номера слайдов PPTX: 45-50 или 3,7")
    parser.add_argument("--out", default="runs/charts", help="каталог выходов")
    parser.add_argument("--no-model", action="store_true", help="только предфильтр")
    parser.add_argument("--structure", default=None, help="каталог с <ключ>.structure.json")
    parser.add_argument("--budget", type=float, default=240.0, help="секунд на все вызовы модели")
    parser.add_argument("--max-images", type=int, default=40, help="не больше вызовов модели")
    return parser


def _slides(spec: str | None) -> set[int] | None:
    if not spec:
        return None
    out: set[int] = set()
    for part in spec.split(","):
        a, _, b = part.partition("-")
        out.update(range(int(a), int(b or a) + 1))
    return out


def collect(inputs: list[str], slides: set[int] | None) -> list[dict[str, Any]]:
    """Картинки входов: для PPTX — рисунки слайдов (без макетов), повторы по sha256 — один раз."""
    items: list[dict[str, Any]] = []
    seen: set[str] = set()
    for raw in inputs:
        path = pathlib.Path(raw)
        if path.suffix.lower() == ".pptx":
            prs = Presentation(str(path))
            for number, slide in enumerate(prs.slides, start=1):
                if slides and number not in slides:
                    continue
                for shape in slide.shapes:
                    if shape.shape_type != 13:  # MSO_SHAPE_TYPE.PICTURE
                        continue
                    data = shape.image.blob
                    sha = image_sha(data)
                    if sha in seen:
                        continue
                    seen.add(sha)
                    items.append(
                        {
                            "key": f"s{number:02d}-{sha[:8]}",
                            "slide": number,
                            "name": shape.name,
                            "sha256": sha,
                            "data": data,
                        }
                    )
        elif path.suffix.lower() in IMAGE_SUFFIXES:
            data = path.read_bytes()
            sha = image_sha(data)
            if sha not in seen:
                seen.add(sha)
                items.append(
                    {
                        "key": path.stem,
                        "slide": None,
                        "name": path.name,
                        "sha256": sha,
                        "data": data,
                    }
                )
        else:
            print(f"пропущен вход неизвестного вида: {raw}", file=sys.stderr)
    return items


def _client() -> Any:
    from presentation_designer.llm import build_client

    client = build_client(get_settings())
    if not client.target("vlm").provider.configured():
        raise RuntimeError("провайдер моделей не настроен (PD_QWEN_BASE_URL, PD_QWEN_API_KEY)")
    return client


def _outcomes(items: list[dict[str, Any]], args: argparse.Namespace) -> dict[str, Outcome] | None:
    """Исходы чтения по sha256: только предфильтр, по сохранённым устройствам или моделью.
    None — клиент моделей не создан."""
    outcomes: dict[str, Outcome] = {}
    if getattr(args, "no_model", False):
        for it in items:
            likeness = chart_likeness(load(it["data"]))
            outcomes[it["sha256"]] = Outcome("skipped", likeness.reason or "похоже на диаграмму")
    elif args.structure:
        src = pathlib.Path(args.structure)
        for it in items:
            file = src / f"{it['key']}.structure.json"
            if not file.is_file():
                outcomes[it["sha256"]] = Outcome("skipped", "нет сохранённого устройства")
                continue
            st = ChartStructure.model_validate_json(file.read_text(encoding="utf-8"))
            outcomes[it["sha256"]] = read_with_structure(load(it["data"]), st)
    else:
        try:
            client = _client()
        except Exception as e:
            print(f"клиент моделей не создан: {e}", file=sys.stderr)
            return None
        outcomes = read_chart_images(
            {it["sha256"]: it["data"] for it in items},
            client,
            budget_s=args.budget,
            max_images=args.max_images,
        )
    return outcomes


def run(args: argparse.Namespace) -> int:
    items = collect(args.inputs, _slides(args.slides))
    if not items:
        print("картинок не найдено", file=sys.stderr)
        return 1
    out = pathlib.Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    started = time.perf_counter()
    found = _outcomes(items, args)
    if found is None:
        return 2
    outcomes = found
    rows = []
    for it in items:
        o = outcomes[it["sha256"]]
        if o.structure is not None:
            (out / f"{it['key']}.structure.json").write_text(
                o.structure.model_dump_json(indent=1), encoding="utf-8"
            )
        if o.measured is not None:
            draw_overlay(load(it["data"]), o.measured).save(out / f"{it['key']}.overlay.png")
        rows.append({k: v for k, v in it.items() if k != "data"} | o.as_dict())
        print(_summary(it, o))
    (out / "readings.json").write_text(
        json.dumps(rows, ensure_ascii=False, indent=1), encoding="utf-8"
    )
    counts: dict[str, int] = {}
    for o in outcomes.values():
        counts[o.status] = counts.get(o.status, 0) + 1
    print(f"итог: {counts}, {time.perf_counter() - started:.1f} с, выходы в {out}")
    return 0


def build_rebuild_parser(parser: argparse.ArgumentParser | None = None) -> argparse.ArgumentParser:
    parser = parser or argparse.ArgumentParser(prog="presentation-designer rebuild-charts")
    parser.add_argument("pptx", help="готовая презентация")
    parser.add_argument("--out", required=True, help="PPTX с нативными диаграммами")
    parser.add_argument("--structure", default=None, help="каталог с <ключ>.structure.json")
    parser.add_argument("--budget", type=float, default=240.0, help="секунд на все вызовы модели")
    parser.add_argument("--max-images", type=int, default=40, help="не больше вызовов модели")
    return parser


def rebuild(args: argparse.Namespace) -> int:
    """Диаграммы-картинки файла → нативные диаграммы, сводка — в stdout."""
    source = pathlib.Path(args.pptx)
    items = collect([str(source)], None)
    outcomes = _outcomes(items, args)
    if outcomes is None:
        return 2
    prs = Presentation(str(source))
    _, where = deck_pictures(prs)
    readings = {sha: o.reading for sha, o in outcomes.items() if o.reading is not None}
    swaps = swap_pictures(prs, readings)
    out = pathlib.Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    for swap in swaps:
        print(f"слайд {swap.slide:>3}  {swap.status:8} {swap.name} {swap.reason}".rstrip())
    print(swap_summary(swaps, unread_charts(outcomes, where)) or "диаграмм-картинок не найдено")
    print(f"сохранено: {out}")
    return 0


def _summary(it: dict[str, Any], o: Outcome) -> str:
    head = f"{it['key']:>16}  {o.status:9}"
    if o.reading is None:
        return f"{head} {o.reason}"
    r = o.reading
    counts = r.counts()
    parts = [f"{r.kind}", f"{len(r.series)} ряд.", f"{len(r.categories)} кат."]
    parts.append(", ".join(f"{k} {v}" for k, v in counts.items() if v))
    series = "; ".join(
        f"{s.name}: "
        + " ".join(
            f"{p.value:g}{'' if p.basis in ('label', 'measured') else '*'}" for p in s.points[:8]
        )
        + (" …" if len(s.points) > 8 else "")
        for s in r.series
    )
    return f"{head} {' · '.join(parts)}\n{'':>28}{series}"
