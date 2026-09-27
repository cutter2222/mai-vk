"""Local ONLYOFFICE integration. JWT authenticates service traffic, not application users."""

from __future__ import annotations

import asyncio
import io
import json
import logging
import tempfile
import time
import uuid
from pathlib import Path
from typing import Any, Literal
from urllib.parse import quote, urlencode

import httpx
from fastapi import APIRouter, Request
from fastapi.responses import FileResponse, Response
from pydantic import BaseModel, Field, model_validator

from presentation_designer.api.deps import Orch
from presentation_designer.api.errors import ApiError
from presentation_designer.export.office_download import export_revision
from presentation_designer.export.office_ooxml import validate_saved_pptx
from presentation_designer.export.office_preview import cache_path, preview_page, preview_revision
from presentation_designer.export.office_source import saved_url, source_path
from presentation_designer.export.pdf import ConversionError, RendererUnavailableError
from presentation_designer.generation import (
    office_edit,
    office_image,
    office_logo,
    office_object_edit,
    office_objects,
    office_ops,
)
from presentation_designer.generation.office_objects import (
    LiveTarget,
    ObjectTarget,
    objects,
    resolve_live,
)
from presentation_designer.layout.merge import MergeError, replace_slide
from presentation_designer.llm.types import LlmError
from presentation_designer.parsing.template.embedded_fonts import prepare_fonts
from presentation_designer.parsing.template.package import PackageError
from presentation_designer.pipeline.artifacts import content_type_for
from presentation_designer.pipeline.office import OfficeStore, sign, verify
from presentation_designer.pipeline.results import build_generation_result
from presentation_designer.pipeline.snapshots import office_snapshot
from presentation_designer.pipeline.state import NotFound
from presentation_designer.shared.text import plural

log = logging.getLogger(__name__)
router = APIRouter(prefix="/office", tags=["onlyoffice"])


class OpenRequest(BaseModel):
    job_id: str = Field(max_length=100)
    artifact: str = Field(max_length=300)


class LogoRequest(BaseModel):
    """Знак шаблона на всех слайдах: заменить картинкой из файлов проекта или убрать."""

    action: Literal["replace", "remove"]
    file_id: str | None = Field(default=None, max_length=100)


class ImagePlacement(BaseModel):
    """Картинка из сообщения — на слайд: файл проекта и номер слайда с единицы; место —
    словами в `instruction` или выделенный объект (`live_target`)."""

    file_id: str = Field(min_length=1, max_length=100)
    slide: int = Field(ge=1)


class ChartPlacement(BaseModel):
    """Новая диаграмма на слайд (этап 40): из приложенного xlsx/csv (`file`), по картинке
    графика (`image`) или «сделай редактируемой» для диаграмм из фигур и картинок слайда."""

    source: Literal["file", "image", "editable"]
    slide: int = Field(ge=1)
    file_id: str | None = Field(default=None, max_length=100)

    @model_validator(mode="after")
    def validate_file(self) -> ChartPlacement:
        if (self.source == "editable") != (self.file_id is None):
            raise ValueError("Файл нужен для диаграммы из файла или картинки, и только для неё")
        return self


class EditRequest(BaseModel):
    revision: int = Field(ge=0)
    instruction: str = Field(min_length=1, max_length=4000)
    target: ObjectTarget | None = None
    targets: list[ObjectTarget] | None = Field(default=None, min_length=1, max_length=100)
    # Объект, выделенный в живом редакторе: сервер находит его в сохранённой копии сам.
    live_target: LiveTarget | None = None
    logo: LogoRequest | None = None
    image: ImagePlacement | None = None
    # Таблица из приложенного xlsx или csv — на слайд, в место, названное словами.
    table: ImagePlacement | None = None
    chart: ChartPlacement | None = None
    # Текстовая правка только на этих слайдах (с единицы); без объектов.
    slides: list[int] | None = Field(default=None, min_length=1, max_length=200)

    @model_validator(mode="after")
    def validate_targets(self) -> EditRequest:
        if self.logo is not None and (self.target is not None or self.targets is not None):
            raise ValueError("Логотип меняется на всех слайдах: объекты не передаются")
        if self.image is not None and (
            self.logo is not None or self.target is not None or self.targets is not None
        ):
            raise ValueError("Картинка ставится по словам или в выделенный объект")
        if self.table is not None and (
            self.logo is not None
            or self.image is not None
            or self.target is not None
            or self.targets is not None
        ):
            raise ValueError("Таблица из файла ставится по словам")
        if self.chart is not None and any(
            v is not None for v in (self.logo, self.image, self.table, self.target, self.targets)
        ):
            raise ValueError("Диаграмма ставится по словам, без других объектов")
        if self.live_target is not None and (
            self.logo is not None or self.target is not None or self.targets is not None
        ):
            raise ValueError("Выделение живого редактора передаётся без других объектов")
        if self.targets is not None:
            if self.target is not None:
                raise ValueError("Передайте target или targets, не оба поля")
            keys = {(target.slide, target.shape_id) for target in self.targets}
            if len(keys) != len(self.targets):
                raise ValueError("Объекты не должны повторяться")
        return self


class ImageRequest(BaseModel):
    """Картинка для текущего слайда живого редактора: файл проекта или ресурс шаблона."""

    project_id: str | None = Field(default=None, max_length=100)
    file_id: str | None = Field(default=None, max_length=100)
    template_id: str | None = Field(default=None, max_length=100)
    asset_id: str | None = Field(default=None, max_length=200)

    @model_validator(mode="after")
    def validate_source(self) -> ImageRequest:
        from_file = bool(self.project_id and self.file_id)
        from_template = bool(self.template_id and self.asset_id)
        if from_file == from_template:
            raise ValueError("Передайте project_id и file_id или template_id и asset_id")
        return self


# Картинки, которые Document Server вставляет по ссылке; тип — по расширению проверенного файла.
INSERTABLE = {"png": "png", "jpg": "jpg", "jpeg": "jpg", "gif": "gif", "bmp": "bmp"}


@router.post("/documents/{document_id}/images")
def insert_image(document_id: str, body: ImageRequest, orch: Orch) -> dict[str, Any]:
    """Подписанная команда `docEditor.insertImage`: Document Server сам скачивает картинку
    по внутреннему адресу API (`storage_url`) и ставит её на текущий слайд. Сервер только
    проверяет источник и подписывает ссылку, документ не меняет."""
    store(orch).get(document_id)
    cfg = orch.settings.onlyoffice
    base = cfg.storage_url.rstrip("/") + "/api"
    if body.project_id and body.file_id:
        try:
            row = orch.state.get_file(body.file_id, body.project_id)
        except NotFound as exc:
            raise ApiError(404, "file_not_found", "Файл не найден в проекте") from exc
        ext = _extension(str(row["name"]))
        check = row.get("check") or {}
        if check.get("status") != "ok" or check.get("format") != "image" or ext not in INSERTABLE:
            raise ApiError(422, "file_not_image", "На слайд вставляются картинки PNG или JPEG")
        if not orch.files.path_for(str(row["sha256"])).is_file():
            raise ApiError(404, "file_missing", "Байты файла удалены из хранилища")
        url = f"{base}/projects/{quote(body.project_id)}/files/{quote(body.file_id)}/content"
    else:
        template_id, asset_id = str(body.template_id), str(body.asset_id)
        try:
            template = orch.state.get_template(template_id)
        except NotFound as exc:
            raise ApiError(404, "template_not_found", "Шаблон не найден") from exc
        asset = next(
            (
                a
                for a in (template.get("profile") or {}).get("assets") or []
                if a.get("asset_id") == asset_id
            ),
            None,
        )
        if asset is None:
            raise ApiError(404, "asset_not_found", "Ресурс не найден в профиле шаблона")
        ext = _extension(str(asset.get("media_path") or ""))
        if ext not in INSERTABLE:
            raise ApiError(422, "file_not_image", "Этот ресурс шаблона нельзя вставить картинкой")
        try:
            orch.extract_template_media(template_id, asset)
        except FileNotFoundError as exc:
            raise ApiError(404, "asset_not_found", "Файл ресурса отсутствует") from exc
        url = f"{base}/templates/{quote(template_id)}/media/{quote(asset_id)}"
    command: dict[str, Any] = {"c": "add", "images": [{"fileType": INSERTABLE[ext], "url": url}]}
    return {**command, "token": sign(command, cfg.jwt_secret)}


def _extension(name: str) -> str:
    return name.rsplit(".", 1)[-1].lower() if "." in name else ""


@router.get("/documents/{document_id}/objects/{revision}")
def document_objects(document_id: str, revision: int, orch: Orch) -> dict[str, Any]:
    data = store(orch).read(document_id, revision)
    return {"revision": revision, "objects": [obj.model_dump() for obj in objects(data)]}


@router.get("/documents/{document_id}/snapshot/{revision}")
def document_snapshot(document_id: str, revision: int, orch: Orch) -> dict[str, Any]:
    """Снимок колоды копии (`deck_snapshot`): слайды, объекты с адресами, оглавление."""
    office = store(orch)
    try:
        return office_snapshot(orch, office, document_id, revision)
    except (PackageError, ValueError) as exc:
        raise ApiError(422, "office_snapshot_failed", "Снимок не построен: " + str(exc)) from exc


@router.post("/documents/{document_id}/edit")
async def edit(document_id: str, body: EditRequest, orch: Orch) -> dict[str, Any]:
    office = store(orch)
    token = office.begin_edit(document_id, body.revision)
    try:
        original = office.read(document_id, body.revision)
        if body.logo is not None:
            updated, explanation = _apply_logo(orch, office.get(document_id), original, body.logo)
        elif body.image is not None:
            updated, explanation = await _place_image(orch, original, body)
        elif body.table is not None:
            updated, explanation = await _place_table(orch, office.get(document_id), original, body)
        elif body.chart is not None:
            updated, explanation = await _place_chart(orch, office.get(document_id), original, body)
        elif body.targets is not None:
            multi_plan = await office_object_edit.propose_many(
                original, body.instruction, orch.settings, body.targets
            )
            updated = office_object_edit.patch_objects(original, body.targets, multi_plan)
            explanation = multi_plan.explanation
        elif body.target is not None or body.live_target is not None:
            target = body.target or resolve_live(original, body.live_target)  # type: ignore[arg-type]
            tokens = _tokens(orch, office.get(document_id), original)
            object_plan = await office_object_edit.propose(
                original, body.instruction, orch.settings, target, tokens
            )
            updated, notes = office_object_edit.apply_plan(original, target, object_plan, tokens)
            explanation = " ".join(
                [object_plan.explanation.strip(), *(n[:1].upper() + n[1:] + "." for n in notes)]
            ).strip()
        else:
            plan = await office_edit.propose(original, body.instruction, orch.settings, body.slides)
            updated = office_edit.patch_pptx(original, plan)
            explanation = plan.explanation
        office.commit_edit(document_id, token, body.revision, updated)
        return {
            "document": office.get(document_id),
            "changed": original != updated,
            "message": explanation,
        }
    except (ValueError, LlmError) as exc:
        raise ApiError(422, "office_edit_failed", "Правка не применена: " + str(exc)) from exc
    finally:
        office.end_edit(document_id, token)


class UndoRequest(BaseModel):
    """Отмена правки копии: `revision` — ревизия, которую дала правка, `to_revision` — до неё."""

    revision: int = Field(ge=1)
    to_revision: int = Field(ge=0)


@router.post("/documents/{document_id}/undo")
def undo(document_id: str, body: UndoRequest, orch: Orch) -> dict[str, Any]:
    """Прежние байты новой ревизией: история копии не переписывается, отмену можно отменить.
    Текущие байты должны совпадать с результатом правки, в том числе после отмены
    более поздних правок. begin_edit повторно сверяет текущую ревизию под блокировкой."""
    office = store(orch)
    if body.to_revision >= body.revision:
        raise ApiError(422, "office_undo_invalid", "Вернуться можно только к прежней ревизии")
    current = office.get(document_id)
    revision = int(current["revision"])
    if revision != body.revision and (
        body.revision > revision
        or office.read(document_id, revision) != office.read(document_id, body.revision)
    ):
        raise ApiError(
            409,
            "office_undo_stale",
            "После этой правки презентацию уже меняли — отменить её отдельно нельзя. "
            "Верните нужное в редакторе (Ctrl+Z) или попросите исправить словами.",
        )
    previous = office.read(document_id, body.to_revision)
    token = office.begin_edit(document_id, revision)
    try:
        office.commit_edit(document_id, token, revision, previous)
        return {"document": office.get(document_id), "changed": True, "message": "Правка отменена."}
    finally:
        office.end_edit(document_id, token)


class ApplySlideRequest(BaseModel):
    """Слайд новой ревизии варианта — на место слайда офисной копии."""

    revision: int = Field(ge=0)
    job_id: str = Field(min_length=1, max_length=64, pattern=r"^[A-Za-z0-9_-]+$")
    variant_id: str = Field(min_length=1, max_length=64, pattern=r"^[A-Za-z0-9_-]+$")
    artifact_revision: int = Field(ge=1)
    # Номера слайдов с единицы: в офисной копии и в ревизии варианта.
    slide: int = Field(ge=1)
    source_slide: int = Field(ge=1)


@router.post("/documents/{document_id}/apply-slide")
def apply_slide(document_id: str, body: ApplySlideRequest, orch: Orch) -> dict[str, Any]:
    """Правка слайда из чата пересобирает слайд в новой ревизии варианта; человек тем временем
    правит офисную копию руками. Сюда переносится только этот слайд — ручные правки остальных
    слайдов остаются, правки самого слайда заменяются пересобранным."""
    office = store(orch)
    document = office.get(document_id)
    if not str(document.get("source") or "").startswith(f"{body.job_id}/{body.variant_id}/"):
        raise ApiError(409, "office_source_mismatch", "Офисная копия открыта из другой сборки")
    artifact = f"{body.variant_id}/r{body.artifact_revision}/deck.pptx"
    result = build_generation_result(orch.state, body.job_id)
    if artifact not in result["artifacts_manifest"]:
        raise ApiError(404, "artifact_not_found", "PPTX новой ревизии не найден в манифесте")
    source_path = orch.artifacts.resolve(body.job_id, artifact)
    token = office.begin_edit(document_id, body.revision)
    try:
        original = office.read(document_id, body.revision)
        updated = _replace_slide(original, body.slide, source_path.read_bytes(), body.source_slide)
        office.commit_edit(document_id, token, body.revision, updated)
        return {
            "document": office.get(document_id),
            "changed": original != updated,
            "message": f"Слайд {body.slide} обновлён.",
        }
    except (MergeError, ValueError) as exc:
        raise ApiError(422, "office_apply_failed", "Слайд не перенесён: " + str(exc)) from exc
    finally:
        office.end_edit(document_id, token)


def _replace_slide(original: bytes, slide: int, source: bytes, source_slide: int) -> bytes:
    from pptx import Presentation

    base = Presentation(io.BytesIO(original))
    replace_slide(base, slide - 1, Presentation(io.BytesIO(source)), source_slide - 1)
    out = io.BytesIO()
    base.save(out)
    return out.getvalue()


def _document_template(orch: Orch, source: str) -> str | None:
    """Шаблон офисной копии по её источнику: копия шаблона проекта
    (`project/<id>/template/<template_id>/<sha>`) или артефакт задания (`<job_id>/<путь>`)."""
    parts = source.split("/")
    if len(parts) >= 4 and parts[0] == "project" and parts[2] == "template":
        return parts[3]
    try:
        return str(orch.state.get_generation(parts[0])["template_id"])
    except (NotFound, KeyError):
        return None


def _apply_logo(
    orch: Orch, document: dict[str, Any], original: bytes, logo: LogoRequest
) -> tuple[bytes, str]:
    """Знак шаблона заменяется или убирается на образцах, макетах и слайдах сразу; картинки
    знака берутся из профиля шаблона, в документе находятся по содержимому."""
    template_id = _document_template(orch, str(document.get("source") or ""))
    try:
        template = orch.state.get_template(template_id) if template_id else None
    except NotFound:
        template = None
    if template is None:
        raise ValueError("не найден шаблон этой презентации")
    if template.get("status") != "succeeded" or not template.get("profile"):
        raise ValueError("шаблон ещё разбирается — повторите через минуту")
    hashes = office_logo.logo_hashes(template["profile"])
    if not hashes:
        raise ValueError("в шаблоне нет логотипа на макетах")
    image: bytes | None = None
    if logo.action == "replace":
        if not logo.file_id:
            raise ValueError("прикрепите картинку нового логотипа (PNG или JPG)")
        try:
            row = orch.state.get_file(logo.file_id)
        except NotFound as exc:
            raise ValueError("картинка логотипа не найдена в файлах проекта") from exc
        image = orch.files.path_for(str(row["sha256"])).read_bytes()
    updated, count = office_logo.replace_logo(original, hashes, image)
    if not count:
        raise ValueError("логотипа шаблона в этой презентации уже нет")
    done = "заменён" if image is not None else "убран"
    places = f"{count} {plural(count, 'место', 'места', 'мест')}"
    return updated, f"Логотип шаблона {done} на всех слайдах ({places})."


async def _place_image(orch: Orch, original: bytes, body: EditRequest) -> tuple[bytes, str]:
    """Картинка из сообщения — в выделенный объект (картинку или пустую рамку) или в место,
    названное словами: его модель находит по снимку слайда (блок бывает нарисован фоном)."""
    image = body.image
    assert image is not None
    try:
        row = orch.state.get_file(image.file_id)
    except NotFound as exc:
        raise ValueError("картинка не найдена в файлах проекта") from exc
    blob = orch.files.path_for(str(row["sha256"])).read_bytes()
    target = None
    if body.live_target is not None:
        try:
            picked = office_objects.selected(original, resolve_live(original, body.live_target))
        except ValueError:  # выделение не нашлось в сохранённой копии — место по словам
            picked = None
        # Выделен текст — картинка его не закрывает: место тоже ищется по словам.
        if picked is not None and (picked.kind == "pic" or not picked.runs):
            target = picked
    if target is not None:
        placement = office_image.Placement(explanation="Поставил картинку в выделенный объект.")
        return office_image.place_image(original, target.slide, blob, placement, target=target), (
            placement.explanation
        )
    snapshot = None
    try:
        root, manifest = await asyncio.to_thread(preview_revision, original, orch.settings)
        page = await asyncio.to_thread(preview_page, root, manifest["slides"][image.slide - 1])
        snapshot = page.read_bytes()
    except (ConversionError, RendererUnavailableError, IndexError, OSError):
        log.warning("снимок слайда %d для картинки не получен", image.slide, exc_info=True)
    placement = await office_image.propose(
        original, body.instruction, orch.settings, image.slide, snapshot
    )
    updated = office_image.place_image(original, image.slide, blob, placement, snapshot=snapshot)
    return updated, placement.explanation


async def _snapshot_png(orch: Orch, original: bytes, slide: int) -> bytes | None:
    try:
        root, manifest = await asyncio.to_thread(preview_revision, original, orch.settings)
        page = await asyncio.to_thread(preview_page, root, manifest["slides"][slide - 1])
        return page.read_bytes()
    except (ConversionError, RendererUnavailableError, IndexError, OSError):
        log.warning("снимок слайда %d не получен", slide, exc_info=True)
        return None


async def _place_table(
    orch: Orch, document: dict[str, Any], original: bytes, body: EditRequest
) -> tuple[bytes, str]:
    """Таблица из xlsx или csv — в место, названное словами, в стиле шаблона."""
    from presentation_designer.generation import office_chart, office_table

    table = body.table
    assert table is not None
    try:
        row = orch.state.get_file(table.file_id)
    except NotFound as exc:
        raise ValueError("файл с таблицей не найден в проекте") from exc
    raw = orch.files.path_for(str(row["sha256"])).read_bytes()
    header, rows, truncated = office_table.read_table(raw, str(row.get("name") or ""))
    style = office_chart.table_style(_tokens(orch, document, original))
    snapshot = await _snapshot_png(orch, original, table.slide)
    placement = await office_image.propose(
        original, body.instruction, orch.settings, table.slide, snapshot, what="таблицу"
    )
    updated = office_table.place_table(
        original, table.slide, header, rows, placement, style, snapshot=snapshot
    )
    note = f"Поставил таблицу: {len(rows)} строк, {len(header)} столбцов."
    if truncated:
        note += " В файле больше данных — взял первые строки и столбцы."
    return updated, " ".join(_sentence(t) for t in (placement.explanation, note) if t.strip())


async def _place_chart(
    orch: Orch, document: dict[str, Any], original: bytes, body: EditRequest
) -> tuple[bytes, str]:
    """Диаграмма из xlsx/csv или по картинке графика — в место, названное словами, в стиле
    шаблона; «сделай редактируемой» — диаграммы из фигур и картинок слайда на месте."""
    from presentation_designer.generation import office_chart, office_table
    from presentation_designer.layout.chart_images import describe

    chart = body.chart
    assert chart is not None
    tokens = _tokens(orch, document, original)
    if chart.source == "editable":
        updated, notes = await office_chart.make_editable_all(original, chart.slide, orch.settings)
        return updated, " ".join(n[:1].upper() + n[1:] + "." for n in notes)
    try:
        row = orch.state.get_file(str(chart.file_id))
    except NotFound as exc:
        raise ValueError("файл для диаграммы не найден в проекте") from exc
    raw = orch.files.path_for(str(row["sha256"])).read_bytes()
    kind = office_chart.kind_in(body.instruction)
    note = ""
    if chart.source == "file":
        header, rows, truncated = office_table.read_table(raw, str(row.get("name") or ""))
        spec = office_chart.spec_from_table(header, rows, kind)
        if truncated:
            note = "В файле больше данных — взял первые строки и столбцы."
    else:
        reading = await office_chart.read_picture(raw, orch.settings)
        spec = office_chart.spec_from_reading(reading)
        if kind and kind != spec.chart_type:
            spec.chart_type = kind
            if office_chart.FAMILY[kind] == "pie":
                spec.series = spec.series[:1]
                spec.show_legend = True
        note = describe(reading)
    snapshot = await _snapshot_png(orch, original, chart.slide)
    placement = await office_image.propose(
        original, body.instruction, orch.settings, chart.slide, snapshot, what="диаграмму"
    )
    area = office_image.area_for(original, chart.slide, placement, snapshot, what="диаграмму")
    dark = snapshot is not None and office_chart.dark_area(snapshot, area)
    updated = office_chart.place_chart(original, chart.slide, spec, area, tokens, dark=dark)
    what = f"Поставил диаграмму ({office_chart.KIND_NAMES[spec.chart_type]})."
    return updated, " ".join(_sentence(t) for t in (placement.explanation, what, note) if t.strip())


def _sentence(text: str) -> str:
    """Фраза для ответа: с заглавной буквы и с точкой в конце."""
    text = text.strip()
    if text and text[-1] not in ".!?…":
        text += "."
    return text[:1].upper() + text[1:]


def _tokens(orch: Orch, document: dict[str, Any], original: bytes) -> office_ops.Tokens:
    """Стиль правок на месте: токены профиля шаблона копии, без профиля — тема файла."""
    template_id = _document_template(orch, str(document.get("source") or ""))
    try:
        template = orch.state.get_template(template_id) if template_id else None
    except NotFound:
        template = None
    return office_ops.Tokens.from_profile(
        (template or {}).get("profile")
    ) or office_ops.Tokens.from_pptx(original)


def store(orch: Orch) -> OfficeStore:
    cfg = orch.settings.onlyoffice
    if not cfg.enabled or len(cfg.jwt_secret) < 32:
        raise ApiError(503, "office_disabled", "ONLYOFFICE не настроен")
    return OfficeStore(orch.settings.paths.data_dir)


@router.get("/capabilities")
def capabilities(orch: Orch) -> dict[str, Any]:
    """Включён ли редактор и адрес его SDK: интерфейс прогревает редактор заранее, пока
    презентация ещё готовится, а не после того, как файл появился."""
    cfg = orch.settings.onlyoffice
    enabled = cfg.enabled and len(cfg.jwt_secret) >= 32
    if not enabled:
        return {"enabled": False}
    return {"enabled": True, "script_url": _script_url(cfg.public_url)}


def _script_url(public_url: str) -> str:
    return f"{public_url.rstrip('/')}/web-apps/apps/api/documents/api.js?rendering=pd16v2"


@router.post("/documents")
def create(body: OpenRequest, orch: Orch) -> dict[str, Any]:
    office = store(orch)
    result = build_generation_result(orch.state, body.job_id)
    if body.artifact not in result["artifacts_manifest"] or not body.artifact.endswith(".pptx"):
        raise ApiError(404, "artifact_not_found", "PPTX не найден в манифесте")
    path = orch.artifacts.resolve(body.job_id, body.artifact)
    if not path.is_file():
        raise ApiError(404, "artifact_not_found", "PPTX отсутствует")
    if path.stat().st_size > orch.files.max_bytes:
        raise ApiError(413, "office_file_too_large", "PPTX превышает лимит загрузки")
    return office.create(
        f"{body.job_id}/{body.artifact}",
        f"{body.job_id}-{body.artifact.replace('/', '-')}",
        path.read_bytes(),
    )


def template_source(template_id: str, orch: Orch) -> tuple[dict[str, Any], Path]:
    try:
        template = orch.state.get_template(template_id)
    except NotFound as exc:
        raise ApiError(404, "template_not_found", "Шаблон не найден") from exc
    path = orch.files.path_for(template["sha256"])
    if not path.is_file():
        raise ApiError(404, "template_source_missing", "Исходный PPTX отсутствует")
    return template, path


class TemplateCopyRequest(BaseModel):
    template_id: str = Field(max_length=100)


@router.post("/projects/{project_id}/template")
def project_template_copy(project_id: str, body: TemplateCopyRequest, orch: Orch) -> dict[str, Any]:
    office = store(orch)
    try:
        project = orch.state.get_project(project_id)
    except NotFound as exc:
        raise ApiError(404, "project_not_found", "Проект не найден") from exc
    if project.get("template_id") != body.template_id:
        raise ApiError(409, "template_changed", "Шаблон проекта изменился. Обновите страницу")
    template, path = template_source(body.template_id, orch)
    if path.stat().st_size > orch.files.max_bytes:
        raise ApiError(413, "office_file_too_large", "PPTX превышает лимит загрузки")
    # A project-local copy: never edit the library template or another project's copy.
    source = f"project/{project_id}/template/{body.template_id}/{template['sha256']}"
    existing = office.find(source)
    if existing is not None:
        return existing
    # Знак шаблона — с макетов на слайды, как в собранных вариантах: его двигают на слайде.
    try:
        promoted, _ = office_logo.promote_file(path)
    except Exception:
        log.exception("знак шаблона %s не перенесён на слайды копии", body.template_id)
        promoted = None
    return office.create(source, template["name"], promoted or path.read_bytes())


@router.post("/templates/{template_id}/config")
def template_config(template_id: str, orch: Orch) -> dict[str, Any]:
    # Просмотр не создаёт OfficeStore/ревизий и не меняет профиль анализатора.
    cfg = orch.settings.onlyoffice
    if not cfg.enabled or len(cfg.jwt_secret) < 32:
        raise ApiError(503, "office_disabled", "ONLYOFFICE не настроен")
    template, path = template_source(template_id, orch)
    font_report = prepare_fonts(path, orch.settings.paths.data_dir)
    token = sign(
        {
            "scope": "office-template",
            "template_id": template_id,
            "sha256": template["sha256"],
            "exp": int(time.time()) + 86400,
        },
        cfg.jwt_secret,
    )
    editor = {
        "documentType": "slide",
        "type": "desktop",
        "width": "100%",
        "height": "100%",
        "document": {
            "fileType": "pptx",
            "key": f"template-{template['sha256']}",
            "title": template["name"],
            "url": (
                f"{cfg.storage_url.rstrip('/')}/api/office/templates/{template_id}/source"
                f"?{urlencode({'token': token})}"
            ),
            "permissions": {
                "edit": False,
                "download": False,
                "print": True,
                "comment": False,
                "chat": False,
            },
        },
        "editorConfig": {
            "mode": "view",
            "lang": "ru",
            "user": {"id": uuid.uuid4().hex, "name": "Просмотр шаблона"},
            "customization": {
                "anonymous": {"request": False},
                "compactHeader": True,
                "toolbarHideFileName": True,
                "uiTheme": "theme-white",
                "compactToolbar": True,
                "hideRightMenu": True,
                "hideRulers": True,
                "hideNotes": True,
                "comments": False,
                "features": {"featuresTips": False},
                "zoom": -1,
                "plugins": False,
                "macros": False,
                "help": False,
                "feedback": {"visible": False},
                "suggestFeature": False,
            },
        },
    }
    return {
        "font_report": font_report,
        "script_url": _script_url(cfg.public_url),
        "config": {**editor, "token": sign(editor, cfg.jwt_secret)},
    }


@router.get("/templates/{template_id}/source")
def signed_template_source(template_id: str, token: str, orch: Orch) -> FileResponse:
    cfg = orch.settings.onlyoffice
    if not cfg.enabled or len(cfg.jwt_secret) < 32:
        raise ApiError(503, "office_disabled", "ONLYOFFICE не настроен")
    claims = verify(token, orch.settings.onlyoffice.jwt_secret)
    if (
        claims.get("scope") != "office-template"
        or claims.get("template_id") != template_id
        or "exp" not in claims
    ):
        raise ApiError(403, "office_token_invalid", "Недействительная ссылка шаблона")
    template, path = template_source(template_id, orch)
    if claims.get("sha256") != template["sha256"]:
        raise ApiError(403, "office_token_invalid", "Версия шаблона изменилась")
    return FileResponse(
        path, media_type=content_type_for("deck.pptx"), headers={"Cache-Control": "no-store"}
    )


@router.get("/documents/{document_id}")
def detail(document_id: str, orch: Orch) -> dict[str, Any]:
    return store(orch).get(document_id)


@router.post("/documents/{document_id}/config")
def config(document_id: str, orch: Orch) -> dict[str, Any]:
    office = store(orch)
    doc = office.open(document_id)
    font_report = prepare_fonts(
        office.read(document_id, doc["seed_revision"]),
        orch.settings.paths.data_dir,
    )
    cfg = orch.settings.onlyoffice
    revision = doc["seed_revision"]
    token = sign(
        {
            "document_id": document_id,
            "revision": revision,
            "scope": "office-file",
            "exp": int(time.time()) + 7 * 86400,
        },
        cfg.jwt_secret,
    )
    base = f"{cfg.storage_url.rstrip('/')}/api/office/documents/{document_id}"
    editor = {
        "documentType": "slide",
        "type": "desktop",
        "width": "100%",
        "height": "100%",
        "document": {
            "fileType": "pptx",
            "key": doc["active_key"],
            "title": doc["title"],
            "url": f"{base}/source/{revision}?{urlencode({'token': token})}",
            "permissions": {
                "edit": True,
                "download": False,
                "print": True,
                "chat": False,
                "comment": False,
            },
        },
        "editorConfig": {
            "mode": "edit",
            "lang": "ru",
            "callbackUrl": f"{base}/callback",
            "user": {"id": uuid.uuid4().hex, "name": "Локальный редактор"},
            "customization": {
                "forcesave": True,
                "autosave": True,
                # The host project already shows the title; do not repeat a raw artifact ID.
                "compactHeader": True,
                "toolbarHideFileName": True,
                # Standard customization only: no paid White Label layout overrides.
                # Users can expand tools again via View; their saved preferences win.
                "uiTheme": "theme-white",
                "compactToolbar": True,
                "hideRightMenu": True,
                "hideRulers": True,
                "hideNotes": True,
                "comments": False,
                "features": {"featuresTips": False},
                "zoom": -1,
                # AI is provided by the host; do not expose a second plugin workflow.
                "plugins": False,
                "macros": False,
                "help": False,
                "feedback": {"visible": False},
                "suggestFeature": False,
            },
        },
    }
    return {
        "font_report": font_report,
        "script_url": _script_url(cfg.public_url),
        "config": {**editor, "token": sign(editor, cfg.jwt_secret)},
    }


@router.get("/documents/{document_id}/preview/{revision}")
def preview(document_id: str, revision: int, orch: Orch) -> dict[str, Any]:
    content = store(orch).read(document_id, revision)
    try:
        _, manifest = preview_revision(content, orch.settings)
    except (ConversionError, RendererUnavailableError, OSError, ValueError) as exc:
        raise ApiError(503, "office_preview_failed", "Не удалось сформировать превью") from exc
    return {"revision": revision, **manifest}


@router.get("/documents/{document_id}/preview/{revision}/{name}")
def preview_image(document_id: str, revision: int, name: str, orch: Orch) -> FileResponse:
    content = store(orch).read(document_id, revision)
    root = cache_path(content, orch.settings)
    manifest = root / "manifest.json"
    try:
        slides = json.loads(manifest.read_text(encoding="utf-8"))["slides"]
    except (OSError, ValueError, KeyError, TypeError) as exc:
        raise ApiError(404, "office_preview_missing", "Превью не найдено") from exc
    if name not in slides or Path(name).name != name:
        raise ApiError(404, "office_preview_missing", "Превью не найдено")
    try:
        image = preview_page(root, name)
    except (ConversionError, OSError, ValueError) as exc:
        raise ApiError(503, "office_preview_failed", "Не удалось отрисовать слайд") from exc
    return FileResponse(
        image,
        media_type="image/png",
        headers={"Cache-Control": "private, max-age=31536000, immutable"},
    )


@router.get("/documents/{document_id}/source/{revision}")
def source(document_id: str, revision: int, token: str, orch: Orch) -> Response:
    office = store(orch)
    claims = verify(token, orch.settings.onlyoffice.jwt_secret)
    if (claims.get("scope"), claims.get("document_id"), claims.get("revision")) != (
        "office-file",
        document_id,
        revision,
    ):
        raise ApiError(403, "office_token_invalid", "Подпись не соответствует файлу")
    return Response(office.read(document_id, revision), media_type=content_type_for("deck.pptx"))


@router.get("/documents/{document_id}/download/{revision}")
def download(
    document_id: str, revision: int, orch: Orch, format: Literal["pptx", "pdf", "html"] = "pptx"
) -> Response:
    content = store(orch).read(document_id, revision)
    if format != "pptx":
        try:
            content = export_revision(content, format, orch.settings)
        except (ConversionError, RendererUnavailableError) as exc:
            raise ApiError(
                503, "office_export_failed", "Не удалось экспортировать презентацию"
            ) from exc
    return Response(
        content,
        media_type=content_type_for(f"deck.{format}"),
        headers={
            "Content-Disposition": (
                f'attachment; filename="office-{document_id}-v{revision}.{format}"'
            ),
            "Cache-Control": "no-store",
        },
    )


@router.get("/conversions/{key}")
def conversion_source(key: str, orch: Orch, token: str = "") -> FileResponse:
    cfg = orch.settings.onlyoffice
    if not cfg.enabled or len(cfg.jwt_secret) < 32:
        raise ApiError(503, "office_disabled", "ONLYOFFICE не настроен")
    claims = verify(token, cfg.jwt_secret)
    if (
        claims.get("scope") != "office-conversion"
        or claims.get("key") != key
        or "exp" not in claims
    ):
        raise ApiError(403, "office_token_invalid", "Недействительная ссылка конвертации")
    try:
        path = source_path(orch.settings, key)
    except ValueError as exc:
        raise ApiError(404, "conversion_not_found", "Файл конвертации не найден") from exc
    if not path.is_file():
        raise ApiError(404, "conversion_not_found", "Файл конвертации удалён")
    return FileResponse(
        path, media_type=content_type_for("deck.pptx"), headers={"Cache-Control": "no-store"}
    )


@router.post("/documents/{document_id}/callback")
async def callback(document_id: str, request: Request, orch: Orch) -> dict[str, int]:
    office = store(orch)
    if len(await request.body()) > 131072:
        raise ApiError(413, "office_callback_too_large", "Слишком большой callback")
    try:
        body = await request.json()
    except ValueError as exc:
        raise ApiError(400, "office_callback_invalid", "Некорректный JSON") from exc
    if not isinstance(body, dict):
        raise ApiError(400, "office_callback_invalid", "Ожидался объект callback")
    # JWT_IN_BODY=true: use ONLY signed claims, not unsigned siblings from the HTTP body.
    claims = verify(body.get("token", ""), orch.settings.onlyoffice.jwt_secret)
    key, status = claims.get("key"), claims.get("status")
    if not isinstance(key, str) or type(status) is not int or status not in {1, 2, 3, 4, 6, 7}:
        return {"error": 1}
    document = office.get(document_id)
    if document["active_key"] != key:
        return {"error": 1}
    data = None
    raw_data = None
    normalized_parts: list[str] = []
    if status in {2, 6}:
        try:
            cfg = orch.settings.onlyoffice
            url = saved_url(claims["url"], cfg.internal_url, cfg.public_url)
            async with httpx.AsyncClient(
                timeout=60, follow_redirects=False, trust_env=False
            ) as client:
                async with client.stream("GET", url) as response:
                    response.raise_for_status()
                    chunks = bytearray()
                    async for chunk in response.aiter_bytes():
                        chunks.extend(chunk)
                        if len(chunks) > orch.files.max_bytes:
                            raise ValueError("save exceeds size limit")
                    data = bytes(chunks)
            raw_data = data
            with tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / "deck.pptx"
                path.write_bytes(data)
                # Check ZIP paths, CRC and uncompressed size before reading XML parts.
                if orch.files.check(path, "pptx")["status"] != "ok":
                    raise ValueError("invalid saved PPTX")
            data, normalized_parts = validate_saved_pptx(
                data, office.read(document_id, document["seed_revision"])
            )
        except (httpx.HTTPError, ValueError, KeyError, TypeError):
            office.callback(document_id, key, 7, None)
            return {"error": 1}
    return {
        "error": 0
        if office.callback(
            document_id,
            key,
            status,
            data,
            raw_pptx=raw_data if normalized_parts else None,
            normalized_parts=normalized_parts,
        )
        else 1
    }
