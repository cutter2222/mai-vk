"""Разбор DOCX (python-docx): заголовки по стилям, абзацы, списки, цитаты, таблицы,
подписи, встроенные картинки в порядке документа."""

from __future__ import annotations

import pathlib
from typing import Any

from presentation_designer.parsing.content.parsers.base import (
    Location,
    ParsedBlock,
    ParsedDocument,
    ParsedImage,
    ParsedTable,
    ParserError,
    clean_text,
    guess_image_kind,
    image_dimensions,
    split_long,
)

VERSION = "0.1.0"

_NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}


def parse(
    path: pathlib.Path, *, max_block_chars: int = 2000, min_image_px: int = 64
) -> ParsedDocument:
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as e:  # pragma: no cover
        raise ParserError("parser_unavailable", "python-docx не установлен") from e
    try:
        with path.open("rb") as handle:
            doc = Document(handle)
    except Exception as e:
        raise ParserError("docx_unreadable", f"DOCX не открывается: {e}") from e
    out = ParsedDocument("docx")
    body = doc.element.body
    pending_bullets: list[str] = []
    pending_caption: str | None = None
    offset = 0
    tables = 0

    def flush_bullets() -> None:
        nonlocal pending_bullets
        if pending_bullets:
            out.blocks.append(ParsedBlock("bullets", items=pending_bullets, location=Location()))
            pending_bullets = []

    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(child, doc)
            style = (paragraph.style.name if paragraph.style is not None else "") or ""
            text = clean_text(paragraph.text)
            for image in _images_of(paragraph, doc, out, min_image_px):
                flush_bullets()
                image.caption = pending_caption
                out.images.append(image)
                out.blocks.append(
                    ParsedBlock(
                        "figure",
                        image_ref=len(out.images) - 1,
                        caption=pending_caption,
                        location=Location(char_offset=offset),
                    )
                )
                pending_caption = None
            if not text:
                continue
            loc = Location(char_offset=offset)
            offset += len(text) + 1
            lowered = style.lower()
            if (
                lowered.startswith("heading")
                or lowered == "title"
                or lowered.startswith("заголовок")
            ):
                flush_bullets()
                level = _heading_level(style)
                out.blocks.append(ParsedBlock("heading", text=text, level=level, location=loc))
            elif lowered.startswith("caption") or lowered.startswith("название"):
                flush_bullets()
                # Подпись относится к предыдущей таблице/рисунку без подписи или к следующему:
                # «Рисунок…» не приписывается таблице, «Таблица…» — рисунку.
                caption_lower = text.lower()
                figure_caption = caption_lower.startswith(("рис", "figure", "fig"))
                table_caption = caption_lower.startswith(("табл", "table"))
                prev = out.blocks[-1] if out.blocks else None
                if (
                    prev is not None
                    and prev.kind == "table"
                    and not prev.caption
                    and not figure_caption
                    and prev.table_ref is not None
                ):
                    prev.caption = text
                    out.tables[prev.table_ref].title = text
                elif (
                    prev is not None
                    and prev.kind == "figure"
                    and not prev.caption
                    and not table_caption
                ):
                    prev.caption = text
                    if prev.image_ref is not None:
                        out.images[prev.image_ref].caption = text
                else:
                    pending_caption = text
            elif _is_list(paragraph, lowered):
                pending_bullets.append(text)
            elif "quote" in lowered or "цитат" in lowered:
                flush_bullets()
                out.blocks.append(ParsedBlock("quote", text=text, location=loc))
            else:
                flush_bullets()
                for piece in split_long(text, max_block_chars):
                    out.blocks.append(ParsedBlock("paragraph", text=piece, location=loc))
        elif tag == "tbl":
            flush_bullets()
            table = Table(child, doc)
            rows = _table_rows(table)
            if not rows:
                continue
            tables += 1
            out.tables.append(ParsedTable(rows=rows, title=pending_caption, location=Location()))
            out.blocks.append(
                ParsedBlock(
                    "table",
                    table_ref=len(out.tables) - 1,
                    caption=pending_caption,
                    location=Location(char_offset=offset),
                )
            )
            pending_caption = None
    flush_bullets()
    out.units = {
        "tables": tables,
        "images": len(out.images),
        "chars": out.text_chars,
    }
    if not out.blocks and not out.images:
        out.extracted = False
        out.warn("no_content", "в документе не найдено текста, таблиц и изображений")
    return out


def _heading_level(style: str) -> int:
    digits = "".join(ch for ch in style if ch.isdigit())
    if style.lower() == "title":
        return 1
    return max(1, min(6, int(digits))) if digits else 1


def _is_list(paragraph: Any, style_lower: str) -> bool:
    if "list" in style_lower or "список" in style_lower:
        return True
    ppr = paragraph._p.pPr
    return ppr is not None and ppr.numPr is not None


def _table_rows(table: Any) -> list[list[Any]]:
    rows: list[list[Any]] = []
    for row in table.rows:
        cells: list[Any] = []
        seen: set[int] = set()
        for cell in row.cells:
            # Объединённые ячейки python-docx повторяет одним объектом: пропускаем дубли.
            key = id(cell._tc)
            if key in seen:
                continue
            seen.add(key)
            cells.append(clean_text(cell.text) or None)
        if any(c is not None for c in cells):
            rows.append(cells)
    width = max((len(r) for r in rows), default=0)
    return [r + [None] * (width - len(r)) for r in rows]


def _images_of(
    paragraph: Any, doc: Any, out: ParsedDocument, min_image_px: int
) -> list[ParsedImage]:
    images: list[ParsedImage] = []
    for rid in paragraph._p.xpath(".//a:blip/@r:embed"):
        part = doc.part.related_parts.get(rid)
        if part is None:
            continue
        blob = part.blob
        try:
            width, height, mime = image_dimensions(blob)
        except ParserError as e:
            out.warn("image_skipped", f"{getattr(part, 'partname', rid)}: {e}")
            continue
        if max(width, height) < min_image_px:
            continue
        name = pathlib.PurePosixPath(str(part.partname)).name
        images.append(
            ParsedImage(
                data=blob,
                mime=mime,
                name=name,
                width_px=width,
                height_px=height,
                kind=guess_image_kind(blob, width, height, mime),
            )
        )
    return images
