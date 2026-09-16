"""Разбор XLSX (openpyxl) и CSV: лист → таблица с числовыми форматами колонок.

Пустые строки и колонки отбрасываются, объединённые ячейки читаются из верхней левой,
даты приводятся к ISO-строкам. Форматы ячеек (`%`, валюта, дата) запоминаются по колонкам,
чтобы наборы данных получили типы percent/money/date без угадывания по тексту.
"""

from __future__ import annotations

import csv
import io
import pathlib
from collections import Counter
from datetime import date, datetime, time
from typing import Any

from presentation_designer.parsing.content.parsers.base import (
    Location,
    ParsedBlock,
    ParsedDocument,
    ParsedTable,
    ParserError,
    clean_text,
    decode_text,
)

VERSION = "0.1.0"

_MONEY_MARKERS = ("₽", "$", "€", "£", "руб", "[$", "р.")


def format_kind(number_format: str | None) -> str | None:
    """percent | money | date | number по формату ячейки Excel; None для General и текста."""
    if not number_format or number_format in ("General", "@"):
        return None
    fmt = number_format
    if "%" in fmt:
        return "percent"
    if any(marker in fmt for marker in _MONEY_MARKERS):
        return "money"
    lowered = fmt.lower().replace("\\", "")
    if any(token in lowered for token in ("yy", "dd", "mm", "d.m", "mmm")) and "0" not in lowered:
        return "date"
    if any(ch in fmt for ch in "0#"):
        return "number"
    return None


def parse_xlsx(path: pathlib.Path, *, max_rows: int = 5000) -> ParsedDocument:
    try:
        from openpyxl import load_workbook
    except ImportError as e:  # pragma: no cover
        raise ParserError("parser_unavailable", "openpyxl не установлен") from e
    try:
        # По файловому объекту: в хранилище загрузок у файлов нет расширения, а openpyxl
        # проверяет его у пути.
        handle = path.open("rb")
        wb = load_workbook(handle, data_only=True, read_only=True)
    except Exception as e:
        raise ParserError("xlsx_unreadable", f"XLSX не открывается: {e}") from e
    out = ParsedDocument("xlsx")
    sheets = 0
    try:
        for ws in wb.worksheets:
            sheets += 1
            if getattr(ws, "sheet_state", "visible") != "visible":
                out.warn("sheet_hidden", f"лист «{ws.title}» скрыт и пропущен")
                continue
            rows: list[list[Any]] = []
            formats: list[list[str | None]] = []
            first_row = None
            last_row = 0
            for r_index, row in enumerate(ws.iter_rows(), start=1):
                values = [_cell_value(c) for c in row]
                if not any(v is not None for v in values):
                    continue
                if first_row is None:
                    first_row = r_index
                last_row = r_index
                rows.append(values)
                formats.append(
                    [
                        format_kind(getattr(c, "number_format", None))
                        if getattr(c, "value", None) is not None
                        else None
                        for c in row
                    ]
                )
                if len(rows) >= max_rows:
                    out.warn("sheet_truncated", f"лист «{ws.title}»: прочитано {max_rows} строк")
                    break
            if not rows:
                continue
            rows, formats, first_col, last_col = _drop_empty_columns(rows, formats)
            if not rows:
                continue
            title = None
            # Строка-заголовок листа: одна заполненная ячейка над таблицей.
            if (
                len(rows) > 1
                and sum(v is not None for v in rows[0]) == 1
                and isinstance(next(v for v in rows[0] if v is not None), str)
            ):
                title = str(next(v for v in rows[0] if v is not None))
                rows, formats = rows[1:], formats[1:]
                first_row = (first_row or 1) + 1
            column_formats = _column_formats(formats)
            cell_range = f"{_col_letter(first_col)}{first_row}:{_col_letter(last_col)}{last_row}"
            table = ParsedTable(
                rows=rows,
                title=title or ws.title,
                location=Location(sheet=ws.title, cell_range=cell_range),
                column_formats=column_formats,
            )
            out.tables.append(table)
            out.blocks.append(
                ParsedBlock(
                    "table",
                    table_ref=len(out.tables) - 1,
                    caption=title,
                    location=Location(sheet=ws.title, cell_range=cell_range),
                )
            )
    finally:
        wb.close()
        handle.close()
    out.units = {"sheets": sheets, "tables": len(out.tables)}
    if not out.tables:
        out.extracted = False
        out.warn("no_content", "в книге нет заполненных листов")
    return out


def parse_csv(
    path: pathlib.Path, *, max_rows: int = 5000, name: str | None = None
) -> ParsedDocument:
    raw = path.read_bytes()
    text, encoding = decode_text(raw)
    out = ParsedDocument("csv")
    if encoding != "utf-8":
        out.warn("encoding_guessed", f"кодировка определена как {encoding}")
    sample = text[:4096]
    try:
        dialect: Any = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
        if sample.count(";") > sample.count(","):
            dialect = csv.excel_tab if sample.count("\t") > sample.count(";") else _Semicolon
    rows: list[list[Any]] = []
    for row in csv.reader(io.StringIO(text), dialect):
        values = [clean_text(v) or None for v in row]
        if not any(v is not None for v in values):
            continue
        rows.append(values)
        if len(rows) >= max_rows:
            out.warn("sheet_truncated", f"прочитано {max_rows} строк")
            break
    if not rows:
        out.extracted = False
        out.warn("no_content", "в CSV нет заполненных строк")
        return out
    width = max(len(r) for r in rows)
    rows = [r + [None] * (width - len(r)) for r in rows]
    title = pathlib.Path(name or path.name).stem
    out.tables.append(ParsedTable(rows=rows, title=title, location=Location()))
    out.blocks.append(ParsedBlock("table", table_ref=0, location=Location()))
    out.units = {"tables": 1}
    return out


class _Semicolon(csv.excel):
    delimiter = ";"


def _cell_value(cell: Any) -> Any:
    value = getattr(cell, "value", None)
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0) else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    if isinstance(value, bool):
        return "да" if value else "нет"
    if isinstance(value, int | float):
        return value
    text = clean_text(str(value))
    return text or None


def _drop_empty_columns(
    rows: list[list[Any]], formats: list[list[str | None]]
) -> tuple[list[list[Any]], list[list[str | None]], int, int]:
    width = max(len(r) for r in rows)
    rows = [r + [None] * (width - len(r)) for r in rows]
    formats = [f + [None] * (width - len(f)) for f in formats]
    keep = [i for i in range(width) if any(r[i] is not None for r in rows)]
    if not keep:
        return [], [], 1, 1
    first_col, last_col = keep[0] + 1, keep[-1] + 1
    new_rows = [[r[i] for i in keep] for r in rows]
    new_formats = [[f[i] for i in keep] for f in formats]
    return new_rows, new_formats, first_col, last_col


def _column_formats(formats: list[list[str | None]]) -> list[str | None]:
    if not formats:
        return []
    width = max(len(f) for f in formats)
    out: list[str | None] = []
    for i in range(width):
        counter = Counter(f[i] for f in formats[1:] if i < len(f) and f[i])
        out.append(counter.most_common(1)[0][0] if counter else None)
    return out


def _col_letter(index: int) -> str:
    letters = ""
    while index > 0:
        index, rem = divmod(index - 1, 26)
        letters = chr(65 + rem) + letters
    return letters or "A"
