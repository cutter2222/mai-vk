"""PPTX как материал (python-pptx): текст слайдов и заметок, таблицы, данные нативных
диаграмм, картинки. Файл не становится шаблоном оформления — стили, макеты и мастера
не читаются; скрытые слайды пропускаются с предупреждением."""

from __future__ import annotations

import pathlib
from typing import Any

from presentation_designer.parsing.content.parsers.base import (
    Location,
    ParsedBlock,
    ParsedDocument,
    ParsedImage,
    ParsedTable,
    ParserError,
    clean_text,
    guess_image_kind,
    image_dimensions,
    split_long,
)

VERSION = "0.1.0"


def parse(
    path: pathlib.Path, *, max_block_chars: int = 2000, min_image_px: int = 64
) -> ParsedDocument:
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError as e:  # pragma: no cover
        raise ParserError("parser_unavailable", "python-pptx не установлен") from e
    try:
        with path.open("rb") as handle:
            prs = Presentation(handle)
    except Exception as e:
        raise ParserError("pptx_unreadable", f"PPTX не открывается: {e}") from e
    out = ParsedDocument("pptx")
    slides = 0
    hidden: list[int] = []
    for index, slide in enumerate(prs.slides, start=1):
        if slide._element.get("show") == "0":
            hidden.append(index)
            continue
        slides += 1
        loc = Location(slide=index)
        title_shape = slide.shapes.title
        title_id = id(title_shape._element) if title_shape is not None else None
        title = (
            clean_text(title_shape.text_frame.text)
            if title_shape is not None and title_shape.has_text_frame
            else ""
        )
        if title:
            out.blocks.append(ParsedBlock("heading", text=title, level=1, location=loc))
        for shape in _iter_shapes(slide.shapes):
            if id(shape._element) == title_id:
                continue
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [
                    [clean_text(cell.text) or None for cell in row.cells]
                    for row in shape.table.rows
                ]
                rows = [r for r in rows if any(c is not None for c in r)]
                if rows:
                    out.tables.append(ParsedTable(rows=rows, title=title or None, location=loc))
                    out.blocks.append(
                        ParsedBlock(
                            "table",
                            table_ref=len(out.tables) - 1,
                            caption=title or None,
                            location=loc,
                        )
                    )
                continue
            if getattr(shape, "has_chart", False) and shape.has_chart:
                table = _chart_table(shape.chart, title)
                if table is not None:
                    table.location = loc
                    out.tables.append(table)
                    out.blocks.append(
                        ParsedBlock(
                            "table",
                            table_ref=len(out.tables) - 1,
                            caption=table.title,
                            location=loc,
                            tags=["chart"],
                        )
                    )
                continue
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image = _picture(shape, out, min_image_px, loc)
                if image is not None:
                    out.images.append(image)
                    out.blocks.append(
                        ParsedBlock("figure", image_ref=len(out.images) - 1, location=loc)
                    )
                continue
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                out.blocks.extend(_text_blocks(shape.text_frame, loc, max_block_chars))
        if slide.has_notes_slide:
            notes = clean_text(slide.notes_slide.notes_text_frame.text)
            if notes:
                for piece in split_long(notes.replace("\n", " "), max_block_chars):
                    out.blocks.append(
                        ParsedBlock(
                            "paragraph",
                            text=piece,
                            location=Location(slide=index, notes=True),
                            tags=["notes"],
                        )
                    )
    if hidden:
        out.warn("slides_hidden", f"скрытые слайды пропущены: {', '.join(map(str, hidden))}")
    out.units = {
        "slides": slides,
        "tables": len(out.tables),
        "images": len(out.images),
        "chars": out.text_chars,
    }
    if not out.blocks and not out.images:
        out.extracted = False
        out.warn("no_content", "в презентации не найдено текста, таблиц и изображений")
    return out


def _iter_shapes(shapes: Any) -> Any:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


def _text_blocks(text_frame: Any, loc: Location, max_block_chars: int) -> list[ParsedBlock]:
    paragraphs = [clean_text(p.text) for p in text_frame.paragraphs]
    paragraphs = [p for p in paragraphs if p]
    if not paragraphs:
        return []
    if len(paragraphs) == 1:
        text = paragraphs[0]
        if len(text) <= 90 and not text.endswith("."):
            # Короткая одиночная строка — подзаголовок или подпись.
            return [ParsedBlock("paragraph", text=text, location=loc, tags=["short"])]
        return [
            ParsedBlock("paragraph", text=piece, location=loc)
            for piece in split_long(text, max_block_chars)
        ]
    # Несколько абзацев в одной рамке — список пунктов слайда.
    return [ParsedBlock("bullets", items=paragraphs, location=loc)]


def _chart_table(chart: Any, slide_title: str) -> ParsedTable | None:
    """Категории и ряды диаграммы как таблица: первая колонка — категория, далее ряды."""
    try:
        plots = list(chart.plots)
    except Exception:
        return None
    if not plots:
        return None
    plot = plots[0]
    try:
        categories = [clean_text(str(c)) for c in plot.categories]
        series = list(plot.series)
    except Exception:
        return None
    if not categories or not series:
        return None
    header: list[Any] = [
        "Категория",
        *[clean_text(s.name) or f"Ряд {i + 1}" for i, s in enumerate(series)],
    ]
    rows: list[list[Any]] = [header]
    values = [list(s.values) for s in series]
    for i, cat in enumerate(categories):
        rows.append([cat, *[(v[i] if i < len(v) else None) for v in values]])
    title = None
    try:
        if chart.has_title and chart.chart_title.has_text_frame:
            title = clean_text(chart.chart_title.text_frame.text) or None
    except Exception:
        title = None
    return ParsedTable(rows=rows, title=title or slide_title or None)


def _picture(
    shape: Any, out: ParsedDocument, min_image_px: int, loc: Location
) -> ParsedImage | None:
    try:
        blob = shape.image.blob
    except Exception:
        return None
    try:
        width, height, mime = image_dimensions(blob)
    except ParserError as e:
        out.warn("image_skipped", f"слайд {loc.slide}: {e}")
        return None
    if max(width, height) < min_image_px:
        return None
    name = f"slide-{loc.slide}-{clean_text(shape.name) or 'image'}.{shape.image.ext}"
    return ParsedImage(
        data=blob,
        mime=mime,
        name=name.replace(" ", "_"),
        width_px=width,
        height_px=height,
        location=loc,
        kind=guess_image_kind(blob, width, height, mime),
    )
