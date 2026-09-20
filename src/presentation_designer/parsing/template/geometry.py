"""Геометрия объектов слайда: нормализованные координаты с учётом групп, поворотов, crop и порядка.

Обход дерева фигур python-pptx с собственным пересчётом координат вложенных групп: у группы
`a:xfrm` задаёт положение (off/ext) и систему координат детей (chOff/chExt), дети хранят
координаты в этой системе. Результат — плоский список объектов с bbox в долях слайда, исходными
EMU, поворотом, путём групп, z-order и ссылкой на элемент (`p:cNvPr@id`), по которой вёрстка
находит объект в клоне образца.
"""

from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass, field
from typing import Any

from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.util import Emu

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}

PLACEHOLDER_TYPES = {
    PP_PLACEHOLDER.TITLE: "title",
    PP_PLACEHOLDER.CENTER_TITLE: "ctrTitle",
    PP_PLACEHOLDER.SUBTITLE: "subTitle",
    PP_PLACEHOLDER.BODY: "body",
    PP_PLACEHOLDER.OBJECT: "obj",
    PP_PLACEHOLDER.PICTURE: "pic",
    PP_PLACEHOLDER.CHART: "chart",
    PP_PLACEHOLDER.TABLE: "tbl",
    PP_PLACEHOLDER.DATE: "dt",
    PP_PLACEHOLDER.FOOTER: "ftr",
    PP_PLACEHOLDER.SLIDE_NUMBER: "sldNum",
}


@dataclass
class Paragraph:
    text: str
    level: int = 0
    bullet: str | None = None  # none | char | number | picture
    bullet_char: str | None = None
    runs: list[dict[str, Any]] = field(default_factory=list)  # локальные свойства фрагментов
    align: str | None = None
    line_spacing: float | None = None
    space_before_pt: float | None = None
    space_after_pt: float | None = None
    indent_emu: int | None = None
    fields: list[str] = field(default_factory=list)  # типы a:fld: slidenum, datetime…


@dataclass
class ShapeInfo:
    element_id: str
    name: str
    kind: str  # text | picture | table | chart | graphic | shape | connector | group | unknown
    x: float
    y: float
    width: float
    height: float
    left_emu: int
    top_emu: int
    width_emu: int
    height_emu: int
    z_order: int
    rotation_deg: float = 0.0
    group_path: list[str] = field(default_factory=list)
    placeholder_type: str | None = None
    placeholder_idx: int | None = None
    text: str = ""
    paragraphs: list[Paragraph] = field(default_factory=list)
    has_text_frame: bool = False
    autofit: str | None = None  # none | shrink | resize_shape
    insets_emu: tuple[int, int, int, int] | None = None
    anchor: str | None = None  # t | ctr | b
    wrap: bool = True
    fill_hex: str | None = None
    fill_kind: str | None = None  # solid | gradient | picture | none | inherit
    # Ссылка на заливку темы: p:style/a:fillRef@idx. У фигур, нарисованных в PowerPoint
    # кнопкой «прямоугольник», своей заливки в файле нет — она вся здесь.
    fill_ref_idx: int | None = None
    fill_ref_element: Any = field(default=None, repr=False)
    line_hex: str | None = None
    # Толщина контура в пунктах (a:ln@w). None — не задана у самой фигуры: её берут из темы.
    line_width_pt: float | None = None
    # Ссылка на стиль темы: p:style/a:lnRef@idx. Контур карточек шаблонов чаще всего только
    # здесь и живёт, у самой фигуры ни цвета, ни толщины нет.
    line_ref_idx: int | None = None
    line_ref_element: Any = field(default=None, repr=False)
    # Скругление roundRect: «adj» из a:prstGeom/a:avLst долей от меньшей стороны фигуры.
    geometry_adjust: float | None = None
    media_part: str | None = None
    media_sha256: str | None = None
    media_ext: str | None = None
    media_blob: bytes | None = field(default=None, repr=False)
    crop: tuple[float, float, float, float] | None = None  # left, top, right, bottom (доли)
    geometry: str | None = None  # prstGeom: rect, ellipse, roundRect…
    table_size: tuple[int, int] | None = None
    chart_type: str | None = None
    graphic_uri: str | None = None
    hyperlink: bool = False
    is_hidden: bool = False
    element: Any = field(default=None, repr=False)

    @property
    def area(self) -> float:
        return max(self.width, 0.0) * max(self.height, 0.0)

    @property
    def bbox(self) -> dict[str, float]:
        return {
            "x": round(self.x, 4),
            "y": round(self.y, 4),
            "width": round(self.width, 4),
            "height": round(self.height, 4),
        }

    @property
    def center(self) -> tuple[float, float]:
        return (self.x + self.width / 2, self.y + self.height / 2)

    @property
    def is_placeholder(self) -> bool:
        return self.placeholder_type is not None


@dataclass
class _Transform:
    """Аффинное отображение координат детей группы в координаты слайда (без поворота группы)."""

    off_x: float = 0.0
    off_y: float = 0.0
    scale_x: float = 1.0
    scale_y: float = 1.0
    ch_off_x: float = 0.0
    ch_off_y: float = 0.0

    def apply(self, x: float, y: float, w: float, h: float) -> tuple[float, float, float, float]:
        return (
            self.off_x + (x - self.ch_off_x) * self.scale_x,
            self.off_y + (y - self.ch_off_y) * self.scale_y,
            w * self.scale_x,
            h * self.scale_y,
        )

    def compose(self, xfrm: Any) -> _Transform:
        """Преобразование для детей вложенной группы с данным a:xfrm."""
        off = xfrm.find("a:off", NS)
        ext = xfrm.find("a:ext", NS)
        ch_off = xfrm.find("a:chOff", NS)
        ch_ext = xfrm.find("a:chExt", NS)
        if off is None or ext is None:
            return self
        gx, gy = float(off.get("x", 0)), float(off.get("y", 0))
        gw, gh = float(ext.get("cx", 0)), float(ext.get("cy", 0))
        cx = float(ch_off.get("x", 0)) if ch_off is not None else 0.0
        cy = float(ch_off.get("y", 0)) if ch_off is not None else 0.0
        cw = float(ch_ext.get("cx", 0)) if ch_ext is not None else gw
        ch = float(ch_ext.get("cy", 0)) if ch_ext is not None else gh
        sx = gw / cw if cw else 1.0
        sy = gh / ch if ch else 1.0
        # Сначала переводим положение группы в координаты слайда через внешнее преобразование.
        px, py, pw, ph = self.apply(gx, gy, gw, gh)
        return _Transform(
            px, py, (pw / gw if gw else 1.0) * sx, (ph / gh if gh else 1.0) * sy, cx, cy
        )


def _shape_kind(shape: Any) -> str:
    st = shape.shape_type
    if st == MSO_SHAPE_TYPE.GROUP:
        return "group"
    if st == MSO_SHAPE_TYPE.PICTURE or (getattr(shape, "image", None) is not None and st is None):
        return "picture"
    if getattr(shape, "has_table", False) and shape.has_table:
        return "table"
    if getattr(shape, "has_chart", False) and shape.has_chart:
        return "chart"
    if st == MSO_SHAPE_TYPE.LINE:
        return "connector"
    if shape._element.tag.endswith("}graphicFrame"):
        return "graphic"
    if shape._element.tag.endswith("}pic"):
        return "picture"
    if shape.has_text_frame:
        return "text"
    return "shape"


def _xfrm_of(shape: Any) -> Any:
    el = shape._element
    for path in ("p:spPr/a:xfrm", "p:grpSpPr/a:xfrm", "p:xfrm"):
        found = el.find(path, NS)
        if found is not None:
            return found
    return None


def _raw_geometry(shape: Any) -> tuple[float, float, float, float] | None:
    """Координаты в EMU из xfrm; для плейсхолдеров без xfrm python-pptx наследует их от макета."""
    xfrm = _xfrm_of(shape)
    if xfrm is not None:
        off, ext = xfrm.find("a:off", NS), xfrm.find("a:ext", NS)
        if off is not None and ext is not None:
            return (
                float(off.get("x", 0)),
                float(off.get("y", 0)),
                float(ext.get("cx", 0)),
                float(ext.get("cy", 0)),
            )
    try:
        left, top, width, height = shape.left, shape.top, shape.width, shape.height
    except Exception:
        return None
    if left is None or top is None or width is None or height is None:
        return None
    return float(left), float(top), float(width), float(height)


def _rotation(shape: Any) -> float:
    xfrm = _xfrm_of(shape)
    if xfrm is None:
        return 0.0
    try:
        return float(xfrm.get("rot", 0)) / 60000.0
    except ValueError:
        return 0.0


def _hex_of(color_el: Any) -> str | None:
    if color_el is None:
        return None
    srgb = color_el.find("a:srgbClr", NS)
    if srgb is not None:
        return f"#{srgb.get('val', '000000').upper()}"
    return None


def _fill(shape: Any) -> tuple[str | None, str | None]:
    """Заливка самой фигуры. Ссылка на заливку темы читается отдельно — `_fill_ref`."""
    sp_pr = shape._element.find("p:spPr", NS)
    if sp_pr is None:
        return None, "inherit"
    if sp_pr.find("a:noFill", NS) is not None:
        return None, "none"
    solid = sp_pr.find("a:solidFill", NS)
    if solid is not None:
        scheme = solid.find("a:schemeClr", NS)
        return _hex_of(solid) or (
            f"scheme:{scheme.get('val')}" if scheme is not None else None
        ), "solid"
    if sp_pr.find("a:gradFill", NS) is not None:
        return None, "gradient"
    if sp_pr.find("a:blipFill", NS) is not None:
        return None, "picture"
    return None, "inherit"


def _fill_ref(shape: Any) -> tuple[int | None, Any]:
    """Заливка из стиля темы: `p:style/a:fillRef@idx` и её цвет.

    То же место, где живёт контур карточек (`a:lnRef`): PowerPoint пишет сюда заливку фигур,
    нарисованных инструментом, и в `p:spPr` тогда пусто. Пока это не читалось, белая плашка
    на половину слайда приезжала в ComposedDeck без заливки — на картинке слайда она была,
    на холсте редактора её не было.
    """
    ref = shape._element.find("p:style/a:fillRef", NS)
    if ref is None:
        return None, None
    try:
        idx = int(ref.get("idx", "0")) or None
    except ValueError:
        idx = None
    return idx, ref


def _line(shape: Any) -> tuple[str | None, float | None, int | None, Any]:
    """Контур фигуры: цвет, толщина в пунктах и ссылка на линию темы.

    У фигур из шаблонов контур обычно не записан в самой фигуре: `a:ln` пустой, а цвет и
    толщина берутся из стиля темы по `p:style/a:lnRef@idx`. Пока это не читалось, карточки
    шаблона приезжали в ComposedDeck без контура и пропадали на холсте редактора и в HTML.
    """
    element = shape._element
    sp_pr = element.find("p:spPr", NS)
    ln = sp_pr.find("a:ln", NS) if sp_pr is not None else None
    if ln is not None and ln.find("a:noFill", NS) is not None:
        return None, None, None, None
    color: str | None = None
    width: float | None = None
    if ln is not None:
        solid = ln.find("a:solidFill", NS)
        if solid is not None:
            scheme = solid.find("a:schemeClr", NS)
            ref = f"scheme:{scheme.get('val')}" if scheme is not None else None
            color = _hex_of(solid) or ref
        raw = ln.get("w")
        if raw:
            try:
                width = round(int(raw) / 12700, 3)
            except ValueError:
                width = None
    ref = element.find("p:style/a:lnRef", NS)
    idx: int | None = None
    if ref is not None:
        try:
            idx = int(ref.get("idx", "0")) or None
        except ValueError:
            idx = None
    return color, width, idx, ref


def _geometry_adjust(shape: Any) -> float | None:
    """Скругление углов: «adj» из a:avLst долей от меньшей стороны фигуры.

    В PowerPoint радиус скругления считается от меньшей стороны, поэтому в отличие от
    процентов CSS одно и то же значение на широкой и на узкой фигуре даёт разный вид.
    """
    geom = shape._element.find("p:spPr/a:prstGeom", NS)
    if geom is None:
        return None
    for gd in geom.findall("a:avLst/a:gd", NS):
        if gd.get("name") != "adj":
            continue
        formula = gd.get("fmla") or ""
        if not formula.startswith("val "):
            continue
        try:
            return round(int(formula[4:]) / 100000, 5)
        except ValueError:
            return None
    return None


_FIELD_RE = re.compile(r"\s+")


def _paragraphs(shape: Any) -> list[Paragraph]:
    out: list[Paragraph] = []
    tx_body = shape._element.find(".//p:txBody", NS)
    if tx_body is None:
        tx_body = shape._element.find(".//a:txBody", NS)
    if tx_body is None:
        return out
    for p in tx_body.findall("a:p", NS):
        p_pr = p.find("a:pPr", NS)
        level = int(p_pr.get("lvl", 0)) if p_pr is not None else 0
        bullet: str | None = None
        bullet_char: str | None = None
        align: str | None = None
        line_spacing: float | None = None
        before: float | None = None
        after: float | None = None
        indent: int | None = None
        if p_pr is not None:
            align = p_pr.get("algn")
            if p_pr.find("a:buNone", NS) is not None:
                bullet = "none"
            elif p_pr.find("a:buChar", NS) is not None:
                bullet = "char"
                bullet_char = p_pr.find("a:buChar", NS).get("char")
            elif p_pr.find("a:buAutoNum", NS) is not None:
                bullet = "number"
            elif p_pr.find("a:buBlip", NS) is not None:
                bullet = "picture"
            ln = p_pr.find("a:lnSpc/a:spcPct", NS)
            if ln is not None:
                line_spacing = float(ln.get("val", 100000)) / 100000.0
            ln_pts = p_pr.find("a:lnSpc/a:spcPts", NS)
            if ln_pts is not None:
                line_spacing = -float(ln_pts.get("val", 0)) / 100.0  # отрицательное = в пунктах
            sb = p_pr.find("a:spcBef/a:spcPts", NS)
            if sb is not None:
                before = float(sb.get("val", 0)) / 100.0
            sa = p_pr.find("a:spcAft/a:spcPts", NS)
            if sa is not None:
                after = float(sa.get("val", 0)) / 100.0
            if p_pr.get("indent") is not None:
                indent = int(p_pr.get("indent"))
        runs: list[dict[str, Any]] = []
        texts: list[str] = []
        fields: list[str] = []
        for child in p:
            tag = child.tag.split("}")[-1]
            if tag in ("r", "fld"):
                t = child.find("a:t", NS)
                text = t.text if t is not None and t.text else ""
                texts.append(text)
                r_pr = child.find("a:rPr", NS)
                run: dict[str, Any] = {"text": text}
                if r_pr is not None:
                    if r_pr.get("sz"):
                        run["size_pt"] = int(r_pr.get("sz")) / 100.0
                    if r_pr.get("b") is not None:
                        run["bold"] = r_pr.get("b") in ("1", "true")
                    if r_pr.get("i") is not None:
                        run["italic"] = r_pr.get("i") in ("1", "true")
                    if r_pr.get("cap"):
                        run["cap"] = r_pr.get("cap")
                    latin = r_pr.find("a:latin", NS)
                    if latin is not None and latin.get("typeface"):
                        run["family"] = latin.get("typeface")
                    solid = r_pr.find("a:solidFill", NS)
                    if solid is not None:
                        run["color"] = _hex_of(solid)
                        scheme = solid.find("a:schemeClr", NS)
                        if scheme is not None:
                            run["scheme_color"] = scheme.get("val")
                    if r_pr.find("a:hlinkClick", NS) is not None:
                        run["hyperlink"] = True
                if tag == "fld":
                    fields.append(child.get("type", "unknown"))
                    run["field"] = child.get("type")
                runs.append(run)
            elif tag == "br":
                texts.append("\n")
        text = "".join(texts)
        out.append(
            Paragraph(
                text=text,
                level=level,
                bullet=bullet,
                bullet_char=bullet_char,
                runs=runs,
                align=align,
                line_spacing=line_spacing,
                space_before_pt=before,
                space_after_pt=after,
                indent_emu=indent,
                fields=fields,
            )
        )
    return out


def _body_props(
    shape: Any,
) -> tuple[str | None, tuple[int, int, int, int] | None, str | None, bool]:
    body = shape._element.find(".//a:bodyPr", NS)
    if body is None:
        return None, None, None, True
    autofit = "none"
    if body.find("a:normAutofit", NS) is not None:
        autofit = "shrink"
    elif body.find("a:spAutoFit", NS) is not None:
        autofit = "resize_shape"
    insets = (
        int(body.get("lIns", 91440)),
        int(body.get("tIns", 45720)),
        int(body.get("rIns", 91440)),
        int(body.get("bIns", 45720)),
    )
    return autofit, insets, body.get("anchor"), body.get("wrap", "square") != "none"


def _picture_info(
    shape: Any, part: Any
) -> tuple[str | None, str | None, str | None, Any, bytes | None]:
    """Часть изображения, sha256 байтов, расширение, crop из a:srcRect и сами байты."""
    blip = shape._element.find(".//a:blip", NS)
    if blip is None:
        return None, None, None, None, None
    r_id = blip.get(f"{{{NS['r']}}}embed")
    if not r_id:
        return None, None, None, None, None
    try:
        image_part = part.related_part(r_id)
    except KeyError:
        return None, None, None, None, None
    blob = bytes(image_part.blob)
    src_rect = shape._element.find(".//a:srcRect", NS)
    crop = None
    if src_rect is not None:
        crop = tuple(float(src_rect.get(side, 0)) / 100000.0 for side in ("l", "t", "r", "b"))
    return (
        str(image_part.partname),
        hashlib.sha256(blob).hexdigest(),
        str(image_part.partname).rsplit(".", 1)[-1].lower(),
        crop,
        blob,
    )


def walk_shapes(container: Any, part: Any, slide_w: int, slide_h: int) -> list[ShapeInfo]:
    """Плоский список объектов слайда/макета/мастера в порядке z-order (снизу вверх)."""
    out: list[ShapeInfo] = []
    counter = [0]

    def visit(shapes: Any, transform: _Transform, group_path: list[str]) -> None:
        for shape in shapes:
            kind = _shape_kind(shape)
            raw = _raw_geometry(shape)
            element_id = str(shape.shape_id)
            if raw is None:
                x, y, w, h = 0.0, 0.0, 0.0, 0.0
            else:
                x, y, w, h = transform.apply(*raw)
            counter[0] += 1
            info = ShapeInfo(
                element_id=element_id,
                name=shape.name or "",
                kind=kind,
                x=x / slide_w if slide_w else 0.0,
                y=y / slide_h if slide_h else 0.0,
                width=w / slide_w if slide_w else 0.0,
                height=h / slide_h if slide_h else 0.0,
                left_emu=int(x),
                top_emu=int(y),
                width_emu=int(w),
                height_emu=int(h),
                z_order=counter[0],
                rotation_deg=_rotation(shape),
                group_path=list(group_path),
                element=shape._element,
            )
            if shape.is_placeholder:
                ph = shape.placeholder_format
                info.placeholder_type = PLACEHOLDER_TYPES.get(ph.type, "other")
                info.placeholder_idx = int(ph.idx)
            if kind == "group":
                out.append(info)
                xfrm = _xfrm_of(shape)
                child_t = transform.compose(xfrm) if xfrm is not None else transform
                visit(shape.shapes, child_t, [*group_path, element_id])
                continue
            if shape.has_text_frame:
                info.has_text_frame = True
                info.paragraphs = _paragraphs(shape)
                info.text = "\n".join(p.text for p in info.paragraphs).strip()
                info.autofit, info.insets_emu, info.anchor, info.wrap = _body_props(shape)
            if kind == "picture":
                (
                    info.media_part,
                    info.media_sha256,
                    info.media_ext,
                    info.crop,
                    info.media_blob,
                ) = _picture_info(shape, part)
            if kind == "table":
                tbl = shape.table
                info.table_size = (len(tbl.rows), len(tbl.columns))
                info.text = " ".join(c.text for r in tbl.rows for c in r.cells if c.text.strip())[
                    :500
                ]
            if kind == "chart":
                try:
                    info.chart_type = str(shape.chart.chart_type).split(".")[-1].split(" ")[0]
                except Exception:
                    info.chart_type = "unknown"
            if kind == "graphic":
                uri = shape._element.find(".//a:graphicData", NS)
                info.graphic_uri = uri.get("uri") if uri is not None else None
                if info.graphic_uri and "smartArt" in info.graphic_uri.lower():
                    info.kind = "smartart"
            geom = shape._element.find("p:spPr/a:prstGeom", NS)
            if geom is not None:
                info.geometry = geom.get("prst")
            elif shape._element.find("p:spPr/a:custGeom", NS) is not None:
                # Нарисованная от руки фигура (мокап телефона, стрелка, кольцо диаграммы):
                # прямоугольником её рисовать нельзя — вместо тонкого контура получается
                # залитый блок во всю рамку. Кто рисует, тот и решает, что с ней делать.
                info.geometry = "custom"
            else:
                info.geometry = None
            info.fill_hex, info.fill_kind = _fill(shape)
            info.fill_ref_idx, info.fill_ref_element = _fill_ref(shape)
            info.line_hex, info.line_width_pt, info.line_ref_idx, info.line_ref_element = _line(
                shape
            )
            info.geometry_adjust = _geometry_adjust(shape)
            info.hyperlink = shape._element.find(".//a:hlinkClick", NS) is not None
            out.append(info)

    visit(container.shapes, _Transform(), [])
    return out


def slide_size(prs: Any) -> tuple[int, int]:
    return int(prs.slide_width or Emu(12192000)), int(prs.slide_height or Emu(6858000))


def bbox_iou(a: ShapeInfo, b: ShapeInfo) -> float:
    ax2, ay2 = a.x + a.width, a.y + a.height
    bx2, by2 = b.x + b.width, b.y + b.height
    iw = max(0.0, min(ax2, bx2) - max(a.x, b.x))
    ih = max(0.0, min(ay2, by2) - max(a.y, b.y))
    inter = iw * ih
    union = a.area + b.area - inter
    return inter / union if union > 0 else 0.0


def same_place(a: ShapeInfo, b: ShapeInfo, tol: float = 0.01) -> bool:
    return (
        abs(a.x - b.x) <= tol
        and abs(a.y - b.y) <= tol
        and abs(a.width - b.width) <= tol
        and abs(a.height - b.height) <= tol
    )


def normalize_text(text: str) -> str:
    return _FIELD_RE.sub(" ", text or "").strip().lower()
