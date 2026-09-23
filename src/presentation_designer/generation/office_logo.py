"""Логотип шаблона в офисной копии: заменить или убрать на всех слайдах одним действием.

Знак шаблона лежит не на слайдах, а на макетах и образцах (в VK Education — на девяти
макетах, светлый и тёмный), поэтому в редакторе его не выделить, а в режиме образцов его
пришлось бы менять на каждом макете. Картинки знака известны по профилю шаблона (sha256
ресурсов вида logo); в документе они находятся по содержимому, а не по id фигур: после
сохранения в ONLYOFFICE части и фигуры переименовываются, а байты картинки остаются.
"""

from __future__ import annotations

import hashlib
import io
from typing import Any

from PIL import Image
from pptx import Presentation

from presentation_designer.layout.shapes import NS, drop_unreferenced_rels, remove_shape

RASTER = {"PNG", "JPEG", "GIF", "BMP"}
R_EMBED = f"{{{NS['r']}}}embed"


def logo_hashes(profile: dict[str, Any] | None) -> set[str]:
    """sha256 картинок знака шаблона по профилю: только логотипы, закреплённые на макетах и
    образцах (`fixed_elements`). Логотипы на слайдах — партнёры, соцсети — чужие знаки, их
    замена логотипа шаблона не касается."""
    if not profile:
        return set()
    wanted = {
        str(f.get("asset_id"))
        for f in profile.get("fixed_elements") or []
        if f.get("kind") == "logo" and f.get("asset_id")
    }
    return {
        str(a["sha256"])
        for a in profile.get("assets") or []
        if a.get("sha256") and a.get("asset_id") in wanted
    }


def _containers(prs: Any) -> list[Any]:
    masters = list(prs.slide_masters)
    layouts = [layout for master in masters for layout in master.slide_layouts]
    return [*masters, *layouts, *prs.slides]


def _image_size(image: bytes) -> tuple[int, int]:
    try:
        with Image.open(io.BytesIO(image)) as raw:
            if raw.format not in RASTER:
                raise ValueError("нужна картинка PNG или JPG")
            return raw.size
    except (OSError, SyntaxError) as exc:
        raise ValueError("картинка логотипа не читается: нужна PNG или JPG") from exc


Box = tuple[int, int, int, int]  # x, y, ширина, высота в EMU


def _box(pic: Any) -> Box | None:
    xfrm = pic.find("p:spPr/a:xfrm", NS)
    off = xfrm.find("a:off", NS) if xfrm is not None else None
    ext = xfrm.find("a:ext", NS) if xfrm is not None else None
    if off is None or ext is None:
        return None
    return (
        int(off.get("x", "0")),
        int(off.get("y", "0")),
        int(ext.get("cx", "0")),
        int(ext.get("cy", "0")),
    )


def _set_box(pic: Any, box: Box) -> None:
    xfrm = pic.find("p:spPr/a:xfrm", NS)
    off, ext = xfrm.find("a:off", NS), xfrm.find("a:ext", NS)
    off.set("x", str(box[0]))
    off.set("y", str(box[1]))
    ext.set("cx", str(box[2]))
    ext.set("cy", str(box[3]))


def _near(a: Box, b: Box) -> bool:
    """Части одного знака: рядом по горизонтали и на одной высоте (у VK Tech квадрат VK
    и надпись tech — две картинки вплотную)."""
    gap = max(a[0], b[0]) - min(a[0] + a[2], b[0] + b[2])
    overlap = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
    return gap <= max(a[3], b[3]) and overlap > 0.4 * min(a[3], b[3])


def _marks(pics: list[tuple[Any, Any]]) -> list[list[tuple[Any, Any]]]:
    """Картинки знака, собранные в знаки: соседние в одной группе фигур — один знак."""
    marks: list[list[tuple[Any, Any]]] = []
    for item in pics:
        box = _box(item[0])
        for mark in marks:
            if box and any(
                other[0].getparent() is item[0].getparent()
                and (ob := _box(other[0])) is not None
                and _near(box, ob)
                for other in mark
            ):
                mark.append(item)
                break
        else:
            marks.append([item])
    return marks


def _union(boxes: list[Box]) -> Box:
    x0 = min(b[0] for b in boxes)
    y0 = min(b[1] for b in boxes)
    x1 = max(b[0] + b[2] for b in boxes)
    y1 = max(b[1] + b[3] for b in boxes)
    return x0, y0, x1 - x0, y1 - y0


def _fit(box: Box, size: tuple[int, int], slide_width: int) -> Box:
    """Новый знак занимает место старого: та же высота и сторона прижатия (левый край у
    знака слева, правый — у знака справа); слишком широкий вписывается по ширине."""
    x, y, w, h = box
    if not w or not h or not size[1]:
        return box
    aspect = size[0] / size[1]
    new_w, new_h = round(h * aspect), h
    if new_w > 1.6 * w:
        new_w, new_h = w, round(w / aspect)
    right_side = x + w / 2 > slide_width / 2
    new_x = x + w - new_w if right_side else x
    return new_x, y + (h - new_h) // 2, new_w, new_h


def replace_logo(pptx: bytes, hashes: set[str], image: bytes | None) -> tuple[bytes, int]:
    """Заменяет знак шаблона картинкой `image` (или убирает, если её нет) на образцах, макетах
    и слайдах. Знак из нескольких картинок становится одной картинкой в их общей рамке.
    Возвращает новый PPTX и число заменённых или убранных знаков."""
    size = _image_size(image) if image is not None else None
    prs = Presentation(io.BytesIO(pptx))
    changed = 0
    for container in _containers(prs):
        found = []
        for pic in container.shapes._spTree.iter(f"{{{NS['p']}}}pic"):
            blip = pic.find(".//a:blip", NS)
            rid = blip.get(R_EMBED) if blip is not None else None
            if not rid or rid not in container.part.rels:
                continue
            blob = container.part.related_part(rid).blob
            if hashlib.sha256(blob).hexdigest() in hashes:
                found.append((pic, blip))
        for mark in _marks(found):
            changed += 1
            if image is None or size is None:
                for pic, _ in mark:
                    remove_shape(container, pic)
                continue
            boxes = [b for pic, _ in mark if (b := _box(pic)) is not None]
            (pic, blip), rest = mark[0], mark[1:]
            _, rid = container.part.get_or_add_image_part(io.BytesIO(image))
            blip.set(R_EMBED, rid)
            crop = pic.find(".//a:srcRect", NS)
            if crop is not None:
                crop.getparent().remove(crop)
            if boxes:
                _set_box(pic, _fit(_union(boxes), size, int(prs.slide_width or 0)))
            for other, _ in rest:
                remove_shape(container, other)
        if found and image is not None:
            drop_unreferenced_rels(container)
    if not changed:
        return pptx, 0
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue(), changed


__all__ = ["logo_hashes", "replace_logo"]
