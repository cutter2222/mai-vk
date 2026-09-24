"""Снимок колоды для чата: что на каждом слайде и где — одной формой для ревизии варианта и
для офисной копии (этап 36, `contracts/schemas/deck_snapshot.schema.json`).

Снимок строится по байтам PPTX тем же обходом фигур и тем же разрешением стилей, что и
ComposedDeck (`parsing/template/package.open_template`, `geometry.walk_shapes`,
`styles.StyleResolver`), поэтому ревизия варианта и её копия в ONLYOFFICE дают одно и то же:
объекты в группах с путём, рамку плейсхолдера из макета, поворот, таблицы ячейками, диаграммы
рядами. ComposedDeck и план ревизии добавляют то, чего в файле нет: слайд плана, слот объекта,
роль слайда. Связь идёт по именам фигур и положению, а не по номерам: ONLYOFFICE при сохранении
перенумеровывает `p:sldId` и `cNvPr id`, а имена сохраняет.

Без ComposedDeck роль объекта определяется по типу плейсхолдера и геометрии, признак элемента
шаблона — по номеру слайда, дате и колонтитулу, логотипу (хеш из профиля) и `fixed_elements`
профиля на своём месте.
"""

from __future__ import annotations

import hashlib
import io
import pathlib
import re
import tempfile
from typing import Any

from presentation_designer.generation.office_logo import logo_hashes
from presentation_designer.layout import background as slide_background
from presentation_designer.layout.composed import object_kind
from presentation_designer.layout.images import image_size
from presentation_designer.parsing.template.geometry import NS, ShapeInfo
from presentation_designer.parsing.template.package import TemplatePackage, open_template
from presentation_designer.parsing.template.styles import StyleResolver

JsonDict = dict[str, Any]

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"

PLACEHOLDER_ROLES = {
    "title": "title",
    "ctrTitle": "title",
    "subTitle": "subtitle",
    "body": "body",
    "obj": "body",
    "pic": "image",
    "chart": "chart",
    "tbl": "table",
    "dt": "date",
    "ftr": "footer",
    "sldNum": "page_number",
}
ROLES = {
    "title", "subtitle", "body", "bullets", "number", "label", "caption", "date", "name",
    "position", "image", "icon", "table", "chart", "diagram", "qr", "code", "footer", "logo",
    "page_number", "decoration", "background", "group", "other",
}  # fmt: skip
FIXED_ROLES = {"date", "footer", "page_number", "logo"}
# Виды содержательных объектов в оглавлении — по роли.
OUTLINE_KINDS = {
    "title": "text", "subtitle": "text", "body": "text", "bullets": "text", "number": "text",
    "label": "text", "caption": "text", "name": "text", "position": "text", "code": "text",
    "image": "picture", "icon": "picture", "qr": "picture", "table": "table", "chart": "chart",
    "diagram": "diagram", "group": "group",
}  # fmt: skip
KIND_ORDER = ["text", "picture", "table", "chart", "diagram", "group"]
_NUMBER = re.compile(
    r"^[\s+\-−–~≈<>]*\d[\d\s.,]*\s*(%|₽|\$|€|x|х|k|к|млн|млрд|тыс\.?)?[\s+]*$", re.I
)


def deck_snapshot(
    data: bytes,
    *,
    source: JsonDict | None = None,
    composed: JsonDict | None = None,
    plan: JsonDict | None = None,
    profile: JsonDict | None = None,
    base: JsonDict | None = None,
) -> JsonDict:
    """Снимок колоды по байтам PPTX. `composed` и `plan` — ComposedDeck и SlidePlan ревизии, из
    которой файл получен (для копии — ревизии-источника); `profile` — профиль шаблона для
    логотипа и фиксированных элементов; `base` — снимок исходной ревизии копии, с ним слайды
    получают признак ручных правок."""
    with tempfile.TemporaryDirectory() as directory:
        path = pathlib.Path(directory) / "deck.pptx"
        path.write_bytes(data)
        pkg = open_template(path)
        slides = _slides(pkg, profile or {})
    snapshot: JsonDict = {
        "schema_version": "1.0",
        "source": {**(source or {"kind": "file"}), "pptx_sha256": hashlib.sha256(data).hexdigest()},
        "slide_size": {
            "width_emu": pkg.width_emu,
            "height_emu": pkg.height_emu,
            "aspect_ratio": round(pkg.aspect_ratio, 4),
        },
        "slides": slides,
    }
    if composed:
        _enrich(slides, composed, plan or {})
    for slide in slides:
        _finish_slide(slide)
    if base is not None:
        _mark_edited(slides, base.get("slides") or [])
    snapshot["outline"] = [_outline_entry(slide) for slide in slides]
    return snapshot


def _slides(pkg: TemplatePackage, profile: JsonDict) -> list[JsonDict]:
    logos = {h.split(":", 1)[-1] for h in logo_hashes(profile)}
    fixed_boxes = [
        (str(f.get("kind")), f.get("bbox") or {})
        for f in profile.get("fixed_elements") or []
        if f.get("kind") in ("logo", "page_number", "footer", "navigation_dots")
    ]
    sld_ids = [int(item.id) for item in pkg.prs.slides._sldIdLst]
    out = []
    for slide_info, sld_id in zip(pkg.slides, sld_ids, strict=True):
        master = pkg.master(slide_info.master_id)
        layout = pkg.layout(slide_info.layout_id)
        resolver = StyleResolver(
            master.theme if master else None,
            master.element if master else None,
            layout.element if layout else None,
            master_part=master.part if master else "",
            layout_part=layout.part if layout else "",
        )
        part = slide_info.slide.part if slide_info.slide is not None else None
        objects = [_object(info, resolver, part, logos, fixed_boxes) for info in slide_info.shapes]
        back = slide_background.read(slide_info.element, part)
        background: JsonDict = {"kind": str(back.get("kind") or "inherited")}
        if back.get("color"):
            background["color"] = str(back["color"])
        if back.get("sha256"):
            background["image_sha256"] = str(back["sha256"])
        out.append(
            {
                "index": slide_info.index,
                "sld_id": sld_id,
                "layout": layout.name if layout else "",
                "title": "",
                "hidden": slide_info.hidden,
                "background": background,
                "notes": slide_info.notes_text,
                "objects": objects,
            }
        )
    return out


def _object(
    info: ShapeInfo,
    resolver: StyleResolver,
    part: Any,
    logos: set[str],
    fixed_boxes: list[tuple[str, JsonDict]],
) -> JsonDict:
    kind = object_kind(info)
    obj: JsonDict = {
        "address": {"object_id": info.element_id, "group_path": list(info.group_path)},
        "name": info.name,
        "kind": kind,
        "role": "other",
        "bbox": info.bbox,
        "fixed": False,
    }
    if abs(info.rotation_deg) > 0.01:
        obj["rotation_deg"] = round(info.rotation_deg, 2)
    if info.is_placeholder:
        own = info.element.find("p:spPr/a:xfrm", NS) if info.element is not None else None
        if own is None and info.element is not None:
            own = info.element.find("p:xfrm", NS)
        obj["placeholder"] = {
            "type": str(info.placeholder_type),
            **({"idx": info.placeholder_idx} if info.placeholder_idx is not None else {}),
            "inherited_geometry": own is None,
        }
    paragraphs = [p for p in info.paragraphs if p.text.strip()]
    if kind != "table" and info.has_text_frame and paragraphs:
        obj["text"] = {
            "plain": info.text,
            "paragraphs": [
                {
                    "text": p.text,
                    "level": p.level,
                    "bullet": p.bullet in ("char", "number", "picture"),
                }
                for p in paragraphs
            ],
        }
        style = _style(info, resolver)
        if style:
            obj["style"] = style
    if kind == "table" and info.element is not None:
        obj["table"] = _table(info.element)
    if kind == "chart" and info.element is not None and part is not None:
        chart = _chart(info.element, part)
        if chart:
            obj["chart"] = chart
    if kind == "picture":
        picture: JsonDict = {}
        if info.media_sha256:
            picture["sha256"] = info.media_sha256
        if info.media_blob:
            width, height = image_size(info.media_blob)
            if width and height:
                picture["width_px"], picture["height_px"] = width, height
        obj["picture"] = picture
    if info.is_hidden:
        obj["hidden"] = True
    obj["role"] = _role(info, kind, obj, logos)
    obj["fixed"] = obj["role"] in FIXED_ROLES or _in_fixed_place(info, kind, fixed_boxes)
    return obj


def _style(info: ShapeInfo, resolver: StyleResolver) -> JsonDict | None:
    for paragraph in info.paragraphs:
        run = next((r for r in paragraph.runs if str(r.get("text", "")).strip()), None)
        if run is None and not paragraph.text.strip():
            continue
        try:
            spec = resolver.resolve(info, paragraph, run).font_spec()
        except Exception:  # стиль — подсказка роутеру, его отсутствие снимок не роняет
            return None
        if not re.fullmatch(r"#[0-9A-Fa-f]{6}", str(spec.get("color") or "")):
            spec.pop("color", None)
        if not spec.get("family"):
            spec.pop("family", None)
        if not spec.get("size_pt") or spec["size_pt"] < 1:
            spec.pop("size_pt", None)
        return spec
    return None


def _role(info: ShapeInfo, kind: str, obj: JsonDict, logos: set[str]) -> str:
    """Роль по самому файлу: тип плейсхолдера, вид и геометрия объекта."""
    if info.placeholder_type in PLACEHOLDER_ROLES:
        role = PLACEHOLDER_ROLES[info.placeholder_type]
        if role == "body" and _bullets(obj):
            return "bullets"
        return role
    if kind in ("table", "chart"):
        return kind
    if kind == "group":
        return "group"
    if kind == "connector":
        return "decoration"
    if kind == "picture":
        if (obj.get("picture") or {}).get("sha256") in logos:
            return "logo"
        if info.area >= 0.85:
            return "background"
        return "icon" if info.area < 0.012 and 0.5 <= _aspect(info) <= 2 else "image"
    text = (obj.get("text") or {}).get("plain", "").strip()
    if not text:
        return "background" if info.area >= 0.85 else "decoration"
    size = (obj.get("style") or {}).get("size_pt") or 0
    if len(text) <= 16 and _NUMBER.match(text) and size >= 20:
        return "number"
    if _bullets(obj):
        return "bullets"
    return "body"


def _bullets(obj: JsonDict) -> bool:
    paragraphs = (obj.get("text") or {}).get("paragraphs") or []
    return sum(1 for p in paragraphs if p["bullet"]) >= 2


def _aspect(info: ShapeInfo) -> float:
    return info.width / info.height if info.height else 0.0


def _in_fixed_place(info: ShapeInfo, kind: str, boxes: list[tuple[str, JsonDict]]) -> bool:
    """Элемент шаблона из профиля (логотип, номер, колонтитул, навигация) на своём месте."""
    wanted = {
        "logo": ("picture",),
        "page_number": ("text", "placeholder_empty"),
        "footer": ("text", "placeholder_empty"),
        "navigation_dots": ("group", "shape"),
    }
    for fixed_kind, box in boxes:
        if kind not in wanted.get(fixed_kind, ()) or not box:
            continue
        if all(
            abs(float(box.get(key, -9)) - value) <= 0.01
            for key, value in (
                ("x", info.x),
                ("y", info.y),
                ("width", info.width),
                ("height", info.height),
            )
        ):
            return True
    return False


def _table(element: Any) -> JsonDict:
    table = element.find(f".//{A}tbl")
    rows: list[list[str]] = []
    merged: list[JsonDict] = []
    if table is None:
        return {"rows": rows}
    for r, row in enumerate(table.findall(f"{A}tr"), 1):
        cells = []
        for c, cell in enumerate(row.findall(f"{A}tc"), 1):
            if cell.get("hMerge") in ("1", "true") or cell.get("vMerge") in ("1", "true"):
                cells.append("")
                continue
            cells.append(
                "\n".join(
                    "".join(t.text or "" for t in p.iter(f"{A}t")) for p in cell.iter(f"{A}p")
                ).strip()
            )
            cols, span = int(cell.get("gridSpan", "1")), int(cell.get("rowSpan", "1"))
            if cols > 1 or span > 1:
                merged.append({"row": r, "col": c, "rows": span, "cols": cols})
        rows.append(cells)
    out: JsonDict = {"rows": rows}
    if merged:
        out["merged"] = merged
    return out


def _chart(element: Any, part: Any) -> JsonDict | None:
    """Тип, категории и ряды диаграммы: из кэша части chart, без кэша — из встроенной книги."""
    ref = element.find(f".//{C}chart")
    rid = ref.get(f"{R}id") if ref is not None else None
    if not rid:
        return None
    try:
        chart_part = part.related_part(rid)
        chart = chart_part.chart
    except (KeyError, AttributeError, ValueError):
        return None
    try:
        chart_type = str(chart.chart_type).split(".")[-1].split(" ")[0].lower()
    except (ValueError, AttributeError, KeyError, NotImplementedError):
        chart_type = "unknown"
    categories: list[str] = []
    series: list[JsonDict] = []
    try:
        for plot in chart.plots:
            if not categories:
                categories = [str(c) for c in plot.categories]
            for item in plot.series:
                values = [None if v is None else float(v) for v in item.values]
                series.append({"name": str(item.name or ""), "values": values})
    except (ValueError, AttributeError, KeyError, TypeError):
        pass
    out: JsonDict = {
        "type": chart_type,
        "categories": categories,
        "series": series,
        "values_from": "cache",
    }
    empty = not series or all(v is None for s in series for v in s["values"])
    if empty and _from_workbook(chart_part, out):
        out["values_from"] = "workbook"
    return out


def _from_workbook(chart_part: Any, out: JsonDict) -> bool:
    """Значения рядов по ссылкам `c:f` из встроенной книги, когда кэша в части нет."""
    try:
        from openpyxl import load_workbook

        blob = chart_part.chart_workbook.xlsx_part.blob
        book = load_workbook(io.BytesIO(blob), data_only=True, read_only=True)
    except Exception:
        return False
    try:
        root = chart_part._element
        series = []
        for ser in root.iter(f"{C}ser"):
            name_ref = ser.find(f"{C}tx/{C}strRef/{C}f")
            value_ref = ser.find(f"{C}val/{C}numRef/{C}f")
            values = _range(book, value_ref.text) if value_ref is not None else []
            names = _range(book, name_ref.text) if name_ref is not None else []
            series.append(
                {
                    "name": str(names[0]) if names and names[0] is not None else "",
                    "values": [float(v) if isinstance(v, int | float) else None for v in values],
                }
            )
        cat_ref = next(
            (
                f
                for f in root.iter(f"{C}f")
                if f.getparent() is not None
                and f.getparent().getparent() is not None
                and f.getparent().getparent().tag == f"{C}cat"
            ),
            None,
        )
        categories = _range(book, cat_ref.text) if cat_ref is not None and cat_ref.text else []
    except Exception:
        return False
    finally:
        book.close()
    if not any(v is not None for s in series for v in s["values"]):
        return False
    out["series"] = series
    if categories:
        out["categories"] = ["" if c is None else str(c) for c in categories]
    return True


def _range(book: Any, formula: str | None) -> list[Any]:
    if not formula or "!" not in formula:
        return []
    sheet, _, cells = formula.rpartition("!")
    sheet = sheet.strip("'")
    if sheet not in book.sheetnames:
        return []
    rows = book[sheet][cells.replace("$", "")]
    if not isinstance(rows, tuple):
        return [rows.value]
    flat: list[Any] = []
    for row in rows:
        if isinstance(row, tuple):
            flat.extend(cell.value for cell in row)
        else:
            flat.append(row.value)
    return flat


def _enrich(slides: list[JsonDict], composed: JsonDict, plan: JsonDict) -> None:
    """Слайд плана, слоты и роли из ComposedDeck ревизии: слайды и объекты сопоставляются по
    именам фигур и положению (номера в копии ONLYOFFICE меняет)."""
    plan_slides = {str(s.get("slide_id")): s for s in plan.get("slides") or []}
    pairs = _pair_slides(slides, composed.get("slides") or [])
    for slide, other in pairs:
        slide_id = str(other.get("slide_id") or "")
        if slide_id:
            slide["slide_id"] = slide_id
        role = (plan_slides.get(slide_id) or {}).get("role")
        if role:
            slide["role"] = str(role)
        if other.get("title") and not slide.get("title"):
            slide["title"] = str(other["title"])
        for obj, source in _pair_objects(slide["objects"], other.get("objects") or []):
            slot_id = source.get("slot_id")
            if slot_id:
                obj["slot"] = {
                    "slot_id": str(slot_id),
                    **({"slot_kind": str(source["slot_kind"])} if source.get("slot_kind") else {}),
                    **(
                        {"source_object_id": str(source["source_object_id"])}
                        if source.get("source_object_id")
                        else {}
                    ),
                }
                kind = str(source.get("block_kind") or source.get("slot_kind") or "")
                if kind in ("table", "chart", "diagram") or (
                    kind in ROLES and obj["role"] not in FIXED_ROLES
                ):
                    obj["role"] = kind
            elif source.get("role") == "background":
                obj["role"] = "background"
            if (source.get("diagram") or {}).get("kind") and obj["kind"] == "group":
                obj["role"] = "diagram"
            if (source.get("picture") or {}).get("asset_id") and "picture" in obj:
                obj["picture"]["asset_id"] = str(source["picture"]["asset_id"]).split(":", 1)[0]


def _names(slide: JsonDict) -> set[str]:
    return {str(o.get("name")) for o in slide.get("objects") or [] if o.get("name")}


def _pair_slides(slides: list[JsonDict], others: list[JsonDict]) -> list[tuple[JsonDict, JsonDict]]:
    """Пары «слайд снимка — слайд ComposedDeck»: больше общих имён фигур и текстов, при
    равенстве — тот же номер."""
    free = list(range(len(others)))
    pairs = []
    for position, slide in enumerate(slides):
        names = _names(slide)
        texts = {o["text"]["plain"] for o in slide["objects"] if o.get("text")}

        def score(
            j: int, names: set[str] = names, texts: set[str] = texts, position: int = position
        ) -> tuple[int, int]:
            other = others[j]
            other_texts = {
                str((o.get("text") or {}).get("plain"))
                for o in other.get("objects") or []
                if o.get("text")
            }
            return (
                len(names & _names(other)) + len(texts & other_texts),
                -abs(j - position),
            )

        if not free:
            break
        best = max(free, key=score)
        if score(best)[0] == 0 and names:
            continue
        free.remove(best)
        pairs.append((slide, others[best]))
    return pairs


def _pair_objects(
    objects: list[JsonDict], others: list[JsonDict]
) -> list[tuple[JsonDict, JsonDict]]:
    by_name: dict[str, list[JsonDict]] = {}
    for other in others:
        by_name.setdefault(str(other.get("name") or ""), []).append(other)
    pairs = []
    for obj in objects:
        candidates = by_name.get(obj["name"]) or []
        if not candidates:
            continue
        box = obj["bbox"]
        best = min(
            candidates,
            key=lambda o: sum(
                abs(float((o.get("bbox") or {}).get(k, 9)) - float(box[k]))
                for k in ("x", "y", "width", "height")
            ),
        )
        candidates.remove(best)
        pairs.append((obj, best))
    return pairs


def _finish_slide(slide: JsonDict) -> None:
    """Заголовок слайда и его роль, если их не дал план."""
    objects = slide["objects"]
    if not slide["title"]:
        titled = [o for o in objects if o["role"] == "title" and o.get("text")]
        if not titled:
            # Заголовок без плейсхолдера: самый крупный короткий текст в верхней трети.
            candidates = [
                o
                for o in objects
                if o.get("text")
                and o["role"] in ("body", "other")
                and not o["fixed"]
                and o["bbox"]["y"] < 0.35
                and len(o["text"]["plain"]) <= 120
                and (o.get("style") or {}).get("size_pt")
            ]
            if candidates:
                top = max(candidates, key=lambda o: (o["style"]["size_pt"], -o["bbox"]["y"]))
                sizes = sorted(
                    ((o.get("style") or {}).get("size_pt") or 0 for o in objects if o.get("text")),
                    reverse=True,
                )
                if len(sizes) == 1 or top["style"]["size_pt"] > sizes[1] or len(candidates) == 1:
                    top["role"] = "title"
                    titled = [top]
        if titled:
            slide["title"] = " ".join(titled[0]["text"]["plain"].split())[:200]
    if not slide.get("role"):
        types = {(o.get("placeholder") or {}).get("type") for o in objects}
        roles = {o["role"] for o in objects}
        if "ctrTitle" in types:
            slide["role"] = "title"
        elif "table" in roles:
            slide["role"] = "table"
        elif "chart" in roles:
            slide["role"] = "chart"
        else:
            slide["role"] = "freeform"


def _signature(slide: JsonDict) -> list[tuple[str, str, tuple[float, ...]]]:
    return sorted(
        (
            str(o.get("name")),
            str((o.get("text") or {}).get("plain", ""))
            + repr((o.get("table") or {}).get("rows"))
            + repr([s.get("values") for s in (o.get("chart") or {}).get("series") or []]),
            tuple(round(float(o["bbox"][k]), 3) for k in ("x", "y", "width", "height")),
        )
        for o in slide.get("objects") or []
    )


def _mark_edited(slides: list[JsonDict], base: list[JsonDict]) -> None:
    paired = {id(slide): other for slide, other in _pair_slides(slides, base)}
    for slide in slides:
        other = paired.get(id(slide))
        slide["edited"] = other is None or _signature(slide) != _signature(other)


def _outline_entry(slide: JsonDict) -> JsonDict:
    # Пустой плейсхолдер (подсказка макета) содержимым не считается.
    kinds = {
        OUTLINE_KINDS[o["role"]]
        for o in slide["objects"]
        if o["role"] in OUTLINE_KINDS
        and not o["fixed"]
        and (o.get("text") or o["kind"] in ("picture", "table", "chart", "group"))
    }
    entry: JsonDict = {
        "index": slide["index"],
        "sld_id": slide["sld_id"],
        "role": slide.get("role") or "freeform",
        "title": slide["title"],
        "kinds": [k for k in KIND_ORDER if k in kinds],
    }
    if slide.get("slide_id"):
        entry["slide_id"] = slide["slide_id"]
    if slide["hidden"]:
        entry["hidden"] = True
    if "edited" in slide:
        entry["edited"] = slide["edited"]
    return entry


KIND_NAMES = {
    "text": "текст",
    "picture": "картинки",
    "table": "таблица",
    "chart": "диаграмма",
    "diagram": "схема",
    "group": "группы фигур",
}


def outline_lines(snapshot: JsonDict) -> list[str]:
    """Оглавление строками для модели: номер, заголовок, что на слайде."""
    lines = []
    for entry in snapshot.get("outline") or []:
        kinds = ", ".join(KIND_NAMES[k] for k in entry.get("kinds") or [])
        title = entry.get("title") or "без заголовка"
        marks = " (скрыт)" if entry.get("hidden") else ""
        marks += " (правлен вручную)" if entry.get("edited") else ""
        lines.append(f"{entry['index']}. «{title}»{marks}" + (f" — {kinds}" if kinds else ""))
    return lines


def slide_lines(slide: JsonDict, limit: int = 2000) -> list[str]:
    """Содержимое слайда строками для модели: тексты, таблицы, диаграммы, картинки; без
    номера слайда, колонтитула и логотипа."""
    lines: list[str] = []
    for obj in slide.get("objects") or []:
        if obj["fixed"] or obj["role"] in ("decoration", "background"):
            continue
        role = obj["role"]
        if obj.get("text"):
            text = " / ".join(p["text"].strip() for p in obj["text"]["paragraphs"] if p["text"])
            lines.append(f"[{role}] {text}")
        elif obj.get("table"):
            rows = obj["table"]["rows"]
            lines.append(f"[таблица {len(rows)}×{max((len(r) for r in rows), default=0)}]")
            lines.extend(" | ".join(row) for row in rows[:12])
        elif obj.get("chart"):
            chart = obj["chart"]
            lines.append(f"[диаграмма {chart['type']}] категории: {', '.join(chart['categories'])}")
            for series in chart["series"][:6]:
                values = ", ".join("—" if v is None else f"{v:g}" for v in series["values"])
                lines.append(f"  {series['name'] or 'ряд'}: {values}")
        elif role in ("image", "icon"):
            lines.append(f"[{'картинка' if role == 'image' else 'иконка'}]")
    if slide.get("notes"):
        lines.append(f"[заметки] {slide['notes']}")
    out, total = [], 0
    for line in lines:
        total += len(line)
        if total > limit:
            out.append("…")
            break
        out.append(line)
    return out


__all__ = ["deck_snapshot", "outline_lines", "slide_lines"]
