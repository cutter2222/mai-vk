"""Диаграммы в офисной копии (этап 40): данные, тип, оформление и смена подачи на месте.

У диаграммы данные лежат в своей части `chart*.xml` и во встроенной книге, поэтому правка идёт
через python-pptx (`replace_data` переписывает и кэш, и книгу), а не заменой XML слайда.
Стиль — тот же, что у диаграмм сборки (`layout/charts.ChartStyle`): из профиля шаблона копии,
без профиля — из темы и шрифтов самого файла. Новые числа берутся только из просьбы или из
самой диаграммы: пересчёт (проценты, суммы) модель не делает — это проверяется до записи.

Отсюда же — новая диаграмма в месте, названном словами: из приложенного xlsx/csv и из
картинки графика (чтение `chart_reader`, этапы 26–28), и «сделай диаграмму редактируемой»
для составных диаграмм из фигур (этап 34).
"""

from __future__ import annotations

import io
import re
from typing import Any

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.util import Pt

from presentation_designer.generation.office_ops import CHART_OPS, ObjectOp, Tokens
from presentation_designer.generation.office_table import center_cells, compact_box
from presentation_designer.layout.charts import (
    DEFAULT_ACCENTS,
    FAMILY,
    ChartSpec,
    ChartStyle,
    add_chart,
    luminance,
)
from presentation_designer.layout.tables import TableSpec, TableStyle, add_table, format_number

MUTED = "#BFBFBF"
# Больше рядов аудит считает перегрузкой (`density.chart_series`).
MAX_SERIES = 5
KIND_NAMES = {
    "column": "столбцы", "bar": "полосы", "stacked_column": "столбцы с накоплением",
    "line": "линия", "area": "области", "pie": "круговая", "doughnut": "кольцо",
}  # fmt: skip
LEGEND = {
    "below": XL_LEGEND_POSITION.BOTTOM,
    "above": XL_LEGEND_POSITION.TOP,
    "right": XL_LEGEND_POSITION.RIGHT,
    "left": XL_LEGEND_POSITION.LEFT,
}
_XL_KIND: dict[Any, str] = {
    XL_CHART_TYPE.COLUMN_CLUSTERED: "column",
    XL_CHART_TYPE.COLUMN_STACKED: "stacked_column",
    XL_CHART_TYPE.COLUMN_STACKED_100: "stacked_column",
    XL_CHART_TYPE.BAR_CLUSTERED: "bar",
    XL_CHART_TYPE.BAR_STACKED: "bar",
    XL_CHART_TYPE.BAR_STACKED_100: "bar",
    XL_CHART_TYPE.LINE: "line",
    XL_CHART_TYPE.LINE_MARKERS: "line",
    XL_CHART_TYPE.LINE_STACKED: "line",
    XL_CHART_TYPE.LINE_MARKERS_STACKED: "line",
    XL_CHART_TYPE.AREA: "area",
    XL_CHART_TYPE.AREA_STACKED: "area",
    XL_CHART_TYPE.PIE: "pie",
    XL_CHART_TYPE.PIE_EXPLODED: "pie",
    XL_CHART_TYPE.DOUGHNUT: "doughnut",
    XL_CHART_TYPE.DOUGHNUT_EXPLODED: "doughnut",
}
# Тип диаграммы словами просьбы: «кольцом», «круговую», «линией», «горизонтально».
KIND_WORDS = (
    (re.compile(r"кольц", re.I), "doughnut"),
    (re.compile(r"кругов|пирог", re.I), "pie"),
    (re.compile(r"накоплен", re.I), "stacked_column"),
    (re.compile(r"горизонтал|полос", re.I), "bar"),
    (re.compile(r"област", re.I), "area"),
    (re.compile(r"линейн|линией|линии|линия|линию", re.I), "line"),
    (re.compile(r"столб", re.I), "column"),
)
_NUMBER = re.compile(r"[-−]?\d[\d\s  ]*(?:[.,]\d+)?")


# --- стиль -------------------------------------------------------------------------------------


def _vivid(color: str) -> bool:
    """Цвет годится для ряда: не почти белый, не почти чёрный и не серый."""
    r, g, b = (int(color[i : i + 2], 16) for i in (1, 3, 5))
    return 0.05 < luminance(color) < 0.85 and max(r, g, b) - min(r, g, b) > 40


def chart_style(tokens: Tokens) -> ChartStyle:
    """Стиль диаграмм копии: профиль шаблона, без него — палитра и шрифты файла."""
    if tokens.profile:
        return ChartStyle.from_profile(tokens.profile)
    accents = [c for c in tokens.palette if _vivid(c)]
    size = min(tokens.body_pt or 12.0, 14.0)
    return ChartStyle(
        font_family=tokens.fonts[0] if tokens.fonts else None,
        font_size_pt=size,
        title_size_pt=min(size + 2, 16.0),
        text_color=tokens.text_color or "#000000",
        accents=accents or list(DEFAULT_ACCENTS),
    )


def table_style(tokens: Tokens) -> TableStyle:
    """Стиль таблиц копии: профиль шаблона, без него — первый яркий цвет палитры в шапке."""
    if tokens.profile:
        return TableStyle.from_profile(tokens.profile)
    return TableStyle(
        font_family=tokens.fonts[0] if tokens.fonts else None,
        font_size_pt=12.0,
        header_fill=next((c for c in tokens.palette if _vivid(c)), "#0077FF"),
        text_color=tokens.text_color or "#000000",
    )


def _rgb(color: str) -> Any:
    return RGBColor.from_string(color.lstrip("#").upper())  # type: ignore[no-untyped-call]


def _font(font: Any, style: ChartStyle, size: float | None = None) -> None:
    font.size = Pt(size or style.font_size_pt)
    font.color.rgb = _rgb(style.text_color)
    if style.font_family:
        font.name = style.font_family


def _dark(chart: Any) -> bool:
    """Диаграмма стоит на тёмном фоне: её текст светлый (так её оформила сборка или автор)."""
    for node in chart._chartSpace.iter(
        "{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr"
    ):
        parent = node.getparent()
        if parent is not None and parent.getparent() is not None:
            owner = parent.getparent().tag.rsplit("}", 1)[-1]
            if owner in ("defRPr", "rPr"):
                return luminance("#" + str(node.get("val"))) >= 0.7
    return False


# --- чтение ------------------------------------------------------------------------------------


def _shape(slide: Any, shape_id: str) -> Any:
    def walk(shapes: Any) -> Any:
        for shape in shapes:
            if str(shape.shape_id) == str(shape_id):
                return shape
            if shape.shape_type is not None and hasattr(shape, "shapes"):
                found = walk(shape.shapes)
                if found is not None:
                    return found
        return None

    return walk(slide.shapes)


def kind_of(chart: Any) -> str:
    try:
        return _XL_KIND.get(chart.chart_type, "column")
    except (ValueError, AttributeError, KeyError, NotImplementedError):
        return "column"


def read_data(chart: Any) -> tuple[list[str], list[tuple[str, list[float | None]]]]:
    """Категории и ряды диаграммы (из кэша части chart)."""
    categories: list[str] = []
    series: list[tuple[str, list[float | None]]] = []
    for plot in chart.plots:
        if not categories:
            categories = [str(c) for c in plot.categories]
        for item in plot.series:
            values = [None if v is None else float(v) for v in item.values]
            series.append((str(item.name or f"Ряд {len(series) + 1}"), values))
    return categories, series


def _fractional(series: list[tuple[str, list[float | None]]]) -> bool:
    return any(v is not None and abs(v - round(v)) > 1e-9 for _, values in series for v in values)


def _spec_of(chart: Any, kind: str) -> ChartSpec:
    """Спецификация по действующей диаграмме: данные и включённые подписи, легенда, заголовок."""
    categories, series = read_data(chart)
    title = None
    if chart.has_title:
        title = chart.chart_title.text_frame.text.strip() or None
    return ChartSpec(
        chart_type=kind,
        categories=categories,
        series=series,
        title=title,
        units=None,
        show_legend=bool(chart.has_legend),
        show_axis_labels=True,
        show_data_labels=any(p.has_data_labels for p in chart.plots),
        number_format="#,##0.0" if _fractional(series) else "#,##0",
        dataset_id="chat",
    )


def _data(categories: list[str], series: list[tuple[str, list[float | None]]]) -> Any:
    data = CategoryChartData()  # type: ignore[no-untyped-call]
    data.categories = categories
    for name, values in series:
        fractional = any(v is not None and abs(v - round(v)) > 1e-9 for v in values)
        data.add_series(  # type: ignore[no-untyped-call]
            name, values, number_format="#,##0.0" if fractional else "#,##0"
        )
    return data


# --- проверка чисел ----------------------------------------------------------------------------


def numbers_in(text: str) -> set[float]:
    out: set[float] = set()
    for match in _NUMBER.findall(text):
        raw = "".join(match.split()).replace("−", "-").replace(",", ".")
        try:
            out.add(round(float(raw), 6))
        except ValueError:
            continue
    return out


def check_numbers(instruction: str, before: list[float | None], ops: list[ObjectOp]) -> None:
    """Новые значения диаграммы — только из просьбы или из прежних данных: пересчёт не делаем."""
    allowed = numbers_in(instruction) | {round(v, 6) for v in before if v is not None}
    for op in ops:
        if op.op != "chart.set_data":
            continue
        for s in op.series or []:
            for v in s.values:
                if v is not None and round(v, 6) not in allowed:
                    raise ValueError(
                        f"значения {v:g} нет ни в просьбе, ни в диаграмме — не пересчитывай "
                        "и не выдумывай числа, бери их из просьбы"
                    )


# --- операции ----------------------------------------------------------------------------------


def _box(frame: Any) -> tuple[int, int, int, int]:
    """Рамка объекта на слайде. У таблицы настоящий размер — сумма ширин столбцов и высот
    строк: рамка graphicFrame бывает меньше (таблицы из Google Slides)."""
    width, height = int(frame.width), int(frame.height)
    if getattr(frame, "has_table", False):
        width = max(width, sum(int(c.width) for c in frame.table.columns))
        height = max(height, sum(int(r.height) for r in frame.table.rows))
    return int(frame.left), int(frame.top), width, height


def _replace_frame(slide: Any, old: Any, spec: ChartSpec, style: ChartStyle) -> Any:
    """Новая диаграмма на месте объекта `old` (та же рамка, имя и порядок наложения)."""
    if old._element.getparent() is not slide.shapes._spTree:
        raise ValueError("объект внутри группы — сменить подачу можно в редакторе")
    frame = add_chart(slide, _box(old), spec, style)
    _swap(slide, old, frame._element)
    return frame


def _few_points(categories: list[str], series: list[Any]) -> bool:
    """Подписи значений читаются, когда точек немного: иначе они налезают друг на друга."""
    return len(categories) * max(len(series), 1) <= 12


def _label_format(chart: Any) -> str | None:
    """Формат подписей значений, заданный у диаграммы (например «0%»), — сохраняется при смене
    типа."""
    for plot in chart.plots:
        if plot.has_data_labels and not plot.data_labels.number_format_is_linked:
            return str(plot.data_labels.number_format)
    return None


def _texts_over(slide: Any, frame: Any) -> list[str]:
    """Надписи слайда, лежащие на середине рамки (подпись в центре кольца)."""
    x, y, w, h = _box(frame)
    out = []
    for shape in slide.shapes:
        if shape is frame or not getattr(shape, "has_text_frame", False):
            continue
        text = shape.text_frame.text.strip()
        if not text or shape.width is None:
            continue
        cx, cy = shape.left + shape.width / 2, shape.top + shape.height / 2
        if abs(cx - (x + w / 2)) < w * 0.2 and abs(cy - (y + h / 2)) < h * 0.2:
            out.append(text.replace("\n", " ")[:40])
    return out


def _swap(slide: Any, old: Any, new_element: Any) -> None:
    old_element = old._element
    name = old.name
    old_element.addprevious(new_element)
    rids = [str(r) for r in old_element.xpath(".//@r:id | .//@r:embed")]
    old_element.getparent().remove(old_element)
    new_element.xpath("./*[1]/p:cNvPr")[0].set("name", name)
    refs = {str(r) for r in slide._element.xpath("//@r:id | //@r:embed | //@r:link")}
    for rid in rids:
        if rid not in refs and rid in slide.part.rels:
            slide.part.drop_rel(rid)


def _color_points(chart: Any, kind: str, style: ChartStyle) -> None:
    """Все ряды (у круга — сектора) заново цветами шаблона."""
    accents = style.accents or list(DEFAULT_ACCENTS)
    for plot in chart.plots:
        for i, series in enumerate(plot.series):
            if FAMILY.get(kind) == "pie":
                for k in range(len(list(series.values))):
                    point = series.points[k]
                    point.format.fill.solid()
                    point.format.fill.fore_color.rgb = _rgb(accents[k % len(accents)])
                continue
            color = _rgb(accents[i % len(accents)])
            if kind == "line":
                series.format.line.color.rgb = color
                series.marker.format.fill.solid()
                series.marker.format.fill.fore_color.rgb = color
            else:
                series.format.fill.solid()
                series.format.fill.fore_color.rgb = color


def _index(names: list[str], wanted: str) -> int | None:
    key = wanted.strip().lower().replace("ё", "е")
    for i, name in enumerate(names):
        if name.strip().lower().replace("ё", "е") == key:
            return i
    for i, name in enumerate(names):
        if key and key in name.lower().replace("ё", "е"):
            return i
    if key.isdigit() and 1 <= int(key) <= len(names):
        return int(key) - 1
    return None


def _highlight(chart: Any, kind: str, category: str, style: ChartStyle) -> str:
    """Одна категория (или ряд) — акцентом, остальные приглушены (как выделение в этапе 27)."""
    categories, series = read_data(chart)
    accent = (style.accents or list(DEFAULT_ACCENTS))[0]
    names = [n for n, _ in series]
    plots = list(chart.plots)
    if len(series) > 1 and FAMILY.get(kind) != "pie":
        index = _index(names, category)
        if index is not None:
            k = 0
            for plot in plots:
                for s in plot.series:
                    color = _rgb(accent if k == index else MUTED)
                    if kind == "line":
                        s.format.line.color.rgb = color
                    else:
                        s.format.fill.solid()
                        s.format.fill.fore_color.rgb = color
                    k += 1
            return f"выделен ряд «{names[index]}»"
    index = _index(categories, category)
    if index is None:
        raise ValueError(f"категории «{category}» в диаграмме нет: {', '.join(categories[:12])}")
    for plot in plots:
        for s in plot.series:
            for k in range(len(categories)):
                point = s.points[k]
                color = _rgb(accent if k == index else MUTED)
                if kind == "line":
                    point.marker.format.fill.solid()
                    point.marker.format.fill.fore_color.rgb = color
                else:
                    point.format.fill.solid()
                    point.format.fill.fore_color.rgb = color
    return f"выделено «{categories[index]}»"


def _labels(chart: Any, kind: str, on: bool, style: ChartStyle, fmt: str | None = None) -> None:
    if on and fmt is None and _label_format(chart) is None:
        # Без заданного формата подписи идут с тремя знаками («16.675»): округляем по данным.
        fmt = "#,##0.0" if _fractional(read_data(chart)[1]) else "#,##0"
    for plot in chart.plots:
        plot.has_data_labels = on
        if not on:
            continue
        labels = plot.data_labels
        _font(labels.font, style)
        if fmt:
            labels.number_format = fmt
            labels.number_format_is_linked = False
        if FAMILY.get(kind) == "pie":
            labels.show_value = True
            labels.show_percentage = False
        elif kind in ("column", "bar"):
            labels.position = XL_LABEL_POSITION.OUTSIDE_END
        elif kind == "line":
            labels.position = XL_LABEL_POSITION.ABOVE


def _table_data(frame: Any) -> tuple[list[str], list[tuple[str, list[float | None]]]]:
    """Таблица → категории (первый столбец) и ряды (числовые столбцы, имя — шапка)."""
    from presentation_designer.layout.charts import _to_number

    table = frame.table
    rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
    if len(rows) < 2 or len(rows[0]) < 2:
        raise ValueError("в таблице нет данных для диаграммы: нужна шапка и строки")
    header, body = rows[0], rows[1:]
    series: list[tuple[str, list[float | None]]] = []
    for c in range(1, len(header)):
        values = [_to_number(r[c]) if c < len(r) else None for r in body]
        if any(v is not None for v in values):
            series.append((header[c] or f"Ряд {c}", values))
    if not series:
        raise ValueError("в таблице нет числовых столбцов — диаграмму строить не из чего")
    categories = [r[0] for r in body]
    if len(series) > MAX_SERIES and len(categories) < len(series):
        # Широкая таблица (кварталы по столбцам): рядами становятся строки — их меньше.
        names = [name for name, _ in series]
        series = [(c, [values[k] for _, values in series]) for k, c in enumerate(categories)]
        categories = names
    return categories, series


def apply_chart_ops(
    data: bytes, slide_number: int, shape_id: str, ops: list[ObjectOp], tokens: Tokens
) -> tuple[bytes, list[str]]:
    """PPTX с операциями над диаграммой (или таблицей для `object.to_chart`) и заметки."""
    ops = [op for op in ops if op.op in CHART_OPS]
    if not ops:
        return data, []
    prs = Presentation(io.BytesIO(data))
    if not 1 <= slide_number <= len(prs.slides):
        raise ValueError(f"слайда {slide_number} нет")
    slide = prs.slides[slide_number - 1]
    frame = _shape(slide, shape_id)
    if frame is None:
        raise ValueError("объект не найден на слайде")
    style = chart_style(tokens)
    notes: list[str] = []
    for op in ops:
        if op.op == "object.to_chart":
            if not getattr(frame, "has_table", False):
                raise ValueError("в диаграмму превращается таблица, а это не таблица")
            categories, series = _table_data(frame)
            kind: str = op.chart_type or "column"
            limit = 1 if FAMILY[kind] == "pie" else MAX_SERIES
            if len(series) > limit:
                shown = ", ".join(f"«{n}»" for n, _ in series[:limit])
                notes.append(
                    f"в диаграмме {limit} из {len(series)} рядов: {shown} — больше не читается"
                )
                series = series[:limit]
            spec = ChartSpec(
                chart_type=kind, categories=categories, series=series, title=None, units=None,
                show_legend=len(series) > 1 or FAMILY[kind] == "pie", show_axis_labels=True,
                show_data_labels=_few_points(categories, series),
                number_format="#,##0.0" if _fractional(series) else "#,##0", dataset_id="chat",
            )  # fmt: skip
            frame = _replace_frame(slide, frame, spec, style)
            same = "" if notes and notes[-1].startswith("в диаграмме ") else ", данные те же"
            notes.append(f"таблица показана диаграммой ({KIND_NAMES[kind]}){same}")
            continue
        if not getattr(frame, "has_chart", False):
            raise ValueError("это не диаграмма")
        chart = frame.chart
        kind = kind_of(chart)
        if _dark(chart) and not style.dark:
            style = style.on_dark()
        if op.op == "object.to_table":
            categories, series = read_data(chart)
            spec_t = TableSpec(
                header=[""] + [n for n, _ in series],
                rows=[
                    [c] + ["" if s[k] is None else format_number(s[k]) for _, s in series]
                    for k, c in enumerate(categories)
                ],
                numeric_columns=[False] + [True] * len(series),
                dataset_id="chat", row_offset=0, truncated=False, highlight_row=None,
            )  # fmt: skip
            if frame._element.getparent() is not slide.shapes._spTree:
                raise ValueError("объект внутри группы — сменить подачу можно в редакторе")
            t_style = table_style(tokens)
            box = compact_box(_box(frame), len(categories) + 1, t_style)
            table_frame = add_table(slide, box, spec_t, t_style)
            center_cells(table_frame)
            _swap(slide, frame, table_frame._element)
            notes.append("диаграмма показана таблицей, данные те же")
            break
        if op.op == "chart.set_data":
            categories = [c.strip() for c in op.categories or []]
            series = [(s.name.strip(), list(s.values)) for s in op.series or []]
            if not categories or not series:
                raise ValueError("для новых данных нужны категории и хотя бы один ряд")
            for name, values in series:
                if len(values) != len(categories):
                    raise ValueError(
                        f"в ряду «{name}» {len(values)} значений, а категорий {len(categories)}"
                    )
            if FAMILY.get(kind) == "pie" and len(series) > 1:
                series = series[:1]
                notes.append(f"у круговой диаграммы один ряд — показан «{series[0][0]}»")
            old_categories, old_series = read_data(chart)
            chart.replace_data(_data(categories, series))
            if len(series) > len(old_series) or (
                FAMILY.get(kind) == "pie" and len(categories) > len(old_categories)
            ):
                _color_points(chart, kind, style)
            notes.append("данные диаграммы обновлены, числа — по вашему сообщению")
        elif op.op == "chart.type":
            if not op.chart_type:
                raise ValueError("не указан тип диаграммы")
            spec = _spec_of(chart, op.chart_type)
            if FAMILY[op.chart_type] == "pie" and len(spec.series) > 1:
                notes.append(f"у круговой диаграммы один ряд — показан «{spec.series[0][0]}»")
                spec.series = spec.series[:1]
                spec.show_legend = True
            if FAMILY[op.chart_type] == "pie" and any(
                v is not None and v < 0 for _, values in spec.series for v in values
            ):
                raise ValueError("в данных есть отрицательные значения — круговой их не показать")
            label_format = _label_format(chart)
            covered = _texts_over(slide, frame) if kind == "doughnut" else []
            frame = _replace_frame(slide, frame, spec, style)
            if label_format and spec.show_data_labels:
                _labels(frame.chart, op.chart_type, True, style, label_format)
            notes.append(f"тип диаграммы — {KIND_NAMES[op.chart_type]}")
            if covered and op.chart_type != "doughnut":
                notes.append(
                    f"надпись «{covered[0]}» была в центре кольца и теперь лежит на диаграмме — "
                    "передвиньте или удалите её"
                )
        elif op.op == "chart.legend":
            if op.on is False and not chart.has_legend:
                notes.append(
                    "у самой диаграммы легенды нет — подписи рядов на слайде отдельные надписи: "
                    "выделите их и напишите «удали»"
                )
            chart.has_legend = op.on is not False
            if chart.has_legend:
                chart.legend.position = LEGEND[op.place or "below"]
                chart.legend.include_in_layout = False
                _font(chart.legend.font, style)
        elif op.op == "chart.labels":
            _labels(chart, kind, op.on is not False, style, op.number_format)
        elif op.op == "chart.gridlines":
            try:
                axis = chart.value_axis
            except (ValueError, AttributeError):
                raise ValueError("у этой диаграммы нет оси значений и сетки") from None
            axis.has_major_gridlines = op.on is not False
            if axis.has_major_gridlines:
                axis.major_gridlines.format.line.color.rgb = _rgb(
                    "#5A5A5A" if style.dark else "#E0E0E0"
                )
        elif op.op == "chart.title":
            text = (op.text or "").strip()
            chart.has_title = bool(text)
            if text:
                chart.chart_title.text_frame.text = text
                for p in chart.chart_title.text_frame.paragraphs:
                    _font(p.font, style, style.title_size_pt)
        elif op.op == "chart.highlight":
            notes.append(_highlight(chart, kind, op.category or "", style))
        elif op.op == "chart.colors":
            _color_points(chart, kind, style)
            notes.append("цвета — из палитры шаблона")
        elif op.op == "chart.number_format":
            if not op.number_format:
                raise ValueError("не указан формат чисел")
            _labels(chart, kind, True, style, op.number_format)
            try:
                chart.value_axis.tick_labels.number_format = op.number_format
                chart.value_axis.tick_labels.number_format_is_linked = False
            except (ValueError, AttributeError):
                pass
            if "%" in op.number_format:
                notes.append("формат процентов не пересчитывает значения: 0,25 станет 25 %")
        elif op.op == "chart.sort":
            categories, series = read_data(chart)
            if not series:
                raise ValueError("в диаграмме нет данных")
            first = [float("-inf") if v is None else v for v in series[0][1]]
            order = sorted(
                range(len(categories)), key=lambda k: first[k], reverse=op.descending is not False
            )
            chart.replace_data(
                _data(
                    [categories[k] for k in order],
                    [(n, [v[k] for k in order]) for n, v in series],
                )
            )
            if FAMILY.get(kind) == "pie":
                _color_points(chart, kind, style)
            notes.append("категории отсортированы " + ("по убыванию" if op.descending is not False
                                                       else "по возрастанию"))  # fmt: skip
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue(), notes


def dark_area(snapshot: bytes, box: tuple[float, float, float, float]) -> bool:
    """Место на снимке слайда тёмное: подписи диаграммы там нужны светлые."""
    from PIL import Image, ImageStat

    image = Image.open(io.BytesIO(snapshot)).convert("L")
    w, h = image.size
    x, y, bw, bh = box
    crop = image.crop(
        (int(x * w), int(y * h), max(int((x + bw) * w), 1), max(int((y + bh) * h), 1))
    )
    return ImageStat.Stat(crop).mean[0] < 100


def kind_in(text: str) -> str | None:
    return next((kind for rx, kind in KIND_WORDS if rx.search(text)), None)


# --- новая диаграмма в месте -------------------------------------------------------------------


def spec_from_table(header: list[str], rows: list[list[str]], kind: str | None = None) -> ChartSpec:
    """Таблица из файла → диаграмма: первый столбец — категории, числовые — ряды."""
    from presentation_designer.layout.charts import _to_number

    series: list[tuple[str, list[float | None]]] = []
    for c in range(1, len(header)):
        values = [_to_number(r[c]) if c < len(r) else None for r in rows]
        if any(v is not None for v in values):
            series.append((header[c] or f"Ряд {c}", values))
    if not series:
        raise ValueError("в файле нет числовых столбцов — диаграмму строить не из чего")
    categories = [r[0] for r in rows]
    kind = kind or ("line" if len(categories) > 8 else "column")
    if FAMILY[kind] == "pie":
        series = series[:1]
    return ChartSpec(
        chart_type=kind, categories=categories, series=series[:MAX_SERIES], title=None, units=None,
        show_legend=len(series) > 1 or FAMILY[kind] == "pie", show_axis_labels=True,
        show_data_labels=len(categories) <= 8,
        number_format="#,##0.0" if _fractional(series) else "#,##0", dataset_id="chat",
    )  # fmt: skip


def spec_from_reading(reading: Any) -> ChartSpec:
    """Чтение картинки графика → диаграмма в стиле шаблона (вид и данные с картинки)."""
    kind = reading.kind if reading.kind in ("pie", "doughnut") else None
    if kind is None:
        base = next((s.type or "column" for s in reading.series if s.type != "line"), "line")
        kind = {"column": "column", "bar": "bar", "area": "area", "line": "line"}.get(
            base, "column"
        )
        if kind == "column" and reading.stacked:
            kind = "stacked_column"
    series = [
        (s.name or f"Ряд {i + 1}", [p.value for p in s.points])
        for i, s in enumerate(reading.series)
    ]
    if FAMILY[kind] == "pie":
        series = series[:1]
    return ChartSpec(
        chart_type=kind,
        categories=list(reading.categories),
        series=series,
        title=reading.title,
        units=reading.unit,
        show_legend=len(series) > 1 or FAMILY[kind] == "pie",
        show_axis_labels=True,
        show_data_labels=bool(reading.data_labels) or len(reading.categories) <= 8,
        number_format="#,##0.0" if _fractional(series) else "#,##0",
        dataset_id="chat",
    )


def place_chart(
    data: bytes, slide: int, spec: ChartSpec, box: tuple[float, float, float, float],
    tokens: Tokens, *, name: str = "Диаграмма из чата", dark: bool = False,
) -> bytes:  # fmt: skip
    """PPTX с новой диаграммой на слайде `slide` в рамке `box` (доли слайда)."""
    prs = Presentation(io.BytesIO(data))
    if not 1 <= slide <= len(prs.slides):
        raise ValueError(f"слайда {slide} нет")
    width, height = int(prs.slide_width or 0), int(prs.slide_height or 0)
    style = chart_style(tokens)
    if dark:
        style = style.on_dark()
    x, y, w, h = box
    frame = add_chart(
        prs.slides[slide - 1],
        (round(x * width), round(y * height), round(w * width), round(h * height)),
        spec,
        style,
    )
    frame.name = name
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def make_editable(data: bytes, slide: int) -> tuple[bytes, list[str]]:
    """Составные диаграммы из фигур на слайде → нативные (этап 34) по просьбе из чата."""
    from presentation_designer.layout.composite_charts import swap_composites

    prs = Presentation(io.BytesIO(data))
    if not 1 <= slide <= len(prs.slides):
        raise ValueError(f"слайда {slide} нет")
    swaps = swap_composites(prs, slides=[prs.slides[slide - 1]])
    done = [s for s in swaps if s.status == "replaced"]
    if not done:
        return data, []
    out = io.BytesIO()
    prs.save(out)
    if len(done) == 1:
        return out.getvalue(), ["диаграмма из фигур стала редактируемой"]
    return out.getvalue(), [f"диаграммы из фигур стали редактируемыми: {len(done)}"]


NOT_READ = (
    "не разобрал значения на картинке — напишите их, например: «график: 2023 — 75, "
    "2024 — 120, 2025 — 140»"
)


async def read_picture(blob: bytes, settings: Any) -> Any:
    """Чтение картинки графика моделью vlm (этапы 26–28): ChartReading или ValueError."""
    from presentation_designer.llm.client import build_client
    from presentation_designer.llm.types import Deadline
    from presentation_designer.parsing.raster_charts.rebuild import read_chart_image

    client = build_client(settings)
    try:
        outcome = await read_chart_image(blob, client, deadline=Deadline.after(150))
    finally:
        await client.aclose()
    if outcome.status == "read" and outcome.reading is not None:
        return outcome.reading
    if outcome.status in ("not_chart", "skipped"):
        raise ValueError("на картинке не видно диаграммы — пришлите график или числа текстом")
    raise ValueError(NOT_READ)


async def make_editable_all(data: bytes, slide: int, settings: Any) -> tuple[bytes, list[str]]:
    """«Сделай диаграмму редактируемой»: сначала диаграммы из фигур (без модели), затем
    картинки-графики слайда (чтение моделью; не больше трёх)."""
    from presentation_designer.layout.chart_images import (
        describe,
        picture_sha,
        slide_pictures,
        swap_pictures,
    )
    from presentation_designer.parsing.raster_charts.pixels import chart_likeness, load

    data, notes = make_editable(data, slide)
    prs = Presentation(io.BytesIO(data))
    page = prs.slides[slide - 1]
    blobs: dict[str, bytes] = {}
    for picture in slide_pictures(page):
        try:
            blob = picture.image.blob
            if chart_likeness(load(blob)).ok:
                blobs.setdefault(picture_sha(picture), blob)
        except Exception:  # связанная или битая картинка остаётся как есть
            continue
    readings: dict[str, Any] = {}
    failed = 0
    for sha, blob in list(blobs.items())[:3]:
        try:
            readings[sha] = await read_picture(blob, settings)
        except ValueError:
            failed += 1
    swaps = swap_pictures(prs, readings, slides=[page]) if readings else []
    done = [s for s in swaps if s.status == "replaced"]
    if done:
        out = io.BytesIO()
        prs.save(out)
        data = out.getvalue()
        notes.append(describe(next(iter(readings.values()))))
    kept = [s.reason for s in swaps if s.status == "kept"]
    if kept:
        notes.append("картинку оставил: " + kept[0])
    if failed:
        notes.append(NOT_READ)
    if not notes:
        if any(getattr(shape, "has_chart", False) for shape in page.shapes):
            raise ValueError(
                "диаграммы на этом слайде уже редактируемые — выделите диаграмму и напишите, "
                "что в ней поменять"
            )
        raise ValueError(
            "на слайде нет диаграмм из фигур или картинок-графиков — редактируемым делать нечего"
        )
    return data, notes


__all__ = [
    "FAMILY", "KIND_NAMES", "apply_chart_ops", "chart_style", "check_numbers", "dark_area",
    "kind_in", "kind_of", "make_editable", "make_editable_all", "numbers_in", "place_chart",
    "read_data", "read_picture", "spec_from_reading", "spec_from_table", "table_style",
]  # fmt: skip
