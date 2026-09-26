"""Таблица из приложенного файла (xlsx, csv) — на слайд открытой копии (этап 39).

Файл читается первым листом: первая непустая строка — шапка, дальше строки данных (не больше
15 строк и 8 столбцов, остальное — пометка). Место называет человек словами, его находит модель
по снимку слайда, как для картинки (`office_image`); таблица строится тем же построителем, что
и в сборке (`layout/tables.add_table`), в стиле шаблона: шрифт, кегль и цвет шапки из профиля.
"""

from __future__ import annotations

import csv
import io
from typing import Any

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Pt

from presentation_designer.generation.office_image import Placement, area_for
from presentation_designer.layout.tables import TableSpec, TableStyle, add_table, format_number

MAX_ROWS = 15
MAX_COLS = 8
ROW_HEIGHT_FACTOR = 2.4  # высота строки в кеглях: строка текста и поля ячейки


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, (int, float)):
        return format_number(value)
    return str(value).strip()


def read_table(data: bytes, name: str) -> tuple[list[str], list[list[str]], bool]:
    """Шапка, строки и признак «обрезано» из xlsx или csv."""
    lower = name.lower()
    rows: list[list[str]] = []
    if lower.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = book.worksheets[0]
        for values in sheet.iter_rows(values_only=True):
            cells = [_cell(v) for v in values]
            if any(cells):
                rows.append(cells)
            if len(rows) > MAX_ROWS + 1:
                break
    elif lower.endswith((".csv", ".tsv", ".txt")):
        text = data.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
            delimiter = dialect.delimiter
        except csv.Error:
            delimiter = ";" if text.count(";") > text.count(",") else ","
        for values in csv.reader(io.StringIO(text), delimiter=delimiter):
            cells = [v.strip() for v in values]
            if any(cells):
                rows.append(cells)
            if len(rows) > MAX_ROWS + 1:
                break
    else:
        raise ValueError("таблицу можно взять из xlsx или csv")
    if len(rows) < 2:
        raise ValueError("в файле нет строк с данными: нужна шапка и хотя бы одна строка")
    width = min(MAX_COLS, max(len(r) for r in rows))
    # Пустые хвостовые столбцы не нужны.
    while width > 1 and not any(r[width - 1] if len(r) >= width else "" for r in rows):
        width -= 1
    table = [(r + [""] * width)[:width] for r in rows]
    truncated = len(rows) > MAX_ROWS + 1 or max(len(r) for r in rows) > width
    return table[0], table[1 : MAX_ROWS + 1], truncated


def _numeric(column: list[str]) -> bool:
    filled = [c for c in column if c]
    # Разделители разрядов бывают любыми пробелами: обычный, неразрывный, узкий (U+202F).
    return bool(filled) and all(
        "".join(c.split())
        .replace(",", ".")
        .rstrip("%₽$€")
        .lstrip("-−")
        .replace(".", "", 1)
        .isdigit()
        for c in filled
    )


def compact_box(
    box: tuple[int, int, int, int], n_rows: int, style: TableStyle
) -> tuple[int, int, int, int]:
    """Место бывает выше, чем нужно таблице: строки не растягиваются на всю высоту, таблица
    встаёт по центру места."""
    x, y, cx, cy = box
    need = n_rows * int(Pt(style.font_size_pt * ROW_HEIGHT_FACTOR))
    return x, y + max(0, (cy - need) // 2), cx, min(cy, need)


def center_cells(frame: Any) -> None:
    """Текст в ячейках — по середине строки."""
    for row in frame.table.rows:
        for cell in row.cells:
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE


def place_table(
    data: bytes,
    slide: int,
    header: list[str],
    rows: list[list[str]],
    placement: Placement,
    style: TableStyle,
    snapshot: bytes | None = None,
) -> bytes:
    """PPTX с таблицей на слайде `slide` в месте `placement` (объект — его рамка)."""
    prs = Presentation(io.BytesIO(data))
    if not 1 <= slide <= len(prs.slides):
        raise ValueError(f"слайда {slide} нет")
    width, height = int(prs.slide_width or 0), int(prs.slide_height or 0)
    area = area_for(data, slide, placement, snapshot)
    columns = list(zip(*rows, strict=False)) if rows else [()] * len(header)
    spec = TableSpec(
        header=header,
        rows=rows,
        numeric_columns=[_numeric(list(c)) for c in columns],
        dataset_id="chat",
        row_offset=0,
        truncated=False,
        highlight_row=None,
    )
    box = compact_box(
        (
            round(area[0] * width),
            round(area[1] * height),
            round(area[2] * width),
            round(area[3] * height),
        ),
        len(rows) + 1,
        style,
    )
    frame = add_table(prs.slides[slide - 1], box, spec, style)
    frame.name = "Таблица из чата"
    center_cells(frame)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


__all__ = ["center_cells", "compact_box", "place_table", "read_table"]
