"""Диаграмма-картинка → нативная диаграмма PowerPoint по чтению с картинки (`ChartReading`).

Картинка заменяется на том же месте и в том же порядке наложения. Повторяется устройство
картинки: тип, ряды и их цвета, цвета сегментов, отверстие и поворот кольца, сглаживание
линий, вторая ось, порядок полос сверху вниз, шкала, легенда и подписи данных. Область
построения ставится ручной раскладкой туда, где данные лежали на картинке; толщина линий,
столбцов и кегль подписей пересчитываются из пикселей картинки в размеры на слайде. Данные —
во встроенной книге: их правят в редакторе.

Значения, снятые по пикселям, приблизительны: описание диаграммы (alt text) говорит, сколько
точек измерено, а отчёт о заменах (`swap_summary`) уходит в чат.
"""

from __future__ import annotations

import hashlib
import re
from collections.abc import Iterable, Mapping
from dataclasses import dataclass
from typing import Any

from pptx.chart.axis import CategoryAxis, ValueAxis
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_MARK
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Pt

from presentation_designer.layout.charts import explicit_booleans
from presentation_designer.parsing.raster_charts.model import ChartReading
from presentation_designer.parsing.template.styles import (
    Theme,
    clr_map_of,
    parse_theme,
    resolve_color,
)
from presentation_designer.shared.text import plural

TEXT_COLOR = "#595959"
AXIS_COLOR = "#BFBFBF"
# Имена рядов, которые модель ставит, когда подписи ряда на картинке нет.
GENERIC_NAME = re.compile(r"^(ряд|series|значения|сегмент)(\s*\d+)?$", re.IGNORECASE)
LEGEND = {
    "bottom": XL_LEGEND_POSITION.BOTTOM,
    "top": XL_LEGEND_POSITION.TOP,
    "right": XL_LEGEND_POSITION.RIGHT,
    "left": XL_LEGEND_POSITION.LEFT,
}


@dataclass
class ChartSwap:
    """Картинка слайда, прочитанная как диаграмма: заменена или оставлена с причиной."""

    slide: int
    name: str
    sha256: str
    status: str  # replaced | kept
    reason: str = ""
    measured: int = 0


def picture_sha(shape: Any) -> str:
    return hashlib.sha256(shape.image.blob).hexdigest()


def slide_pictures(slide: Any) -> list[Any]:
    """Рисунки слайда верхнего уровня (не заполнители и не рисунки внутри групп)."""
    return [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]


# ---------- размеры ----------


@dataclass
class _Scale:
    """Пересчёт пикселей картинки в пункты слайда по рамке, в которой она стоит."""

    x: float
    y: float

    @property
    def mean(self) -> float:
        return (self.x + self.y) / 2


def _scale(reading: ChartReading, cx: int, cy: int) -> _Scale:
    return _Scale(Emu(cx).pt / max(1, reading.width), Emu(cy).pt / max(1, reading.height))


def _font_pt(reading: ChartReading, scale: _Scale) -> float:
    font = reading.geometry.font
    if not font:
        return 11.0
    return round(min(20.0, max(7.0, font * scale.y)) * 2) / 2


# ---------- постройка ----------


def _base_type(reading: ChartReading) -> str:
    """Тип основной группы: у комбо — первый нелинейный ряд, линии идут своей группой."""
    if reading.kind in ("pie", "doughnut"):
        return reading.kind
    types = [s.type or "column" for s in reading.series]
    return next((t for t in types if t != "line"), "line")


def _xl_type(base: str, stacked: bool) -> Any:
    return {
        ("column", False): XL_CHART_TYPE.COLUMN_CLUSTERED,
        ("column", True): XL_CHART_TYPE.COLUMN_STACKED,
        ("bar", False): XL_CHART_TYPE.BAR_CLUSTERED,
        ("bar", True): XL_CHART_TYPE.BAR_STACKED,
        ("line", False): XL_CHART_TYPE.LINE,
        ("line", True): XL_CHART_TYPE.LINE_STACKED,
        ("area", False): XL_CHART_TYPE.AREA,
        ("area", True): XL_CHART_TYPE.AREA_STACKED,
        ("pie", False): XL_CHART_TYPE.PIE,
        ("doughnut", False): XL_CHART_TYPE.DOUGHNUT,
    }[(base, stacked and base not in ("pie", "doughnut"))]


def _data(reading: ChartReading, names: list[str]) -> Any:
    data = CategoryChartData()  # type: ignore[no-untyped-call]
    data.categories = list(reading.categories)
    for s, name in zip(reading.series, names, strict=True):
        data.add_series(name, [p.value for p in s.points])  # type: ignore[no-untyped-call]
    return data


def _rgb(hex_color: str) -> Any:
    return RGBColor.from_string(hex_color.lstrip("#").upper())  # type: ignore[no-untyped-call]


def _child(parent: Any, tag: str, val: str | None = None) -> Any:
    el = OxmlElement(tag)
    if val is not None:
        el.set("val", val)
    parent.append(el)
    return el


def _set_val(parent: Any, tag: str, val: str) -> None:
    el = parent.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        parent.append(el)
    el.set("val", val)


def _no_fill(parent: Any, after: Any = None) -> None:
    """spPr без заливки и обводки — фон слайда виден, как у картинки на белом.

    У chartSpace spPr по схеме идёт сразу за c:chart (`after`), у plotArea — последним."""
    sp = parent.find(qn("c:spPr"))
    if sp is None:
        sp = OxmlElement("c:spPr")
        if after is not None:
            after.addnext(sp)
        else:
            parent.append(sp)
    for child in list(sp):
        sp.remove(child)
    _child(sp, "a:noFill")
    _child(_child(sp, "a:ln"), "a:noFill")


def _manual_layout(plot_area: Any, reading: ChartReading) -> None:
    """Внутренняя область построения — там же, где данные лежали на картинке."""
    box = reading.geometry.plot
    if box is None:
        return
    left, top, right, bottom = box
    w, h = max(1, reading.width), max(1, reading.height)
    x0, y0 = max(0.0, left / w), max(0.0, top / h)
    x1, y1 = min(1.0, (right + 1) / w), min(1.0, (bottom + 1) / h)
    if x1 - x0 < 0.2 or y1 - y0 < 0.2:
        return
    layout = plot_area.find(qn("c:layout"))
    if layout is None:
        layout = OxmlElement("c:layout")
        plot_area.insert(0, layout)
    for child in list(layout):
        layout.remove(child)
    manual = _child(layout, "c:manualLayout")
    _child(manual, "c:layoutTarget", "inner")
    _child(manual, "c:xMode", "edge")
    _child(manual, "c:yMode", "edge")
    for tag, value in (("c:x", x0), ("c:y", y0), ("c:w", x1 - x0), ("c:h", y1 - y0)):
        _child(manual, tag, f"{value:.4f}")


def _to_line_series(ser: Any, smooth: bool) -> None:
    """Ряд столбцов → ряд линии: у линии нет invertIfNegative и shape, есть marker и smooth."""
    for tag in ("c:invertIfNegative", "c:pictureOptions", "c:shape"):
        for el in ser.findall(qn(tag)):
            ser.remove(el)
    marker = OxmlElement("c:marker")
    _child(marker, "c:symbol", "none")
    anchor = ser.find(qn("c:spPr"))
    if anchor is None:
        anchor = ser.find(qn("c:tx"))
    anchor.addnext(marker)
    smooth_el = OxmlElement("c:smooth")
    smooth_el.set("val", "1" if smooth else "0")
    ext = ser.find(qn("c:extLst"))
    if ext is not None:
        ext.addprevious(smooth_el)
    else:
        ser.append(smooth_el)


def _axis_ids(plot_area: Any) -> list[int]:
    return [
        int(el.get("val"))
        for el in plot_area.iter(qn("c:axId"))
        if el.getparent().tag.endswith("Ax")
    ]


def _secondary_axes(plot_area: Any, cat_id: int, val_id: int, scale: Any, font_pt: float) -> None:
    """Скрытая ось категорий и видимая ось значений справа для второй группы рядов; шкала,
    подписи и линия — как на картинке. Порядок элементов — по схеме CT_ValAx."""
    cat = OxmlElement("c:catAx")
    _child(cat, "c:axId", str(cat_id))
    scaling = _child(cat, "c:scaling")
    _child(scaling, "c:orientation", "minMax")
    _child(cat, "c:delete", "1")
    _child(cat, "c:axPos", "b")
    _child(cat, "c:majorTickMark", "none")
    _child(cat, "c:minorTickMark", "none")
    _child(cat, "c:tickLblPos", "nextTo")
    _child(cat, "c:crossAx", str(val_id))
    _child(cat, "c:crosses", "autoZero")
    _child(cat, "c:auto", "1")
    _child(cat, "c:lblAlgn", "ctr")
    _child(cat, "c:lblOffset", "100")
    _child(cat, "c:noMultiLvlLbl", "0")
    val = OxmlElement("c:valAx")
    _child(val, "c:axId", str(val_id))
    scaling = _child(val, "c:scaling")
    _child(scaling, "c:orientation", "minMax")
    if scale is not None:
        _child(scaling, "c:max", f"{scale.maximum:g}")
        _child(scaling, "c:min", f"{scale.minimum:g}")
    _child(val, "c:delete", "0")
    _child(val, "c:axPos", "r")
    fmt = _child(val, "c:numFmt")
    fmt.set("formatCode", scale.number_format if scale is not None else "General")
    fmt.set("sourceLinked", "0")
    _child(val, "c:majorTickMark", "none")
    _child(val, "c:minorTickMark", "none")
    _child(val, "c:tickLblPos", "nextTo")
    sp = _child(val, "c:spPr")
    _child(_child(_child(sp, "a:ln"), "a:solidFill"), "a:srgbClr", AXIS_COLOR.lstrip("#"))
    tx = _child(val, "c:txPr")
    _child(tx, "a:bodyPr")
    _child(tx, "a:lstStyle")
    paragraph = _child(tx, "a:p")
    run = _child(_child(paragraph, "a:pPr"), "a:defRPr")
    run.set("sz", str(int(font_pt * 100)))
    _child(_child(run, "a:solidFill"), "a:srgbClr", TEXT_COLOR.lstrip("#"))
    _child(paragraph, "a:endParaRPr").set("lang", "ru-RU")
    _child(val, "c:crossAx", str(cat_id))
    _child(val, "c:crosses", "max")
    _child(val, "c:crossBetween", "between")
    if scale is not None and scale.major_unit:
        _child(val, "c:majorUnit", f"{scale.major_unit:g}")
    last_axis = [el for el in plot_area if el.tag.endswith("Ax")][-1]
    last_axis.addnext(cat)
    cat.addnext(val)


def _split_lines(chart: Any, reading: ChartReading, base: str, font_pt: float) -> None:
    """Комбо: ряды-линии переезжают из основной группы в свою группу линий; если они на
    второй оси — со своей парой осей (категории скрыты, значения справа)."""
    lines = [
        i for i, s in enumerate(reading.series) if (s.type or base) == "line" and base != "line"
    ]
    if not lines:
        return
    plot_area = chart._chartSpace.chart.plotArea
    group = next(el for el in plot_area if el.tag.endswith("Chart"))
    serieses = group.findall(qn("c:ser"))
    secondary = any(reading.series[i].axis == "secondary" for i in lines)
    ids = _axis_ids(plot_area)
    if secondary:
        cat_id, val_id = max(ids) + 1, max(ids) + 2
        _secondary_axes(plot_area, cat_id, val_id, reading.secondary, font_pt)
        axes = [cat_id, val_id]
    else:
        axes = ids[:2]
    line_chart = OxmlElement("c:lineChart")
    _child(line_chart, "c:grouping", "standard")
    _child(line_chart, "c:varyColors", "0")
    for i in lines:
        ser = serieses[i]
        _to_line_series(ser, reading.series[i].smooth)
        line_chart.append(ser)
    _child(line_chart, "c:marker", "1")
    for axis_id in axes:
        _child(line_chart, "c:axId", str(axis_id))
    group.addnext(line_chart)


def _scale_axis(val_ax: Any, scale: Any) -> None:
    """Шкала оси значений (минимум, максимум, шаг, формат) как на картинке."""
    val_ax.minimum_scale = scale.minimum
    val_ax.maximum_scale = scale.maximum
    if scale.major_unit:
        val_ax.major_unit = scale.major_unit
    val_ax.tick_labels.number_format = scale.number_format or "General"
    val_ax.tick_labels.number_format_is_linked = False


def _style_axes(chart: Any, reading: ChartReading, font_pt: float, base: str) -> None:
    # Основные оси — первые в plotArea: `chart.value_axis` при двух осях значений отдаёт
    # вторую (у точечных диаграмм первая — ось X), а у комбо вторая — правая.
    plot_area = chart._chartSpace.chart.plotArea
    cats, vals = plot_area.findall(qn("c:catAx")), plot_area.findall(qn("c:valAx"))
    if not cats or not vals:
        return
    cat_ax = CategoryAxis(cats[0])  # type: ignore[no-untyped-call]
    val_ax = ValueAxis(vals[0])  # type: ignore[no-untyped-call]
    for axis in (cat_ax, val_ax):
        axis.has_major_gridlines = False
        axis.has_minor_gridlines = False
        axis.major_tick_mark = XL_TICK_MARK.NONE
        axis.minor_tick_mark = XL_TICK_MARK.NONE
        axis.format.line.color.rgb = _rgb(AXIS_COLOR)
        axis.tick_labels.font.size = Pt(font_pt)
        axis.tick_labels.font.color.rgb = _rgb(TEXT_COLOR)
    if base == "bar":
        # Полосы сверху вниз в порядке картинки; ось значений остаётся внизу.
        cat_ax.reverse_order = True
        _set_val(val_ax._element, "c:crosses", "max")
    if reading.primary is None:
        # На картинке нет подписей шкалы (часто — подписи у самих столбцов): оси не видно,
        # а её конец там же, где кончалась область данных на картинке.
        val_ax.visible = False
        if reading.geometry.value_max:
            val_ax.minimum_scale = 0
            val_ax.maximum_scale = reading.geometry.value_max
    else:
        _scale_axis(val_ax, reading.primary)
    if base == "area":
        # Области идут от края до края области построения, как на картинке.
        _set_val(val_ax._element, "c:crossBetween", "midCat")


def _style_series(chart: Any, reading: ChartReading, scale: _Scale, base: str) -> None:
    plots = list(chart.plots)
    # Порядок рядов — по c:idx: комбо разносит их по группам, а номер остаётся прежним.
    by_idx = {
        int(series._element.find(qn("c:idx")).get("val")): series
        for plot in plots
        for series in plot.series
    }
    for i, rs in enumerate(reading.series):
        series = by_idx[i]
        stroke = reading.geometry.stroke[i] if i < len(reading.geometry.stroke) else None
        kind = rs.type or base
        if reading.kind in ("pie", "doughnut"):
            colors = reading.slice_colors or [rs.color]
            for k, point in enumerate(series.points):
                if k >= len(reading.categories):
                    break
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = _rgb(colors[k % len(colors)])
                point.format.line.fill.background()
            continue
        if kind == "line":
            width = ((stroke or 2.0) + 0.75) * scale.mean
            series.format.line.color.rgb = _rgb(rs.color)
            series.format.line.width = Pt(min(8.0, max(0.75, width)))
            series.smooth = rs.smooth
        else:
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = _rgb(rs.color)
            series.format.line.fill.background()
            # Выделение цветом: у каждого столбца свой цвет, как на картинке.
            for point, color in zip(series.points, rs.point_colors, strict=False):
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = _rgb(color)
                point.format.line.fill.background()
    if base in ("column", "bar") and reading.geometry.bar and reading.geometry.pitch:
        bar_plot = plots[0]
        n = 1 if reading.stacked else sum(1 for s in reading.series if (s.type or base) == base)
        gap = (reading.geometry.pitch - n * reading.geometry.bar) / reading.geometry.bar * 100
        bar_plot.gap_width = int(min(500, max(0, round(gap))))
        bar_plot.overlap = 100 if reading.stacked else 0


def _labels(chart: Any, reading: ChartReading, font_pt: float, base: str) -> None:
    for plot in chart.plots:
        plot.has_data_labels = reading.data_labels
        if not reading.data_labels:
            continue
        labels = plot.data_labels
        labels.number_format = reading.label_format or "General"
        labels.number_format_is_linked = False
        labels.show_value = True
        labels.font.size = Pt(font_pt)
        labels.font.color.rgb = _rgb(TEXT_COLOR)
        if base in ("column", "bar"):
            labels.position = XL_LABEL_POSITION.OUTSIDE_END


def _round_shape(chart: Any, reading: ChartReading) -> None:
    group = next(el for el in chart._chartSpace.chart.plotArea if el.tag.endswith("Chart"))
    if reading.first_angle is not None:
        _set_val(group, "c:firstSliceAng", str(round(reading.first_angle) % 360))
    if reading.kind == "doughnut" and reading.hole:
        _set_val(group, "c:holeSize", str(min(90, max(10, round(reading.hole * 100)))))


def series_names(reading: ChartReading, slide: Any = None, picture: Any = None) -> list[str]:
    """Имена рядов: с картинки, а если модель дала «Ряд N» — из легенды, нарисованной на
    слайде фигурами (кружок цвета ряда и подпись справа от него)."""
    names = [s.name for s in reading.series]
    if slide is None or reading.kind in ("pie", "doughnut"):
        return names
    generic = [i for i, n in enumerate(names) if GENERIC_NAME.match(n.strip())]
    if not generic:
        return names
    found = legend_texts(slide, picture)
    for i in generic:
        color = reading.series[i].color
        best = min(found, key=lambda t: _color_distance(t[0], color), default=None)
        if best is not None and _color_distance(best[0], color) <= 60:
            names[i] = best[1]
    return names


def legend_texts(slide: Any, picture: Any = None) -> list[tuple[str, str]]:
    """Пары «цвет маркера — текст» легенды, нарисованной на слайде: небольшая фигура со
    сплошной заливкой и надпись справа от неё на той же строке. Координаты фигур внутри
    групп переводятся в координаты слайда: у каждой группы своя система (chOff/chExt)."""
    theme = _slide_theme(slide)
    markers: list[tuple[Box, str]] = []
    texts: list[tuple[Box, str]] = []
    limit = int(slide.part.package.presentation_part.presentation.slide_width or 0) // 20
    for sh, box in _flat(slide.shapes):
        if picture is not None and sh.shape_id == picture.shape_id:
            continue
        if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip():
            text = " ".join(sh.text_frame.text.split())
            if len(text) <= 40:
                texts.append((box, text))
            continue
        color = _solid_fill(sh, theme)
        if color and 0 < box[2] <= limit and 0 < box[3] <= limit:
            markers.append((box, color))
    out: list[tuple[str, str]] = []
    for (left, top, width, height), color in markers:
        mid = top + height / 2
        right = left + width
        near = [
            (t[0] - right, text)
            for t, text in texts
            if t[1] <= mid <= t[1] + t[3] and -width < t[0] - right < 8 * width
        ]
        if near:
            out.append((color, min(near)[1]))
    return out


Box = tuple[float, float, float, float]


def _flat(shapes: Iterable[Any], outer: Any = None) -> Iterable[tuple[Any, Box]]:
    """Фигуры без групп с рамкой в координатах слайда."""
    for sh in shapes:
        if sh.left is None or sh.top is None or sh.width is None or sh.height is None:
            continue
        box: Box = (float(sh.left), float(sh.top), float(sh.width), float(sh.height))
        if outer is not None:
            box = outer(box)
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _flat(sh.shapes, _group_transform(sh, box))
        else:
            yield sh, box


def _group_transform(group: Any, box: Box) -> Any:
    """Перевод из системы координат детей группы (a:chOff/a:chExt) в рамку группы на слайде."""
    xfrm = group._element.grpSpPr.find(qn("a:xfrm"))
    ch_off = xfrm.find(qn("a:chOff")) if xfrm is not None else None
    ch_ext = xfrm.find(qn("a:chExt")) if xfrm is not None else None
    if ch_off is None or ch_ext is None:
        return None
    cx0, cy0 = float(ch_off.get("x", 0)), float(ch_off.get("y", 0))
    cw, ch = float(ch_ext.get("cx", 0)) or 1.0, float(ch_ext.get("cy", 0)) or 1.0
    kx, ky = box[2] / cw, box[3] / ch

    def apply(child: Box) -> Box:
        return (
            box[0] + (child[0] - cx0) * kx,
            box[1] + (child[1] - cy0) * ky,
            child[2] * kx,
            child[3] * ky,
        )

    return apply


def _slide_theme(slide: Any) -> Theme | None:
    """Тема мастера слайда с картой цветов: маркеры легенды часто залиты цветом темы."""
    try:
        master = slide.slide_layout.slide_master
        theme = parse_theme(master.part.part_related_by(RT.THEME))
        theme.clr_map = clr_map_of(master._element)
        return theme
    except (AttributeError, KeyError, ValueError):
        return None


def _solid_fill(shape: Any, theme: Theme | None) -> str | None:
    """Сплошная заливка фигуры в #RRGGBB: RGB, цвет темы с модификаторами или системный."""
    sp_pr = getattr(shape._element, "spPr", None)
    fill = sp_pr.find(qn("a:solidFill")) if sp_pr is not None else None
    resolved = resolve_color(fill, theme) if fill is not None else None
    return resolved.hex if resolved else None


def _color_distance(a: str, b: str) -> int:
    pa = [int(a.lstrip("#")[k : k + 2], 16) for k in (0, 2, 4)]
    pb = [int(b.lstrip("#")[k : k + 2], 16) for k in (0, 2, 4)]
    return max(abs(x - y) for x, y in zip(pa, pb, strict=True))


def describe(reading: ChartReading) -> str:
    """Описание (alt text): откуда данные и что проверить."""
    counts = reading.counts()
    approx = counts["measured"] + counts["inferred"] + counts["model"]
    if not approx:
        return "Диаграмма восстановлена по картинке: значения перенесены из подписей."
    return (
        "Диаграмма восстановлена по картинке: "
        f"{approx} {plural(approx, 'значение снято', 'значения сняты', 'значений снято')} "
        "по пикселям — проверьте перед показом."
    )


def add_native_chart(
    slide: Any, box: tuple[int, int, int, int], reading: ChartReading, names: list[str]
) -> Any:
    """Нативная диаграмма в рамке `box` (EMU), похожая на картинку; возвращает graphicFrame."""
    x, y, cx, cy = box
    base = _base_type(reading)
    frame = slide.shapes.add_chart(
        _xl_type(base, reading.stacked), x, y, cx, cy, _data(reading, names)
    )
    chart = frame.chart
    scale = _scale(reading, cx, cy)
    font_pt = _font_pt(reading, scale)
    chart.font.size = Pt(font_pt)
    chart.font.color.rgb = _rgb(TEXT_COLOR)
    chart.has_title = False
    space = chart._chartSpace
    _no_fill(space, after=space.chart)
    plot_area = space.chart.plotArea
    _no_fill(plot_area)
    if reading.kind in ("pie", "doughnut"):
        _round_shape(chart, reading)
    else:
        _split_lines(chart, reading, base, font_pt)
    _manual_layout(plot_area, reading)
    _style_series(chart, reading, scale, base)
    if reading.kind not in ("pie", "doughnut"):
        _style_axes(chart, reading, font_pt, base)
    _labels(chart, reading, font_pt, base)
    chart.has_legend = reading.legend in LEGEND
    if chart.has_legend:
        chart.legend.position = LEGEND[reading.legend]
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(font_pt)
        chart.legend.font.color.rgb = _rgb(TEXT_COLOR)
    explicit_booleans(chart)
    return frame


def swap_picture(slide: Any, picture: Any, reading: ChartReading) -> Any:
    """Картинка → нативная диаграмма на том же месте и в том же порядке наложения."""
    box = (int(picture.left), int(picture.top), int(picture.width), int(picture.height))
    names = series_names(reading, slide, picture)
    frame = add_native_chart(slide, box, reading, names)
    pic_el = picture._element
    rid = pic_el.blipFill.blip.rEmbed if pic_el.blipFill.blip is not None else None
    pic_el.addprevious(frame._element)
    c_nv_pr = frame._element.nvGraphicFramePr.cNvPr
    c_nv_pr.set("name", f"Диаграмма {frame.shape_id} (по картинке)")
    c_nv_pr.set("descr", describe(reading))
    pic_el.getparent().remove(pic_el)
    if rid and not _referenced(slide, rid):
        slide.part.rels.pop(rid)
    return frame


def _referenced(slide: Any, rid: str) -> bool:
    """Ссылается ли ещё что-то на слайде на связь: рисунки — через r:embed и r:link."""
    refs = slide._element.xpath("//@r:embed | //@r:link | //@r:id")
    return rid in {str(r) for r in refs}


def _blocked(picture: Any) -> str:
    """Почему картинку нельзя честно заменить диаграммой того же размера."""
    if picture.rotation:
        return "картинка повёрнута"
    crop = [picture.crop_left, picture.crop_right, picture.crop_top, picture.crop_bottom]
    if any(abs(c) > 0.01 for c in crop):
        return "картинка обрезана на слайде"
    return ""


def swap_pictures(
    prs: Any, readings: Mapping[str, ChartReading], *, slides: Iterable[Any] | None = None
) -> list[ChartSwap]:
    """Все рисунки слайдов, для которых есть чтение, заменяются нативными диаграммами."""
    out: list[ChartSwap] = []
    # Номер слайда — по slide_id: обёртки слайдов python-pptx создаёт заново при каждом обходе.
    order = {s.slide_id: k for k, s in enumerate(prs.slides, start=1)}
    for slide in slides if slides is not None else prs.slides:
        number = order.get(slide.slide_id, 0)
        for picture in slide_pictures(slide):
            sha = picture_sha(picture)
            reading = readings.get(sha)
            if reading is None:
                continue
            name = picture.name
            blocked = _blocked(picture)
            if blocked:
                out.append(ChartSwap(number, name, sha, "kept", blocked))
                continue
            try:
                swap_picture(slide, picture, reading)
            except Exception as e:  # замена необязательна: картинка остаётся как была
                out.append(ChartSwap(number, name, sha, "kept", f"не построена: {e}"[:200]))
                continue
            counts = reading.counts()
            measured = counts["measured"] + counts["inferred"] + counts["model"]
            out.append(ChartSwap(number, name, sha, "replaced", measured=measured))
    return out


def deck_pictures(prs: Any) -> tuple[dict[str, bytes], dict[str, list[int]]]:
    """Картинки слайдов для чтения: байты по sha256 (повторы — один раз) и номера слайдов."""
    blobs: dict[str, bytes] = {}
    where: dict[str, list[int]] = {}
    for number, slide in enumerate(prs.slides, start=1):
        for picture in slide_pictures(slide):
            try:
                blob = picture.image.blob
            except (AttributeError, KeyError, ValueError):
                continue  # связанная, а не встроенная картинка
            sha = hashlib.sha256(blob).hexdigest()
            blobs.setdefault(sha, blob)
            where.setdefault(sha, []).append(number)
    return blobs, where


# Пропуски чтения, о которых стоит сказать: диаграмма, возможно, была, но до неё не дошли.
LATE = ("время на чтение диаграмм вышло", "лимит чтения диаграмм")


def unread_charts(
    outcomes: Mapping[str, Any], where: Mapping[str, list[int]]
) -> list[tuple[int, str]]:
    """Диаграммы, которые не удалось прочитать: (первый слайд, причина). Не диаграммы
    (фото, логотипы, отказ модели «это не диаграмма») в отчёт не попадают."""
    out: list[tuple[int, str]] = []
    for sha, outcome in outcomes.items():
        late = outcome.status == "skipped" and outcome.reason in LATE
        if outcome.status != "rejected" and not late:
            continue
        reason = "не успел прочитать" if late else str(outcome.reason or "не прочитана")
        if reason.startswith("модель не ответила"):
            reason = "модель не ответила"
        out.append((min(where.get(sha) or [0]), reason[:120]))
    return out


def charts_report(
    outcomes: Mapping[str, Any],
    where: Mapping[str, list[int]],
    swaps: list[ChartSwap],
    message: str | None,
) -> dict[str, Any]:
    """charts.json ревизии: что читалось, что заменено; чтения по sha256 переиспользует
    полная сборка варианта original, чтобы получить ту же колоду, что и предварительная."""
    return {
        "version": 1,
        "pictures": [
            {
                "sha256": sha,
                "slides": list(where.get(sha) or []),
                "status": outcome.status,
                "reason": outcome.reason,
                "reading": outcome.reading.model_dump() if outcome.reading else None,
            }
            for sha, outcome in outcomes.items()
            if outcome.status != "skipped" or outcome.reason in LATE
        ],
        "swaps": [
            {"slide": s.slide, "sha256": s.sha256, "status": s.status, "reason": s.reason}
            for s in swaps
        ],
        "message": message,
    }


def report_readings(report: Mapping[str, Any]) -> dict[str, ChartReading]:
    """Чтения из charts.json: sha256 картинки → ChartReading."""
    out: dict[str, ChartReading] = {}
    for item in report.get("pictures") or []:
        if item.get("reading"):
            out[str(item["sha256"])] = ChartReading.model_validate(item["reading"])
    return out


def swap_summary(swaps: list[ChartSwap], unread: list[tuple[int, str]]) -> str | None:
    """Фраза для чата: сколько диаграмм стали редактируемыми и что осталось картинками.

    `unread` — диаграммы, которые не удалось прочитать: (номер слайда, причина)."""
    replaced = [s for s in swaps if s.status == "replaced"]
    kept = [(s.slide, s.reason) for s in swaps if s.status == "kept"] + list(unread)
    if not replaced and not kept:
        return None
    parts = []
    if replaced:
        slides = sorted({s.slide for s in replaced})
        parts.append(f"Сделал диаграммы редактируемыми: {len(replaced)} ({_slides(slides)})")
        if any(s.measured for s in replaced):
            parts[-1] += "; значения сняты по картинке — проверьте их"
    if kept:
        n = len(kept)
        reasons = "; ".join(f"слайд {slide} — {reason}" for slide, reason in sorted(set(kept))[:4])
        parts.append(f"{n} {plural(n, 'осталась', 'остались', 'остались')} картинкой: {reasons}")
    return ". ".join(parts) + "."


def _slides(numbers: list[int]) -> str:
    """[45, 46, 47, 49] → «слайды 45–47, 49»: подряд идущие номера — диапазоном."""
    if len(numbers) == 1:
        return f"слайд {numbers[0]}"
    runs: list[list[int]] = []
    for n in numbers:
        if runs and n == runs[-1][-1] + 1:
            runs[-1].append(n)
        else:
            runs.append([n])
    return "слайды " + ", ".join(
        f"{r[0]}–{r[-1]}" if len(r) > 2 else ", ".join(str(n) for n in r) for r in runs
    )


__all__ = [
    "ChartSwap",
    "add_native_chart",
    "charts_report",
    "deck_pictures",
    "describe",
    "legend_texts",
    "picture_sha",
    "report_readings",
    "series_names",
    "slide_pictures",
    "swap_picture",
    "swap_pictures",
    "swap_summary",
    "unread_charts",
]
