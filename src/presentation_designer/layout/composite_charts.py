"""Диаграммы, собранные на слайде из кусков, → нативные диаграммы PowerPoint.

Шаблоны из Google Slides и Figma рисуют диаграмму отдельными объектами: каждый столбец — своя
картинка (часто с градиентом в прозрачность) или прямоугольник, числа у столбцов, подписи
категорий и легенда — отдельные надписи, кольцо — картинка с прозрачной серединой (иногда две
наложенные), а число в центре — надпись поверх. В редакторе это набор фигур: число правится,
а столбец или дуга остаются прежними.

Здесь такие наборы находятся по геометрии слайда, без модели: столбцы и полосы с общим
основанием и числами у концов, кольца с числом-процентом в центре. На их месте встаёт одна
нативная диаграмма того же вида: заливки (градиент, прозрачность) повторяют картинки, числа
становятся подписями данных самой диаграммы, подписи категорий — её осью, легенда — её
легендой. Данные — во встроенной книге: меняется число — меняется рисунок. У кольца остаток —
формула `=100-B2`, правится одно число.

Значения: числа на слайде сходятся с рисунком — берутся числа; не сходятся (в шаблонах на
всех столбцах «10» при разной высоте, «10%» в каждом кольце) — значения снимаются по рисунку:
у кольца — доля окружности, у столбцов — пропорционально длине, самый длинный равен самому
большому числу. Такая диаграмма помечается «значения сняты по рисунку — проверьте».
"""

from __future__ import annotations

import io
import math
from collections import Counter
from collections.abc import Iterable
from dataclasses import dataclass, field
from itertools import pairwise, permutations
from typing import Any, Literal

import xlsxwriter  # type: ignore[import-untyped]
from PIL import Image
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_MARK
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

from presentation_designer.layout.chart_images import ChartSwap, _slide_theme
from presentation_designer.layout.charts import explicit_booleans
from presentation_designer.parsing.raster_charts.labels import LabelNumber, label_number
from presentation_designer.parsing.raster_charts.pieces import (
    Paint,
    Ring,
    RingError,
    Stop,
    bar_paint,
    measure_ring,
    rgba,
    ring_shape,
)
from presentation_designer.parsing.template.styles import Theme, resolve_color

Box = tuple[float, float, float, float]  # x, y, w, h в EMU
EMU_PT = 12700
EMU_IN = 914400
# Единицы, при которых надпись всё ещё «просто число»: «10%», «259 ₽», «12 млн».
UNITS = {"%", "‰", "₽", "$", "€", "млн", "млрд", "тыс", "тыс.", "k", "K", "M", "x", "×"}
BAR_GEOMETRY = {"rect", "roundRect", "round1Rect", "round2SameRect", "snip1Rect", "snipRoundRect"}
MARKER_GEOMETRY = {"ellipse", "rect", "roundRect", "flowChartConnector"}
LABEL_POSITIONS = {
    "inEnd": XL_LABEL_POSITION.INSIDE_END,
    "outEnd": XL_LABEL_POSITION.OUTSIDE_END,
    "ctr": XL_LABEL_POSITION.CENTER,
}


# ---------- объекты слайда ----------


@dataclass(frozen=True)
class TextStyle:
    size: float  # пункты
    color: str  # #RRGGBB
    font: str | None = None
    bold: bool = False


@dataclass
class Item:
    """Объект верхнего уровня слайда: рамка в EMU, порядок наложения, текст и число."""

    shape: Any
    box: Box
    z: int
    text: str = ""
    number: LabelNumber | None = None
    kind: Literal["picture", "shape", "text", "other"] = "other"

    @property
    def cx(self) -> float:
        return self.box[0] + self.box[2] / 2

    @property
    def cy(self) -> float:
        return self.box[1] + self.box[3] / 2


def _items(slide: Any) -> list[Item]:
    out: list[Item] = []
    for z, sh in enumerate(slide.shapes):
        if None in (sh.left, sh.top, sh.width, sh.height):
            continue
        box = (float(sh.left), float(sh.top), float(sh.width), float(sh.height))
        item = Item(sh, box, z)
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            item.kind = "picture"
        elif sh.is_placeholder or sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            item.kind = "other"
        elif getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip():
            item.kind = "text"
            item.text = sh.text_frame.text.strip()
            item.number = _plain_number(item.text)
        elif sh.shape_type in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.FREEFORM):
            item.kind = "shape"
        out.append(item)
    return out


def _plain_number(text: str) -> LabelNumber | None:
    """Надпись — одно число с единицей («10», «10%», «259 ₽»), без слов и переносов."""
    if "\n" in text or "\v" in text or len(text) > 24:
        return None
    number = label_number(text)
    if number is None:
        return None
    if number.unit is not None and number.unit not in UNITS:
        return None
    return number


def _rotated_or_cropped(sh: Any) -> bool:
    if getattr(sh, "rotation", 0):
        return True
    if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
        crop = [sh.crop_left, sh.crop_right, sh.crop_top, sh.crop_bottom]
        return any(abs(c) > 0.01 for c in crop)
    return False


def _image(item: Item) -> Image.Image | None:
    try:
        return rgba(item.shape.image.blob)
    except Exception:  # связанная или битая картинка — не кусок диаграммы
        return None


# ---------- текст надписи ----------


def text_style(shape: Any, theme: Theme | None) -> TextStyle:
    """Кегль, цвет, гарнитура и жирность первого фрагмента надписи; по умолчанию — 18 пт,
    чёрный, шрифт темы."""
    size, color, font, bold = 18.0, "#000000", None, False
    body = shape.text_frame._txBody
    rpr = None
    for r in body.iter(qn("a:r")):
        t = r.find(qn("a:t"))
        if t is not None and (t.text or "").strip():
            rpr = r.find(qn("a:rPr"))
            break
    if rpr is None:
        rpr = next(body.iter(qn("a:endParaRPr")), None)
    if rpr is not None:
        if rpr.get("sz"):
            size = int(rpr.get("sz")) / 100
        bold = rpr.get("b") in ("1", "true")
        fill = rpr.find(qn("a:solidFill"))
        resolved = resolve_color(fill, theme) if fill is not None else None
        if resolved is not None:
            color = resolved.hex
        latin = rpr.find(qn("a:latin"))
        if latin is not None and latin.get("typeface"):
            font = latin.get("typeface")
    if font in ("+mn-lt", "+mj-lt") and theme is not None:
        font = theme.minor_font if font == "+mn-lt" else theme.major_font
    return TextStyle(size, color, font, bold)


def text_box(item: Item, style: TextStyle) -> Box:
    """Где на слайде сам текст надписи: по полям, привязке по вертикали и выравниванию.
    Ширина — оценка по числу знаков: у чисел ширина знака около 0,6 кегля."""
    sh = item.shape
    body = sh.text_frame._txBody
    bpr = body.find(qn("a:bodyPr"))

    def inset(name: str, default: int) -> float:
        v = bpr.get(name) if bpr is not None else None
        return float(v) if v is not None else float(default)

    left, top, width, height = item.box
    l_in, r_in = inset("lIns", 91440), inset("rIns", 91440)
    t_in, b_in = inset("tIns", 45720), inset("bIns", 45720)
    lines = max(1, len(item.text.splitlines()))
    spacing = 1.2
    ln = next(body.iter(qn("a:lnSpc")), None)
    pct = ln.find(qn("a:spcPct")) if ln is not None else None
    if pct is not None and pct.get("val"):
        spacing = 1.2 * int(pct.get("val")) / 100000
    text_h = lines * style.size * spacing * EMU_PT
    text_w = min(width - l_in - r_in, len(item.text) * style.size * 0.6 * EMU_PT)
    text_w = max(text_w, style.size * 0.6 * EMU_PT)
    anchor = bpr.get("anchor", "t") if bpr is not None else "t"
    inner_top, inner_h = top + t_in, max(0.0, height - t_in - b_in)
    y = {
        "ctr": inner_top + (inner_h - text_h) / 2,
        "b": inner_top + inner_h - text_h,
    }.get(anchor, inner_top)
    ppr = next(body.iter(qn("a:pPr")), None)
    algn = ppr.get("algn", "l") if ppr is not None else "l"
    inner_left, inner_w = left + l_in, max(0.0, width - l_in - r_in)
    x = {
        "ctr": inner_left + (inner_w - text_w) / 2,
        "r": inner_left + inner_w - text_w,
    }.get(algn, inner_left)
    return (x, y, text_w, text_h)


# ---------- заливки фигур ----------


def _shape_paint(sh: Any, theme: Theme | None, along: Literal["down", "right"]) -> Paint | None:
    """Заливка фигуры-столбца: сплошная или линейный градиент (с прозрачностью точек)."""
    sp_pr = sh._element.find(qn("p:spPr"))
    if sp_pr is None:
        return None
    solid = sp_pr.find(qn("a:solidFill"))
    if solid is not None:
        stop = _stop(solid, 0.0, theme)
        return Paint((stop,), 90.0 if along == "down" else 0.0) if stop else None
    grad = sp_pr.find(qn("a:gradFill"))
    if grad is None:
        return None
    stops = [
        s
        for gs in grad.iter(qn("a:gs"))
        if (s := _stop(gs, int(gs.get("pos", "0")) / 100000, theme)) is not None
    ]
    if not stops:
        return None
    lin = grad.find(qn("a:lin"))
    angle = int(lin.get("ang", "0")) / 60000 if lin is not None else 90.0
    return Paint(tuple(sorted(stops, key=lambda s: s.pos)), angle)


def _stop(parent: Any, pos: float, theme: Theme | None) -> Stop | None:
    resolved = resolve_color(parent, theme)
    if resolved is None:
        return None
    alpha = 1.0
    for el in parent.iter(qn("a:alpha")):
        alpha = int(el.get("val", "100000")) / 100000
    h = resolved.hex.lstrip("#")
    return Stop(pos, (int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)), alpha)


def _hex(c: tuple[int, int, int]) -> str:
    return "#{:02X}{:02X}{:02X}".format(*c)


def _distance(a: tuple[int, int, int], b: tuple[int, int, int]) -> int:
    return max(abs(x - y) for x, y in zip(a, b, strict=True))


# ---------- найденная диаграмма ----------


@dataclass
class Label:
    """Число в центре кольца: где стоял текст, где подпись стоит без сдвига и каким шрифтом."""

    target: tuple[float, float]  # центр текста на слайде, EMU
    anchor: tuple[float, float]  # место подписи без сдвига, EMU
    size: tuple[float, float]  # ширина и высота подписи, EMU
    style: TextStyle


@dataclass
class Legend:
    box: Box
    style: TextStyle


@dataclass
class Assembled:
    """Диаграмма, собранная из кусков: что построить и какие фигуры убрать."""

    kind: Literal["doughnut", "column", "bar"]
    frame: Box
    plot: Box  # область построения на слайде, EMU
    categories: list[str]
    names: list[str]
    values: list[list[float]]
    paints: list[Paint]  # по ряду (столбцы) или по точке (кольцо)
    number_format: str
    basis: Literal["label", "measured"]
    pieces: list[Item]
    # кольцо
    first_angle: float = 0.0
    hole: float = 0.5
    ring_labels: list[Label] = field(default_factory=list)
    # столбцы и полосы
    label_style: TextStyle | None = None
    label_position: str = "outEnd"
    category_style: TextStyle | None = None
    legend: Legend | None = None
    gap: int = 150
    overlap: int = 0
    reverse: bool = False  # полосы растут влево
    point_paints: list[list[Paint | None]] = field(default_factory=list)  # свои у столбцов

    @property
    def measured(self) -> int:
        return sum(len(v) for v in self.values) if self.basis == "measured" else 0


def _union(boxes: Iterable[Box]) -> Box:
    bs = list(boxes)
    x0 = min(b[0] for b in bs)
    y0 = min(b[1] for b in bs)
    x1 = max(b[0] + b[2] for b in bs)
    y1 = max(b[1] + b[3] for b in bs)
    return (x0, y0, x1 - x0, y1 - y0)


def _within(inner: Box, outer: Box, slack: float) -> bool:
    return (
        inner[0] >= outer[0] - slack
        and inner[1] >= outer[1] - slack
        and inner[0] + inner[2] <= outer[0] + outer[2] + slack
        and inner[1] + inner[3] <= outer[1] + outer[3] + slack
    )


# ---------- кольца ----------


def _compose(base: Item, image: Image.Image, overlays: list[Item]) -> Image.Image:
    """Картинка кольца с наложенными сверху кусками — в пикселях основной картинки."""
    out = image.copy()
    sx, sy = image.width / base.box[2], image.height / base.box[3]
    for q in overlays:
        pic = _image(q)
        if pic is None:
            continue
        w, h = max(1, round(q.box[2] * sx)), max(1, round(q.box[3] * sy))
        layer = Image.new("RGBA", out.size, (0, 0, 0, 0))
        layer.paste(
            pic.resize((w, h), Image.Resampling.LANCZOS),
            (round((q.box[0] - base.box[0]) * sx), round((q.box[1] - base.box[1]) * sy)),
        )
        out = Image.alpha_composite(out, layer)
    return out


def _ring_share(ring: Ring, image: Image.Image, q: Item, base: Item) -> float:
    """Доля непрозрачных пикселей куска `q`, попавших в полосу кольца."""
    pic = _image(q)
    if pic is None:
        return 0.0
    sx, sy = image.width / base.box[2], image.height / base.box[3]
    ox, oy = (q.box[0] - base.box[0]) * sx, (q.box[1] - base.box[1]) * sy
    kx, ky = q.box[2] * sx / pic.width, q.box[3] * sy / pic.height
    px = pic.load()
    assert px is not None
    step = max(1, min(pic.size) // 60)
    inside = total = 0
    for y in range(0, pic.height, step):
        for x in range(0, pic.width, step):
            if px[x, y][3] < 32:  # type: ignore[index]
                continue
            total += 1
            r = math.hypot(ox + x * kx - ring.cx, oy + y * ky - ring.cy)
            if ring.inner * 0.9 <= r <= ring.outer * 1.08:
                inside += 1
    return inside / total if total else 0.0


def _find_rings(items: list[Item], theme: Theme | None, used: set[int]) -> list[Assembled]:
    out: list[Assembled] = []
    numbers = [it for it in items if it.kind == "text" and it.number is not None]
    for base in items:
        if base.kind != "picture" or id(base) in used or _rotated_or_cropped(base.shape):
            continue
        w, h = base.box[2], base.box[3]
        if min(w, h) < 0.4 * EMU_IN or not 0.7 <= w / h <= 1.45:
            continue
        image = _image(base)
        if image is None:
            continue
        try:
            rx, ry, outer, _ = ring_shape(image)
        except RingError:
            continue
        sx, sy = base.box[2] / image.width, base.box[3] / image.height
        if abs(sx - sy) > 0.05 * max(sx, sy):
            continue  # картинка растянута — на слайде эллипс, а не кольцо
        ring_box = (
            base.box[0] + (rx - outer) * sx,
            base.box[1] + (ry - outer) * sy,
            2 * outer * sx,
            2 * outer * sy,
        )
        overlays = [
            q
            for q in items
            if q.kind == "picture"
            and q.z > base.z
            and id(q) not in used
            and not _rotated_or_cropped(q.shape)
            and _within(q.box, ring_box, 0.05 * ring_box[2])
        ]
        try:
            ring = measure_ring(image, _compose(base, image, overlays) if overlays else None)
        except RingError:
            continue
        overlays = [q for q in overlays if _ring_share(ring, image, q, base) >= 0.5]
        found = _ring_chart(base, ring, ring_box, sx, overlays, numbers, theme, used)
        if found is not None:
            out.append(found)
            used.update(id(p) for p in found.pieces)
    return out


def _ring_chart(
    base: Item,
    ring: Ring,
    ring_box: Box,
    scale: float,
    overlays: list[Item],
    numbers: list[Item],
    theme: Theme | None,
    used: set[int],
) -> Assembled | None:
    cx, cy = ring_box[0] + ring_box[2] / 2, ring_box[1] + ring_box[3] / 2
    inner = ring.inner * scale
    center: Item | None = None
    for it in numbers:
        if id(it) in used:
            continue
        style = text_style(it.shape, theme)
        tb = text_box(it, style)
        if math.hypot(tb[0] + tb[2] / 2 - cx, tb[1] + tb[3] / 2 - cy) <= 0.6 * inner:
            center = it
            break
    # Кольцо без числа в центре и из одной картинки — дело чтения картинки моделью: там
    # подписи бывают внутри самой картинки.
    if center is None and not overlays:
        return None
    if center is not None and (center.number is None or center.number.unit != "%"):
        return None
    segments = list(ring.segments)
    labels: list[Label] = []
    if len(segments) == 2:
        main = max(segments, key=lambda s: s.saturation)
        rest = segments[1 - segments.index(main)]
        share = main.sweep / 3.6
        number = center.number if center is not None else None
        decimals = number.decimals if number is not None else 0
        if number is not None and abs(number.value - share) <= 2.0:
            value, basis = number.value, "label"
        else:
            value, basis = round(share, decimals), "measured"
        values = [value, round(100 - value, max(decimals, 2))]
        ordered = [main, rest]
        categories = ["Значение", "Остаток"]
        fmt = "0" + ("." + "0" * decimals if decimals else "") + '"%"'
    else:
        ordered = segments
        values = [round(s.sweep / 3.6, 1) for s in segments]
        basis = "measured"
        categories = [f"Доля {k + 1}" for k in range(len(segments))]
        fmt = '0.0"%"'
        center = None  # общее число в центре к одной доле не относится
    pieces = [base, *overlays]
    plot, hole = ring_box, ring.hole
    frame_parts = [ring_box]
    if center is not None:
        style = text_style(center.shape, theme)
        tb = text_box(center, style)
        size = (tb[2] * 1.25 + 0.1 * EMU_IN, style.size * 1.6 * EMU_PT)
        target = (tb[0] + tb[2] / 2, tb[1] + tb[3] / 2)
        pieces.append(center)
        # Число несёт второе, невидимое кольцо снаружи — полный круг: его подпись стоит на
        # середине круга (первый угол + 180°) при любом значении, и один постоянный сдвиг
        # ведёт её в центр. Подпись самой дуги ездила бы вслед за серединой дуги.
        outer = ring_box[2] / 2
        total = (2 - ring.hole) * outer  # кольца двух рядов одной толщины
        plot = (cx - total, cy - total, 2 * total, 2 * total)
        hole = ring.hole / (2 - ring.hole)
        r_mid = outer + (1 - ring.hole) * outer / 2
        angle = math.radians(ordered[0].start + 180)
        anchor = (cx + r_mid * math.sin(angle), cy - r_mid * math.cos(angle))
        labels.append(Label(target, anchor, size, style))
        # Рамка с запасом: подпись и без сдвига, и в центре целиком внутри рамки — иначе
        # ONLYOFFICE прижимает её к краю и сдвиг теряется.
        frame_parts.append(plot)
        for x, y in (anchor, target):
            frame_parts.append((x - size[0] / 2, y - size[1] / 2, size[0], size[1]))
    frame = _pad(_union(frame_parts), 0.04 * EMU_IN)
    return Assembled(
        kind="doughnut",
        frame=frame,
        plot=plot,
        categories=categories,
        names=["Доля"],
        values=[values],
        paints=[s.paint for s in ordered],
        number_format=fmt,
        basis=basis,  # type: ignore[arg-type]
        pieces=pieces,
        first_angle=ordered[0].start,
        hole=hole,
        ring_labels=labels,
    )


# ---------- кольцо-показатель, прочитанное с картинки ----------


def indicator_ring(picture: Any, reading: Any, theme: Theme | None = None) -> Assembled | None:
    """Кольцо-показатель с картинки («35%» нарисовано в центре самой картинки) — тем же
    построением, что кольцо из кусков: число встаёт в центр подписью диаграммы и меняется
    вместе с данными, а не мелкой подписью на дуге. None — не кольцо-показатель или число
    в центре по пикселям не найдено."""
    if reading.kind != "doughnut" or len(reading.series) != 1 or len(reading.slice_colors) != 2:
        return None
    points = reading.series[0].points
    if len(points) != 2 or points[0].basis != "label" or points[1].label:
        return None
    number = label_number(points[0].label or "")
    if number is None or number.unit != "%" or not reading.hole or reading.geometry.plot is None:
        return None
    try:
        image = rgba(picture.image.blob)
    except Exception:  # связанная или битая картинка
        return None
    sx, sy = picture.width / image.width, picture.height / image.height
    if abs(sx - sy) > 0.05 * max(sx, sy):
        return None  # картинка растянута — на слайде эллипс
    left, top, right, bottom = reading.geometry.plot
    inner_px = (right - left) / 2 * reading.hole
    ink = _center_ink(image, (left + right) / 2, (top + bottom) / 2, 0.9 * inner_px)
    if ink is None:
        return None
    (x0, y0, x1, y1), color = ink
    # Высота цифр — около 0,72 кегля; по ширине — не шире числа на картинке (цифра около
    # 0,56 кегля, «%» — 0,9): шрифт темы бывает плотнее тонкого шрифта картинки.
    text = (points[0].label or "").strip()
    em = sum(0.9 if ch == "%" else 0.28 if ch in ".,  " else 0.56 for ch in text) or 1.0
    by_height = (y1 - y0) * sy / EMU_PT / 0.72
    by_width = (x1 - x0) * sx / EMU_PT / em
    size_pt = round(min(by_height, by_width) * 2) / 2
    font = theme.minor_font if theme is not None else None
    style = TextStyle(size=size_pt, color=color, font=font or None)
    px, py = int(picture.left), int(picture.top)
    ring_box = (px + left * sx, py + top * sy, (right - left) * sx, (bottom - top) * sy)
    cx, cy = ring_box[0] + ring_box[2] / 2, ring_box[1] + ring_box[3] / 2
    target = (px + (x0 + x1) / 2 * sx, py + (y0 + y1) / 2 * sy)
    size = ((x1 - x0) * sx * 1.25 + 0.1 * EMU_IN, size_pt * 1.6 * EMU_PT)
    # Невидимое внешнее кольцо несёт число, как в `_ring_chart`.
    hole = float(reading.hole)
    outer = ring_box[2] / 2
    total = (2 - hole) * outer
    plot = (cx - total, cy - total, 2 * total, 2 * total)
    first_angle = float(reading.first_angle or 0.0)
    r_mid = outer + (1 - hole) * outer / 2
    angle = math.radians(first_angle + 180)
    anchor = (cx + r_mid * math.sin(angle), cy - r_mid * math.cos(angle))
    parts = [ring_box, plot]
    for x, y in (anchor, target):
        parts.append((x - size[0] / 2, y - size[1] / 2, size[0], size[1]))
    decimals = number.decimals
    value = points[0].value
    paints = [Paint((Stop(0.0, _rgb_of(c), 1.0),)) for c in reading.slice_colors]
    return Assembled(
        kind="doughnut",
        frame=_pad(_union(parts), 0.04 * EMU_IN),
        plot=plot,
        categories=["Значение", "Остаток"],
        names=["Доля"],
        values=[[value, round(100 - value, max(decimals, 2))]],
        paints=paints,
        number_format="0" + ("." + "0" * decimals if decimals else "") + '"%"',
        basis="label",
        pieces=[],
        first_angle=first_angle,
        hole=hole / (2 - hole),
        ring_labels=[Label(target, anchor, size, style)],
    )


def _rgb_of(hex_color: str) -> tuple[int, int, int]:
    h = hex_color.lstrip("#")
    return (int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _center_ink(
    image: Image.Image, cx: float, cy: float, radius: float
) -> tuple[tuple[int, int, int, int], str] | None:
    """Число в отверстии кольца: рамка тёмных (отличных от фона отверстия) пикселей и их
    цвет. None — в отверстии пусто."""
    px: Any = image.load()
    r = int(radius)
    cells: list[tuple[int, int, tuple[int, int, int, int]]] = [
        (x, y, px[x, y])
        for y in range(max(0, int(cy) - r), min(image.height, int(cy) + r + 1))
        for x in range(max(0, int(cx) - r), min(image.width, int(cx) + r + 1))
        if (x - cx) ** 2 + (y - cy) ** 2 <= r * r
    ]
    solid = [(x, y, (c[0], c[1], c[2])) for x, y, c in cells if c[3] >= 128]
    if not solid:
        return None
    # Отверстие прозрачное — всё непрозрачное в нём и есть число; иначе фон — частый цвет.
    clear = len(solid) < 0.5 * len(cells)
    back = Counter(c for _, _, c in solid).most_common(1)[0][0]

    def far(c: tuple[int, int, int]) -> int:
        if clear:
            return 255 * 3 - sum(c)  # чем темнее, тем дальше от сглаженного края
        return sum(abs(a - b) for a, b in zip(c, back, strict=True))

    ink = solid if clear else [(x, y, c) for x, y, c in solid if far(c) > 120]
    if len(ink) < 20:
        return None
    xs, ys = [x for x, _, _ in ink], [y for _, y, _ in ink]
    # Цвет — по самым непохожим на фон пикселям: края букв сглажены в сторону фона.
    ink.sort(key=lambda t: -far(t[2]))
    core = [c for _, _, c in ink[: max(1, len(ink) // 3)]]
    mean = tuple(round(sum(c[k] for c in core) / len(core)) for k in range(3))
    color = "#{:02X}{:02X}{:02X}".format(*mean)
    return (min(xs), min(ys), max(xs) + 1, max(ys) + 1), color


def _pad(box: Box, margin: float) -> Box:
    return (box[0] - margin, box[1] - margin, box[2] + 2 * margin, box[3] + 2 * margin)


# ---------- столбцы и полосы ----------


@dataclass
class Mark:
    item: Item
    paint: Paint
    base: float  # координата основания (EMU)
    end: float  # координата конца
    lo: float  # поперечные границы
    hi: float

    @property
    def length(self) -> float:
        return abs(self.end - self.base)

    @property
    def mid(self) -> float:
        return (self.lo + self.hi) / 2

    @property
    def thick(self) -> float:
        return self.hi - self.lo


Orient = Literal["up", "left", "right"]


def _mark(item: Item, orient: Orient, theme: Theme | None) -> Mark | None:
    """Кусок как столбец: картинка ровного цвета поперёк или прямоугольник с заливкой."""
    x, y, w, h = item.box
    along: Literal["down", "right"] = "down" if orient == "up" else "right"
    if item.kind == "picture":
        if _rotated_or_cropped(item.shape):
            return None
        image = _image(item)
        paint = bar_paint(image, along) if image is not None else None
    elif item.kind == "shape":
        geom = item.shape._element.find(".//" + qn("a:prstGeom"))
        if geom is None or geom.get("prst") not in BAR_GEOMETRY or _rotated_or_cropped(item.shape):
            return None
        paint = _shape_paint(item.shape, theme, along)
    else:
        return None
    if paint is None:
        return None
    if orient == "up":
        return Mark(item, paint, y + h, y, x, x + w)
    if orient == "right":
        return Mark(item, paint, x, x + w, y, y + h)
    return Mark(item, paint, x + w, x, y, y + h)


def _find_bars(
    items: list[Item], theme: Theme | None, used: set[int], slide_h: float
) -> list[Assembled]:
    found: list[tuple[Assembled, list[Item | None]]] = []
    for orient in ("up", "right", "left"):
        candidates = [
            m
            for it in items
            if it.kind in ("picture", "shape") and id(it) not in used
            if (m := _mark(it, orient, theme)) is not None
        ]
        tol = max(0.005 * slide_h, 0.02 * EMU_IN)
        groups: list[list[Mark]] = []
        for m in sorted(candidates, key=lambda m: m.base):
            for g in groups:
                if abs(g[0].base - m.base) <= tol:
                    g.append(m)
                    break
            else:
                groups.append([m])
        for g in groups:
            if any(id(m.item) in used for m in g):
                continue
            got = _bar_chart(g, orient, items, theme, used)
            if got is not None:
                found.append(got)
                used.update(id(p) for p in got[0].pieces)
    # Подписи категорий, которые нашли две группы полос («бабочка» с подписями посередине),
    # остаются надписями слайда: у каждой из диаграмм своя ось была бы лишней.
    claims: dict[int, int] = {}
    for _, cats in found:
        for it in cats:
            if it is not None:
                claims[id(it)] = claims.get(id(it), 0) + 1
    out: list[Assembled] = []
    for spec, cats in found:
        own = [it for it in cats if it is not None]
        if own and all(claims[id(it)] == 1 for it in own) and spec.category_style is not None:
            spec.pieces += own
            spec.frame = _union([spec.frame, *(it.box for it in own)])
            used.update(id(it) for it in own)
        else:
            spec.category_style = None
        spec.frame = _pad(spec.frame, 0.02 * EMU_IN)
        out.append(spec)
    return out


def _categories(marks: list[Mark]) -> list[list[Mark]] | None:
    """Столбцы по категориям: одна категория — одно место по поперечной оси (ряды одной
    категории стоят в одном месте с перекрытием или вплотную друг к другу)."""
    thick = sorted(m.thick for m in marks)[len(marks) // 2]
    marks = [m for m in marks if 0.85 * thick <= m.thick <= 1.18 * thick]
    cats: list[list[Mark]] = []
    for m in sorted(marks, key=lambda m: m.mid):
        if cats and abs(cats[-1][0].mid - m.mid) <= 0.25 * thick:
            cats[-1].append(m)
        else:
            cats.append([m])
    if len(cats) < 3:
        return None
    mids = [sum(m.mid for m in c) / len(c) for c in cats]
    steps = [b - a for a, b in pairwise(mids)]
    if min(steps) <= 0 or max(steps) > 1.35 * min(steps):
        return None
    return cats


def _bar_chart(
    group: list[Mark], orient: Orient, items: list[Item], theme: Theme | None, used: set[int]
) -> tuple[Assembled, list[Item | None]] | None:
    if len(group) < 3:
        return None
    cats = _categories(group)
    if cats is None:
        return None
    marks = [m for c in cats for m in c]
    lengths = [m.length for m in marks]
    if min(lengths) <= 0 or max(lengths) < 1.08 * min(lengths):
        return None  # все одной длины — ряд карточек или ячеек, а не диаграмма
    per_cat = max(len(c) for c in cats)
    if per_cat > 4:
        return None
    series = _series(cats, per_cat)
    if series is None:
        return None
    labels = _value_labels(marks, orient, items, theme, used)
    if len(labels) < 0.6 * len(marks):
        return None  # без чисел у концов это может быть и декор
    styles = [text_style(it.shape, theme) for it, _ in labels.values()]
    label_style = max(set(styles), key=styles.count)
    inside = sum(1 for _, pos in labels.values() if pos == "inEnd")
    position = "inEnd" if inside * 2 >= len(labels) else "outEnd"
    # Значения: подписи, если сходятся с длинами; иначе по рисунку.
    pairs = [(labels[id(m)][0].number, m.length) for m in marks if id(m) in labels]
    numbers = [n for n, _ in pairs if n is not None]
    decimals = max(n.decimals for n in numbers)
    k = sum(n.value * ln for n, ln in pairs if n) / sum(ln * ln for _, ln in pairs)
    consistent = all(
        abs(k * ln - n.value) <= max(0.04 * abs(n.value), 0.51 * 10**-n.decimals)
        for n, ln in pairs
        if n is not None
    )
    if consistent:
        basis: Literal["label", "measured"] = "label"
        fmt = _format(numbers)
    else:
        basis = "measured"
        k = max(n.value for n in numbers) / max(lengths)
        fmt = "General"
    values: list[list[float]] = []
    for s in series:
        row: list[float] = []
        for m in s:
            if m is None:
                row.append(0.0)
                continue
            lb = labels.get(id(m))
            if basis == "label" and lb is not None and lb[0].number is not None:
                row.append(lb[0].number.value)
            else:
                row.append(round(k * m.length, decimals + 1))
        values.append(row)
    # Заливка ряда — у самого длинного столбца; свои заливки у столбцов (блик в разных местах
    # полосы) — по точкам, если отличаются.
    paints = [max((m for m in s if m is not None), key=lambda m: m.length).paint for s in series]
    point_paints = [
        [m.paint if m is not None and not _same_paint(m.paint, p) else None for m in s]
        for s, p in zip(series, paints, strict=True)
    ]
    overlap, order = _overlap(series, paints)
    series = [series[i] for i in order]
    values = [values[i] for i in order]
    paints = [paints[i] for i in order]
    point_paints = [point_paints[i] for i in order]
    cat_items, cat_style = _category_labels(cats, orient, items, theme, used)
    names_cat = [" ".join(it.text.split()) if it is not None else "" for it in cat_items]
    legend_items, names = _legend(paints, items, theme, used)
    label_items = [it for it, _ in labels.values()]
    pieces = [m.item for m in marks] + label_items + legend_items
    plot = _stretch(_union(m.item.box for m in marks), orient, max(max(r) for r in values))
    # Рамка — по самому тексту подписей: сами надписи бывают много шире числа.
    frame_parts = [plot] + [text_box(it, text_style(it.shape, theme)) for it in label_items]
    if legend_items:
        frame_parts.append(_legend_box(legend_items, theme))
    frame = _union(frame_parts)
    thick = sum(m.thick for m in marks) / len(marks)
    mids = [sum(m.mid for m in c) / len(c) for c in cats]
    pitch = (mids[-1] - mids[0]) / (len(mids) - 1)
    side = 1 if overlap == 100 else len(series)
    gap = round((pitch - side * thick) / thick * 100)
    legend = None
    if legend_items:
        texts = [it for it in legend_items if it.kind == "text"]
        legend = Legend(_legend_box(legend_items, theme), text_style(texts[0].shape, theme))
    kind: Literal["column", "bar"] = "column" if orient == "up" else "bar"
    spec = Assembled(
        kind=kind,
        frame=frame,
        plot=plot,
        categories=names_cat,
        names=names,
        values=values,
        paints=paints,
        point_paints=point_paints,
        number_format=fmt,
        basis=basis,
        pieces=pieces,
        label_style=label_style,
        label_position=position,
        category_style=cat_style,
        legend=legend,
        gap=max(0, min(500, gap)),
        overlap=overlap,
        reverse=orient == "left",
    )
    return spec, cat_items


def auto_max(top: float) -> float:
    """Конец оси значений, который выбирают ONLYOFFICE и Excel при минимуме 0: наибольшее
    значение плюс 5 %, округлённое вверх до шага 1, 2 или 5 × 10ⁿ так, чтобы делений было не
    больше десяти (10 → 12, 9,1 → 10, 58 → 70, 100 → 120, 3,3 → 3,5 — замер 24.09.2026)."""
    if top <= 0:
        return 1.0
    target = top * 1.05
    exp = math.floor(math.log10(target)) - 2
    while True:
        for base in (1, 2, 5):
            major = base * 10.0**exp
            n = math.ceil(target / major - 1e-9)
            if n <= 10:
                return round(n * major, 10)
        exp += 1


def _stretch(plot: Box, orient: Orient, top: float) -> Box:
    """Область построения, продолжённая от основания так, чтобы при автоматическом конце оси
    самый длинный столбец (значение `top`) кончался там же, где на рисунке."""
    k = auto_max(top) / top if top > 0 else 1.0
    x, y, w, h = plot
    if orient == "up":
        return (x, y + h - h * k, w, h * k)
    if orient == "right":
        return (x, y, w * k, h)
    return (x + w - w * k, y, w * k, h)


def _legend_box(items: list[Item], theme: Theme | None) -> Box:
    """Рамка легенды: маркеры и сам текст надписей (с запасом по ширине на перенос)."""
    boxes: list[Box] = []
    for it in items:
        if it.kind == "text":
            x, y, w, h = text_box(it, text_style(it.shape, theme))
            boxes.append((x, y, w * 1.4 + 0.1 * EMU_IN, h))
        else:
            boxes.append(it.box)
    return _union(boxes)


def _format(numbers: list[LabelNumber]) -> str:
    decimals = max(n.decimals for n in numbers)
    base = "#,##0" if any(n.thousands for n in numbers) else "0"
    if decimals:
        base += "." + "0" * decimals
    units = {n.unit for n in numbers if n.unit}
    if len(units) == 1:
        unit = units.pop()
        return base + f'"{unit}"' if unit == "%" else base + f'" {unit}"'
    return base


def _series(cats: list[list[Mark]], per_cat: int) -> list[list[Mark | None]] | None:
    """Ряды по цвету: в каждой категории не больше одного столбца ряда."""
    if per_cat == 1:
        return [[c[0] for c in cats]]
    seeds = next(c for c in cats if len(c) == per_cat)
    colors = [m.paint.main for m in seeds]
    if min(_distance(a, b) for i, a in enumerate(colors) for b in colors[i + 1 :]) < 25:
        return None
    out: list[list[Mark | None]] = [[] for _ in range(per_cat)]
    for c in cats:
        slots: list[Mark | None] = [None] * per_cat
        for m in c:
            k = min(range(per_cat), key=lambda i: _distance(m.paint.main, colors[i]))
            if slots[k] is not None or _distance(m.paint.main, colors[k]) > 60:
                return None
            slots[k] = m
        for k in range(per_cat):
            out[k].append(slots[k])
    return out


def _overlap(series: list[list[Mark | None]], paints: list[Paint]) -> tuple[int, list[int]]:
    """Перекрытие рядов и порядок отрисовки. Ряды одной категории на одном месте — перекрытие
    100. На картинке спереди обычно короткий столбец, а в нативной диаграмме порядок рядов
    один на все категории: берётся порядок, при котором передний столбец не длиннее заднего
    в наибольшем числе категорий; при равенстве полупрозрачный ряд — спереди."""
    n = len(series)
    if n == 1:
        return 0, [0]
    full = [c for c in zip(*series, strict=True) if all(m is not None for m in c)]
    if not full:
        return 0, list(range(n))
    same = all(
        max(m.mid for m in c if m) - min(m.mid for m in c if m) <= 0.25 * c[0].thick for c in full
    )
    if not same:
        return 0, sorted(range(n), key=lambda i: next(m.mid for m in series[i] if m))
    translucent = [min(s.alpha for s in p.stops) < 0.95 for p in paints]

    def score(order: tuple[int, ...]) -> tuple[int, int]:
        fits = sum(
            1
            for c in full
            if all(
                c[order[a]].length >= c[order[b]].length - 1
                for a in range(n)
                for b in range(a + 1, n)
            )
        )
        front = sum(k for k, i in enumerate(order) if translucent[i])
        return fits, front

    best = max(permutations(range(n)), key=score)
    return 100, list(best)


def _same_paint(a: Paint, b: Paint) -> bool:
    if len(a.stops) != len(b.stops):
        return False
    return all(
        abs(x.pos - y.pos) <= 0.08
        and _distance(x.color, y.color) <= 10
        and abs(x.alpha - y.alpha) <= 0.06
        for x, y in zip(a.stops, b.stops, strict=True)
    )


def _value_labels(
    marks: list[Mark], orient: Orient, items: list[Item], theme: Theme | None, used: set[int]
) -> dict[int, tuple[Item, str]]:
    """Числа у концов столбцов: центр текста в пределах толщины столбца и недалеко от
    конца — внутри (inEnd) или снаружи (outEnd)."""
    found: dict[int, tuple[Item, str]] = {}
    taken: set[int] = set()
    numbers = [it for it in items if it.kind == "text" and it.number and id(it) not in used]
    for m in sorted(marks, key=lambda m: -m.length):
        best: tuple[float, Item, str] | None = None
        for it in numbers:
            if id(it) in taken:
                continue
            style = text_style(it.shape, theme)
            tb = text_box(it, style)
            tx, ty = tb[0] + tb[2] / 2, tb[1] + tb[3] / 2
            along, across = (ty, tx) if orient == "up" else (tx, ty)
            if not m.lo - 0.15 * m.thick <= across <= m.hi + 0.15 * m.thick:
                continue
            reach = max(1.6 * style.size * EMU_PT, 0.12 * EMU_IN)
            direction = 1 if m.end > m.base else -1  # куда растёт столбец
            offset = (along - m.end) * direction  # > 0 — за концом, < 0 — внутри
            if not -reach <= offset <= reach:
                continue
            score = abs(offset)
            if best is None or score < best[0]:
                best = (score, it, "outEnd" if offset > 0 else "inEnd")
        if best is not None:
            found[id(m)] = (best[1], best[2])
            taken.add(id(best[1]))
    return found


def _category_labels(
    cats: list[list[Mark]], orient: Orient, items: list[Item], theme: Theme | None, used: set[int]
) -> tuple[list[Item | None], TextStyle | None]:
    """Подписи категорий за основанием, по одной на место."""
    base = cats[0][0].base
    mids = [sum(m.mid for m in c) / len(c) for c in cats]
    pitch = (mids[-1] - mids[0]) / max(1, len(mids) - 1)
    texts = [it for it in items if it.kind == "text" and id(it) not in used]
    out: list[Item | None] = []
    for mid in mids:
        best: tuple[float, Item] | None = None
        for it in texts:
            if orient == "up":
                across, near = it.cx, it.box[1] - base  # верх подписи под основанием
                ok = -0.05 * EMU_IN <= near <= 3 * it.box[3]
            elif orient == "right":
                across, near = it.cy, base - (it.box[0] + it.box[2])
                ok = -0.6 * EMU_IN <= near <= 1.5 * EMU_IN
            else:
                across, near = it.cy, it.box[0] - base
                ok = -0.6 * EMU_IN <= near <= 1.5 * EMU_IN
            if not ok or abs(across - mid) > 0.45 * pitch:
                continue
            score = abs(across - mid)
            if best is None or score < best[0]:
                best = (score, it)
        out.append(best[1] if best else None)
    if sum(1 for it in out if it is not None) < 0.7 * len(out):
        return [None] * len(out), None
    styles = [text_style(it.shape, theme) for it in out if it is not None]
    return out, max(set(styles), key=styles.count)


def _legend(
    paints: list[Paint], items: list[Item], theme: Theme | None, used: set[int]
) -> tuple[list[Item], list[str]]:
    """Легенда, нарисованная фигурами: маленький маркер цвета ряда и надпись справа от него.
    Если нашлась у каждого ряда — маркеры и надписи уходят в легенду диаграммы."""
    names = [f"Ряд {k + 1}" for k in range(len(paints))]
    limit = 0.35 * EMU_IN
    markers = []
    for it in items:
        if it.kind != "shape" or id(it) in used or it.box[2] > limit or it.box[3] > limit:
            continue
        geom = it.shape._element.find(".//" + qn("a:prstGeom"))
        if geom is None or geom.get("prst") not in MARKER_GEOMETRY:
            continue
        paint = _shape_paint(it.shape, theme, "down")
        if paint is not None and paint.solid:
            markers.append((it, paint.main))
    texts = [it for it in items if it.kind == "text" and id(it) not in used and it.number is None]
    chosen: list[Item] = []
    for k, paint in enumerate(paints):
        best = None
        for marker, color in markers:
            if _distance(color, paint.main) > 60 or marker in chosen:
                continue
            mid = marker.cy
            right = marker.box[0] + marker.box[2]
            near = [
                (t.box[0] - right, t)
                for t in texts
                if t.box[1] <= mid <= t.box[1] + t.box[3]
                and -marker.box[2] < t.box[0] - right < 8 * marker.box[2]
            ]
            if near:
                text = min(near, key=lambda p: p[0])[1]
                if best is None or _distance(color, paint.main) < best[0]:
                    best = (_distance(color, paint.main), marker, text)
        if best is None:
            return [], names
        names[k] = " ".join(best[2].text.split())
        chosen += [best[1], best[2]]
    return chosen, names


# ---------- постройка ----------


def _el(tag: str, **attrs: Any) -> Any:
    el = OxmlElement(tag)
    for k, v in attrs.items():
        el.set(k, str(v))
    return el


def _color_el(rgb: tuple[int, int, int], alpha: float) -> Any:
    c = _el("a:srgbClr", val=_hex(rgb).lstrip("#"))
    if alpha < 0.995:
        c.append(_el("a:alpha", val=max(0, round(alpha * 100000))))
    return c


def _paint_sp(paint: Paint) -> Any:
    """c:spPr с заливкой куска и без обводки."""
    sp = _el("c:spPr")
    if paint.solid:
        fill = _el("a:solidFill")
        fill.append(_color_el(paint.stops[0].color, paint.stops[0].alpha))
    else:
        fill = _el("a:gradFill", rotWithShape="1")
        lst = _el("a:gsLst")
        for s in paint.stops:
            gs = _el("a:gs", pos=round(s.pos * 100000))
            gs.append(_color_el(s.color, s.alpha))
            lst.append(gs)
        fill.append(lst)
        fill.append(_el("a:lin", ang=round(paint.angle % 360 * 60000), scaled="0"))
    sp.append(fill)
    ln = _el("a:ln")
    ln.append(_el("a:noFill"))
    sp.append(ln)
    return sp


def _set_sp(parent: Any, sp: Any, before: tuple[str, ...]) -> None:
    old = parent.find(qn("c:spPr"))
    if old is not None:
        parent.remove(old)
    for tag in before:
        anchor = parent.find(qn(tag))
        if anchor is not None:
            anchor.addprevious(sp)
            return
    parent.append(sp)


def _tx_pr(style: TextStyle) -> Any:
    tx = _el("c:txPr")
    tx.append(_el("a:bodyPr"))
    tx.append(_el("a:lstStyle"))
    p = _el("a:p")
    ppr = _el("a:pPr")
    d = _el("a:defRPr", sz=round(style.size * 100), b="1" if style.bold else "0")
    fill = _el("a:solidFill")
    fill.append(_el("a:srgbClr", val=style.color.lstrip("#").upper()))
    d.append(fill)
    if style.font:
        d.append(_el("a:latin", typeface=style.font))
        d.append(_el("a:cs", typeface=style.font))
    ppr.append(d)
    p.append(ppr)
    p.append(_el("a:endParaRPr", lang="ru-RU"))
    tx.append(p)
    return tx


def _no_fill(parent: Any, after: Any = None) -> None:
    sp = parent.find(qn("c:spPr"))
    if sp is None:
        sp = _el("c:spPr")
        if after is not None:
            after.addnext(sp)
        else:
            parent.append(sp)
    for child in list(sp):
        sp.remove(child)
    sp.append(_el("a:noFill"))
    ln = _el("a:ln")
    ln.append(_el("a:noFill"))
    sp.append(ln)


def _manual(parent: Any, box: Box, frame: Box, *, inner: bool) -> None:
    """Ручная раскладка элемента (область построения, легенда) по рамке на слайде."""
    layout = parent.find(qn("c:layout"))
    if layout is None:
        layout = _el("c:layout")
        if parent.tag == qn("c:plotArea"):
            parent.insert(0, layout)
        else:
            _after_first(parent, layout)
    for child in list(layout):
        layout.remove(child)
    m = _el("c:manualLayout")
    if inner:
        m.append(_el("c:layoutTarget", val="inner"))
    m.append(_el("c:xMode", val="edge"))
    m.append(_el("c:yMode", val="edge"))
    fx, fy, fw, fh = frame
    for tag, v in (
        ("c:x", (box[0] - fx) / fw),
        ("c:y", (box[1] - fy) / fh),
        ("c:w", box[2] / fw),
        ("c:h", box[3] / fh),
    ):
        m.append(_el(tag, val=f"{max(0.0, min(1.0, v)):.4f}"))
    layout.append(m)


def _after_first(parent: Any, el: Any) -> None:
    """c:layout легенды — сразу за c:legendPos/c:legendEntry (порядок схемы CT_Legend)."""
    anchor = None
    for tag in ("c:legendEntry", "c:legendPos"):
        found = parent.findall(qn(tag))
        if found:
            anchor = found[-1]
            break
    if anchor is not None:
        anchor.addnext(el)
    else:
        parent.insert(0, el)


def _chart_text(chart: Any, style: TextStyle | None) -> None:
    """Шрифт диаграммы по умолчанию — как у надписей слайда."""
    if style is None:
        return
    chart.font.size = Pt(style.size)
    if style.font:
        chart.font.name = style.font


def build(slide: Any, spec: Assembled) -> Any:
    """Нативная диаграмма по найденному набору; возвращает graphicFrame."""
    if spec.kind == "doughnut":
        return _build_doughnut(slide, spec)
    return _build_bars(slide, spec)


def _frame(spec: Assembled) -> tuple[int, int, int, int]:
    x, y, w, h = spec.frame
    return (round(x), round(y), max(1, round(w)), max(1, round(h)))


def _build_doughnut(slide: Any, spec: Assembled) -> Any:
    data = CategoryChartData()  # type: ignore[no-untyped-call]
    data.categories = spec.categories
    data.add_series(spec.names[0], spec.values[0])  # type: ignore[no-untyped-call]
    if spec.ring_labels:
        # Невидимое внешнее кольцо — полный круг с тем же числом (см. `_ring_chart`).
        data.add_series(CENTER_SERIES, [max(spec.values[0][0], 0.001), 0])  # type: ignore[no-untyped-call]
    frame = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, *_frame(spec), data)
    chart = frame.chart
    chart.has_legend = False
    chart.has_title = False
    _chart_text(chart, spec.ring_labels[0].style if spec.ring_labels else None)
    space = chart._chartSpace
    _no_fill(space, after=space.chart)
    plot_area = space.chart.plotArea
    _no_fill(plot_area)
    _manual(plot_area, spec.plot, spec.frame, inner=True)
    group = next(el for el in plot_area if el.tag == qn("c:doughnutChart"))
    group.find(qn("c:firstSliceAng")).set("val", str(round(spec.first_angle) % 360))
    group.find(qn("c:holeSize")).set("val", str(min(90, max(10, round(spec.hole * 100)))))
    visible, *helper = chart.plots[0].series
    for k, paint in enumerate(spec.paints):
        visible.points[k].format.fill.solid()  # создаёт c:dPt
        dpt = visible._element.findall(qn("c:dPt"))[k]
        _set_sp(dpt, _paint_sp(paint), ("c:pictureOptions", "c:extLst"))
    if helper:
        ser = helper[0]._element
        for k in range(len(spec.categories)):
            helper[0].points[k].format.fill.background()
            helper[0].points[k].format.line.fill.background()
        ser.find(qn("c:cat")).addprevious(_ring_labels(spec))
    if len(spec.values[0]) == 2:
        # Остаток — формула: правится одно число, кольцо пересчитывается само.
        chart.part.chart_workbook.update_from_xlsx_blob(_ring_workbook(spec))
    explicit_booleans(chart)
    return frame


CENTER_SERIES = "Число в центре"


def _ring_labels(spec: Assembled) -> Any:
    """Подписи невидимого кольца: у полного круга — сдвиг туда, где стоял текст, и размер
    (без размера ONLYOFFICE ломает крупное число по знакам); пустая точка — без подписи."""
    _, _, fw, fh = spec.frame
    (lb,) = spec.ring_labels
    dls = _el("c:dLbls")
    for k in range(len(spec.categories)):
        d = _el("c:dLbl")
        d.append(_el("c:idx", val=k))
        if k:
            d.append(_el("c:delete", val="1"))
            dls.append(d)
            continue
        layout = _el("c:layout")
        m = _el("c:manualLayout")
        m.append(_el("c:x", val=f"{(lb.target[0] - lb.anchor[0]) / fw:.4f}"))
        m.append(_el("c:y", val=f"{(lb.target[1] - lb.anchor[1]) / fh:.4f}"))
        m.append(_el("c:w", val=f"{lb.size[0] / fw:.4f}"))
        m.append(_el("c:h", val=f"{lb.size[1] / fh:.4f}"))
        layout.append(m)
        d.append(layout)
        d.append(_el("c:numFmt", formatCode=spec.number_format, sourceLinked="0"))
        sp = _el("c:spPr")
        sp.append(_el("a:noFill"))
        ln = _el("a:ln")
        ln.append(_el("a:noFill"))
        sp.append(ln)
        d.append(sp)
        d.append(_tx_pr(lb.style))
        for tag, v in _SHOW_VALUE:
            d.append(_el(tag, val=v))
        dls.append(d)
    for tag, _ in _SHOW_VALUE:
        dls.append(_el(tag, val="0"))
    return dls


_SHOW_VALUE = (
    ("c:showLegendKey", "0"),
    ("c:showVal", "1"),
    ("c:showCatName", "0"),
    ("c:showSerName", "0"),
    ("c:showPercent", "0"),
    ("c:showBubbleSize", "0"),
)


def _ring_workbook(spec: Assembled) -> bytes:
    """Книга кольца: B — доли (остаток — формула), C — число в центре (формула от B2)."""
    buf = io.BytesIO()
    book = xlsxwriter.Workbook(buf, {"in_memory": True})
    sheet = book.add_worksheet("Sheet1")
    value = spec.values[0][0]
    sheet.write(0, 1, spec.names[0])
    sheet.write(1, 0, spec.categories[0])
    sheet.write(2, 0, spec.categories[1])
    sheet.write_number(1, 1, value)
    sheet.write_formula(2, 1, "=100-B2", None, round(100 - value, 6))
    if spec.ring_labels:
        sheet.write(0, 2, CENTER_SERIES)
        sheet.write_formula(1, 2, "=MAX(B2,0.001)", None, max(value, 0.001))
        sheet.write_number(2, 2, 0)
    book.close()
    return buf.getvalue()


def _build_bars(slide: Any, spec: Assembled) -> Any:
    data = CategoryChartData()  # type: ignore[no-untyped-call]
    data.categories = spec.categories
    for name, row in zip(spec.names, spec.values, strict=True):
        data.add_series(name, row)  # type: ignore[no-untyped-call]
    kind = XL_CHART_TYPE.COLUMN_CLUSTERED if spec.kind == "column" else XL_CHART_TYPE.BAR_CLUSTERED
    frame = slide.shapes.add_chart(kind, *_frame(spec), data)
    chart = frame.chart
    chart.has_title = False
    _chart_text(chart, spec.category_style or spec.label_style)
    space = chart._chartSpace
    _no_fill(space, after=space.chart)
    plot_area = space.chart.plotArea
    _no_fill(plot_area)
    plot = chart.plots[0]
    plot.gap_width = spec.gap
    plot.overlap = spec.overlap
    for k, (series, paint) in enumerate(zip(plot.series, spec.paints, strict=True)):
        _set_sp(
            series._element,
            _paint_sp(paint),
            ("c:invertIfNegative", "c:pictureOptions", "c:dPt", "c:dLbls", "c:cat", "c:val"),
        )
        own = spec.point_paints[k] if k < len(spec.point_paints) else []
        for idx, point_paint in enumerate(own):
            if point_paint is None:
                continue
            series.points[idx].format.fill.solid()  # создаёт c:dPt точки
            dpt = next(
                d
                for d in series._element.findall(qn("c:dPt"))
                if d.find(qn("c:idx")).get("val") == str(idx)
            )
            _set_sp(dpt, _paint_sp(point_paint), ("c:pictureOptions", "c:extLst"))
    if spec.label_style is not None:
        plot.has_data_labels = True
        labels = plot.data_labels
        labels.number_format = spec.number_format
        labels.number_format_is_linked = False
        labels.show_value = True
        position = spec.label_position
        if spec.reverse and position == "inEnd":
            # При развёрнутой оси ONLYOFFICE зеркалит сдвиг подписи: «внутри у конца» рисует
            # за концом полосы (белое на светлом фоне пропадает). По центру полосы подпись
            # стоит одинаково в ONLYOFFICE и PowerPoint.
            position = "ctr"
        labels.position = LABEL_POSITIONS.get(position, XL_LABEL_POSITION.OUTSIDE_END)
        dls = plot._element.find(qn("c:dLbls"))
        old = dls.find(qn("c:txPr"))
        if old is not None:
            dls.remove(old)
        anchor = dls.find(qn("c:dLblPos"))
        if anchor is None:
            anchor = dls.find(qn("c:showLegendKey"))
        anchor.addprevious(_tx_pr(spec.label_style))
        if position == "ctr" and spec.label_position == "inEnd":
            _contrast_labels(plot, spec, position)
    cat_ax, val_ax = chart.category_axis, chart.value_axis
    for axis in (cat_ax, val_ax):
        axis.has_major_gridlines = False
        axis.has_minor_gridlines = False
        axis.major_tick_mark = XL_TICK_MARK.NONE
        axis.minor_tick_mark = XL_TICK_MARK.NONE
        axis.format.line.fill.background()
    if spec.category_style is not None:
        cat_ax.tick_labels.font.size = Pt(spec.category_style.size)
        cat_ax.tick_labels.font.color.rgb = _rgb(spec.category_style.color)
        if spec.category_style.font:
            cat_ax.tick_labels.font.name = spec.category_style.font
    else:
        cat_ax.visible = False
    if spec.kind == "bar":
        cat_ax.reverse_order = True  # полосы сверху вниз в порядке слайда
    # Конец оси — автоматический: впишут числа крупнее — столбцы перемасштабируются, а не
    # упрутся в потолок. Совпадение с рисунком даёт область построения (см. `_stretch`).
    val_ax.visible = False
    val_ax.minimum_scale = 0
    if spec.reverse:
        val_ax.reverse_order = True
    if spec.kind == "bar":
        # Ось значений полос — внизу при обратном порядке категорий.
        _set_crosses(val_ax._element, "max")
    _manual(plot_area, spec.plot, spec.frame, inner=True)
    chart.has_legend = spec.legend is not None
    if spec.legend is not None:
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(spec.legend.style.size)
        chart.legend.font.color.rgb = _rgb(spec.legend.style.color)
        if spec.legend.style.font:
            chart.legend.font.name = spec.legend.style.font
        _manual(chart.legend._element, spec.legend.box, spec.frame, inner=False)
    explicit_booleans(chart)
    return frame


def _contrast_labels(plot: Any, spec: Assembled, position: str) -> None:
    """Подписи, перенесённые с конца полосы в её середину: светлые (белые) числа на светлой
    середине — блик градиента — не видны. Таким точкам — цвет самой полосы; цвет середины
    при правке значения не меняется: градиент тянется вместе с полосой."""
    style = spec.label_style
    if style is None or _luminance(_hex_rgb(style.color)) < 0.6:
        return
    for k, series in enumerate(plot.series):
        own = spec.point_paints[k] if k < len(spec.point_paints) else []
        dark: list[tuple[int, Paint]] = []
        for idx in range(len(spec.categories)):
            paint = (own[idx] if idx < len(own) else None) or spec.paints[k]
            if _luminance(_blend(paint, 0.5)) > 0.6:
                dark.append((idx, paint))
        if not dark:
            continue
        dls = _el("c:dLbls")
        for idx, paint in dark:
            d = _el("c:dLbl")
            d.append(_el("c:idx", val=idx))
            d.append(_el("c:numFmt", formatCode=spec.number_format, sourceLinked="0"))
            sp = _el("c:spPr")
            sp.append(_el("a:noFill"))
            ln = _el("a:ln")
            ln.append(_el("a:noFill"))
            sp.append(ln)
            d.append(sp)
            d.append(_tx_pr(TextStyle(style.size, _hex(paint.main), style.font, style.bold)))
            d.append(_el("c:dLblPos", val=position))
            for tag, v in _SHOW_VALUE:
                d.append(_el(tag, val=v))
            dls.append(d)
        dls.append(_el("c:numFmt", formatCode=spec.number_format, sourceLinked="0"))
        dls.append(_tx_pr(style))
        dls.append(_el("c:dLblPos", val=position))
        for tag, v in _SHOW_VALUE:
            dls.append(_el(tag, val=v))
        ser = series._element
        before = next(
            (
                ser.find(qn(t))
                for t in ("c:trendline", "c:errBars", "c:cat", "c:val")
                if ser.find(qn(t)) is not None
            ),
            None,
        )
        if before is not None:
            before.addprevious(dls)
        else:
            ser.append(dls)


def _blend(paint: Paint, t: float) -> tuple[int, int, int]:
    """Видимый цвет заливки в точке `t` (0–1 вдоль градиента) на белом фоне."""
    stops = paint.stops
    lo = max((s for s in stops if s.pos <= t), key=lambda s: s.pos, default=stops[0])
    hi = min((s for s in stops if s.pos >= t), key=lambda s: s.pos, default=stops[-1])
    k = 0.0 if hi.pos == lo.pos else (t - lo.pos) / (hi.pos - lo.pos)
    alpha = lo.alpha + (hi.alpha - lo.alpha) * k
    out = []
    for i in range(3):
        c = lo.color[i] + (hi.color[i] - lo.color[i]) * k
        out.append(round(c * alpha + 255 * (1 - alpha)))
    return (out[0], out[1], out[2])


def _luminance(rgb: tuple[int, int, int]) -> float:
    return (0.2126 * rgb[0] + 0.7152 * rgb[1] + 0.0722 * rgb[2]) / 255


def _hex_rgb(value: str) -> tuple[int, int, int]:
    h = value.lstrip("#")
    return (int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _set_crosses(axis: Any, value: str) -> None:
    el = axis.find(qn("c:crosses"))
    if el is None:
        el = _el("c:crosses")
        axis.find(qn("c:crossAx")).addnext(el)
    el.set("val", value)


def _rgb(hex_color: str) -> Any:
    from pptx.dml.color import RGBColor

    return RGBColor.from_string(hex_color.lstrip("#").upper())  # type: ignore[no-untyped-call]


# ---------- замена на слайде ----------


def find_charts(slide: Any) -> list[Assembled]:
    """Диаграммы из кусков на слайде: сначала кольца, потом столбцы и полосы."""
    items = _items(slide)
    theme = _slide_theme(slide)
    used: set[int] = set()
    slide_h = float(slide.part.package.presentation_part.presentation.slide_height or 0)
    return _find_rings(items, theme, used) + _find_bars(items, theme, used, slide_h)


def describe(spec: Assembled) -> str:
    if spec.basis == "label":
        return "Диаграмма собрана из фигур слайда: значения перенесены из подписей."
    return (
        "Диаграмма собрана из фигур слайда: числа на слайде не совпадали с рисунком, "
        "значения сняты по рисунку — проверьте перед показом."
    )


def swap(slide: Any, spec: Assembled) -> Any:
    """Куски → нативная диаграмма на месте нижнего куска. Если постройка сорвалась, слайд
    остаётся как был: недостроенная диаграмма и её часть убираются."""
    before = set(map(id, slide.shapes._spTree))
    try:
        frame = build(slide, spec)
    except Exception:
        for el in list(slide.shapes._spTree):
            if id(el) not in before and el.tag == qn("p:graphicFrame"):
                for rid in el.xpath(".//c:chart/@r:id"):
                    slide.part.drop_rel(str(rid))
                el.getparent().remove(el)
        raise
    anchor = min(spec.pieces, key=lambda p: p.z).shape._element
    anchor.addprevious(frame._element)
    c_nv_pr = frame._element.nvGraphicFramePr.cNvPr
    c_nv_pr.set("name", f"Диаграмма {frame.shape_id} (из фигур)")
    c_nv_pr.set("descr", describe(spec))
    rids: list[str] = []
    for piece in spec.pieces:
        el = piece.shape._element
        if piece.kind == "picture" and el.blipFill.blip is not None:
            rid = el.blipFill.blip.rEmbed
            if rid:
                rids.append(rid)
        el.getparent().remove(el)
    refs = {str(r) for r in slide._element.xpath("//@r:embed | //@r:link | //@r:id")}
    for rid in set(rids):
        if rid not in refs:
            slide.part.rels.pop(rid)
    return frame


def swap_composites(prs: Any, *, slides: Iterable[Any] | None = None) -> list[ChartSwap]:
    """Все диаграммы из кусков на слайдах → нативные диаграммы; отчёт по каждой."""
    out: list[ChartSwap] = []
    order = {s.slide_id: k for k, s in enumerate(prs.slides, start=1)}
    for slide in slides if slides is not None else prs.slides:
        number = order.get(slide.slide_id, 0)
        try:
            found = find_charts(slide)
        except Exception as e:  # поиск необязателен: слайд остаётся как был
            reason = f"не разобраны: {e}"[:200]
            out.append(ChartSwap(number, "фигуры слайда", "", "kept", reason, kind="pieces"))
            continue
        for spec in found:
            name = f"{_KIND_NAMES[spec.kind]} из фигур"
            try:
                swap(slide, spec)
            except Exception as e:  # замена необязательна: куски остаются как были
                reason = f"не построена: {e}"[:200]
                out.append(ChartSwap(number, name, "", "kept", reason, kind="pieces"))
                continue
            out.append(
                ChartSwap(number, name, "", "replaced", measured=spec.measured, kind="pieces")
            )
    return out


_KIND_NAMES = {"doughnut": "кольцо", "column": "столбцы", "bar": "полосы"}

__all__ = [
    "Assembled",
    "build",
    "describe",
    "find_charts",
    "indicator_ring",
    "swap",
    "swap_composites",
]
