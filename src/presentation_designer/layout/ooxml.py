"""Операции над пакетом PPTX, проверенные на этапе 0A: клонирование слайда и диаграммы,
замена текста с сохранением оформления, удаление слайдов, проверка связей пакета.

python-pptx не даёт публичного API для клонирования слайда, поэтому копия собирается из
сериализованного XML исходной части и её связей: неизменяемые ресурсы (картинки, макет)
разделяются, изменяемые (диаграмма и её книга данных) получают независимые копии.
Идентификаторы связей в копии выдаёт python-pptx, ссылки в XML переписываются по карте.
Удалённые слайды и их исключительные ресурсы не попадают в файл: при сохранении
python-pptx записывает только части, достижимые от корня пакета.
"""

from __future__ import annotations

import copy
import posixpath
import re
import zipfile
from dataclasses import dataclass, field
from typing import Any

from lxml import etree
from pptx.opc.constants import CONTENT_TYPE as CT
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part, XmlPart
from pptx.parts.chart import ChartPart
from pptx.parts.embeddedpackage import EmbeddedXlsxPart
from pptx.parts.slide import SlidePart

NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
NS_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
NS_CT = "http://schemas.openxmlformats.org/package/2006/content-types"

# Атрибуты, в которых части ссылаются на свои связи.
_REL_ATTRS = ("id", "embed", "link", "pict", "dm", "lo", "qs", "cs", "href")
_PARTNAME_NUMBER = re.compile(r"^(?P<prefix>.*?)(?P<num>\d+)(?P<ext>\.\w+)$")
_XML_PARSER = etree.XMLParser(remove_blank_text=False, resolve_entities=False, huge_tree=True)


# --- клонирование -------------------------------------------------------------------


def clone_slide(prs: Any, source: Any, index: int | None = None) -> Any:
    """Клонирует слайд `source` в ту же презентацию и возвращает новый слайд.

    Копируются фигуры, группы, обрезка картинок, стили текста, связи с макетом и
    ресурсами; заметки докладчика не копируются. Диаграммы клонируются вместе с книгами
    данных, чтобы копии редактировались независимо. `index` — позиция в колоде.
    """
    package = prs.part.package
    src_part = source.part
    new_part = SlidePart.load(
        package.next_partname("/ppt/slides/slide%d.xml"), CT.PML_SLIDE, package, src_part.blob
    )
    _copy_rels(new_part, src_part, skip_reltypes={RT.NOTES_SLIDE})
    rid = prs.part.relate_to(new_part, RT.SLIDE)
    sld_id_lst = prs.slides._sldIdLst  # публичного API для списка слайдов у python-pptx нет
    sld_id = sld_id_lst.add_sldId(rid)
    if index is not None:
        sld_id_lst.remove(sld_id)
        sld_id_lst.insert(index, sld_id)
    return prs.slides.get(int(sld_id.id))


def clone_chart_part(package: Any, chart_part: Any) -> Any:
    """Независимая копия части диаграммы с её книгой данных и стилевыми частями."""
    new_part = ChartPart.load(
        package.next_partname(ChartPart.partname_template), CT.DML_CHART, package, chart_part.blob
    )
    _copy_rels(new_part, chart_part, skip_reltypes=set())
    return new_part


def _copy_rels(new_part: Any, src_part: Any, skip_reltypes: set[str]) -> None:
    """Переносит связи `src_part` в `new_part`, переписывая идентификаторы в XML."""
    package = new_part.package
    # Одна карта старый → новый идентификатор; пропущенные связи получают "" и удаляются
    # тем же проходом, иначе удаление могло бы задеть уже переписанные новые идентификаторы.
    mapping: dict[str, str] = dict.fromkeys(src_part.rels.keys(), "")
    for old_rid, rel in src_part.rels.items():
        if rel.reltype in skip_reltypes:
            continue
        if rel.is_external:
            mapping[old_rid] = new_part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
            continue
        target = rel.target_part
        if rel.reltype == RT.CHART:
            target = clone_chart_part(package, target)
        elif rel.reltype == RT.PACKAGE and isinstance(target, EmbeddedXlsxPart):
            target = EmbeddedXlsxPart.load(
                package.next_partname(EmbeddedXlsxPart.partname_template),
                target.content_type,
                package,
                target.blob,
            )
        elif isinstance(src_part, ChartPart):
            # Стиль, цвета и переопределение темы диаграммы принадлежат ей одной.
            target = Part.load(
                _next_partname_like(package, target.partname),
                target.content_type,
                package,
                target.blob,
            )
        mapping[old_rid] = new_part.relate_to(target, rel.reltype)
    _remap_rel_ids(new_part._element, mapping)


def _remap_rel_ids(element: Any, mapping: dict[str, str]) -> None:
    """Переписывает ссылки `r:*` по карте за один проход; пустое значение удаляет атрибут."""
    for node in element.iter():
        for name in _REL_ATTRS:
            key = f"{{{NS_R}}}{name}"
            value = node.get(key)
            if value is None or value not in mapping:
                continue
            if mapping[value] == "":
                del node.attrib[key]
            else:
                node.set(key, mapping[value])


def _next_partname_like(package: Any, partname: str) -> Any:
    m = _PARTNAME_NUMBER.match(str(partname))
    if not m:
        return package.next_partname(str(partname).replace(".", "%d.", 1))
    return package.next_partname(f"{m['prefix']}%d{m['ext']}")


# --- текст --------------------------------------------------------------------------


def replace_paragraph_text(paragraph: Any, text: str) -> None:
    """Меняет текст абзаца, сохраняя оформление первого фрагмента и свойства абзаца.

    Остальные фрагменты удаляются; переносы строк внутри абзаца задаются `\\n`
    и становятся `a:br` с тем же оформлением.
    """
    p = paragraph._p
    runs = p.findall(f"{{{NS_A}}}r")
    if runs:
        template = runs[0]
        rpr = template.find(f"{{{NS_A}}}rPr")
    else:
        template = None
        end = p.find(f"{{{NS_A}}}endParaRPr")
        rpr = None
        if end is not None:
            rpr = copy.deepcopy(end)
            rpr.tag = f"{{{NS_A}}}rPr"
    for child in list(p):
        if child.tag in (f"{{{NS_A}}}r", f"{{{NS_A}}}br", f"{{{NS_A}}}fld"):
            p.remove(child)
    end_para = p.find(f"{{{NS_A}}}endParaRPr")
    insert_at = list(p).index(end_para) if end_para is not None else len(p)
    for i, line in enumerate(text.split("\n")):
        if i > 0:
            br = etree.SubElement(p, f"{{{NS_A}}}br")
            if rpr is not None:
                br.append(copy.deepcopy(rpr))
            p.remove(br)
            p.insert(insert_at, br)
            insert_at += 1
        r = etree.Element(f"{{{NS_A}}}r")
        if rpr is not None:
            r.append(copy.deepcopy(rpr))
        t = etree.SubElement(r, f"{{{NS_A}}}t")
        t.text = line
        p.insert(insert_at, r)
        insert_at += 1


def set_run_texts(paragraph: Any, texts: list[str]) -> None:
    """Подставляет текст в существующие фрагменты по порядку: оформление каждого сохраняется.

    Лишние фрагменты получают пустой текст, лишние строки не добавляются.
    """
    runs = list(paragraph.runs)
    for run, text in zip(runs, texts, strict=False):
        run.text = text
    for run in runs[len(texts) :]:
        run.text = ""


def paragraph_run_styles(paragraph: Any) -> list[dict[str, Any]]:
    """Локальные свойства фрагментов абзаца: для проверки, что оформление не потеряно."""
    result = []
    for run in paragraph.runs:
        rpr = run._r.find(f"{{{NS_A}}}rPr")
        result.append(
            {
                "text": run.text,
                "attrs": dict(rpr.attrib) if rpr is not None else {},
                "children": [etree.QName(c).localname for c in rpr] if rpr is not None else [],
            }
        )
    return result


# --- удаление -----------------------------------------------------------------------


def delete_slide(prs: Any, slide: Any) -> None:
    """Убирает слайд из колоды; его части и исключительные ресурсы не сохраняются в файл."""
    sld_id_lst = prs.slides._sldIdLst
    for sld_id in list(sld_id_lst):
        if prs.part.related_part(sld_id.rId) is slide.part:
            sld_id_lst.remove(sld_id)
            prs.part.drop_rel(sld_id.rId)
            _drop_from_sections(prs.part._element, sld_id.id)
            return
    raise ValueError("слайд не принадлежит этой презентации")


def _drop_from_sections(presentation_element: Any, slide_id: str) -> None:
    for sld_id in presentation_element.iter(f"{{{NS_P14}}}sldId"):
        if sld_id.get("id") == str(slide_id):
            sld_id.getparent().remove(sld_id)


def keep_only_slides(prs: Any, keep: list[Any]) -> int:
    """Удаляет все слайды, кроме перечисленных; возвращает число удалённых."""
    keep_parts = {s.part for s in keep}
    removed = 0
    for slide in list(prs.slides):
        if slide.part not in keep_parts:
            delete_slide(prs, slide)
            removed += 1
    return removed


def part_is_xml(part: Any) -> bool:
    return isinstance(part, XmlPart)


# --- проверка пакета ----------------------------------------------------------------


@dataclass
class PackageReport:
    path: str
    parts: int = 0
    slides: int = 0
    unreachable_parts: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.errors

    def as_dict(self) -> dict[str, Any]:
        return {
            "path": self.path,
            "parts": self.parts,
            "slides": self.slides,
            "ok": self.ok,
            "unreachable_parts": self.unreachable_parts,
            "errors": self.errors,
            "warnings": self.warnings,
        }


def check_package(path: Any) -> PackageReport:
    """Проверяет ZIP-пакет PPTX без python-pptx: типы содержимого, цели связей,
    ссылки `r:*` в XML, достижимость частей, уникальность идентификаторов фигур на слайде.
    """
    report = PackageReport(path=str(path))
    with zipfile.ZipFile(path) as zf:
        names = set(zf.namelist())
        if zf.testzip() is not None:
            report.errors.append("повреждён ZIP")
            return report
        parts = {n for n in names if not n.endswith("/") and n != "[Content_Types].xml"}
        parts = {n for n in parts if not n.endswith(".rels")}
        report.parts = len(parts)
        types = _content_types(zf)
        for name in sorted(parts):
            if _content_type_for(types, name) is None:
                report.errors.append(f"нет типа содержимого для /{name}")
        rels_of: dict[str, dict[str, tuple[str, str, bool]]] = {}
        for name in sorted(n for n in names if n.endswith(".rels")):
            source = _rels_source(name)
            if source and source not in parts:
                report.errors.append(f"{name} принадлежит несуществующей части")
                continue
            rels: dict[str, tuple[str, str, bool]] = {}
            for rel in etree.fromstring(zf.read(name), _XML_PARSER):
                rid, reltype = rel.get("Id"), rel.get("Type")
                target, external = rel.get("Target"), rel.get("TargetMode") == "External"
                if rid in rels:
                    report.errors.append(f"{name}: повтор идентификатора {rid}")
                if not external:
                    resolved = _resolve(source, target)
                    if resolved not in parts:
                        report.errors.append(f"{name}: {rid} указывает на отсутствующую {target}")
                    target = resolved
                rels[rid] = (reltype, target, external)
            rels_of[source] = rels
        reachable = _reachable(rels_of)
        report.unreachable_parts = sorted(parts - reachable)
        slides = sorted(
            (p for p in parts if re.fullmatch(r"ppt/slides/slide\d+\.xml", p)),
            key=lambda p: int(re.findall(r"\d+", p)[0]),
        )
        report.slides = len([s for s in slides if s in reachable])
        for name in sorted(parts):
            if not name.endswith(".xml"):
                continue
            root = etree.fromstring(zf.read(name), _XML_PARSER)
            available = rels_of.get(name, {})
            for node in root.iter():
                for attr in _REL_ATTRS:
                    value = node.get(f"{{{NS_R}}}{attr}")
                    if value is not None and value not in available:
                        report.errors.append(f"/{name}: r:{attr}={value} без связи")
            if name in slides:
                ids = [n.get("id") for n in root.iter(f"{{{NS_P}}}cNvPr")]
                dupes = {i for i in ids if ids.count(i) > 1}
                if dupes:
                    report.warnings.append(f"/{name}: повторяются id фигур {sorted(dupes)}")
        pres = etree.fromstring(zf.read("ppt/presentation.xml"), _XML_PARSER)
        pres_rels = rels_of.get("ppt/presentation.xml", {})
        for sld_id in pres.iter(f"{{{NS_P}}}sldId"):
            rid = sld_id.get(f"{{{NS_R}}}id")
            if rid not in pres_rels or pres_rels[rid][0] != RT.SLIDE:
                report.errors.append(f"presentation.xml: sldId {sld_id.get('id')} без слайда")
        for sld_id in pres.iter(f"{{{NS_P14}}}sldId"):
            if not any(s.get("id") == sld_id.get("id") for s in pres.iter(f"{{{NS_P}}}sldId")):
                report.warnings.append(f"раздел ссылается на удалённый слайд {sld_id.get('id')}")
    return report


def _content_types(zf: zipfile.ZipFile) -> tuple[dict[str, str], dict[str, str]]:
    root = etree.fromstring(zf.read("[Content_Types].xml"), _XML_PARSER)
    defaults = {
        n.get("Extension").lower(): n.get("ContentType") for n in root.iter(f"{{{NS_CT}}}Default")
    }
    overrides = {n.get("PartName"): n.get("ContentType") for n in root.iter(f"{{{NS_CT}}}Override")}
    return defaults, overrides


def _content_type_for(types: tuple[dict[str, str], dict[str, str]], name: str) -> str | None:
    defaults, overrides = types
    if f"/{name}" in overrides:
        return overrides[f"/{name}"]
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    return defaults.get(ext)


def _rels_source(rels_name: str) -> str:
    """`ppt/slides/_rels/slide1.xml.rels` → `ppt/slides/slide1.xml`; корень → ''."""
    if rels_name == "_rels/.rels":
        return ""
    head, tail = rels_name.rsplit("_rels/", 1)
    return f"{head}{tail[: -len('.rels')]}"


def _resolve(source: str, target: str) -> str:
    if target.startswith("/"):
        return target[1:]
    base = posixpath.dirname(source) if source else ""
    return posixpath.normpath(posixpath.join(base, target)) if base else posixpath.normpath(target)


def _reachable(rels_of: dict[str, dict[str, tuple[str, str, bool]]]) -> set[str]:
    seen: set[str] = set()
    stack = [""]
    while stack:
        current = stack.pop()
        for _reltype, target, external in rels_of.get(current, {}).values():
            if external or target in seen:
                continue
            seen.add(target)
            stack.append(target)
    return seen


# --- сравнение пакетов --------------------------------------------------------------


def compare_packages(a: Any, b: Any) -> dict[str, Any]:
    """Сравнивает два PPTX по частям: списки, XML по канонической форме, бинарные по байтам."""
    with zipfile.ZipFile(a) as za, zipfile.ZipFile(b) as zb:
        names_a, names_b = set(za.namelist()), set(zb.namelist())
        only_a, only_b = sorted(names_a - names_b), sorted(names_b - names_a)
        changed_xml: list[str] = []
        changed_binary: list[str] = []
        for name in sorted(names_a & names_b):
            blob_a, blob_b = za.read(name), zb.read(name)
            if blob_a == blob_b:
                continue
            if name.endswith((".xml", ".rels")):
                if _canonical(blob_a) != _canonical(blob_b):
                    changed_xml.append(name)
            else:
                changed_binary.append(name)
    return {
        "only_in_a": only_a,
        "only_in_b": only_b,
        "changed_xml": changed_xml,
        "changed_binary": changed_binary,
        "identical": not (only_a or only_b or changed_xml or changed_binary),
    }


def _canonical(blob: bytes) -> bytes:
    root = etree.fromstring(blob, _XML_PARSER)
    for node in root.iter():
        if len(node) and node.text is not None and not node.text.strip():
            node.text = None
        if node.tail is not None and not node.tail.strip():
            node.tail = None
    # Списки связей и типов содержимого — множества: порядок записей не имеет значения.
    if etree.QName(root).localname in ("Relationships", "Types"):
        children = sorted(root, key=lambda n: sorted(n.attrib.items()))
        for child in children:
            root.remove(child)
        root.extend(children)
    # Порядок атрибутов и объявления префиксов не влияют на смысл.
    return bytes(etree.tostring(root, method="c14n"))


def shape_tree_signature(slide: Any) -> bytes:
    """Каноническое дерево фигур слайда, где ссылки на связи заменены их целями.

    Две копии слайда с одинаковой подписью содержат те же фигуры и ссылаются на те же
    ресурсы, даже если идентификаторы связей у них разные. Диаграммы у копии свои,
    поэтому для них подпись хранит только тип связи.
    """
    tree = copy.deepcopy(slide.shapes._spTree)
    rels = slide.part.rels
    for node in tree.iter():
        for name in _REL_ATTRS:
            key = f"{{{NS_R}}}{name}"
            value = node.get(key)
            if value is None:
                continue
            rel = rels.get(value)
            if rel is None:
                node.set(key, f"missing:{value}")
            elif rel.is_external:
                node.set(key, f"external:{rel.target_ref}")
            elif rel.reltype == RT.CHART:
                node.set(key, "chart")
            else:
                node.set(key, f"part:{rel.target_part.partname}")
    return _canonical(bytes(etree.tostring(tree)))
