"""Слияние пакетов PPTX: дизайн-система из нескольких файлов в одном сводном пакете.

Зачем. Шаблон приходит не один раз: у человека есть прошлая презентация, потом вторая,
потом обновлённая версия. Держать их отдельными шаблонами нельзя — вёрстка клонирует
образец **внутри** пакета и сверяет `template_hash`, поэтому композиции из разных файлов в
одну колоду не соберутся.

Решение. Дизайн-система хранит один растущий PPTX. Каждый источник вливается в него вместе
со своим мастером и темой, поэтому его слайды сохраняют вид, а не перекрашиваются в палитру
первого файла: несколько мастеров в одном файле — штатный OOXML (у шаблона VK Tech их уже
два). Профиль строится по сводному пакету, и вёрстка, ёмкость, аудит и экспорт работают без
единой правки.

Как. Части переносятся рекурсивно по графу связей: мастер → его макеты и тема → картинки.
Каждая часть создаётся через `PartFactory`, чтобы python-pptx получил типизированный объект
(`SlideMasterPart`, `SlideLayoutPart`, `ImagePart`), а не безликий blob. Картинки
дедуплицируются по sha256: три файла одного бренда обычно несут один и тот же логотип, и без
этого сводный пакет рос бы втрое.
"""

from __future__ import annotations

import hashlib
import logging
from dataclasses import dataclass, field
from typing import Any

from pptx.opc.constants import CONTENT_TYPE as CT
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import PartFactory

from presentation_designer.layout.ooxml import _remap_rel_ids, part_is_xml

log = logging.getLogger(__name__)

# Части, которые переносятся под своим именем с нумерацией: шаблон имени для нового пакета.
_PARTNAME_TEMPLATES = {
    CT.PML_SLIDE_MASTER: "/ppt/slideMasters/slideMaster%d.xml",
    CT.PML_SLIDE_LAYOUT: "/ppt/slideLayouts/slideLayout%d.xml",
    CT.PML_SLIDE: "/ppt/slides/slide%d.xml",
    CT.OFC_THEME: "/ppt/theme/theme%d.xml",
}
_MEDIA_PREFIX = "/ppt/media/"


@dataclass
class MergeResult:
    """Что принесло слияние: по этому строится запись версии дизайн-системы."""

    masters: int = 0
    layouts: int = 0
    slides: int = 0
    media_added: int = 0
    media_reused: int = 0
    warnings: list[str] = field(default_factory=list)

    def as_dict(self) -> dict[str, Any]:
        return {
            "masters": self.masters,
            "layouts": self.layouts,
            "slides": self.slides,
            "media_added": self.media_added,
            "media_reused": self.media_reused,
            "warnings": list(self.warnings),
        }


class MergeError(RuntimeError):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


@dataclass
class _Ctx:
    """Состояние одного слияния: что уже перенесено и какие картинки можно переиспользовать."""

    package: Any
    imported: dict[str, Any] = field(default_factory=dict)
    media_by_sha: dict[str, Any] = field(default_factory=dict)
    # Имена, уже выданные в этом слиянии. Пакет python-pptx обходит части по графу связей от
    # корня, поэтому только что созданная часть, которую родитель ещё не привязал, для
    # `next_partname` невидима — и второй картинке достаётся то же имя.
    taken: set[str] = field(default_factory=set)
    result: MergeResult = field(default_factory=MergeResult)


def merge_presentation(base: Any, source: Any, *, slides: bool = True) -> MergeResult:
    """Вливает `source` в презентацию `base`: мастера с макетами и темами, затем слайды.

    `slides=False` переносит только оформление — это случай, когда из файла нужны макеты,
    а его собственные слайды образцами не считаются.
    """
    ctx = _Ctx(package=base.part.package)
    _index_media(ctx, base)

    for master in source.slide_masters:
        new_master = _import_part(ctx, master.part)
        _attach_master(base, new_master)
        ctx.result.masters += 1
        ctx.result.layouts += len(master.slide_layouts)

    if slides:
        for slide in source.slides:
            new_slide = _import_part(ctx, slide.part)
            _attach_slide(base, new_slide)
            ctx.result.slides += 1
    return ctx.result


def _index_media(ctx: _Ctx, base: Any) -> None:
    """Картинки, которые в сводном пакете уже есть: по ним идёт дедупликация."""
    for part in base.part.package.iter_parts():
        if str(part.partname).startswith(_MEDIA_PREFIX):
            ctx.media_by_sha.setdefault(hashlib.sha256(part.blob).hexdigest(), part)


def _import_part(ctx: _Ctx, src_part: Any) -> Any:
    """Переносит часть и всё, на что она ссылается; повторный вызов отдаёт уже перенесённое."""
    key = str(src_part.partname)
    done = ctx.imported.get(key)
    if done is not None:
        return done

    blob = src_part.blob
    if key.startswith(_MEDIA_PREFIX):
        sha = hashlib.sha256(blob).hexdigest()
        existing = ctx.media_by_sha.get(sha)
        if existing is not None:
            ctx.imported[key] = existing
            ctx.result.media_reused += 1
            return existing

    new_part = PartFactory(
        _reserve_partname(ctx, _partname_template(src_part)),
        src_part.content_type,
        ctx.package,
        blob,
    )
    # Запись до обхода связей: у мастера и макета ссылки взаимные, иначе будет бесконечность.
    ctx.imported[key] = new_part
    if key.startswith(_MEDIA_PREFIX):
        ctx.media_by_sha[hashlib.sha256(blob).hexdigest()] = new_part
        ctx.result.media_added += 1

    mapping: dict[str, str] = {}
    for rid, rel in src_part.rels.items():
        if rel.is_external:
            mapping[rid] = new_part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
            continue
        try:
            target = _import_part(ctx, rel.target_part)
        except Exception as e:  # одна битая связь не должна ронять слияние целиком
            log.warning("связь %s части %s не перенесена: %s", rel.reltype, key, e)
            ctx.result.warnings.append(f"{key}: связь {rel.reltype} не перенесена")
            continue
        mapping[rid] = new_part.relate_to(target, rel.reltype)
    element = getattr(new_part, "_element", None)
    if mapping and element is not None and part_is_xml(new_part):
        _remap_rel_ids(element, mapping)
    return new_part


def _reserve_partname(ctx: _Ctx, template: str) -> Any:
    """Свободное имя части с учётом уже выданных в этом слиянии."""
    name = ctx.package.next_partname(template)
    while str(name) in ctx.taken:
        # `next_partname` вернул занятое: номер увеличивается, пока имя не станет свободным.
        head, _, tail = str(name).rpartition("/")
        stem, dot, ext = tail.rpartition(".")
        digits = len(stem) - len(stem.rstrip("0123456789"))
        base_stem = stem[: len(stem) - digits] if digits else stem
        number = int(stem[len(base_stem) :] or 0) + 1
        name = type(name)(f"{head}/{base_stem}{number}{dot}{ext}")
    ctx.taken.add(str(name))
    return name


def _partname_template(src_part: Any) -> str:
    """Имя для части в новом пакете: с номером, чтобы не столкнуться с существующими.

    Для картинок номер ставится вместо номера исходного файла, а не приписывается к нему:
    из `image12.png` шаблон `image12%d.png` даёт `image121.png`, которое может совпасть с
    уже существующим `image121.png` — и в ZIP появляются две записи с одним именем.
    """
    template = _PARTNAME_TEMPLATES.get(src_part.content_type)
    if template is not None:
        return template
    name = str(src_part.partname)
    head, _, ext = name.rpartition(".")
    if not _:
        return f"{name}%d"
    stem = head.rstrip("0123456789") or head
    return f"{stem}%d.{ext}"


_SLD_LAYOUT_ID = "{http://schemas.openxmlformats.org/presentationml/2006/main}sldLayoutId"
_SLD_LAYOUT_ID_LST = "{http://schemas.openxmlformats.org/presentationml/2006/main}sldLayoutIdLst"
# Нижняя граница id мастеров и макетов по ECMA-376: PowerPoint выдаёт их от 2147483648.
_MIN_MASTER_ID = 2147483648


def _attach_master(base: Any, master_part: Any) -> None:
    """Регистрирует мастер в презентации: связь и запись в `sldMasterIdLst`.

    Id мастера и его макетов — один общий счётчик на всю презентацию. Импортированный мастер
    приносит номера из своего файла, и у двух файлов PowerPoint они почти наверняка
    совпадают (оба начинают с 2147483648): PowerPoint тогда предлагает «восстановить» файл.
    Поэтому макетам перенесённого мастера выдаются номера следом за занятыми.
    """
    rid = base.part.relate_to(master_part, RT.SLIDE_MASTER)
    lst = base.part._element.get_or_add_sldMasterIdLst()
    # У python-pptx нет публичного добавления мастера (в отличие от слайдов): элемент
    # создаётся напрямую, id берётся на единицу больше максимального — так делает PowerPoint.
    entry = lst._add_sldMasterId()
    entry.rId = rid
    taken = _master_and_layout_ids(base, exclude=master_part)
    next_id = max(taken) + 1 if taken else _MIN_MASTER_ID
    entry.set("id", str(next_id))
    next_id += 1
    for layout_id in master_part._element.iter(_SLD_LAYOUT_ID):
        layout_id.set("id", str(next_id))
        next_id += 1


def _master_and_layout_ids(base: Any, *, exclude: Any) -> list[int]:
    """Id мастеров и макетов, уже занятые в презентации (кроме только что добавленного)."""
    lst = base.part._element.get_or_add_sldMasterIdLst()
    ids = [int(e.get("id")) for e in lst.sldMasterId_lst if e.get("id")]
    for master in base.slide_masters:
        if master.part is exclude:
            continue
        ids += [int(e.get("id")) for e in master.part._element.iter(_SLD_LAYOUT_ID) if e.get("id")]
    return ids


def _attach_slide(base: Any, slide_part: Any) -> None:
    """Добавляет перенесённый слайд в конец колоды сводного пакета."""
    rid = base.part.relate_to(slide_part, RT.SLIDE)
    base.slides._sldIdLst.add_sldId(rid)
