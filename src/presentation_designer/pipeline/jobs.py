"""Задания и очередь: постановка, зависимости, задачи воркеров, отмена, сверка.

Граф генерации: анализ шаблона идёт параллельно импорту содержания; смысловой план ждёт
импорт; варианты ждут план и анализ; финализатор ждёт все варианты и не занимает воркер
в ожидании — зависимости планирует очередь. Отказ варианта не отменяет остальные
(`Dependency(allow_failure=True)`), состояние всегда читается из SQLite, а не из RQ.

Исполнитель RQ работает поверх Valkey; встроенный исполнитель выполняет задачи сразу
и служит тестам и CLI. Функции задач принимают только идентификаторы и достают
оркестратор процесса через `get_orchestrator()`.
"""

from __future__ import annotations

import hashlib
import json
import logging
import pathlib
import shutil
import socket
import threading
import time
from collections.abc import Callable
from datetime import UTC, datetime, timedelta
from typing import Any, Protocol

from presentation_designer import __version__
from presentation_designer.contracts import CONTRACTS_VERSION
from presentation_designer.generation.original import VARIANT_ID as ORIGINAL_VARIANT
from presentation_designer.generation.original import original_story
from presentation_designer.generation.variants import upgrade_plan_schema
from presentation_designer.llm import model_refs, version_refs
from presentation_designer.pipeline.artifacts import ArtifactStore, Staging
from presentation_designer.pipeline.files import FileStore, format_for
from presentation_designer.pipeline.run import (
    VARIANT_AXIS,
    AnalyzeInput,
    BriefInput,
    EditContext,
    ExportInput,
    ImportInput,
    Layers,
    RepairInput,
    SourceFile,
    StageError,
    StoryInput,
    VariantContext,
    audit_summary,
    run_edit,
    run_repair,
    run_variant,
)
from presentation_designer.pipeline.state import (
    TERMINAL,
    JsonDict,
    NotFound,
    State,
    new_id,
    now_iso,
)
from presentation_designer.shared.settings import Settings, get_models_config, get_settings
from presentation_designer.shared.text import plural

log = logging.getLogger(__name__)

STAGE_ORDER = [
    "queued",
    "analyze",
    "import",
    "story",
    "plan",
    "compose",
    "export",
    "audit",
    "repair",
    "finalize",
    "done",
]


class ConflictError(ValueError):
    def __init__(self, code: str, message: str, details: JsonDict | None = None) -> None:
        super().__init__(message)
        self.code = code
        self.details = details or {}


# ---------- исполнители ----------


class Executor(Protocol):
    name: str

    def enqueue(
        self,
        queue: str,
        func: Callable[..., Any],
        args: tuple[Any, ...],
        *,
        depends_on: list[str],
        timeout: int,
        description: str,
    ) -> str: ...

    def status(self, rq_id: str) -> str | None: ...

    def workers(self) -> dict[str, int]: ...

    def renderer_ok(self) -> bool | None: ...

    def ping(self) -> bool: ...

    def try_lock(self, key: str, ttl_s: int) -> bool: ...


class InlineExecutor:
    """Выполняет задачу сразу при постановке. Порядок постановки повторяет порядок зависимостей."""

    name = "inline"

    def __init__(self) -> None:
        self.done: dict[str, str] = {}
        self._locks: dict[str, float] = {}

    def enqueue(
        self,
        queue: str,
        func: Callable[..., Any],
        args: tuple[Any, ...],
        *,
        depends_on: list[str],
        timeout: int,
        description: str,
    ) -> str:
        rq_id = new_id("inl")
        try:
            func(*args)
            self.done[rq_id] = "finished"
        except Exception:
            log.exception("задача %s завершилась ошибкой", description)
            self.done[rq_id] = "failed"
        return rq_id

    def status(self, rq_id: str) -> str | None:
        return self.done.get(rq_id)

    def workers(self) -> dict[str, int]:
        return {"analysis": 1, "generation": 1}

    def renderer_ok(self) -> bool | None:
        return None

    def ping(self) -> bool:
        return True

    def try_lock(self, key: str, ttl_s: int) -> bool:
        import time

        now = time.monotonic()
        if self._locks.get(key, 0) > now:
            return False
        self._locks[key] = now + ttl_s
        return True


class RQExecutor:
    name = "rq"

    def __init__(self, url: str, settings: Settings) -> None:
        from rq import Queue

        from presentation_designer.shared.valkey import connect

        self.settings = settings
        self.redis = connect(url)
        self._queues: dict[str, Queue] = {}
        self._queue_cls = Queue

    def queue(self, name: str) -> Any:
        if name not in self._queues:
            self._queues[name] = self._queue_cls(name, connection=self.redis)
        return self._queues[name]

    def enqueue(
        self,
        queue: str,
        func: Callable[..., Any],
        args: tuple[Any, ...],
        *,
        depends_on: list[str],
        timeout: int,
        description: str,
    ) -> str:
        from rq.job import Dependency

        q = self.settings.queue
        kwargs: dict[str, Any] = {
            "args": args,
            "job_timeout": timeout,
            "result_ttl": q.result_ttl_s,
            "failure_ttl": q.failure_ttl_s,
            "description": description,
        }
        if depends_on:
            kwargs["depends_on"] = Dependency(jobs=list(depends_on), allow_failure=True)
        job = self.queue(queue).enqueue(func, **kwargs)
        return str(job.id)

    def status(self, rq_id: str) -> str | None:
        from rq.exceptions import NoSuchJobError
        from rq.job import Job

        try:
            job = Job.fetch(rq_id, connection=self.redis)
        except NoSuchJobError:
            return None
        status = job.get_status(refresh=False)
        return str(getattr(status, "value", status))

    def workers(self) -> dict[str, int]:
        from rq import Worker

        counts = {"analysis": 0, "generation": 0}
        try:
            for worker in Worker.all(connection=self.redis):
                names = set(worker.queue_names())
                if self.settings.queue.generation_queue in names:
                    counts["generation"] += 1
                if self.settings.queue.analysis_queue in names:
                    counts["analysis"] += 1
        except Exception:
            log.exception("не удалось получить список воркеров")
        return counts

    def renderer_ok(self) -> bool | None:
        """Все живые воркеры генерации прошли проверку рендерера при старте."""
        from rq import Worker

        try:
            statuses: list[bytes | None] = []
            for worker in Worker.all(connection=self.redis):
                if self.settings.queue.generation_queue in set(worker.queue_names()):
                    statuses.append(self.redis.get(f"pd:renderer:{worker.name}"))
        except Exception:
            return None
        if not statuses:
            return None
        return all(s == b"ok" for s in statuses)

    def ping(self) -> bool:
        try:
            return bool(self.redis.ping())
        except Exception:
            return False

    def try_lock(self, key: str, ttl_s: int) -> bool:
        """Ключ с TTL: True у того, кто поставил его первым (SET NX)."""
        try:
            return bool(self.redis.set(key, "1", nx=True, ex=max(1, ttl_s)))
        except Exception:
            log.exception("не удалось взять блокировку %s", key)
            return False


# ---------- оркестратор ----------


class Orchestrator:
    def __init__(
        self,
        settings: Settings,
        state: State,
        files: FileStore,
        artifacts: ArtifactStore,
        layers: Layers,
        executor: Executor,
    ) -> None:
        self.settings = settings
        self.state = state
        self.files = files
        self.artifacts = artifacts
        self.layers = layers
        self.executor = executor
        set_current(self)

    # ----- шаблоны -----

    def submit_template(self, *, sha256: str, name: str, size_bytes: int) -> tuple[JsonDict, bool]:
        job_id = new_id("job")
        template, cached = self.state.get_or_create_template(
            sha256=sha256, name=name, size_bytes=size_bytes, job_id=job_id
        )
        key = self.template_profile_key(sha256)
        if cached:
            degraded = {w.get("code") for w in (template.get("profile") or {}).get("warnings", [])}
            stale = template["status"] == "succeeded" and (
                template.get("profile_key") != key
                # Прошлый анализ прошёл без рендерера: профиль без превью не считается готовым.
                or "previews_unavailable" in degraded
            )
            if template["status"] == "failed" or stale:
                # Ошибка прошлого анализа или устаревший ключ (сменились анализатор, контракты,
                # модель, промпты или шрифты): тот же template_id анализируется заново.
                self.state.update_template(
                    template["template_id"], status="queued", error=None, job_id=job_id
                )
                template["job_id"] = job_id
                cached = False
            else:
                return template, True
        self.state.update_template(template["template_id"], profile_key=key)
        job = self.state.create_job(
            kind="template_analysis",
            job_id=job_id,
            stage="analyze",
            result={
                "template_id": template["template_id"],
                "template_profile_url": f"/api/templates/{template['template_id']}",
            },
            progress={"percent": 0, "message": "Анализ образцов шаблона"},
        )
        rq_id = self.executor.enqueue(
            self.settings.queue.analysis_queue,
            task_analyze,
            (template["template_id"],),
            depends_on=[],
            timeout=self.settings.timeouts.stage_analyze_s,
            description=f"analyze {template['template_id']}",
        )
        self.state.update_job(job["job_id"], rq_ids=[rq_id])
        return template, cached

    def reanalyze_templates(self, *, force: bool = False) -> list[JsonDict]:
        """Ставит в очередь повторный анализ шаблонов библиотеки, чей ключ профиля устарел
        (сменились анализатор, контракты, скилл, модель или шрифты) или чей анализ упал; с
        `force` — всех. Тот же путь, что при повторной загрузке файла (`submit_template`),
        только без загрузки: файл уже в хранилище по sha256."""
        out: list[JsonDict] = []
        for t in self.state.list_templates():
            entry: JsonDict = {
                "template_id": t["template_id"],
                "name": t["name"],
                "status": t["status"],
            }
            if t["status"] in ("queued", "running"):
                entry["action"] = "in_progress"
                out.append(entry)
                continue
            key = self.template_profile_key(t["sha256"])
            stale = t.get("profile_key") != key or t["status"] == "failed"
            if not stale and not force:
                entry["action"] = "up_to_date"
                out.append(entry)
                continue
            if not self.files.exists(t["sha256"]):
                entry["action"] = "file_missing"
                out.append(entry)
                continue
            if not stale:
                # Ключ совпадает, но анализ запрошен явно: ключ сбрасывается, чтобы
                # submit_template счёл профиль устаревшим.
                self.state.update_template(t["template_id"], profile_key=None)
            template, cached = self.submit_template(
                sha256=t["sha256"], name=t["name"], size_bytes=t["size_bytes"]
            )
            entry["action"] = "kept" if cached else "queued"
            entry["job_id"] = template.get("job_id")
            out.append(entry)
        return out

    def delete_template(self, template_id: str) -> JsonDict:
        """Убирает шаблон из библиотеки: строку, миниатюры и ссылки проектов на него.
        Незавершённый анализ отменяется; его задание и байты файла подберёт сборка мусора."""
        template = self.state.get_template(template_id)
        if template["status"] in {"queued", "running"}:
            try:
                self.cancel_job(template["job_id"])
            except NotFound:
                pass
        self.state.delete_template(template_id)
        if self.artifacts.template_dir(template_id).exists():
            self.artifacts.delete_dir("templates", template_id)
        return template

    def extract_template_media(self, template_id: str, asset: JsonDict) -> pathlib.Path:
        """Файл ресурса профиля (`assets[].media_path`) в каталоге шаблона: извлекается из
        загруженного PPTX при первом запросе (`media/<asset_id>.<ext>`), дальше отдаётся с
        диска. Только части `ppt/media/`; отсутствие — FileNotFoundError."""
        import zipfile

        from presentation_designer.layout.media import media_file_name

        template = self.state.get_template(template_id)
        media_path = str(asset.get("media_path") or "").lstrip("/")
        asset_id = str(asset.get("asset_id") or "")
        if not media_path.startswith("ppt/media/") or not asset_id:
            raise FileNotFoundError(media_path)
        target = (
            self.artifacts.template_dir(template_id)
            / "media"
            / media_file_name(asset_id, media_path)
        )
        if target.is_file():
            return target
        source = self.files.path_for(str(template["sha256"]))
        if not source.is_file():
            raise FileNotFoundError(str(source))
        try:
            with zipfile.ZipFile(source) as zf:
                if media_path not in zf.namelist():
                    raise FileNotFoundError(media_path)
                data = zf.read(media_path)
        except zipfile.BadZipFile as e:
            raise FileNotFoundError(media_path) from e
        target.parent.mkdir(parents=True, exist_ok=True)
        tmp = target.with_name(f".{target.name}.tmp")
        tmp.write_bytes(data)
        tmp.replace(target)
        return target

    # ----- содержание -----

    def submit_package(
        self, *, sources: list[JsonDict], brief: JsonDict | None
    ) -> tuple[JsonDict, bool]:
        """sources: записи файлов проекта (file_id, name, sha256, size_bytes, check)."""
        file_ids = [s["file_id"] for s in sources]
        mode = "mixed" if sources and brief else "brief" if brief else "package"
        # Ключ пакета: содержимое файлов в их порядке, бриф и версия импортёра; сам разбор
        # файлов кэшируется слоем отдельно, поэтому смена брифа их не перечитывает.
        key_src = json.dumps(
            {
                "files": [s["sha256"] for s in sources],
                "brief": brief or {},
                "importer": self.layers.import_version(),
            },
            sort_keys=True,
            ensure_ascii=False,
        )
        idem_key = hashlib.sha256(key_src.encode("utf-8")).hexdigest()
        job_id = new_id("job")
        package, cached = self.state.get_or_create_package(
            idem_key=idem_key, mode=mode, file_ids=file_ids, brief=brief, job_id=job_id
        )
        if cached:
            return package, True
        self.state.create_job(
            kind="content_import",
            job_id=job_id,
            stage="import",
            result={
                "package_id": package["package_id"],
                "content_package_url": f"/api/content/{package['package_id']}",
            },
            progress={"percent": 0, "message": "Извлечение блоков и фактов"},
        )
        rq_id = self.executor.enqueue(
            self.settings.queue.generation_queue,
            task_import,
            (package["package_id"],),
            depends_on=[],
            timeout=self.settings.timeouts.stage_story_s,
            description=f"import {package['package_id']}",
        )
        self.state.update_job(job_id, rq_ids=[rq_id])
        return package, False

    # ----- генерация -----

    def submit_generation(self, request: JsonDict) -> JsonDict:
        key = request.get("idempotency_key")
        if key:
            same = self.state.find_generation_by_key(key)
            if same:
                return self.state.get_job(same["job_id"])
        template = self.state.get_template(request["template_id"])
        package = self.state.get_package(request["package_id"])
        settings = request.get("settings") or {}
        variant_ids = list(settings.get("variants") or ["compact", "balanced", "detailed"])
        job_id = new_id("job")
        deadline = (
            datetime.now(UTC) + timedelta(seconds=self.settings.timeouts.job_total_s)
        ).strftime("%Y-%m-%dT%H:%M:%S.%f")[:-3] + "Z"
        depends_on = [template["job_id"], package["job_id"]]
        self.state.create_job(
            kind="generation",
            job_id=job_id,
            stage="queued",
            deadline_at=deadline,
            depends_on=depends_on,
            result={
                "template_id": template["template_id"],
                "package_id": package["package_id"],
                "generation_result_url": f"/api/generations/{job_id}",
            },
            progress={"percent": 0, "message": "Задание в очереди"},
        )
        self.state.create_generation(
            job_id=job_id,
            template_id=template["template_id"],
            package_id=package["package_id"],
            request=request,
            idempotency_key=key,
            execution_mode=self.layers.execution_mode(),
            versions=self.versions(),
            variants=[
                {
                    "variant_id": v,
                    "axis": VARIANT_AXIS.get(v, ("density", v))[0],
                    "value": VARIANT_AXIS.get(v, ("density", v))[1],
                }
                for v in variant_ids
            ],
        )
        if variant_ids == ["original"]:
            # Готовая презентация показывается сразу: редактору нужны только байты PPTX, поэтому
            # копия файла публикуется ревизией r1 здесь же, без очереди и без модели. Ошибка
            # копии запрос не роняет: копию сделает фоновая задача.
            try:
                _publish_original_copy(self, job_id, template)
            except Exception:
                log.exception("копия презентации %s не опубликована в запросе", job_id)
        self.enqueue_generation(job_id, template, package, variant_ids)
        return self.state.get_job(job_id)

    def enqueue_generation(
        self, job_id: str, template: JsonDict, package: JsonDict, variant_ids: list[str]
    ) -> None:
        q = self.settings.queue
        t = self.settings.timeouts
        rq_ids: list[str] = []
        original = variant_ids == ["original"]
        # Разбор готовой презентации без проверки моделью — смысловой план по профилю и сборка
        # «как есть», вместе около 8 с на 55 слайдов и без модели. Это фон уже открытой
        # презентации: в очереди generation он ждал бы чужих сборок до десяти минут, а правки
        # из чата — вместе с ним. С контекстным аудитом (минуты модели) — очередь generation.
        request_settings = self.state.get_generation(job_id)["request"].get("settings") or {}
        light = original and not bool(request_settings.get("run_contextual_audit", True))
        parse_queue = q.interactive_queue if light else q.generation_queue
        deps_story = self._pending_rq_ids(package["job_id"])
        if original:
            # Смысловой план исходной презентации строится по профилю: ждём и анализ.
            deps_story = [*deps_story, *self._pending_rq_ids(template["job_id"])]
        story_id = self.executor.enqueue(
            parse_queue,
            task_story,
            (job_id,),
            depends_on=deps_story,
            timeout=t.stage_story_s,
            description=f"story {job_id}",
        )
        rq_ids.append(story_id)
        deps_variant = [story_id, *self._pending_rq_ids(template["job_id"])]
        if original:
            # Копия уже показана (r1 опубликовал запрос). Подписи, диаграммы-картинки и рендер
            # открытой презентации идут в очереди interactive — воркеры генерации берут её
            # раньше чужих сборок; полная сборка ждёт эту задачу и рендер не повторяет.
            preview_id = self.executor.enqueue(
                q.interactive_queue,
                task_original_preview,
                (job_id,),
                depends_on=[],
                timeout=t.stage_export_s + int(self.settings.layout.chart_images.budget_s) + 60,
                description=f"preview {job_id}",
            )
            rq_ids.append(preview_id)
            deps_variant.append(preview_id)
        # Первый вариант собирается один: планы вариантов упираются в модель, и три плана
        # параллельно делят её между собой — первый готовый появлялся через минуты. Остальные
        # ставятся за ним (и при его ошибке тоже) и идут в фоне, пока человек смотрит первый.
        variant_rq: list[str] = []
        for variant_id in ordered_variants(variant_ids):
            vid = self.executor.enqueue(
                parse_queue,
                task_variant,
                (job_id, variant_id),
                depends_on=deps_variant if not variant_rq else [variant_rq[0]],
                timeout=t.stage_plan_s + t.stage_compose_s + t.stage_export_s + t.stage_audit_s,
                description=f"variant {job_id} {variant_id}",
            )
            variant_rq.append(vid)
        rq_ids.extend(variant_rq)
        # Итог готовой презентации — мгновенная запись статуса: в очереди generation он ждал бы
        # чужих сборок минутами, и открытая презентация числилась бы «в работе».
        fin = self.executor.enqueue(
            q.interactive_queue if original else q.generation_queue,
            task_finalize,
            (job_id,),
            depends_on=variant_rq,
            timeout=60,
            description=f"finalize {job_id}",
        )
        rq_ids.append(fin)
        self.state.update_job(job_id, rq_ids=rq_ids)

    def _pending_rq_ids(self, job_id: str) -> list[str]:
        """RQ-идентификаторы незавершённого задания: завершённые не нужны как зависимости."""
        try:
            job = self.state.get_job(job_id)
        except NotFound:
            return []
        if job["status"] in TERMINAL:
            return []
        return [rid for rid in job["rq_ids"] if self.executor.status(rid) is not None]

    def cancel_job(self, job_id: str) -> JsonDict:
        job = self.state.get_job(job_id)
        if job["status"] in TERMINAL:
            return job
        ts = now_iso()
        if job["kind"] == "generation":
            self.state.update_generation(job_id, canceled=True)
            for variant in self.state.get_variants(job_id):
                if variant["status"] in {"pending", "running"}:
                    self.state.update_variant(
                        job_id,
                        variant["variant_id"],
                        status="failed",
                        error={"code": "canceled", "message": "Задание отменено пользователем"},
                    )
        self.state.update_job(
            job_id,
            status="canceled",
            stage="done",
            finished_at=ts,
            error={"code": "canceled", "message": "Задание отменено пользователем"},
        )
        return self.state.get_job(job_id)

    def retry_job(self, job_id: str) -> JsonDict:
        job = self.state.get_job(job_id)
        if job["kind"] != "generation":
            raise ConflictError("retry_unsupported", "Повтор доступен только для заданий генерации")
        gen = self.state.get_generation(job_id)
        request = {k: v for k, v in gen["request"].items() if k != "idempotency_key"}
        return self.submit_generation(request)

    # ----- исправления -----

    def submit_repair(
        self, job_id: str, variant_id: str, base_revision: int, issue_ids: list[str]
    ) -> JsonDict:
        variant = self.state.get_variant(job_id, variant_id)
        if variant["revision"] != base_revision:
            raise ConflictError(
                "revision_stale",
                f"Ревизия {base_revision} устарела: текущая ревизия {variant['revision']}. Обновите отчёт и выберите находки заново.",  # noqa: E501
                {"current_revision": variant["revision"]},
            )
        if variant["status"] == "failed":
            raise ConflictError("variant_failed", "Вариант не собран, исправлять нечего")
        active = self.state.active_repair(job_id, variant_id)
        if active:
            raise ConflictError(
                "repair_in_progress",
                "Исправление этой ревизии уже выполняется",
                {"repair_job_id": active["repair_job_id"]},
            )
        repair_id = new_id("rep")
        self.state.create_job(
            kind="repair",
            job_id=repair_id,
            stage="repair",
            parent_job_id=job_id,
            result={"generation_result_url": f"/api/generations/{job_id}"},
            progress={"percent": 0, "message": "Исправление выбранных находок"},
        )
        self.state.create_repair(
            repair_job_id=repair_id,
            job_id=job_id,
            variant_id=variant_id,
            base_revision=base_revision,
            issue_ids=issue_ids,
        )
        rq_id = self.executor.enqueue(
            self.settings.queue.repair_queue,
            task_repair,
            (repair_id,),
            depends_on=[],
            timeout=self.settings.timeouts.stage_compose_s
            + self.settings.timeouts.stage_export_s
            + self.settings.timeouts.stage_audit_s,
            description=f"repair {repair_id}",
        )
        self.state.update_job(repair_id, rq_ids=[rq_id])
        return self.state.get_job(repair_id)

    # ----- отмена правки слайда -----

    def submit_revert(
        self, job_id: str, variant_id: str, base_revision: int, to_revision: int, slide_index: int
    ) -> JsonDict:
        """Отмена правки новой ревизией, без потери последующих изменений.

        После отмены более поздних правок номер ревизии уже другой. В этом случае
        должны совпасть все артефакты результата отменяемой правки, включая план:
        одного совпадения PPTX недостаточно для безопасного восстановления колоды.
        """
        variant = self.state.get_variant(job_id, variant_id)
        current_revision = variant["revision"]
        if current_revision != base_revision and not self._same_revision_files(
            job_id, variant_id, current_revision, base_revision
        ):
            raise ConflictError(
                "revision_stale",
                "После этой правки вариант уже меняли — отменить её отдельно нельзя.",
                {"current_revision": variant["revision"]},
            )
        if not 1 <= to_revision < base_revision:
            raise ConflictError("revert_invalid", "Вернуться можно только к прежней ревизии")
        try:
            self.state.get_revision(job_id, variant_id, to_revision)
        except NotFound as exc:
            raise ConflictError("revert_invalid", f"Ревизии {to_revision} нет") from exc
        active = self.state.active_repair(job_id, variant_id)
        if active:
            raise ConflictError(
                "repair_in_progress",
                "Правка этого варианта ещё выполняется",
                {"repair_job_id": active["repair_job_id"]},
            )
        undone = next(
            (
                r
                for r in self.state.list_repairs(job_id, kind=("edit", "patch", "repair"))
                if r["variant_id"] == variant_id and r["new_revision"] == base_revision
            ),
            None,
        )
        if (
            undone is None
            or undone["result"] != "applied"
            or undone["base_revision"] != to_revision
            or int(undone["slide_index"] or 0) != slide_index
        ):
            raise ConflictError("revert_invalid", "Ревизии и слайд не соответствуют правке")
        changed = list((undone or {}).get("changed_slide_ids") or [])
        edit_id = new_id("edit")
        new_rev = current_revision + 1
        self.state.create_job(
            kind="slide_edit",
            job_id=edit_id,
            stage="compose",
            parent_job_id=job_id,
            progress={"percent": 0, "message": "Отменяю правку"},
        )
        self.state.create_repair(
            repair_job_id=edit_id,
            job_id=job_id,
            variant_id=variant_id,
            base_revision=current_revision,
            issue_ids=[],
            kind="edit",
            slide_index=slide_index,
            instruction="Отмена правки",
        )
        source = self.artifacts.revision_dir(job_id, variant_id, to_revision)
        with self.artifacts.stage_revision(job_id, variant_id, new_rev) as staging:
            for path in sorted(source.rglob("*")):
                if (
                    path.is_file()
                    and path.name != "manifest.json"
                    and not path.name.startswith(".")
                ):
                    staging.write_bytes(path.relative_to(source).as_posix(), path.read_bytes())
        prefix = self.artifacts.prefix(variant_id, new_rev)
        manifest = self.artifacts.read_manifest(job_id, variant_id, new_rev)
        self.state.add_revision(
            job_id=job_id,
            variant_id=variant_id,
            revision=new_rev,
            artifacts_prefix=prefix,
            manifest=manifest,
            repair_job_id=edit_id,
            changed_slide_ids=changed,
            pptx_hash=_pptx_hash(manifest, prefix),
        )
        audit = dict(variant.get("audit") or {})
        if audit:
            audit["report_artifact"] = f"{prefix}audit.json"
            self.state.update_variant(job_id, variant_id, audit=audit)
        note = "Правка отменена: слайд как до неё"
        self.state.update_repair(
            edit_id,
            result="applied",
            new_revision=new_rev,
            slide_id=(undone or {}).get("slide_id"),
            change_note=note,
            changed_slide_ids=changed,
        )
        self.state.update_job(
            edit_id,
            status="succeeded",
            stage="done",
            finished_at=now_iso(),
            result={
                "revision": new_rev,
                "change_note": note,
                "generation_result_url": f"/api/generations/{job_id}",
            },
            progress={"percent": 100, "message": note},
        )
        return self.state.get_job(edit_id)

    def _same_revision_files(
        self, job_id: str, variant_id: str, current: int, expected: int
    ) -> bool:
        """Сравнение содержимого, а не номеров/путей ревизий; отсутствующее — не равно."""
        if expected >= current:
            return False
        snapshots: list[dict[str, str]] = []
        for revision in (current, expected):
            try:
                row = self.state.get_revision(job_id, variant_id, revision)
            except NotFound:
                return False
            prefix = self.artifacts.prefix(variant_id, revision)
            manifest = row["manifest"]
            if f"{prefix}deck.pptx" not in manifest:
                return False
            snapshot: dict[str, str] = {}
            for name in manifest:
                if not name.startswith(prefix):
                    return False
                path = self.artifacts.resolve(job_id, name)
                if not path.is_file():
                    return False
                snapshot[name.removeprefix(prefix)] = hashlib.sha256(path.read_bytes()).hexdigest()
            snapshots.append(snapshot)
        return snapshots[0] == snapshots[1]

    # ----- правка слайда по запросу -----

    def submit_edit(
        self,
        job_id: str,
        variant_id: str,
        base_revision: int,
        slide_index: int,
        instruction: str,
    ) -> JsonDict:
        """Правка одного слайда варианта по инструкции: то же задание ревизии, что и
        исправление, с видом `edit`; одна правка ревизии за раз.

        Готовая презентация открыта раньше, чем разобрана: пока у последней ревизии нет
        плана, правка принимается и ждёт разбора — задача ставится за заданиями генерации,
        а база перед запуском переносится на последнюю ревизию (`task_edit`)."""
        variant = self.state.get_variant(job_id, variant_id)
        waiting = self._edit_waits_for_plan(job_id, variant)
        if waiting is None and variant["revision"] != base_revision:
            raise ConflictError(
                "revision_stale",
                f"Ревизия {base_revision} устарела: текущая ревизия {variant['revision']}. Обновите результат и повторите просьбу.",  # noqa: E501
                {"current_revision": variant["revision"]},
            )
        if variant["status"] == "failed" and variant_id == "original":
            raise ConflictError("variant_failed", PARSE_FAILED)
        if waiting is None and variant["status"] in {"failed", "pending", "running"}:
            raise ConflictError("variant_failed", "Вариант ещё не собран, править нечего")
        active = self.state.active_repair(job_id, variant_id)
        if active:
            raise ConflictError(
                "repair_in_progress",
                "Предыдущая правка этой ревизии ещё применяется",
                {"repair_job_id": active["repair_job_id"]},
            )
        count = int(variant.get("slide_count") or 0)
        if count and not 0 <= slide_index < count:
            raise ConflictError(
                "slide_index_out_of_range",
                f"В варианте {count} слайдов, слайда с номером {slide_index + 1} нет",
            )
        edit_id = new_id("edit")
        message = f"Правка слайда {slide_index + 1}"
        if waiting is not None:
            from presentation_designer.pipeline.results import deferred_edit_message

            message = deferred_edit_message(self.state, job_id)
        self.state.create_job(
            kind="slide_edit",
            job_id=edit_id,
            stage="queued",
            parent_job_id=job_id,
            # Правка ждёт разбора: задание генерации в зависимостях — и знак для task_edit
            # перенести базу на последнюю ревизию.
            depends_on=[job_id] if waiting is not None else None,
            result={"generation_result_url": f"/api/generations/{job_id}"},
            progress={"percent": 0, "message": message},
        )
        self.state.create_repair(
            repair_job_id=edit_id,
            job_id=job_id,
            variant_id=variant_id,
            base_revision=variant["revision"] if waiting is not None else base_revision,
            issue_ids=[],
            kind="edit",
            slide_index=slide_index,
            instruction=instruction,
        )
        rq_id = self.executor.enqueue(
            self.settings.queue.repair_queue,
            task_edit,
            (edit_id,),
            depends_on=waiting or [],
            timeout=self.settings.timeouts.stage_plan_s
            + self.settings.timeouts.stage_compose_s
            + self.settings.timeouts.stage_export_s
            + self.settings.timeouts.stage_audit_s,
            description=f"edit {edit_id}",
        )
        self.state.update_job(edit_id, rq_ids=[rq_id])
        return self.state.get_job(edit_id)

    def _edit_waits_for_plan(self, job_id: str, variant: JsonDict) -> list[str] | None:
        """Задачи генерации, за которыми встаёт правка, если ревизия уже показана, а плана у
        неё ещё нет (разбор готовой презентации идёт в фоне); None — ждать нечего, правка
        проверяется как обычно."""
        if variant["status"] not in {"pending", "running"}:
            return None
        try:
            job = self.state.get_job(job_id)
            self.state.get_revision(job_id, variant["variant_id"], int(variant["revision"]))
        except NotFound:
            return None
        if job["status"] in TERMINAL:
            return None
        plan = self.artifacts.revision_dir(job_id, variant["variant_id"], int(variant["revision"]))
        if (plan / "plan.json").is_file():
            return None
        # Итог задания (последняя задача, см. enqueue_generation) правке не нужен: план готов,
        # как только собран вариант.
        final = job["rq_ids"][-1:] if job["rq_ids"] else []
        return [rid for rid in self._pending_rq_ids(job_id) if rid not in final]

    def submit_patch(
        self,
        job_id: str,
        variant_id: str,
        base_revision: int,
        slides: list[JsonDict],
        order: list[str] | None = None,
        template_logo: str | None = None,
    ) -> JsonDict:
        """Ручные правки из визуального редактора (документ slide_patch): то же задание
        ревизии, что и правка из чата, с видом `patch`, без модели. Правки проверяются
        против плана и ComposedDeck базовой ревизии до постановки в очередь; файлы проекта
        (источник `file`) должны принадлежать проекту задания и быть картинками."""
        from presentation_designer.generation.patch import (
            PatchError,
            normalize_patch,
            patch_is_noop,
            validate_patch,
        )

        variant = self.state.get_variant(job_id, variant_id)
        if variant["revision"] != base_revision:
            raise ConflictError(
                "revision_stale",
                f"Ревизия {base_revision} устарела: текущая ревизия {variant['revision']}. Обновите результат и повторите правки.",  # noqa: E501
                {"current_revision": variant["revision"]},
            )
        if variant["status"] in {"failed", "pending", "running"}:
            raise ConflictError("variant_failed", "Вариант ещё не собран, править нечего")
        active = self.state.active_repair(job_id, variant_id)
        if active:
            raise ConflictError(
                "repair_in_progress",
                "Предыдущая правка этой ревизии ещё применяется",
                {"repair_job_id": active["repair_job_id"]},
            )
        base_dir = self.artifacts.revision_dir(job_id, variant_id, base_revision)
        plan_path = base_dir / "plan.json"
        deck_path = base_dir / "composed.json"
        if not plan_path.is_file():
            raise ConflictError("plan_missing", "План базовой ревизии не найден")
        plan = upgrade_plan_schema(json.loads(plan_path.read_text(encoding="utf-8")))
        deck = json.loads(deck_path.read_text(encoding="utf-8")) if deck_path.is_file() else None
        patch = normalize_patch(
            {
                "job_id": job_id,
                "variant_id": variant_id,
                "base_revision": base_revision,
                "slides": slides,
                **({"order": order} if order is not None else {}),
                **({"template_logo": template_logo} if template_logo is not None else {}),
            }
        )
        if (
            not patch["slides"]
            and patch.get("order") is None
            and patch.get("template_logo") is None
        ):
            raise ConflictError("patch_empty", "В запросе нет ни правок, ни нового порядка")
        from presentation_designer.pipeline.stubs import explicit_plan

        plan = explicit_plan(plan)
        gen = self.state.get_generation(job_id)
        template = self.state.get_template(gen["template_id"])
        package = self.state.get_package(gen["package_id"])
        violations = validate_patch(
            patch, plan, deck, template.get("profile"), package.get("package")
        )
        if violations:
            raise ConflictError(
                "patch_invalid",
                "Правки не применимы к этой ревизии: " + "; ".join(str(v) for v in violations[:6]),
                {"violations": [str(v) for v in violations[:20]]},
            )
        if patch_is_noop(patch, plan):
            raise ConflictError("patch_empty", "Правки совпадают с текущей ревизией: менять нечего")
        try:
            self._resolve_patch_files(job_id, patch)
        except PatchError as e:
            raise ConflictError(e.code, str(e), e.details) from e
        patch_id = new_id("patch")
        touched = [str(s["slide_id"]) for s in patch["slides"]]
        self.state.create_job(
            kind="slide_patch",
            job_id=patch_id,
            stage="queued",
            parent_job_id=job_id,
            result={"generation_result_url": f"/api/generations/{job_id}"},
            progress={"percent": 0, "message": "Применяю правки редактора"},
        )
        first_index = next(
            (
                i
                for i, s in enumerate(sorted(plan["slides"], key=lambda s: int(s["order"])))
                if str(s["slide_id"]) in touched
            ),
            0,
        )
        self.state.create_repair(
            repair_job_id=patch_id,
            job_id=job_id,
            variant_id=variant_id,
            base_revision=base_revision,
            issue_ids=[],
            kind="patch",
            slide_index=first_index,
            instruction=None,
            patch=patch,
        )
        rq_id = self.executor.enqueue(
            self.settings.queue.repair_queue,
            task_edit,
            (patch_id,),
            depends_on=[],
            timeout=self.settings.timeouts.stage_compose_s
            + self.settings.timeouts.stage_export_s
            + self.settings.timeouts.stage_audit_s
            + 30,
            description=f"patch {patch_id}",
        )
        self.state.update_job(patch_id, rq_ids=[rq_id])
        return self.state.get_job(patch_id)

    def _resolve_patch_files(self, job_id: str, patch: JsonDict) -> None:
        """Источники `file` в правках: файл проекта задания, картинка; sha256 дописывается
        в патч, чтобы сборка нашла байты в хранилище без базы."""
        from presentation_designer.generation.patch import PatchError

        project = self.state.find_project_by_job(job_id)
        project_id = str(project["project_id"]) if project else None
        for entry in patch.get("slides") or []:
            for override in entry.get("overrides") or []:
                for spec in (override.get("picture"), override.get("background")):
                    source = (spec or {}).get("source") or {}
                    if source.get("kind") != "file":
                        continue
                    file_id = str(source.get("file_id") or "")
                    try:
                        record = self.state.get_file(file_id, project_id)
                    except NotFound as e:
                        raise PatchError(
                            "file_not_found", f"Файл {file_id} не найден в проекте"
                        ) from e
                    if str((record.get("check") or {}).get("format") or "") != "image":
                        raise PatchError(
                            "file_not_image", f"Файл {record.get('name') or file_id} не картинка"
                        )
                    source["sha256"] = str(record["sha256"])
                    source.setdefault("name", str(record.get("name") or ""))

    # ----- бриф из сообщения -----

    async def extract_brief(
        self, text: str, brief: JsonDict | None = None, settings: JsonDict | None = None
    ) -> JsonDict:
        """Слой `brief` с тайм-аутом `timeouts.brief_s`: модель или эвристика (`source`)."""
        import asyncio

        from presentation_designer.parsing.content.brief import extract_brief as heuristic

        try:
            return await asyncio.wait_for(
                self.layers.extract_brief(BriefInput(text, brief, settings)),
                timeout=self.settings.timeouts.brief_s + 2,
            )
        except TimeoutError:
            log.warning(
                "извлечение брифа не уложилось в %s с: эвристика", self.settings.timeouts.brief_s
            )
        except Exception:
            log.exception("слой brief завершился ошибкой: эвристика")
        return heuristic(text, brief)

    # ----- служебное -----

    def template_profile_key(self, sha256: str) -> str:
        """Ключ кэша профиля: заглушки зависят только от файла и версии приложения."""
        if self.layers.modes.get("parsing.template") != "real":
            return f"stub:{__version__}:{sha256}"
        from presentation_designer.parsing.template.analyzer import profile_key

        skill = None
        vlm_model = None
        try:
            from presentation_designer.llm.skills import get_skill

            skill = get_skill("template_analyzer")
            vlm_model = get_models_config().roles["vlm"].model
        except Exception:
            log.warning("скилл template_analyzer или конфиг моделей недоступны для ключа профиля")
        return profile_key(sha256, settings=self.settings, skill=skill, vlm_model=vlm_model)

    def versions(self) -> JsonDict:
        """Версии для воспроизведения: скиллы и промпты из skills/, модели из models.yaml.
        Слои пока заглушечные, но версии записываются уже сейчас — их читает ключ кэша."""
        skills, prompts = version_refs()
        return {
            "app": __version__,
            "contracts": CONTRACTS_VERSION,
            "skills": skills,
            "prompts": prompts,
            "models": model_refs(),
            "renderer": {"name": "stub", "version": __version__},
        }

    def reconcile(self) -> list[str]:
        """Задания без живого исполнителя и с истёкшим сроком помечаются ошибкой;
        возвращает их идентификаторы."""
        touched: list[str] = []
        now = datetime.now(UTC)
        for job in self.state.list_active_jobs():
            reason: JsonDict | None = None
            if (
                job["deadline_at"]
                and datetime.fromisoformat(job["deadline_at"].replace("Z", "+00:00")) < now
            ):
                reason = {
                    "code": "deadline_exceeded",
                    "message": "Общий срок задания истёк",
                    "retryable": True,
                }
            elif job["rq_ids"]:
                statuses = [self.executor.status(rid) for rid in job["rq_ids"]]
                if all(s in (None, "failed", "stopped", "canceled") for s in statuses):
                    reason = {
                        "code": "worker_lost",
                        "message": "Исполнитель задания потерян; повторите запуск",
                        "retryable": True,
                    }
            if reason is None:
                continue
            ts = now_iso()
            if job["kind"] == "generation":
                for variant in self.state.get_variants(job["job_id"]):
                    if variant["status"] in {"pending", "running"}:
                        self.state.update_variant(
                            job["job_id"], variant["variant_id"], status="failed", error=reason
                        )
            self.state.update_job(
                job["job_id"], status="failed", stage="done", finished_at=ts, error=reason
            )
            touched.append(job["job_id"])
        return touched

    def is_canceled(self, job_id: str) -> bool:
        try:
            return bool(self.state.get_generation(job_id)["canceled"])
        except NotFound:
            return True

    # ----- сборка мусора -----

    GC_LOCK = "pd:gc:scheduled"

    @property
    def gc_report_path(self) -> pathlib.Path:
        return self.settings.runs_dir / "gc" / "last.json"

    def ensure_gc_scheduled(self) -> bool:
        """Раз в `retention.gc_interval_s` ставит сборку мусора в очередь анализа.
        Вызывается из периодической сверки API; ключ в Valkey не даёт поставить задачу дважды."""
        interval = max(60, self.settings.retention.gc_interval_s)
        if not self.executor.try_lock(self.GC_LOCK, interval):
            return False
        self.executor.enqueue(
            self.settings.queue.analysis_queue,
            task_gc,
            (),
            depends_on=[],
            timeout=600,
            description="gc",
        )
        return True

    def collect_garbage(
        self, *, dry_run: bool = False, grace_hours: float | None = None
    ) -> JsonDict:
        from presentation_designer.pipeline.gc import collect_garbage

        report = collect_garbage(
            self.state,
            self.files,
            self.artifacts,
            self.settings.retention,
            dry_run=dry_run,
            grace_override=timedelta(hours=grace_hours) if grace_hours is not None else None,
            report_path=self.gc_report_path,
        )
        return report.to_dict()


_current: Orchestrator | None = None
_lock = threading.Lock()


def set_current(orchestrator: Orchestrator) -> None:
    global _current
    _current = orchestrator


def build_orchestrator(
    settings: Settings | None = None, executor: Executor | None = None
) -> Orchestrator:
    """Собирает оркестратор из настроек: SQLite, хранилища и слои по execution.mode."""
    from presentation_designer.pipeline.real import build_layers

    settings = settings or get_settings()
    state = State(settings.db_path)
    files = FileStore(
        settings.uploads_dir,
        max_upload_mb=settings.limits.max_upload_mb,
        max_unzipped_mb=settings.limits.max_unzipped_mb,
    )
    artifacts = ArtifactStore(settings.artifacts_dir)
    layers = build_layers(settings)
    if executor is None:
        import os

        mode = os.environ.get("PD_QUEUE_MODE", "rq")
        executor = (
            InlineExecutor()
            if mode == "inline"
            else RQExecutor(os.environ.get("PD_VALKEY_URL", "redis://localhost:6379/0"), settings)
        )
    return Orchestrator(settings, state, files, artifacts, layers, executor)


def _record_llm(o: Orchestrator, job_id: str) -> None:
    """Вызовы модели закончившегося этапа — в метрики задания.

    Регистратор живёт вместе с клиентом, то есть на весь процесс воркера. Если его не
    опустошать после каждого этапа, вызовы разных заданий смешаются, а в отчёте останется
    ноль вызовов при работающей модели — ровно это и было.
    """
    try:
        usage = o.layers.take_llm_usage()
        if usage:
            o.state.add_llm_usage(job_id, usage)
    except Exception:  # учёт не должен ронять задание
        log.exception("метрики вызовов модели не записаны: %s", job_id)


def get_orchestrator() -> Orchestrator:
    global _current
    with _lock:
        if _current is None:
            _current = build_orchestrator()
        return _current


# ---------- задачи воркеров ----------


# Вариант, который собирается первым: средняя плотность — то, что человек видит по умолчанию.
PRIMARY_VARIANT = "balanced"

# Готовая презентация открыта, а разбор в фоне не удался: правки из чата ждать нечего.
PARSE_FAILED = (
    "Разбор презентации не удался: правки из чата недоступны, редактор и скачивание работают. "
    "Повторите разбор кнопкой в чате"
)


def ordered_variants(variant_ids: list[str]) -> list[str]:
    """Порядок сборки: основной вариант первым, остальные в порядке запроса."""
    if PRIMARY_VARIANT not in variant_ids:
        return list(variant_ids)
    return [PRIMARY_VARIANT, *(v for v in variant_ids if v != PRIMARY_VARIANT)]


def task_analyze(template_id: str) -> None:
    o = get_orchestrator()
    try:
        template = o.state.get_template(template_id)
    except NotFound:
        log.info("шаблон %s удалён до начала анализа", template_id)
        return
    job_id = template["job_id"]
    o.state.job_started(job_id)
    o.state.update_template(template_id, status="running")
    started = now_iso()
    try:
        path = o.files.open(template["sha256"])
        out = o.layers.analyze(
            AnalyzeInput(
                template_id, template["sha256"], template["name"], template["size_bytes"], path
            )
        )
        try:
            o.state.get_template(template_id)
        except NotFound:
            log.info("шаблон %s удалён во время анализа, результат не сохраняется", template_id)
            return
        with o.artifacts.stage_dir(o.artifacts.template_dir(template_id)) as tmp:
            for name, data in out.previews.items():
                target = tmp / name
                target.parent.mkdir(parents=True, exist_ok=True)
                target.write_bytes(data)
        previews = sorted(out.previews)
        o.state.update_template(
            template_id, status="succeeded", profile=out.profile, previews=previews
        )
        _record_llm(o, job_id)
        o.state.add_stage(
            job_id,
            {
                "stage": "analyze",
                "status": "done",
                "started_at": started,
                "duration_ms": _ms_since(started),
            },
        )
        o.state.update_job(
            job_id,
            status="succeeded",
            stage="done",
            finished_at=now_iso(),
            progress={
                "percent": 100,
                "message": f"Профиль готов: {len(out.profile.get('patterns', []))} паттернов",
            },
        )
    except Exception as e:
        log.exception("анализ шаблона %s не удался", template_id)
        if isinstance(e, StageError):
            error = e.as_dict("analyze")  # неподдерживаемый файл: код и текст из слоя
        else:
            error = {
                "code": "analyze_failed",
                "message": str(e) or e.__class__.__name__,
                "stage": "analyze",
                "retryable": True,
            }
        o.state.update_template(template_id, status="failed", error=error)
        o.state.add_stage(
            job_id,
            {
                "stage": "analyze",
                "status": "failed",
                "started_at": started,
                "duration_ms": _ms_since(started),
            },
        )
        o.state.update_job(
            job_id, status="failed", stage="done", finished_at=now_iso(), error=error
        )


def task_import(package_id: str) -> None:
    o = get_orchestrator()
    package = o.state.get_package(package_id)
    job_id = package["job_id"]
    o.state.job_started(job_id)
    o.state.update_package(package_id, status="running")
    started = now_iso()
    try:
        sources: list[SourceFile] = []
        for file_id in package["file_ids"]:
            row = o.state.get_file(file_id)
            fmt = row.get("check", {}).get("format") or format_for(row["name"])
            path = o.files.path_for(row["sha256"])
            sources.append(
                SourceFile(
                    file_id,
                    row["name"],
                    row["sha256"],
                    row["size_bytes"],
                    fmt,
                    path if path.is_file() else None,
                )
            )
        out = o.layers.import_content(
            ImportInput(package_id, package["mode"], sources, package["brief"])
        )
        if out.assets or out.report:
            with o.artifacts.stage_dir(o.artifacts.package_dir(package_id)) as tmp:
                for name, data in out.assets.items():
                    target = tmp / name
                    target.parent.mkdir(parents=True, exist_ok=True)
                    target.write_bytes(data)
                if out.report:
                    (tmp / "import-report.json").write_text(
                        json.dumps(out.report, ensure_ascii=False, indent=2), encoding="utf-8"
                    )
        o.state.update_package(package_id, status="succeeded", package=out.package)
        _record_llm(o, job_id)
        pk = out.package
        o.state.add_stage(
            job_id,
            {
                "stage": "import",
                "status": "done",
                "started_at": started,
                "duration_ms": _ms_since(started),
            },
        )
        message = f"Импорт готов: {len(pk.get('blocks', []))} блоков, {len(pk.get('facts', []))} фактов, {len(pk.get('datasets', []))} таблиц, {len(pk.get('assets', []))} изображений"  # noqa: E501
        o.state.update_job(
            job_id,
            status="succeeded",
            stage="done",
            finished_at=now_iso(),
            progress={"percent": 100, "message": message},
        )
    except Exception as e:
        log.exception("импорт пакета %s не удался", package_id)
        error = (
            e.as_dict("import")
            if isinstance(e, StageError)
            else {
                "code": "import_failed",
                "message": str(e) or e.__class__.__name__,
                "stage": "import",
                "retryable": True,
            }
        )
        o.state.update_package(package_id, status="failed", error=error)
        o.state.update_job(
            job_id, status="failed", stage="done", finished_at=now_iso(), error=error
        )


def _fail_generation(o: Orchestrator, job_id: str, error: JsonDict) -> None:
    for variant in o.state.get_variants(job_id):
        if variant["status"] in {"pending", "running"}:
            o.state.update_variant(job_id, variant["variant_id"], status="failed", error=error)
    o.state.update_job(job_id, status="failed", stage="done", finished_at=now_iso(), error=error)


def task_story(job_id: str) -> None:
    o = get_orchestrator()
    job = o.state.get_job(job_id)
    gen = o.state.get_generation(job_id)
    if gen["canceled"] or job["status"] in TERMINAL:
        return
    o.state.job_started(job_id)
    package = o.state.get_package(gen["package_id"])
    template = o.state.get_template(gen["template_id"])
    if package["status"] == "failed":
        _fail_generation(
            o,
            job_id,
            {
                "code": "import_failed",
                "message": "Импорт содержания завершился ошибкой",
                "stage": "import",
                "retryable": True,
            },
        )
        return
    if package["status"] != "succeeded" or package["package"] is None:
        _fail_generation(
            o,
            job_id,
            {
                "code": "package_not_ready",
                "message": "Контент-пакет ещё не импортирован",
                "stage": "import",
                "retryable": True,
            },
        )
        return
    import_job = o.state.get_job(package["job_id"])
    o.state.add_stage(
        job_id,
        {
            "stage": "import",
            "status": "done",
            "duration_ms": _duration(import_job),
            "cache_hit": import_job["created_at"] < job["created_at"],
        },
    )
    if template["status"] in TERMINAL:
        template_job = o.state.get_job(template["job_id"])
        o.state.add_stage(
            job_id,
            {
                "stage": "analyze",
                "status": "done" if template["status"] == "succeeded" else "failed",
                "duration_ms": _duration(template_job),
                "cache_hit": template_job["created_at"] < job["created_at"],
            },
        )
    o.state.update_job(
        job_id, stage="story", progress={"percent": 5, "message": "Строится общий смысловой план"}
    )
    started = now_iso()
    story_input = StoryInput(package["package"], gen["request"].get("settings") or {})
    story: JsonDict | None = None
    hit = False
    original = list((gen["request"].get("settings") or {}).get("variants") or []) == ["original"]
    if original and (template["status"] != "succeeded" or template["profile"] is None):
        _fail_generation(
            o,
            job_id,
            {
                "code": "template_failed",
                "message": "Анализ презентации не завершился успешно",
                "stage": "analyze",
                "retryable": True,
            },
        )
        return
    try:
        if original:
            # Исходная презентация: тезис на слайд по профилю и пакету, без модели.
            story = original_story(package["package"], template["profile"], story_input.settings)
        # Готовый план переиспользуется по смысловому ключу (содержание, эффективный бриф,
        # модель, промпт), а не по package_id: смена языка или брифа даёт новый ключ.
        elif not story_input.settings.get("force_regenerate"):
            story = o.state.find_story(o.layers.story_key(story_input))
            hit = story is not None
        if story is None:
            story = o.layers.story(story_input)
    except Exception as e:
        log.exception("смысловой план %s не удался", job_id)
        _fail_generation(
            o,
            job_id,
            e.as_dict("story")
            if isinstance(e, StageError)
            else {
                "code": "story_failed",
                "message": str(e) or e.__class__.__name__,
                "stage": "story",
                "retryable": True,
            },
        )
        return
    _record_llm(o, job_id)
    o.state.update_generation(job_id, story=story, story_hit=1 if hit else 0)
    o.state.add_stage(
        job_id,
        {
            "stage": "story",
            "status": "done",
            "started_at": started,
            "duration_ms": _ms_since(started),
            "cache_hit": hit,
        },
    )
    # Без текста: ход вариантов («Собираю первый вариант», «Готово 1 из 3…») строка результата
    # считает сама по их состоянию (results._progress_message).
    o.state.update_job(
        job_id,
        stage="plan",
        progress={"percent": 20, "message": "Планы вариантов"} if original else {"percent": 20},
    )


def _publish_original_copy(o: Orchestrator, job_id: str, template: JsonDict) -> bool:
    """Копия загруженной презентации — ревизия r1 варианта original, байт в байт. Редактору
    нужен только файл, поэтому копия публикуется сразу: без очереди, без модели и без разбора
    пакета (число слайдов — по списку в presentation.xml). True — r1 есть."""
    variant_id, revision = ORIGINAL_VARIANT, 1
    try:
        o.state.get_revision(job_id, variant_id, revision)
        return True
    except NotFound:
        pass
    source = o.files.path_for(str(template["sha256"]))
    if not source.is_file():
        log.warning("копия презентации %s: файл %s не найден", job_id, template["sha256"])
        return False
    started = time.perf_counter()
    prefix = o.artifacts.prefix(variant_id, revision)
    with o.artifacts.stage_revision(job_id, variant_id, revision) as staging:
        shutil.copyfile(source, staging.path("deck.pptx"))
    manifest = staging.manifest
    o.state.add_revision(
        job_id=job_id,
        variant_id=variant_id,
        revision=revision,
        artifacts_prefix=prefix,
        manifest=manifest,
        pptx_hash=_pptx_hash(manifest, prefix),
    )
    # ready_at — первая публикация: дальше его не переписывает никто (полная сборка тоже).
    o.state.update_variant(
        job_id,
        variant_id,
        status="running",
        slide_count=_slide_count(source) or None,
        ready_at=now_iso(),
    )
    o.state.job_started(job_id)
    o.state.update_job(
        job_id, stage="analyze", progress={"percent": 5, "message": "Открываю презентацию"}
    )
    log.info(
        "копия презентации %s опубликована за %d мс",
        job_id,
        int((time.perf_counter() - started) * 1000),
    )
    return True


def task_original_preview(job_id: str) -> None:
    """Фон уже открытой презентации (очередь interactive): подписи пустых плейсхолдеров и
    диаграммы-картинки делают ревизию r2, если в копии что-то изменилось, затем PDF и
    миниатюры последней ревизии — их переиспользует полная сборка. r1 публикует сам запрос
    создания задания; здесь копия делается, только если её там не вышло."""
    o = get_orchestrator()
    job = o.state.get_job(job_id)
    gen = o.state.get_generation(job_id)
    if gen["canceled"] or job["status"] in TERMINAL:
        return
    template = o.state.get_template(gen["template_id"])
    try:
        if not _publish_original_copy(o, job_id, template):
            return
    except Exception:
        log.exception("копия презентации %s не опубликована", job_id)
        return
    o.state.job_started(job_id)
    started = time.perf_counter()
    settings = gen["request"].get("settings") or {}
    try:
        revision = _refine_original(o, job_id, settings)
    except Exception:  # подписи и диаграммы необязательны: остаётся показанная копия
        log.exception("подписи и диаграммы копии %s не обработаны", job_id)
        revision = 1
    refined_ms = int((time.perf_counter() - started) * 1000)
    try:
        pages = _render_revision(o, job_id, ORIGINAL_VARIANT, revision, template)
    except Exception:  # без рендера полная сборка отрендерит сама
        log.exception("рендер копии %s r%d не удался", job_id, revision)
        return
    log.info(
        "фон копии %s: r%d, подписи и диаграммы %d мс, с рендером %d мс, %d страниц",
        job_id,
        revision,
        refined_ms,
        int((time.perf_counter() - started) * 1000),
        pages,
    )


def _refine_original(o: Orchestrator, job_id: str, settings: JsonDict) -> int:
    """Подписи пустых плейсхолдеров и диаграммы-картинки копии → ревизия r2; номер ревизии,
    которую теперь показывать. r1 не переписывается: офисная копия r1 создаётся один раз по
    своему source, и её байты менять бессмысленно. Ничего не изменилось — r2 нет, отчёт о
    чтении диаграмм (если было что читать) дописывается в r1."""
    variant_id = ORIGINAL_VARIANT
    latest = int(o.state.get_variant(job_id, variant_id)["revision"] or 1)
    if latest > 1:  # повторный запуск задачи: r2 уже опубликована
        return latest
    shown = o.artifacts.revision_dir(job_id, variant_id, 1)
    out: Any = None
    with o.artifacts.stage_revision(job_id, variant_id, 2) as staging:
        deck = staging.path("deck.pptx")
        shutil.copyfile(shown / "deck.pptx", deck)
        prompted = _preview_prompts(deck, settings)
        out = _preview_charts(o, job_id, deck)
        logos = _preview_logos(job_id, deck)
        if prompted or logos or (out is not None and out.replaced):
            if out is not None and out.report:
                staging.write_json("charts.json", out.report)
        else:
            staging.discard = True
    if staging.published:
        prefix = o.artifacts.prefix(variant_id, 2)
        o.state.add_revision(
            job_id=job_id,
            variant_id=variant_id,
            revision=2,
            artifacts_prefix=prefix,
            manifest=staging.manifest,
            pptx_hash=_pptx_hash(staging.manifest, prefix),
        )
        _chart_warning(o, job_id, out)
        log.info(
            "копия %s: r2 — подписей %d, диаграмм заменено %d, знаков на слайдах %d",
            job_id,
            prompted,
            out.replaced if out is not None else 0,
            logos,
        )
        return 2
    if out is not None and out.report:
        prefix = o.artifacts.prefix(variant_id, 1)
        republish = Staging(dir=shown, final=shown, prefix=prefix)
        republish.write_json("charts.json", out.report)
        manifest = o.artifacts.publish(republish)
        o.state.update_revision(job_id, variant_id, 1, manifest=manifest)
    _chart_warning(o, job_id, out)
    return 1


def _chart_warning(o: Orchestrator, job_id: str, out: Any) -> None:
    """Фраза о диаграммах-картинках — в предупреждения задания в момент публикации ревизии:
    чат говорит её сразу, а не в конце разбора."""
    if out is None or not out.message:
        return
    current = o.state.get_generation(job_id)["warnings"]
    if any(w.get("code") == "chart_images" for w in current):
        return
    o.state.update_generation(
        job_id, warnings=[*current, {"code": "chart_images", "message": out.message}]
    )


def _render_revision(
    o: Orchestrator, job_id: str, variant_id: str, revision: int, template: JsonDict
) -> int:
    """PDF и миниатюры опубликованной ревизии: файлы дописываются в её каталог, манифест
    переписывается атомарно. Возвращает число страниц."""
    final = o.artifacts.revision_dir(job_id, variant_id, revision)
    prefix = o.artifacts.prefix(variant_id, revision)
    staging = Staging(dir=final, final=final, prefix=prefix)
    titles = _slide_titles(final / "deck.pptx")
    exported = o.layers.export(
        ExportInput(
            job_id,
            variant_id,
            revision,
            staging,
            titles or [f"Слайд {i + 1}" for i in range(_slide_count(final / "deck.pptx"))],
            str(template["name"]).rsplit(".", 1)[0] or "Презентация",
        )
    )
    manifest = o.artifacts.publish(staging)
    o.state.update_revision(
        job_id, variant_id, revision, manifest=manifest, pptx_hash=_pptx_hash(manifest, prefix)
    )
    o.state.update_variant(job_id, variant_id, slide_count=len(exported.thumbnails))
    return len(exported.thumbnails)


def _slide_titles(path: pathlib.Path) -> list[str]:
    try:
        from pptx import Presentation

        titles: list[str] = []
        for slide in Presentation(str(path)).slides:
            title = slide.shapes.title.text if slide.shapes.title is not None else ""
            titles.append(" ".join(title.split())[:120])
        return titles
    except Exception:
        return []


def _preview_prompts(path: pathlib.Path, settings: JsonDict) -> int:
    """Пустые плейсхолдеры копии получают текст подсказки, как в полной сборке: ONLYOFFICE
    показывает эту копию, и «Заголовок слайда» должен быть текстом, а не подсказкой, которая
    исчезает по щелчку и не видна в миниатюрах. Возвращает число заполненных плейсхолдеров.

    В синхронный путь не входит: сохранение пакета на 55 слайдах на сервере — 0,7 с, с разбором
    около секунды (замер 24.09.2026), а копия должна публиковаться мгновенно."""
    from pptx import Presentation

    from presentation_designer.layout.text import fill_empty_placeholders

    prs = Presentation(str(path))
    filled = int(fill_empty_placeholders(prs.slides, str(settings.get("language") or "ru")) or 0)
    if filled:
        prs.save(str(path))
    return filled


def _preview_logos(job_id: str, deck: pathlib.Path) -> int:
    """Знак шаблона копии — с макетов на слайды, как в полной сборке: в редакторе его двигают
    и меняют на слайде. Профиля ещё нет (копия показывается раньше, чем кончается разбор),
    поэтому знак ищется по самому файлу тем же правилом анализатора. Идёт после диаграмм:
    картинки знака на слайдах им читать незачем. Возвращает число поставленных картинок."""
    from presentation_designer.generation.office_logo import promote_file

    try:
        data, placed = promote_file(deck)
    except Exception:  # перенос необязателен: знак остаётся на макетах
        log.exception("знак шаблона копии %s не перенесён на слайды", job_id)
        return 0
    if data is not None:
        deck.write_bytes(data)
    return placed


def _preview_charts(o: Orchestrator, job_id: str, deck: pathlib.Path) -> Any:
    """Диаграммы-картинки копии → нативные диаграммы на месте; полная сборка повторит те же
    замены по charts.json. None — чтение сломалось: остаётся копия как есть."""

    def progress(message: str, share: float) -> None:
        percent = 10 + round(25 * max(0.0, min(1.0, share)))
        o.state.update_job(job_id, progress={"percent": percent, "message": message})

    try:
        out = o.layers.chart_images(deck, progress)
    except Exception:  # замена необязательна: покажется исходная презентация
        log.exception("диаграммы-картинки %s не заменены", job_id)
        return None
    finally:
        _record_llm(o, job_id)
    log.info("диаграммы-картинки %s: заменено %d; %s", job_id, out.replaced, out.message or "—")
    return out


def _preview_chart_report(preview_dir: pathlib.Path | None) -> JsonDict | None:
    """charts.json предварительной ревизии. Без него (чтения не было или оно сломалось)
    полная сборка ничего не заменяет, иначе колода разойдётся с уже показанной копией и её
    рендером; None — предварительной ревизии нет, сборка читает сама."""
    if preview_dir is None:
        return None
    try:
        report = json.loads((preview_dir / "charts.json").read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}
    return report if isinstance(report, dict) else {}


def _slide_count(path: pathlib.Path) -> int:
    """Число слайдов по списку sldIdLst в ppt/presentation.xml: без разбора всего пакета,
    который на 55 слайдах стоит десятые доли секунды."""
    import zipfile

    from lxml import etree

    try:
        with zipfile.ZipFile(path) as zf:
            root = etree.fromstring(zf.read("ppt/presentation.xml"))
        return sum(1 for el in root.iter() if etree.QName(el).localname == "sldId")
    except Exception:
        try:
            from pptx import Presentation

            return len(Presentation(str(path)).slides)
        except Exception:
            return 0


def task_variant(job_id: str, variant_id: str) -> None:
    o = get_orchestrator()
    job = o.state.get_job(job_id)
    gen = o.state.get_generation(job_id)
    if gen["canceled"] or job["status"] in TERMINAL:
        return
    template = o.state.get_template(gen["template_id"])
    if template["status"] != "succeeded" or template["profile"] is None:
        error = {
            "code": "template_failed",
            "message": "Анализ шаблона не завершился успешно",
            "stage": "analyze",
            "retryable": True,
        }
        o.state.update_variant(job_id, variant_id, status="failed", error=error)
        return
    if gen["story"] is None:
        o.state.update_variant(
            job_id,
            variant_id,
            status="failed",
            error={
                "code": "story_missing",
                "message": "Смысловой план не построен",
                "stage": "story",
                "retryable": True,
            },
        )
        return
    package = o.state.get_package(gen["package_id"])
    o.state.update_variant(job_id, variant_id, status="running")
    o.state.add_stage(job_id, {"stage": "analyze", "status": "done", "cache_hit": True})
    template_path = o.files.path_for(template["sha256"])
    revision = 1
    # Готовая презентация уже показана: сборка пишет в последнюю её ревизию (r2 с диаграммами,
    # если она есть, иначе r1) — запись обновляется, а не создаётся заново, замены диаграмм
    # повторяются по charts.json этой ревизии, рендер переиспользуется, если успел.
    shown_dir: pathlib.Path | None = None
    if variant_id == ORIGINAL_VARIANT:
        latest = int(o.state.get_variant(job_id, variant_id)["revision"] or 1)
        try:
            o.state.get_revision(job_id, variant_id, latest)
            revision = latest
            shown_dir = o.artifacts.revision_dir(job_id, variant_id, revision)
        except NotFound:
            shown_dir = None
    prefix = o.artifacts.prefix(variant_id, revision)
    preview_dir = shown_dir if shown_dir and (shown_dir / "deck.pdf").is_file() else None

    def emit(event: str, data: JsonDict) -> None:
        if event == "stage":
            o.state.add_variant_stage(job_id, variant_id, data)
            if data.get("status") == "running":
                o.state.update_job(job_id, stage=data["stage"])
        elif event == "files_ready":
            manifest = o.artifacts.publish(staging)
            if shown_dir is not None:
                # Ревизия уже опубликована копией: запись обновляется.
                o.state.update_revision(
                    job_id,
                    variant_id,
                    revision,
                    manifest=manifest,
                    pptx_hash=_pptx_hash(manifest, prefix),
                )
            else:
                o.state.add_revision(
                    job_id=job_id,
                    variant_id=variant_id,
                    revision=revision,
                    artifacts_prefix=prefix,
                    manifest=manifest,
                    pptx_hash=_pptx_hash(manifest, prefix),
                )
            # ready_at — первая публикация файла; у готовой презентации это копия в запросе,
            # и сборка его не переписывает (иначе first_file_ready_ms врёт).
            first = (
                {}
                if o.state.get_variant(job_id, variant_id).get("ready_at")
                else {"ready_at": now_iso()}
            )
            o.state.update_variant(
                job_id,
                variant_id,
                status="ready",
                slide_count=data.get("slide_count"),
                **first,
                audit={
                    "status": "running",
                    "coverage_complete": False,
                    "issues_total": 0,
                    "blocking": 0,
                },
            )

    with o.artifacts.stage_revision(job_id, variant_id, revision) as staging:
        ctx = VariantContext(
            job_id=job_id,
            variant_id=variant_id,
            template_profile=template["profile"],
            template_path=template_path if template_path.is_file() else None,
            package=package["package"] or {},
            story=gen["story"],
            settings=gen["request"].get("settings") or {},
            staging=staging,
            revision=revision,
            is_canceled=lambda: o.is_canceled(job_id),
            package_dir=o.artifacts.package_dir(gen["package_id"]),
            prerendered=preview_dir,
            chart_report=_preview_chart_report(shown_dir) if variant_id == "original" else None,
        )
        outcome = run_variant(o.layers, ctx, emit)
        _record_llm(o, job_id)
        if outcome.status == "failed":
            o.state.update_variant(
                job_id, variant_id, status="failed", error=outcome.error, stages=outcome.stages
            )
            return
    manifest = o.artifacts.read_manifest(job_id, variant_id, revision)
    o.state.update_revision(job_id, variant_id, revision, manifest=manifest)
    if outcome.warnings:
        # Единственный вариант original пишет предупреждения плана в задание сам: варианты
        # вёрстки идут параллельно и в общий список не пишут.
        current = o.state.get_generation(job_id)["warnings"]
        o.state.update_generation(job_id, warnings=[*current, *outcome.warnings])
    audit = dict(outcome.audit or {})
    audit["report_artifact"] = f"{prefix}audit.json"
    o.state.update_variant(
        job_id,
        variant_id,
        status=outcome.status,
        slide_count=outcome.slide_count,
        audited_at=now_iso(),
        audit=audit,
        stages=outcome.stages,
    )


def task_finalize(job_id: str) -> None:
    o = get_orchestrator()
    job = o.state.get_job(job_id)
    if job["status"] in TERMINAL and not _deadline_only(job):
        return
    gen = o.state.get_generation(job_id)
    variants = o.state.get_variants(job_id)
    unfinished = [v for v in variants if v["status"] in {"pending", "running"}]
    if unfinished and not gen["canceled"]:
        for v in unfinished:
            o.state.update_variant(
                job_id,
                v["variant_id"],
                status="failed",
                error={
                    "code": "variant_lost",
                    "message": "Вариант не завершился",
                    "retryable": True,
                },
            )
        variants = o.state.get_variants(job_id)
    failed = [v for v in variants if v["status"] == "failed"]
    needs_review = any(
        v["status"] == "needs_review"
        or (v.get("audit") and not v["audit"].get("coverage_complete"))
        for v in variants
    )
    ts = now_iso()
    status: str
    error: JsonDict | None
    warnings = list(gen["warnings"])
    if failed and len(failed) < len(variants):
        warnings.append(
            {
                "code": "variant_failed",
                "message": f"Вариант {', '.join(v['variant_id'] for v in failed)} не собран; доступны остальные варианты",  # noqa: E501
            }
        )
    warnings.extend(_plan_distinctness(o, job_id, variants))
    warnings.extend(_slide_count_short(o, job_id, variants, gen.get("request") or {}))
    if gen["canceled"]:
        status, error = (
            "canceled",
            {"code": "canceled", "message": "Задание отменено пользователем"},
        )
    elif len(failed) == len(variants):
        status, error = (
            "failed",
            job.get("error")
            if _deadline_only(job)
            else {
                "code": "all_variants_failed",
                "message": "Ни один вариант не собран",
                "retryable": True,
            },
        )
    elif failed or needs_review:
        status, error = "needs_review", None
    else:
        status, error = "succeeded", None
    o.state.update_generation(job_id, warnings=warnings)
    o.state.add_stage(job_id, {"stage": "finalize", "status": "done", "duration_ms": 0})
    o.state.update_job(
        job_id,
        status=status,
        stage="done",
        finished_at=ts,
        error=error,
        progress={"percent": 100, "message": _final_message(variants)},
    )


def _deadline_only(job: JsonDict) -> bool:
    """Задание помечено ошибкой только по общему сроку (`Orchestrator.reconcile`), а варианты
    ещё собирались: итог подводится по ним. Колода Education на мосте к Claude собралась за
    15 мин 31 с при сроке 15 мин — три готовых варианта, а проект показывал ошибку."""
    return job["status"] == "failed" and (job.get("error") or {}).get("code") == (
        "deadline_exceeded"
    )


def _slide_count_short(
    o: Orchestrator, job_id: str, variants: list[JsonDict], request: JsonDict
) -> list[JsonDict]:
    """Содержания меньше, чем просил пользователь: варианты собраны короче, и задание говорит
    об этом одной фразой. Диапазон по умолчанию сокращается молча — его никто не просил."""
    counts: list[int] = []
    for v in variants:
        if v["status"] in ("failed", "pending", "running"):
            continue
        path = o.artifacts.revision_dir(job_id, v["variant_id"], v["revision"]) / "plan.json"
        try:
            plan = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            continue
        if any(w.get("code") == "slide_count_short" for w in plan.get("warnings") or []):
            counts.append(int(v.get("slide_count") or len(plan.get("slides") or [])))
    if not counts:
        return []
    spec = (request.get("settings") or {}).get("slide_count") or {}
    got = f"{min(counts)}–{max(counts)}" if min(counts) != max(counts) else str(counts[0])
    if spec.get("exact"):
        exact = int(spec["exact"])
        asked = f"Просили {exact} {plural(exact, 'слайд', 'слайда', 'слайдов')}"
    elif spec.get("min") and spec.get("max"):
        asked = f"Просили {spec['min']}–{spec['max']} слайдов"
    else:
        asked = "Просили больше слайдов"
    return [
        {
            "code": "slide_count_short",
            "message": f"{asked} — содержания хватило на {got}. Добавьте материалы или "
            "расскажите о теме подробнее, и я соберу больше.",
        }
    ]


def _plan_distinctness(o: Orchestrator, job_id: str, variants: list[JsonDict]) -> list[JsonDict]:
    """Различимость настоящих планов вариантов по последовательности композиций и
    визуализации; пары без различий — предупреждение задания, а не переименование."""
    if o.layers.modes.get("generation.plan") != "real":
        return []
    plans: dict[str, JsonDict] = {}
    for v in variants:
        if v["status"] in ("failed", "pending", "running"):
            continue
        path = o.artifacts.revision_dir(job_id, v["variant_id"], v["revision"]) / "plan.json"
        try:
            plans[v["variant_id"]] = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            continue
    if len(plans) < 2:
        return []
    from presentation_designer.generation.variants import compare_plans, indistinct_warning

    return [indistinct_warning(entry) for entry in compare_plans(plans)["indistinct"]]


def _repair_overrides(
    o: Orchestrator,
    job_id: str,
    base_dir: pathlib.Path,
    base_report: JsonDict,
    issue_ids: list[str],
    plan: JsonDict,
) -> Any:
    """Правки, которыми чинятся выбранные находки.

    `None` — исправлять нечем по входам: нет описания собранной колоды или профиля шаблона
    (режим заглушек). Тогда задание идёт прежним путём и честно показывает, что это заглушка.
    """
    # Настоящие правки к плану имеют смысл только при настоящей вёрстке: в режиме заглушек
    # колода и профиль ненастоящие, и «исправление» там остаётся заглушкой, как и весь путь.
    if o.layers.modes.get("layout") != "real":
        return None
    deck_path = base_dir / "composed.json"
    if not deck_path.is_file():
        return None
    gen = o.state.get_generation(job_id)
    try:
        profile = (o.state.get_template(gen["template_id"]) or {}).get("profile")
    except NotFound:
        profile = None
    if not profile:
        return None
    from presentation_designer.audit.repair import build_repair
    from presentation_designer.generation.design_mode import profile_for_mode

    settings = gen["request"].get("settings") or {}
    if (plan.get("variant") or {}).get("variant_id") == "original":
        settings = {**settings, "design_mode": "mixed"}
    profile = profile_for_mode(profile, settings)
    deck = json.loads(deck_path.read_text(encoding="utf-8"))
    return build_repair(base_report, list(issue_ids), deck, profile, plan, settings=settings)


def _finish_repair_unchanged(o: Orchestrator, repair_job_id: str, job_id: str, fix: Any) -> None:
    """Ни одна выбранная находка кодом не чинится: ревизии нет, и в задании написано почему."""
    reasons = "; ".join(dict.fromkeys(f.note for f in fix.fixes if f.note)) or "исправлять нечего"
    o.state.update_repair(repair_job_id, result="unchanged", change_note=reasons)
    o.state.update_job(
        repair_job_id,
        status="succeeded",
        stage="done",
        finished_at=now_iso(),
        result={
            "unchanged": True,
            "change_note": reasons,
            "generation_result_url": f"/api/generations/{job_id}",
        },
        progress={"percent": 100, "message": "Исправлять нечего"},
    )


def task_repair(repair_job_id: str) -> None:
    o = get_orchestrator()
    repair = o.state.get_repair(repair_job_id)
    job_id, variant_id = repair["job_id"], repair["variant_id"]
    o.state.job_started(repair_job_id)
    o.state.update_job(
        repair_job_id,
        stage="repair",
        progress={"percent": 10, "message": "Исправление выбранных находок"},
    )
    try:
        variant = o.state.get_variant(job_id, variant_id)
        base = repair["base_revision"]
        if variant["revision"] != base:
            raise ConflictError("revision_stale", "Ревизия устарела до начала исправления")
        base_dir = o.artifacts.revision_dir(job_id, variant_id, base)
        base_report = json.loads((base_dir / "audit.json").read_text(encoding="utf-8"))
        plan = json.loads((base_dir / "plan.json").read_text(encoding="utf-8"))
        gen = o.state.get_generation(job_id)
        package = o.state.get_package(gen["package_id"])
        deck_title = (package["package"] or {}).get("brief", {}).get("title") or "Презентация"
        from presentation_designer.pipeline.stubs import plan_slides

        _, titles = plan_slides(plan)
        new_rev = base + 1
        # Исправление — это правки к плану того же вида, что и ручные из редактора, поэтому
        # ревизия собирается тем же путём (план → вёрстка → экспорт → аудит), а не своим.
        fix = _repair_overrides(o, job_id, base_dir, base_report, repair["issue_ids"], plan)
        if fix is not None:
            if not fix.patch["slides"]:
                _finish_repair_unchanged(o, repair_job_id, job_id, fix)
                return
            # Патч проверяется той же схемой, что и правки редактора: поля задания обязательны.
            o.state.update_repair(
                repair_job_id,
                patch={
                    "job_id": job_id,
                    "variant_id": variant_id,
                    "base_revision": base,
                    "slides": fix.patch["slides"],
                },
            )
            task_edit(repair_job_id)
            return
        prefix = o.artifacts.prefix(variant_id, new_rev)
        with o.artifacts.stage_revision(job_id, variant_id, new_rev) as staging:
            out = run_repair(
                o.layers,
                RepairInput(
                    job_id,
                    variant_id,
                    new_rev,
                    base,
                    base_report,
                    base_dir,
                    repair["issue_ids"],
                    staging,
                    titles,
                    deck_title,
                ),
                lambda e, d: None,
            )
        manifest = o.artifacts.read_manifest(job_id, variant_id, new_rev)
        o.state.add_revision(
            job_id=job_id,
            variant_id=variant_id,
            revision=new_rev,
            artifacts_prefix=prefix,
            manifest=manifest,
            repair_job_id=repair_job_id,
            changed_slide_ids=out.changed_slide_ids,
            pptx_hash=_pptx_hash(manifest, prefix),
        )
        summary = audit_summary(out.report)
        summary["report_artifact"] = f"{prefix}audit.json"
        status = (
            "needs_review"
            if summary["issues_total"] > 0 or not summary["coverage_complete"]
            else "ready"
        )
        o.state.update_variant(
            job_id, variant_id, status=status, audited_at=now_iso(), audit=summary
        )
        o.state.update_repair(
            repair_job_id,
            result="applied",
            new_revision=new_rev,
            changed_slide_ids=out.changed_slide_ids,
        )
        o.state.update_job(
            repair_job_id,
            status="succeeded",
            stage="done",
            finished_at=now_iso(),
            result={"revision": new_rev, "generation_result_url": f"/api/generations/{job_id}"},
            progress={
                "percent": 100,
                "message": "Исправления применены, затронутые слайды перепроверены",
            },
        )
    except Exception as e:
        log.exception("исправление %s не удалось", repair_job_id)
        message = str(e) or e.__class__.__name__
        code = getattr(e, "code", "repair_failed")
        o.state.update_repair(repair_job_id, result="failed", message=message)
        o.state.update_job(
            repair_job_id,
            status="failed",
            stage="done",
            finished_at=now_iso(),
            error={"code": code, "message": message, "stage": "repair", "retryable": True},
        )


def _edit_message(is_patch: bool, is_repair: bool, slide_index: int, stage: str) -> str:
    """Слова о ходе ревизии: у исправления находок, ручной правки и просьбы в чат они разные,
    а путь один — поэтому текст выбирается здесь, а не ветвлением по всему заданию."""
    plan = (
        "Исправляю выбранные находки"
        if is_repair
        else "Применяю правки редактора"
        if is_patch
        else f"Переделываю слайд {slide_index + 1}"
    )
    return {
        "plan": plan,
        "compose": "Собираю новую ревизию",
        "export": "Экспортирую PDF и миниатюры",
        "audit": "Проверяю изменённые слайды" if is_patch else "Проверяю изменённый слайд",
    }.get(stage, "")


def task_edit(edit_job_id: str) -> None:
    """Правка слайда: план базовой ревизии → модель → сборка, экспорт и аудит новой ревизии.
    Отказ модели завершает задание успешно без ревизии (`result.unchanged`)."""
    o = get_orchestrator()
    edit = o.state.get_repair(edit_job_id)
    job_id, variant_id = edit["job_id"], edit["variant_id"]
    slide_index = int(edit["slide_index"] or 0)
    # Патч приходит и от редактора, и от исправления находок: путь ревизии у них общий,
    # различаются только слова в карточке задания.
    patch = edit.get("patch")
    is_patch = patch is not None
    is_repair = edit.get("kind") == "repair"
    o.state.job_started(edit_job_id)
    o.state.update_job(
        edit_job_id,
        stage="plan",
        progress={
            "percent": 10,
            "message": _edit_message(is_patch, is_repair, slide_index, "plan"),
        },
    )
    try:
        variant = o.state.get_variant(job_id, variant_id)
        base = edit["base_revision"]
        if o.state.get_job(edit_job_id).get("depends_on"):
            # Правка ждала разбора готовой презентации: за это время копия могла получить r2,
            # а сборка — записать план в последнюю ревизию. База — она.
            latest = int(variant["revision"])
            if (
                variant["status"] == "failed"
                or not (
                    o.artifacts.revision_dir(job_id, variant_id, latest) / "plan.json"
                ).is_file()
            ):
                raise StageError("plan_unavailable", PARSE_FAILED, stage="plan")
            if latest != base:
                base = latest
                o.state.update_repair(edit_job_id, base_revision=base)
        if variant["revision"] != base:
            raise ConflictError("revision_stale", "Ревизия устарела до начала правки")
        base_dir = o.artifacts.revision_dir(job_id, variant_id, base)
        plan_path = base_dir / "plan.json"
        if not plan_path.is_file():
            raise StageError("plan_missing", "План базовой ревизии не найден", stage="plan")
        plan = upgrade_plan_schema(json.loads(plan_path.read_text(encoding="utf-8")))
        gen = o.state.get_generation(job_id)
        template = o.state.get_template(gen["template_id"])
        package = o.state.get_package(gen["package_id"])
        if template["profile"] is None or gen["story"] is None:
            raise StageError(
                "inputs_missing", "Профиль шаблона или смысловой план недоступны", stage="plan"
            )
        template_path = o.files.path_for(template["sha256"])
        new_rev = base + 1
        prefix = o.artifacts.prefix(variant_id, new_rev)
        progress = {"plan": 15, "compose": 45, "export": 65, "audit": 85}
        base_deck = None
        extra_assets: dict[str, pathlib.Path] = {}
        if is_patch:
            deck_path = base_dir / "composed.json"
            if deck_path.is_file():
                base_deck = json.loads(deck_path.read_text(encoding="utf-8"))
            extra_assets = _patch_files(o, patch or {})

        def emit(event: str, data: JsonDict) -> None:
            if event == "stage":
                o.state.add_stage(edit_job_id, {k: v for k, v in data.items() if k != "variant_id"})
                if data.get("status") == "running":
                    o.state.update_job(
                        edit_job_id,
                        stage=data["stage"],
                        progress={
                            "percent": progress.get(data["stage"], 10),
                            "message": _edit_message(
                                is_patch, is_repair, slide_index, str(data["stage"])
                            ),
                        },
                    )

        with o.artifacts.stage_revision(job_id, variant_id, new_rev) as staging:
            ctx = EditContext(
                job_id=job_id,
                variant_id=variant_id,
                revision=new_rev,
                base_revision=base,
                plan=plan,
                slide_index=slide_index,
                instruction=str(edit["instruction"] or ""),
                template_profile=template["profile"],
                template_path=template_path if template_path.is_file() else None,
                package=package["package"] or {},
                story=gen["story"],
                settings=gen["request"].get("settings") or {},
                staging=staging,
                package_dir=o.artifacts.package_dir(gen["package_id"]),
                nonce=edit_job_id,
                patch=patch,
                base_deck=base_deck,
                extra_assets=extra_assets,
                chart_report=_preview_chart_report(base_dir)
                if variant_id == ORIGINAL_VARIANT
                else None,
            )
            outcome = run_edit(o.layers, ctx, emit)
            if outcome.status == "failed":
                raise StageError(
                    str((outcome.error or {}).get("code") or "edit_failed"),
                    str((outcome.error or {}).get("message") or "Правка не выполнена"),
                    retryable=bool((outcome.error or {}).get("retryable", True)),
                    stage=str((outcome.error or {}).get("stage") or "plan"),
                )
            if outcome.status == "unchanged":
                # Ревизия не создаётся: незавершённый каталог staging удаляется без публикации.
                staging.discard = True
        if outcome.status == "unchanged":
            o.state.update_repair(
                edit_job_id,
                result="unchanged",
                slide_id=outcome.slide_id,
                change_note=outcome.reason,
            )
            o.state.update_job(
                edit_job_id,
                status="succeeded",
                stage="done",
                finished_at=now_iso(),
                result={
                    "unchanged": True,
                    "change_note": outcome.reason,
                    "generation_result_url": f"/api/generations/{job_id}",
                },
                progress={"percent": 100, "message": "Слайд оставлен без изменений"},
            )
            return
        manifest = o.artifacts.read_manifest(job_id, variant_id, new_rev)
        changed_ids = list(outcome.changed_slide_ids) or [outcome.slide_id]
        o.state.add_revision(
            job_id=job_id,
            variant_id=variant_id,
            revision=new_rev,
            artifacts_prefix=prefix,
            manifest=manifest,
            repair_job_id=edit_job_id,
            changed_slide_ids=changed_ids,
            pptx_hash=_pptx_hash(manifest, prefix),
        )
        summary = dict(outcome.audit or {})
        summary["report_artifact"] = f"{prefix}audit.json"
        status = (
            "needs_review"
            if summary.get("issues_total", 0) > 0 or not summary.get("coverage_complete")
            else "ready"
        )
        o.state.update_variant(
            job_id,
            variant_id,
            status=status,
            slide_count=outcome.slide_count,
            audited_at=now_iso(),
            audit=summary,
        )
        dropped = sum(
            len(s.get("overrides_dropped") or [])
            for s in ((outcome.report or {}).get("deck") or {}).get("slides") or []
        )
        o.state.update_repair(
            edit_job_id,
            result="applied",
            new_revision=new_rev,
            slide_id=outcome.slide_id,
            change_note=outcome.change_note,
            changed_slide_ids=changed_ids,
        )
        if is_patch:
            message = f"Правки применены, ревизия {new_rev} собрана"
            if dropped:
                message += f"; правок отброшено: {dropped}"
        else:
            message = f"Слайд {slide_index + 1} изменён, ревизия {new_rev} собрана"
        o.state.update_job(
            edit_job_id,
            status="succeeded",
            stage="done",
            finished_at=now_iso(),
            result={
                "revision": new_rev,
                "change_note": outcome.change_note,
                "generation_result_url": f"/api/generations/{job_id}",
                **({"changed_slide_ids": changed_ids} if is_patch else {}),
            },
            progress={"percent": 100, "message": message},
        )
    except Exception as e:
        log.exception("правка %s не удалась", edit_job_id)
        message = str(e) or e.__class__.__name__
        code = getattr(e, "code", "edit_failed")
        o.state.update_repair(edit_job_id, result="failed", message=message)
        o.state.update_job(
            edit_job_id,
            status="failed",
            stage="done",
            finished_at=now_iso(),
            error={
                "code": code,
                "message": message,
                "stage": getattr(e, "stage", None) or "plan",
                "retryable": bool(getattr(e, "retryable", True)),
            },
        )


def _patch_files(o: Orchestrator, patch: JsonDict) -> dict[str, pathlib.Path]:
    """Пути файлов проекта для правок с источником `file` (sha256 записан при постановке)."""
    out: dict[str, pathlib.Path] = {}
    for entry in patch.get("slides") or []:
        for override in entry.get("overrides") or []:
            for spec in (override.get("picture"), override.get("background")):
                source = (spec or {}).get("source") or {}
                if source.get("kind") == "file" and source.get("file_id") and source.get("sha256"):
                    path = o.files.path_for(str(source["sha256"]))
                    if path.is_file():
                        out[str(source["file_id"])] = path
    return out


def task_gc() -> None:
    """Периодическая сборка мусора на воркере анализа; сроки — из config/app.yaml."""
    o = get_orchestrator()
    try:
        o.collect_garbage()
    except Exception:
        log.exception("сборка мусора не удалась")


# ---------- вспомогательное ----------


def _ms_since(started_iso: str) -> int:
    started = datetime.fromisoformat(started_iso.replace("Z", "+00:00"))
    return int((datetime.now(UTC) - started).total_seconds() * 1000)


def _duration(job: JsonDict) -> int | None:
    if not job.get("started_at") or not job.get("finished_at"):
        return None
    a = datetime.fromisoformat(job["started_at"].replace("Z", "+00:00"))
    b = datetime.fromisoformat(job["finished_at"].replace("Z", "+00:00"))
    return int((b - a).total_seconds() * 1000)


def _pptx_hash(manifest: JsonDict, prefix: str) -> str | None:
    entry = manifest.get(f"{prefix}deck.pptx")
    return f"sha256:{entry['sha256']}" if entry and entry.get("sha256") else None


def _final_message(variants: list[JsonDict]) -> str:
    audited = sum(1 for v in variants if v.get("audited_at"))
    word = plural(len(variants), "вариант", "варианта", "вариантов")
    return f"Готово: {len(variants)} {word}, аудит завершён у {audited}"


def renderer_check(
    redis_url: str, worker_name: str, fixture: pathlib.Path | None = None, timeout_s: int = 90
) -> bool:
    """Пробный рендер собственной фикстуры при старте воркера; результат читает /api/health.

    Использует тот же слой экспорта: ONLYOFFICE под общим слотом рендера
    с временным подписанным источником и тайм-аутом, затем миниатюра первой страницы.
    """
    import tempfile

    from presentation_designer.export import pdf, thumbnails
    from presentation_designer.shared.valkey import connect

    ok = False
    fixture = (
        fixture
        or pathlib.Path(__file__).resolve().parents[3]
        / "tests"
        / "fixtures"
        / "pptx"
        / "mini_template.pptx"
    )
    if fixture.is_file():
        with tempfile.TemporaryDirectory(prefix="render-check-") as tmp:
            try:
                result = pdf.convert_to_pdf(fixture, pathlib.Path(tmp), timeout_s=timeout_s)
                thumbs = thumbnails.render_thumbnails(
                    result.pdf_path, pathlib.Path(tmp), width_px=320, pages=[1]
                )
                ok = bool(thumbs.paths)
                log.info(
                    "проверка рендерера: ONLYOFFICE за %.1f с",
                    result.seconds,
                )
            except Exception:
                log.exception("проверка рендерера не прошла")
    else:
        log.error("нет фикстуры для проверки рендерера: %s", fixture)
    try:
        # Срок жизни отметки — неделя: после падения воркера ключ не остаётся навсегда.
        connect(redis_url).set(
            f"pd:renderer:{worker_name}", "ok" if ok else "fail", ex=7 * 24 * 3600
        )
    except Exception:
        log.exception("не удалось записать результат проверки рендерера")
    return ok


def worker_name(role: str) -> str:
    return f"{role}-{socket.gethostname()}"
