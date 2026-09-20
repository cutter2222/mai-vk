"""Явные заглушки слоёв конвейера.

Возвращают документы, построенные из примеров contracts/examples, и настоящие
открывающиеся файлы: PPTX через python-pptx, PDF и миниатюры через Pillow, HTML текстом.
Каждый результат помечен `execution_mode=stub`; заглушечный результат не выдаётся за
генерацию. Этапы 4–10 заменяют эти функции реальными реализациями слоёв, не меняя
интерфейс ядра в pipeline/run.py.

Сценарии для сквозных тестов: имя шаблона со словом «fail» роняет вариант detailed на сборке.
"""

from __future__ import annotations

import copy
import html
import io
import json
import pathlib
import time
from typing import Any

from presentation_designer.parsing.content.brief import extract_brief as heuristic_brief
from presentation_designer.pipeline.run import (
    AnalyzeInput,
    AnalyzeOutput,
    AuditInput,
    BriefInput,
    ComposeInput,
    ComposeOutput,
    EditInput,
    EditOutput,
    ExportInput,
    ExportOutput,
    ImportInput,
    ImportOutput,
    Layers,
    PlanInput,
    RepairInput,
    RepairOutput,
    StageError,
    StoryInput,
)
from presentation_designer.pipeline.state import now_iso

ROOT = pathlib.Path(__file__).resolve().parents[3]
EXAMPLES = ROOT / "contracts" / "examples"

SLIDE_TITLES = [
    "Умные уведомления: пилот и план запуска",
    "Проблема",
    "Пользователи теряют почти половину важных уведомлений",
    "Решение",
    "Три механизма возвращают внимание пользователей",
    "Как работает приоритизация: четыре шага",
    "Открываемость выросла с 31 до 44 % за три месяца пилота",
    "Метрики пилота по месяцам",
    "Пилот окупается за год",
    "Просим одобрить расширение пилота",
    "Риски и как мы их снимаем",
    "План на четвёртый квартал",
    "Команда пилота",
    "Что нужно от руководителей",
    "Приложение: методика измерений",
]

VARIANT_SLIDES = {"compact": 10, "balanced": 12, "detailed": 15, "original": 10}
VARIANT_COLORS = {
    "compact": (0, 119, 255),
    "balanced": (255, 56, 133),
    "detailed": (82, 9, 119),
    "original": (60, 60, 60),
}
VARIANT_RATIONALE = {
    "compact": "Минимум текста, один факт на слайд, таблицы заменены графиками",
    "balanced": "Тезис и пояснение, графики с подписями, таблицы до 5 строк",
    "detailed": "Подробные буллеты в пределах порогов, таблицы и схемы, разделители секций",
    "original": "Исходная презентация как есть: слайды, тексты и оформление файла сохранены",
}

KIND_BY_FORMAT = {
    "docx": "docx",
    "xlsx": "xlsx",
    "csv": "csv",
    "pdf": "pdf",
    "markdown": "markdown",
    "text": "text",
    "image": "image",
    "pptx": "pptx",
}

_FONT_CANDIDATES = [
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    "/System/Library/Fonts/Supplemental/Arial.ttf",
    "/Library/Fonts/Arial.ttf",
    "C:/Windows/Fonts/arial.ttf",
]


def _example(name: str) -> dict[str, Any]:
    return json.loads((EXAMPLES / f"{name}.example.json").read_text(encoding="utf-8"))  # type: ignore[no-any-return]


def _font(size: int) -> Any:
    from PIL import ImageFont

    for candidate in _FONT_CANDIDATES:
        if pathlib.Path(candidate).is_file():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default(size)


def _pause(delay_ms: int) -> None:
    if delay_ms > 0:
        time.sleep(delay_ms / 1000)


def slide_image(title: str, subtitle: str, color: tuple[int, int, int], width: int = 1280) -> bytes:
    """PNG слайда: цветная полоса, заголовок и подзаголовок."""
    from PIL import Image, ImageDraw

    height = round(width * 9 / 16)
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    draw.rectangle([0, 0, width, round(height * 0.02)], fill=color)
    draw.text(
        (round(width * 0.06), round(height * 0.16)),
        title,
        fill=(20, 20, 20),
        font=_font(round(width * 0.045)),
    )
    draw.text(
        (round(width * 0.06), round(height * 0.34)),
        subtitle,
        fill=(90, 90, 90),
        font=_font(round(width * 0.025)),
    )
    draw.rectangle(
        [round(width * 0.06), round(height * 0.86), round(width * 0.12), round(height * 0.88)],
        fill=color,
    )
    buf = io.BytesIO()
    img.save(buf, format="PNG", optimize=True)
    return buf.getvalue()


def build_pptx(title: str, slide_titles: list[str], subtitle: str) -> bytes:
    """Настоящая презентация: титульный слайд и слайды с заголовком и текстом."""
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    first = prs.slides.add_slide(prs.slide_layouts[0])
    first.shapes.title.text = title
    if len(first.placeholders) > 1:
        first.placeholders[1].text = subtitle
    for text in slide_titles[1:]:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = text
        box = slide.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.5), Inches(3.5))
        frame = box.text_frame
        frame.word_wrap = True
        frame.text = subtitle
        for paragraph in frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(20)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def build_pdf(pages: list[bytes]) -> bytes:
    """PDF из изображений слайдов: по странице на слайд."""
    from PIL import Image

    images = [Image.open(io.BytesIO(png)).convert("RGB") for png in pages]
    buf = io.BytesIO()
    images[0].save(buf, format="PDF", save_all=True, append_images=images[1:])
    return buf.getvalue()


def build_html(title: str, slide_titles: list[str], subtitle: str) -> str:
    sections = "\n".join(
        f'<section class="slide"><h2>{html.escape(t)}</h2><p>{html.escape(subtitle)}</p><footer>{i + 1} / {len(slide_titles)}</footer></section>'  # noqa: E501
        for i, t in enumerate(slide_titles)
    )
    return (
        '<!doctype html><html lang="ru"><head><meta charset="utf-8">'
        f"<title>{html.escape(title)}</title>"
        "<style>body{margin:0;font-family:system-ui,sans-serif;background:#f3f4f6}"
        ".slide{aspect-ratio:16/9;max-width:960px;margin:24px auto;background:#fff;padding:48px;box-sizing:border-box;position:relative}"  # noqa: E501
        "h2{margin:0 0 16px;font-size:28px}p{color:#555}footer{position:absolute;bottom:16px;right:24px;color:#999;font-size:12px}"  # noqa: E501
        f"</style></head><body>{sections}</body></html>"
    )


def _logo_png() -> bytes:
    from PIL import Image

    img = Image.new("RGBA", (64, 64), (0, 119, 255, 255))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def stub_deck(plan: dict[str, Any], titles: list[str], *, prefix: str = "") -> dict[str, Any]:
    """ComposedDeck заглушки по числу слайдов плана: на каждом слайде заголовок, текст и
    логотип с теми же идентификаторами, что в примере контракта, эхо ручных правок плана.
    Нужен ревизии-патчу: адреса объектов проверяются по описанию базовой ревизии."""
    sample = _example("composed_deck")
    slides_raw = plan.get("slides") or []
    explicit = bool(slides_raw) and not isinstance(plan.get("slide_count"), int)
    ordered = sorted(slides_raw, key=lambda s: int(s.get("order", 0))) if explicit else []
    slides: list[dict[str, Any]] = []
    for index, title in enumerate(titles):
        plan_slide = ordered[index] if index < len(ordered) else {}
        slide_id = str(plan_slide.get("slide_id") or f"s{index + 1}")
        objects: list[dict[str, Any]] = [
            {
                "object_id": "2",
                "name": "Title 1",
                "kind": "text",
                "geometry": "rect",
                "bbox": {"x": 0.05, "y": 0.06, "width": 0.9, "height": 0.12},
                "z_order": 1,
                "slot_id": "title",
                "slot_kind": "title",
                "block_kind": "title",
                "source_object_id": "2",
                "role": "content",
                "content_source": "plan",
                "text": {
                    "plain": title,
                    "paragraphs": [{"text": title, "level": 0, "bullet": False, "align": "left"}],
                    "computed_style": {
                        "font": {"family": "Play", "size_pt": 32, "bold": True, "color": "#000000"}
                    },
                    "autofit": "none",
                },
            },
            {
                "object_id": "3",
                "name": "Body 2",
                "kind": "text",
                "geometry": "rect",
                "bbox": {"x": 0.05, "y": 0.3, "width": 0.6, "height": 0.4},
                "z_order": 2,
                "slot_id": "body",
                "slot_kind": "body",
                "block_kind": "body",
                "source_object_id": "3",
                "role": "content",
                "content_source": "plan",
                "text": {
                    "plain": f"Вариант {plan.get('variant', {}).get('variant_id', '')}",
                    "paragraphs": [
                        {
                            "text": f"Вариант {plan.get('variant', {}).get('variant_id', '')}",
                            "level": 0,
                            "bullet": False,
                            "align": "left",
                        }
                    ],
                    "computed_style": {
                        "font": {
                            "family": "Arial",
                            "size_pt": 16,
                            "bold": False,
                            "color": "#202020",
                        }
                    },
                    "autofit": "none",
                },
            },
            {
                "object_id": "5",
                "name": "Logo",
                "kind": "picture",
                "bbox": {"x": 0.87, "y": 0.9, "width": 0.08, "height": 0.06},
                "z_order": 0,
                "role": "fixed",
                "source_object_id": "5",
                "content_source": "template",
                "picture": {
                    "asset_id": "asset_logo",
                    "natural_width_px": 64,
                    "natural_height_px": 64,
                    "origin": "template",
                    "fit": "as_is",
                },
            },
        ]
        # Порядок правок в плане уже упорядочен редактором: своя надпись появляется до правок
        # на неё, удаление — последним.
        for override in plan_slide.get("overrides") or []:
            op = str(override.get("op") or "")
            target = override.get("target") or {}
            object_id = str(target.get("object_id") or "")
            if op == "add_text":
                text = str(override.get("text") or "")
                objects.append(
                    {
                        "object_id": object_id,
                        "name": "Своя надпись",
                        "kind": "text",
                        "bbox": dict((override.get("geometry") or {}).get("bbox") or {}),
                        "z_order": len(objects) + 1,
                        "role": "content",
                        "content_source": "user",
                        "text": {
                            "plain": text,
                            "paragraphs": [{"text": line} for line in text.split("\n")],
                        },
                        "user_overrides": [copy.deepcopy(override)],
                    }
                )
                continue
            if op == "delete":
                objects = [o for o in objects if o["object_id"] != object_id]
                continue
            for obj in objects:
                if obj["object_id"] == object_id:
                    obj.setdefault("user_overrides", []).append(copy.deepcopy(override))
                    if op == "text" and override.get("text") is not None:
                        text = str(override["text"])
                        obj["text"]["plain"] = text
                        obj["text"]["paragraphs"][0]["text"] = text
                        obj["content_source"] = "user"
                    if op == "geometry":
                        obj["bbox"] = dict(
                            (override.get("geometry") or {}).get("bbox") or obj["bbox"]
                        )
        slide: dict[str, Any] = {
            "slide_id": slide_id,
            "index": index,
            "pptx_slide_part": f"ppt/slides/slide{index + 1}.xml",
            "layout_id": "slideLayout6",
            "pattern_id": str(plan_slide.get("pattern_id") or "pat_title"),
            "background": {"kind": "inherited"},
            "objects": objects,
            "title": title,
            "removed_object_ids": [],
        }
        if plan_slide.get("overrides"):
            slide["overrides"] = copy.deepcopy(plan_slide["overrides"])
            slide["overrides_dropped"] = []
            bg = next((o for o in plan_slide["overrides"] if o.get("op") == "background"), None)
            if bg:
                spec = bg.get("background") or {}
                slide["background"] = {
                    k: v
                    for k, v in {"kind": spec.get("kind"), "color": spec.get("color")}.items()
                    if v
                }
        slides.append(slide)
    deck = {k: v for k, v in sample.items() if k not in ("slides", "assets", "stats")}
    deck["slides"] = slides
    deck["assets"] = [
        {
            "asset_id": "asset_logo",
            "media_path": "ppt/media/image1.png",
            "sha256": "0" * 64,
            "content_type": "image/png",
            "origin": "template",
            "shared_with_template": True,
            "artifact": f"{prefix}media/asset_logo.png",
        }
    ]
    deck["stats"] = {
        "slides": len(slides),
        "objects": 3 * len(slides),
        "text_objects": 2 * len(slides),
        "pictures": len(slides),
        "tables": 0,
        "charts": 0,
        "diagrams": 0,
        "removed_objects": 0,
    }
    deck["warnings"] = []
    return deck


def explicit_plan(plan: dict[str, Any]) -> dict[str, Any]:
    """План заглушки с явным списком слайдов: заглушечный план примера несёт лишь число
    слайдов и примерные заголовки, а правкам нужны slide_id и order."""
    plan = copy.deepcopy(plan)
    count, titles = plan_slides(plan)
    if isinstance(plan.get("slide_count"), int) or len(plan.get("slides") or []) != count:
        sample = (plan.get("slides") or [{}])[0]
        plan["slides"] = [
            {
                "slide_id": f"s{i + 1}",
                "order": i + 1,
                "role": "title" if i == 0 else "bullets",
                "pattern_id": sample.get("pattern_id", "pat_title"),
                "title": title,
                "key_message": title,
                "blocks": [{"slot_id": "title", "kind": "title", "text": title}],
                "thesis_refs": [],
            }
            for i, title in enumerate(titles)
        ]
        plan["slide_count"] = {"target": count}
    return plan


def plan_slides(plan: dict[str, Any]) -> tuple[int, list[str]]:
    """Число слайдов и заголовки по плану: настоящий план несёт свои слайды, заглушечный —
    число слайдов и примерные заголовки."""
    slides = plan.get("slides") or []
    if slides and not isinstance(plan.get("slide_count"), int):
        ordered = sorted(slides, key=lambda s: int(s.get("order", 0)))
        return len(ordered), [str(s.get("title") or "Слайд") for s in ordered]
    count = int(plan.get("slide_count") or len(slides) or 10)
    return count, (SLIDE_TITLES * 4)[:count]


class StubLayers(Layers):
    """Набор заглушек со временем этапа из настроек."""

    def __init__(self, stage_delay_ms: int = 0) -> None:
        self.delay_ms = stage_delay_ms
        self.modes = {
            "parsing.template": "stub",
            "parsing.content": "stub",
            "brief": "stub",
            "generation.story": "stub",
            "generation.plan": "stub",
            "generation.edit": "stub",
            "layout": "stub",
            "export": "stub",
            "audit.deterministic": "stub",
            "audit.contextual": "stub",
        }

    # ---------- анализ шаблона ----------

    def analyze(self, inp: AnalyzeInput) -> AnalyzeOutput:
        _pause(self.delay_ms)
        profile = _example("template_profile")
        profile["template_id"] = inp.template_id
        profile["template_hash"] = f"sha256:{inp.sha256}"
        profile["source_file"] = {"name": inp.name, "size_bytes": inp.size_bytes, "format": "pptx"}
        profile["created_at"] = now_iso()
        previews: dict[str, bytes] = {}
        for pattern in profile["patterns"]:
            path = pattern.get("preview_path")
            if path and path not in previews:
                previews[path] = slide_image(
                    pattern.get("name", "Образец"),
                    f"Образец шаблона: {pattern.get('role', '')}",
                    (0, 119, 255),
                    640,
                )
        for sample in profile.get("sample_slides", []):
            path = sample.get("preview_path")
            if path and path not in previews:
                previews[path] = slide_image(
                    "Образец", "Образцовый слайд шаблона", (0, 119, 255), 640
                )
        # Пустой слайд каждого макета — подложка холста редактора (этап 23).
        for layout in profile.get("layouts", []):
            layout_id = str(layout.get("layout_id") or "")
            if layout_id:
                previews[f"previews/layout-{layout_id}.png"] = slide_image(
                    "", str(layout.get("name") or layout_id), (200, 205, 215), 640
                )
        return AnalyzeOutput(profile=profile, previews=previews)

    # ---------- импорт содержания ----------

    def import_content(self, inp: ImportInput) -> ImportOutput:
        _pause(self.delay_ms)
        package = _example("content_package")
        package["package_id"] = inp.package_id
        package["created_at"] = now_iso()
        package["mode"] = inp.mode
        warnings: list[dict[str, Any]] = []
        brief = dict(package["brief"])
        supplied = {k: v for k, v in (inp.brief or {}).items() if v not in ("", None, [])}
        brief.update(supplied)
        missing = [k for k in ("purpose", "title") if k not in supplied]
        if missing:
            brief.setdefault("purpose", "other")
            brief.setdefault("title", inp.sources[0].name if inp.sources else "Презентация")
            warnings.append(
                {
                    "code": "brief_incomplete",
                    "message": f"В брифе не заданы поля: {', '.join(missing)}; применены умолчания",
                }
            )
        brief.setdefault("language", "ru")
        if not brief.get("purpose"):
            brief["purpose"] = "other"
        if not brief.get("title"):
            brief["title"] = "Презентация"
        package["brief"] = brief
        sources: list[dict[str, Any]] = (
            [{"source_id": "src_brief", "kind": "user_input", "name": "бриф", "extracted": True}]
            if inp.brief
            else []
        )
        for i, src in enumerate(inp.sources, start=1):
            sources.append(
                {
                    "source_id": f"src_{i}",
                    "kind": KIND_BY_FORMAT.get(src.format, "text"),
                    "name": src.name,
                    "sha256": src.sha256,
                    "size_bytes": src.size_bytes,
                    "extracted": True,
                    "file_id": src.file_id,
                }
            )
        if sources:
            package["sources"] = sources
        package["warnings"] = warnings
        return ImportOutput(package=package, warnings=warnings)

    # ---------- смысловой план ----------

    def story_key(self, inp: StoryInput) -> str:
        return f"stub:{inp.package['package_id']}"

    async def extract_brief(self, inp: BriefInput) -> dict[str, Any]:
        return heuristic_brief(inp.text, inp.brief)

    def story(self, inp: StoryInput) -> dict[str, Any]:
        _pause(self.delay_ms)
        story = _example("story_plan")
        story["package_id"] = inp.package["package_id"]
        story["story_id"] = f"story_{inp.package['package_id']}"
        story["content_hash"] = self.story_key(inp)
        brief = inp.package.get("brief", {})
        story["purpose"] = brief.get("purpose", story["purpose"])
        story["audience"] = brief.get("audience") or story["audience"]
        story["goal"] = brief.get("goal") or story["goal"]
        story["language"] = brief.get("language", "ru")
        story["generation_meta"]["created_at"] = now_iso()
        return story

    # ---------- план варианта ----------

    def plan(self, inp: PlanInput) -> dict[str, Any]:
        _pause(self.delay_ms)
        plan = _example("slide_plan")
        plan["plan_id"] = f"plan_{inp.job_id}_{inp.variant_id}"
        plan["template_id"] = inp.template_profile["template_id"]
        plan["package_id"] = inp.package["package_id"]
        plan["story_id"] = inp.story["story_id"]
        plan["variant"] = {
            **plan["variant"],
            "variant_id": inp.variant_id,
            "rationale": VARIANT_RATIONALE[inp.variant_id],
        }
        count = inp.slide_count or VARIANT_SLIDES[inp.variant_id]
        plan["slide_count"] = count
        plan["generation_meta"]["created_at"] = now_iso()
        return plan

    # ---------- сборка ----------

    def compose(self, inp: ComposeInput) -> ComposeOutput:
        _pause(self.delay_ms)
        if (
            "fail" in inp.template_profile.get("source_file", {}).get("name", "").lower()
            and inp.variant_id == "detailed"
        ):
            raise StageError(
                "compose_failed",
                "Не удалось клонировать образец slide29: битая ссылка на медиа",
                retryable=True,
            )
        count, titles = plan_slides(inp.plan)
        subtitle = f"Вариант {inp.variant_id}, ревизия {inp.revision}"
        title = (
            inp.story.get("theses", [{}])[0].get("statement", "Презентация")
            if inp.story
            else "Презентация"
        )
        deck_title = inp.package.get("brief", {}).get("title") or title
        pptx = build_pptx(deck_title, titles, subtitle)
        inp.staging.write_bytes("deck.pptx", pptx)
        deck = stub_deck(inp.plan, titles, prefix=inp.staging.prefix)
        deck["deck_id"] = f"deck_{inp.job_id}_{inp.variant_id}_r{inp.revision}"
        deck["job_id"] = inp.job_id
        deck["variant_id"] = inp.variant_id
        deck["revision"] = inp.revision
        deck["plan_id"] = inp.plan["plan_id"]
        deck["template_id"] = inp.template_profile["template_id"]
        deck["pptx_artifact"] = f"{inp.staging.prefix}deck.pptx"
        inp.staging.write_bytes("media/asset_logo.png", _logo_png())
        inp.staging.write_json("composed.json", deck)
        inp.staging.write_json("plan.json", inp.plan)
        return ComposeOutput(slide_count=count, slide_titles=titles, composed_deck=deck)

    # ---------- экспорт ----------

    def export(self, inp: ExportInput) -> ExportOutput:
        _pause(self.delay_ms)
        color = VARIANT_COLORS.get(inp.variant_id, (0, 119, 255))
        subtitle = f"Вариант {inp.variant_id}, ревизия {inp.revision}"
        pages: list[bytes] = []
        thumbnails: list[dict[str, Any]] = []
        for index, title in enumerate(inp.slide_titles):
            marks = inp.slide_marks.get(index, "")
            png = slide_image(title, marks or subtitle, color)
            pages.append(png)
            name = f"thumbs/slide-{index + 1:02d}.png"
            inp.staging.write_bytes(name, png)
            thumbnails.append(
                {
                    "slide_index": index,
                    "name": f"{inp.staging.prefix}{name}",
                    "width_px": 1280,
                    "height_px": 720,
                }
            )
        inp.staging.write_bytes("deck.pdf", build_pdf(pages))
        inp.staging.write_bytes(
            "deck.html", build_html(inp.deck_title, inp.slide_titles, subtitle).encode("utf-8")
        )
        return ExportOutput(thumbnails=thumbnails)

    # ---------- аудит ----------

    def audit(self, inp: AuditInput) -> dict[str, Any]:
        _pause(self.delay_ms)
        report = _example("audit_report")
        report["job_id"] = inp.job_id
        report["variant_id"] = inp.variant_id
        report["revision"] = inp.revision
        report["report_id"] = f"audit_{inp.job_id}_{inp.variant_id}_r{inp.revision}"
        report["created_at"] = now_iso()
        report["deck"]["slide_count"] = inp.slide_count
        report["deck"]["pptx_artifact"] = f"{inp.staging.prefix}deck.pptx"
        report["deck"]["composed_deck_artifact"] = f"{inp.staging.prefix}composed.json"
        if not inp.contextual:
            report["coverage"] = {
                "complete": False,
                "checked": 5,
                "not_checked": 2,
                "not_applicable": 1,
                "missing_inputs": ["contextual_disabled"],
            }
            report["contextual_answers"] = []
        if inp.variant_id == "compact":
            report["issues"] = []
            report["results"] = [
                {**r, "issue_ids": []} for r in report["results"] if r.get("outcome") != "failed"
            ]
            report["coverage"] = {
                "complete": inp.contextual,
                "checked": 6,
                "not_checked": 0 if inp.contextual else 1,
                "not_applicable": 1,
                "missing_inputs": [] if inp.contextual else ["contextual_disabled"],
            }
            report["summary"] = {
                "issues_total": 0,
                "by_severity": {"blocking": 0, "error": 0, "warning": 0, "info": 0},
                "by_category": {},
                "by_kind": {},
                "slides_with_issues": 0,
                "score": 100,
            }
            report["contextual_answers"] = []
        inp.staging.write_json("audit.json", report)
        return report

    # ---------- исправление ----------

    def repair(self, inp: RepairInput) -> RepairOutput:
        _pause(self.delay_ms)
        fixed = set(inp.issue_ids)
        prev = inp.base_report
        remaining = [
            dict(i, revision=inp.revision) for i in prev["issues"] if i["issue_id"] not in fixed
        ]
        changed = sorted(
            {f"s{i['slide_index'] + 1}" for i in prev["issues"] if i["issue_id"] in fixed}
        )
        changed_indexes = {i["slide_index"] for i in prev["issues"] if i["issue_id"] in fixed}
        report = copy.deepcopy(prev)
        report["report_id"] = f"audit_{inp.job_id}_{inp.variant_id}_r{inp.revision}"
        report["revision"] = inp.revision
        report["created_at"] = now_iso()
        report["issues"] = remaining
        report["results"] = [
            (
                {**r, "outcome": "passed", "issue_ids": []}
                if any(i in fixed for i in r.get("issue_ids", []))
                else r
            )
            for r in prev["results"]
        ]
        report["rechecked_after_repair"] = {
            "base_revision": inp.base_revision,
            "changed_slide_ids": changed,
            "dependent_slide_ids": [],
            "deck_checks_rerun": ["integrity.duplicate_slides"],
        }
        report["summary"] = {
            **prev["summary"],
            "issues_total": len(remaining),
            "by_severity": {
                "blocking": sum(1 for i in remaining if i["severity"] == "blocking"),
                "error": sum(1 for i in remaining if i["severity"] == "error"),
                "warning": sum(1 for i in remaining if i["severity"] == "warning"),
                "info": sum(1 for i in remaining if i["severity"] == "info"),
            },
            "slides_with_issues": len({i["slide_index"] for i in remaining}),
            "score": min(100, int(prev["summary"].get("score", 88)) + 6),
        }
        report["deck"]["pptx_artifact"] = f"{inp.staging.prefix}deck.pptx"
        report["deck"]["composed_deck_artifact"] = f"{inp.staging.prefix}composed.json"
        # Файлы новой ревизии: копия предыдущей с пометкой исправленных слайдов на миниатюрах.
        for name in ("deck.pptx", "composed.json", "plan.json"):
            src = inp.base_dir / name
            if src.is_file():
                inp.staging.write_bytes(name, src.read_bytes())
        color = VARIANT_COLORS.get(inp.variant_id, (0, 119, 255))
        subtitle = f"Вариант {inp.variant_id}, ревизия {inp.revision}"
        pages: list[bytes] = []
        thumbnails: list[dict[str, Any]] = []
        for index, title in enumerate(inp.slide_titles):
            mark = (
                "Слайд исправлен в новой ревизии"
                if index in changed_indexes
                else (
                    "На этом слайде есть находки аудита"
                    if any(i["slide_index"] == index for i in remaining)
                    else subtitle
                )
            )
            png = slide_image(title, mark, color)
            pages.append(png)
            name = f"thumbs/slide-{index + 1:02d}.png"
            inp.staging.write_bytes(name, png)
            thumbnails.append(
                {
                    "slide_index": index,
                    "name": f"{inp.staging.prefix}{name}",
                    "width_px": 1280,
                    "height_px": 720,
                }
            )
        inp.staging.write_bytes("deck.pdf", build_pdf(pages))
        inp.staging.write_bytes(
            "deck.html", build_html(inp.deck_title, inp.slide_titles, subtitle).encode("utf-8")
        )
        inp.staging.write_json("audit.json", report)
        return RepairOutput(report=report, changed_slide_ids=changed, thumbnails=thumbnails)

    # ---------- правка слайда по запросу ----------

    def edit(self, inp: EditInput) -> EditOutput:
        """Детерминированная правка: заголовок слайда — из инструкции. Инструкция со словом
        «невозможно» отклоняется, как отказ модели по нехватке данных. План заглушки без
        явного списка слайдов сначала получает его по своим примерным заголовкам, чтобы
        заголовок новой ревизии был виден на миниатюре."""
        _pause(self.delay_ms)
        if inp.patch is not None:
            return self._apply_patch(inp)
        plan = copy.deepcopy(inp.plan)
        count, _titles = plan_slides(plan)
        if inp.slide_index < 0 or inp.slide_index >= count:
            raise StageError(
                "slide_index_out_of_range",
                f"В плане {count} слайдов, слайда с номером {inp.slide_index + 1} нет",
                stage="plan",
            )
        plan = explicit_plan(plan)
        slide = sorted(plan["slides"], key=lambda s: int(s.get("order", 0)))[inp.slide_index]
        slide_id = str(slide["slide_id"])
        instruction = " ".join(inp.instruction.split())
        if "невозможно" in instruction.lower():
            return EditOutput(
                plan=None,
                changed=False,
                slide_id=slide_id,
                reason="В материалах нет данных для такой правки; выдумывать не буду",
                report={"stub": True},
            )
        title = instruction[:80].rstrip(" .")
        slide["title"] = title
        slide["key_message"] = title
        for block in slide.get("blocks") or []:
            if block.get("kind") == "title":
                block["text"] = title
        slide["revision_note"] = "Заголовок заменён по просьбе"
        slide.pop("overrides", None)
        return EditOutput(
            plan=plan,
            changed=True,
            slide_id=slide_id,
            change_note="Заголовок заменён по просьбе",
            report={"stub": True},
        )

    def _apply_patch(self, inp: EditInput) -> EditOutput:
        """Ручные правки редактора: детерминированно, без модели, одинаково в заглушках и в
        настоящих слоях (`generation/patch.py`)."""
        from presentation_designer.generation.patch import PatchError, apply_patch

        plan = explicit_plan(inp.plan)
        try:
            result = apply_patch(
                plan,
                inp.patch or {},
                inp.base_deck,
                inp.template_profile,
                inp.package,
                inp.story,
            )
        except PatchError as e:
            raise StageError(e.code, str(e), stage="plan") from e
        return EditOutput(
            plan=result.plan,
            changed=True,
            slide_id=result.changed_slide_ids[0] if result.changed_slide_ids else "",
            change_note=result.summary,
            report={**result.report, "warnings": result.warnings},
            changed_slide_ids=list(result.changed_slide_ids),
            summary=result.summary,
        )
