"""Отделка слайда собственной композиции после заполнения: вид, который делает слайд
«нарисованным», а не собранным из рамок.

Композиция размечает места в долях холста ещё до того, как известен текст, поэтому без
отделки карточки растянуты на всю высоту и пустуют, а текст в них мелкий. Здесь, когда
текст уже стоит, слайд доводится по правилам дизайна:

* кегль — по роли (заголовок карточки, текст, число) и растёт, пока есть место; высота
  карточки — по тексту, а группа карточек ставится по вертикали в рабочей области;
* у каждой карточки значок: иконка по смыслу заголовка (`library.iconset`) или номер, если
  иконки нашлись не для всех — разнобой читается хуже, чем ровный ряд номеров;
* пункты списка получают акцентный маркер и жирное начало до двоеточия, рядом — панель с
  иконкой темы, если список занимает не всю ширину;
* число показателя — крупно в акценте, с короткой чертой над ним;
* под заголовком — короткая акцентная черта, если у шаблона нет своего декора шапки.

Всё видимое берётся из дизайн-кода шаблона: палитра, шрифты, скругление, тень. Вид
(`Look`) у трёх вариантов свой: «панели», «акцент» и «линии», — чтобы варианты отличались
подачей, а не только количеством текста.

Объекты слотов не пересоздаются: текстовые рамки только двигаются и переоформляются, их id
и текст остаются, поэтому аудит, правка из чата и редактор видят те же объекты. Новые
фигуры (значки, иконки, черты) — оформление без текста.
"""

from __future__ import annotations

import dataclasses
from dataclasses import dataclass
from typing import Any

from lxml import etree
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

from presentation_designer.layout.ooxml import NS_A, NS_P
from presentation_designer.library import iconset
from presentation_designer.library.spec import Composition, CompositionSlot
from presentation_designer.library.tokens import DesignCode
from presentation_designer.parsing.template.tone import relative_luminance

JsonDict = dict[str, Any]
EMU_PT = 12700

# Семейства, которые отделываются как ряд карточек: значок, заголовок, текст.
CARD_FAMILIES = ("cards_grid", "icon_cards", "matrix", "comparison", "process")
# Доля свободной высоты над группой карточек: чуть выше середины — так группа держится
# заголовка и не выглядит упавшей.
TOP_SHARE = 0.42


@dataclass(frozen=True)
class Look:
    """Подача варианта: вид карточек, значков и маркеров."""

    name: str
    card: str  # panel — мягкая подложка; accent — первая карточка залита; line — без плашек
    badge: str  # soft — бледный круг и иконка акцентом; solid — круг акцента и белая иконка
    title_mark: bool = True
    colorful: bool = False  # карточки ряда в разных акцентах палитры


LOOKS: dict[str, Look] = {
    "balanced": Look("panels", card="panel", badge="soft"),
    "compact": Look("accent", card="accent", badge="solid"),
    "detailed": Look("lines", card="line", badge="soft"),
}


def look_for(variant_id: str | None, code: DesignCode) -> Look:
    look = LOOKS.get(str(variant_id or ""), LOOKS["balanced"])
    if len(set(code.accents)) >= 3 and code.accents_from_slides:
        # Шаблон сам красит ряды разными цветами (01 красный, 02 оранжевый…): так же и здесь.
        look = dataclasses.replace(look, colorful=True)
    return look


@dataclass
class Frame:
    """Геометрия слайда и доступ к объектам по id."""

    slide: Any
    width: int
    height: int
    code: DesignCode
    look: Look
    backdrop: str
    shapes: dict[str, Any]

    @property
    def dark(self) -> bool:
        return relative_luminance(self.backdrop) < 0.35

    def get(self, shape_id: str | None) -> Any | None:
        return self.shapes.get(str(shape_id)) if shape_id else None

    def text_of(self, shape_id: str | None) -> str:
        shape = self.get(shape_id)
        if shape is None or not getattr(shape, "has_text_frame", False):
            return ""
        return "\n".join(p.text for p in shape.text_frame.paragraphs).replace("\v", "\n").strip()


def dress_slide(
    slide: Any,
    composition: Composition,
    refs: dict[str, str],
    card_ids: list[str],
    code: DesignCode,
    look: Look,
    *,
    width: int,
    height: int,
    backdrop: str,
    blocks: list[JsonDict] | None = None,
    title_decor: bool = False,
) -> list[str]:
    """Отделка по семейству композиции; возвращает, что сделано (для отчёта сборки)."""
    frame = Frame(
        slide=slide,
        width=width,
        height=height,
        code=code,
        look=look,
        backdrop=backdrop,
        shapes={str(s.shape_id): s for s in slide.shapes},
    )
    queries = {
        str(b.get("slot_id")): str((b.get("icon") or {}).get("query") or "")
        for b in blocks or []
        if b.get("kind") == "icon"
    }
    done: list[str] = []
    family = composition.family
    if family in CARD_FAMILIES:
        if _dress_cards(frame, composition, refs, card_ids, queries):
            done.append("cards")
    elif family == "bullets_pane":
        if _dress_bullets(frame, composition, refs):
            done.append("bullets")
    elif family in ("kpi_row", "hero_number", "chart_kpi"):
        if _dress_numbers(frame, composition, refs, card_ids):
            done.append("numbers")
    elif family in ("statement", "quote"):
        if _dress_statement(frame, composition, refs):
            done.append("statement")
    service = family in ("title_slide", "section", "closing", "agenda")
    if (
        look.title_mark
        and not title_decor
        and not service
        and _title_mark(frame, refs.get("title"))
    ):
        done.append("title_mark")
    return done


# ---------- измерение текста ----------


def _font(family: str | None, bold: bool = False) -> Any:
    from presentation_designer.shared import text_metrics

    return text_metrics.resolve_font(family, bold=bold)


def _line_pt(family: str | None, size: float, bold: bool = False, spacing: float = 1.0) -> float:
    from presentation_designer.shared import text_metrics

    return float(text_metrics.line_metrics(_font(family, bold), size).line_height_pt) * spacing


def text_height(
    text: str,
    family: str | None,
    size: float,
    width_emu: int,
    *,
    bold: bool = False,
    spacing: float = 1.0,
    para_gap_pt: float = 0.0,
) -> int:
    """Высота текста в EMU при переносе по словам в рамке шириной `width_emu` без полей."""
    from presentation_designer.generation.capacity import wrap_lines

    if not text.strip():
        return 0
    font = _font(family, bold)
    # Запас 4 %: рендерер переносит чуть раньше, чем жадный подсчёт по ширинам.
    width_pt = width_emu / EMU_PT * 0.96
    paragraphs = [p for p in text.split("\n") if p.strip()] or [text]
    lines = sum(wrap_lines(p, width_pt, font, size) for p in paragraphs)
    line = _line_pt(family, size, bold, spacing)
    gaps = para_gap_pt * max(0, len(paragraphs) - 1)
    return int((lines * line + gaps) * EMU_PT)


# ---------- низкоуровневые приёмы оформления ----------


def _hex(color: str) -> str:
    return color.lstrip("#").upper()


def _mix(fore: str, back: str, share: float) -> str:
    f = [int(_hex(fore)[i : i + 2], 16) for i in (0, 2, 4)]
    b = [int(_hex(back)[i : i + 2], 16) for i in (0, 2, 4)]
    return "#{:02X}{:02X}{:02X}".format(
        *(round(x * share + y * (1 - share)) for x, y in zip(f, b, strict=True))
    )


def _strip_style(shape: Any) -> None:
    style = shape._element.find(f"{{{NS_P}}}style")
    if style is not None:
        shape._element.remove(style)


def _set_fill(shape: Any, color: str | None, alpha: float = 1.0) -> None:
    """Заливка цветом с прозрачностью (alpha 0–1) или без заливки (None)."""
    sp_pr = shape._element.spPr
    for tag in ("solidFill", "noFill", "gradFill", "pattFill", "blipFill"):
        for old in sp_pr.findall(f"{{{NS_A}}}{tag}"):
            sp_pr.remove(old)
    anchor = sp_pr.find(f"{{{NS_A}}}ln")
    if color is None:
        element = etree.Element(f"{{{NS_A}}}noFill")
    else:
        element = etree.Element(f"{{{NS_A}}}solidFill")
        clr = etree.SubElement(element, f"{{{NS_A}}}srgbClr", val=_hex(color))
        if alpha < 1.0:
            etree.SubElement(clr, f"{{{NS_A}}}alpha", val=str(round(alpha * 100000)))
    geom = sp_pr.find(f"{{{NS_A}}}prstGeom")
    if geom is None:
        geom = sp_pr.find(f"{{{NS_A}}}custGeom")
    if anchor is not None:
        anchor.addprevious(element)
    elif geom is not None:
        geom.addnext(element)
    else:
        sp_pr.append(element)


def _set_line(shape: Any, color: str | None, width_pt: float = 0.0, alpha: float = 1.0) -> None:
    sp_pr = shape._element.spPr
    for old in sp_pr.findall(f"{{{NS_A}}}ln"):
        sp_pr.remove(old)
    line = etree.Element(f"{{{NS_A}}}ln")
    if color is None or width_pt <= 0:
        etree.SubElement(line, f"{{{NS_A}}}noFill")
    else:
        line.set("w", str(round(width_pt * EMU_PT)))
        fill = etree.SubElement(line, f"{{{NS_A}}}solidFill")
        clr = etree.SubElement(fill, f"{{{NS_A}}}srgbClr", val=_hex(color))
        if alpha < 1.0:
            etree.SubElement(clr, f"{{{NS_A}}}alpha", val=str(round(alpha * 100000)))
    effect = sp_pr.find(f"{{{NS_A}}}effectLst")
    if effect is not None:
        effect.addprevious(line)
    else:
        sp_pr.append(line)


def _set_shadow(shape: Any, on: bool, dark: bool = False) -> None:
    sp_pr = shape._element.spPr
    for old in sp_pr.findall(f"{{{NS_A}}}effectLst"):
        sp_pr.remove(old)
    effect = etree.SubElement(sp_pr, f"{{{NS_A}}}effectLst")
    if on:
        shadow = etree.SubElement(
            effect,
            f"{{{NS_A}}}outerShdw",
            blurRad="228600",
            dist="38100",
            dir="5400000",
            algn="t",
            rotWithShape="0",
        )
        clr = etree.SubElement(shadow, f"{{{NS_A}}}srgbClr", val="000000")
        etree.SubElement(clr, f"{{{NS_A}}}alpha", val="30000" if dark else "10000")


def _shape(
    frame: Frame, kind: Any, box: tuple[int, int, int, int], name: str, corner: float = 0.0
) -> Any:
    shape = frame.slide.shapes.add_shape(kind, *(Emu(v) for v in box))
    _strip_style(shape)
    if corner > 0:
        try:
            shape.adjustments[0] = max(0.0, min(corner, 0.5))
        except (IndexError, AttributeError):
            pass
    shape.name = name
    if shape.has_text_frame:
        shape.text_frame.text = ""
    return shape


def _to_back(frame: Frame, shape: Any) -> None:
    """Фигура под содержание слайда: сразу за служебными элементами дерева."""
    tree = frame.slide.shapes._spTree
    tree.remove(shape._element)
    tree.insert(2, shape._element)


def _place(shape: Any, box: tuple[int, int, int, int]) -> None:
    shape.left, shape.top, shape.width, shape.height = (Emu(v) for v in box)


def _style_text(
    shape: Any,
    *,
    size: float,
    color: str | None = None,
    bold: bool | None = None,
    align: Any | None = None,
    anchor: Any = MSO_ANCHOR.TOP,
    spacing: float = 1.0,
    para_gap_pt: float = 0.0,
    family: str | None = None,
) -> None:
    """Кегль, цвет, выравнивание и интервалы всех абзацев рамки; поля рамки — нулевые, чтобы
    измерение и отрисовка совпадали."""
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Emu(0)
    frame.margin_top = frame.margin_bottom = Emu(0)
    frame.vertical_anchor = anchor
    body_pr = frame._txBody.find(f"{{{NS_A}}}bodyPr")
    if body_pr is not None:
        for tag in ("spAutoFit", "normAutofit", "noAutofit"):
            for old in body_pr.findall(f"{{{NS_A}}}{tag}"):
                body_pr.remove(old)
        etree.SubElement(body_pr, f"{{{NS_A}}}noAutofit")
    paragraphs = frame.paragraphs
    for index, paragraph in enumerate(paragraphs):
        if align is not None:
            paragraph.alignment = align
        paragraph.line_spacing = spacing
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(para_gap_pt if index < len(paragraphs) - 1 else 0)
        for run in paragraph.runs:
            run.font.size = Pt(size)
            if family:
                run.font.name = family
            if bold is not None:
                run.font.bold = bold
            if color:
                run.font.color.rgb = _rgb(color)


def _rgb(color: str) -> Any:
    from pptx.dml.color import RGBColor

    return RGBColor.from_string(_hex(color))  # type: ignore[no-untyped-call]


def _box(shape: Any) -> tuple[int, int, int, int]:
    return int(shape.left), int(shape.top), int(shape.width), int(shape.height)


def _on(fill: str, code: DesignCode) -> str:
    """Цвет текста и иконки поверх заливки."""
    return code.on_accent(fill)


def _surface(frame: Frame, accent: str) -> tuple[str | None, float]:
    """Подложка карточки: бледный акцент на светлом фоне, полупрозрачный белый на тёмном."""
    if frame.dark:
        return "#FFFFFF", 0.07
    if relative_luminance(frame.backdrop) < 0.9:
        # Светлый, но не белый фон (серый, бежевый): белая карточка отделяется от него сама.
        return "#FFFFFF", 1.0
    return _mix(accent, frame.backdrop, 0.07), 1.0


def _text_colors(frame: Frame, fill: str | None) -> tuple[str, str]:
    """(цвет заголовка, цвет текста) на подложке или фоне слайда."""
    from presentation_designer.library.build import legible

    surface = fill or frame.backdrop
    code = frame.code
    title = legible(code.title_color or code.text_color, surface, code, 4.5, accent=False)
    body = legible(code.muted_color, surface, code, 4.5, accent=False)
    return title, body


def _accent(frame: Frame, index: int) -> str:
    from presentation_designer.library.build import legible

    code = frame.code
    color = code.accent_at(index) if frame.look.colorful else code.accent
    return legible(color, frame.backdrop, code, 3.0, accent=True)


def _badge(
    frame: Frame,
    center_left: int,
    top: int,
    diameter: int,
    accent: str,
    *,
    icon: str | None,
    on_fill: str | None = None,
) -> Any:
    """Круг значка и иконка в нём (или ничего — номер пишется в рамку номера)."""
    look = frame.look
    circle = _shape(frame, MSO_SHAPE.OVAL, (center_left, top, diameter, diameter), "Badge")
    if on_fill is not None:
        # Значок на залитой карточке: светлый круг и иконка цветом заливки.
        _set_fill(circle, _on(on_fill, frame.code), 1.0)
        glyph = on_fill
    elif look.badge == "solid" or frame.dark:
        _set_fill(circle, accent, 1.0)
        glyph = _on(accent, frame.code)
    else:
        _set_fill(circle, _mix(accent, frame.backdrop, 0.14), 1.0)
        glyph = accent
    _set_line(circle, None)
    _set_shadow(circle, False)
    if icon:
        size = int(diameter * 0.54)
        offset = (diameter - size) // 2
        iconset.add_icon(
            frame.slide, icon, center_left + offset, top + offset, size, glyph, weight=1.1
        )
    return circle


# ---------- карточки ----------


@dataclass
class CardItem:
    index: int
    plate: Any | None
    bbox: tuple[float, float, float, float]
    number: Any | None
    icon_slot: str | None
    title: Any | None
    body: Any | None
    title_text: str
    body_text: str


def _card_items(
    frame: Frame, composition: Composition, refs: dict[str, str], card_ids: list[str]
) -> list[CardItem]:
    plates = dict(zip((c.index for c in composition.cards), card_ids, strict=False))
    by_card: dict[int, list[CompositionSlot]] = {}
    for slot in composition.slots:
        if slot.card_index:
            by_card.setdefault(slot.card_index, []).append(slot)
    items: list[CardItem] = []
    for card in composition.cards:
        slots = by_card.get(card.index) or []
        number = icon_slot = title = body = None
        for slot in slots:
            ref = refs.get(slot.slot_id)
            if slot.kind == "icon":
                icon_slot = slot.slot_id
            elif slot.slot_id.endswith("_number"):
                number = frame.get(ref)
            elif slot.kind in ("label", "subtitle") and title is None:
                title = frame.get(ref)
            elif slot.kind in ("caption", "body", "bullets") and body is None:
                body = frame.get(ref)
        plate = frame.get(plates.get(card.index))
        if plate is None and title is None and body is None:
            continue  # незаполненная карточка убрана сборкой
        title_text = frame.text_of(str(title.shape_id)) if title is not None else ""
        body_text = frame.text_of(str(body.shape_id)) if body is not None else ""
        if not title_text and not body_text:
            continue
        items.append(
            CardItem(
                index=card.index,
                plate=plate,
                bbox=card.bbox,
                number=number,
                icon_slot=icon_slot,
                title=title,
                body=body,
                title_text=title_text,
                body_text=body_text,
            )
        )
    return items


def _rows(items: list[CardItem]) -> list[list[CardItem]]:
    rows: list[list[CardItem]] = []
    for item in sorted(items, key=lambda i: (round(i.bbox[1], 2), i.bbox[0])):
        if rows and abs(rows[-1][0].bbox[1] - item.bbox[1]) < 0.02:
            rows[-1].append(item)
        else:
            rows.append([item])
    return rows


def _body_area(frame: Frame, refs: dict[str, str]) -> tuple[int, int]:
    """Верх и низ рабочей области под заголовком (EMU)."""
    title = frame.get(refs.get("title"))
    code = frame.code
    bottom = int(frame.height * (1 - max(code.margins.get("bottom", 0.08), 0.07)))
    if title is None:
        return int(frame.height * 0.2), bottom
    lines_h = _title_text_height(frame, title)
    top = int(title.top) + lines_h + int(frame.height * 0.07)
    return top, bottom


def _title_text_height(frame: Frame, title: Any) -> int:
    runs = [r for p in title.text_frame.paragraphs for r in p.runs]
    size = float(runs[0].font.size.pt) if runs and runs[0].font.size else frame.code.title_pt
    inset_l = int(title.text_frame.margin_left or 0)
    inset_r = int(title.text_frame.margin_right or 0)
    inset_t = int(title.text_frame.margin_top or 0)
    text = frame.text_of(str(title.shape_id))
    height = text_height(
        text,
        frame.code.font_for("title"),
        size,
        int(title.width) - inset_l - inset_r,
        bold=bool(runs and runs[0].font.bold),
    )
    return inset_t + height


@dataclass
class _Grid:
    """Раскладка ряда карточек: размеры в EMU и кегли."""

    cols: int
    col_w: int
    col_gap: int
    row_gap: int
    pad: int
    badge_d: int
    side: bool  # значок слева от текста (широкие карточки), иначе сверху
    title_pt: float
    body_pt: float

    @property
    def text_w(self) -> int:
        inner = self.col_w - 2 * self.pad
        return inner - (self.badge_d + self.pad) if self.side else inner


def _dress_cards(
    frame: Frame,
    composition: Composition,
    refs: dict[str, str],
    card_ids: list[str],
    queries: dict[str, str],
) -> bool:
    items = _card_items(frame, composition, refs, card_ids)
    if not items:
        return False
    _drop_slot_icons(frame, composition)
    code, look = frame.code, frame.look
    rows = _rows(items)
    cols = max(len(r) for r in rows)
    top, bottom = _body_area(frame, refs)
    avail = bottom - top
    width, height = frame.width, frame.height
    left = min(int(i.bbox[0] * width) for i in items)
    right = max(int((i.bbox[0] + i.bbox[2]) * width) for i in items)
    col_gap = int(width * 0.022)
    col_w = (right - left - col_gap * (cols - 1)) // cols
    line = look.card == "line"
    grid = _Grid(
        cols=cols,
        col_w=col_w,
        col_gap=col_gap,
        row_gap=int(height * 0.035),
        pad=0 if line else int(min(col_w * 0.08, height * 0.04)),
        badge_d=int(height * (0.09 if cols <= 3 else 0.075)),
        # Широкая карточка (две колонки и меньше): значок слева, текст рядом — так две строки
        # карточек помещаются по высоте, а строки текста не растягиваются на полслайда.
        side=col_w > width * 0.3,
        title_pt=max(code.body_pt * 1.2, min(code.subtitle_pt, code.body_pt * 1.45)),
        body_pt=code.body_pt,
    )
    if line:
        grid.pad = 0
    scale = {1: 1.1, 2: 1.05, 3: 1.0}.get(cols, 0.9)
    grid.title_pt *= scale
    grid.body_pt *= scale
    title_font, body_font = code.font_for("subtitle"), code.font_for("body")
    numbered = any(i.number is not None for i in items) or composition.family == "process"
    icons = [None] * len(items) if numbered else _card_icons(items, queries, composition)
    use_icons = not numbered and all(icons)
    show_badge = use_icons or numbered or len(items) > 1
    stroke = int(width * 0.004) if line else 0
    title_gap = int(height * 0.012)
    badge_gap = int(height * 0.028)

    def text_block(item: CardItem, t_pt: float, b_pt: float, text_w: int) -> tuple[int, int]:
        t_h = (
            text_height(item.title_text, title_font, t_pt, text_w, bold=True, spacing=0.95)
            if item.title_text
            else 0
        )
        b_h = (
            text_height(item.body_text, body_font, b_pt, text_w, spacing=1.05, para_gap_pt=4)
            if item.body_text
            else 0
        )
        return t_h, b_h

    def card_height(item: CardItem, g: _Grid) -> int:
        t_h, b_h = text_block(item, g.title_pt, g.body_pt, g.text_w - stroke * 3)
        text = t_h + (title_gap if t_h and b_h else 0) + b_h
        if not show_badge:
            return g.pad * 2 + text
        if g.side:
            return g.pad * 2 + max(g.badge_d, text)
        return g.pad * 2 + g.badge_d + badge_gap + text

    def measure(g: _Grid) -> tuple[int, list[int]]:
        heights = [max(card_height(i, g) for i in row) for row in rows]
        return sum(heights) + g.row_gap * (len(rows) - 1), heights

    def scaled(k: float) -> _Grid:
        return dataclasses.replace(grid, title_pt=grid.title_pt * k, body_pt=grid.body_pt * k)

    # Кегль растёт, пока группа занимает меньше 70 % рабочей области (не больше чем в 1,4
    # раза), и уменьшается, пока не поместится; ниже подписи шаблона текст не опускается.
    k = 1.0
    while k < 1.4 and measure(scaled(k + 0.05))[0] <= avail * 0.7:
        k += 0.05
    while measure(scaled(k))[0] > avail and grid.body_pt * k > code.caption_pt:
        k *= 0.94
    grid = scaled(k)
    total, heights = measure(grid)
    if len(rows) == 1 and not line:
        # Однорядные карточки не ниже 45 % области: иначе широкие короткие плашки выглядят
        # полосками. Текст в высокой карточке стоит сверху, как в карточке образца.
        heights = [max(heights[0], int(avail * 0.45))]
        total = heights[0]
    y = top + max(0, int((avail - total) * TOP_SHARE))
    first_row_y = y
    for row, row_h in zip(rows, heights, strict=True):
        row_left = left + (right - left - (len(row) * col_w + (len(row) - 1) * col_gap)) // 2
        for position, item in enumerate(row):
            order = items.index(item)
            x = row_left + position * (col_w + col_gap)
            accent = _accent(frame, order)
            fill = _style_plate(frame, item, (x, y, col_w, row_h), accent, order)
            on_fill = fill if look.card == "accent" and order == 0 and fill else None
            title_color, body_color = _text_colors(frame, fill)
            if on_fill:
                title_color = body_color = _on(on_fill, code)
            text_x = x + grid.pad + stroke * 3
            cursor = y + grid.pad
            if show_badge:
                badge_box = (text_x, cursor, grid.badge_d)
                _badge(frame, text_x, cursor, grid.badge_d, accent,
                       icon=icons[order] if use_icons else None, on_fill=on_fill)  # fmt: skip
                if not use_icons:
                    _number_in_badge(frame, item, order, badge_box, accent, on_fill)
                if grid.side:
                    text_x += grid.badge_d + grid.pad
                else:
                    cursor += grid.badge_d + badge_gap
            elif item.number is not None:
                item.number.text_frame.text = ""
            text_w = grid.text_w - stroke * 3
            t_h, b_h = text_block(item, grid.title_pt, grid.body_pt, text_w)
            if grid.side and show_badge:
                # Короткий текст — по центру значка, длинный — от верха карточки.
                block = t_h + (title_gap if t_h and b_h else 0) + b_h
                if block < grid.badge_d:
                    cursor += (grid.badge_d - block) // 2
            if item.title is not None and t_h:
                _place(item.title, (text_x, cursor, text_w, t_h + int(height * 0.008)))
                _style_text(item.title, size=grid.title_pt, color=title_color, bold=True,
                            spacing=0.95, align=PP_ALIGN.LEFT, family=title_font)  # fmt: skip
                cursor += t_h + title_gap
            if item.body is not None and b_h:
                _place(item.body, (text_x, cursor, text_w, b_h + int(height * 0.012)))
                _style_text(item.body, size=grid.body_pt, color=body_color, spacing=1.05,
                            para_gap_pt=4, align=PP_ALIGN.LEFT, family=body_font)  # fmt: skip
            if item.icon_slot:
                _drop(frame, refs.get(item.icon_slot))
        y += row_h + grid.row_gap
    if composition.family == "process" and len(rows) == 1 and len(items) > 1:
        _process_arrows(frame, rows[0], left, grid, first_row_y)
    return True


def _drop_slot_icons(frame: Frame, composition: Composition) -> None:
    """Иконки, которые вёрстка уже поставила в слоты `icon` карточек: значок рисуется заново
    вместе с кругом, старая иконка в прежнем месте была бы дублем."""
    boxes = [card.bbox for card in composition.cards]
    for shape in list(frame.slide.shapes):
        if not str(shape.name).startswith("Icon "):
            continue
        cx = (int(shape.left) + int(shape.width) / 2) / frame.width
        cy = (int(shape.top) + int(shape.height) / 2) / frame.height
        if any(x <= cx <= x + w and y <= cy <= y + h for x, y, w, h in boxes):
            _drop_shape(shape)


def _card_icons(
    items: list[CardItem], queries: dict[str, str], composition: Composition
) -> list[str | None]:
    texts = []
    for item in items:
        query = queries.get(item.icon_slot or "") if item.icon_slot else ""
        texts.append(query or item.title_text or item.body_text)
    icons = iconset.pick_icons(texts)
    # Запрос модели мог не найтись — тогда по заголовку карточки.
    for position, item in enumerate(items):
        if icons[position] is None and texts[position] != item.title_text:
            icons[position] = iconset.find_icon(item.title_text, exclude=set(filter(None, icons)))
    return icons


def _style_plate(
    frame: Frame, item: CardItem, box: tuple[int, int, int, int], accent: str, order: int
) -> str | None:
    """Плашка карточки по виду варианта; возвращает цвет её заливки (None — фон слайда)."""
    code, look = frame.code, frame.look
    plate = item.plate
    if look.card == "line":
        if plate is not None:
            _drop_shape(plate)
        x, y, _w, h = box
        bar = _shape(
            frame,
            MSO_SHAPE.RECTANGLE,
            (x, y, max(int(frame.width * 0.0035), 25400), h),
            "Accent line",
        )
        _set_fill(bar, accent)
        _set_line(bar, None)
        _set_shadow(bar, False)
        return None
    if plate is None:
        plate = _shape(frame, MSO_SHAPE.ROUNDED_RECTANGLE, box, "Card")
        _to_back(frame, plate)
    _place(plate, box)
    corner = code.corner_ratio if code.card_geometry == "roundRect" else 0.0
    if corner > 0:
        try:
            plate.adjustments[0] = min(corner, 0.5)
        except (IndexError, AttributeError):
            pass
    if look.card == "accent" and order == 0:
        _set_fill(plate, accent)
        _set_line(plate, None)
        _set_shadow(plate, False)
        return accent
    if look.card == "accent":
        _set_fill(plate, None)
        _set_line(plate, "#FFFFFF" if frame.dark else accent, 1.25, 0.35 if frame.dark else 1.0)
        _set_shadow(plate, False)
        return None
    color, alpha = _surface(frame, accent)
    _set_fill(plate, color, alpha)
    if frame.dark:
        _set_line(plate, "#FFFFFF", 0.75, 0.14)
    else:
        _set_line(plate, None)
    _set_shadow(plate, code.shadow and not frame.dark)
    if not frame.dark:
        # Акцентная полоса по верху плашки.
        x, y, w, _h = box
        bar_h = max(int(frame.height * 0.007), 38100)
        inset = int(w * min(corner, 0.5) * 0.5) if corner else 0
        bar = _shape(frame, MSO_SHAPE.RECTANGLE, (x + inset, y, w - 2 * inset, bar_h), "Accent bar")
        _set_fill(bar, accent)
        _set_line(bar, None)
        _set_shadow(bar, False)
    return color if alpha >= 1.0 else None


def _number_in_badge(
    frame: Frame,
    item: CardItem,
    order: int,
    badge: tuple[int, int, int],
    accent: str,
    on_fill: str | None,
) -> None:
    """Номер карточки в круге значка: рамка номера (или новая) по центру круга."""
    x, y, d = badge
    shape = item.number
    text = frame.text_of(str(shape.shape_id)) if shape is not None else ""
    text = text or f"{order + 1:02d}"
    if shape is None:
        shape = frame.slide.shapes.add_textbox(Emu(x), Emu(y), Emu(d), Emu(d))
        shape.name = "Badge number"
        shape.text_frame.text = text
    else:
        _place(shape, (x, y, d, d))
    if on_fill is not None:
        color = on_fill
    elif frame.look.badge == "solid" or frame.dark:
        color = _on(accent, frame.code)
    else:
        color = accent
    size = d / EMU_PT * (0.36 if len(text) <= 2 else 0.28)
    _style_text(
        shape, size=size, color=color, bold=True, align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE, family=frame.code.font_for("title"),
    )  # fmt: skip
    # Номер поверх круга.
    tree = frame.slide.shapes._spTree
    tree.remove(shape._element)
    tree.append(shape._element)


def _process_arrows(frame: Frame, row: list[CardItem], left: int, grid: _Grid, y: int) -> None:
    """Шаги процесса: шеврон между карточками на уровне значков."""
    size = int(grid.col_gap * 0.9)
    cy = y + grid.pad + grid.badge_d // 2
    for position in range(len(row) - 1):
        x = left + (position + 1) * grid.col_w + position * grid.col_gap
        x += (grid.col_gap - size) // 2
        try:
            iconset.add_icon(
                frame.slide, "chevron-right", x, cy - size // 2, size, _accent(frame, position)
            )
        except KeyError:
            return


def _drop(frame: Frame, shape_id: str | None) -> None:
    shape = frame.get(shape_id)
    if shape is not None:
        _drop_shape(shape)


def _drop_shape(shape: Any) -> None:
    element = shape._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


# ---------- список ----------


LEAD_SEPARATORS = (": ", " — ", " – ")


def _dress_bullets(frame: Frame, composition: Composition, refs: dict[str, str]) -> bool:
    boxes: list[Any] = [
        frame.get(refs.get(slot.slot_id)) for slot in composition.slots if slot.kind == "bullets"
    ]
    boxes = [b for b in boxes if b is not None and frame.text_of(str(b.shape_id))]
    if not boxes:
        return False
    for box in boxes:
        split_breaks(box)
    code = frame.code
    top, bottom = _body_area(frame, refs)
    avail = bottom - top
    single = len(boxes) == 1
    title = frame.get(refs.get("title"))
    left = int(boxes[0].left)
    right = max(int(b.left + b.width) for b in boxes)
    panel_w = int((right - left) * 0.34) if single else 0
    gap = int(frame.width * 0.035) if single else 0
    text_w = (right - left - panel_w - gap) if single else int(boxes[0].width)
    family = code.font_for("body")
    items_all = [[p.text.strip() for p in b.text_frame.paragraphs if p.text.strip()] for b in boxes]
    para_gap = 0.55
    size, block = _list_size(frame, items_all, text_w, avail)
    y = top + max(0, int((avail - block) * TOP_SHARE))
    accent = _accent(frame, 0)
    title_color, body_color = _text_colors(frame, None)
    for box in boxes:
        x = left if single else int(box.left)
        _place(box, (x, y, text_w, block + int(frame.height * 0.02)))
        _style_text(
            box, size=size, color=body_color, spacing=1.1, para_gap_pt=size * para_gap,
            align=PP_ALIGN.LEFT, family=family,
        )  # fmt: skip
        _bullets_with_leads(box, accent, title_color, size)
    if single and panel_w > 0:
        _topic_panel(frame, (right - panel_w, y, panel_w, max(block, int(avail * 0.62))), title)
    return True


def _typed_glyph(text: str) -> str | None:
    stripped = text.lstrip()
    for glyph in BULLET_GLYPHS:
        after = stripped[len(glyph) : len(glyph) + 1]
        if stripped.startswith(glyph) and after in (" ", "\t", " "):
            return glyph
    return None


def _bullets_with_leads(box: Any, accent: str, lead_color: str, size: float) -> None:
    """Маркер акцентом с висячим отступом и жирное начало пункта до двоеточия или тире.

    Текст не меняется: если маркеры набраны в самом тексте («• Октябрь: …»), набранный знак
    перекрашивается в акцент, а пробел после него становится табуляцией до отступа; абзац
    без знака среди таких — вводная фраза, она идёт без маркера. Если знаков в тексте нет,
    маркер ставится оформлением абзаца."""
    indent = int(size * 1.6 * EMU_PT)
    paragraphs = [p for p in box.text_frame.paragraphs if p.text.strip()]
    typed = any(_typed_glyph(p.text) for p in paragraphs)
    for paragraph in paragraphs:
        p_pr = paragraph._p.get_or_add_pPr()
        for tag in ("buNone", "buChar", "buAutoNum", "buClr", "buSzPct", "buFont", "buBlip"):
            for old in p_pr.findall(f"{{{NS_A}}}{tag}"):
                p_pr.remove(old)
        glyph = _typed_glyph(paragraph.text)
        if typed and glyph is None:
            p_pr.set("marL", "0")
            p_pr.set("indent", "0")
            etree.SubElement(p_pr, f"{{{NS_A}}}buNone")
            continue
        p_pr.set("marL", str(indent))
        p_pr.set("indent", str(-indent))
        skip = 0
        if glyph is not None:
            etree.SubElement(p_pr, f"{{{NS_A}}}buNone")
            skip = _color_glyph(paragraph, glyph, accent)
        else:
            clr = etree.SubElement(p_pr, f"{{{NS_A}}}buClr")
            etree.SubElement(clr, f"{{{NS_A}}}srgbClr", val=_hex(accent))
            etree.SubElement(p_pr, f"{{{NS_A}}}buSzPct", val="100000")
            etree.SubElement(p_pr, f"{{{NS_A}}}buFont", typeface="Arial")
            etree.SubElement(p_pr, f"{{{NS_A}}}buChar", char="■")
        _bold_lead(paragraph, lead_color, skip=skip)


def _color_glyph(paragraph: Any, glyph: str, accent: str) -> int:
    """Набранный маркер — отдельным прогоном цвета акцента, пробел после него — табуляция.
    Возвращает число прогонов маркера (их пропускает выделение начала пункта)."""
    runs = list(paragraph.runs)
    if not runs:
        return 0
    first = runs[0]
    text = first.text
    start = len(text) - len(text.lstrip())
    end = start + len(glyph)
    if text[start:end] != glyph:
        return 0
    rest = text[end:].lstrip("  \t")
    if rest:
        tail = copy_run(first)
        tail.text = rest
    first.text = text[:start] + glyph + "\t"
    first.font.color.rgb = _rgb(accent)
    return 1


def _bold_lead(paragraph: Any, color: str, skip: int = 0) -> None:
    runs = list(paragraph.runs)[skip:]
    if not runs:
        return
    text = "".join(r.text for r in runs)
    cut = -1
    for sep in LEAD_SEPARATORS:
        position = text.find(sep)
        if 0 < position <= 48 and (cut < 0 or position < cut):
            cut = position + len(sep.rstrip())
    if cut <= 0:
        return
    # Разрез по первой рамке прогона, где кончается начало пункта.
    consumed = 0
    for run in runs:
        end = consumed + len(run.text)
        if end <= cut:
            run.font.bold = True
            run.font.color.rgb = _rgb(color)
        elif consumed < cut:
            # Хвост копируется до перекраски: у него остаются цвет и начертание текста.
            head, rest = run.text[: cut - consumed], run.text[cut - consumed :]
            tail = copy_run(run)
            tail.text = rest
            run.text = head
            run.font.bold = True
            run.font.color.rgb = _rgb(color)
            break
        consumed = end


def copy_run(run: Any) -> Any:
    """Копия прогона сразу за ним (тот же вид), возвращается как объект python-pptx."""
    import copy

    from pptx.text.text import _Run

    element = copy.deepcopy(run._r)
    run._r.addnext(element)
    return _Run(element, run._parent)


def _topic_panel(frame: Frame, box: tuple[int, int, int, int], title: Any | None) -> None:
    """Панель с крупной иконкой темы слайда рядом со списком."""
    text = frame.text_of(str(title.shape_id)) if title is not None else ""
    icon = iconset.find_icon(text)
    accent = _accent(frame, 0)
    x, y, w, h = box
    corner = frame.code.corner_ratio if frame.code.card_geometry == "roundRect" else 0.04
    panel = _shape(frame, MSO_SHAPE.ROUNDED_RECTANGLE, box, "Topic panel", corner=max(corner, 0.04))
    if frame.dark:
        _set_fill(panel, "#FFFFFF", 0.06)
        _set_line(panel, "#FFFFFF", 0.75, 0.14)
    else:
        _set_fill(panel, _mix(accent, frame.backdrop, 0.09))
        _set_line(panel, None)
    _set_shadow(panel, False)
    _to_back(frame, panel)
    if icon:
        size = int(min(w, h) * 0.42)
        iconset.add_icon(
            frame.slide, icon, x + (w - size) // 2, y + (h - size) // 2, size, accent, weight=0.7
        )
    else:
        # Без иконки — крупный акцентный круг-знак, чтобы панель не пустовала.
        d = int(min(w, h) * 0.36)
        ring = _shape(
            frame, MSO_SHAPE.OVAL, (x + (w - d) // 2, y + (h - d) // 2, d, d), "Topic mark"
        )
        _set_fill(ring, None)
        _set_line(ring, accent, max(2.0, d / EMU_PT * 0.06))
        _set_shadow(ring, False)


# ---------- числа ----------


def _dress_numbers(
    frame: Frame, composition: Composition, refs: dict[str, str], card_ids: list[str]
) -> bool:
    code = frame.code
    values: list[tuple[Any, Any]] = [
        (slot, frame.get(refs.get(slot.slot_id)))
        for slot in composition.slots
        if slot.kind == "number"
    ]
    values = [(s, v) for s, v in values if v is not None and frame.text_of(str(v.shape_id))]
    if not values:
        return False
    number_font = code.font_for("number")
    for order, (_slot, shape) in enumerate(values):
        text = frame.text_of(str(shape.shape_id))
        width = int(shape.width)
        size = code.number_pt * (1.5 if composition.family == "hero_number" else 1.15)
        # Число растёт до ширины своей рамки, но не выше высоты рамки.
        from presentation_designer.shared import text_metrics

        font = _font(number_font, True)
        while size > code.title_pt:
            w_pt = text_metrics.text_width_pt(text, font, size)
            h_pt = _line_pt(number_font, size, True)
            if w_pt * EMU_PT <= width * 0.94 and h_pt * EMU_PT <= int(shape.height) * 1.05:
                break
            size *= 0.94
        accent = _accent(frame, order)
        _style_text(
            shape, size=size, color=accent, bold=True, align=None,
            anchor=MSO_ANCHOR.BOTTOM, family=number_font,
        )  # fmt: skip
        # Короткая акцентная черта над числом.
        bar_w = int(frame.width * 0.035)
        bar_h = max(int(frame.height * 0.008), 38100)
        align = shape.text_frame.paragraphs[0].alignment
        x = int(shape.left)
        if align == PP_ALIGN.CENTER:
            x = int(shape.left + (shape.width - bar_w) / 2)
        y = int(shape.top) - bar_h - int(frame.height * 0.012)
        bar = _shape(frame, MSO_SHAPE.RECTANGLE, (x, max(y, 0), bar_w, bar_h), "Accent bar")
        _set_fill(bar, accent)
        _set_line(bar, None)
        _set_shadow(bar, False)
    for slot in composition.slots:
        if slot.kind == "label" and slot.slot_id.endswith("_label"):
            shape = frame.get(refs.get(slot.slot_id))
            if shape is not None and frame.text_of(str(shape.shape_id)):
                _, body_color = _text_colors(frame, None)
                runs = [r for p in shape.text_frame.paragraphs for r in p.runs]
                size = max(
                    code.body_pt, float(runs[0].font.size.pt) if runs and runs[0].font.size else 0
                )
                for run in runs:
                    run.font.size = Pt(size)
                    run.font.color.rgb = _rgb(body_color)
    for card_id in card_ids:
        plate = frame.get(card_id)
        if plate is not None:
            color, alpha = _surface(frame, code.accent)
            _set_fill(plate, color, alpha)
            _set_line(plate, "#FFFFFF" if frame.dark else None, 0.75, 0.14)
            _set_shadow(plate, code.shadow and not frame.dark)
    return True


# ---------- одна мысль ----------


def _dress_statement(frame: Frame, composition: Composition, refs: dict[str, str]) -> bool:
    main = next(
        (
            frame.get(refs.get(slot.slot_id))
            for slot in composition.slots
            if slot.slot_id in ("statement", "quote")
        ),
        None,
    )
    if main is None or not frame.text_of(str(main.shape_id)):
        return False
    code = frame.code
    if split_breaks(main) >= 3:
        frame.shapes = {str(s.shape_id): s for s in frame.slide.shapes}
    text = frame.text_of(str(main.shape_id))
    top, bottom = _body_area(frame, refs)
    support = next(
        (
            frame.get(refs.get(slot.slot_id))
            for slot in composition.slots
            if slot.slot_id in ("support", "author")
        ),
        None,
    )
    s_text = frame.text_of(str(support.shape_id)) if support is not None else ""
    s_size = code.body_pt
    s_gap = int(frame.height * 0.035)
    paragraphs = [p for p in text.split("\n") if p.strip()]
    accent = _accent(frame, 0)
    title_color, body_color = _text_colors(frame, None)
    bar_w = max(int(frame.width * 0.006), 50800)
    x = int(main.left)
    indent = bar_w * 4
    width = int(main.width) - indent
    s_h = text_height(s_text, code.font_for("body"), s_size, width) if s_text else 0
    avail = bottom - top - (s_h + s_gap if s_text else 0)
    if len(paragraphs) >= 3:
        # «Одна мысль», в которую модель сложила пункты: это список, и оформляется он как список.
        size, h = _list_size(frame, [paragraphs], width, avail)
        y = top + max(0, int((avail - h) * TOP_SHARE))
        _place(main, (x, y, width + indent, h + int(frame.height * 0.02)))
        _style_text(main, size=size, color=body_color, spacing=1.1, para_gap_pt=size * 0.55,
                    align=PP_ALIGN.LEFT, family=code.font_for("body"))  # fmt: skip
        _bullets_with_leads(main, accent, title_color, size)
        end = y + h
    else:
        family = code.font_for("subtitle")
        size = code.subtitle_pt

        def fits(pt: float, share: float) -> bool:
            return text_height(text, family, pt, width, spacing=1.05) <= avail * share

        while fits(size * 1.06, 0.55) and size < code.title_pt * 1.1:
            size *= 1.06
        while not fits(size, 0.8) and size > code.body_pt:
            size *= 0.94
        h = text_height(text, family, size, width, spacing=1.05)
        y = top + max(0, int((avail - h) * TOP_SHARE))
        _place(main, (x + indent, y, width, h + int(frame.height * 0.02)))
        _style_text(main, size=size, color=title_color, spacing=1.05, family=family,
                    align=PP_ALIGN.LEFT)  # fmt: skip
        bar = _shape(frame, MSO_SHAPE.RECTANGLE, (x, y, bar_w, h), "Accent line")
        _set_fill(bar, accent)
        _set_line(bar, None)
        _set_shadow(bar, False)
        end = y + h
    if support is not None and s_text:
        _place(support, (x + indent, end + s_gap, width, s_h + 25400))
        _style_text(support, size=s_size, color=body_color, family=code.font_for("body"))
    return True


BULLET_GLYPHS = ("•", "·", "▪", "■", "–", "—", "-", "*")


def split_breaks(shape: Any) -> int:
    """Разрывы строк внутри абзаца (`a:br`) → отдельные абзацы: так каждый пункт получает
    свой маркер и отступ. Возвращает число непустых абзацев."""
    import copy

    body = shape.text_frame._txBody
    for p in list(body.findall(f"{{{NS_A}}}p")):
        if p.find(f"{{{NS_A}}}br") is None:
            continue
        p_pr = p.find(f"{{{NS_A}}}pPr")
        end = p.find(f"{{{NS_A}}}endParaRPr")
        groups: list[list[Any]] = [[]]
        for child in list(p):
            if child.tag == f"{{{NS_A}}}br":
                groups.append([])
            elif child.tag in (f"{{{NS_A}}}r", f"{{{NS_A}}}fld"):
                groups[-1].append(child)
        anchor = p
        for group in groups:
            if not group:
                continue
            # Копия абзаца, а не etree.Element: python-pptx узнаёт свои элементы по классу.
            new_p = copy.deepcopy(p)
            for child in list(new_p):
                new_p.remove(child)
            if p_pr is not None:
                new_p.append(copy.deepcopy(p_pr))
            for run in group:
                new_p.append(run)
            if end is not None:
                new_p.append(copy.deepcopy(end))
            anchor.addnext(new_p)
            anchor = new_p
        body.remove(p)
    count = 0
    for paragraph in shape.text_frame.paragraphs:
        count += 1 if paragraph.text.strip() else 0
    return count


def _list_size(frame: Frame, lists: list[list[str]], width: int, avail: int) -> tuple[float, int]:
    """Кегль списка: растёт до 75 % высоты области (не больше 1,45 основного), уменьшается,
    пока не поместится. Возвращает кегль и высоту самого длинного столбца."""
    code = frame.code
    family = code.font_for("body")

    def height_at(size: float) -> int:
        indent = int(size * 1.6 * EMU_PT)
        return max(
            text_height(
                "\n".join(items),
                family,
                size,
                width - indent,
                spacing=1.1,
                para_gap_pt=size * 0.55,
            )
            for items in lists
        )

    size = code.body_pt * 1.05
    while height_at(size * 1.06) < avail * 0.75 and size < code.body_pt * 1.45:
        size *= 1.06
    while height_at(size) > avail and size > code.caption_pt:
        size *= 0.94
    return size, height_at(size)


# ---------- заголовок ----------


def _title_mark(frame: Frame, title_id: str | None) -> bool:
    title = frame.get(title_id)
    if title is None or not frame.text_of(str(title.shape_id)):
        return False
    paragraph = title.text_frame.paragraphs[0]
    if paragraph.alignment not in (None, PP_ALIGN.LEFT):
        return False
    height = _title_text_height(frame, title)
    x = int(title.left) + int(title.text_frame.margin_left or 0)
    y = int(title.top) + height + int(frame.height * 0.018)
    bar = _shape(
        frame,
        MSO_SHAPE.RECTANGLE,
        (x, y, int(frame.width * 0.045), max(int(frame.height * 0.008), 38100)),
        "Title mark",
    )
    _set_fill(bar, _accent(frame, 0))
    _set_line(bar, None)
    _set_shadow(bar, False)
    return True


__all__ = ["LOOKS", "Look", "dress_slide", "look_for", "text_height"]
