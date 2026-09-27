"""Bounded emphasis for native text: measure first, style only if the whole group fits.

Unlike card reflow this pass does not move frames, replace runs or shrink planned text.
It can therefore share a fitting policy across leads, cover titles and KPI labels without
changing the neighbouring diagram or the template's alignment and paragraph styling.
"""

from __future__ import annotations

import math
from typing import Any

from pptx.util import Pt

from presentation_designer.generation.capacity import wrap_lines
from presentation_designer.shared import text_metrics

EMU_PT = 12700


def required_height(shape: Any, size: float, family: str | None) -> int:
    """Measure existing paragraphs, breaks, insets and spacing at a uniform candidate size.

    Mixed runs use the most conservative font measurement. Explicit paragraph line heights
    cannot make the glyphs shorter. A small renderer reserve is kept in both dimensions.
    """
    frame = shape.text_frame
    width = (int(shape.width) - frame.margin_left - frame.margin_right) / EMU_PT * 0.96
    if width <= 0:
        return int(shape.height) + 1
    total = 0.0
    for paragraph in frame.paragraphs:
        runs = list(paragraph.runs)
        fonts = [
            text_metrics.resolve_font(
                run.font.name or family, bold=bool(run.font.bold), italic=bool(run.font.italic)
            )
            for run in runs
        ] or [text_metrics.resolve_font(family)]
        ppr = paragraph._p.pPr
        indent = 0
        if ppr is not None:
            indent = max(0, int(ppr.get("marL", "0"))) + max(0, int(ppr.get("marR", "0")))
            indent += max(0, int(ppr.get("indent", "0")))
        available = width - indent / EMU_PT
        if available <= 0:
            return int(shape.height) + 1
        lines = max(
            sum(wrap_lines(line, available, font, size) for line in paragraph.text.split("\v"))
            for font in fonts
        )
        line_height = max(text_metrics.line_metrics(font, size).line_height_pt for font in fonts)
        spacing = paragraph.line_spacing
        if isinstance(spacing, float):
            line_height *= max(1.0, spacing)
        elif spacing is not None:
            line_height = max(line_height, spacing.pt)
        total += lines * line_height
        total += (paragraph.space_before or 0) / EMU_PT
        total += (paragraph.space_after or 0) / EMU_PT
    return math.ceil(total * EMU_PT * 1.04 + int(frame.margin_top) + int(frame.margin_bottom))


def emphasize(shapes: list[Any], target_pt: float, family: str | None) -> bool:
    """Apply one fitting size to peers, or leave all untouched if no safe size exists.

    The lower bound is their largest planned run: a presentation should not gain visual
    emphasis by making its difficult labels smaller. Half-point steps keep sizes stable.
    """
    shapes = [s for s in shapes if s is not None and s.text_frame.text.strip()]
    if not shapes:
        return False
    sizes: list[float] = []
    for shape in shapes:
        runs = [run for paragraph in shape.text_frame.paragraphs for run in paragraph.runs]
        # Unknown inheritance may be larger than the target. Mixed explicit sizes
        # carry hierarchy (e.g. a footnote) which uniform emphasis must not flatten.
        if not runs or any(run.font.size is None for run in runs):
            return False
        local_sizes = {run.font.size.pt for run in runs}
        if len(local_sizes) != 1:
            return False
        sizes.extend(local_sizes)
    if not sizes:
        return False
    floor = max(sizes)
    target = max(floor, target_pt)
    candidates = sorted(
        {floor, target, *(step / 2 for step in range(math.ceil(floor * 2), int(target * 2) + 1))},
        reverse=True,
    )
    for size in candidates:
        if not all(required_height(shape, size, family) <= shape.height for shape in shapes):
            continue
        for shape in shapes:
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(size)
        return True
    return False
