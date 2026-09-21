"""Проверка движка PPTX (этап 0A): сохранность, клонирование, текст, диаграммы, рендер, замеры.

На каждом шаблоне из --templates-dir (и на собственной фикстуре) выполняются операции из
критерия FRAMEWORKS.md §6 и сохраняются выходные файлы в --out/<машина>/<шаблон>/:

  01-roundtrip.pptx     открытие и сохранение без изменений
  02-clone.pptx         клоны образцов с группами, обрезкой картинок, смешанными стилями, таблицей
  03-text.pptx          замена текста в клонах с сохранением оформления фрагментов
  04-native.pptx        нативная таблица и диаграмма на клоне
  05-chart-copies.pptx  две независимо отредактированные копии слайда с диаграммой
  06-pruned.pptx        удалены все образцы, остались только новые слайды
  *.pdf, thumbs/        рендер ONLYOFFICE и миниатюры pypdfium2, если ONLYOFFICE доступен

Результаты — results.json на машину; --report собирает все results.json из --out в блок
docs/pptx-capabilities.md между маркерами `<!-- probe:begin -->` и `<!-- probe:end -->`.
Ручная проверка в PowerPoint остаётся за пользователем: скрипт её не заменяет.

Запуск: uv run scripts/probe_pptx.py --templates-dir data/organizers --out runs/pptx-probe \\
            --report docs/pptx-capabilities.md
"""

from __future__ import annotations

import argparse
import json
import os
import pathlib
import platform
import re
import resource
import shutil
import socket
import sys
import time
from datetime import UTC, datetime
from typing import Any

from PIL import Image, ImageChops
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Emu

ROOT = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "scripts"))
from inspect_pptx import inspect_presentation  # noqa: E402
from presentation_designer.export import pdf as pdf_export  # noqa: E402
from presentation_designer.export import thumbnails  # noqa: E402
from presentation_designer.layout import ooxml  # noqa: E402

FIXTURE = ROOT / "tests" / "fixtures" / "pptx" / "mini_template.pptx"
DIFF_WIDTH = 640
PIXEL_THRESHOLD = 40  # разница яркости, с которой пиксель считается изменённым


# --- окружение ----------------------------------------------------------------------


def machine_info(label: str | None) -> dict[str, Any]:
    import pptx

    mem_mb: int | None = None
    if pathlib.Path("/proc/meminfo").exists():
        for line in pathlib.Path("/proc/meminfo").read_text().splitlines():
            if line.startswith("MemTotal:"):
                mem_mb = int(line.split()[1]) // 1024
    elif sys.platform == "darwin":
        try:
            import subprocess

            mem_mb = int(subprocess.check_output(["sysctl", "-n", "hw.memsize"])) // (1024 * 1024)
        except Exception:
            mem_mb = None
    in_docker = pathlib.Path("/.dockerenv").exists()
    default_label = f"{platform.system().lower()}-{platform.machine()}"
    if in_docker:
        default_label = f"docker-{platform.machine()}"
    return {
        "label": label or default_label,
        "hostname": socket.gethostname(),
        "platform": platform.platform(),
        "machine": platform.machine(),
        "in_docker": in_docker,
        "cpu_count": os.cpu_count(),
        "memory_mb": mem_mb,
        "python": platform.python_version(),
        "python_pptx": pptx.__version__,
        "renderer": "ONLYOFFICE" if pdf_export.renderer_configured() else None,
        "measured_at": datetime.now(UTC).isoformat(timespec="seconds"),
    }


def _rss_mb() -> float:
    usage = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
    return round(usage / (1024 * 1024) if sys.platform == "darwin" else usage / 1024, 1)


def _slug(name: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9А-Яа-яЁё]+", "-", name).strip("-")
    return slug[:60] or "template"


# --- выбор образцов -------------------------------------------------------------------


def pick_samples(details: list[dict[str, Any]]) -> dict[str, int | None]:
    """Индексы образцов по признакам; None, если признака в шаблоне нет."""

    def best(key: str, *, require: str | None = None) -> int | None:
        candidates = [d for d in details if d[key] > 0 and (require is None or d[require] > 0)]
        if not candidates:
            return None
        return max(candidates, key=lambda d: (d[key], d["pictures"], d["shapes"]))["index"]

    text_slides = [d for d in details if d["text_shapes"] > 0]
    content = min(text_slides, key=lambda d: d["shapes"])["index"] if text_slides else 0
    return {
        "group": best("groups"),
        "crop": best("cropped_pictures") or best("pictures"),
        "mixed": best("mixed_style_paragraphs"),
        "table": best("tables"),
        "chart": best("charts"),
        "content": content,
    }


# --- шаги проверки ------------------------------------------------------------------


class TemplateProbe:
    def __init__(
        self, path: pathlib.Path, out_dir: pathlib.Path, render: bool, timeout_s: int
    ) -> None:
        self.path = path
        self.out_dir = out_dir
        self.render = render
        self.timeout_s = timeout_s
        self.result: dict[str, Any] = {"name": path.name, "path": str(path), "steps": {}}
        self.out_dir.mkdir(parents=True, exist_ok=True)

    def run(self) -> dict[str, Any]:
        summary = inspect_presentation(self.path)
        self.result["inspect"] = {
            k: v for k, v in summary.items() if k not in ("slide_details", "package_check")
        }
        self.result["source_check"] = summary["package_check"]
        samples = pick_samples(summary["slide_details"])
        self.result["samples"] = samples
        self.step_load()
        self.step_roundtrip()
        clones = self.step_clone(samples)
        self.step_text(clones)
        native_index = self.step_native(samples)
        chart_pages = self.step_chart_copies(samples, native_index)
        self.step_prune(clones, native_index, chart_pages)
        if self.render:
            self.step_render(clones, chart_pages)
        return self.result

    def _save(self, prs: Any, name: str) -> pathlib.Path:
        target = self.out_dir / name
        started = time.perf_counter()
        prs.save(str(target))
        self.result["steps"].setdefault("save_seconds", {})[name] = round(
            time.perf_counter() - started, 3
        )
        return target

    def step_load(self) -> None:
        before = _rss_mb()
        started = time.perf_counter()
        prs = Presentation(str(self.path))
        opened = time.perf_counter() - started
        shapes = sum(len(slide.shapes) for slide in prs.slides)
        walked = time.perf_counter() - started
        self.result["steps"]["load"] = {
            "open_seconds": round(opened, 3),
            "open_and_walk_seconds": round(walked, 3),
            "shapes": shapes,
            "python_peak_rss_before_mb": before,
            "python_peak_rss_after_mb": _rss_mb(),
        }

    def step_roundtrip(self) -> None:
        prs = Presentation(str(self.path))
        out = self._save(prs, "01-roundtrip.pptx")
        diff = ooxml.compare_packages(self.path, out)
        check = ooxml.check_package(out)
        self.result["steps"]["roundtrip"] = {
            "size_before": self.path.stat().st_size,
            "size_after": out.stat().st_size,
            "diff": diff,
            "check": check.as_dict(),
            "ok": check.ok and not diff["changed_binary"] and not diff["only_in_b"],
        }

    def step_clone(self, samples: dict[str, int | None]) -> dict[str, tuple[int, int]]:
        """Клонирует выбранные образцы в конец; возвращает {признак: (исходный, клон)}."""
        prs = Presentation(str(self.path))
        original_count = len(prs.slides)
        clones: dict[str, tuple[int, int]] = {}
        details: dict[str, Any] = {}
        seen: dict[int, int] = {}
        for key in ("group", "crop", "mixed", "table"):
            index = samples.get(key)
            if index is None:
                continue
            if index in seen:
                clones[key] = (index, seen[index])
                continue
            source = prs.slides[index]
            clone = ooxml.clone_slide(prs, source)
            clone_index = len(prs.slides) - 1
            seen[index] = clone_index
            clones[key] = (index, clone_index)
            details[key] = {
                "source": index,
                "clone": clone_index,
                "shapes_equal": [s.name for s in source.shapes] == [s.name for s in clone.shapes],
                "layout_shared": clone.slide_layout.part is source.slide_layout.part,
                "sp_tree_equal": ooxml.shape_tree_signature(source)
                == ooxml.shape_tree_signature(clone),
            }
        out = self._save(prs, "02-clone.pptx")
        check = ooxml.check_package(out)
        self.result["steps"]["clone"] = {
            "original_slides": original_count,
            "slides_after": len(prs.slides),
            "clones": details,
            "check": check.as_dict(),
            "ok": check.ok
            and all(d["shapes_equal"] and d["sp_tree_equal"] for d in details.values()),
        }
        return clones

    def step_text(self, clones: dict[str, tuple[int, int]]) -> None:
        prs = Presentation(str(self.out_dir / "02-clone.pptx"))
        outcome: dict[str, Any] = {}
        mixed = clones.get("mixed")
        if mixed:
            slide = prs.slides[mixed[1]]
            paragraph = _first_mixed_paragraph(slide)
            if paragraph is not None:
                before = ooxml.paragraph_run_styles(paragraph)
                texts = [f"{s['text'].strip() or 'фрагмент'}·{i + 1}" for i, s in enumerate(before)]
                ooxml.set_run_texts(paragraph, texts)
                after = ooxml.paragraph_run_styles(paragraph)
                outcome["mixed_runs"] = {
                    "runs": len(before),
                    "styles_preserved": [b["attrs"] for b in before] == [a["attrs"] for a in after]
                    and [b["children"] for b in before] == [a["children"] for a in after],
                    "text": paragraph.text,
                }
        target = clones.get("group") or clones.get("crop") or mixed
        if target:
            slide = prs.slides[target[1]]
            paragraph = _first_text_paragraph(slide)
            if paragraph is not None:
                before = ooxml.paragraph_run_styles(paragraph)
                ooxml.replace_paragraph_text(paragraph, "Проверка замены текста\nвторая строка")
                after = ooxml.paragraph_run_styles(paragraph)
                outcome["paragraph"] = {
                    "first_run_style_preserved": (not before)
                    or (after and after[0]["attrs"] == before[0]["attrs"]),
                    "text": paragraph.text,
                }
        out = self._save(prs, "03-text.pptx")
        check = ooxml.check_package(out)
        outcome["check"] = check.as_dict()
        outcome["ok"] = check.ok and all(
            v.get("styles_preserved", v.get("first_run_style_preserved", True))
            for k, v in outcome.items()
            if k in ("mixed_runs", "paragraph")
        )
        self.result["steps"]["text"] = outcome

    def step_native(self, samples: dict[str, int | None]) -> int:
        prs = Presentation(str(self.out_dir / "03-text.pptx"))
        source = prs.slides[samples["content"] or 0]
        slide = ooxml.clone_slide(prs, source)
        index = len(prs.slides) - 1
        w, h = prs.slide_width, prs.slide_height
        table = slide.shapes.add_table(
            4, 3, Emu(int(w * 0.06)), Emu(int(h * 0.30)), Emu(int(w * 0.40)), Emu(int(h * 0.35))
        ).table
        for r in range(4):
            for c in range(3):
                table.cell(r, c).text = (
                    ("Показатель", "Май", "Июнь")[c] if r == 0 else f"{r * c + 1}"
                )
        data = CategoryChartData()
        data.categories = ["Май", "Июнь", "Июль"]
        data.add_series("Открываемость", (31, 38, 44))
        data.add_series("Отписки", (4.1, 3.2, 2.7))
        slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            Emu(int(w * 0.52)),
            Emu(int(h * 0.30)),
            Emu(int(w * 0.42)),
            Emu(int(h * 0.55)),
            data,
        )
        out = self._save(prs, "04-native.pptx")
        check = ooxml.check_package(out)
        reopened = Presentation(str(out))
        shapes = reopened.slides[index].shapes
        self.result["steps"]["native"] = {
            "slide": index,
            "table_present": any(getattr(s, "has_table", False) and s.has_table for s in shapes),
            "chart_present": any(getattr(s, "has_chart", False) and s.has_chart for s in shapes),
            "check": check.as_dict(),
            "ok": check.ok,
        }
        return index

    def step_chart_copies(self, samples: dict[str, int | None], native_index: int) -> list[int]:
        prs = Presentation(str(self.out_dir / "04-native.pptx"))
        base_index = samples["chart"] if samples["chart"] is not None else native_index
        base = prs.slides[base_index]
        copy_a = ooxml.clone_slide(prs, base)
        copy_b = ooxml.clone_slide(prs, base)
        pages = [len(prs.slides) - 2, len(prs.slides) - 1]
        data_a = CategoryChartData()
        data_a.categories = ["Q1", "Q2", "Q3"]
        data_a.add_series("Копия А", (5, 9, 12))
        data_b = CategoryChartData()
        data_b.categories = ["Янв", "Фев", "Мар", "Апр", "Май"]
        data_b.add_series("Копия Б", (100, 80, 60, 40, 20))
        _first_chart(copy_a).replace_data(data_a)
        _first_chart(copy_b).replace_data(data_b)
        out = self._save(prs, "05-chart-copies.pptx")
        check = ooxml.check_package(out)
        reopened = Presentation(str(out))
        described = []
        for index in (base_index, *pages):
            chart = _first_chart(reopened.slides[index])
            plot = chart.plots[0]
            described.append(
                {
                    "slide": index,
                    "chart_part": str(chart.part.partname),
                    "workbook_part": str(chart.part.chart_workbook.xlsx_part.partname),
                    "series": [s.name for s in plot.series],
                    "categories": list(plot.categories),
                }
            )
        parts = {d["chart_part"] for d in described}
        workbooks = {d["workbook_part"] for d in described}
        self.result["steps"]["chart_copies"] = {
            "base_slide": base_index,
            "base_is_template_chart": samples["chart"] is not None,
            "copies": pages,
            "charts": described,
            "independent": len(parts) == 3
            and len(workbooks) == 3
            and described[1]["categories"] != described[2]["categories"],
            "check": check.as_dict(),
            "ok": check.ok and len(parts) == 3 and len(workbooks) == 3,
        }
        return pages

    def step_prune(
        self, clones: dict[str, tuple[int, int]], native_index: int, chart_pages: list[int]
    ) -> None:
        prs = Presentation(str(self.out_dir / "05-chart-copies.pptx"))
        keep_indexes = sorted({c[1] for c in clones.values()} | {native_index, *chart_pages})
        keep = [prs.slides[i] for i in keep_indexes]
        before_parts = sum(1 for _ in prs.part.package.iter_parts())
        started = time.perf_counter()
        removed = ooxml.keep_only_slides(prs, keep)
        delete_seconds = time.perf_counter() - started
        out = self._save(prs, "06-pruned.pptx")
        check = ooxml.check_package(out)
        reopened = Presentation(str(out))
        layouts = sum(len(m.slide_layouts) for m in reopened.slide_masters)
        self.result["steps"]["prune"] = {
            "removed": removed,
            "kept": keep_indexes,
            "slides_after": len(reopened.slides),
            "delete_seconds": round(delete_seconds, 3),
            "parts_before": before_parts,
            "parts_after": check.parts,
            "size_before": (self.out_dir / "05-chart-copies.pptx").stat().st_size,
            "size_after": out.stat().st_size,
            "layouts_kept": layouts,
            "check": check.as_dict(),
            "ok": check.ok and len(reopened.slides) == len(keep_indexes),
        }

    # --- рендер ---

    def step_render(self, clones: dict[str, tuple[int, int]], chart_pages: list[int]) -> None:
        render: dict[str, Any] = {"conversions": {}}
        thumbs_dir = self.out_dir / "thumbs"
        if thumbs_dir.exists():
            shutil.rmtree(thumbs_dir)
        try:
            # Independent conversion keys avoid Document Server's result cache.
            render["conversions"]["original_first"] = self._convert(self.path, "00-original")
            render["conversions"]["original_second"] = self._convert(self.path, "00-original")
            render["conversions"]["original_repeat"] = self._convert(self.path, "00-original")
            for name in (
                "01-roundtrip",
                "02-clone",
                "03-text",
                "04-native",
                "05-chart-copies",
                "06-pruned",
            ):
                render["conversions"][name] = self._convert(self.out_dir / f"{name}.pptx", name)
        except pdf_export.ConversionError as exc:
            render["error"] = str(exc)
            self.result["render"] = render
            return
        original_pdf = self.out_dir / "00-original.pdf"
        render["pdf_fonts"] = _pdf_fonts(original_pdf)
        render["thumbnails"] = {}
        for name in ("00-original", "01-roundtrip", "02-clone", "05-chart-copies", "06-pruned"):
            result = thumbnails.render_thumbnails(
                self.out_dir / f"{name}.pdf", thumbs_dir / name, width_px=1280
            )
            render["thumbnails"][name] = {
                "pages": len(result.paths),
                "seconds": round(result.seconds, 3),
                "per_page_ms": round(result.seconds * 1000 / max(1, len(result.paths)), 1),
            }
        render["visual"] = {
            "roundtrip_vs_original": _compare_pdfs(
                original_pdf, self.out_dir / "01-roundtrip.pdf", thumbs_dir / "diff-roundtrip"
            ),
            "clones_vs_sources": _compare_pages(
                self.out_dir / "02-clone.pdf",
                [(src + 1, dst + 1) for src, dst in clones.values()],
                thumbs_dir / "diff-clone",
            ),
            "chart_copies_differ": _pages_differ(
                self.out_dir / "05-chart-copies.pdf", chart_pages[0] + 1, chart_pages[1] + 1
            ),
        }
        self.result["render"] = render

    def _convert(self, source: pathlib.Path, name: str) -> dict[str, Any]:
        target = self.out_dir / f"{name}.pdf"
        result = pdf_export.convert_to_pdf(
            source,
            self.out_dir / "pdf-tmp",
            timeout_s=self.timeout_s,
        )
        shutil.move(str(result.pdf_path), target)
        shutil.rmtree(self.out_dir / "pdf-tmp", ignore_errors=True)
        return {
            "seconds": round(result.seconds, 2),
            "pages": thumbnails.pdf_page_count(target),
            "pdf_bytes": target.stat().st_size,
        }


# --- вспомогательные --------------------------------------------------------------------


def _iter_text_frames(shapes: Any) -> Any:
    for shape in shapes:
        if shape.shape_type is not None and "GROUP" in str(shape.shape_type):
            yield from _iter_text_frames(shape.shapes)
        elif shape.has_text_frame:
            yield shape.text_frame


def _first_mixed_paragraph(slide: Any) -> Any:
    from lxml import etree

    for frame in _iter_text_frames(slide.shapes):
        for paragraph in frame.paragraphs:
            keys = set()
            for run in paragraph.runs:
                if not run.text.strip():
                    continue
                rpr = run._r.find(f"{{{ooxml.NS_A}}}rPr")
                keys.add(etree.tostring(rpr, method="c14n") if rpr is not None else b"")
            if len(keys) > 1:
                return paragraph
    return None


def _first_text_paragraph(slide: Any) -> Any:
    for frame in _iter_text_frames(slide.shapes):
        for paragraph in frame.paragraphs:
            if paragraph.text.strip():
                return paragraph
    return None


def _first_chart(slide: Any) -> Any:
    for shape in slide.shapes:
        if getattr(shape, "has_chart", False) and shape.has_chart:
            return shape.chart
    raise RuntimeError("на слайде нет диаграммы")


def _pdf_fonts(pdf_path: pathlib.Path) -> list[str]:
    names = set(re.findall(rb"/BaseFont\s*/([A-Za-z0-9+#,._-]+)", pdf_path.read_bytes()))
    cleaned = {re.sub(r"^[A-Z]{6}\+", "", n.decode("latin-1")) for n in names}
    return sorted(cleaned)


def _page_images(pdf_path: pathlib.Path, pages: list[int]) -> dict[int, Image.Image]:
    import pypdfium2 as pdfium

    images: dict[int, Image.Image] = {}
    with thumbnails.PDFIUM_LOCK:
        doc = pdfium.PdfDocument(str(pdf_path))
        try:
            for number in pages:
                if not 1 <= number <= len(doc):
                    continue
                page = doc[number - 1]
                bitmap = page.render(scale=DIFF_WIDTH / page.get_width())
                images[number] = bitmap.to_pil().convert("L")
                bitmap.close()
                page.close()
        finally:
            doc.close()
    return images


def _diff_ratio(a: Image.Image, b: Image.Image) -> float:
    if a.size != b.size:
        b = b.resize(a.size)
    diff = ImageChops.difference(a, b).point(lambda v: 255 if v > PIXEL_THRESHOLD else 0)
    changed = diff.histogram()[255]
    return round(changed / (a.size[0] * a.size[1]), 5)


def _compare_pdfs(a: pathlib.Path, b: pathlib.Path, diff_dir: pathlib.Path) -> dict[str, Any]:
    count = min(thumbnails.pdf_page_count(a), thumbnails.pdf_page_count(b))
    pages = list(range(1, count + 1))
    images_a, images_b = _page_images(a, pages), _page_images(b, pages)
    ratios = {p: _diff_ratio(images_a[p], images_b[p]) for p in pages}
    worst = sorted(ratios.items(), key=lambda kv: -kv[1])[:3]
    diff_dir.mkdir(parents=True, exist_ok=True)
    for page, ratio in worst:
        if ratio > 0:
            ImageChops.difference(images_a[page], images_b[page]).save(
                diff_dir / f"page-{page}.png"
            )
    return {
        "pages": count,
        "pages_a": thumbnails.pdf_page_count(a),
        "pages_b": thumbnails.pdf_page_count(b),
        "max_ratio": max(ratios.values()) if ratios else 0.0,
        "mean_ratio": round(sum(ratios.values()) / len(ratios), 5) if ratios else 0.0,
        "worst_pages": [{"page": p, "ratio": r} for p, r in worst],
    }


def _compare_pages(
    pdf_path: pathlib.Path, pairs: list[tuple[int, int]], diff_dir: pathlib.Path
) -> list[dict[str, Any]]:
    wanted = sorted({p for pair in pairs for p in pair})
    images = _page_images(pdf_path, wanted)
    result = []
    diff_dir.mkdir(parents=True, exist_ok=True)
    for src, dst in pairs:
        if src not in images or dst not in images:
            continue
        ratio = _diff_ratio(images[src], images[dst])
        if ratio > 0:
            ImageChops.difference(images[src], images[dst]).save(diff_dir / f"{src}-vs-{dst}.png")
        result.append({"source_page": src, "clone_page": dst, "ratio": ratio})
    return result


def _pages_differ(pdf_path: pathlib.Path, a: int, b: int) -> dict[str, Any]:
    images = _page_images(pdf_path, [a, b])
    ratio = _diff_ratio(images[a], images[b]) if a in images and b in images else None
    return {"page_a": a, "page_b": b, "ratio": ratio, "differ": bool(ratio and ratio > 0.001)}


# --- отчёт --------------------------------------------------------------------------


def load_templates(templates_dir: pathlib.Path, only: str | None) -> list[pathlib.Path]:
    manifest = templates_dir / "manifest.json"
    if manifest.exists():
        entries = json.loads(manifest.read_text())["templates"]
        paths = [ROOT / e["path"] for e in entries]
    else:
        paths = sorted(templates_dir.glob("*.pptx"))
    paths = [p for p in paths if p.exists()]
    if only:
        paths = [p for p in paths if only.lower() in p.name.lower()]
    return paths


def _fmt_ratio(value: float | None) -> str:
    return "—" if value is None else f"{value * 100:.2f} %"


def _status(ok: bool | None) -> str:
    return "—" if ok is None else ("да" if ok else "нет")


def _ok(value: bool) -> str:
    return "ок" if value else "НЕТ"


def _engine_row(t: dict[str, Any]) -> str:
    s = t["steps"]
    diff = s["roundtrip"]["diff"]
    steps = ("roundtrip", "clone", "text", "native", "chart_copies", "prune")
    checks_ok = all(s[k]["check"]["ok"] for k in steps)
    native_ok = s["native"]["table_present"] and s["native"]["chart_present"]
    cells = [
        t["name"],
        f"{t['inspect']['size_bytes'] / 1e6:.1f}",
        str(t["inspect"]["slides"]),
        f"{s['load']['open_and_walk_seconds']:.2f}",
        f"{s['save_seconds']['01-roundtrip.pptx']:.2f}",
        f"{len(diff['only_in_a'])}/{len(diff['only_in_b'])}, {len(diff['changed_xml'])}, "
        f"{len(diff['changed_binary'])}",
        _status(checks_ok),
        f"{_status(s['clone']['ok'])} ({len(s['clone']['clones'])})",
        _status(s["text"]["ok"]),
        _status(native_ok),
        _status(s["chart_copies"]["independent"]),
        f"{_status(s['prune']['ok'])} ({s['prune']['slides_after']}, "
        f"{s['prune']['size_after'] / 1e6:.1f})",
    ]
    return "| " + " | ".join(cells) + " |"


def _render_row(t: dict[str, Any]) -> str:
    r = t["render"]
    if "error" in r:
        return f"| {t['name']} | ошибка рендера: {r['error']} |" + " |" * 10
    c = r["conversions"]
    v = r["visual"]
    clone_max = max((item["ratio"] for item in v["clones_vs_sources"]), default=None)
    cells = [
        t["name"],
        str(c["original_first"]["pages"]),
        f"{c['original_first']['seconds']:.1f}",
        f"{c['original_second']['seconds']:.1f}",
        f"{c['original_repeat']['seconds']:.1f}",
        "—",
        f"{c['06-pruned']['seconds']:.1f}",
        f"{r['thumbnails']['00-original']['per_page_ms']:.0f}",
        f"{_fmt_ratio(v['roundtrip_vs_original']['max_ratio'])} / "
        f"{_fmt_ratio(v['roundtrip_vs_original']['mean_ratio'])}",
        _fmt_ratio(clone_max),
        _status(v["chart_copies_differ"]["differ"]),
        ", ".join(r["pdf_fonts"])[:120],
    ]
    return "| " + " | ".join(cells) + " |"


ENGINE_HEADER = (
    "| Шаблон | МБ | Слайдов | Открытие, с | Сохранение, с | Roundtrip: части −/+, XML, бин. "
    "| Связи | Клоны | Текст | Таблица+диаграмма | Копии диаграмм независимы "
    "| Удаление образцов (осталось слайдов, МБ) |\n"
    "| --- | ---: | ---: | ---: | ---: | --- | --- | --- | --- | --- | --- | --- |"
)
RENDER_HEADER = (
    "| Шаблон | Страниц | PDF первый, с | PDF второй, с | PDF повтор, с "
    "| Память сервиса (внешний замер) | PDF после удаления образцов, с "
    "| Миниатюры 1280 px, мс/стр. "
    "| Roundtrip: разница max/сред. | Клоны: разница max | Копии диаграмм различаются "
    "| Шрифты в PDF |\n"
    "| --- | ---: | ---: | ---: | ---: | ---: | ---: | ---: | --- | --- | --- | --- |"
)


def render_report_block(out_root: pathlib.Path) -> str:
    lines: list[str] = ["<!-- probe:begin -->", "", "### Замеры", ""]
    lines.append(
        "Блок генерируется `scripts/probe_pptx.py --report` из `runs/pptx-probe/*/results.json`; "
        "правки руками не сохраняются."
    )
    lines.append("")
    results = sorted(out_root.glob("*/results.json"))
    if not results:
        lines += ["Замеров пока нет.", "", "<!-- probe:end -->"]
        return "\n".join(lines)
    for result_path in results:
        data = json.loads(result_path.read_text())
        m = data["machine"]
        measured = m["measured_at"][:16].replace("T", " ")
        lines.append(
            f"#### {m['label']} — {m['platform']}, {m['cpu_count']} CPU, "
            f"{m['memory_mb'] or '?'} МБ, python-pptx {m['python_pptx']}, "
            f"Рендерер: {m.get('renderer') or 'нет'}, {measured} UTC"
        )
        lines += ["", ENGINE_HEADER]
        lines += [_engine_row(t) for t in data["templates"]]
        lines.append("")
        rendered = [t for t in data["templates"] if t.get("render")]
        if rendered:
            lines.append(RENDER_HEADER)
            lines += [_render_row(t) for t in rendered]
            lines.append("")
    lines.append("<!-- probe:end -->")
    return "\n".join(lines)


def write_checklist(out_root: pathlib.Path, results: list[dict[str, Any]]) -> pathlib.Path:
    """Чек-лист ручной проверки в PowerPoint с номерами слайдов из этого прогона."""
    lines = [
        "# Ручная проверка в PowerPoint",
        "",
        f"Файлы: `{out_root}`. Для каждого файла: открыть в PowerPoint без сообщения "
        "«восстановить», пролистать, затем сохранить под новым именем и открыть снова. "
        "Номера слайдов — как в панели слайдов PowerPoint (с 1).",
        "",
    ]
    for t in results:
        s = t["steps"]
        clones = s["clone"]["clones"]
        lines.append(f"## {t['name']}")
        lines.append("")
        lines.append(
            f"- [ ] `01-roundtrip.pptx`: {t['inspect']['slides']} слайдов, вид совпадает "
            "с исходным файлом (фон, картинки, шрифты Play/Arial, заметки докладчика на месте)."
        )
        pairs = ", ".join(
            f"{v['clone'] + 1} — копия {v['source'] + 1} ({k})" for k, v in clones.items()
        )
        lines.append(
            f"- [ ] `02-clone.pptx`: в конце добавлены слайды {pairs}; копии неотличимы "
            "от оригиналов, группы выделяются и разгруппировываются, обрезка картинок сохранена."
        )
        text = s["text"]
        details = []
        if "paragraph" in text:
            details.append(
                "абзац «Проверка замены текста / вторая строка» сохранил шрифт, кегль и цвет"
            )
        if "mixed_runs" in text:
            details.append(
                f"абзац «{text['mixed_runs']['text'][:60]}» сохранил разные стили фрагментов"
            )
        lines.append(
            f"- [ ] `03-text.pptx`: на клонах {'; '.join(details) or 'текст без изменений'}."
        )
        n = s["native"]["slide"] + 1
        lines.append(
            f"- [ ] `04-native.pptx`: слайд {n} — таблица 4×3 и диаграмма «Открываемость/Отписки»; "
            "ячейки редактируются, у диаграммы работает «Изменить данные» (открывается Excel-лист)."
        )
        a, b = (c + 1 for c in s["chart_copies"]["copies"])
        base = s["chart_copies"]["base_slide"] + 1
        lines.append(
            f"- [ ] `05-chart-copies.pptx`: слайды {a} и {b} — копии слайда {base}; "
            f"на {a} диаграмма «Копия А» (Q1–Q3), на {b} — «Копия Б» (Янв–Май); "
            "изменение данных одной копии не меняет другую и исходный слайд."
        )
        lines.append(
            f"- [ ] `06-pruned.pptx`: осталось {s['prune']['slides_after']} слайдов "
            f"({s['prune']['size_after'] / 1e6:.1f} МБ), макеты в меню «Макет» на месте, "
            "открывается без восстановления."
        )
        lines.append("")
    lines.append(
        "Результат записать в docs/pptx-capabilities.md, раздел «Ручная проверка в PowerPoint»."
    )
    path = out_root / "CHECKLIST.md"
    path.write_text("\n".join(lines) + "\n")
    return path


def update_report(report_path: pathlib.Path, block: str) -> None:
    begin, end = "<!-- probe:begin -->", "<!-- probe:end -->"
    if report_path.exists():
        text = report_path.read_text()
        if begin in text and end in text:
            head = text[: text.index(begin)]
            tail = text[text.index(end) + len(end) :]
            report_path.write_text(head + block + tail)
            return
        report_path.write_text(text.rstrip("\n") + "\n\n" + block + "\n")
        return
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text("# Возможности движка PPTX\n\n" + block + "\n")


# --- запуск -------------------------------------------------------------------------


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Проверка движка PPTX и рендера")
    parser.add_argument("--templates-dir", type=pathlib.Path, default=ROOT / "data" / "organizers")
    parser.add_argument("--out", type=pathlib.Path, default=ROOT / "runs" / "pptx-probe")
    parser.add_argument("--report", type=pathlib.Path, help="обновить блок замеров в этом файле")
    parser.add_argument("--machine", help="метка машины для results.json и отчёта")
    parser.add_argument("--only", help="подстрока имени шаблона")
    parser.add_argument("--skip-render", action="store_true", help="без ONLYOFFICE и миниатюр")
    parser.add_argument("--no-fixture", action="store_true", help="не проверять свою фикстуру")
    parser.add_argument("--report-only", action="store_true", help="только пересобрать отчёт")
    parser.add_argument("--timeout", type=int, default=300, help="тайм-аут одной конвертации, с")
    args = parser.parse_args(argv)

    if not args.report_only:
        machine = machine_info(args.machine)
        render = not args.skip_render
        if render and not machine["renderer"]:
            print("ONLYOFFICE не настроен: рендер пропущен (--skip-render)")
            render = False
        paths = load_templates(args.templates_dir, args.only)
        if not args.no_fixture and (not args.only or args.only.lower() in FIXTURE.name):
            paths.insert(0, FIXTURE)
        if not paths:
            print("нет шаблонов для проверки")
            return 1
        out_root = args.out / machine["label"]
        out_root.mkdir(parents=True, exist_ok=True)
        results: list[dict[str, Any]] = []
        for path in paths:
            print(f"== {path.name}")
            started = time.perf_counter()
            probe = TemplateProbe(path, out_root / _slug(path.stem), render, args.timeout)
            result = probe.run()
            result["total_seconds"] = round(time.perf_counter() - started, 1)
            results.append(result)
            _print_result(result)
        (out_root / "results.json").write_text(
            json.dumps({"machine": machine, "templates": results}, ensure_ascii=False, indent=2)
            + "\n"
        )
        print(f"результаты: {out_root / 'results.json'}")
        print(f"чек-лист ручной проверки: {write_checklist(out_root, results)}")
    if args.report:
        update_report(args.report, render_report_block(args.out))
        print(f"отчёт обновлён: {args.report}")
    return 0


def _print_result(result: dict[str, Any]) -> None:
    s = result["steps"]
    diff = s["roundtrip"]["diff"]
    print(
        f"  открытие {s['load']['open_and_walk_seconds']} с, roundtrip {_ok(s['roundtrip']['ok'])} "
        f"(XML изменено {len(diff['changed_xml'])}, частей −{len(diff['only_in_a'])}), "
        f"клоны {_ok(s['clone']['ok'])}, текст {_ok(s['text']['ok'])}, "
        f"нативные {_ok(s['native']['ok'])}, "
        f"копии диаграмм {'независимы' if s['chart_copies']['independent'] else 'НЕТ'}, "
        f"удаление {_ok(s['prune']['ok'])} ({s['prune']['slides_after']} слайдов, "
        f"{s['prune']['size_after'] / 1e6:.1f} МБ)"
    )
    render = result.get("render")
    if render and "error" in render:
        print("  рендер: ошибка", render["error"])
    elif render:
        c = render["conversions"]
        v = render["visual"]
        print(
            f"  PDF: первый {c['original_first']['seconds']} с, второй "
            f"{c['original_second']['seconds']} с, "
            f"повтор {c['original_repeat']['seconds']} с, "
            f"страниц {c['original_first']['pages']}; roundtrip разница "
            f"max {_fmt_ratio(v['roundtrip_vs_original']['max_ratio'])}; "
            f"шрифты {render['pdf_fonts']}"
        )
    print(f"  всего {result['total_seconds']} с")


if __name__ == "__main__":
    sys.exit(main())
