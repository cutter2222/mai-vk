"""Create a fresh sandbox OfficeStore document for opt-in live SDK tests.

Run in the API environment. This writes a NEW sandbox document to its OfficeStore,
not an existing office copy. --project also publishes a synthetic terminal job,
artifact and project using the normal stores. No editor, model or worker is called.
--source copies existing PPTX bytes without modifying the source. With --project
it publishes a new sandbox project; standalone copies use sdk-brand namespace.
The output directory must not exist. Retain seed.json for test diagnostics/cleanup.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import uuid
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from presentation_designer.api.routes.projects import DEFAULT_BRIEF, DEFAULT_SETTINGS
from presentation_designer.pipeline.artifacts import ArtifactStore
from presentation_designer.pipeline.office import OfficeStore
from presentation_designer.pipeline.state import State, now_iso
from presentation_designer.shared.settings import get_settings


def seed(
    data_dir: Path,
    output: Path,
    artifacts_dir: Path | None = None,
    *,
    source_path: Path | None = None,
) -> dict[str, object]:
    supplied = source_path.read_bytes() if source_path is not None else None
    if supplied is not None:
        Presentation(io.BytesIO(supplied))  # Validate before writing any sandbox state.
    output.mkdir(parents=True, exist_ok=False)
    identifier = uuid.uuid4().hex
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(
        Inches(1), Inches(1), Inches(8), Inches(1)
    ).text = f"SDK sandbox {identifier}"
    buffer = io.BytesIO()
    deck.save(buffer)
    original = supplied if supplied is not None else buffer.getvalue()
    slide_count = len(Presentation(io.BytesIO(original)).slides)
    (output / "before.pptx").write_bytes(original)
    source = f"sdk-{'brand' if supplied is not None else 'sandbox'}/{identifier}"
    project_fields: dict[str, object] = {}
    if artifacts_dir is not None:
        # Terminal synthetic job, never enqueued and never sent to a model/worker.
        state = State(data_dir / "state.sqlite3")
        artifacts = ArtifactStore(artifacts_dir)
        job_id = f"job_sdk_sandbox_{identifier}"
        state.create_job(kind="sdk-sandbox", job_id=job_id)
        state.update_job(job_id, status="succeeded", stage="done", finished_at=now_iso())
        state.create_generation(
            job_id=job_id,
            template_id="sdk-sandbox",
            package_id="sdk-sandbox",
            request={},
            idempotency_key=None,
            execution_mode={"mode": "real", "layers": {}},
            versions={},
            variants=[{"variant_id": "compact", "axis": "density", "value": "compact"}],
        )
        with artifacts.stage_revision(job_id, "compact", 1) as staging:
            staging.write_bytes("deck.pptx", original)
        state.add_revision(
            job_id=job_id,
            variant_id="compact",
            revision=1,
            artifacts_prefix=staging.prefix,
            manifest=staging.manifest,
            pptx_hash=hashlib.sha256(original).hexdigest(),
        )
        state.update_variant(
            job_id, "compact", status="ready", slide_count=slide_count, ready_at=now_iso()
        )
        project = state.create_project(
            f"SDK sandbox {identifier}",
            DEFAULT_BRIEF,
            DEFAULT_SETTINGS,
            job_id=job_id,
            chosen_variant="compact",
        )
        artifact = f"{staging.prefix}deck.pptx"
        source = f"{job_id}/{artifact}"
        project_fields = {
            "project_id": project["project_id"],
            "job_id": job_id,
            "artifact": artifact,
        }
    doc = OfficeStore(data_dir).create(source, f"SDK sandbox {identifier}.pptx", original)
    report = {
        "document_id": doc["id"],
        "source": doc["source"],
        "revision": doc["revision"],
        "sha256": hashlib.sha256(original).hexdigest(),
        "slides": slide_count,
        "synthetic": supplied is None,
        "source_file": str(source_path.resolve()) if source_path is not None else None,
        **project_fields,
    }
    (output / "seed.json").write_text(json.dumps(report, indent=2), encoding="utf-8")
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument(
        "--source", type=Path, help="Copy a PPTX into a new isolated Office document"
    )
    parser.add_argument(
        "--project", action="store_true", help="Also seed a synthetic project/artifact"
    )
    args = parser.parse_args()
    settings = get_settings()
    print(
        json.dumps(
            seed(
                settings.data_dir,
                args.out,
                settings.artifacts_dir if args.project else None,
                source_path=args.source,
            ),
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
