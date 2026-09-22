"""Audit an SDK round-trip copy, without opening sessions or modifying source files.

Render both PPTX files with the same ONLYOFFICE converter. Pixel differences are
diagnostics, not proof of PowerPoint fidelity. Exit nonzero if other slides change.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from collections.abc import Iterable
from pathlib import Path
from typing import Any
from zipfile import ZipFile

from lxml import etree
from PIL import Image, ImageChops, ImageStat
from pptx import Presentation

from presentation_designer.audit.office_roundtrip import numeric_xml_deltas
from presentation_designer.export.pdf import convert_to_pdf
from presentation_designer.export.thumbnails import render_thumbnails
from presentation_designer.shared.settings import get_settings


def xml_errors(path: Path) -> list[dict[str, str]]:
    errors = []
    with ZipFile(path) as archive:
        for name in archive.namelist():
            if name.endswith((".xml", ".rels")):
                try:
                    etree.fromstring(
                        archive.read(name),
                        etree.XMLParser(resolve_entities=False, no_network=True),
                    )
                except etree.XMLSyntaxError as exc:
                    errors.append({"part": name, "error": str(exc)})
    return errors


def snapshot(path: Path) -> dict[str, Any]:
    deck = Presentation(str(path))

    def shapes(items: Iterable[Any]) -> list[dict[str, Any]]:
        result = []
        for shape in items:
            item: dict[str, Any] = {
                "type": int(shape.shape_type),
                "box": [shape.left, shape.top, shape.width, shape.height, shape.rotation],
            }
            if shape.has_text_frame:
                item["text"] = shape.text
            if shape.has_table:
                item["table"] = [[cell.text for cell in row.cells] for row in shape.table.rows]
            if hasattr(shape, "shapes"):
                item["children"] = shapes(shape.shapes)
            result.append(item)
        return result

    return {
        "size": [deck.slide_width, deck.slide_height],
        "slides": [shapes(slide.shapes) for slide in deck.slides],
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("before", type=Path)
    parser.add_argument("after", type=Path)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--slide", type=int, required=True)
    parser.add_argument("--old", required=True)
    parser.add_argument("--new", required=True)
    args = parser.parse_args()
    original, updated = args.before.read_bytes(), args.after.read_bytes()
    args.out.mkdir(parents=True, exist_ok=False)
    invalid = {"before": xml_errors(args.before), "after": xml_errors(args.after)}
    if any(invalid.values()):
        (args.out / "report.json").write_text(
            json.dumps(
                {
                    "accepted": False,
                    "reason": "invalid_ooxml",
                    "xml_errors": invalid,
                    "before_sha256": hashlib.sha256(original).hexdigest(),
                    "after_sha256": hashlib.sha256(updated).hexdigest(),
                    "visual_comparison_performed": False,
                },
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )
        raise SystemExit("Invalid OOXML; see report.json. No repair attempted.")
    left, right = snapshot(args.before), snapshot(args.after)
    if not 1 <= args.slide <= len(left["slides"]):
        raise ValueError("Invalid target slide")
    expected = json.loads(json.dumps(left))
    matches = 0

    def replace(items: list[dict[str, Any]]) -> None:
        nonlocal matches
        for item in items:
            if "text" in item:
                matches += item["text"].count(args.old)
                item["text"] = item["text"].replace(args.old, args.new)
            if "children" in item:
                replace(item["children"])

    replace(expected["slides"][args.slide - 1])
    if matches != 1:
        raise ValueError(f"Expected exactly one source match, got {matches}")
    with ZipFile(args.before) as old_zip, ZipFile(args.after) as new_zip:
        old_names, new_names = set(old_zip.namelist()), set(new_zip.namelist())
        parts = {
            "removed": sorted(old_names - new_names),
            "added": sorted(new_names - old_names),
            "changed": sorted(
                n for n in old_names & new_names if old_zip.read(n) != new_zip.read(n)
            ),
        }
    (args.out / "before-structure.json").write_text(
        json.dumps(left, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (args.out / "after-structure.json").write_text(
        json.dumps(right, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    rendered = []
    for name, path in [("before", args.before), ("after", args.after)]:
        pdf = convert_to_pdf(path, args.out / name, settings=get_settings(), timeout_s=240)
        rendered.append(render_thumbnails(pdf.pdf_path, args.out / name / "slides", width_px=1600))
    images = []
    if len(rendered[0].paths) != len(rendered[1].paths):
        raise ValueError("Rendered page count changed")
    for i, (a, b) in enumerate(zip(rendered[0].paths, rendered[1].paths, strict=True), 1):
        with Image.open(a) as x, Image.open(b) as y:
            if x.size != y.size:
                raise ValueError("Rendered dimensions changed")
            diff = ImageChops.difference(x.convert("RGB"), y.convert("RGB"))
            box = diff.getbbox()
            red, green, blue = diff.split()
            mask = ImageChops.lighter(ImageChops.lighter(red, green), blue).point(
                lambda p: 255 if p else 0
            )
            images.append(
                {
                    "slide": i,
                    "identical": box is None,
                    "bbox": box,
                    "mean_absolute_channel_delta": ImageStat.Stat(diff).mean,
                    "changed_pixel_fraction": ImageStat.Stat(mask).mean[0] / 255,
                }
            )
            if box is not None:
                diff.save(args.out / f"diff-{i:02}.png")
    semantic_changes = [
        i
        for i, (a, b) in enumerate(zip(expected["slides"], right["slides"], strict=False), 1)
        if a != b
    ]
    report = {
        "before_sha256": hashlib.sha256(original).hexdigest(),
        "after_sha256": hashlib.sha256(updated).hexdigest(),
        "slides_before": len(left["slides"]),
        "slides_after": len(right["slides"]),
        "expected_text_and_geometry_only": expected == right,
        "numeric_xml_deltas": numeric_xml_deltas(args.before, args.after),
        "unexpected_structure_slides": semantic_changes,
        "zip_parts": parts,
        "pixels": images,
        "unchanged_other_slide_pixels": all(
            x["identical"] for x in images if x["slide"] != args.slide
        ),
        "inputs_unchanged": (
            args.before.read_bytes() == original and args.after.read_bytes() == updated
        ),
        "limitations": "Same-renderer comparison, not PowerPoint fidelity or AI/plugin acceptance",
    }
    (args.out / "report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(
        json.dumps(
            {
                k: v
                for k, v in report.items()
                if k not in {"zip_parts", "pixels", "numeric_xml_deltas"}
            }
        )
    )
    if not (
        expected == right
        and report["unchanged_other_slide_pixels"]
        and not report["numeric_xml_deltas"]
    ):
        raise SystemExit(1)


if __name__ == "__main__":
    main()
